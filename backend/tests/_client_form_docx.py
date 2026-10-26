"""Fikcyjny formularz CV klienta w kształcie formularzy SWZ Centrum e-Zdrowia.

Prawdziwych plików klienta nie trzymamy w publicznym repo — ten buduje ten sam
układ: etykiety nagłówka, ramki 1×1, tabelę doświadczenia z „…”, kompetencje
z wierszami-sekcjami i drugim nagłówkiem w środku oraz tabelę z przesuniętą
siatką (jak „Analityk biznesowo-systemowy”: „Lp.” scalone w nagłówku na dwie
kolumny, a w wierszach inaczej).
"""

from __future__ import annotations

import io

import docx


def _row(table, index: int, texts: list[str]) -> None:
    for position, value in enumerate(texts):
        if value is not None:
            table.cell(index, position).text = value


def _merge(table, row: int, first: int, last: int, value: str = "") -> None:
    merged = table.cell(row, first).merge(table.cell(row, last))
    merged.text = value


def form_docx() -> bytes:
    document = docx.Document()
    document.add_paragraph("Załącznik nr 4b do SWZ")
    document.add_paragraph("Formularz CV\nTester automatyczny")
    document.add_paragraph("Imię i nazwisko: ")
    document.add_paragraph("Łączna liczba lat doświadczenia zawodowego:")
    document.add_paragraph("Miejsce zamieszkania:")

    document.add_paragraph("Cel zawodowy:")
    document.add_table(rows=1, cols=1)

    document.add_paragraph("Doświadczenie zawodowe*:")
    experience = document.add_table(rows=6, cols=5)
    _row(
        experience,
        0,
        ["Lp.", "Miejsce pracy", "Okres pracy", "Stanowisko", "Opis projektów"],
    )
    for number in (1, 2, 3):
        _row(experience, number, [str(number), "", "", "", ""])
    _row(experience, 4, ["4", "…", "", "", ""])

    document.add_paragraph("Kompetencje*:")
    skills = document.add_table(rows=7, cols=5)
    _row(
        skills,
        0,
        [
            "Lp.",
            "Nazwa kompetencji",
            "Spełnienie kompetencji (TAK/NIE)",
            "Liczba lat doświadczenia",
            "Opis wykonywanych zadań",
        ],
    )
    _merge(skills, 1, 0, 4, "Kompetencje obligatoryjne")
    _row(skills, 2, ["1.", "Znajomość języka SQL;", "", "", ""])
    _row(
        skills,
        3,
        [
            "Lp.",
            "Nazwa kompetencji",
            "Liczba lat doświadczenia",
            "Stopień zaawansowania w skali 1-10**",
            "Opis wykonywanych zadań",
        ],
    )
    _merge(
        skills,
        4,
        0,
        4,
        "Kompetencje dodatkowe – podlegające ocenie przez Zamawiającego",
    )
    _row(skills, 5, ["", "Długość doświadczenia z frameworkiem Selenium", "", "", ""])
    _row(
        skills,
        6,
        ["", "Staż pracy w projektach związanych ze służbą zdrowia", "", "", ""],
    )

    document.add_paragraph("Doświadczenie w projektach publicznych")
    shifted = document.add_table(rows=3, cols=7)
    _merge(shifted, 0, 0, 1, "Lp.")
    _merge(shifted, 0, 2, 3, "Miejsce pracy")
    _row(
        shifted,
        0,
        [None, None, None, None, "Okres pracy", "Stanowisko", "Opis projektów"],
    )
    for number in (1, 2):
        _row(shifted, number, [str(number)])
        _merge(shifted, number, 1, 2)
        _merge(shifted, number, 3, 4)

    document.add_paragraph("Certyfikaty (jeśli posiada):")
    certificates = document.add_table(rows=2, cols=4)
    _row(certificates, 0, ["Lp.", "Certyfikat", "Data wydania", "Instytucja wydająca"])
    _row(certificates, 1, ["1", "Certyfikat ISTQB na poziomie podstawowym", "", ""])

    document.add_paragraph("Informacje dodatkowe:")
    document.add_table(rows=1, cols=1)
    document.add_paragraph("Wyrażam zgodę na przetwarzanie danych osobowych…")
    return _bytes(document)


def filled_table_docx() -> bytes:
    """Opis zamówienia: tabela bez pustych pól — to nie formularz."""
    document = docx.Document()
    table = document.add_table(rows=3, cols=2)
    _row(table, 0, ["Rola", "Tester automatyczny"])
    _row(table, 1, ["Tryb pracy", "Zdalnie"])
    _row(table, 2, ["Etapy rekrutacji", "CV + rozmowa"])
    return _bytes(document)


def _bytes(document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
