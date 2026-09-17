"""Jednorazowa korekta 0316 — kontakt z generatora na powiązane umowy.

Kryteria akceptacji ticketu w części „backfill + arkusz braków":

* kopiuje e-mail i telefon z `render_payload` podpisanego dokumentu na umowę;
* jest FILL-ONLY — ręczna poprawka w kontrakcie wygrywa z dokumentem;
* jest idempotentna — drugi przebieg nic nie robi;
* NIE materializuje profilu kandydata (fallback ma zostać żywy);
* paragon ma kształt paragonu migracji, a wartości (PII) leżą pod kluczem,
  którego publiczny `show_migration_receipts` nie wydrukuje;
* raport braków odmawia przed wykonaniem korekty i wypisuje wyłącznie
  kontrakty bez kontaktu w OBU źródłach.
"""

from __future__ import annotations

import io
import uuid
from datetime import date

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook

from app.services.contract_candidate_contact_backfill import (
    DETAILS_KEY,
    REPAIR_MARKER,
    details_key,
    run_candidate_contact_backfill,
    summarize_for_log,
)

pytestmark = pytest.mark.asyncio


#: `candidates.email` jest UNIQUE — seed dostaje flagę, nie gotowy adres,
#: inaczej drugi test w pliku wywraca się na wstawianiu kandydata.
async def _seed(
    *,
    profile_email: bool = False,
    profile_phone: str | None,
    payload_email: str | None,
    payload_phone: str | None,
    contract_email: str | None = None,
    contract_phone: str | None = None,
    link_document: bool = True,
) -> int:
    """Kandydat + klient + umowa + wygenerowany dokument. Zwraca contract_id."""

    from app.core.database import AsyncSessionLocal
    from app.models.b2b_generated_contract import B2BGeneratedContract
    from app.models.candidate import Candidate
    from app.models.client import Client
    from app.models.contract import Contract, ContractStatus, ContractType

    unique = uuid.uuid4().hex[:8]
    async with AsyncSessionLocal() as db:
        candidate = Candidate(
            name="Jan",
            lastname=f"Backfill{unique}",
            email=f"backfill-{unique}@example.com" if profile_email else None,
            phone=profile_phone,
        )
        client = Client(name=f"Klient Backfill {unique}")
        db.add_all([candidate, client])
        await db.flush()
        contract = Contract(
            candidate_id=candidate.id,
            client_id=client.id,
            contract_type=ContractType.b2b,
            status=ContractStatus.active,
            start_date=date(2026, 1, 1),
            candidate_email=contract_email,
            candidate_phone=contract_phone,
        )
        db.add(contract)
        await db.flush()
        payload: dict[str, object] = {"partner_name": "Jan Backfill"}
        if payload_email is not None:
            payload["partner_email"] = payload_email
        if payload_phone is not None:
            payload["partner_phone"] = payload_phone
        db.add(
            B2BGeneratedContract(
                # `seq` jest unikalne w obrębie roku — trzymamy testy w roku,
                # którego produkcja nie używa, żeby nie kolidować z sąsiadami.
                year=2999,
                seq=int(uuid.uuid4().int % 1_000_000),
                contract_number=f"{unique}/2999",
                partner_name="Jan Backfill",
                client_name=client.name,
                language="pl",
                signing_date=date(2026, 1, 1),
                start_date=date(2026, 1, 1),
                candidate_id=candidate.id,
                client_id=client.id,
                contract_id=contract.id if link_document else None,
                render_payload=payload,
            )
        )
        await db.commit()
        return contract.id


async def _run(contract_ids: set[int], marker: str):
    from app.core.database import AsyncSessionLocal

    async with AsyncSessionLocal() as db:
        summary = await run_candidate_contact_backfill(
            db,
            marker=marker,
            key_for_details=details_key(marker),
            only_contract_ids=contract_ids,
        )
        await db.commit()
    return summary


def _marker() -> str:
    # Marker produkcyjny jest jednorazowy globalnie, a baza testowa jest
    # współdzielona — każdy test dostaje własny, tak samo ukształtowany.
    return f"0320_test_{uuid.uuid4().hex[:8]}"


async def test_copies_the_generator_contact_onto_the_linked_contract():
    from app.core.database import AsyncSessionLocal
    from app.models.contract import Contract

    contract_id = await _seed(
        profile_email=False,
        profile_phone=None,
        payload_email="  Z.Dokumentu@Example.COM ",
        payload_phone=" +48 600 100 200 ",
    )
    summary = await _run({contract_id}, _marker())
    assert summary is not None
    assert summary["contracts_email_filled"] == 1
    assert summary["contracts_phone_filled"] == 1

    async with AsyncSessionLocal() as db:
        contract = await db.get(Contract, contract_id)
        assert contract.candidate_email == "z.dokumentu@example.com"
        assert contract.candidate_phone == "+48 600 100 200"


async def test_never_overwrites_a_value_already_on_the_contract():
    """Ręczna poprawka Delivery wygrywa z dokumentem podpisanym wcześniej."""

    from app.core.database import AsyncSessionLocal
    from app.models.contract import Contract

    contract_id = await _seed(
        profile_email=False,
        profile_phone=None,
        payload_email="z.dokumentu@example.com",
        payload_phone="+48 600 100 200",
        contract_email="recznie@example.com",
        contract_phone="+48 700 700 700",
    )
    summary = await _run({contract_id}, _marker())
    assert summary["contracts_email_filled"] == 0
    assert any(item["reason"] == "already_set" for item in summary["skipped"])

    async with AsyncSessionLocal() as db:
        contract = await db.get(Contract, contract_id)
        assert contract.candidate_email == "recznie@example.com"
        assert contract.candidate_phone == "+48 700 700 700"


async def test_does_not_materialize_the_candidate_profile():
    """Poziom drugi kolejności źródeł ma zostać ŻYWY, nie skopiowany."""

    from app.core.database import AsyncSessionLocal
    from app.models.contract import Contract

    contract_id = await _seed(
        profile_email=True,
        profile_phone="+48 500 500 500",
        payload_email=None,
        payload_phone=None,
    )
    summary = await _run({contract_id}, _marker())
    assert summary["contracts_email_filled"] == 0
    assert any(
        item["reason"] == "no_contact_in_payload" for item in summary["skipped"]
    )

    async with AsyncSessionLocal() as db:
        contract = await db.get(Contract, contract_id)
        assert contract.candidate_email is None, "profil zostaje w profilu"


async def test_an_unlinked_document_is_never_guessed_onto_a_contract():
    """Bez `contract_id` nie zgadujemy po kandydacie — to cudzy numer."""

    from app.core.database import AsyncSessionLocal
    from app.models.contract import Contract

    contract_id = await _seed(
        profile_email=False,
        profile_phone=None,
        payload_email="z.dokumentu@example.com",
        payload_phone="+48 600 100 200",
        link_document=False,
    )
    summary = await _run({contract_id}, _marker())
    assert summary["linked_documents_found"] == 0

    async with AsyncSessionLocal() as db:
        contract = await db.get(Contract, contract_id)
        assert contract.candidate_email is None


async def test_second_run_does_nothing():
    contract_id = await _seed(
        profile_email=False,
        profile_phone=None,
        payload_email="z.dokumentu@example.com",
        payload_phone=None,
    )
    marker = _marker()
    assert await _run({contract_id}, marker) is not None
    assert await _run({contract_id}, marker) is None
    assert summarize_for_log(None) == "already done"


def test_the_receipt_is_receipt_shaped_and_the_details_key_is_not():
    """Log GitHub Actions jest publiczny — kontakt nie może do niego trafić."""

    from scripts.show_migration_receipts import is_receipt_key

    assert is_receipt_key(REPAIR_MARKER)
    assert not is_receipt_key(DETAILS_KEY)
    assert len(DETAILS_KEY) <= 100


def test_the_log_line_carries_only_numbers_and_reason_codes():
    line = summarize_for_log(
        {
            "linked_documents_found": 3,
            "contracts_email_filled": 2,
            "contracts_phone_filled": 1,
            "skipped": [{"contract_id": 7, "reason": "already_set"}],
        }
    )
    assert "already_set=1" in line
    assert "@" not in line, "log kontenera idzie do Loki — bez adresów"


async def test_the_gap_report_refuses_before_the_backfill_has_run(
    app_client: AsyncClient, app_auth_headers: dict[str, str]
):
    from app.core.database import AsyncSessionLocal
    from app.models.app_setting import AppSetting

    async with AsyncSessionLocal() as db:
        existing = await db.get(AppSetting, REPAIR_MARKER)
        if existing is not None:
            await db.delete(existing)
            await db.commit()

    refused = await app_client.get(
        "/api/contracts/candidate-contact-report", headers=app_auth_headers
    )
    assert refused.status_code == 409


async def test_the_gap_report_lists_only_contracts_missing_in_both_sources(
    app_client: AsyncClient, app_auth_headers: dict[str, str]
):
    from app.core.database import AsyncSessionLocal
    from app.models.app_setting import AppSetting

    covered = await _seed(
        profile_email=True,
        profile_phone="+48 500 500 500",
        payload_email=None,
        payload_phone=None,
    )
    missing = await _seed(
        profile_email=False,
        profile_phone=None,
        payload_email=None,
        payload_phone=None,
    )

    async with AsyncSessionLocal() as db:
        if await db.get(AppSetting, REPAIR_MARKER) is None:
            db.add(
                AppSetting(
                    key=REPAIR_MARKER,
                    value={
                        "executed_at": "2026-09-16T00:00:00+00:00",
                        "linked_documents_found": 0,
                        "contracts_email_filled": 0,
                        "contracts_phone_filled": 0,
                        "skipped": [],
                    },
                )
            )
            await db.commit()

    response = await app_client.get(
        "/api/contracts/candidate-contact-report", headers=app_auth_headers
    )
    assert response.status_code == 200, response.text
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames[0] == "Braki"
    ids = {row[0] for row in book["Braki"].iter_rows(min_row=2, values_only=True)}
    assert missing in ids
    assert covered not in ids, "kontakt z profilu to nie brak"
