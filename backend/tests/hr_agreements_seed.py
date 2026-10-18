"""Zasiew testów umów z ludźmi (etap 5, E5). Dane fikcyjne — repo jest publiczne.

Fixture ``agreements_on`` włącza moduły „Ludzi” i przełącznik umów, ustawia
minimalne stawki na ten i przyszły rok (31,40 zł) i czyści tabele umów oraz
stawki osób z zasiewu. E-teczka (moduł dokumentów osoby, E4) jest zaślepiona:
``attached`` zbiera wywołania ``attach_document``. Autenti to fałszywy klient
(``FakeAutenti``) podstawiany pod ``agreements_signing.client_factory``.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date
from typing import Any, Optional

import docx
import pytest
import pytest_asyncio
from sqlalchemy import delete, select

import app.models  # noqa: F401  (rejestracja wszystkich mapperów)
from app.core.config import settings as app_settings
from app.core.database import AsyncSessionLocal
from app.core.scheduling import business_today
from app.models.app_setting import AppSetting
from app.models.hr_agreement import (
    HrAgreement,
    HrAgreementTemplate,
    HrAgreementTemplateVersion,
)
from app.models.hr_document import HrDocument
from app.models.hr_payroll import EmployeeRate
from app.models.user import User
from app.services import cv_pdf_preview
from app.services.autenti.client import AutentiError
from app.services.hr import agreements_effects, agreements_signing
from app.services.hr import agreements_templates as templates
from app.services.hr.agreements import SETTINGS_KEY
from app.services.hr.payroll_periods import add_months
from app.services.hr.payroll_rates import MINIMUM_RATES_SETTING
from app.services.hr_foundation_schema import HR_MODULE_KEYS, HR_MODULES_SETTING
from tests.hr_settlements_seed import make_user, set_setting

PREFIX = "/api/hr/agreements"
FAKE_PDF = b"%PDF-1.4\n% podglad umowy (fikcja)\n"
SIGNED_PDF = b"%PDF-1.4\n% umowa podpisana w Autenti (fikcja)\n"
SCAN_PDF = b"%PDF-1.7\n% skan podpisanej umowy (fikcja)\n"

BODY = (
    "{{ employment_form }} zawarta {{ agreement_date }} z {{ person_name }}.",
    "Od {{ start_date }}{% if indefinite %} bez daty końca{% else %} do {{ end_date }}{% endif %}.",
    "Stawka {{ rate }} {{ currency }} {{ rate_basis }} za godzinę ({{ rate_words }}).",
    "{% if hours_confirmation %}Godziny potwierdza {{ coordinator_name }}.{% endif %}",
    "Wypłata do {{ payout_day }}. dnia następnego miesiąca.",
)


def make_docx(*lines: str) -> bytes:
    document = docx.Document()
    for line in lines:
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def docx_text(content: bytes) -> str:
    return "\n".join(p.text for p in docx.Document(io.BytesIO(content)).paragraphs)


def first_of_next_month(months: int = 1) -> date:
    today = business_today()
    return date(*add_months(today.year, today.month, months), 1)


async def make_template(kind: str = "zlecenie", *, name: Optional[str] = None) -> int:
    """Wzór z wersją 1 przez serwis (autor — konto Kadr z zasiewu)."""

    author = await make_user("recruiter", grants=("people_admin",), profile=False)
    async with AsyncSessionLocal() as db:
        template = await templates.create_template(
            db,
            kind=kind,
            name=name or f"Wzór {kind} {author.id}",
            description=None,
            filename="wzor.docx",
            content=make_docx(*BODY),
            actor=author,
        )
        await db.commit()
        return template.id


async def wipe(user_ids: list[int]) -> None:
    async with AsyncSessionLocal() as db:
        await db.execute(delete(HrDocument).where(HrDocument.source == "agreement"))
        await db.execute(
            delete(HrAgreement).where(HrAgreement.base_agreement_id.is_not(None))
        )
        await db.execute(delete(HrAgreement))
        await db.execute(delete(HrAgreementTemplateVersion))
        await db.execute(delete(HrAgreementTemplate))
        if user_ids:
            await db.execute(
                delete(EmployeeRate).where(EmployeeRate.user_id.in_(user_ids))
            )
        await db.commit()


async def _setting(key: str) -> Any:
    async with AsyncSessionLocal() as db:
        return await db.scalar(select(AppSetting.value).where(AppSetting.key == key))


@dataclass
class Attached:
    calls: list[dict[str, Any]] = field(default_factory=list)
    fail: Optional[Exception] = None


@dataclass
class Team:
    hr: User  # Kadry
    hr2: User  # druga osoba z Kadr
    payroll: User  # Płace
    boss: User  # przełożony / koordynator
    mandate: User  # zlecenie
    b2b: User
    attached: Attached


@pytest_asyncio.fixture(name="agreements_on")
async def agreements_on_fixture(monkeypatch):
    previous = {
        key: await _setting(key)
        for key in (HR_MODULES_SETTING, SETTINGS_KEY, MINIMUM_RATES_SETTING)
    }
    year = business_today().year
    await set_setting(HR_MODULES_SETTING, {key: True for key in HR_MODULE_KEYS})
    await set_setting(SETTINGS_KEY, {"enabled": True})
    await set_setting(
        MINIMUM_RATES_SETTING, {str(year): "31.40", str(year + 1): "31.40"}
    )

    hr = await make_user(
        "recruiter", grants=("people_admin",), profile=False, name="Ewa Kadrowa"
    )
    hr2 = await make_user(
        "recruiter", grants=("people_admin",), profile=False, name="Kamila Kadrowa"
    )
    payroll = await make_user(
        "recruiter", grants=("payroll",), profile=False, name="Filip Finansowy"
    )
    boss = await make_user("recruiter", employment_type="uop", name="Marta Przełożona")
    mandate = await make_user(
        "employee",
        employment_type="zlecenie",
        manager_user_id=boss.id,
        name="Celina Testowa",
    )
    b2b = await make_user(
        "employee", employment_type="b2b", manager_user_id=boss.id, name="Ola Nowak"
    )
    attached = Attached()

    async def fake_attach(db, agreement, *, content, mime, actor_id):
        if attached.fail is not None:
            raise attached.fail
        attached.calls.append(
            {
                "user_id": agreement.user_id,
                "kind": agreements_effects.DOCUMENT_KINDS[agreement.kind],
                "source": agreements_effects.DOCUMENT_SOURCE,
                "source_ref": str(agreement.id),
                "content": content,
                "mime": mime,
                "valid_until": agreement.end_date,
                "actor_id": actor_id,
            }
        )
        return 9000 + len(attached.calls)

    async def fake_pdf(document: bytes) -> bytes:
        return FAKE_PDF

    monkeypatch.setattr(agreements_effects, "attach_document", fake_attach)
    monkeypatch.setattr(cv_pdf_preview, "pdf_from_docx", fake_pdf)
    people = [hr.id, hr2.id, payroll.id, boss.id, mandate.id, b2b.id]
    await wipe(people)
    yield Team(hr, hr2, payroll, boss, mandate, b2b, attached)
    await wipe(people)
    for key, value in previous.items():
        await set_setting(key, value)


class FakeAutenti:
    """Fałszywy klient Autenti — ta sama powierzchnia co ``AutentiClient``."""

    def __init__(self) -> None:
        self.calls: list[tuple[str, Any]] = []
        self.status = "IN_PROGRESS"
        self.signed = SIGNED_PDF
        self.fail: dict[str, Exception] = {}
        self.created = 0

    async def __aenter__(self) -> "FakeAutenti":
        return self

    async def __aexit__(self, *exc: Any) -> None:
        return None

    def _maybe_fail(self, name: str) -> None:
        if name in self.fail:
            raise self.fail[name]

    async def create_document_process(self, payload, *, idempotency_key=None):
        self._maybe_fail("create")
        self.created += 1
        self.calls.append(("create", {"payload": payload, "key": idempotency_key}))
        return {"id": f"proc-test-{self.created}-{id(self) % 10000}"}

    async def upload_file(
        self, process_id, *, pdf_bytes, filename, idempotency_key=None
    ):
        self._maybe_fail("upload")
        self.calls.append(
            (
                "upload",
                {"process": process_id, "bytes": pdf_bytes, "filename": filename},
            )
        )
        return {}

    async def send(self, process_id, *, idempotency_key=None):
        self._maybe_fail("send")
        self.calls.append(("send", {"process": process_id, "key": idempotency_key}))
        return {}

    async def withdraw(self, process_id):
        self._maybe_fail("withdraw")
        self.calls.append(("withdraw", process_id))
        return {}

    async def get_process(self, process_id):
        self._maybe_fail("get")
        self.calls.append(("get", process_id))
        return {"status": self.status}

    async def download_signed_file(self, process_id):
        self._maybe_fail("download")
        self.calls.append(("download", process_id))
        return self.signed

    def names(self) -> list[str]:
        return [name for name, _ in self.calls]


@pytest.fixture(name="autenti")
def autenti_fixture(monkeypatch) -> FakeAutenti:
    fake = FakeAutenti()
    monkeypatch.setattr(app_settings, "AUTENTI_ENABLED", True)
    monkeypatch.setattr(app_settings, "AUTENTI_CLIENT_ID", "test-client")
    monkeypatch.setattr(app_settings, "AUTENTI_CLIENT_SECRET", "test-secret")
    monkeypatch.setattr(app_settings, "AUTENTI_DEFAULT_SIGNATURE_TYPE", "SES")
    monkeypatch.setattr(agreements_signing, "client_factory", lambda: fake)
    return fake


def autenti_error(status: int = 500) -> AutentiError:
    return AutentiError(status, "odpowiedź z danymi, których nie pokazujemy")


__all__ = [
    "BODY",
    "FAKE_PDF",
    "PREFIX",
    "SCAN_PDF",
    "SIGNED_PDF",
    "FakeAutenti",
    "Team",
    "agreements_on_fixture",
    "autenti_fixture",
    "autenti_error",
    "docx_text",
    "first_of_next_month",
    "make_docx",
    "make_template",
]
