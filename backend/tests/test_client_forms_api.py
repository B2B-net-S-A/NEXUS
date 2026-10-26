"""0522: trasy formularza CV klienta pary osoba × rekrutacja (z bazą, CI).

Model zastąpiony atrapą (``client_form_fill._call_model``); wzór to fikcyjny
formularz w kształcie SWZ CEZ (``tests/_client_form_docx.py``).
"""

from __future__ import annotations

import io
import json
import uuid
from contextlib import asynccontextmanager
from pathlib import Path

import docx
import pytest
from sqlalchemy import text

from app.core.database import AsyncSessionLocal
from app.models.candidate import Candidate
from app.models.client import Client
from app.models.job import Job, JobStatus
from app.models.recruitment_pipeline import CandidateStage, PipelineStage
from app.models.user import UserRole
from app.services import client_form_fill, storage_service
from tests._client_form_docx import form_docx
from tests._jarvis_helpers import make_user

CV = (
    "Anna Fikcyjna\n07.2023 – obecnie Automation Test Engineer / Bank Przykładowy S.A.\n"
    "Umiejętności: SQL, Selenium WebDriver"
)


@pytest.fixture
def storage(tmp_path, monkeypatch) -> Path:
    monkeypatch.setattr(storage_service, "STORAGE_ROOT", tmp_path)
    monkeypatch.setattr(storage_service, "JOB_FILES_DIR", tmp_path / "job_files")
    return tmp_path


@asynccontextmanager
async def _no_quota(*_args, **_kwargs):
    yield


@pytest.fixture
def model(monkeypatch):
    answers: list[dict] = []
    monkeypatch.setattr(client_form_fill, "api_key_configured", lambda _m: True)
    monkeypatch.setattr(client_form_fill, "ai_feature", _no_quota)
    monkeypatch.setattr(
        client_form_fill, "_call_model", lambda *_a: json.dumps(answers.pop(0))
    )
    return answers


async def _pair(document: str = "client_form") -> tuple[int, int]:
    async with AsyncSessionLocal() as db:
        client = Client(
            name=f"Formularz {uuid.uuid4().hex[:6]}", cv_client_document=document
        )
        db.add(client)
        await db.flush()
        job = Job(
            title=f"Tester {uuid.uuid4().hex[:6]}",
            status=JobStatus.published,
            client_id=client.id,
        )
        candidate = Candidate(
            name="Anna",
            lastname=f"Fikcyjna-{uuid.uuid4().hex[:6]}",
            email=f"form-{uuid.uuid4().hex[:6]}@example.com",
            raw_cv_text=CV,
        )
        db.add_all([job, candidate])
        await db.flush()
        stage = CandidateStage(
            candidate_id=candidate.id, job_id=job.id, stage=PipelineStage.verified
        )
        db.add(stage)
        await db.commit()
        return job.id, stage.id


async def _add_template(app_client, headers, job_id: int) -> None:
    resp = await app_client.post(
        f"/api/jobs/{job_id}/files",
        files={
            "file": (
                "Zal 4b Formularz CV.docx",
                form_docx(),
                "application/octet-stream",
            )
        },
        headers=headers,
    )
    assert resp.status_code == 201 and resp.json()["client_form"] is True, resp.text


@pytest.mark.asyncio
async def test_without_template_the_form_says_where_to_add_it(
    app_client, app_auth_headers, storage
):
    _, stage_id = await _pair()
    url = f"/api/candidates/stages/{stage_id}/client-form"
    got = await app_client.get(url, headers=app_auth_headers)
    assert got.status_code == 200, got.text
    assert got.json()["status"] == "no_template"
    assert got.json()["client_document"] == "client_form"
    refused = await app_client.post(f"{url}/fill", headers=app_auth_headers)
    assert refused.status_code == 409
    assert "Plikach rekrutacji" in refused.json()["detail"]


@pytest.mark.asyncio
async def test_fill_edit_refill_and_download(
    app_client, app_auth_headers, storage, model
):
    job_id, stage_id = await _pair()
    await _add_template(app_client, app_auth_headers, job_id)
    url = f"/api/candidates/stages/{stage_id}/client-form"
    assert (await app_client.get(url, headers=app_auth_headers)).json()[
        "status"
    ] == "not_filled"

    model.append(
        {
            "fields": [
                {
                    "id": "t1.r1.c1",
                    "value": "Bank Przykładowy S.A.",
                    "quote": "Bank Przykładowy S.A.",
                },
                {"id": "t2.r2.c2", "value": "TAK", "quote": "SQL"},
                {"id": "t2.r5.c3", "value": "9", "quote": "Selenium WebDriver"},
            ]
        }
    )
    filled = await app_client.post(f"{url}/fill", headers=app_auth_headers)
    assert filled.status_code == 200, filled.text
    body = filled.json()
    values = {item["id"]: item for item in body["fields"]}
    assert body["status"] == "filled" and body["message"] is None
    assert values["p2"]["value"].startswith("Anna Fikcyjna-")
    assert values["t1.r1.c1"]["source"] == "cv"
    assert values["t2.r5.c3"]["value"] is None  # ocena bez liczby w cytacie
    assert "t2.r5.c3" in body["missing"]

    edited = await app_client.put(
        url,
        json={"fields": {"t2.r5.c3": "6", "t1.r1.c1": ""}},
        headers=app_auth_headers,
    )
    assert edited.status_code == 200, edited.text
    edited_values = {item["id"]: item for item in edited.json()["fields"]}
    assert edited_values["t2.r5.c3"]["source"] == "recruiter"
    assert edited_values["t1.r1.c1"]["value"] is None
    unknown = await app_client.put(
        url, json={"fields": {"t99.r1.c1": "x"}}, headers=app_auth_headers
    )
    assert unknown.status_code == 422

    # Ponowne wypełnienie nie zdejmuje tego, co wpisał rekruter.
    model.append({"fields": [{"id": "t2.r5.c3", "value": "2", "quote": "SQL"}]})
    refilled = (await app_client.post(f"{url}/fill", headers=app_auth_headers)).json()
    assert {item["id"]: item for item in refilled["fields"]}["t2.r5.c3"]["value"] == "6"

    download = await app_client.get(f"{url}/docx", headers=app_auth_headers)
    assert download.status_code == 200, download.text
    assert download.headers["content-disposition"].startswith("attachment")
    document = docx.Document(io.BytesIO(download.content))
    assert document.tables[2].rows[5].cells[3].text == "6"
    assert (
        document.tables[2].rows[2].cells[2].text == ""
    )  # TAK z 1. wypełnienia zniknął

    async with AsyncSessionLocal() as db:
        actions = (
            (
                await db.execute(
                    text(
                        "SELECT action FROM activities WHERE entity_type = 'candidate' "
                        "AND action LIKE 'client_form_%' AND details->>'job_id' = :job "
                        "ORDER BY id"
                    ),
                    {"job": str(job_id)},
                )
            )
            .scalars()
            .all()
        )
    assert actions == ["client_form_filled", "client_form_saved", "client_form_filled"]


@pytest.mark.asyncio
async def test_model_down_leaves_a_form_to_fill_by_hand(
    app_client, app_auth_headers, storage, monkeypatch
):
    monkeypatch.setattr(client_form_fill, "api_key_configured", lambda _m: False)
    job_id, stage_id = await _pair()
    await _add_template(app_client, app_auth_headers, job_id)
    resp = await app_client.post(
        f"/api/candidates/stages/{stage_id}/client-form/fill", headers=app_auth_headers
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["message"] == client_form_fill.MSG_UNAVAILABLE
    filled = [item for item in resp.json()["fields"] if item["value"]]
    assert [item["id"] for item in filled] == ["p2"]


@pytest.mark.asyncio
async def test_roles_without_candidate_documents_are_refused(
    app_client, app_auth_headers, storage
):
    _, stage_id = await _pair()
    _, viewer = await make_user(UserRole.user)
    url = f"/api/candidates/stages/{stage_id}/client-form"
    assert (await app_client.get(url, headers=viewer)).status_code == 403
    assert (await app_client.post(f"{url}/fill", headers=viewer)).status_code == 403
    assert (
        await app_client.get(
            "/api/candidates/stages/999999999/client-form", headers=app_auth_headers
        )
    ).status_code == 404
