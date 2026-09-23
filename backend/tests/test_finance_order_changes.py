"""Finanse → Zmiany w zamówieniach: dziennik zmian, Braki i widok miesiąca.

Baza testowa jest wspólna dla całego przebiegu i nie jest czyszczona, więc
każdy test pracuje na własnym kliencie i na datach losowanych z dalekich lat —
a asercje filtrują wyniki po swoim kliencie.
"""

from __future__ import annotations

import io
import random
import uuid
from dataclasses import replace
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.models.client_order import ClientOrder, ClientOrderStatus
from app.models.contract import Contract, ContractStatus, ContractType, RateUnit
from app.services.order_facts import OrderFact, covers_after, previous_of, successor_of

pytestmark = pytest.mark.asyncio


def _far_day(first_year: int, last_year: int) -> date:
    """Losowy dzień z dalekich lat — sąsiednie testy nie trafią w ten sam."""
    start = date(first_year, 1, 1).toordinal()
    end = date(last_year, 12, 20).toordinal()
    return date.fromordinal(random.randint(start, end))


# ── Reguły następcy (bez bazy) ──────────────────────────────────────────────


def _fact(order_id: int, **overrides) -> OrderFact:
    values = dict(
        order_id=order_id,
        order_group_id=None,
        contract_id=1,
        candidate_id=5,
        client_id=10,
        client_name="Klient",
        consultant_name="Jan Nowak",
        status="active",
        start=date(2026, 1, 1),
        end=date(2026, 8, 31),
        number=f"NB-{order_id}",
        order_type="periodic",
        rate_cost=Decimal("100"),
        rate_revenue=Decimal("150"),
        rate_unit="hourly",
        currency="PLN",
        created_at=datetime(2026, 1, 1, tzinfo=timezone.utc),
    )
    values.update(overrides)
    return OrderFact(**values)


def test_successor_counts_active_future_and_draft_but_not_cancelled_or_finished():
    ended = _fact(1)
    future = _fact(2, start=date(2026, 10, 1), end=date(2026, 12, 31))
    draft = _fact(3, status="draft", start=None, end=None)
    cancelled = _fact(4, status="cancelled", start=date(2026, 9, 1), end=None)
    finished_undated = _fact(5, status="completed", start=None, end=None)

    assert successor_of(ended, [ended, future]) == future
    assert successor_of(ended, [ended, draft]) == draft
    assert successor_of(ended, [ended, cancelled]) is None
    assert successor_of(ended, [ended, finished_undated]) is None
    assert not covers_after(_fact(6, end=date(2026, 8, 31)), date(2026, 8, 31))


def test_successor_ignores_the_order_itself_unless_it_was_extended():
    ended = _fact(1)
    extended = replace(ended, end=date(2026, 12, 31))
    assert successor_of(ended, [ended]) is None
    assert successor_of(ended, [extended], include_self=True) == extended


def test_continuation_means_previous_order_ended_within_a_month_before_start():
    new = _fact(2, start=date(2026, 9, 1), end=date(2026, 12, 31))
    recent = _fact(1, end=date(2026, 8, 31))
    long_ago = _fact(3, end=date(2026, 6, 30))
    assert previous_of(new, [new, recent]) == recent
    assert previous_of(new, [new, long_ago]) is None


# ── Dane w bazie ────────────────────────────────────────────────────────────


async def _seed(
    *,
    start: date,
    end: date,
    status: ClientOrderStatus = ClientOrderStatus.active,
    contract_status: ContractStatus = ContractStatus.active,
    client_name: str | None = None,
    contract_rate_client: Decimal | None = Decimal("1340.000"),
) -> dict[str, int]:
    """Klient + kandydat + kontrakt + jedno zamówienie.

    ``contract_rate_client=None`` zostawia kontrakt bez stawki przychodowej —
    stan realny (podpisana umowa B2B ma samą stawkę kosztową do czasu
    zamówienia) i jedyny, który NIE wchodzi do globalnego rankingu
    `/api/contract-analytics/margin-by-client`. Ranking tnie się na 20
    klientach o najwyższej marży, a baza testowa nie jest czyszczona, więc
    każdy dodatkowy wyceniony klient wypycha z niego cudzy wiersz i wywraca
    `test_contract_analytics`. Widok Finansów i tak czyta stawki
    z ZAMÓWIENIA, nie z kontraktu — testy tego modułu nic na tym nie tracą.
    """

    from app.core.database import AsyncSessionLocal
    from app.models.candidate import Candidate
    from app.models.client import Client

    suffix = uuid.uuid4().hex[:8]
    async with AsyncSessionLocal() as db:
        client = Client(name=client_name or f"FinOrders {suffix}")
        candidate = Candidate(
            name="Jan", lastname=f"Nowak-{suffix}", email=f"foc-{suffix}@example.com"
        )
        db.add_all([client, candidate])
        await db.flush()
        contract = Contract(
            candidate_id=candidate.id,
            client_id=client.id,
            contract_type=ContractType.b2b,
            status=contract_status,
            start_date=start,
            rate_candidate=Decimal("960.000"),
            rate_client=contract_rate_client,
            rate_unit=RateUnit.daily,
            currency="PLN",
            rate_candidate_currency="PLN",
        )
        db.add(contract)
        await db.flush()
        order = _order(client.id, contract.id, start=start, end=end, status=status)
        db.add(order)
        await db.commit()
        return {
            "client_id": client.id,
            "candidate_id": candidate.id,
            "contract_id": contract.id,
            "order_id": order.id,
        }


def _order(
    client_id: int,
    contract_id: int,
    *,
    start: date | None,
    end: date | None,
    status: ClientOrderStatus = ClientOrderStatus.active,
    title: str | None = None,
    **extra,
) -> ClientOrder:
    return ClientOrder(
        client_id=client_id,
        contract_id=contract_id,
        title=title or f"NB-{uuid.uuid4().hex[:6]}",
        status=status,
        start_date=start,
        end_date=end,
        rate_unit=RateUnit.daily,
        rate_client=Decimal("1340.000"),
        rate_candidate=Decimal("960.000"),
        rate_client_currency="PLN",
        rate_candidate_currency="PLN",
        currency="PLN",
        **extra,
    )


async def _add_order(ids: dict[str, int], **kwargs) -> int:
    from app.core.database import AsyncSessionLocal

    async with AsyncSessionLocal() as db:
        order = _order(ids["client_id"], ids["contract_id"], **kwargs)
        db.add(order)
        await db.commit()
        return order.id


async def _gap_for(order_id: int):
    from app.core.database import AsyncSessionLocal
    from app.models.order_gap import OrderGap

    async with AsyncSessionLocal() as db:
        return await db.scalar(select(OrderGap).where(OrderGap.order_id == order_id))


async def _detect(monkeypatch, *, tracking_start: date, today: date) -> None:
    from app.core.config import settings
    from app.core.database import AsyncSessionLocal
    from app.services.order_gaps import detect_order_gaps

    monkeypatch.setattr(settings, "ORDER_GAP_TRACKING_START", tracking_start)
    async with AsyncSessionLocal() as db:
        await detect_order_gaps(db, today=today)
        await db.commit()


async def _seed_dl(client_id: int) -> int:
    from app.core.database import AsyncSessionLocal
    from app.core.security import hash_password
    from app.models.team_structure import DeliveryLeadClientAssignment
    from app.models.user import User, UserRole

    suffix = uuid.uuid4().hex[:8]
    async with AsyncSessionLocal() as db:
        user = User(
            email=f"foc-dl-{suffix}@example.com",
            password_hash=hash_password(f"T3st_{suffix}!PassX"),
            name="DL Zamówień",
            role=UserRole.delivery_lead,
            is_active=True,
            profile_completed=True,
        )
        db.add(user)
        await db.flush()
        db.add(
            DeliveryLeadClientAssignment(
                client_id=client_id, delivery_lead_user_id=user.id
            )
        )
        await db.commit()
        return user.id


async def _login(app_client: AsyncClient, role: str) -> dict[str, str]:
    from app.core.database import AsyncSessionLocal
    from app.core.security import hash_password
    from app.models.user import User, UserRole

    suffix = uuid.uuid4().hex[:8]
    email = f"foc-{role}-{suffix}@example.com"
    password = f"T3st_{suffix}!PassX"
    async with AsyncSessionLocal() as db:
        db.add(
            User(
                email=email,
                password_hash=hash_password(password),
                name=f"Foc {role}",
                role=UserRole(role),
                is_active=True,
                profile_completed=True,
            )
        )
        await db.commit()
    resp = await app_client.post(
        "/api/auth/login", json={"email": email, "password": password}
    )
    assert resp.status_code == 200, resp.text
    return {"Authorization": f"Bearer {resp.json()['access_token']}"}


# ── Dziennik zmian ──────────────────────────────────────────────────────────


async def _events(order_id: int):
    from app.core.database import AsyncSessionLocal
    from app.models.order_change_event import OrderChangeEvent

    async with AsyncSessionLocal() as db:
        return list(
            (
                await db.scalars(
                    select(OrderChangeEvent)
                    .where(OrderChangeEvent.order_id == order_id)
                    .order_by(OrderChangeEvent.id)
                )
            ).all()
        )


async def test_rate_and_end_date_change_is_recorded_with_old_and_new_value(
    app_client: AsyncClient, app_auth_headers: dict
):
    day = _far_day(2081, 2084)
    ids = await _seed(start=day, end=day + timedelta(days=90))

    resp = await app_client.patch(
        f"/api/clients/{ids['client_id']}/orders/{ids['order_id']}",
        json={
            "rate_client": "1400",
            "end_date": (day + timedelta(days=180)).isoformat(),
        },
        headers=app_auth_headers,
    )
    assert resp.status_code == 200, resp.text

    events = {event.field: event for event in await _events(ids["order_id"])}
    revenue = events["rate_revenue"]
    assert (revenue.old_amount, revenue.new_amount) == (
        Decimal("1340"),
        Decimal("1400"),
    )
    assert (revenue.old_unit, revenue.new_unit) == ("daily", "daily")
    assert revenue.source == "user" and revenue.created_by_user_id is not None
    end = events["end_date"]
    assert (end.old_date, end.new_date) == (
        day + timedelta(days=90),
        day + timedelta(days=180),
    )
    assert "rate_cost" not in events, "stawka kosztowa się nie zmieniła"


async def test_filling_a_draft_is_not_a_change(
    app_client: AsyncClient, app_auth_headers: dict
):
    day = _far_day(2081, 2084)
    ids = await _seed(
        start=day, end=day + timedelta(days=90), status=ClientOrderStatus.draft
    )

    resp = await app_client.patch(
        f"/api/clients/{ids['client_id']}/orders/{ids['order_id']}",
        json={
            "rate_client": "1500",
            "end_date": (day + timedelta(days=120)).isoformat(),
        },
        headers=app_auth_headers,
    )
    assert resp.status_code == 200, resp.text
    assert await _events(ids["order_id"]) == []


async def test_system_write_without_logged_user_is_marked_as_system():
    from app.core.database import AsyncSessionLocal

    day = _far_day(2081, 2084)
    ids = await _seed(start=day, end=day + timedelta(days=90))
    async with AsyncSessionLocal() as db:
        order = await db.get(ClientOrder, ids["order_id"])
        order.rate_candidate = Decimal("1000")
        await db.commit()

    [event] = await _events(ids["order_id"])
    assert event.field == "rate_cost" and event.source == "system"
    assert event.created_by_user_id is None


async def test_rate_reverted_by_contract_sync_in_the_same_request_is_not_a_change(
    app_client: AsyncClient, app_auth_headers: dict
):
    """Stan pośredni (100 → 120 → 100 przez synchronizację) nie jest zmianą."""
    day = _far_day(2081, 2084)
    ids = await _seed(start=day, end=day + timedelta(days=90))

    resp = await app_client.patch(
        f"/api/clients/{ids['client_id']}/orders/{ids['order_id']}",
        # Koszt należy do umowy — synchronizacja przywraca 960 przed commitem.
        json={"rate_candidate": "999"},
        headers=app_auth_headers,
    )
    assert resp.status_code == 200, resp.text
    assert Decimal(str(resp.json()["rate_candidate"])) == Decimal("960")
    assert [event.field for event in await _events(ids["order_id"])] == []


async def test_rolled_back_savepoint_does_not_drop_earlier_changes_of_the_transaction():
    """FIN-CHG-3 (audyt 22.09): wycofany SAVEPOINT nie kasuje zmian sprzed niego.

    Writerzy (poczta zamówień, synchronizacja) otwierają ``begin_nested`` i przy
    odmowie wycofują savepoint. Dziennik zapamiętywał stare wartości dla całej
    transakcji, a wycofanie savepointu czyściło je w całości — zmiana stawki
    zapisana PRZED savepointem ginęła z „Zmian w zamówieniach".
    """
    from app.core.database import AsyncSessionLocal

    day = _far_day(2081, 2084)
    ids = await _seed(start=day, end=day + timedelta(days=90))
    async with AsyncSessionLocal() as db:
        order = await db.get(ClientOrder, ids["order_id"])
        order.rate_candidate = Decimal("1000")
        await db.flush()
        async with db.begin_nested() as savepoint:
            order.end_date = day + timedelta(days=120)
            await db.flush()
            await savepoint.rollback()
        await db.commit()

    events = await _events(ids["order_id"])
    assert [event.field for event in events] == ["rate_cost"], events
    assert (events[0].old_amount, events[0].new_amount) == (
        Decimal("960"),
        Decimal("1000"),
    )


# ── Braki ───────────────────────────────────────────────────────────────────


async def test_order_ended_without_successor_becomes_an_open_gap_with_dl_alert(
    monkeypatch,
):
    from app.core.database import AsyncSessionLocal
    from app.models.dl_alert import ALERT_ORDER_MISSING_SUCCESSOR, DlAlert
    from app.models.notification import Notification, NotificationType

    end = _far_day(2085, 2088)
    ids = await _seed(start=end - timedelta(days=60), end=end)
    dl_id = await _seed_dl(ids["client_id"])

    await _detect(monkeypatch, tracking_start=end, today=end + timedelta(days=1))

    gap = await _gap_for(ids["order_id"])
    assert gap is not None and gap.status == "open"
    assert gap.detected_on == end + timedelta(days=1)
    async with AsyncSessionLocal() as db:
        alerts = (
            await db.scalars(
                select(DlAlert).where(
                    DlAlert.user_id == dl_id,
                    DlAlert.alert_type == ALERT_ORDER_MISSING_SUCCESSOR,
                )
            )
        ).all()
        bell = (
            await db.scalars(
                select(Notification).where(
                    Notification.user_id == dl_id,
                    Notification.notification_type
                    == NotificationType.order_missing_successor,
                )
            )
        ).all()
    assert (
        len(alerts) == 1
        # Panel „Moi klienci": CTA prowadzi do konkretnego zamówienia.
        and alerts[0].link
        == f"/clients/{ids['client_id']}?tab=zamowienia&order={alerts[0].order_id}"
    )
    assert len(bell) == 1

    # Drugi przebieg nie dubluje ani braku, ani powiadomień.
    await _detect(monkeypatch, tracking_start=end, today=end + timedelta(days=2))
    async with AsyncSessionLocal() as db:
        again = (
            await db.scalars(select(Notification).where(Notification.user_id == dl_id))
        ).all()
    assert len(again) == 1


async def test_gap_is_not_detected_on_the_last_day_of_the_order(monkeypatch):
    end = _far_day(2085, 2088)
    ids = await _seed(start=end - timedelta(days=60), end=end)
    await _detect(monkeypatch, tracking_start=end, today=end)
    assert await _gap_for(ids["order_id"]) is None


async def test_draft_successor_counts_as_planned(monkeypatch):
    end = _far_day(2085, 2088)
    ids = await _seed(start=end - timedelta(days=60), end=end)
    await _add_order(ids, start=None, end=None, status=ClientOrderStatus.draft)
    await _detect(monkeypatch, tracking_start=end, today=end + timedelta(days=1))
    assert await _gap_for(ids["order_id"]) is None


async def test_terminated_contract_is_a_deliberate_end_not_a_gap(monkeypatch):
    from app.core.database import AsyncSessionLocal

    end = _far_day(2085, 2088)
    ids = await _seed(start=end - timedelta(days=60), end=end)
    async with AsyncSessionLocal() as db:
        contract = await db.get(Contract, ids["contract_id"])
        contract.terminated_at = end - timedelta(days=20)
        contract.end_date = end
        await db.commit()
    await _detect(monkeypatch, tracking_start=end, today=end + timedelta(days=1))
    assert await _gap_for(ids["order_id"]) is None


async def test_replaced_consultant_is_not_a_gap(monkeypatch):
    end = _far_day(2085, 2088)
    ids = await _seed(start=end - timedelta(days=60), end=end)
    other = await _seed(start=end + timedelta(days=1), end=end + timedelta(days=90))
    from app.core.database import AsyncSessionLocal

    async with AsyncSessionLocal() as db:
        replacement = await db.get(ClientOrder, other["order_id"])
        replacement.predecessor_order_id = ids["order_id"]
        await db.commit()
    await _detect(monkeypatch, tracking_start=end, today=end + timedelta(days=1))
    assert await _gap_for(ids["order_id"]) is None


async def test_successor_added_after_the_day_of_detection_is_recorded_as_late(
    monkeypatch,
):
    """Przebieg po deployu: następca jest, ale dodany po terminie."""
    end = _far_day(2001, 2020)
    ids = await _seed(start=end - timedelta(days=60), end=end)
    late_id = await _add_order(
        ids, start=end + timedelta(days=10), end=end + timedelta(days=100)
    )
    await _detect(monkeypatch, tracking_start=end, today=end + timedelta(days=30))

    gap = await _gap_for(ids["order_id"])
    assert gap is not None and gap.status == "filled_late"
    assert gap.resolved_order_id == late_id


async def test_successor_added_on_the_day_of_detection_is_on_time(monkeypatch):
    """UAT B69: następca z dnia wykrycia to nie opóźnienie „0 dni po terminie”."""
    from app.services.order_gaps import local_day_start

    end = _far_day(2085, 2088)
    ids = await _seed(start=end - timedelta(days=60), end=end)
    await _add_order(
        ids,
        start=end + timedelta(days=1),
        end=end + timedelta(days=90),
        created_at=local_day_start(end + timedelta(days=1)) + timedelta(hours=15),
    )
    await _detect(monkeypatch, tracking_start=end, today=end + timedelta(days=5))

    assert await _gap_for(ids["order_id"]) is None


async def test_gap_filled_on_the_day_of_detection_is_not_reported(
    monkeypatch, app_client: AsyncClient, app_auth_headers: dict
):
    """Pętla założyła brak rano, DL dodał zamówienie tego samego dnia — raport
    Finansów nie pokazuje tego jako opóźnienia (UAT B69)."""
    from app.core.database import AsyncSessionLocal
    from app.models.order_gap import OrderGap
    from app.services.order_gaps import local_day_start

    end = _far_day(2085, 2088)
    ids = await _seed(start=end - timedelta(days=60), end=end)
    await _detect(monkeypatch, tracking_start=end, today=end + timedelta(days=1))
    open_gap = await _gap_for(ids["order_id"])
    assert open_gap is not None and open_gap.status == "open"
    async with AsyncSessionLocal() as db:
        gap = await db.get(OrderGap, open_gap.id)
        gap.status = "filled_late"
        gap.resolved_order_number = "NB-SAME-DAY"
        gap.resolved_at = local_day_start(gap.detected_on) + timedelta(hours=11)
        await db.commit()

    detected = end + timedelta(days=1)
    resp = await app_client.get(
        f"/api/finance/order-changes?year={detected.year}&month={detected.month}",
        headers=app_auth_headers,
    )
    assert resp.status_code == 200, resp.text
    assert ids["order_id"] not in {g["order_id"] for g in resp.json()["gaps"]}


async def test_adding_the_order_later_keeps_the_gap_as_filled_late(
    monkeypatch, app_client: AsyncClient, app_auth_headers: dict
):
    """Zapis zamówienia przez API zamyka brak od razu — wpis zostaje."""
    from app.core.database import AsyncSessionLocal
    from app.models.dl_alert import ALERT_ORDER_MISSING_SUCCESSOR, DlAlert

    end = _far_day(2085, 2088)
    ids = await _seed(start=end - timedelta(days=60), end=end)
    await _seed_dl(ids["client_id"])
    await _detect(monkeypatch, tracking_start=end, today=end + timedelta(days=1))
    assert (await _gap_for(ids["order_id"])).status == "open"

    resp = await app_client.patch(
        f"/api/clients/{ids['client_id']}/orders/{ids['order_id']}",
        json={"end_date": (end + timedelta(days=90)).isoformat()},
        headers=app_auth_headers,
    )
    assert resp.status_code == 200, resp.text

    gap = await _gap_for(ids["order_id"])
    assert gap.status == "filled_late" and gap.resolved_at is not None
    async with AsyncSessionLocal() as db:
        statuses = set(
            (
                await db.scalars(
                    select(DlAlert.status).where(
                        DlAlert.order_id == ids["order_id"],
                        DlAlert.alert_type == ALERT_ORDER_MISSING_SUCCESSOR,
                    )
                )
            ).all()
        )
    assert statuses == {"handled"}


# ── Widok miesiąca ──────────────────────────────────────────────────────────


async def test_month_view_lists_entries_exits_and_gaps_with_matching_counts(
    monkeypatch, app_client: AsyncClient, app_auth_headers: dict
):
    month_day = _far_day(2089, 2094)
    first = month_day.replace(day=1)
    # Osoba A: zamówienie kończy się 10. dnia, kolejne od 11. (kontynuacja).
    a = await _seed(start=first - timedelta(days=60), end=first + timedelta(days=9))
    a_next = await _add_order(
        a, start=first + timedelta(days=10), end=first + timedelta(days=120)
    )
    # Osoba B: kończy się 5. dnia i nic dalej (brak).
    b = await _seed(start=first - timedelta(days=60), end=first + timedelta(days=4))
    # Osoba A zaczyna w tym miesiącu projekt u innego klienta.
    from app.core.database import AsyncSessionLocal
    from app.models.client import Client

    async with AsyncSessionLocal() as db:
        other_client = Client(name=f"FinOrders inny {uuid.uuid4().hex[:6]}")
        db.add(other_client)
        await db.flush()
        other_contract = Contract(
            candidate_id=a["candidate_id"],
            client_id=other_client.id,
            contract_type=ContractType.b2b,
            status=ContractStatus.active,
            start_date=first,
            rate_candidate=Decimal("960.000"),
            rate_unit=RateUnit.daily,
            currency="PLN",
        )
        db.add(other_contract)
        await db.flush()
        extra = _order(
            other_client.id,
            other_contract.id,
            start=first + timedelta(days=14),
            end=first + timedelta(days=60),
        )
        db.add(extra)
        await db.commit()
        extra_id = extra.id

    await _detect(monkeypatch, tracking_start=first, today=first + timedelta(days=20))

    resp = await app_client.get(
        f"/api/finance/order-changes?year={first.year}&month={first.month}",
        headers=app_auth_headers,
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    counts = body["counts"]
    assert counts == {
        "changes": len(body["changes"]),
        "entries": len(body["entries"]),
        "exits": len(body["exits"]),
        "ending": len(body["ending_orders"]),
        "gaps": len(body["gaps"]),
    }

    ours = {a["client_id"], b["client_id"], other_client.id}
    # Żadna z tych osób nie jest nowa — wszystkie trzy zamówienia to zmiany
    # w trwającej współpracy, a nie wejścia.
    assert [e["order_id"] for e in body["entries"] if e["client_id"] in ours] == []

    changes = {
        c["order_id"]: c
        for c in body["changes"]
        if c["client_id"] in ours
        and c["kind"] in ("order_continuation", "additional_project")
    }
    assert changes[a_next]["kind"] == "order_continuation"
    assert changes[a_next]["previous_order_number"]
    assert changes[a_next]["previous_end_date"]
    assert changes[a_next]["rate_revenue"] == 1340.0
    assert changes[extra_id]["kind"] == "additional_project"
    assert changes[extra_id]["other_client_names"]
    assert changes[extra_id]["effective_date"]

    exits = {e["order_id"]: e for e in body["exits"] if e["client_id"] in ours}
    ending = {e["order_id"]: e for e in body["ending_orders"] if e["client_id"] in ours}
    # Kontynuacja to NIE zejście — osoba pracuje dalej pod nowym numerem.
    assert a["order_id"] not in exits
    # Osobie B skończyło się zamówienie, ale nikt nie zapisał końca współpracy:
    # to kończące się zamówienie, nie zejście.
    assert b["order_id"] not in exits
    assert ending[b["order_id"]]["verdict"] == "no_successor"

    gaps = [g for g in body["gaps"] if g["client_id"] in ours]
    assert [g["order_id"] for g in gaps] == [b["order_id"]]
    assert gaps[0]["status"] == "open"

    export = await app_client.get(
        f"/api/finance/order-changes/export?year={first.year}&month={first.month}",
        headers=app_auth_headers,
    )
    assert export.status_code == 200, export.text
    workbook = load_workbook(io.BytesIO(export.content))
    assert [name.split(" (")[0] for name in workbook.sheetnames] == [
        "Zmiany",
        "Wejścia",
        "Zejścia",
        "Kończące się zam.",
        "Braki",
    ]

    # Eksport jednej zakładki: jeden arkusz i nazwa pliku mówiąca która.
    one = await app_client.get(
        f"/api/finance/order-changes/export"
        f"?year={first.year}&month={first.month}&tab=exits",
        headers=app_auth_headers,
    )
    assert one.status_code == 200, one.text
    assert "Zejscia" in one.headers["content-disposition"]
    assert [
        name.split(" (")[0]
        for name in load_workbook(io.BytesIO(one.content)).sheetnames
    ] == ["Zejścia"]


# ── Kwalifikacja zakładek (korekta 09.2026) ─────────────────────────────────


async def _month(app_client: AsyncClient, headers: dict, day: date, **params) -> dict:
    query = "".join(f"&{key}={value}" for key, value in params.items())
    resp = await app_client.get(
        f"/api/finance/order-changes?year={day.year}&month={day.month}{query}",
        headers=headers,
    )
    assert resp.status_code == 200, resp.text
    return resp.json()


async def _other_client_order(
    ids: dict[str, int], *, start: date, end: date | None, client_name: str
) -> tuple[int, int]:
    """Zamówienie tej samej OSOBY u innego klienta (nowy klient + kontrakt)."""

    from app.core.database import AsyncSessionLocal
    from app.models.client import Client

    async with AsyncSessionLocal() as db:
        client = Client(name=f"{client_name} {uuid.uuid4().hex[:6]}")
        db.add(client)
        await db.flush()
        contract = Contract(
            candidate_id=ids["candidate_id"],
            client_id=client.id,
            contract_type=ContractType.b2b,
            status=ContractStatus.active,
            start_date=start,
            rate_candidate=Decimal("960.000"),
            rate_unit=RateUnit.daily,
            currency="PLN",
        )
        db.add(contract)
        await db.flush()
        order = _order(client.id, contract.id, start=start, end=end)
        db.add(order)
        await db.commit()
        return client.id, order.id


async def test_first_ever_order_is_an_entry_and_carries_the_draft_status(
    app_client: AsyncClient, app_auth_headers: dict
):
    """Wejście = osoba, która zaczyna z nami po raz pierwszy."""

    first = _far_day(2095, 2098).replace(day=1)
    fresh = await _seed(
        contract_rate_client=None,
        start=first + timedelta(days=3),
        end=first + timedelta(days=200),
        status=ClientOrderStatus.draft,
    )

    body = await _month(app_client, app_auth_headers, first)

    entries = {e["order_id"]: e for e in body["entries"]}
    assert entries[fresh["order_id"]]["status"] == "draft"
    assert fresh["order_id"] not in {c["order_id"] for c in body["changes"]}


async def test_continuation_leaves_entries_and_lands_in_changes(
    app_client: AsyncClient, app_auth_headers: dict
):
    first = _far_day(2050, 2053).replace(day=1)
    ids = await _seed(
        contract_rate_client=None,
        start=first - timedelta(days=90),
        end=first + timedelta(days=2),
    )
    nxt = await _add_order(
        ids, start=first + timedelta(days=3), end=first + timedelta(days=200)
    )

    body = await _month(app_client, app_auth_headers, first)

    assert nxt not in {e["order_id"] for e in body["entries"]}
    row = next(c for c in body["changes"] if c["order_id"] == nxt)
    assert row["kind"] == "order_continuation"
    assert row["previous_order_number"]
    assert row["previous_end_date"] == (first + timedelta(days=2)).isoformat()


async def test_moving_to_another_client_is_a_client_change_not_an_entry(
    app_client: AsyncClient, app_auth_headers: dict
):
    """Konsultant kończy u klienta A i startuje u B — to zmiana klienta."""

    first = _far_day(2054, 2057).replace(day=1)
    ids = await _seed(
        contract_rate_client=None,
        start=first - timedelta(days=90),
        end=first + timedelta(days=2),
    )
    _, moved = await _other_client_order(
        ids,
        start=first + timedelta(days=3),
        end=first + timedelta(days=200),
        client_name="FinOrders docelowy",
    )

    body = await _month(app_client, app_auth_headers, first)

    assert moved not in {e["order_id"] for e in body["entries"]}
    row = next(c for c in body["changes"] if c["order_id"] == moved)
    assert row["kind"] == "client_change"
    assert row["previous_client_name"]
    assert row["previous_client_name"] != row["client_name"]


async def test_two_clients_the_same_day_still_shows_a_new_person_once(
    app_client: AsyncClient, app_auth_headers: dict
):
    """Osoba naprawdę nowa nie może zniknąć z Wejść przez własne drugie zamówienie."""

    first = _far_day(2058, 2061).replace(day=1)
    day = first + timedelta(days=5)
    ids = await _seed(
        contract_rate_client=None,
        start=day,
        end=first + timedelta(days=200),
    )
    _, twin = await _other_client_order(
        ids, start=day, end=first + timedelta(days=200), client_name="FinOrders drugi"
    )

    body = await _month(app_client, app_auth_headers, first)

    entered = {e["order_id"] for e in body["entries"]}
    assert ids["order_id"] in entered
    assert twin not in entered
    assert next(c for c in body["changes"] if c["order_id"] == twin)["kind"] == (
        "additional_project"
    )


async def test_termination_with_a_live_successor_is_not_an_exit(
    app_client: AsyncClient, app_auth_headers: dict
):
    """Wypowiedziana umowa z żywym następcą — osoba pracuje dalej."""

    from app.core.database import AsyncSessionLocal

    first = _far_day(2062, 2065).replace(day=1)
    ends = first + timedelta(days=4)
    ids = await _seed(
        contract_rate_client=None,
        start=first - timedelta(days=90),
        end=ends,
    )
    await _add_order(
        ids, start=ends + timedelta(days=1), end=first + timedelta(days=200)
    )
    async with AsyncSessionLocal() as db:
        contract = await db.get(Contract, ids["contract_id"])
        contract.end_date = ends
        contract.terminated_at = datetime.now(timezone.utc)
        contract.termination_reason = "consultant_resigned"
        await db.commit()

    body = await _month(app_client, app_auth_headers, first)

    assert ids["order_id"] not in {e["order_id"] for e in body["exits"]}


async def test_termination_without_a_successor_stays_an_exit(
    app_client: AsyncClient, app_auth_headers: dict
):
    from app.core.database import AsyncSessionLocal

    first = _far_day(2066, 2069).replace(day=1)
    ends = first + timedelta(days=4)
    ids = await _seed(
        contract_rate_client=None,
        start=first - timedelta(days=90),
        end=ends,
    )
    async with AsyncSessionLocal() as db:
        contract = await db.get(Contract, ids["contract_id"])
        contract.end_date = ends
        contract.terminated_at = datetime.now(timezone.utc)
        contract.termination_reason = "consultant_resigned"
        await db.commit()

    body = await _month(app_client, app_auth_headers, first)

    row = next(e for e in body["exits"] if e["order_id"] == ids["order_id"])
    assert row["verdict"] == "ended_intent"


async def test_contract_older_than_its_first_order_is_not_a_new_consultant(
    app_client: AsyncClient, app_auth_headers: dict
):
    """Rejestr zamówień jest młodszy niż współpraca, którą opisuje.

    Zgłoszenie 09.2026: konsultantka z umową bezterminową od poprzedniego roku
    dostała pierwszy wiersz zamówienia dopiero teraz i pokazała się w Wejściach
    jako „Nowy konsultant". Jedynym dowodem wcześniejszej pracy jest UMOWA.
    """

    from app.core.database import AsyncSessionLocal

    first = _far_day(2030, 2033).replace(day=1)
    ids = await _seed(
        contract_rate_client=None,
        start=first + timedelta(days=0),
        end=first + timedelta(days=200),
    )
    # Umowa trwa od poprzedniego roku, choć zamówienie zaczyna się dziś.
    since = first - timedelta(days=270)
    async with AsyncSessionLocal() as db:
        contract = await db.get(Contract, ids["contract_id"])
        contract.start_date = since
        await db.commit()

    body = await _month(app_client, app_auth_headers, first)

    assert ids["order_id"] not in {e["order_id"] for e in body["entries"]}
    row = next(c for c in body["changes"] if c["order_id"] == ids["order_id"])
    assert row["kind"] == "order_continuation"
    # Nie ma poprzedniego zamówienia — dowodem jest data startu umowy, a nie
    # puste „—", które czyta się jak utrata danych.
    assert row["previous_order_number"] is None
    assert row["engagement_since"] == since.isoformat()


async def test_contract_starting_in_the_same_month_still_counts_as_an_entry(
    app_client: AsyncClient, app_auth_headers: dict
):
    """Próg to początek MIESIĄCA, nie dzień startu zamówienia.

    Osobie faktycznie nowej zakłada się umowę razem z pierwszym zamówieniem,
    często z datą o kilka dni wcześniejszą. Próg „ściśle przed startem
    zamówienia" zdjąłby z Wejść także takie osoby.
    """

    from app.core.database import AsyncSessionLocal

    first = _far_day(2034, 2037).replace(day=1)
    ids = await _seed(
        contract_rate_client=None,
        start=first + timedelta(days=14),
        end=first + timedelta(days=200),
    )
    async with AsyncSessionLocal() as db:
        contract = await db.get(Contract, ids["contract_id"])
        contract.start_date = first + timedelta(days=1)
        await db.commit()

    body = await _month(app_client, app_auth_headers, first)

    assert ids["order_id"] in {e["order_id"] for e in body["entries"]}


async def test_draft_contract_is_not_proof_of_earlier_work(
    app_client: AsyncClient, app_auth_headers: dict
):
    """Szkic to plan, nie praca — lustro reguły dla szkiców zamówień."""

    from app.core.database import AsyncSessionLocal

    first = _far_day(2038, 2041).replace(day=1)
    ids = await _seed(
        contract_rate_client=None,
        start=first + timedelta(days=2),
        end=first + timedelta(days=200),
    )
    async with AsyncSessionLocal() as db:
        contract = await db.get(Contract, ids["contract_id"])
        contract.start_date = first - timedelta(days=200)
        contract.status = ContractStatus.draft
        await db.commit()

    body = await _month(app_client, app_auth_headers, first)

    assert ids["order_id"] in {e["order_id"] for e in body["entries"]}


async def test_earlier_contract_at_another_client_is_a_client_change(
    app_client: AsyncClient, app_auth_headers: dict
):
    """Zakończona umowa u innego klienta, bez ani jednego zamówienia w NEXUSIE."""

    from app.core.database import AsyncSessionLocal
    from app.models.client import Client

    first = _far_day(2042, 2045).replace(day=1)
    ids = await _seed(
        contract_rate_client=None,
        start=first + timedelta(days=3),
        end=first + timedelta(days=200),
    )
    previous_name = f"FinOrders poprzedni {uuid.uuid4().hex[:6]}"
    async with AsyncSessionLocal() as db:
        old_client = Client(name=previous_name)
        db.add(old_client)
        await db.flush()
        db.add(
            Contract(
                candidate_id=ids["candidate_id"],
                client_id=old_client.id,
                contract_type=ContractType.b2b,
                status=ContractStatus.ended,
                start_date=first - timedelta(days=400),
                end_date=first - timedelta(days=100),
                rate_candidate=Decimal("960.000"),
                rate_unit=RateUnit.daily,
                currency="PLN",
            )
        )
        await db.commit()

    body = await _month(app_client, app_auth_headers, first)

    assert ids["order_id"] not in {e["order_id"] for e in body["entries"]}
    row = next(c for c in body["changes"] if c["order_id"] == ids["order_id"])
    assert row["kind"] == "client_change"
    assert row["previous_client_name"] == previous_name


async def test_order_ending_at_a_live_contract_is_not_an_exit(
    app_client: AsyncClient, app_auth_headers: dict
):
    """Kończy się ZAMÓWIENIE, nie współpraca — osobna zakładka, nie Zejścia."""

    first = _far_day(2046, 2049).replace(day=1)
    ends = first + timedelta(days=25)
    ids = await _seed(
        contract_rate_client=None,
        start=first - timedelta(days=200),
        end=ends,
    )

    body = await _month(app_client, app_auth_headers, first)

    assert ids["order_id"] not in {e["order_id"] for e in body["exits"]}
    row = next(e for e in body["ending_orders"] if e["order_id"] == ids["order_id"])
    assert row["verdict"] in ("ending_pending", "no_successor")
    assert row["intent"] is None


async def test_a_live_successor_keeps_the_order_out_of_both_lists(
    app_client: AsyncClient, app_auth_headers: dict
):
    """Osoba pracuje dalej — nie jest ani zejściem, ani kończącym się zamówieniem."""

    first = _far_day(2021, 2024).replace(day=1)
    ends = first + timedelta(days=9)
    ids = await _seed(
        contract_rate_client=None,
        start=first - timedelta(days=200),
        end=ends,
    )
    await _add_order(
        ids, start=ends + timedelta(days=1), end=first + timedelta(days=200)
    )

    body = await _month(app_client, app_auth_headers, first)

    assert ids["order_id"] not in {e["order_id"] for e in body["exits"]}
    assert ids["order_id"] not in {e["order_id"] for e in body["ending_orders"]}


async def test_ending_orders_tab_filters_and_exports_like_the_others(
    app_client: AsyncClient, app_auth_headers: dict
):
    first = _far_day(2078, 2080).replace(day=1)
    mine = await _seed(
        contract_rate_client=None,
        start=first - timedelta(days=200),
        end=first + timedelta(days=20),
        client_name=f"FinOrders Kaszuby {uuid.uuid4().hex[:6]}",
    )
    other = await _seed(
        contract_rate_client=None,
        start=first - timedelta(days=200),
        end=first + timedelta(days=20),
        client_name=f"FinOrders Mazury {uuid.uuid4().hex[:6]}",
    )

    by_client = await _month(
        app_client, app_auth_headers, first, client_id=mine["client_id"]
    )
    ids = {e["order_id"] for e in by_client["ending_orders"]}
    assert mine["order_id"] in ids and other["order_id"] not in ids
    assert by_client["counts"]["ending"] == len(by_client["ending_orders"])

    export = await app_client.get(
        f"/api/finance/order-changes/export?year={first.year}&month={first.month}"
        f"&tab=ending&client_id={mine['client_id']}",
        headers=app_auth_headers,
    )
    assert export.status_code == 200, export.text
    assert "Konczace_sie_zamowienia" in export.headers["content-disposition"]
    sheet = load_workbook(io.BytesIO(export.content))["Kończące się zam. (1)"]
    assert sheet.max_row == 2  # nagłówek + jeden wiersz


# ── Filtry ──────────────────────────────────────────────────────────────────


async def test_filters_narrow_every_tab_and_its_counter(
    app_client: AsyncClient, app_auth_headers: dict
):
    first = _far_day(2070, 2073).replace(day=1)
    mine = await _seed(
        contract_rate_client=None,
        start=first + timedelta(days=2),
        end=first + timedelta(days=200),
        client_name=f"FinOrders Zagłębie {uuid.uuid4().hex[:6]}",
    )
    other = await _seed(
        contract_rate_client=None,
        start=first + timedelta(days=2),
        end=first + timedelta(days=200),
        client_name=f"FinOrders Pomorze {uuid.uuid4().hex[:6]}",
    )

    by_client = await _month(
        app_client, app_auth_headers, first, client_id=mine["client_id"]
    )
    ids = {e["order_id"] for e in by_client["entries"]}
    assert mine["order_id"] in ids and other["order_id"] not in ids
    assert by_client["counts"]["entries"] == len(by_client["entries"])
    assert all(
        item["client_id"] == mine["client_id"]
        for key in ("changes", "entries", "exits", "ending_orders", "gaps")
        for item in by_client[key]
    )

    # Szukanie po kliencie bez ogonków — „Zaglebie" ma znaleźć „Zagłębie".
    by_name = await _month(app_client, app_auth_headers, first, q="zaglebie")
    assert mine["order_id"] in {e["order_id"] for e in by_name["entries"]}
    assert other["order_id"] not in {e["order_id"] for e in by_name["entries"]}

    # Zakres dat: start obu zamówień wypada przed tą granicą.
    later = (first + timedelta(days=20)).isoformat()
    by_date = await _month(app_client, app_auth_headers, first, date_from=later)
    assert mine["order_id"] not in {e["order_id"] for e in by_date["entries"]}
    assert by_date["counts"]["entries"] == len(by_date["entries"])


async def test_export_carries_the_same_filters_as_the_screen(
    app_client: AsyncClient, app_auth_headers: dict
):
    first = _far_day(2074, 2077).replace(day=1)
    mine = await _seed(
        contract_rate_client=None,
        start=first + timedelta(days=2),
        end=first + timedelta(days=200),
    )
    await _seed(
        contract_rate_client=None,
        start=first + timedelta(days=2),
        end=first + timedelta(days=200),
    )

    export = await app_client.get(
        f"/api/finance/order-changes/export?year={first.year}&month={first.month}"
        f"&tab=entries&client_id={mine['client_id']}",
        headers=app_auth_headers,
    )
    assert export.status_code == 200, export.text
    sheet = load_workbook(io.BytesIO(export.content))["Wejścia (1)"]
    assert sheet.max_row == 2  # nagłówek + jeden wiersz


async def test_unknown_export_tab_is_refused(
    app_client: AsyncClient, app_auth_headers: dict
):
    resp = await app_client.get(
        "/api/finance/order-changes/export?year=2026&month=9&tab=wejscia",
        headers=app_auth_headers,
    )
    assert resp.status_code == 422


async def test_reversed_date_range_is_refused(
    app_client: AsyncClient, app_auth_headers: dict
):
    resp = await app_client.get(
        "/api/finance/order-changes?date_from=2026-09-30&date_to=2026-09-01",
        headers=app_auth_headers,
    )
    assert resp.status_code == 422


async def test_month_view_requires_the_finance_section(app_client: AsyncClient):
    finance = await _login(app_client, "finance")
    recruiter = await _login(app_client, "recruiter")

    allowed = await app_client.get("/api/finance/order-changes", headers=finance)
    denied = await app_client.get("/api/finance/order-changes", headers=recruiter)

    assert allowed.status_code == 200, allowed.text
    assert denied.status_code == 403


async def test_md_line_still_billing_after_its_end_date_is_not_a_gap(monkeypatch):
    """Linia MD pracuje po dacie końca, dopóki ma budżet — to nie brak."""
    from app.core.database import AsyncSessionLocal

    end = _far_day(2085, 2088)
    ids = await _seed(start=end - timedelta(days=60), end=end)
    async with AsyncSessionLocal() as db:
        order = await db.get(ClientOrder, ids["order_id"])
        order.md_total = Decimal("100")
        order.md_remaining = Decimal("20")
        order.md_rate_revenue = Decimal("1340")
        order.md_input_mode = "md"
        order.md_input_value = Decimal("100")
        await db.commit()
    await _detect(monkeypatch, tracking_start=end, today=end + timedelta(days=1))
    assert await _gap_for(ids["order_id"]) is None


async def _as_md_line(order_id: int, **extra) -> None:
    from app.core.database import AsyncSessionLocal

    async with AsyncSessionLocal() as db:
        order = await db.get(ClientOrder, order_id)
        order.md_total = Decimal("100")
        order.md_remaining = Decimal("20")
        order.md_rate_revenue = Decimal("1340")
        order.md_input_mode = "md"
        order.md_input_value = Decimal("100")
        for key, value in extra.items():
            setattr(order, key, value)
        await db.commit()


async def test_md_line_closed_forward_is_completed_once_its_day_passed():
    """FIN-CHG-1 (audyt 22.09): linia MD zamknięta „w przód" nie wisi ``active``.

    Zamiana kontraktora z datą w przyszłości i zamknięcie grupy z przyszłą datą
    zostawiały poprzednika ``active`` na zawsze — skaner celowo nie domyka linii
    MD po dacie. Świadomy koniec (następca, zakończona grupa) domyka ją teraz;
    linia z samą datą nadal pracuje do wyczerpania budżetu.
    """
    from app.core.database import AsyncSessionLocal
    from app.core.scheduling import business_today
    from app.models.client_order_group import ClientOrderGroup
    from app.tasks.dl_portal_expiry_scanner import _promote_statuses

    today = business_today()
    end = today - timedelta(days=3)
    swapped = await _seed(start=end - timedelta(days=60), end=end)
    await _as_md_line(swapped["order_id"])
    await _add_order(
        swapped,
        start=end + timedelta(days=1),
        end=today + timedelta(days=60),
        predecessor_order_id=swapped["order_id"],
    )
    still_billing = await _seed(start=end - timedelta(days=60), end=end)
    await _as_md_line(still_billing["order_id"])
    closed = await _seed(start=end - timedelta(days=60), end=end)
    async with AsyncSessionLocal() as db:
        group = ClientOrderGroup(
            client_id=closed["client_id"],
            order_number=f"G-{uuid.uuid4().hex[:6]}",
            order_type="md",
            status="completed",
            start_date=end - timedelta(days=60),
            end_date=end,
            closure_date=end,
        )
        db.add(group)
        await db.commit()
        group_id = group.id
    await _as_md_line(closed["order_id"], order_group_id=group_id)

    async with AsyncSessionLocal() as db:
        await _promote_statuses(db)
        await db.commit()
        statuses = {
            key: (await db.get(ClientOrder, ids["order_id"])).status
            for key, ids in (
                ("swapped", swapped),
                ("still_billing", still_billing),
                ("closed", closed),
            )
        }
    assert statuses == {
        "swapped": ClientOrderStatus.completed,
        "still_billing": ClientOrderStatus.active,
        "closed": ClientOrderStatus.completed,
    }


async def test_successor_on_a_contract_without_a_person_is_seen(monkeypatch):
    from app.core.database import AsyncSessionLocal

    end = _far_day(2085, 2088)
    ids = await _seed(start=end - timedelta(days=60), end=end)
    async with AsyncSessionLocal() as db:
        contract = await db.get(Contract, ids["contract_id"])
        contract.candidate_id = None
        await db.commit()
    await _add_order(ids, start=end + timedelta(days=1), end=end + timedelta(days=90))
    await _detect(monkeypatch, tracking_start=end, today=end + timedelta(days=1))
    assert await _gap_for(ids["order_id"]) is None


async def test_old_endings_outside_the_lookback_window_are_not_backfilled(monkeypatch):
    end = _far_day(2085, 2088)
    ids = await _seed(start=end - timedelta(days=60), end=end)
    await _detect(
        monkeypatch,
        tracking_start=end - timedelta(days=400),
        today=end + timedelta(days=90),
    )
    assert await _gap_for(ids["order_id"]) is None


# ── Zejścia nie wracają w Brakach ────────────────────────────────────────────


async def _terminate(contract_id: int, end: date) -> None:
    from app.core.database import AsyncSessionLocal

    async with AsyncSessionLocal() as db:
        contract = await db.get(Contract, contract_id)
        contract.end_date = end
        contract.terminated_at = datetime.now(timezone.utc)
        contract.termination_reason = "consultant_resigned"
        await db.commit()


def _gap_order_ids(body: dict, client_id: int) -> set[int]:
    return {g["order_id"] for g in body["gaps"] if g["client_id"] == client_id}


async def test_cooperation_ended_after_the_gap_leaves_braki_for_zejscia(
    monkeypatch, app_client: AsyncClient, app_auth_headers: dict
):
    """DL wypowiada umowę dopiero, gdy zobaczy brak — osoba idzie do Zejść."""

    first = _far_day(2040, 2043).replace(day=1)
    end = first + timedelta(days=9)
    ended = await _seed(
        contract_rate_client=None, start=end - timedelta(days=60), end=end
    )
    # Kontrola: ta sama sytuacja bez wypowiedzenia zostaje brakiem.
    kept = await _seed(
        contract_rate_client=None, start=end - timedelta(days=60), end=end
    )
    await _detect(monkeypatch, tracking_start=end, today=end + timedelta(days=1))
    assert (await _gap_for(ended["order_id"])).status == "open"

    await _terminate(ended["contract_id"], end)

    body = await _month(app_client, app_auth_headers, first)
    assert ended["order_id"] in {
        e["order_id"] for e in body["exits"] if e["verdict"] == "ended_intent"
    }
    assert _gap_order_ids(body, ended["client_id"]) == set()
    assert _gap_order_ids(body, kept["client_id"]) == {kept["order_id"]}

    # Licznik liczy to samo co lista.
    scoped = await _month(
        app_client, app_auth_headers, first, client_id=ended["client_id"]
    )
    assert scoped["counts"]["gaps"] == 0 and scoped["gaps"] == []
    assert scoped["counts"]["exits"] == 1

    export = await app_client.get(
        f"/api/finance/order-changes/export?year={first.year}&month={first.month}"
        f"&tab=gaps&client_id={ended['client_id']}",
        headers=app_auth_headers,
    )
    assert export.status_code == 200, export.text
    assert "Braki (0)" in load_workbook(io.BytesIO(export.content)).sheetnames

    # Wpis braku nie znika z bazy — znika z raportu.
    assert await _gap_for(ended["order_id"]) is not None


async def test_person_in_the_months_exits_is_hidden_even_on_another_order(
    app_client: AsyncClient, app_auth_headers: dict
):
    """Wypowiedzenie przypięte do późniejszego zamówienia tej samej osoby."""

    from app.core.database import AsyncSessionLocal
    from app.models.order_gap import GAP_STATUS_OPEN, OrderGap

    first = _far_day(2044, 2047).replace(day=1)
    early_end = first + timedelta(days=4)
    ids = await _seed(
        contract_rate_client=None, start=first - timedelta(days=60), end=early_end
    )
    late_end = first + timedelta(days=20)
    late = await _add_order(ids, start=first + timedelta(days=8), end=late_end)
    async with AsyncSessionLocal() as db:
        db.add(
            OrderGap(
                order_id=ids["order_id"],
                contract_id=ids["contract_id"],
                client_id=ids["client_id"],
                order_number="NB-early",
                ended_on=early_end,
                detected_on=early_end + timedelta(days=1),
                status=GAP_STATUS_OPEN,
            )
        )
        await db.commit()
    await _terminate(ids["contract_id"], late_end)

    body = await _month(app_client, app_auth_headers, first, client_id=ids["client_id"])
    assert [e["order_id"] for e in body["exits"]] == [late]
    assert body["gaps"] == [] and body["counts"]["gaps"] == 0


async def test_reminders_stop_and_card_closes_once_cooperation_ended(monkeypatch):
    from app.core.database import AsyncSessionLocal
    from app.models.dl_alert import (
        ALERT_ORDER_MISSING_SUCCESSOR,
        DL_ALERT_STATUS_HANDLED,
        DlAlert,
    )
    from app.services.order_gaps import remind_open_gaps

    end = _far_day(2048, 2051)
    ids = await _seed(
        contract_rate_client=None, start=end - timedelta(days=60), end=end
    )
    dl_id = await _seed_dl(ids["client_id"])
    await _detect(monkeypatch, tracking_start=end, today=end + timedelta(days=1))
    await _terminate(ids["contract_id"], end)

    async with AsyncSessionLocal() as db:
        await remind_open_gaps(db)
        await db.commit()
        alerts = (
            await db.scalars(
                select(DlAlert).where(
                    DlAlert.user_id == dl_id,
                    DlAlert.alert_type == ALERT_ORDER_MISSING_SUCCESSOR,
                )
            )
        ).all()
    assert alerts and all(a.status == DL_ALERT_STATUS_HANDLED for a in alerts)
    assert all(a.handled_by_user_id is None for a in alerts)


async def test_gap_resolved_by_a_later_run_keeps_the_successor_creation_time(
    monkeypatch,
):
    """FIN-CHG-2: moment uzupełnienia = założenie następcy, nie przebieg pętli."""
    from app.core.database import AsyncSessionLocal
    from app.models.order_gap import GAP_STATUS_FILLED_LATE
    from app.services.order_gaps import resolve_order_gaps

    end = _far_day(2085, 2088)
    ids = await _seed(start=end - timedelta(days=60), end=end)
    await _detect(monkeypatch, tracking_start=end, today=end + timedelta(days=1))
    successor_id = await _add_order(
        ids, start=end + timedelta(days=1), end=end + timedelta(days=90)
    )
    async with AsyncSessionLocal() as db:
        successor = await db.get(ClientOrder, successor_id)
        created = successor.created_at
        await resolve_order_gaps(
            db,
            contract_ids=[ids["contract_id"]],
            now=created + timedelta(days=5),
        )
        await db.commit()
    gap = await _gap_for(ids["order_id"])
    assert gap.status == GAP_STATUS_FILLED_LATE
    assert gap.resolved_at == created


async def test_deleted_order_leaves_no_gap_in_the_report_and_closes_the_card(
    monkeypatch, app_client: AsyncClient, app_auth_headers: dict
):
    """FIN-CHG-5: brak usuniętego zamówienia nie jest pokazywany ani przypominany."""
    from app.core.database import AsyncSessionLocal
    from app.models.dl_alert import (
        ALERT_ORDER_MISSING_SUCCESSOR,
        DL_ALERT_STATUS_NEW,
        DlAlert,
    )

    end = _far_day(2085, 2088)
    ids = await _seed(start=end - timedelta(days=60), end=end)
    await _seed_dl(ids["client_id"])
    await _detect(monkeypatch, tracking_start=end, today=end + timedelta(days=1))
    gap = await _gap_for(ids["order_id"])
    assert gap is not None

    resp = await app_client.delete(
        f"/api/clients/{ids['client_id']}/orders/{ids['order_id']}",
        headers=app_auth_headers,
    )
    assert resp.status_code in (200, 204), resp.text

    async with AsyncSessionLocal() as db:
        open_cards = await db.scalar(
            select(func.count())
            .select_from(DlAlert)
            .where(
                DlAlert.alert_type == ALERT_ORDER_MISSING_SUCCESSOR,
                DlAlert.client_id == ids["client_id"],
                DlAlert.status == DL_ALERT_STATUS_NEW,
            )
        )
    assert open_cards == 0
    body = await _month(
        app_client,
        app_auth_headers,
        end + timedelta(days=1),
        client_id=ids["client_id"],
    )
    assert _gap_order_ids(body, ids["client_id"]) == set()


async def test_md_line_billing_at_another_client_is_an_additional_project():
    """FIN-CHG-6: linia MD z budżetem po dacie końca to równoległa praca."""
    from app.services.finance_order_changes import _runs_on

    md = _fact(1, md_total=Decimal("100"), md_remaining=Decimal("10"))
    assert _runs_on(md, date(2026, 12, 1))
    assert not _runs_on(_fact(2), date(2026, 12, 1))


async def test_draft_and_order_of_the_same_person_are_one_entry(
    app_client: AsyncClient, app_auth_headers: dict
):
    """FIN-CHG-7: szkic + zamówienie tej samej współpracy = jedno Wejście."""
    day = _far_day(2081, 2084).replace(day=10)
    ids = await _seed(start=day, end=day + timedelta(days=90))
    await _add_order(
        ids,
        start=day + timedelta(days=5),
        end=day + timedelta(days=60),
        status=ClientOrderStatus.draft,
    )
    body = await _month(app_client, app_auth_headers, day, client_id=ids["client_id"])
    entries = [e for e in body["entries"] if e["client_id"] == ids["client_id"]]
    assert [e["order_id"] for e in entries] == [ids["order_id"]]
