"""Zamówienia MD/kosztowe i okresowe są NIEZALEŻNE — jedno nie kończy drugiego.

Zgłoszenie (BNP / Polkomtel / BIK / Lotte Wedel): u konsultantów obsadzonych na
zamówieniu rozliczanym w MD powstawały równolegle zamówienia okresowe, a
zakończenie tego drugiego kasowało budżet MD tej samej osoby. Ten plik przybija
oba końce naprawy — wejście (nie da się już zrobić duplikatu) i wyjście
(zakończenie zamówienia dotyczy dokładnie jednego wiersza).
"""

from __future__ import annotations

import uuid
from datetime import date, timedelta
from decimal import Decimal

from httpx import AsyncClient
from sqlalchemy import select

from tests.test_multi_consultant_orders import (
    _TODAY,
    _create_group,
    _enable_for,
    _line_payload,
    _seed_client_with_contracts,
)


async def _standalone_order(
    client_id: int,
    contract_id: int,
    *,
    title: str = "ZAM-OKRESOWE",
    status: str = "active",
    start_date: date | None = None,
    end_date: date | None = None,
    order_group_id: int | None = None,
) -> int:
    from app.core.database import AsyncSessionLocal
    from app.models.client_order import ClientOrder, ClientOrderStatus

    async with AsyncSessionLocal() as db:
        order = ClientOrder(
            client_id=client_id,
            contract_id=contract_id,
            title=title,
            order_type="periodic" if order_group_id is None else None,
            order_group_id=order_group_id,
            status=ClientOrderStatus(status),
            start_date=start_date if start_date is not None else _TODAY - timedelta(days=5),
            end_date=end_date,
            rate_client=Decimal("150.000"),
        )
        db.add(order)
        await db.commit()
        await db.refresh(order)
        return order.id


async def _order_row(order_id: int):
    from app.core.database import AsyncSessionLocal
    from app.models.client_order import ClientOrder

    async with AsyncSessionLocal() as db:
        return await db.get(ClientOrder, order_id)


async def _contract_row(contract_id: int):
    from app.core.database import AsyncSessionLocal
    from app.models.contract import Contract

    async with AsyncSessionLocal() as db:
        return await db.get(Contract, contract_id)


# ── Wyjście: zakończenie zamówienia nie dotyka niczego obok ─────────────────


async def test_closing_a_periodic_order_leaves_the_md_line_and_contract_alone(
    app_client: AsyncClient, app_auth_headers: dict, monkeypatch
):
    """Kryterium akceptacji ticketu, dosłownie.

    Wypowiedzenie umowy (dawna jedyna droga „Zakończ" z karty kontraktora)
    domyka WSZYSTKIE zamówienia kontraktu, bo umowa opisuje całą współpracę
    u klienta. Linia MD opisuje jednak inne, trwające zaangażowanie tej samej
    osoby — stąd osobna, wąska akcja na jednym wierszu.
    """
    from app.models.client_order import ClientOrderStatus
    from app.models.contract import ContractStatus

    client_id, contracts, _ = await _seed_client_with_contracts(1)
    _enable_for(monkeypatch, client_id)
    group = await _create_group(
        app_client, app_auth_headers, client_id, [_line_payload(contracts[0])]
    )
    md_line_id = group["lines"][0]["id"]
    periodic_id = await _standalone_order(client_id, contracts[0])

    before = await _contract_row(contracts[0])
    assert before is not None

    resp = await app_client.post(
        f"/api/clients/{client_id}/orders/{periodic_id}/close",
        json={"closure_date": _TODAY.isoformat(), "closure_reason": "koniec projektu"},
        headers=app_auth_headers,
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["status"] == "completed"

    closed = await _order_row(periodic_id)
    assert closed is not None
    assert closed.status == ClientOrderStatus.completed
    assert closed.end_date == _TODAY

    md_line = await _order_row(md_line_id)
    assert md_line is not None
    assert md_line.status == ClientOrderStatus.active
    assert md_line.order_group_id == group["id"]

    after = await _contract_row(contracts[0])
    assert after is not None
    assert after.status == before.status == ContractStatus.active
    assert after.end_date == before.end_date


async def test_future_closure_date_keeps_the_order_running(
    app_client: AsyncClient, app_auth_headers: dict
):
    """Lustro `close_order_group` i syncu terminacji: data tak, status jeszcze nie."""
    from app.models.client_order import ClientOrderStatus

    client_id, contracts, _ = await _seed_client_with_contracts(1)
    order_id = await _standalone_order(client_id, contracts[0])
    when = _TODAY + timedelta(days=30)

    resp = await app_client.post(
        f"/api/clients/{client_id}/orders/{order_id}/close",
        json={"closure_date": when.isoformat()},
        headers=app_auth_headers,
    )
    assert resp.status_code == 200, resp.text

    row = await _order_row(order_id)
    assert row is not None
    assert row.end_date == when
    assert row.status == ClientOrderStatus.active


async def test_group_line_cannot_be_closed_through_the_standalone_route(
    app_client: AsyncClient, app_auth_headers: dict, monkeypatch
):
    """Grupa ma własne zakończenie — dwie drogi do jednego wiersza rozjechałyby się."""
    client_id, contracts, _ = await _seed_client_with_contracts(1)
    _enable_for(monkeypatch, client_id)
    group = await _create_group(
        app_client, app_auth_headers, client_id, [_line_payload(contracts[0])]
    )

    resp = await app_client.post(
        f"/api/clients/{client_id}/orders/{group['lines'][0]['id']}/close",
        json={"closure_date": _TODAY.isoformat()},
        headers=app_auth_headers,
    )
    assert resp.status_code == 409, resp.text
    assert resp.json()["detail"]["code"] == "order_belongs_to_group"


# ── Wejście: duplikat nie powstanie ─────────────────────────────────────────


async def test_periodic_order_is_refused_next_to_a_live_group_line(
    app_client: AsyncClient, app_auth_headers: dict, monkeypatch
):
    """Przyczyna zgłoszenia: dwa równoległe zapisy tej samej współpracy."""
    client_id, contracts, _ = await _seed_client_with_contracts(1)
    _enable_for(monkeypatch, client_id)
    await _create_group(
        app_client, app_auth_headers, client_id, [_line_payload(contracts[0])]
    )

    resp = await app_client.post(
        f"/api/clients/{client_id}/orders",
        data={
            "contract_id": str(contracts[0]),
            "title": "ZAM_1453_2026",
            "order_type": "periodic",
            "start_date": _TODAY.isoformat(),
        },
        headers=app_auth_headers,
    )
    assert resp.status_code == 409, resp.text
    assert resp.json()["detail"]["code"] == "consultant_already_on_group_order"


async def test_adding_a_group_line_absorbs_the_auto_draft_shell(
    app_client: AsyncClient, app_auth_headers: dict, monkeypatch
):
    """Szkic „(bez numeru)" po hooku zatrudnienia znika, gdy osoba wchodzi na MD.

    To druga połowa przyczyny: hook zakłada zaślepkę PRZED obsadą, więc bramka
    przy tworzeniu zamówienia nie miała czego odrzucić.
    """
    from app.core.database import AsyncSessionLocal
    from app.models.client_order import ClientOrder

    client_id, contracts, _ = await _seed_client_with_contracts(1)
    _enable_for(monkeypatch, client_id)
    shell_id = await _standalone_order(
        client_id, contracts[0], title="(bez numeru)", status="draft"
    )
    kept_id = await _standalone_order(
        client_id, contracts[0], title="ZAM-Z-NUMEREM", status="draft"
    )

    await _create_group(
        app_client, app_auth_headers, client_id, [_line_payload(contracts[0])]
    )

    async with AsyncSessionLocal() as db:
        remaining = set(
            (
                await db.execute(
                    select(ClientOrder.id).where(
                        ClientOrder.contract_id == contracts[0],
                        ClientOrder.order_group_id.is_(None),
                    )
                )
            ).scalars()
        )
    assert shell_id not in remaining
    assert kept_id in remaining

    # Kasowanie wiersza, którego nikt nie widzi w historii, jest dla operatora
    # nieodróżnialne od danych, które zniknęły same. Ten sam kształt wpisu ma
    # migracja 0261.
    from app.models.activity import Activity

    async with AsyncSessionLocal() as db:
        audit = (
            await db.execute(
                select(Activity).where(
                    Activity.entity_type == "client_order",
                    Activity.entity_id == shell_id,
                    Activity.action == "order_deleted",
                )
            )
        ).scalars().all()
    assert len(audit) == 1
    assert audit[0].details["reason"] == "absorbed_auto_draft_shell"
    assert audit[0].details["contract_id"] == contracts[0]


# ── Kontrakt nie dziedziczy daty końca zamówienia ───────────────────────────


async def test_line_for_a_person_from_the_nexus_base_creates_an_open_ended_contract(
    app_client: AsyncClient, app_auth_headers: dict, monkeypatch
):
    """Umowa B2B jest bezterminowa, dopóki człowiek nie powie inaczej.

    Szkic kontraktu zakładany przy obsadzie dziedziczył datę końca linii, więc
    nocny `_promote_statuses` przestawiał umowę na „Kończąca się", a potem
    „Zakończona" — mimo trwającego zamówienia.
    """
    from app.core.database import AsyncSessionLocal
    from app.models.candidate import Candidate
    from app.models.contract import Contract

    client_id, contracts, _ = await _seed_client_with_contracts(1)
    _enable_for(monkeypatch, client_id)
    group = await _create_group(
        app_client, app_auth_headers, client_id, [_line_payload(contracts[0])]
    )

    suffix = uuid.uuid4().hex[:6]
    async with AsyncSessionLocal() as db:
        candidate = Candidate(
            name="Nowy",
            lastname=f"Konsultant-{suffix}",
            email=f"sep-{suffix}@example.com",
        )
        db.add(candidate)
        await db.commit()
        await db.refresh(candidate)
        candidate_id = candidate.id

    line_end = _TODAY + timedelta(days=90)
    payload = _line_payload(contracts[0])
    payload.pop("contract_id")
    payload["candidate_id"] = candidate_id
    payload["end_date"] = line_end.isoformat()
    resp = await app_client.post(
        f"/api/clients/{client_id}/order-groups/{group['id']}/lines",
        json=payload,
        headers=app_auth_headers,
    )
    assert resp.status_code == 201, resp.text

    async with AsyncSessionLocal() as db:
        contract = (
            await db.execute(
                select(Contract).where(
                    Contract.candidate_id == candidate_id,
                    Contract.client_id == client_id,
                )
            )
        ).scalar_one()
    assert contract.end_date is None
    assert contract.start_date is not None


# ── Zakończenie u jednego klienta nie sięga do drugiego ─────────────────────


async def test_contract_offboarding_never_touches_another_clients_order():
    """Rozjechany wiersz spinający dwóch klientów nie może być domykany „przy okazji".

    ``client_orders.client_id`` to WŁASNA kolumna, a baza nie ma więzu
    wiążącego ją z klientem kontraktu (rozjazd zna też ``contract_merge``).
    Bez złączenia po kliencie zakończenie współpracy u jednego klienta domykało
    zamówienie u drugiego — zgłoszony objaw „zamknęło kontrakt u innego klienta".
    """
    from app.core.database import AsyncSessionLocal
    from app.models.client_order import ClientOrderStatus
    from app.services.contract_order_offboarding import (
        apply_contract_order_offboarding,
    )

    client_a, contracts_a, _ = await _seed_client_with_contracts(1)
    client_b, _, _ = await _seed_client_with_contracts(1)

    own_order = await _standalone_order(client_a, contracts_a[0], title="U-KLIENTA-A")
    foreign_order = await _standalone_order(
        client_b, contracts_a[0], title="ROZJECHANY-DO-B"
    )

    async with AsyncSessionLocal() as db:
        await apply_contract_order_offboarding(
            db,
            contract_id=contracts_a[0],
            effective_date=_TODAY,
            actor_id=None,
        )
        await db.commit()

    closed = await _order_row(own_order)
    untouched = await _order_row(foreign_order)
    assert closed is not None and closed.status == ClientOrderStatus.completed
    assert untouched is not None
    assert untouched.status == ClientOrderStatus.active
    assert untouched.end_date is None


# ── Eksport: stan NA DZIŚ, jeden wiersz na konsultanta ──────────────────────


async def test_excel_export_takes_only_the_currently_valid_order(
    app_client: AsyncClient, app_auth_headers: dict
):
    """Ta sama osoba pojawiała się w arkuszu raz na każde zamówienie w historii."""
    import io

    from openpyxl import load_workbook

    client_id, contracts, _ = await _seed_client_with_contracts(1)
    past = await _standalone_order(
        client_id,
        contracts[0],
        title="PRZESZLE",
        status="completed",
        start_date=_TODAY - timedelta(days=400),
        end_date=_TODAY - timedelta(days=200),
    )
    current = await _standalone_order(
        client_id,
        contracts[0],
        title="BIEZACE",
        start_date=_TODAY - timedelta(days=10),
        end_date=_TODAY + timedelta(days=20),
    )
    future = await _standalone_order(
        client_id,
        contracts[0],
        title="PRZYSZLE",
        start_date=_TODAY + timedelta(days=60),
        end_date=_TODAY + timedelta(days=200),
    )

    resp = await app_client.post(
        f"/api/clients/{client_id}/orders/export",
        json={"order_ids": [past, current, future]},
        headers=app_auth_headers,
    )
    assert resp.status_code == 200, resp.text

    sheet = load_workbook(io.BytesIO(resp.content)).active
    numbers = [row[1] for row in sheet.iter_rows(min_row=2, values_only=True)]
    assert numbers == ["BIEZACE"]
