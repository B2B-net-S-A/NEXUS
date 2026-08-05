"""Tests for the per-client contract-register filter + export.

Two features land together and share the same filter helper, so they are tested
together:

1. „Okres" overlap filter on GET /api/contracts — a contract appears when its
   validity [start_date, end_date] INTERSECTS [period_from, period_to] (not just
   by start or just by end), with open-ended (NULL) bounds handled.
2. GET /api/contracts/register/export — the client-scoped XLSX export whose 7
   columns mirror the visible ClientContractRegister table (no rates/margin, so
   it is TacPlus not Admin-only).

Layers:
- Pure unit tests for the row builder + period cell — no DB.
- In-process integration (app_client) for the endpoints end-to-end.
"""

from __future__ import annotations

import uuid
from datetime import date
from io import BytesIO
from types import SimpleNamespace

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook

from app.api.contracts import (
    _ENGAGEMENT_MODEL_LABELS,
    _REGISTER_EXPORT_COLUMNS,
    _REGISTER_PROLONGATION_LABELS,
    _register_export_row,
    _register_period_cell,
)
from app.models.contract import (
    Contract,
    ContractStatus,
    ContractType,
    EngagementModel,
    ProlongationStatus,
)

XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ── Pure unit tests (no DB) ──────────────────────────────────────────────────


def _mem_contract(**kw) -> Contract:
    # SQLAlchemy stores loaded relationships in ``instance.__dict__[key]``; the
    # instrumented setter rejects a plain SimpleNamespace (no _sa_instance_state).
    # Writing __dict__ is exactly how an eager-loaded relation looks.
    candidate = kw.pop("candidate", SimpleNamespace(name="Jan", lastname="Kowalski"))
    c = Contract(candidate_id=kw.pop("candidate_id", 1), client_id=1)
    c.id = kw.pop("id", 101)
    for k, v in kw.items():
        setattr(c, k, v)
    c.__dict__["candidate"] = candidate
    return c


def test_register_export_columns_order_matches_spec():
    # Kolejność kolumn jest częścią kontraktu (warunek zamknięcia #6).
    assert _REGISTER_EXPORT_COLUMNS == [
        "Nr projektu",
        "Projekt",
        "Konsultant",
        "Model",
        "Okres / Pula godzin",
        "Prolongata",
        "Status",
    ]


def test_period_cell_time_based_shows_range_pl_dates():
    # Daty w formacie PL (lustro Intl.DateTimeFormat('pl-PL')): dzień bez zera
    # wiodącego, miesiąc z zerem — jak na ekranie (formatDate).
    c = _mem_contract(
        engagement_model=EngagementModel.time_based,
        start_date=date(2026, 1, 1),
        end_date=date(2026, 12, 31),
    )
    assert _register_period_cell(c) == "1.01.2026 → 31.12.2026"


def test_period_cell_time_based_open_ended_is_bezterminowo():
    c = _mem_contract(
        engagement_model=EngagementModel.time_based,
        start_date=date(2026, 1, 1),
        end_date=None,
    )
    assert _register_period_cell(c) == "1.01.2026 → bezterminowo"


def test_period_cell_time_based_missing_start_dash():
    c = _mem_contract(
        engagement_model=EngagementModel.time_based,
        start_date=None,
        end_date=date(2026, 6, 30),
    )
    assert _register_period_cell(c) == "— → 30.06.2026"


def test_period_cell_hours_pool_shows_consumption_and_pct():
    c = _mem_contract(
        engagement_model=EngagementModel.hours_pool,
        hours_pool_total=100,
        hours_pool_consumed=25,
    )
    assert _register_period_cell(c) == "25 / 100 h (25%)"


def test_period_cell_hours_pool_rounds_half_up_like_js_math_round():
    # 25/200 = 12.5% — JS Math.round(12.5)=13 (half-up). Python round(12.5)=12
    # (banker's) rozjeżdżałoby eksport z ekranem; komórka musi pokazać 13%.
    c = _mem_contract(
        engagement_model=EngagementModel.hours_pool,
        hours_pool_total=200,
        hours_pool_consumed=25,
    )
    assert _register_period_cell(c) == "25 / 200 h (13%)"


def test_period_cell_hours_pool_zero_total_shows_zero_pct():
    c = _mem_contract(
        engagement_model=EngagementModel.hours_pool,
        hours_pool_total=0,
        hours_pool_consumed=0,
    )
    # Ekran zawsze renderuje procent (brak/zerowa pula → 0%); eksport też.
    assert _register_period_cell(c) == "0 / 0 h (0%)"


def test_register_export_row_shape_and_labels():
    c = _mem_contract(
        project_code="PRJ-7",
        project_name="Migracja Core",
        engagement_model=EngagementModel.time_based,
        start_date=date(2026, 1, 1),
        end_date=date(2026, 12, 31),
        prolongation_status=ProlongationStatus.yes,
        status=ContractStatus.active,
    )
    row = _register_export_row(c)
    assert len(row) == len(_REGISTER_EXPORT_COLUMNS)
    by = dict(zip(_REGISTER_EXPORT_COLUMNS, row))
    assert by["Nr projektu"] == "PRJ-7"
    assert by["Projekt"] == "Migracja Core"
    assert by["Konsultant"] == "Jan Kowalski"
    assert by["Model"] == "Czasowy"
    assert by["Okres / Pula godzin"] == "1.01.2026 → 31.12.2026"
    assert by["Prolongata"] == "Tak"
    assert by["Status"] == "Aktywny"


def test_register_export_row_fallbacks_when_project_and_names_missing():
    c = _mem_contract(
        id=555,
        candidate=None,
        candidate_id=999,
        project_code=None,
        project_name=None,
        engagement_model=EngagementModel.time_based,
        start_date=None,
        end_date=None,
        prolongation_status=ProlongationStatus.unknown,
        status=ContractStatus.draft,
    )
    by = dict(zip(_REGISTER_EXPORT_COLUMNS, _register_export_row(c)))
    assert by["Nr projektu"] == "#555"  # brak project_code → #id (jak w UI)
    assert by["Projekt"] == ""
    assert by["Konsultant"] == "#999"
    assert by["Prolongata"] == "Nieznany"
    assert by["Status"] == "Szkic"


def test_engagement_and_prolongation_label_maps_are_complete():
    for m in EngagementModel:
        assert m.value in _ENGAGEMENT_MODEL_LABELS
    for p in ProlongationStatus:
        assert p.value in _REGISTER_PROLONGATION_LABELS


# ── Integration helpers ──────────────────────────────────────────────────────


async def _seed_client() -> int:
    from app.core.database import AsyncSessionLocal
    from app.models.client import Client

    async with AsyncSessionLocal() as db:
        cl = Client(name=f"RegExpCo-{uuid.uuid4().hex[:6]}")
        db.add(cl)
        await db.commit()
        await db.refresh(cl)
        return cl.id


async def _seed_candidate(marker: str) -> int:
    from app.core.database import AsyncSessionLocal
    from app.models.candidate import Candidate

    async with AsyncSessionLocal() as db:
        c = Candidate(
            name="RegExp",
            lastname=f"X-{marker}",
            email=f"regexp-{uuid.uuid4().hex[:8]}@example.com",
        )
        db.add(c)
        await db.commit()
        await db.refresh(c)
        return c.id


async def _seed_contract(
    *,
    client_id: int,
    marker: str,
    status: str = "active",
    start_date: date | None = None,
    end_date: date | None = None,
    engagement_model: str = "time_based",
    hours_pool_total: int | None = None,
    hours_pool_consumed: int | None = None,
    project_code: str | None = None,
    project_name: str | None = None,
    prolongation_status: str = "unknown",
    subcategory: str | None = None,
) -> tuple[int, int, int]:
    """Seed one contract under `client_id`. Returns (contract, cand, client).

    When ``subcategory`` is given, also seeds a linked Job carrying that
    ``Job.subcategory`` and points the contract's ``job_id`` at it (so the
    register subcategory filter has something to match); otherwise no job.
    """
    from app.core.database import AsyncSessionLocal
    from app.models.job import Job

    cand_id = await _seed_candidate(marker)
    async with AsyncSessionLocal() as db:
        job_id = None
        if subcategory is not None:
            j = Job(
                title=f"RegExpJob-{marker}",
                client_id=client_id,
                subcategory=subcategory,
            )
            db.add(j)
            await db.commit()
            await db.refresh(j)
            job_id = j.id
        c = Contract(
            candidate_id=cand_id,
            client_id=client_id,
            job_id=job_id,
            status=ContractStatus(status),
            contract_type=ContractType.b2b,
            start_date=start_date,
            end_date=end_date,
            engagement_model=EngagementModel(engagement_model),
            hours_pool_total=hours_pool_total,
            hours_pool_consumed=hours_pool_consumed,
            project_code=project_code,
            project_name=project_name,
            prolongation_status=ProlongationStatus(prolongation_status),
        )
        db.add(c)
        await db.commit()
        await db.refresh(c)
        return c.id, cand_id, client_id


async def _cleanup(rows: list[tuple[int, int, int]], client_ids: set[int]) -> None:
    from app.core.database import AsyncSessionLocal
    from app.models.candidate import Candidate
    from app.models.client import Client
    from app.models.job import Job
    from sqlalchemy import delete

    async with AsyncSessionLocal() as db:
        for contract_id, _, _ in rows:
            await db.execute(delete(Contract).where(Contract.id == contract_id))
        for _, cand_id, _ in rows:
            await db.execute(delete(Candidate).where(Candidate.id == cand_id))
        for cid in client_ids:
            # Jobs seeded for the subcategory filter (FK Job.client_id) go before
            # the client row.
            await db.execute(delete(Job).where(Job.client_id == cid))
            await db.execute(delete(Client).where(Client.id == cid))
        await db.commit()


def _rows_by_project(ws) -> dict[str, dict]:
    """Map each data row to {column: value} keyed by the (unique) Nr projektu."""
    header = [c.value for c in ws[1]]
    out: dict[str, dict] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        by = dict(zip(header, row))
        out[by["Nr projektu"]] = by
    return out


# ── „Okres" overlap on the list endpoint ─────────────────────────────────────


@pytest.mark.asyncio
async def test_period_overlap_includes_intersecting_excludes_disjoint(
    app_client: AsyncClient, app_auth_headers: dict
):
    marker = uuid.uuid4().hex[:8]
    client_id = await _seed_client()
    # Window [2026-06-01, 2026-06-30].
    a = await _seed_contract(  # overlaps: end inside window
        client_id=client_id,
        marker=marker,
        start_date=date(2026, 5, 1),
        end_date=date(2026, 6, 15),
    )
    b = await _seed_contract(  # disjoint: starts after window
        client_id=client_id,
        marker=marker,
        start_date=date(2026, 7, 1),
        end_date=date(2026, 8, 1),
    )
    d = await _seed_contract(  # disjoint: ends before window
        client_id=client_id,
        marker=marker,
        start_date=date(2026, 1, 1),
        end_date=date(2026, 3, 1),
    )
    e = await _seed_contract(  # open-ended: end NULL, starts inside window
        client_id=client_id,
        marker=marker,
        start_date=date(2026, 6, 10),
        end_date=None,
    )
    f = await _seed_contract(  # open start: start NULL, ends inside window
        client_id=client_id,
        marker=marker,
        start_date=None,
        end_date=date(2026, 6, 20),
    )
    rows = [a, b, d, e, f]
    try:
        r = await app_client.get(
            "/api/contracts",
            params={
                "q": marker,
                "period_from": "2026-06-01",
                "period_to": "2026-06-30",
                "page_size": 100,
            },
            headers=app_auth_headers,
        )
        assert r.status_code == 200, r.text
        ids = {it["id"] for it in r.json()["items"]}
        assert ids == {a[0], e[0], f[0]}
    finally:
        await _cleanup(rows, {client_id})


@pytest.mark.asyncio
async def test_period_from_only_keeps_open_ended_and_future(
    app_client: AsyncClient, app_auth_headers: dict
):
    marker = uuid.uuid4().hex[:8]
    client_id = await _seed_client()
    past = await _seed_contract(  # ends before period_from → excluded
        client_id=client_id,
        marker=marker,
        start_date=date(2026, 1, 1),
        end_date=date(2026, 3, 1),
    )
    future = await _seed_contract(  # ends after → included
        client_id=client_id,
        marker=marker,
        start_date=date(2026, 7, 1),
        end_date=date(2026, 8, 1),
    )
    open_ended = await _seed_contract(  # end NULL → always included
        client_id=client_id,
        marker=marker,
        start_date=date(2025, 1, 1),
        end_date=None,
    )
    rows = [past, future, open_ended]
    try:
        r = await app_client.get(
            "/api/contracts",
            params={"q": marker, "period_from": "2026-06-01", "page_size": 100},
            headers=app_auth_headers,
        )
        assert r.status_code == 200, r.text
        ids = {it["id"] for it in r.json()["items"]}
        assert ids == {future[0], open_ended[0]}
    finally:
        await _cleanup(rows, {client_id})


# ── Register export endpoint ─────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_register_export_returns_xlsx_with_seven_columns(
    app_client: AsyncClient, app_auth_headers: dict
):
    marker = uuid.uuid4().hex[:8]
    client_id = await _seed_client()
    tb = await _seed_contract(
        client_id=client_id,
        marker=marker,
        status="active",
        start_date=date(2026, 1, 1),
        end_date=date(2026, 12, 31),
        engagement_model="time_based",
        project_code="ZAM-11",
        project_name="Core Banking",
        prolongation_status="yes",
    )
    hp = await _seed_contract(
        client_id=client_id,
        marker=marker,
        status="active",
        engagement_model="hours_pool",
        hours_pool_total=200,
        hours_pool_consumed=50,
        project_code="ZAM-12",
        project_name="Wsparcie ad-hoc",
        prolongation_status="negotiate",
    )
    rows = [tb, hp]
    try:
        r = await app_client.get(
            "/api/contracts/register/export",
            params={"client_id": client_id},
            headers=app_auth_headers,
        )
        assert r.status_code == 200, r.text
        assert r.headers["content-type"].startswith(XLSX_MEDIA)
        assert ".xlsx" in r.headers.get("content-disposition", "")

        ws = load_workbook(BytesIO(r.content)).active
        assert [c.value for c in ws[1]] == _REGISTER_EXPORT_COLUMNS
        by_project = _rows_by_project(ws)
        # Scoped to a fresh client → exactly our two rows.
        assert set(by_project) == {"ZAM-11", "ZAM-12"}

        tb_row = by_project["ZAM-11"]
        assert tb_row["Projekt"] == "Core Banking"
        assert tb_row["Konsultant"].startswith("RegExp")
        assert tb_row["Model"] == "Czasowy"
        assert tb_row["Okres / Pula godzin"] == "1.01.2026 → 31.12.2026"
        assert tb_row["Prolongata"] == "Tak"
        assert tb_row["Status"] == "Aktywny"

        hp_row = by_project["ZAM-12"]
        assert hp_row["Model"] == "Pula godzin"
        assert hp_row["Okres / Pula godzin"] == "50 / 200 h (25%)"
        assert hp_row["Prolongata"] == "Negocjacje"
    finally:
        await _cleanup(rows, {client_id})


@pytest.mark.asyncio
async def test_register_export_requires_client_id(
    app_client: AsyncClient, app_auth_headers: dict
):
    # Bez client_id eksport nie ma sensu → 422 (warunek zamknięcia #7).
    r = await app_client.get("/api/contracts/register/export", headers=app_auth_headers)
    assert r.status_code == 422, r.text


@pytest.mark.asyncio
async def test_register_export_honours_status_filter(
    app_client: AsyncClient, app_auth_headers: dict
):
    marker = uuid.uuid4().hex[:8]
    client_id = await _seed_client()
    active = await _seed_contract(
        client_id=client_id,
        marker=marker,
        status="active",
        project_code="A-1",
    )
    draft = await _seed_contract(
        client_id=client_id,
        marker=marker,
        status="draft",
        project_code="D-1",
    )
    rows = [active, draft]
    try:
        r = await app_client.get(
            "/api/contracts/register/export",
            params={"client_id": client_id, "status": ["active"]},
            headers=app_auth_headers,
        )
        assert r.status_code == 200, r.text
        ws = load_workbook(BytesIO(r.content)).active
        codes = {row[0] for row in ws.iter_rows(min_row=2, values_only=True) if row}
        assert codes == {"A-1"}  # draft excluded by the filter
    finally:
        await _cleanup(rows, {client_id})


@pytest.mark.asyncio
async def test_register_export_honours_period_overlap(
    app_client: AsyncClient, app_auth_headers: dict
):
    marker = uuid.uuid4().hex[:8]
    client_id = await _seed_client()
    inside = await _seed_contract(
        client_id=client_id,
        marker=marker,
        project_code="IN-1",
        start_date=date(2026, 6, 5),
        end_date=date(2026, 6, 25),
    )
    outside = await _seed_contract(
        client_id=client_id,
        marker=marker,
        project_code="OUT-1",
        start_date=date(2026, 9, 1),
        end_date=date(2026, 9, 30),
    )
    rows = [inside, outside]
    try:
        r = await app_client.get(
            "/api/contracts/register/export",
            params={
                "client_id": client_id,
                "period_from": "2026-06-01",
                "period_to": "2026-06-30",
            },
            headers=app_auth_headers,
        )
        assert r.status_code == 200, r.text
        ws = load_workbook(BytesIO(r.content)).active
        codes = {row[0] for row in ws.iter_rows(min_row=2, values_only=True) if row}
        assert codes == {"IN-1"}  # outside the window is excluded
    finally:
        await _cleanup(rows, {client_id})


@pytest.mark.asyncio
async def test_register_export_scoped_to_selected_client_only(
    app_client: AsyncClient, app_auth_headers: dict
):
    marker = uuid.uuid4().hex[:8]
    client_a = await _seed_client()
    client_b = await _seed_client()
    a = await _seed_contract(client_id=client_a, marker=marker, project_code="CA-1")
    b = await _seed_contract(client_id=client_b, marker=marker, project_code="CB-1")
    rows = [a, b]
    try:
        r = await app_client.get(
            "/api/contracts/register/export",
            params={"client_id": client_a},
            headers=app_auth_headers,
        )
        assert r.status_code == 200, r.text
        ws = load_workbook(BytesIO(r.content)).active
        codes = {row[0] for row in ws.iter_rows(min_row=2, values_only=True) if row}
        # Only the selected client's row — never client B's.
        assert codes == {"CA-1"}
    finally:
        await _cleanup(rows, {client_a, client_b})


# ── Subcategory (Job.subcategory) filter ─────────────────────────────────────


@pytest.mark.asyncio
async def test_subcategory_filter_on_list_matches_linked_job(
    app_client: AsyncClient, app_auth_headers: dict
):
    marker = uuid.uuid4().hex[:8]
    client_id = await _seed_client()
    be = await _seed_contract(client_id=client_id, marker=marker, subcategory="Backend")
    fe = await _seed_contract(
        client_id=client_id, marker=marker, subcategory="Frontend"
    )
    nojob = await _seed_contract(client_id=client_id, marker=marker)  # job_id NULL
    rows = [be, fe, nojob]
    try:
        r = await app_client.get(
            "/api/contracts",
            params={"q": marker, "subcategory": ["Backend"], "page_size": 100},
            headers=app_auth_headers,
        )
        assert r.status_code == 200, r.text
        ids = {it["id"] for it in r.json()["items"]}
        # Only the Backend-job contract; Frontend and the job-less one excluded.
        assert ids == {be[0]}
    finally:
        await _cleanup(rows, {client_id})


@pytest.mark.asyncio
async def test_register_export_honours_subcategory(
    app_client: AsyncClient, app_auth_headers: dict
):
    marker = uuid.uuid4().hex[:8]
    client_id = await _seed_client()
    be = await _seed_contract(
        client_id=client_id, marker=marker, project_code="BE-1", subcategory="Backend"
    )
    fe = await _seed_contract(
        client_id=client_id, marker=marker, project_code="FE-1", subcategory="Frontend"
    )
    rows = [be, fe]
    try:
        r = await app_client.get(
            "/api/contracts/register/export",
            params={"client_id": client_id, "subcategory": ["Backend"]},
            headers=app_auth_headers,
        )
        assert r.status_code == 200, r.text
        ws = load_workbook(BytesIO(r.content)).active
        codes = {row[0] for row in ws.iter_rows(min_row=2, values_only=True) if row}
        assert codes == {"BE-1"}  # Frontend excluded by the subcategory filter
    finally:
        await _cleanup(rows, {client_id})


@pytest.mark.asyncio
async def test_register_subcategories_endpoint_returns_client_distinct(
    app_client: AsyncClient, app_auth_headers: dict
):
    marker = uuid.uuid4().hex[:8]
    client_id = await _seed_client()
    other = await _seed_client()
    a = await _seed_contract(client_id=client_id, marker=marker, subcategory="Backend")
    # Duplicate subcategory + a void row (must be ignored) + a job-less row.
    b = await _seed_contract(client_id=client_id, marker=marker, subcategory="Backend")
    c = await _seed_contract(
        client_id=client_id, marker=marker, subcategory="DevOps", status="void"
    )
    d = await _seed_contract(client_id=client_id, marker=marker)  # no job
    # A different client's subcategory must not leak in.
    e = await _seed_contract(client_id=other, marker=marker, subcategory="Frontend")
    rows = [a, b, c, d, e]
    try:
        r = await app_client.get(
            "/api/contracts/register/subcategories",
            params={"client_id": client_id},
            headers=app_auth_headers,
        )
        assert r.status_code == 200, r.text
        # Distinct, void-excluded, client-scoped: only "Backend" (DevOps is void,
        # Frontend belongs to the other client, the job-less row contributes none).
        assert r.json()["subcategories"] == ["Backend"]
    finally:
        await _cleanup(rows, {client_id, other})


@pytest.mark.asyncio
async def test_register_subcategories_requires_client_id(
    app_client: AsyncClient, app_auth_headers: dict
):
    r = await app_client.get(
        "/api/contracts/register/subcategories", headers=app_auth_headers
    )
    assert r.status_code == 422, r.text
