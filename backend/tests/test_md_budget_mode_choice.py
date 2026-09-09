"""New orders choose their budget scope; historical orders are unchanged."""

from datetime import timedelta
from decimal import Decimal
from types import SimpleNamespace

import pytest

from app.core.scheduling import business_today
from app.schemas.client_order_group import OrderGroupCreate
from app.services.shared_md_orders import uses_shared_md_pool
from tests.test_explicit_order_types import _seed_client_with_contracts, _md_line
from tests.test_order_group_export_granularity import _group, _line
from app.services.order_excel_export import export_rows_for_group


def test_new_explicit_shared_pool_is_available_for_any_client():
    for client_id in (1, 12, 155, 38339):
        assert uses_shared_md_pool(
            SimpleNamespace(
                client_id=client_id, is_md_budget_based=True, md_budget_mode="shared"
            )
        )
        assert not uses_shared_md_pool(
            SimpleNamespace(
                client_id=client_id,
                is_md_budget_based=False,
                md_budget_mode="per_person",
            )
        )
    assert not uses_shared_md_pool(
        SimpleNamespace(client_id=1, is_md_budget_based=True)
    )


def test_draft_schema_requires_explicit_md_scope():
    from pydantic import ValidationError

    with pytest.raises(ValidationError):
        OrderGroupCreate(
            order_number="draft",
            start_date=business_today(),
            status="draft",
            order_type="cost",
            budget_amount=100,
        )
    draft = OrderGroupCreate(
        order_number="draft",
        start_date=business_today(),
        status="draft",
        order_type="md",
        md_budget_mode="shared",
        md_budget_total=100,
    )
    assert draft.is_md_budget_based
    for total in (None, 0):
        with pytest.raises(ValidationError):
            OrderGroupCreate(
                order_number="invalid-pool",
                start_date=business_today(),
                status="draft",
                order_type="md",
                md_budget_mode="shared",
                md_budget_total=total,
            )


def test_generic_shared_export_aggregates_once():
    group = _group(
        md_budget_mode="shared",
        is_md_budget_based=True,
        md_budget_total=100,
        md_budget_used=25,
        md_budget_remaining=75,
        lines=[_line("One"), _line("Two")],
    )
    rows = export_rows_for_group(group)
    assert [r.allocation for r in rows] == [100, None, None]
    assert [r.consumption for r in rows] == [25, None, None]
    assert [r.remaining_md for r in rows] == [75, None, None]
    import io
    from openpyxl import load_workbook
    from app.services.order_excel_export import build_orders_workbook

    sheet = load_workbook(
        io.BytesIO(build_orders_workbook(rows, include_model_columns=True))
    ).active
    assert sheet.cell(1, 8).value == "Pozostały budżet MD"
    assert sheet.cell(2, 8).value == 75


@pytest.mark.asyncio
async def test_draft_mode_change_activation_consumption_and_lock(
    app_client, app_auth_headers
):
    client_id, contracts = await _seed_client_with_contracts()
    url = f"/api/clients/{client_id}/order-groups"
    created = await app_client.post(
        url,
        headers=app_auth_headers,
        json={
            "order_number": "MD-MODE-CHOICE",
            "start_date": (business_today() - timedelta(days=5)).isoformat(),
            "order_type": "md",
            "md_budget_mode": "per_person",
            "status": "draft",
            "lines": [_md_line(contracts[0], 40), _md_line(contracts[1], 60)],
        },
    )
    assert created.status_code == 201, created.text
    group = created.json()
    assert group["status"] == "draft" and not group["md_budget_mode_locked"]
    assert [line["md_total"] for line in group["lines"]] == [40, 60]
    url += f"/{group['id']}"
    changed = await app_client.patch(
        url,
        headers=app_auth_headers,
        json={"md_budget_mode": "shared", "md_budget_total": 100},
    )
    assert changed.status_code == 200, changed.text
    assert changed.json()["md_budget_total"] == 100
    assert all(line["md_total"] is None for line in changed.json()["lines"])
    activated = await app_client.patch(
        url, headers=app_auth_headers, json={"status": "active"}
    )
    assert activated.status_code == 200, activated.text
    assert activated.json()["status"] == "active"
    assert activated.json()["md_budget_mode_locked"]
    for used in (25, 25, 30):
        consumed = await app_client.patch(
            url,
            headers=app_auth_headers,
            json={
                "md_consumption_month": business_today().strftime("%Y-%m"),
                "md_consumption_value": used,
            },
        )
        assert consumed.status_code == 200, consumed.text
        assert consumed.json()["md_budget_remaining"] == 100 - used
    rejected = await app_client.patch(
        url, headers=app_auth_headers, json={"md_budget_mode": "per_person"}
    )
    assert rejected.status_code == 409, rejected.text
    rejected = await app_client.patch(
        url, headers=app_auth_headers, json={"status": "draft"}
    )
    assert rejected.status_code == 409, rejected.text


@pytest.mark.asyncio
async def test_new_line_currency_round_trip_without_silent_pln(
    app_client, app_auth_headers, monkeypatch
):
    from app.api import client_order_groups

    async def fx(_db, currencies, _date):
        return {
            currency: Decimal("4") if currency == "EUR" else Decimal("1")
            for currency in currencies
        }

    monkeypatch.setattr(client_order_groups, "rates_to_pln", fx)
    client_id, contracts = await _seed_client_with_contracts(1)
    url = f"/api/clients/{client_id}/order-groups"
    created = await app_client.post(
        url,
        headers=app_auth_headers,
        json={
            "order_number": "EUR-ROUND-TRIP",
            "start_date": (business_today() - timedelta(days=5)).isoformat(),
            "order_type": "md",
            "md_budget_mode": "per_person",
            "lines": [
                {
                    **_md_line(contracts[0]),
                    "rate_cost": 176,
                    "rate_revenue": 218.75,
                    "rate_candidate_currency": "EUR",
                    "rate_client_currency": "EUR",
                    "input_mode": "amount",
                    "input_value": 35000,
                }
            ],
        },
    )
    assert created.status_code == 201, created.text
    group = created.json()
    line = group["lines"][0]
    assert line["md_total"] == 40  # PLN amount / canonical 875 PLN per MD
    assert line["source_rate_cost"] == 176
    assert line["rate_cost"] == 704
    assert line["rate_candidate_currency"] == line["rate_client_currency"] == "EUR"
    patched = await app_client.patch(
        f"{url}/{group['id']}/lines/{line['id']}",
        headers=app_auth_headers,
        json={
            "rate_cost": 176,
            "rate_revenue": 218.75,
            "rate_candidate_currency": "EUR",
            "rate_client_currency": "EUR",
        },
    )
    assert patched.status_code == 200, patched.text
    assert patched.json()["source_rate_revenue"] == 218.75
    assert patched.json()["rate_client_currency"] == "EUR"
    date_only = await app_client.patch(
        f"{url}/{group['id']}/lines/{line['id']}",
        headers=app_auth_headers,
        json={"end_date": None},
    )
    assert date_only.status_code == 200, date_only.text
    assert date_only.json()["rate_client_currency"] == "EUR"
    legacy = await app_client.patch(
        f"{url}/{group['id']}/lines/{line['id']}",
        headers=app_auth_headers,
        json={"rate_cost": 704, "rate_revenue": 875},
    )
    assert legacy.status_code == 200, legacy.text
    assert legacy.json()["source_rate_cost"] == 176
    assert legacy.json()["source_rate_revenue"] == 218.75
    assert legacy.json()["rate_client_currency"] == "EUR"


@pytest.mark.asyncio
async def test_generic_shared_import_sums_two_consultants_and_replay_is_idempotent(
    app_client, app_auth_headers
):
    from tests.test_order_lifecycle_and_cost import (
        _seed_client_with_contracts as seed_with_names,
        _cost_line,
        _sheet,
        _finance_headers,
        _import_sheet,
    )

    client_id, contracts, names = await seed_with_names(2)
    number = str(9870000000 + client_id)
    response = await app_client.post(
        f"/api/clients/{client_id}/order-groups",
        headers=app_auth_headers,
        json={
            "order_number": number,
            "start_date": (business_today() - timedelta(days=10)).isoformat(),
            "order_type": "md",
            "md_budget_mode": "shared",
            "md_budget_total": 100,
            "lines": [_cost_line(contract) for contract in contracts],
        },
    )
    assert response.status_code == 201, response.text
    group_id = response.json()["id"]
    finance = await _finance_headers(app_client)
    payload = _sheet([(names[0], 12.5, number, 0), (names[1], 17.5, number, 0)])
    period = business_today().strftime("%Y-%m")
    for _ in range(2):
        imported = await _import_sheet(app_client, finance, payload, period=period)
        assert imported["rows_applied"] == 2
        groups = await app_client.get(
            f"/api/clients/{client_id}/order-groups", headers=app_auth_headers
        )
        group = next(g for g in groups.json()["groups"] if g["id"] == group_id)
        assert group["md_budget_used"] == 30
        assert group["md_budget_remaining"] == 70
        assert all(line["md_total"] is None for line in group["lines"])


@pytest.mark.asyncio
async def test_first_zero_consumption_permanently_locks_draft_mode(
    app_client, app_auth_headers
):
    from app.core.database import AsyncSessionLocal
    from app.models.client_order import ClientOrder
    from app.services.client_order_lines import upsert_consumption

    client_id, contracts = await _seed_client_with_contracts(1)
    response = await app_client.post(
        f"/api/clients/{client_id}/order-groups",
        headers=app_auth_headers,
        json={
            "order_number": f"LOCK-{client_id}",
            "start_date": (business_today() - timedelta(days=5)).isoformat(),
            "order_type": "md",
            "md_budget_mode": "per_person",
            "status": "draft",
            "lines": [_md_line(contracts[0])],
        },
    )
    assert response.status_code == 201, response.text
    group = response.json()
    async with AsyncSessionLocal() as db:
        order = await db.get(ClientOrder, group["lines"][0]["id"])
        await upsert_consumption(
            db,
            order=order,
            period_month=business_today().strftime("%Y-%m"),
            md_reported=Decimal("0"),
        )
        await db.commit()
    response = await app_client.patch(
        f"/api/clients/{client_id}/order-groups/{group['id']}",
        headers=app_auth_headers,
        json={"md_budget_mode": "shared", "md_budget_total": 100},
    )
    assert response.status_code == 409, response.text


def test_startup_safety_net_preserves_new_draft_status_and_scope():
    from pathlib import Path

    entrypoint = (Path(__file__).parents[1] / "entrypoint.sh").read_text()
    assert "md_budget_mode VARCHAR(16) NULL" in entrypoint
    assert "md_budget_mode_locked BOOLEAN NOT NULL DEFAULT FALSE" in entrypoint
    assert (
        "status IN ('active', 'scheduled', 'completed', 'exhausted')" not in entrypoint
    )
    assert (
        "status IN ('draft', 'active', 'scheduled', 'completed', 'exhausted')"
        in entrypoint
    )
