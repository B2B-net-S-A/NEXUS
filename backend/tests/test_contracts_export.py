"""Tests for /api/contracts/export (CSV / XLSX).

Two layers:
1. Pure unit tests for the row builder + cell helpers — no DB, exercise the
   column shape, Polish enum labels, effective-rate resolution and numeric/date
   coercion.
2. In-process integration tests (app_client) that hit the endpoint end-to-end:
   default XLSX, CSV with a UTF-8 BOM, filter parity with the list endpoint, and
   that client / rates / order-dates land in the sheet.
"""

from __future__ import annotations

import uuid
from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook

from app.api.contracts import (
    _CONTRACT_EXPORT_COLUMNS,
    _CONTRACT_STATUS_LABELS,
    _RATE_UNIT_LABELS,
    _contract_export_row,
    _date_cell,
    _enum_label,
    _num_cell,
)
from app.models.contract import (
    Contract,
    ContractStatus,
    ContractType,
    ContractWorkMode,
    ProlongationStatus,
    RateUnit,
)

XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ── Pure unit tests (no DB) ──────────────────────────────────────────────────


def test_num_cell_coerces_decimal_and_none():
    assert _num_cell(None) == ""
    assert _num_cell(Decimal("164.375")) == 164.375
    assert _num_cell(50) == 50.0


def test_date_cell_isoformats_dates_and_datetimes():
    assert _date_cell(None) == ""
    assert _date_cell(date(2026, 6, 30)) == "2026-06-30"
    assert _date_cell(datetime(2026, 1, 2, 12, 0, 0)) == "2026-01-02T12:00:00"


def test_enum_label_maps_known_falls_back_and_handles_none():
    assert _enum_label(None, _CONTRACT_STATUS_LABELS) == ""
    assert _enum_label(ContractStatus.active, _CONTRACT_STATUS_LABELS) == "Aktywny"
    assert _enum_label(RateUnit.hourly, _RATE_UNIT_LABELS) == "godzinowa"
    # Unknown raw value falls back to itself rather than raising.
    assert _enum_label(SimpleNamespace(value="mystery"), {}) == "mystery"


def _stub_relations(c: Contract, **relations) -> None:
    """Inject eager-loaded relationship values without a DB session.

    SQLAlchemy stores relationship values in ``instance.__dict__[key]``, so
    writing there is exactly how a loaded relation looks — and it sidesteps the
    instrumented setter, which rejects a plain ``SimpleNamespace`` (no
    ``_sa_instance_state``). The row builder only reads ``.name/.lastname/.title``
    and iterates the schedules, so lightweight stand-ins are enough.
    """
    c.__dict__.update(relations)


def _in_memory_contract() -> Contract:
    """Build a fully populated Contract without a DB session (hourly rate)."""
    c = Contract(
        candidate_id=1,
        client_id=1,
        status=ContractStatus.active,
        contract_type=ContractType.b2b,
        start_date=date(2026, 1, 1),
        end_date=date(2026, 12, 31),
        client_order_end_date=date(2026, 6, 30),
        rate_candidate=Decimal("150"),
        rate_client=Decimal("200"),
        margin=Decimal("50"),
        rate_unit=RateUnit.hourly,
        billing_hours_per_month=160,
        currency="PLN",
        framework_rate=Decimal("210.50"),
        project_code="PRJ-1",
        project_name="Nordea Core",
        client_pm_name="Anna PM",
        line_manager="Bob LM",
        work_mode=ContractWorkMode.remote,
        prolongation_status=ProlongationStatus.yes,
    )
    _stub_relations(
        c,
        candidate=SimpleNamespace(name="Jan", lastname="Kowalski"),
        client=SimpleNamespace(name="Nordea"),
        job=SimpleNamespace(title="Senior Dev"),
        candidate_rate_schedule=[],
        client_rate_schedule=[],
    )
    c.created_at = datetime(2026, 1, 2, 12, 0, 0)
    return c


def test_contract_export_row_shape_and_values():
    c = _in_memory_contract()
    row = _contract_export_row(c, date(2026, 7, 15), date(2026, 3, 1))
    # One value per declared column.
    assert len(row) == len(_CONTRACT_EXPORT_COLUMNS)
    by = dict(zip(_CONTRACT_EXPORT_COLUMNS, row))
    assert by["Kandydat"] == "Jan Kowalski"
    assert by["Klient"] == "Nordea"
    assert by["Stanowisko / Oferta"] == "Senior Dev"
    assert by["Typ"] == "B2B"
    assert by["Status"] == "Aktywny"
    assert by["Data rozpoczęcia"] == "2026-01-01"
    assert by["Data zakończenia"] == "2026-12-31"
    assert by["Koniec zamówienia u klienta"] == "2026-06-30"
    assert by["Najnowsze zamówienie do"] == "2026-07-15"
    assert by["Stawka kandydata"] == 150.0
    assert by["Stawka klienta"] == 200.0
    assert by["Marża"] == 50.0
    assert by["Jednostka stawki"] == "godzinowa"
    assert by["Waluta"] == "PLN"
    # hourly × 160 billing hours → monthly equivalents.
    assert by["Stawka mies. kandydata"] == 150.0 * 160
    assert by["Stawka mies. klienta"] == 200.0 * 160
    assert by["Marża mies."] == 50.0 * 160
    assert by["Stawka ramowa (MSA)"] == 210.5
    assert by["Numer projektu/zamówienia"] == "PRJ-1"
    assert by["Tryb pracy"] == "zdalnie"
    assert by["Status przedłużenia"] == "tak"


def test_contract_export_row_effective_rate_uses_current_schedule_step():
    """A past rate step overrides the legacy column; a future step does not."""
    c = _in_memory_contract()
    c.rate_candidate = Decimal("150")  # legacy baseline
    _stub_relations(
        c,
        candidate_rate_schedule=[
            SimpleNamespace(rate=Decimal("150"), effective_from=date(2026, 1, 1)),
            SimpleNamespace(rate=Decimal("165"), effective_from=date(2026, 3, 1)),
            SimpleNamespace(rate=Decimal("180"), effective_from=date(2027, 1, 1)),
        ],
    )
    row = _contract_export_row(c, None, date(2026, 6, 1))
    by = dict(zip(_CONTRACT_EXPORT_COLUMNS, row))
    # On 2026-06-01 the 165 step is in effect (180 is future-dated).
    assert by["Stawka kandydata"] == 165.0


# ── Integration tests (in-process app_client) ────────────────────────────────


async def _seed_client_minimal() -> int:
    from app.core.database import AsyncSessionLocal
    from app.models.client import Client

    async with AsyncSessionLocal() as db:
        cl = Client(name=f"ExportCo-{uuid.uuid4().hex[:6]}")
        db.add(cl)
        await db.commit()
        await db.refresh(cl)
        return cl.id


async def _seed_candidate_minimal() -> int:
    from app.core.database import AsyncSessionLocal
    from app.models.candidate import Candidate

    async with AsyncSessionLocal() as db:
        c = Candidate(
            name="ExportCand",
            lastname=f"X-{uuid.uuid4().hex[:6]}",
            email=f"exp-{uuid.uuid4().hex[:8]}@example.com",
        )
        db.add(c)
        await db.commit()
        await db.refresh(c)
        return c.id


async def _seed_contract(
    *,
    status: str = "active",
    with_order: bool = False,
) -> tuple[int, int, int]:
    """Seed a contract (+ optional ClientOrder). Returns (contract, cand, client)."""
    from app.core.database import AsyncSessionLocal
    from app.models.client_order import ClientOrder
    from app.models.contract import Contract as C

    cand_id = await _seed_candidate_minimal()
    client_id = await _seed_client_minimal()
    order_end = date.today() + timedelta(days=90)

    async with AsyncSessionLocal() as db:
        c = C(
            candidate_id=cand_id,
            client_id=client_id,
            status=ContractStatus(status),
            contract_type=ContractType.b2b,
            start_date=date.today() - timedelta(days=30),
            end_date=date.today() + timedelta(days=60),
            client_order_end_date=date.today() + timedelta(days=45),
            rate_client=Decimal("10000"),
            rate_candidate=Decimal("8000"),
            margin=Decimal("2000"),
            currency="PLN",
            project_code="ZAM-77",
        )
        db.add(c)
        await db.commit()
        await db.refresh(c)
        if with_order:
            db.add(
                ClientOrder(
                    client_id=client_id,
                    contract_id=c.id,
                    title="PO-2026",
                    end_date=order_end,
                )
            )
            await db.commit()
        return c.id, cand_id, client_id


async def _cleanup(rows: list[tuple[int, int, int]]) -> None:
    from app.core.database import AsyncSessionLocal
    from app.models.candidate import Candidate
    from app.models.client import Client
    from app.models.contract import Contract as C
    from sqlalchemy import delete

    async with AsyncSessionLocal() as db:
        for contract_id, _, _ in rows:
            # ClientOrder rows cascade via ON DELETE CASCADE on contract_id.
            await db.execute(delete(C).where(C.id == contract_id))
        for _, cand_id, client_id in rows:
            await db.execute(delete(Candidate).where(Candidate.id == cand_id))
            await db.execute(delete(Client).where(Client.id == client_id))
        await db.commit()


def _find_row(ws, contract_id: int) -> list:
    """Return the export row (list of cell values) whose ID cell == contract_id."""
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[0] == contract_id:
            return list(row)
    raise AssertionError(f"contract {contract_id} not found in export sheet")


@pytest.mark.asyncio
async def test_export_defaults_to_xlsx_with_client_rates_and_order_dates(
    app_client: AsyncClient, app_auth_headers: dict
):
    seeded = await _seed_contract(with_order=True)
    contract_id, _, client_id = seeded
    try:
        # No `format` → XLSX by default. Scope to the seeded client so the sheet
        # only carries our row (and stays under the export row cap).
        r = await app_client.get(
            "/api/contracts/export",
            params={"client_id": client_id},
            headers=app_auth_headers,
        )
        assert r.status_code == 200, r.text
        assert r.headers["content-type"].startswith(XLSX_MEDIA)
        assert ".xlsx" in r.headers.get("content-disposition", "")

        wb = load_workbook(BytesIO(r.content))
        ws = wb.active
        assert [c.value for c in ws[1]] == _CONTRACT_EXPORT_COLUMNS
        row = _find_row(ws, contract_id)
        by = dict(zip(_CONTRACT_EXPORT_COLUMNS, row))
        assert by["Klient"].startswith("ExportCo-")
        assert by["Status"] == "Aktywny"
        assert by["Typ"] == "B2B"
        assert by["Stawka klienta"] == 10000
        assert by["Stawka kandydata"] == 8000
        assert by["Marża"] == 2000
        assert by["Waluta"] == "PLN"
        assert by["Numer projektu/zamówienia"] == "ZAM-77"
        # Order dates: contract-level PO end + latest ClientOrder end.
        assert by["Koniec zamówienia u klienta"] == (
            date.today() + timedelta(days=45)
        ).isoformat()
        assert by["Najnowsze zamówienie do"] == (
            date.today() + timedelta(days=90)
        ).isoformat()
    finally:
        await _cleanup([seeded])


@pytest.mark.asyncio
async def test_export_csv_has_utf8_bom_and_header(
    app_client: AsyncClient, app_auth_headers: dict
):
    seeded = await _seed_contract()
    contract_id, _, client_id = seeded
    try:
        r = await app_client.get(
            "/api/contracts/export",
            params={"format": "csv", "client_id": client_id},
            headers=app_auth_headers,
        )
        assert r.status_code == 200, r.text
        assert r.headers["content-type"].startswith("text/csv")
        assert ".csv" in r.headers.get("content-disposition", "")
        # BOM so Excel renders Polish diacritics.
        assert r.content.startswith(b"\xef\xbb\xbf")
        text = r.content.decode("utf-8-sig")
        lines = text.splitlines()
        assert lines[0].split(",")[0] == "ID"
        assert any(line.startswith(f"{contract_id},") for line in lines[1:])
    finally:
        await _cleanup([seeded])


@pytest.mark.asyncio
async def test_export_respects_status_filter(
    app_client: AsyncClient, app_auth_headers: dict
):
    active = await _seed_contract(status="active")
    draft = await _seed_contract(status="draft")
    try:
        r = await app_client.get(
            "/api/contracts/export",
            params={"format": "csv", "status": "active"},
            headers=app_auth_headers,
        )
        assert r.status_code == 200, r.text
        text = r.content.decode("utf-8-sig")
        ids = {line.split(",")[0] for line in text.splitlines()[1:] if line}
        assert str(active[0]) in ids
        assert str(draft[0]) not in ids
    finally:
        await _cleanup([active, draft])
