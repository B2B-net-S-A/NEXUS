"""Import plików recepcji: synonimy nagłówków i dane, których NIE przenosimy (D10).

Pliki recepcji miały pełne klucze BitLocker, kody PUK, PESEL-e i prywatne
e-maile. Test buduje fikcyjny skoroszyt w kształcie tamtych plików (nagłówki
jak u recepcji, wartości zmyślone) i sprawdza, że parser: czyta kolumny po
nagłówkach, nie czyta kolumn z sekretem w nazwie, wycina sekret z kolumn,
które czyta, i pomija arkusze ukryte. Bez bazy.
"""

from __future__ import annotations

import io
from datetime import datetime

import openpyxl
import pytest

from app.services.equipment import importer, sanitize
from app.services.equipment.registry import EquipmentError

# Fikcyjne wartości o kształcie sekretów (nie są niczyimi danymi).
FAKE_BITLOCKER = "111111-222222-333333-444444-555555-666666-777777-888888"
FAKE_PESEL = "44051401458"  # przykład z dokumentacji formatu (suma kontrolna OK)
FAKE_PUK = "12345678"


def test_secret_detectors() -> None:
    assert sanitize.is_bitlocker_key(FAKE_BITLOCKER)
    assert sanitize.is_bitlocker_key(FAKE_BITLOCKER.replace("-", ""))
    assert not sanitize.is_bitlocker_key("123456-123456")
    assert sanitize.is_pesel(FAKE_PESEL)
    assert not sanitize.is_pesel("44051401459")  # zła cyfra kontrolna
    assert not sanitize.is_pesel("500000077")
    assert sanitize.looks_like_puk(FAKE_PUK)
    assert not sanitize.looks_like_puk("1234567")
    assert sanitize.is_private_email("anna.przykladowa@gmail.com")
    assert not sanitize.is_private_email("anna.przykladowa@example.com")


@pytest.mark.parametrize(
    ("header", "reason"),
    [
        ("Klucz do Komputera", "klucz BitLocker"),
        ("BitLocker recovery", "klucz BitLocker"),
        ("PUK 1", "kod PUK"),
        ("PIN", "kod PIN"),
        ("PESEL", "PESEL"),
        ("Hasło", "hasło"),
        ("E-mail prywatny", "prywatny e-mail"),
    ],
)
def test_secret_headers_are_recognised(header: str, reason: str) -> None:
    assert sanitize.secret_header_reason(header) == reason


def test_innocent_headers_are_not_secret() -> None:
    for header in ("Numer seryjny / Service tag", "Uwagi", "Spinacz", "Shipping"):
        assert sanitize.secret_header_reason(header) is None


def test_clean_text_strips_secrets_from_free_text() -> None:
    cleaned = sanitize.clean_text(
        f"oddany, klucz {FAKE_BITLOCKER}, PUK {FAKE_PUK}, kontakt jan@gmail.com, "
        f"pesel {FAKE_PESEL}"
    )
    assert cleaned.value == "oddany, klucz , PUK , kontakt , pesel"
    assert set(cleaned.reasons) == {
        "klucz BitLocker",
        "kod PUK",
        "prywatny e-mail",
        "PESEL",
    }


def test_identifier_column_keeps_eight_digit_card_numbers() -> None:
    cleaned = sanitize.clean_text("12345678", identifier=True)
    assert cleaned.value == "12345678"
    assert cleaned.reasons == ()


@pytest.mark.parametrize(
    ("header", "field"),
    [
        ("Numer seryjny / Service tag", "serial_number"),
        ("Nr\nLaptopa", "inventory_number"),
        ("Rodzaj urządzenia", "category"),
        ("Numer usługi", "phone_number"),
        ("NR. KARTY", "access_card_number"),
        ("KARTA PRZYPISANA", "holder"),
        ("Użytkownik", "holder"),
        ("lnformacja", "status"),
        ("Status laptopa", "status"),
        ("Po kim telefon", "notes"),
        ("Data zakupu", "purchase_date"),
        ("Kolumna1", None),
    ],
)
def test_header_synonyms(header: str, field: str | None) -> None:
    assert importer.match_header(header)[0] == field


def _workbook() -> bytes:
    workbook = openpyxl.Workbook()
    laptops = workbook.active
    laptops.title = "LAPTOPY_01.01.26"
    laptops.append(
        [
            "Status laptopa",
            "Nr\nLaptopa",
            "Rodzaj urządzenia",
            "Model",
            "Numer seryjny / Service tag",
            "Klucz do Komputera",
            "Użytkownik",
            "Projekt",
            "Uwagi",
            "Data zakupu",
            "Kolumna1",
        ]
    )
    laptops.append(
        [
            "wolny",
            901,
            "Laptop",
            "Dell Vostro 3530",
            "TEST0001",
            FAKE_BITLOCKER,
            None,
            "biuro",
            f"po Annie, PUK {FAKE_PUK}",
            datetime(2025, 3, 1),
            FAKE_BITLOCKER,
        ]
    )
    laptops.append(
        [
            None,
            902,
            "Laptop",
            "Dell Vostro 3530",
            "ST:TEST0002",
            None,
            "Osoba Nieznana",
            "Klient Fikcyjny",
            "kontakt anna@gmail.com",
            None,
            None,
        ]
    )
    laptops.append([None, 903, "Laptop", "Lenovo E14", "TEST0003"])
    laptops.append([None, 904, "Laptop", "Lenovo E14", "test0001"])  # duplikat S/N
    laptops.append([None, None, "Laptop", "Lenovo E14", None, None, "Ktoś"])
    phones = workbook.create_sheet("TELEFONY")
    phones.append(
        [
            None,
            "Rodzaj urządzenia",
            "Model",
            "Numer seryjny / Service tag",
            "Użytkownik",
        ]
    )
    phones.append(
        [None, "Telefon", "Redmi A5", "IMEI: 350000000000011", "wolny-potwierdzony"]
    )
    sims = workbook.create_sheet("SIM")
    sims.append(["Lp", "Numer usługi", "Taryfa", "PUK", "Użytkownik"])
    sims.append([1, "+48 500 000 077", "karty głosowe", FAKE_PUK, "wolne"])
    sims.append([2, FAKE_PESEL, "karty głosowe", FAKE_PUK, None])
    cards = workbook.create_sheet("Karty")
    cards.append(["LISTA UŻYTKOWNIKÓW"])
    cards.append(
        ["L.P", "NR. KARTY", "UŻYTKOWNIK", "KARTA PRZYPISANA", "FIRMA", "DOSTĘP"]
    )
    cards.append([1, 1611, "Firma", "Anna Przykładowa", "Firma", "OGÓLNY + BIURO"])
    hidden = workbook.create_sheet("Stare")
    hidden.append(["Numer seryjny", "Model"])
    hidden.append(["HIDDEN1", "Coś"])
    hidden.sheet_state = "hidden"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _rows_by_sheet(result: importer.ParseResult) -> dict[str, list[importer.ParsedRow]]:
    out: dict[str, list[importer.ParsedRow]] = {}
    for row in result.rows:
        out.setdefault(row.sheet, []).append(row)
    return out


def test_parser_reads_by_headers_and_drops_secrets() -> None:
    result = importer.parse_workbook(_workbook())
    rows = _rows_by_sheet(result)

    first = rows["LAPTOPY_01.01.26"][0].data
    assert first["inventory_number"] == "901"
    assert first["serial_number"] == "TEST0001"
    assert first["category"] == "laptop"
    assert first["model"] == "Dell Vostro 3530"
    assert first["purchase_date"] == "2025-03-01"
    assert FAKE_PUK not in (first["notes"] or "")
    serialized = repr([row.data for row in result.rows])
    for secret in (FAKE_BITLOCKER, FAKE_PUK, FAKE_PESEL, "anna@gmail.com"):
        assert secret not in serialized

    second = rows["LAPTOPY_01.01.26"][1].data
    assert second["serial_number"] == "TEST0002"  # prefiks „ST:” zdjęty
    assert second["holder_text"] == "Osoba Nieznana"

    phone = rows["TELEFONY"][0].data
    assert (phone["imei"], phone["serial_number"], phone["category"]) == (
        "350000000000011",
        None,
        "telefon",
    )
    sim_ok, sim_pesel = rows["SIM"]
    assert sim_ok.data["phone_number"] == "500000077"
    assert sim_ok.data["category"] == "karta_sim"
    assert sim_pesel.data["phone_number"] is None
    assert "usunięto: PESEL" in sim_pesel.reasons

    card = rows["Karty"][0].data
    assert card["access_card_number"] == "1611"
    assert card["holder_text"] == "Anna Przykładowa"  # „Karta przypisana” wygrywa
    assert card["category"] == "karta_dostepu"

    removed = {(c["sheet"], c["column"]): c["reason"] for c in result.removed_columns}
    assert removed[("LAPTOPY_01.01.26", "Klucz do Komputera")] == "klucz BitLocker"
    assert removed[("SIM", "PUK")] == "kod PUK"
    # Nieznana kolumna z kluczem w środku: nie czytana, ale zgłoszona recepcji.
    assert removed[("LAPTOPY_01.01.26", "Kolumna1")] == "klucz BitLocker"

    sheets = {s.name: s for s in result.sheets}
    assert sheets["Stare"].skipped == "arkusz ukryty"
    assert "HIDDEN1" not in serialized
    assert sheets["Karty"].header_row == 2


def test_holder_state_reads_free_words_and_disposal() -> None:
    assert importer.holder_state({"holder_text": "wolny - w szafce"}) == (
        "na_stanie",
        None,
    )
    assert importer.holder_state(
        {"status_text": "DO UTYLIZACJI", "holder_text": "X"}
    ) == (
        "wycofany",
        None,
    )
    assert importer.holder_state({"holder_text": "Anna Przykładowa"}) == (
        "wydany",
        "Anna Przykładowa",
    )
    assert importer.holder_state({"status_text": "do czyszczenia"}) == (
        "na_stanie",
        None,
    )
    assert importer.holder_state({}) == ("nieznany", None)


def test_not_an_excel_file_is_refused() -> None:
    with pytest.raises(EquipmentError):
        importer.parse_workbook(b"to nie jest plik excel")
    with pytest.raises(EquipmentError):
        importer.parse_workbook(b"PK\x03\x04uszkodzony")


def test_zip_bomb_workbook_is_refused_before_parsing(monkeypatch) -> None:
    """Mały plik, który rozwija się ponad limit części XML, nie trafia do openpyxl.

    Przegląd bezpieczeństwa 16.10: 7,9 MB rozwijające się do 68 MB sharedStrings
    parsowało się 62 s w pętli zdarzeń. Strażnik ma odmówić ZANIM openpyxl
    zacznie czytać — dlatego atrapa ``load_workbook`` wywraca test, gdy zostanie
    wywołana.
    """
    import io as _io
    import zipfile as _zipfile

    import pytest as _pytest

    from app.services.equipment import importer as _importer
    from app.services.equipment.registry import EquipmentError as _EquipmentError

    def _must_not_parse(*_args, **_kwargs):
        raise AssertionError("openpyxl czytał plik, który strażnik miał odrzucić")

    monkeypatch.setattr(_importer.openpyxl, "load_workbook", _must_not_parse)
    buffer = _io.BytesIO()
    with _zipfile.ZipFile(buffer, "w", _zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("xl/workbook.xml", "<workbook/>")
        archive.writestr(
            "xl/sharedStrings.xml", "<sst>" + "<si><t>a</t></si>" * 900_000 + "</sst>"
        )
    content = buffer.getvalue()
    assert len(content) < _importer.MAX_FILE_BYTES
    with _pytest.raises(_EquipmentError):
        _importer.parse_workbook(content)


def test_status_column_wins_over_state_in_holder_column() -> None:
    """Pliki recepcji 18.10.2026: „Status laptopa: wolny” i „Użytkownik: Gotowy”.

    Import na produkcji brał „Gotowy”, „do czyszczenia” i „UTYLIZACJA” z kolumny
    osoby za nazwiska — 149 wierszy „do wyjaśnienia” w głównym pliku.
    """
    state = importer.holder_state
    assert state({"status_text": "wolny", "holder_text": "do czyszczenia"}) == (
        "na_stanie",
        None,
    )
    assert state({"status_text": "czysty", "holder_text": "Gotowy"}) == (
        "na_stanie",
        None,
    )
    assert state({"holder_text": "Gotowy"}) == ("na_stanie", None)
    assert state({"holder_text": "Stoi na Open"}) == ("na_stanie", None)
    assert state({"holder_text": "do utylizacji"}) == ("wycofany", None)
    assert state({"holder_text": "Do czyszczenia/utylizacja"}) == ("wycofany", None)
    assert state({"holder_text": "Anna Przykładowa - wykupiony"}) == ("wycofany", None)
    assert state({"holder_text": "Zaginiona (Ola Testowa)"}) == ("zgubiony", None)
    # „nowe” to stan egzemplarza, nie dostępność — AirPods z osobą są wydane.
    assert state({"status_text": "nowe", "holder_text": "Anna Przykładowa"}) == (
        "wydany",
        "Anna Przykładowa",
    )
    assert state({"status_text": "w użyciu", "holder_text": "Anna Przykładowa"}) == (
        "wydany",
        "Anna Przykładowa",
    )
    # Recepcja sama oznaczyła wiersz jako niepewny — zostaje dla człowieka.
    assert state({"status_text": "do wyjaśnienia", "holder_text": "Anna X"}) == (
        "do_wyjasnienia",
        "Anna X",
    )


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("Anna Przykładowa", ["Anna Przykładowa"]),
        ("Przykładowa Anna", ["Przykładowa Anna"]),
        ("Anna Przykładowa recepcja", ["Anna Przykładowa"]),
        ("Jan Testowy obecnie posiada", ["Jan Testowy"]),
        ("Ewa Wzorcowa(zamiana)", ["Ewa Wzorcowa"]),
        ("Hubert Próbny?", ["Hubert Próbny"]),
        ("Kinga Fikcyjna do domu", ["Kinga Fikcyjna"]),
        (
            "Jan Pierwszy -> Ewa Druga->Ola Trzecia",
            ["Ola Trzecia", "Ewa Druga", "Jan Pierwszy"],
        ),
        ("Jan Pierwszy -> dac Dianie", ["Jan Pierwszy"]),
        ("Igor do wysyłki maili", []),
        ("Oddany dla Jana Testowego", []),
        ("Oscar", []),
        ("recepcja/zamiast tableta", []),
        # „Ś” rozłożone na literę i znak łączący (tak jest w pliku kart).
        ("Patryk Świętoń", ["Patryk Świętoń"]),
    ],
)
def test_person_candidates(text: str, expected: list[str]) -> None:
    assert importer.person_candidates(text) == expected


def test_summary_rows_and_first_name_column() -> None:
    """Wiersze „RAZEM/ZAPAS” to nie egzemplarze; imię z kolumny bez nagłówka (SIM)."""
    assert importer._is_summary_row({"serial_number": "RAZEM", "holder_text": "51"})
    assert importer._is_summary_row({"serial_number": "BIURO UŻYWA"})
    assert not importer._is_summary_row({"serial_number": "C7DYLG3"})

    workbook = openpyxl.Workbook()
    sim = workbook.active
    sim.title = "na dzień 05.01.2026"
    sim.append([None, None, "Numer usługi", "Taryfa", None, "Użytkownik"])
    sim.append([1, None, "48500000001", "abonament", "Anna", "Przykładowa"])
    sim.append([2, None, "48500000002", "abonament", "wolne", None])
    sim.append([3, None, "48500000003", "abonament", "Jan", "Jan Testowy"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    rows = importer.parse_workbook(buffer.getvalue()).rows
    assert [r.data["holder_text"] for r in rows] == [
        "Anna Przykładowa",
        None,
        "Jan Testowy",
    ]
