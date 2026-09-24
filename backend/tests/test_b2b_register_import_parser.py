"""Parser rejestru umów z Excela działu — syntetyczny XLSX, bez bazy.

Dane są fikcyjne (nazwiska testowe), ale kształt pliku odtwarza prawdziwy
rejestr: nagłówki A–O, numery liczbowe i tekstowe, trzy tryby daty startu,
„nie doszła do skutku”, czerwone tło kolumny A, legenda pod tabelą i arkusz
„Bez dzialalności ” ze spacją na końcu nazwy.
"""

from __future__ import annotations

import io
from datetime import date, datetime
from typing import Any, Optional

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.styles.colors import Color

from app.services.b2b_register_import.parser import (
    RegisterParseError,
    closure_date_from,
    is_green,
    is_orange,
    is_red,
    parse_register,
)

HEADERS = [
    "NAZWISKO, PÓŹNIEJ IMIE",
    "Numer umowy",
    "Klient",
    "Stanowisko",
    "Rodzaj umowy",
    "Data podpisania",
    "Data startu pracy",
    "Data zakończenia umowy",
    "Okres lojalności",
    "Rekruter",
    "czy wysłano informację o rozliczeniach",
    "mail powitalny",
    "UWAGI",
    "ZMIANY W UMOWIE",
    None,
]
NB_HEADERS = [
    "IMIĘ I NAZWISKO",
    "DATA STARTU PRACY",
    "REKRUTER",
    "Aneksy do umów – kiedy zrobione",
    "KLIENT",
    "UWAGI DO WPISANIA W ANEKSIE",
]
RED = PatternFill(fill_type="solid", fgColor="FFFF0000")
ORANGE = PatternFill(fill_type="solid", fgColor=Color(theme=5, tint=0.5999))
GREEN = PatternFill(fill_type="solid", fgColor="FF92D050")


def contract_row(
    name: str,
    number: Any,
    *,
    client: Optional[str] = "Klient Testowy",
    position: Optional[str] = "tester manualny",
    kind: str = "B2B",
    signing: Any = None,
    start: Any = None,
    end: Any = "czas nieokreślony",
    recruiter: Optional[str] = None,
    notes: Optional[str] = None,
    changes: Optional[str] = None,
    red: bool = False,
    orange: bool = False,
) -> dict[str, Any]:
    return {
        "values": [
            name,
            number,
            client,
            position,
            kind,
            signing,
            start,
            end,
            None,
            recruiter,
            "tak",
            "kandydat z bazy.",
            notes,
            changes,
            None,
        ],
        "red": red,
        "orange": orange,
    }


def build_register(
    rows: list[dict[str, Any]],
    *,
    no_business: Optional[list[dict[str, Any]]] = None,
    legend: bool = True,
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Umowy B2B"
    sheet.append(HEADERS)
    for row in rows:
        values = [
            datetime.combine(v, datetime.min.time()) if isinstance(v, date) else v
            for v in row["values"]
        ]
        sheet.append(values)
        index = sheet.max_row
        if row.get("red"):
            sheet.cell(index, 1).fill = RED
        if row.get("orange"):
            for column in range(1, 15):
                sheet.cell(index, column).fill = ORANGE
    if legend:
        sheet.append([])
        sheet.append(["numer kolejnej umowy", 99_999_999])
        sheet.append(["UMOWY ZLECENIE NIE MAJĄ NUMERÓW!"])
        sheet.append(["Numeracja umów jest ciągła, zmieniamy jedynie rok!"])
        sheet.append(["Umowa nie doszła do skutku"])
    nb = workbook.create_sheet("Bez dzialalności ")
    nb.append(NB_HEADERS)
    for row in no_business or []:
        values = [
            datetime.combine(v, datetime.min.time()) if isinstance(v, date) else v
            for v in row["values"]
        ]
        nb.append(values)
        if row.get("green"):
            for column in range(1, 7):
                nb.cell(nb.max_row, column).fill = GREEN
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def no_business_row(
    name: str,
    start: Any,
    *,
    client: str = "Klient Testowy",
    annex: Optional[str] = "aneks do zrobienia",
    green: bool = False,
) -> dict[str, Any]:
    return {"values": [name, start, "Rekruterka", annex, client, None], "green": green}


def test_numbers_dates_and_statuses_are_read():
    payload = build_register(
        [
            contract_row(
                "Testowa Anna", 1517, signing=date(2026, 9, 18), start=date(2026, 10, 1)
            ),
            contract_row(
                "Próbny Jan",
                "264A",
                signing="12.03.2019",
                start="nie później niż 01.04.2019",
            ),
            contract_row(
                "Wzorcowa Ewa",
                "bez numeru",
                kind="zlecenie",
                start="nie wcześniej niż 15.05.2022",
            ),
            contract_row(
                "Fikcyjny Piotr", "zlecenie", kind="zlecenie", start="02.02.2023 r."
            ),
            contract_row(
                "Niedoszła Maria",
                1400,
                signing="umowa nie doszła do skutku",
                start="jak będzie klient",
            ),
            contract_row("Tekstowy Adam", "412", signing=date(2020, 1, 2), start="?"),
            contract_row(
                "Literówka Olga", 1300, signing="30.07.026", end="czas nieokeślony"
            ),
        ]
    )
    parsed = parse_register(payload)
    rows = {row.name_raw: row for row in parsed.contracts}
    assert len(parsed.contracts) == 7
    assert [s["reason"] for s in parsed.skipped].count("legend") == 4

    anna = rows["Testowa Anna"]
    assert (anna.number_int, anna.number_raw) == (1517, "1517")
    assert anna.partner_name == "Anna Testowa"
    assert anna.signing_date == date(2026, 9, 18)
    assert (anna.start_date, anna.start_date_mode) == (date(2026, 10, 1), "exact")
    assert anna.kind == "b2b"
    assert anna.end_date is None

    jan = rows["Próbny Jan"]
    assert jan.number_int is None and jan.number_raw == "264A"
    assert (jan.start_date, jan.start_date_mode) == (date(2019, 4, 1), "not_later")

    ewa = rows["Wzorcowa Ewa"]
    assert ewa.number_raw == "bez numeru" and ewa.kind == "mandate"
    assert ewa.start_date_mode == "not_earlier"

    assert rows["Fikcyjny Piotr"].start_date_mode == "exact"
    assert rows["Fikcyjny Piotr"].start_date == date(2023, 2, 2)

    maria = rows["Niedoszła Maria"]
    assert maria.cancelled is True
    assert maria.start_date is None and "start_date_unknown" in maria.flags

    assert rows["Tekstowy Adam"].number_int == 412
    assert rows["Literówka Olga"].signing_date == date(2026, 7, 30)


def test_red_name_cell_with_and_without_closure_date():
    payload = build_register(
        [
            contract_row(
                "Zakończony Karol",
                900,
                signing=date(2023, 1, 5),
                notes="rozwiązanie za porozumieniem stron z dniem 31.03.2024",
                red=True,
            ),
            contract_row(
                "Bezdatny Leon",
                901,
                signing=date(2023, 1, 6),
                notes="klient zrezygnował",
                red=True,
            ),
            contract_row(
                "Pomarańczowy Tomasz", 902, position="(bez projektu)", orange=True
            ),
        ]
    )
    rows = {row.name_raw: row for row in parse_register(payload).contracts}
    assert rows["Zakończony Karol"].red is True
    assert rows["Zakończony Karol"].closure_date == date(2024, 3, 31)
    assert rows["Bezdatny Leon"].closure_date is None
    assert "likely_ended" in rows["Bezdatny Leon"].flags
    tomasz = rows["Pomarańczowy Tomasz"]
    assert tomasz.orange and not tomasz.red
    assert {"without_project", "row_highlight_orange"} <= set(tomasz.flags)


def test_no_business_sheet_with_trailing_space_is_found():
    payload = build_register(
        [
            contract_row(
                "Bezfirmowa Zofia",
                903,
                signing=date(2024, 5, 1),
                start=date(2024, 5, 6),
            )
        ],
        no_business=[
            no_business_row(
                "Zofia Bezfirmowa", date(2024, 5, 6), annex="aneks zrobiony 12.06.2024"
            ),
            no_business_row("Ktoś Inny", date(2024, 1, 1), green=True, annex=None),
        ],
    )
    parsed = parse_register(payload)
    assert parsed.no_business_sheet_found
    zofia, other = parsed.no_business
    assert zofia.done and zofia.done_date == date(2024, 6, 12)
    assert zofia.tokens == parsed.contracts[0].tokens
    assert other.done and other.done_date is None


def test_file_without_register_sheet_is_refused():
    workbook = Workbook()
    workbook.active.append(["coś", "innego"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    with pytest.raises(RegisterParseError):
        parse_register(buffer.getvalue())
    with pytest.raises(RegisterParseError):
        parse_register(b"not a zip")


def test_color_classification():
    assert is_red("FF0000") and is_red("EE0000") and not is_red("92D050")
    assert is_orange("theme:5:0.6") and not is_orange("theme:5:0.0")
    assert is_green("92D050") and is_green("99FF99") and not is_green("FF0000")


def test_closure_date_takes_latest_date_next_to_keyword():
    text = "zmiana stawki od 01.01.2024; wypowiedzenie złożone 10.02.2024, koniec 31.03.2024"
    assert closure_date_from(text) == date(2024, 3, 31)
    assert closure_date_from("zmiana stawki od 01.01.2024") is None
