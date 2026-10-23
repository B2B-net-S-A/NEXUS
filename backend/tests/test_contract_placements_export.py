"""Eksport „Placementy” (/api/contracts/export/placements).

1. Czyste testy wiersza, uwag i premii — bez bazy.
2. Trasa w procesie (app_client): okres, statusy, nagłówek i typy komórek.
"""

from __future__ import annotations

import uuid
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook

from app.models.contract import (
    Contract,
    ContractStatus,
    ContractTerminationReason,
    ContractType,
    RateUnit,
)
from app.services.contract_placements_export import (
    PLACEMENT_EXPORT_COLUMNS,
    Acceptance,
    PlacementBonus,
    PlacementFacts,
    SignedContracts,
    bonus_text,
    placement_notes,
    placement_row,
    upcoming_rate_steps,
    visible_bonuses,
)
from app.services.hr.payroll_access import PayrollViewer, is_payroll

XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ── Bez bazy ─────────────────────────────────────────────────────────────────


def _user(role, *permissions: str):
    from app.models.user import User
    from app.services import permission_catalog as catalog

    person = User(
        id=77,
        email=f"{role.value}-placements@example.com",
        name=role.value,
        role=role,
        roles=[role.value],
        is_active=True,
        profile_completed=True,
    )
    held = catalog.close(permissions)
    person.effective_action_access = {key: "manage" for key in held}
    person.effective_section_access = catalog.derive_sections(held)
    return person


def test_finance_module_without_payroll_is_not_payroll():
    """Bramka eksportu przepuszcza „Moduł Finanse”, Płace — nie: premie innych
    osób musi odfiltrować sam eksport."""

    from app.models.user import UserRole

    assert is_payroll(_user(UserRole.recruiter, "finance_module")) is False
    assert is_payroll(_user(UserRole.finance, "payroll")) is True


def test_bonuses_follow_payroll_visibility():
    rows = [
        ("recruiter", 77, Decimal("1500"), False),
        ("dl", 12, Decimal("772.80"), True),
        ("dl_extra", 13, Decimal("772.80"), True),
    ]
    finance_module = PayrollViewer(
        user_id=77, payroll=False, admin=False, report_ids=frozenset({13})
    )
    payroll = PayrollViewer(
        user_id=5, payroll=True, admin=False, report_ids=frozenset()
    )

    # Sama osoba i jej podwładny — jak lista premii w Płacach; obcej premii DL brak.
    assert [r[1] for r in visible_bonuses(rows, finance_module.sees_person)] == [
        77,
        13,
    ]
    assert visible_bonuses(rows, payroll.sees_person) == rows


def _contract(**overrides) -> Contract:
    values = dict(
        candidate_id=1,
        client_id=1,
        job_id=7,
        status=ContractStatus.active,
        contract_type=ContractType.b2b,
        start_date=date(2026, 10, 19),
        rate_candidate=Decimal("150"),
        rate_client=Decimal("196"),
        rate_unit=RateUnit.hourly,
        currency="PLN",
    )
    values.update(overrides)
    c = Contract(**values)
    # Relacje „wczytane” bez sesji — tak wygląda obiekt po selectinload.
    c.__dict__.update(
        job=SimpleNamespace(title="Architekt"),
        candidate_rate_schedule=[],
        client_rate_schedule=[],
    )
    return c


def _facts(**overrides) -> PlacementFacts:
    values = dict(
        dl_name="Ewa Kowal",
        recruiter_name="Jan Bąk",
        verified_by_dl=False,
        margin_per_hour=Decimal("46"),
        cost_per_hour=Decimal("150"),
        client_per_hour=Decimal("196"),
        cost_currency="PLN",
        client_currency="PLN",
        acceptance=Acceptance(date(2026, 10, 13), True),
        generated=SignedContracts(1, date(2026, 10, 16)),
        bonuses=(),
        upcoming_rates=(),
    )
    values.update(overrides)
    return PlacementFacts(**values)


def test_row_follows_sheet_columns_and_counts_monthly_margin_on_168_hours():
    row = placement_row(_contract(), _facts(), "Anna Wrona", "Nordea")
    by = dict(zip(PLACEMENT_EXPORT_COLUMNS, row))

    assert len(row) == len(PLACEMENT_EXPORT_COLUMNS) == 14
    assert by["Konsultant"] == "Anna Wrona"
    assert by["DL"] == "Ewa Kowal"
    assert by["Stanowisko"] == "Architekt"
    assert (by["Stawka kosztowa"], by["Stawka przychodowa"], by["Marża"]) == (
        150.0,
        196.0,
        46.0,
    )
    assert by["Marża mies."] == 7728.0
    # Daty zostają datami — arkusz sortuje i filtruje je jak w Excelu Dominika.
    assert by["Data podpisania"] == date(2026, 10, 16)
    assert by["Start day"] == date(2026, 10, 19)
    assert by["Data akceptacji"] == date(2026, 10, 13)
    assert by["Uwagi"] == ""


def test_unknown_margin_leaves_empty_cells_and_says_why():
    row = placement_row(
        _contract(),
        _facts(margin_per_hour=None, cost_per_hour=None, client_per_hour=None),
        "Anna Wrona",
        "Nordea",
    )
    by = dict(zip(PLACEMENT_EXPORT_COLUMNS, row))

    assert by["Marża"] == by["Marża mies."] == by["Stawka kosztowa"] == ""
    assert "Brak stawki albo kursu — marża nieznana" in by["Uwagi"]


def test_notes_for_ended_contract_without_recruitment_or_generator():
    contract = _contract(
        job_id=None,
        status=ContractStatus.ended,
        end_date=date(2026, 11, 30),
        termination_reason=ContractTerminationReason.consultant_resigned,
        returned_from_contract_id=3,
    )
    notes = placement_notes(
        contract,
        _facts(acceptance=None, generated=SignedContracts(0, None)),
    )

    assert notes == (
        "Koniec 30.11.2026 (konsultant zrezygnował); Powrót po przerwie; "
        "Kontrakt bez rekrutacji; Umowa spoza Generatora"
    )


def test_notes_prefer_termination_date_over_contract_end():
    contract = _contract(
        status=ContractStatus.ending,
        end_date=date(2026, 12, 31),
        terminated_at=date(2026, 11, 15),
    )
    assert placement_notes(contract, _facts()) == "Koniec 15.11.2026"


def test_notes_flag_dl_verifier_move_date_acceptance_unsigned_and_currency():
    notes = placement_notes(
        _contract(),
        _facts(
            verified_by_dl=True,
            acceptance=Acceptance(date(2026, 10, 2), False),
            generated=SignedContracts(2, None),
            client_currency="EUR",
        ),
    )

    assert notes.split("; ") == [
        "Zweryfikował DL — bez premii rekrutera",
        "Data akceptacji z ruchu karty",
        "Umowa z Generatora bez podpisu",
        "Stawki: kosztowa PLN, przychodowa EUR",
    ]


def test_bonus_text_orders_roles_and_marks_proposals():
    text = bonus_text(
        [
            PlacementBonus("dl", "Ewa Kowal", Decimal("772.80"), True),
            PlacementBonus("recruiter", "Jan Bąk", Decimal("1500.00"), False),
        ]
    )
    assert text == "rekruter Jan Bąk 1500 zł; DL Ewa Kowal 772,80 zł (propozycja)"
    assert bonus_text([]) == ""


def test_upcoming_rate_steps_take_the_nearest_future_step_per_side():
    contract = _contract()
    contract.__dict__.update(
        candidate_rate_schedule=[
            SimpleNamespace(rate=Decimal("55"), effective_from=date(2026, 4, 1)),
            SimpleNamespace(rate=Decimal("65"), effective_from=date(2027, 10, 1)),
            SimpleNamespace(rate=Decimal("60"), effective_from=date(2027, 4, 1)),
        ],
        client_rate_schedule=[
            SimpleNamespace(rate=Decimal("100"), effective_from=date(2026, 4, 1)),
        ],
    )

    steps = upcoming_rate_steps(contract, date(2026, 10, 23))

    assert steps == (("kosztowa", Decimal("60"), date(2027, 4, 1)),)
    assert placement_notes(contract, _facts(upcoming_rates=steps)) == (
        "Stawka kosztowa 60 od 01.04.2027"
    )


# ── Trasa ────────────────────────────────────────────────────────────────────


async def _seed(status: ContractStatus, start: date) -> tuple[int, int, int]:
    from app.core.database import AsyncSessionLocal
    from app.models.candidate import Candidate
    from app.models.client import Client

    async with AsyncSessionLocal() as db:
        candidate = Candidate(
            name="Placement",
            lastname=f"P-{uuid.uuid4().hex[:6]}",
            email=f"pl-{uuid.uuid4().hex[:8]}@example.com",
        )
        client = Client(name=f"PlacementCo-{uuid.uuid4().hex[:6]}")
        db.add_all([candidate, client])
        await db.flush()
        contract = Contract(
            candidate_id=candidate.id,
            client_id=client.id,
            status=status,
            contract_type=ContractType.b2b,
            start_date=start,
            rate_candidate=Decimal("150"),
            rate_client=Decimal("196"),
            rate_unit=RateUnit.hourly,
            currency="PLN",
        )
        db.add(contract)
        await db.commit()
        return contract.id, candidate.id, client.id


async def _cleanup(rows: list[tuple[int, int, int]]) -> None:
    from sqlalchemy import delete

    from app.core.database import AsyncSessionLocal
    from app.models.candidate import Candidate
    from app.models.client import Client

    async with AsyncSessionLocal() as db:
        for contract_id, candidate_id, client_id in rows:
            await db.execute(delete(Contract).where(Contract.id == contract_id))
            await db.execute(delete(Candidate).where(Candidate.id == candidate_id))
            await db.execute(delete(Client).where(Client.id == client_id))
        await db.commit()


@pytest.mark.asyncio
async def test_export_lists_started_contracts_in_range_only(
    app_client: AsyncClient, app_auth_headers: dict
):
    # Rok 2001: poza danymi innych testów, więc arkusz ma tylko nasze wiersze.
    inside = await _seed(ContractStatus.active, date(2001, 3, 10))
    ended = await _seed(ContractStatus.ended, date(2001, 3, 20))
    draft = await _seed(ContractStatus.draft, date(2001, 3, 15))
    outside = await _seed(ContractStatus.active, date(2001, 4, 1))
    try:
        r = await app_client.get(
            "/api/contracts/export/placements",
            params={"start_from": "2001-03-01", "start_to": "2001-03-31"},
            headers=app_auth_headers,
        )
        assert r.status_code == 200, r.text
        assert r.headers["content-type"].startswith(XLSX_MEDIA)
        assert (
            "placementy_2001-03-01_2001-03-31.xlsx" in r.headers["content-disposition"]
        )

        ws = load_workbook(BytesIO(r.content)).active
        assert ws.title == "Placementy"
        assert tuple(c.value for c in ws[1]) == PLACEMENT_EXPORT_COLUMNS
        rows = [
            dict(zip(PLACEMENT_EXPORT_COLUMNS, row))
            for row in ws.iter_rows(min_row=2, values_only=True)
        ]
        assert [row["Start day"] for row in rows] == [
            datetime(2001, 3, 10),
            datetime(2001, 3, 20),
        ]
        first = rows[0]
        assert first["Konsultant"].startswith("Placement P-")
        assert first["Klient"].startswith("PlacementCo-")
        assert (first["Stawka kosztowa"], first["Stawka przychodowa"]) == (150, 196)
        assert first["Marża"] == 46
        assert first["Marża mies."] == 7728
        assert first["Uwagi"] == "Kontrakt bez rekrutacji; Umowa spoza Generatora"
    finally:
        await _cleanup([inside, ended, draft, outside])


@pytest.mark.asyncio
async def test_export_rejects_reversed_range(
    app_client: AsyncClient, app_auth_headers: dict
):
    r = await app_client.get(
        "/api/contracts/export/placements",
        params={"start_from": "2026-10-31", "start_to": "2026-10-01"},
        headers=app_auth_headers,
    )
    assert r.status_code == 422
    assert r.json()["detail"] == "Data „do” jest wcześniejsza niż „od”."


@pytest.mark.asyncio
async def test_export_requires_both_dates(
    app_client: AsyncClient, app_auth_headers: dict
):
    r = await app_client.get(
        "/api/contracts/export/placements",
        params={"start_from": "2026-10-01"},
        headers=app_auth_headers,
    )
    assert r.status_code == 422


@pytest.mark.asyncio
async def test_signing_day_comes_from_the_dialog_then_from_the_click(
    app_client: AsyncClient, app_auth_headers: dict
):
    """0514: „Data podpisania” = ``signed_on``; wiersz bez niego — dzień
    ``signed_at``. Kilka umów do kontraktu — najwcześniejszy dzień."""
    from datetime import timezone

    from sqlalchemy import delete

    from app.core.database import AsyncSessionLocal
    from app.models.b2b_generated_contract import B2BGeneratedContract

    with_day = await _seed(ContractStatus.active, date(2002, 5, 10))
    legacy = await _seed(ContractStatus.active, date(2002, 5, 20))

    def _row(contract_id: int, signed_on, signed_at) -> B2BGeneratedContract:
        seq = int(uuid.uuid4().hex[:7], 16)
        return B2BGeneratedContract(
            year=2099,
            seq=seq,
            contract_number=f"{seq}/2099",
            contract_id=contract_id,
            signature_status="signed_both",
            signed_on=signed_on,
            signed_at=signed_at,
        )

    async with AsyncSessionLocal() as db:
        db.add_all(
            [
                # Podpis 05.05, oznaczony 09.05.
                _row(
                    with_day[0],
                    date(2002, 5, 5),
                    datetime(2002, 5, 9, 10, tzinfo=timezone.utc),
                ),
                # Późniejsza umowa (np. po korekcie) nie przesuwa daty.
                _row(
                    with_day[0],
                    None,
                    datetime(2002, 5, 7, 10, tzinfo=timezone.utc),
                ),
                # Sprzed 0514: tylko chwila oznaczenia; 23:30 UTC = 18.05 w Warszawie.
                _row(
                    legacy[0], None, datetime(2002, 5, 17, 23, 30, tzinfo=timezone.utc)
                ),
            ]
        )
        await db.commit()
    try:
        r = await app_client.get(
            "/api/contracts/export/placements",
            params={"start_from": "2002-05-01", "start_to": "2002-05-31"},
            headers=app_auth_headers,
        )
        assert r.status_code == 200, r.text
        ws = load_workbook(BytesIO(r.content)).active
        rows = [
            dict(zip(PLACEMENT_EXPORT_COLUMNS, row))
            for row in ws.iter_rows(min_row=2, values_only=True)
        ]
        assert [row["Data podpisania"] for row in rows] == [
            datetime(2002, 5, 5),
            datetime(2002, 5, 18),
        ]
    finally:
        async with AsyncSessionLocal() as db:
            await db.execute(
                delete(B2BGeneratedContract).where(
                    B2BGeneratedContract.contract_id.in_([with_day[0], legacy[0]])
                )
            )
            await db.commit()
        await _cleanup([with_day, legacy])
