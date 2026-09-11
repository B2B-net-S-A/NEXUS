"""The 14 000-character limit belongs to the AI parser only (09.2026).

#1477 applied it to every Champion read: the Word form read from its tables
(no model), the CV-upload builder (no model) and — through the preview — the
Traffit collector. A long but perfectly structured profile was refused, and
the CV upload preflight reported it as "nie można odczytać pliku DOCX".
"""

from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document

from app.services.champion_intake import MAX_TEXT, preview_document
from app.services.champion_document import document_text
from app.services.cv_generator_b2b import upload_preflight as preflight
from app.services.cv_generator_b2b.champion_builder import (
    parse_champion_from_docx_bytes,
)
from app.services.cv_generator_b2b.standalone_service import (
    StandaloneGenerationError,
    UploadGenerationInput,
)
from app.services.cv_generator_b2b.text_extractor import CVTextExtractionError
from tests.test_cv_upload_preflight import docx, pdf

FIXTURE = Path(__file__).parent / "fixtures" / "champion" / "v4-two-questions.docx"


def long_form() -> bytes:
    """The v4 form plus a long free-text note: over the AI limit, still a form."""
    document = Document(FIXTURE)
    document.add_paragraph("Notatka rekrutera. " * 900)
    stream = BytesIO()
    document.save(stream)
    data = stream.getvalue()
    assert len(document_text(data)) > MAX_TEXT
    return data


def long_free_text() -> bytes:
    document = Document()
    document.add_paragraph("x" * (MAX_TEXT + 1))
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


async def test_long_word_form_is_read_without_the_ai_limit(monkeypatch):
    async def forbidden(*args, **kwargs):
        raise AssertionError("the v4 form must not reach the AI parser")

    monkeypatch.setattr(
        "app.services.champion_profile_ingest.parse_champion_document", forbidden
    )

    preview = await preview_document(long_form(), "long.docx")

    assert preview["must_skills"] == ["Java 17", "Spring Boot", "Kafka"]
    assert len(preview["champion_profile"]["screening_questions"]) == 2


async def test_ai_path_rejects_over_limit_before_spending_quota(monkeypatch):
    def quota_entered(*args, **kwargs):
        raise AssertionError("an over-long document must not open the AI quota")

    monkeypatch.setattr("app.services.ai_quota.ai_feature", quota_entered)

    with pytest.raises(ValueError, match="limit"):
        await preview_document(long_free_text(), "free.docx", db=object())


def test_cv_upload_builder_reads_a_long_form():
    champion = parse_champion_from_docx_bytes(long_form(), "Champion.docx")

    assert not champion.is_empty()
    assert "Kafka" in champion.must_have


def test_upload_preflight_accepts_a_long_form_the_client_requires():
    preflight.validate_upload_inputs(
        UploadGenerationInput(
            cv_bytes=docx(),
            cv_filename="CV.docx",
            champion_bytes=long_form(),
            champion_filename="Champion.docx",
            client_rule=SimpleNamespace(require_champion=True),
        )
    )


def test_upload_preflight_reports_the_real_reason(monkeypatch):
    def unreadable(*args, **kwargs):
        raise CVTextExtractionError("Profil ma nieobsługiwany układ tabel.")

    monkeypatch.setattr(preflight, "parse_champion_from_docx_bytes", unreadable)

    with pytest.raises(StandaloneGenerationError) as error:
        preflight.validate_upload_inputs(
            UploadGenerationInput(
                cv_bytes=pdf(),
                cv_filename="CV.pdf",
                champion_bytes=docx("Profil"),
                champion_filename="Champion.docx",
            )
        )

    assert error.value.message == (
        "Profil Championa: Profil ma nieobsługiwany układ tabel."
    )


def test_upload_preflight_names_an_oversized_archive(monkeypatch):
    monkeypatch.setattr(preflight, "MAX_UPLOAD_BYTES", 1_000)

    with pytest.raises(StandaloneGenerationError) as error:
        preflight.validate_upload_inputs(
            UploadGenerationInput(
                cv_bytes=pdf(),
                cv_filename="CV.pdf",
                champion_bytes=long_form(),
                champion_filename="Champion.docx",
            )
        )

    assert "rozmiar" in error.value.message
    assert "nie można odczytać" not in error.value.message
