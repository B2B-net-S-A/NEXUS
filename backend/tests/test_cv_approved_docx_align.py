"""Wyrównanie tekstu CV w pliku Word (justowanie, decyzja J1 z 14.10.2026).

Zespół prosił o justowanie. Sprawdzone 14.10.2026: edytor nie miał wyrównania,
a konwerter brał ``text-align`` tylko ze zwykłych akapitów i nagłówków —
punkt z „wyjustuj” wychodził w Wordzie do lewej. Teraz CV jest domyślnie
wyjustowane (reguła klienta może przełączyć na „do lewej”), a jawne
wyrównanie akapitu z edytora wygrywa w każdym bloku treści.
"""

from io import BytesIO

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.services.cv_approved_docx import cv_text_align, render_approved_docx
from app.services.cv_document_assets import default_template
from app.services.html_sanitizer import sanitize_cv_html

LONG = (
    "Projektowanie mikroserwisów płatności w Java 17 i Spring Boot dla autoryzacji "
    "kart i rozliczeń z kilkoma bankami partnerskimi"
)


def _html(style: str = "") -> str:
    """Kształt po edytorze; ``style`` trafia do każdego akapitu treści."""
    s = f' style="{style}"' if style else ""
    return (
        "<h1>Tester QA – Anna Testowa</h1><hr>"
        '<h2 data-cv-section="why_points">DLACZEGO NASZ KANDYDAT?</h2>'
        f"<ul><li><p{s}>Powód: {LONG}</p></li></ul>"
        '<h2 data-cv-section="education">EDUKACJA</h2>'
        "<p>2015 - 2019 <strong>Politechnika Testowa</strong> — Informatyka</p>"
        '<h2 data-cv-section="skills">UMIEJĘTNOŚCI</h2>'
        f"<p{s}><strong>Języki: </strong>Python, SQL, Kotlin, Java, Go, Rust, C#</p>"
        f'<h2 data-cv-section="languages">JĘZYKI</h2><p{s}>Polski - C2 · Angielski - B2</p>'
        '<h2 data-cv-section="experience">DOŚWIADCZENIE</h2>'
        f'<p data-cv-section="role"{s}><strong>Senior QA Engineer</strong> 12.2025 - obecnie</p>'
        f'<p data-cv-section="employer"{s}>Acme · Telekomunikacja</p>'
        f'<p data-cv-section="duties_label"{s}>Zakres zadań:</p>'
        f"<ul><li><p{s}>Obowiązek: {LONG}</p></li></ul>"
        f'<p data-cv-section="technologies"{s}>Technologie: Python, Jenkins</p>'
        '<p data-cv-section="rodo">Wyrażam zgodę na przetwarzanie moich danych osobowych.</p>'
    )


def _render(content: str, **kwargs) -> Document:
    return Document(
        BytesIO(
            render_approved_docx(
                sanitize_cv_html(content), default_template(), **kwargs
            )
        )
    )


def _align(paragraph):
    """Wyrównanie, które zobaczy Word: akapit, a gdy go nie ma — styl i jego bazy."""
    if paragraph.paragraph_format.alignment is not None:
        return paragraph.paragraph_format.alignment
    style = paragraph.style
    while style is not None:
        if style.paragraph_format.alignment is not None:
            return style.paragraph_format.alignment
        style = style.base_style
    return None


def _paragraph(doc: Document, starts: str):
    found = [p for p in doc.paragraphs if p.text.strip().startswith(starts)]
    assert found, f"brak akapitu „{starts}”"
    return found[0]


BODY = [
    "Powód:",  # punkt „Dlaczego nasz kandydat”
    "Obowiązek:",  # punkt obowiązków
    "Języki:",  # grupa umiejętności
    "Polski - C2",  # język
    "Technologie:",
]


@pytest.mark.parametrize("starts", BODY)
def test_body_text_is_justified_by_default(starts):
    doc = _render(_html())
    assert _align(_paragraph(doc, starts)) == WD_ALIGN_PARAGRAPH.JUSTIFY


def test_headings_stay_left_when_body_is_justified():
    doc = _render(_html())
    for starts in ("Tester QA", "DLACZEGO NASZ KANDYDAT", "DOŚWIADCZENIE"):
        assert _align(_paragraph(doc, starts)) in (None, WD_ALIGN_PARAGRAPH.LEFT)


@pytest.mark.parametrize("starts", BODY)
def test_left_rule_keeps_body_left(starts):
    doc = _render(_html(), text_align="left")
    assert _align(_paragraph(doc, starts)) in (None, WD_ALIGN_PARAGRAPH.LEFT)


@pytest.mark.parametrize("starts", BODY + ["Nazwa firmy:", "Stanowisko:", "12.2025"])
def test_explicit_left_from_editor_wins_over_default(starts):
    doc = _render(_html("text-align: left"))
    assert _align(_paragraph(doc, starts)) == WD_ALIGN_PARAGRAPH.LEFT


@pytest.mark.parametrize("starts", BODY + ["Nazwa firmy:", "Stanowisko:", "12.2025"])
def test_explicit_justify_from_editor_wins_over_left_rule(starts):
    doc = _render(_html("text-align: justify"), text_align="left")
    assert _align(_paragraph(doc, starts)) == WD_ALIGN_PARAGRAPH.JUSTIFY


def test_education_table_is_not_justified():
    doc = _render(_html())
    cells = [
        paragraph
        for table in doc.tables
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
        if paragraph.text.strip()
    ]
    assert cells
    assert all(_align(p) != WD_ALIGN_PARAGRAPH.JUSTIFY for p in cells)


@pytest.mark.parametrize(
    ("metadata", "expected"),
    [
        (None, "justify"),
        ({}, "justify"),
        ({"client_rule_snapshot": None}, "justify"),
        ({"client_rule_snapshot": {"highlight_policy": "must"}}, "justify"),
        ({"client_rule_snapshot": {"text_align": "left"}}, "left"),
        ({"client_rule_snapshot": {"text_align": "justify"}}, "justify"),
        ({"client_rule_snapshot": {"text_align": "center"}}, "justify"),
        ({"client_rule_snapshot": "zepsute"}, "justify"),
    ],
)
def test_text_align_comes_from_frozen_client_rule(metadata, expected):
    assert cv_text_align(metadata) == expected


# ── Plik z generatora (idzie do klienta przy zatwierdzeniu bez zmian) ───────


def _generator_payload(rule: dict | None) -> dict:
    payload = {
        "name": "Jan Testowy",
        "first_name": "Jan",
        "position": "Java Developer",
        "language": "pl",
        "why_points": [f"Powód: {LONG}"],
        "skills": [{"label": "Backend:", "content": "Java, Spring Boot"}],
        "languages": ["Polski – ojczysty"],
        "experience": [
            {
                "dates": "01.2020 – obecnie",
                "company": "Acme",
                "industry": "IT",
                "position": "Java Developer",
                "responsibilities": [f"Obowiązek: {LONG}"],
                "technologies": ["Java", "Spring Boot"],
            }
        ],
    }
    if rule is not None:
        payload["client_rule_snapshot"] = rule
    return payload


def _generator_doc(rule: dict | None) -> Document:
    from app.services.cv_generator_b2b.standalone_service import (
        rerender_docx_from_payload,
    )

    return Document(BytesIO(rerender_docx_from_payload(_generator_payload(rule))))


@pytest.mark.parametrize("rule", [None, {"text_align": "justify"}, {}])
@pytest.mark.parametrize("starts", ["Powód:", "Obowiązek:", "Backend:", "Technologie:"])
def test_generator_docx_justifies_body_by_default(rule, starts):
    assert (
        _align(_paragraph(_generator_doc(rule), starts)) == WD_ALIGN_PARAGRAPH.JUSTIFY
    )


@pytest.mark.parametrize("starts", ["Powód:", "Obowiązek:", "Backend:", "Technologie:"])
def test_generator_docx_follows_left_rule(starts):
    doc = _generator_doc({"text_align": "left"})
    assert _align(_paragraph(doc, starts)) in (None, WD_ALIGN_PARAGRAPH.LEFT)


def test_generator_title_is_not_justified():
    doc = _generator_doc(None)
    assert (
        _align(_paragraph(doc, "Java Developer – Jan Testowy"))
        != WD_ALIGN_PARAGRAPH.JUSTIFY
    )
