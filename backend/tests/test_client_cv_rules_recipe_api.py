"""API pełnej recepty reguły CV (migracja 0267): wersje, historia, kopiowanie,
sygnał zwrotny, lint, CV próbne, wymagane wejścia w generatorze.

Każdy test dowodzi jednego kontraktu, który łatwo cofnąć „przy okazji":

1. Zapis zmieniający treść bumpuje ``version`` i zostawia wpis z DIFFEM pól;
   samo zatwierdzenie tej samej treści wersji NIE zmienia (stempel na CV ma
   mówić o treści, nie o kliknięciach). Flagi karty klienta zapisują się na
   ``clients`` i wchodzą do diffu.
2. Kopia z innego klienta nadpisuje pola, NIE kopiuje flag klienta, NIE
   zatwierdza i zostawia wpis ``copied``.
3. Sygnał zwrotny liczy pominięte instrukcje per tekst z ostrzeżeń
   wygenerowanych CV i widzi stempel wersji.
4. Lint: kwota ``cv_rule_lint`` przed modelem, model zastąpiony atrapą.
5. CV próbne: rekrutacja innego klienta → 422; kwota generatora naliczona
   DWA razy; wiersz „processing" → job w tle (atrapa pipeline'u) → „ready"
   z dwiema wersjami i blokiem promptu.
6. Generator odmawia 422 z listą braków, gdy reguła wymaga stanowiska /
   numeru projektu, PRZED naliczeniem kwoty.
"""

from __future__ import annotations

import io
import uuid
from datetime import datetime, timezone

import pytest
from httpx import AsyncClient
from tests.cv_rule_requests import rule_request
from sqlalchemy import delete, select

RULE_URL = "/api/clients/{cid}/cv-rule"


async def _headers_for(
    app_client: AsyncClient,
    role_value: str,
    *,
    assigned_client_id: int | None = None,
) -> dict[str, str]:
    from app.core.database import AsyncSessionLocal
    from app.core.security import hash_password
    from app.models.team_structure import DeliveryLeadClientAssignment
    from app.models.user import User, UserRole

    email = f"recipe-{role_value}-{uuid.uuid4().hex[:8]}@example.com"
    password = f"P4ss_{uuid.uuid4().hex[:6]}!"
    async with AsyncSessionLocal() as db:
        role = UserRole(role_value)
        user = User(
            email=email,
            password_hash=hash_password(password),
            name=f"Recipe {role_value}",
            role=role,
            roles=[role.value],
            is_active=True,
            profile_completed=True,
        )
        db.add(user)
        await db.flush()
        if assigned_client_id is not None and role is UserRole.delivery_lead:
            db.add(
                DeliveryLeadClientAssignment(
                    delivery_lead_user_id=user.id,
                    client_id=assigned_client_id,
                )
            )
        await db.commit()
    login = await rule_request(
        app_client,
        "POST",
        "/api/auth/login",
        json={"email": email, "password": password},
    )
    assert login.status_code == 200, login.text
    return {"Authorization": f"Bearer {login.json()['access_token']}"}


async def _make_client(name: str) -> int:
    from app.core.database import AsyncSessionLocal
    from app.models.client import Client, ClientStatus

    async with AsyncSessionLocal() as db:
        row = Client(name=name, status=ClientStatus.active, hidden=False)
        db.add(row)
        await db.commit()
        await db.refresh(row)
        return row.id


async def _cleanup(client_ids: list[int]) -> None:
    from app.core.database import AsyncSessionLocal
    from app.models.client import Client
    from app.models.client_cv_rule import ClientCvRule
    from app.models.client_cv_rule_event import ClientCvRuleEvent
    from app.models.client_cv_rule_preview import ClientCvRulePreview
    from app.models.cv_generated_document import CvGeneratedDocument
    from app.models.team_structure import DeliveryLeadClientAssignment

    async with AsyncSessionLocal() as db:
        for model in (
            CvGeneratedDocument,
            ClientCvRulePreview,
            ClientCvRuleEvent,
            ClientCvRule,
        ):
            await db.execute(delete(model).where(model.client_id.in_(client_ids)))
        await db.execute(
            delete(DeliveryLeadClientAssignment).where(
                DeliveryLeadClientAssignment.client_id.in_(client_ids)
            )
        )
        await db.execute(delete(Client).where(Client.id.in_(client_ids)))
        await db.commit()


def _full_payload(**overrides) -> dict:
    payload = {
        "filename_pattern": "B2B_{STANOWISKO}_{IMIE_NAZWISKO}",
        "spaces_to_underscores": True,
        "cv_language": "pl",
        "requires_en_copy": False,
        "requires_rodo_consent_block": False,
        "notes": "Klient ceni bankowość.",
        "generator_instructions": "Bez sekcji zainteresowań.",
        "generator_instructions_en": None,
        "content_mode": "basic",
        "content_mode_locked": True,
        "require_screening_notes_min_chars": 300,
        "require_project_ref": True,
        "require_position": True,
        "require_champion": False,
        "auto_second_language": False,
        "omit_sections": ["languages", "education"],
        "max_roles": 5,
        "max_bullets_per_role": 3,
        "max_bullet_chars": 160,
        "why_points_max": 4,
        "date_format": "MM.YYYY",
        "glossary": [{"from": "Business Analyst", "to": "Analityk Biznesowy"}],
        "cv_content_mode_cap": "polished",
        "cv_interactive_enabled": False,
        "confirm": True,
    }
    payload.update(overrides)
    return payload


@pytest.mark.asyncio
async def test_save_bumps_version_only_on_content_change_and_records_diff(
    app_client: AsyncClient,
):
    cid = await _make_client(f"Recipe versions {uuid.uuid4().hex[:6]}")
    try:
        headers = await _headers_for(
            app_client, "delivery_lead", assigned_client_id=cid
        )

        r = await rule_request(
            app_client,
            "PUT",
            RULE_URL.format(cid=cid),
            json=_full_payload(),
            headers=headers,
        )
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["version"] == 1 and body["is_active"] is True
        assert body["omit_sections"] == ["education", "languages"]  # kolejność katalogu
        assert body["glossary"] == [
            {"from": "Business Analyst", "to": "Analityk Biznesowy"}
        ]
        assert body["content_mode"] == "basic" and body["content_mode_locked"] is True
        assert body["cv_content_mode_cap"] == "polished"
        assert body["cv_interactive_enabled"] is False
        assert "tryb „Przepisanie”" in body["client_policy"]
        assert "polityka prezentacji" in body["client_policy"]
        assert "notatka DL" in body["client_policy"]

        # Flagi klienta trafiły na `clients`.
        from app.core.database import AsyncSessionLocal
        from app.models.client import Client

        async with AsyncSessionLocal() as db:
            client = await db.get(Client, cid)
            assert client.cv_content_mode_cap == "polished"
            assert client.cv_interactive_enabled is False

        # Ta sama treść, tylko zatwierdzenie → wersja bez zmian.
        r = await rule_request(
            app_client,
            "PUT",
            RULE_URL.format(cid=cid),
            json=_full_payload(),
            headers=headers,
        )
        assert r.json()["version"] == 1

        # Zmiana treści → bump + diff.
        r = await rule_request(
            app_client,
            "PUT",
            RULE_URL.format(cid=cid),
            json=_full_payload(max_roles=3, cv_interactive_enabled=True),
            headers=headers,
        )
        assert r.json()["version"] == 2

        h = await app_client.get(RULE_URL.format(cid=cid) + "/history", headers=headers)
        assert h.status_code == 200, h.text
        events = h.json()
        assert [e["action"] for e in events][:3] == [
            "saved_and_confirmed",
            "saved_and_confirmed",
            "saved_and_confirmed",
        ]
        latest = events[0]
        assert latest["rule_version"] == 2
        assert latest["changes"]["max_roles"] == {"from": 5, "to": 3}
        assert latest["changes"]["cv_interactive_enabled"] == {
            "from": False,
            "to": True,
        }
        assert "notes" not in latest["changes"]
        assert events[0]["actor_name"] == "Recipe delivery_lead"

        # Usunięcie i ponowne założenie NIE restartuje numeracji — stempel
        # `client_rule_version` na CV ma pozostać jednoznaczny.
        r = await rule_request(
            app_client, "DELETE", RULE_URL.format(cid=cid), headers=headers
        )
        assert r.status_code == 204
        r = await rule_request(
            app_client,
            "PUT",
            RULE_URL.format(cid=cid),
            json=_full_payload(),
            headers=headers,
        )
        assert r.status_code == 200, r.text
        assert r.json()["version"] == 3

        # TAC bez dostępu Delivery nie może ani zapisać, ani czytać historii.
        tac = await _headers_for(app_client, "tac")
        r = await rule_request(
            app_client,
            "PUT",
            RULE_URL.format(cid=cid),
            json=_full_payload(),
            headers=tac,
        )
        assert r.status_code == 403
        r = await app_client.get(RULE_URL.format(cid=cid) + "/history", headers=tac)
        assert r.status_code == 403
    finally:
        await _cleanup([cid])


@pytest.mark.asyncio
async def test_copy_from_another_client_is_a_proposal_without_client_flags(
    app_client: AsyncClient,
):
    tag = uuid.uuid4().hex[:6]
    source = await _make_client(f"Recipe source {tag}")
    target = await _make_client(f"Recipe target {tag}")
    try:
        headers = await _headers_for(app_client, "admin")
        r = await rule_request(
            app_client,
            "PUT",
            RULE_URL.format(cid=source),
            json=_full_payload(),
            headers=headers,
        )
        assert r.status_code == 200, r.text

        r = await rule_request(
            app_client,
            "POST",
            RULE_URL.format(cid=target) + f"/copy-from/{source}",
            headers=headers,
        )
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["filename_pattern"] == "B2B_{STANOWISKO}_{IMIE_NAZWISKO}"
        assert body["max_roles"] == 5 and body["content_mode_locked"] is True
        assert body["is_active"] is False, "kopia jest PROPOZYCJĄ"
        assert body["version"] == 1
        assert body["cv_content_mode_cap"] is None, "flagi klienta nie są kopiowane"
        assert body["cv_interactive_enabled"] is True

        h = await app_client.get(
            RULE_URL.format(cid=target) + "/history", headers=headers
        )
        assert h.json()[0]["action"] == "copied"
        assert h.json()[0]["changes"]["source_client_id"] == source

        r = await rule_request(
            app_client,
            "POST",
            RULE_URL.format(cid=target) + f"/copy-from/{target}",
            headers=headers,
        )
        assert r.status_code == 422
    finally:
        await _cleanup([source, target])


@pytest.mark.asyncio
async def test_feedback_counts_skipped_instructions_from_generated_warnings(
    app_client: AsyncClient,
):
    cid = await _make_client(f"Recipe feedback {uuid.uuid4().hex[:6]}")
    try:
        headers = await _headers_for(
            app_client, "delivery_lead", assigned_client_id=cid
        )
        from app.core.database import AsyncSessionLocal
        from app.models.cv_generated_document import CvGeneratedDocument

        async with AsyncSessionLocal() as db:
            for warnings, version in (
                (["Pominięto instrukcję klienta: Dopisz Kubernetes"], 2),
                (
                    [
                        "Skipped client instruction: Dopisz Kubernetes",
                        "WERYFIKUJ: domknięto politykę prezentacji klienta w kodzie: x.",
                    ],
                    2,
                ),
                (["WERYFIKUJ: coś innego"], None),
            ):
                db.add(
                    CvGeneratedDocument(
                        client_id=cid,
                        candidate_name="Jan Testowy",
                        language="pl",
                        blind=False,
                        mode="upload",
                        content_mode="polished",
                        filename="x.docx",
                        status="ready",
                        render_payload={"name": "Jan"},
                        warnings=warnings,
                        client_rule_version=version,
                    )
                )
            await db.commit()

        r = await app_client.get(
            RULE_URL.format(cid=cid) + "/feedback", headers=headers
        )
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["generated_total"] == 3
        assert body["with_skipped_instructions"] == 2
        assert body["with_policy_enforced"] == 1
        assert body["skipped_by_instruction"] == [
            {"text": "Dopisz Kubernetes", "count": 2}
        ]
        assert {g["client_rule_version"] for g in body["recent"]} == {2, None}
    finally:
        await _cleanup([cid])


@pytest.mark.asyncio
async def test_lint_charges_quota_then_returns_per_line_verdicts(
    app_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
):
    from app.api import client_cv_rules as api_module
    from app.services.cv_generator_b2b import rule_lint

    calls: list[str] = []

    def fake_lint(text: str, *, request_id: str):
        calls.append(text)
        lines = rule_lint.split_instruction_lines(text)
        return [
            rule_lint.LintFinding(
                index=i,
                line=line,
                verdict="adds_facts" if "Dopisz" in line else "ok",
                reason="atrapa",
                suggestion="" if "Dopisz" not in line else "Umieść wyżej, jeśli jest.",
            )
            for i, line in enumerate(lines)
        ]

    monkeypatch.setattr(rule_lint, "lint_instructions", fake_lint)

    charged: list[str] = []

    async def fake_charge(db, feature, user_id=None):
        charged.append(feature.value)

    monkeypatch.setattr(api_module, "check_and_increment", fake_charge)

    cid = await _make_client(f"Recipe lint {uuid.uuid4().hex[:6]}")
    try:
        headers = await _headers_for(
            app_client, "delivery_lead", assigned_client_id=cid
        )
        r = await rule_request(
            app_client,
            "POST",
            RULE_URL.format(cid=cid) + "/lint",
            json={
                "generator_instructions": "Bez sekcji zainteresowań. Dopisz Kubernetes.",
                "generator_instructions_en": None,
                "notes": "Klient ceni bankowość.",
            },
            headers=headers,
        )
        assert r.status_code == 200, r.text
        body = r.json()
        assert charged == ["cv_rule_lint", "cv_rule_lint"], (
            "jedno pole = jedno obciążenie"
        )
        assert len(calls) == 2  # instrukcje + notatka, osobno
        assert body["adds_facts_count"] == 1 and body["ok_count"] == 2
        bad = next(f for f in body["findings"] if f["verdict"] == "adds_facts")
        assert bad["field"] == "generator_instructions"
        assert bad["line"] == "Dopisz Kubernetes."
        assert bad["suggestion"]

        # Pusty lint nie kosztuje nic.
        charged.clear()
        r = await rule_request(
            app_client,
            "POST",
            RULE_URL.format(cid=cid) + "/lint",
            json={
                "generator_instructions": "",
                "generator_instructions_en": None,
                "notes": None,
            },
            headers=headers,
        )
        assert r.status_code == 200 and r.json()["findings"] == []
        assert charged == []
    finally:
        await _cleanup([cid])


@pytest.mark.asyncio
async def test_preview_rejects_foreign_recruitment_and_runs_both_variants(
    app_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
):
    from app.api import client_cv_rules as api_module
    from app.core.database import AsyncSessionLocal
    from app.models.candidate import Candidate
    from app.models.job import Job
    from app.models.recruitment_pipeline import CandidateStage

    tag = uuid.uuid4().hex[:6]
    cid = await _make_client(f"Recipe preview {tag}")
    other = await _make_client(f"Recipe preview other {tag}")
    candidate_id = job_id = stage_id = other_stage_id = None
    try:
        async with AsyncSessionLocal() as db:
            cand = Candidate(
                name="Jan", lastname=f"Próbny {tag}", email=f"probny-{tag}@example.com"
            )
            db.add(cand)
            await db.flush()
            job = Job(title="Analityk", client_id=cid)
            other_job = Job(title="Analityk u innego", client_id=other)
            db.add_all([job, other_job])
            await db.flush()
            stage = CandidateStage(
                candidate_id=cand.id, job_id=job.id, stage="verified"
            )
            other_stage = CandidateStage(
                candidate_id=cand.id, job_id=other_job.id, stage="verified"
            )
            db.add_all([stage, other_stage])
            await db.commit()
            candidate_id, job_id, stage_id, other_stage_id = (
                cand.id,
                job.id,
                stage.id,
                other_stage.id,
            )

        headers = await _headers_for(
            app_client, "delivery_lead", assigned_client_id=cid
        )
        r = await rule_request(
            app_client,
            "PUT",
            RULE_URL.format(cid=cid),
            json=_full_payload(),
            headers=headers,
        )
        assert r.status_code == 200, r.text

        charged: list[tuple[str, int]] = []
        remaining = 1

        async def fake_charge(db, feature, user_id=None, *, units=1):
            nonlocal remaining
            from app.services.ai_quota import AIQuotaExceeded

            if remaining < units:
                raise AIQuotaExceeded(feature, "Limit", used=9, limit=10)
            remaining -= units
            charged.append((feature.value, units))

        monkeypatch.setattr(api_module, "check_and_increment", fake_charge)

        seen_rules: list[object] = []

        async def fake_generate(source, *, language="pl", **kw):
            seen_rules.append(kw.get("client_rule"))
            from app.services.cv_generator_b2b.standalone_service import (
                GenerationResult,
            )

            return GenerationResult(
                candidate_name="Jan Próbny",
                filename="x.docx",
                docx_bytes=b"DOCX",
                warnings=["w"] if kw.get("client_rule") else [],
                processing_time_ms=1,
                render_payload={"name": "Jan", "why_points": ["a"]},
                job_id=stage_id,
            )

        from unittest.mock import AsyncMock

        from types import SimpleNamespace

        from io import BytesIO
        from docx import Document

        document = Document()
        document.add_paragraph("Synthetic candidate CV with source experience.")
        buffer = BytesIO()
        document.save(buffer)
        from app.services.cv_generator_b2b.standalone_service import (
            CandidateGenerationSource,
        )

        frozen_source = CandidateGenerationSource(
            cv_bytes=buffer.getvalue(),
            cv_filename="cv.docx",
            screening_notes_text="notes",
            champion_json="{}",
            has_champion=True,
            source_warnings=(),
            fallback_name="Synthetic",
            job_id=job_id,
            job_title="Developer",
            client_content_mode_cap=None,
            cv_document_id=None,
            candidate_id=candidate_id,
            stage_id=stage_id,
            client_id=cid,
        )
        monkeypatch.setattr(
            api_module, "prepare_source_facts", lambda **kwargs: object()
        )
        from app.services import object_storage

        stored_inputs = {}

        def upload_snapshot(raw, filename, content_type):
            key = f"test-only/{uuid.uuid4().hex}"
            stored_inputs[key] = raw
            return key

        monkeypatch.setattr(object_storage, "upload_cv", upload_snapshot)
        monkeypatch.setattr(
            object_storage, "download_cv", lambda key: stored_inputs[key]
        )
        load_source = AsyncMock(return_value=frozen_source)
        monkeypatch.setattr(api_module, "load_candidate_generation_source", load_source)
        monkeypatch.setattr(
            api_module, "generate_cv_from_candidate_source", fake_generate
        )

        # Rekrutacja bez CV / Championa / notatek → 422 PRZED kwotą (prawdziwa
        # gotowość: kandydat testowy nie ma nic z tych trzech).
        r = await rule_request(
            app_client,
            "POST",
            RULE_URL.format(cid=cid) + "/preview",
            json={"candidate_id": candidate_id, "stage_id": stage_id, "language": "pl"},
            headers=headers,
        )
        assert r.status_code == 422, r.text
        assert "nie jest gotowa" in r.json()["detail"]
        assert charged == []


        async def fake_readiness(db, candidate_id, **kwargs):
            return [
                SimpleNamespace(
                    stage_id=stage_id,
                    ready=True,
                    has_cv=True,
                    has_champion=True,
                    has_notes=True,
                ),
                SimpleNamespace(
                    stage_id=other_stage_id,
                    ready=True,
                    has_cv=True,
                    has_champion=True,
                    has_notes=True,
                ),
            ]

        monkeypatch.setattr(
            api_module, "list_recruitments_with_readiness", fake_readiness
        )

        # Rekrutacja u INNEGO klienta → 422, bez naliczania kwoty.
        r = await rule_request(
            app_client,
            "POST",
            RULE_URL.format(cid=cid) + "/preview",
            json={
                "candidate_id": candidate_id,
                "stage_id": other_stage_id,
                "language": "pl",
            },
            headers=headers,
        )
        assert r.status_code == 422, r.text
        assert charged == []

        r = await rule_request(
            app_client,
            "POST",
            RULE_URL.format(cid=cid) + "/preview",
            json={"candidate_id": candidate_id, "stage_id": stage_id, "language": "pl"},
            headers=headers,
        )
        assert r.status_code == 503, r.text
        assert charged == [], "one remaining unit must not be consumed"
        assert remaining == 1
        assert seen_rules == [], "neither variant may run after denied admission"

        remaining = 2
        r = await rule_request(
            app_client,
            "POST",
            RULE_URL.format(cid=cid) + "/preview",
            json={"candidate_id": candidate_id, "stage_id": stage_id, "language": "pl"},
            headers=headers,
        )
        assert r.status_code == 202, r.text
        assert charged == [("cv_generator", 2)], (
            "both preview variants must have one admission for two units"
        )
        preview_id = r.json()["id"]

        # Zadanie w tle wykonało się po odpowiedzi (BackgroundTasks w ASGI transport).
        g = await app_client.get(
            RULE_URL.format(cid=cid) + f"/preview/{preview_id}", headers=headers
        )
        assert g.status_code == 200, g.text
        body = g.json()
        assert body["status"] == "ready", body
        assert body["with_rule"]["payload"]["name"] == "Jan"
        assert body["with_rule"]["warnings"] == ["w"]
        assert body["without_rule"]["warnings"] == []
        assert "<client_presentation_rules>" in body["prompt_block"]
        assert seen_rules[0] is not None and seen_rules[1] is None

        # Każdy DL może otworzyć klienta, ale ID podglądu nie przechodzi między
        # klientami i nie ujawnia, do którego z nich naprawdę należy.
        g = await app_client.get(
            RULE_URL.format(cid=other) + f"/preview/{preview_id}", headers=headers
        )
        assert g.status_code == 404
    finally:
        async with AsyncSessionLocal() as db:
            if stage_id:
                await db.execute(
                    delete(CandidateStage).where(
                        CandidateStage.id.in_([stage_id, other_stage_id])
                    )
                )
            if job_id:
                await db.execute(delete(Job).where(Job.client_id.in_([cid, other])))
            if candidate_id:
                await db.execute(delete(Candidate).where(Candidate.id == candidate_id))
            await db.commit()
        await _cleanup([cid, other])


@pytest.mark.asyncio
async def test_generate_upload_refuses_missing_required_inputs_before_quota(
    app_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
):
    from app.api import cv_generator_b2b as gen_module

    charged: list[str] = []

    async def fake_charge(db, user_id):
        charged.append("cv_generator")

    monkeypatch.setattr(gen_module, "_charge_cv_generation_quota", fake_charge)

    cid = await _make_client(f"Recipe required {uuid.uuid4().hex[:6]}")
    try:
        dl = await _headers_for(app_client, "delivery_lead", assigned_client_id=cid)
        r = await rule_request(
            app_client,
            "PUT",
            RULE_URL.format(cid=cid),
            json=_full_payload(require_screening_notes_min_chars=None),
            headers=dl,
        )
        assert r.status_code == 200, r.text

        recruiter = await _headers_for(app_client, "recruiter")
        r = await rule_request(
            app_client,
            "POST",
            "/api/cv-generator/generate-upload",
            data={"client_id": str(cid), "language": "pl", "content_mode": "tailored"},
            files={"cv_file": ("cv.pdf", io.BytesIO(b"%PDF-1.4 x"), "application/pdf")},
            headers=recruiter,
        )
        assert r.status_code == 422, r.text
        detail = r.json()["detail"]
        assert "Numer / nazwa projektu" in detail
        assert "Stanowisko" in detail
        assert charged == [], "odmowa nie może kosztować kwoty"
    finally:
        await _cleanup([cid])


@pytest.mark.asyncio
async def test_unconfirmed_recipe_is_invisible_to_the_generator():
    """Blokady i polityka też czekają na zatwierdzenie — inaczej propozycja
    z kopii u innego klienta zaczęłaby blokować rekruterów."""
    from app.core.database import AsyncSessionLocal
    from app.models.client import Client, ClientStatus
    from app.models.client_cv_rule import ClientCvRule
    from app.services.cv_generator_b2b.client_rules import resolve_client_rule

    async with AsyncSessionLocal() as db:
        client = Client(
            name=f"Recipe hidden {uuid.uuid4().hex[:6]}", status=ClientStatus.active
        )
        db.add(client)
        await db.flush()
        db.add(
            ClientCvRule(
                client_id=client.id,
                content_mode="basic",
                content_mode_locked=True,
                require_position=True,
                confirmed_at=None,
            )
        )
        await db.commit()
        cid = client.id
    try:
        async with AsyncSessionLocal() as db:
            assert await resolve_client_rule(db, cid) is None
            rule = (
                await db.execute(
                    select(ClientCvRule).where(ClientCvRule.client_id == cid)
                )
            ).scalar_one()
            rule.confirmed_at = datetime.now(timezone.utc)
            await db.commit()
        async with AsyncSessionLocal() as db:
            found = await resolve_client_rule(db, cid)
            assert found is not None and found.content_mode_locked is True
    finally:
        await _cleanup([cid])
