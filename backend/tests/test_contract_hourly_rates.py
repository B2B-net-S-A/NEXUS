"""Stawki w module Kontrakty wyłącznie w zł/h (ticket 14.09.2026).

Kryteria akceptacji, każde osobnym testem:

* kontrakt w MD → zł/h (MD ÷ 8) dla stawki kosztowej i przychodowej, bez
  zaokrąglenia zniekształcającego kwotę; miesięczne ekwiwalenty bez zmian
  (176 h/mc = 22 MD × 8 h);
* kontrakty godzinowe i ryczałtowe (decyzja Artura) — bez zmian;
* zamówienia — bez zmian, co do wartości i jednostki;
* nowe zamówienie w MD daje w Kontraktach stawkę godzinową;
* eksport z Kontraktów nie ma już jednostki „dzienna".
"""

from __future__ import annotations

import io
import uuid
from datetime import date
from decimal import Decimal

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import delete, select

from app.models.activity import Activity
from app.models.client_order import ClientOrder, ClientOrderStatus
from app.models.contract import Contract, ContractStatus, ContractType, RateUnit
from app.models.contract_candidate_rate import ContractCandidateRate
from app.models.contract_client_rate import ContractClientRate
from app.models.contract_framework_rate import ContractFrameworkRate
from app.services.contract_order_sync import (
    MD_BILLING_HOURS_PER_MONTH,
    apply_contract_hourly_policy,
    contract_unit_for_order,
    sync_contract_from_orders,
    sync_orders_cost_from_contract,
)

pytestmark = pytest.mark.asyncio

TODAY = date(2026, 9, 20)


class _FakeDb:
    def __init__(self) -> None:
        self.added: list[object] = []

    def add(self, obj: object) -> None:
        self.added.append(obj)


def _contract(**overrides) -> Contract:
    contract = Contract(
        id=overrides.pop("id", 1),
        client_id=10,
        candidate_id=5,
        contract_type=ContractType.b2b,
        status=overrides.pop("status", ContractStatus.active),
        start_date=date(2026, 1, 1),
        end_date=None,
        rate_candidate=overrides.pop("rate_candidate", None),
        rate_client=overrides.pop("rate_client", None),
        framework_rate=overrides.pop("framework_rate", None),
        rate_unit=overrides.pop("rate_unit", RateUnit.hourly),
        billing_hours_per_month=overrides.pop("billing_hours_per_month", 160),
        currency="PLN",
        rate_candidate_currency="PLN",
        rate_client_currency="PLN",
    )
    contract.candidate_rate_schedule = overrides.pop("candidate_rate_schedule", [])
    contract.client_rate_schedule = overrides.pop("client_rate_schedule", [])
    contract.framework_rate_schedule = overrides.pop("framework_rate_schedule", [])
    for key, value in overrides.items():
        setattr(contract, key, value)
    return contract


def _order(order_id: int = 100, **overrides) -> ClientOrder:
    values = {
        "id": order_id,
        "client_id": 10,
        "contract_id": 1,
        "title": "ZAM/1",
        "status": ClientOrderStatus.active,
        "start_date": date(2026, 9, 1),
        "end_date": date(2026, 12, 31),
        "rate_client": Decimal("1340.000"),
        "rate_candidate": None,
        "rate_unit": RateUnit.daily,
        "billing_hours_per_month": 160,
        "rate_client_currency": "PLN",
        "rate_candidate_currency": "PLN",
        "currency": "PLN",
        "order_group_id": None,
    }
    values.update(overrides)
    return ClientOrder(**values)


# ── Reguła jednostki ────────────────────────────────────────────────────────


def test_md_order_maps_to_an_hourly_contract_and_nothing_else_changes():
    assert contract_unit_for_order(RateUnit.daily) == RateUnit.hourly
    assert contract_unit_for_order(RateUnit.hourly) == RateUnit.hourly
    assert contract_unit_for_order(RateUnit.monthly) == RateUnit.monthly


def test_daily_contract_is_converted_exactly_with_all_its_amounts():
    """Stawki, harmonogramy, stawka ramowa, widełki — ÷ 8, bez zaokrąglenia."""
    contract = _contract(
        rate_unit=RateUnit.daily,
        rate_candidate=Decimal("1001.550"),
        rate_client=Decimal("1340.000"),
        framework_rate=Decimal("1500.10"),
        target_rate_min=Decimal("1200.00"),
        target_rate_max=Decimal("1400.00"),
        candidate_rate_schedule=[
            ContractCandidateRate(rate=Decimal("1001.550"), effective_from=date(2026, 1, 1))
        ],
        client_rate_schedule=[
            ContractClientRate(rate=Decimal("1340.000"), effective_from=date(2026, 1, 1))
        ],
        framework_rate_schedule=[
            ContractFrameworkRate(rate=Decimal("1500.10"), effective_from=date(2026, 1, 1))
        ],
    )
    monthly_before = contract.monthly_rate(contract.rate_client)

    assert apply_contract_hourly_policy(contract)

    assert contract.rate_unit == RateUnit.hourly
    assert contract.billing_hours_per_month == MD_BILLING_HOURS_PER_MONTH == 176
    assert contract.rate_candidate == Decimal("125.19375")
    assert contract.rate_client == Decimal("167.5")
    assert contract.framework_rate == Decimal("187.5125")
    assert (contract.target_rate_min, contract.target_rate_max) == (
        Decimal("150"),
        Decimal("175"),
    )
    assert contract.candidate_rate_schedule[0].rate == Decimal("125.19375")
    assert contract.client_rate_schedule[0].rate == Decimal("167.5")
    assert contract.framework_rate_schedule[0].rate == Decimal("187.5125")
    # Prawdziwość kwot: × 8 daje dawną stawkę, miesiąc liczy się tak samo.
    assert contract.rate_candidate * 8 == Decimal("1001.55")
    assert contract.monthly_rate(contract.rate_client) == monthly_before


@pytest.mark.parametrize("unit", [RateUnit.hourly, RateUnit.monthly])
def test_hourly_and_monthly_contracts_are_left_alone(unit):
    contract = _contract(
        rate_unit=unit, rate_candidate=Decimal("120"), rate_client=Decimal("150")
    )

    assert not apply_contract_hourly_policy(contract)

    assert contract.rate_unit == unit
    assert (contract.rate_candidate, contract.rate_client) == (
        Decimal("120"),
        Decimal("150"),
    )
    assert contract.billing_hours_per_month == 160


# ── Synchronizacja z zamówieniami ───────────────────────────────────────────


async def test_md_order_on_a_monthly_contract_gives_hourly_with_the_same_month():
    contract = _contract(
        rate_unit=RateUnit.monthly,
        rate_candidate=Decimal("17600"),
    )

    outcome = await sync_contract_from_orders(
        _FakeDb(), contract, [_order()], actor_id=None, today=TODAY
    )

    assert (outcome.unit_switched_from, outcome.unit_switched_to) == (
        "monthly",
        "hourly",
    )
    assert contract.billing_hours_per_month == 176
    assert contract.rate_candidate == Decimal("100"), "17 600 zł/mc ÷ 176 h"
    assert contract.monthly_rate(contract.rate_candidate) == Decimal("17600")
    assert contract.rate_client == Decimal("167.5")


async def test_monthly_order_keeps_the_monthly_contract():
    contract = _contract(rate_unit=RateUnit.monthly, rate_candidate=Decimal("17600"))
    order = _order(rate_unit=RateUnit.monthly, rate_client=Decimal("22000"))

    await sync_contract_from_orders(
        _FakeDb(), contract, [order], actor_id=None, today=TODAY
    )

    assert contract.rate_unit == RateUnit.monthly
    assert contract.rate_client == Decimal("22000")


async def test_first_md_order_gives_a_new_hourly_contract_176_hours():
    """Kontrakt bez przychodu (świeża umowa B2B) — pierwsze zamówienie w MD
    ustala jego pieniądze, więc miesiąc liczy się jak w zamówieniu (22 MD)."""
    contract = _contract(rate_candidate=Decimal("120"), billing_hours_per_month=168)

    outcome = await sync_contract_from_orders(
        _FakeDb(), contract, [_order()], actor_id=None, today=TODAY
    )
    assert contract.billing_hours_per_month == 176
    assert outcome.billing_hours_from == 168
    assert outcome.as_details()["billing_hours_per_month"] == {"from": 168}

    explicit = _contract(rate_candidate=Decimal("120"), billing_hours_per_month=168)
    await sync_contract_from_orders(
        _FakeDb(),
        explicit,
        [_order()],
        actor_id=None,
        today=TODAY,
        follow_order_unit=False,
    )
    assert explicit.billing_hours_per_month == 168
    assert explicit.rate_client == Decimal("167.5")


async def test_hourly_contract_with_revenue_keeps_its_hours_on_an_md_order():
    """Ticket 1a: kontrakt już godzinowy pozostaje bez zmian — także godziny."""
    contract = _contract(
        rate_candidate=Decimal("120"),
        rate_client=Decimal("150"),
        billing_hours_per_month=160,
    )

    await sync_contract_from_orders(
        _FakeDb(), contract, [_order()], actor_id=None, today=TODAY
    )

    assert contract.rate_unit == RateUnit.hourly
    assert contract.billing_hours_per_month == 160
    assert contract.rate_candidate == Decimal("120")
    assert contract.effective_client_rate(TODAY) == Decimal("167.5")


async def test_176_hours_follow_the_newest_order_back_to_hourly():
    contract = _contract(
        rate_candidate=Decimal("120"), billing_hours_per_month=176
    )
    hourly = _order(
        101,
        rate_unit=RateUnit.hourly,
        rate_client=Decimal("170"),
        start_date=date(2027, 1, 1),
        billing_hours_per_month=168,
    )

    await sync_contract_from_orders(
        _FakeDb(), contract, [_order(), hourly], actor_id=None, today=TODAY
    )

    assert contract.billing_hours_per_month == 168
    assert contract.rate_candidate == Decimal("120")


async def test_monthly_order_uses_the_contract_hours_both_ways():
    """Kontrakt 125 zł/h × 176 h = 22 000 zł/mc — zamówienie miesięczne 22 000
    nie może stać się 20 000 (koszt) ani krokiem 137,5 zł/h (przychód)."""
    contract = _contract(
        rate_candidate=Decimal("125"),
        rate_client=None,
        billing_hours_per_month=176,
    )
    monthly = _order(
        rate_unit=RateUnit.monthly,
        rate_client=Decimal("22000"),
        rate_candidate=Decimal("22000"),
        billing_hours_per_month=160,
    )

    await sync_contract_from_orders(
        _FakeDb(),
        contract,
        [monthly],
        actor_id=None,
        today=TODAY,
        follow_order_unit=False,
    )
    changed = await sync_orders_cost_from_contract(
        _FakeDb(), contract, [monthly], today=TODAY
    )

    assert contract.effective_client_rate(TODAY) == Decimal("125")
    assert changed == []
    assert monthly.rate_candidate == Decimal("22000")


def test_orders_inheriting_from_a_converted_contract_stay_in_md():
    from app.services.order_rate_snapshots import (
        inherited_order_rate_fields,
        order_unit_for_contract,
    )

    converted = _contract(
        rate_candidate=Decimal("125.19375"),
        rate_client=Decimal("167.5"),
        billing_hours_per_month=176,
    )
    b2b = _contract(rate_candidate=Decimal("120"), billing_hours_per_month=160)

    assert order_unit_for_contract(converted) == RateUnit.daily
    assert order_unit_for_contract(b2b) == RateUnit.hourly
    fields = inherited_order_rate_fields(converted)
    assert fields["rate_unit"] == RateUnit.daily
    assert fields["rate_candidate"] == Decimal("1001.550")
    assert fields["rate_client"] == Decimal("1340.000")
    assert inherited_order_rate_fields(b2b)["rate_unit"] == RateUnit.hourly


def test_order_margin_falls_back_to_the_contract_rate_in_the_order_unit():
    from app.api.client_orders import _compute_monthly_margin

    contract = _contract(
        rate_candidate=Decimal("120"),
        rate_client=Decimal("167.5"),
        billing_hours_per_month=176,
    )
    order = _order(rate_client=Decimal("1340"), rate_candidate=None)

    margin = _compute_monthly_margin(order, contract, on=TODAY)

    assert margin == (Decimal("1340") - Decimal("960")) * 22


async def test_second_sync_changes_nothing():
    contract = _contract(rate_candidate=Decimal("120"))
    orders = [_order()]
    await sync_contract_from_orders(
        _FakeDb(), contract, orders, actor_id=None, today=TODAY
    )

    again = await sync_contract_from_orders(
        _FakeDb(), contract, orders, actor_id=None, today=TODAY
    )

    assert not again.changed


async def test_fractional_hourly_cost_returns_to_the_md_order_unchanged():
    """1001,55 zł/MD → 125,19375 zł/h → z powrotem 1001,55 zł/MD w zamówieniu."""
    contract = _contract(rate_candidate=Decimal("125.19375"))
    order = _order(rate_candidate=Decimal("1001.550"))

    changed = await sync_orders_cost_from_contract(
        _FakeDb(), contract, [order], today=TODAY
    )

    assert changed == []
    assert order.rate_candidate == Decimal("1001.550")
    assert order.rate_unit == RateUnit.daily


# ── Integracja: API ─────────────────────────────────────────────────────────


async def _seed_people(suffix: str | None = None) -> tuple[int, int]:
    from app.core.database import AsyncSessionLocal
    from app.models.candidate import Candidate
    from app.models.client import Client

    suffix = suffix or uuid.uuid4().hex[:6]
    async with AsyncSessionLocal() as db:
        client = Client(name=f"Godzinowy Klient {suffix}")
        candidate = Candidate(
            name="Godzina",
            lastname=f"Stawka-{suffix}",
            email=f"hourly-{suffix}@example.com",
        )
        db.add_all([client, candidate])
        await db.commit()
        return client.id, candidate.id


async def _load(contract_id: int) -> Contract:
    from app.core.database import AsyncSessionLocal
    from app.services.contract_rates import RATE_SCHEDULE_LOADS

    async with AsyncSessionLocal() as db:
        return await db.scalar(
            select(Contract)
            .where(Contract.id == contract_id)
            .options(*RATE_SCHEDULE_LOADS)
        )


async def _seed_contract(client_id: int, candidate_id: int, **values) -> int:
    from app.core.database import AsyncSessionLocal

    async with AsyncSessionLocal() as db:
        contract = Contract(
            candidate_id=candidate_id,
            client_id=client_id,
            contract_type=ContractType.b2b,
            status=values.pop("status", ContractStatus.active),
            start_date=date(2026, 1, 5),
            currency="PLN",
            rate_candidate_currency="PLN",
            rate_client_currency="PLN",
            **values,
        )
        db.add(contract)
        await db.commit()
        return contract.id


async def test_new_contract_entered_in_md_is_saved_hourly(
    app_client: AsyncClient, app_auth_headers: dict
):
    client_id, candidate_id = await _seed_people()

    resp = await app_client.post(
        "/api/contracts",
        json={
            "candidate_id": candidate_id,
            "client_id": client_id,
            "start_date": "2026-09-01",
            "rate_unit": "daily",
            "rate_candidate": 960,
            "rate_client": 1340,
        },
        headers=app_auth_headers,
    )

    assert resp.status_code in (200, 201), resp.text
    body = resp.json()
    assert body["rate_unit"] == "hourly"
    assert body["billing_hours_per_month"] == 176
    assert Decimal(str(body["rate_candidate"])) == Decimal("120")
    assert Decimal(str(body["rate_client"])) == Decimal("167.5")


async def test_switching_an_hourly_contract_to_md_is_refused(
    app_client: AsyncClient, app_auth_headers: dict
):
    client_id, candidate_id = await _seed_people()
    contract_id = await _seed_contract(
        client_id,
        candidate_id,
        rate_unit=RateUnit.hourly,
        rate_candidate=Decimal("120"),
        rate_client=Decimal("150"),
    )

    resp = await app_client.patch(
        f"/api/contracts/{contract_id}",
        json={"rate_unit": "daily", "rate_candidate": 960},
        headers=app_auth_headers,
    )

    assert resp.status_code == 422, resp.text
    assert resp.json()["detail"]["reason"] == "contract_rates_are_hourly"
    loaded = await _load(contract_id)
    assert (loaded.rate_unit, loaded.rate_candidate) == (RateUnit.hourly, Decimal("120"))


async def test_relabelling_a_legacy_md_contract_to_hours_is_refused(
    app_client: AsyncClient, app_auth_headers: dict
):
    """Samo przestawienie etykiety MD → godziny zostawiłoby kroki w MD."""
    client_id, candidate_id = await _seed_people()
    contract_id = await _seed_contract(
        client_id,
        candidate_id,
        rate_unit=RateUnit.daily,
        rate_candidate=Decimal("960"),
    )

    resp = await app_client.patch(
        f"/api/contracts/{contract_id}",
        json={"rate_unit": "hourly", "rate_candidate": 960},
        headers=app_auth_headers,
    )

    assert resp.status_code == 422, resp.text
    loaded = await _load(contract_id)
    assert (loaded.rate_unit, loaded.rate_candidate) == (RateUnit.daily, Decimal("960"))


async def test_saving_a_legacy_md_contract_converts_it_whole(
    app_client: AsyncClient, app_auth_headers: dict
):
    """Kontrakt sprzed korekty: formularz odsyła MD — zapis przelicza całość."""
    client_id, candidate_id = await _seed_people()
    contract_id = await _seed_contract(
        client_id,
        candidate_id,
        rate_unit=RateUnit.daily,
        rate_candidate=Decimal("960"),
        rate_client=Decimal("1340"),
    )

    resp = await app_client.patch(
        f"/api/contracts/{contract_id}",
        json={
            "rate_unit": "daily",
            "rate_candidate": 960,
            "rate_client": 1340,
            "project_name": "Po zmianie",
        },
        headers=app_auth_headers,
    )

    assert resp.status_code == 200, resp.text
    loaded = await _load(contract_id)
    assert loaded.rate_unit == RateUnit.hourly
    assert loaded.billing_hours_per_month == 176
    assert (loaded.rate_candidate, loaded.rate_client) == (
        Decimal("120"),
        Decimal("167.5"),
    )


async def test_new_contractor_md_order_keeps_md_and_contract_is_hourly(
    app_client: AsyncClient, app_auth_headers: dict
):
    """Zamówienie w module Klienci zostaje w MD, kontrakt dostaje zł/h."""
    client_id, candidate_id = await _seed_people()

    resp = await app_client.post(
        f"/api/clients/{client_id}/contract-with-order",
        headers=app_auth_headers,
        json={
            "candidate_id": candidate_id,
            "title": f"MD-{uuid.uuid4().hex[:6]}",
            "contract_start_date": "2026-09-01",
            "order_start_date": "2026-09-01",
            "order_end_date": "2026-12-31",
            "rate_client": "1340.000",
            "rate_candidate": "960.000",
            "rate_unit": "daily",
        },
    )

    assert resp.status_code == 201, resp.text
    ids = resp.json()
    contract = await _load(ids["contract_id"])
    assert contract.rate_unit == RateUnit.hourly
    assert contract.effective_candidate_rate(date(2026, 10, 1)) == Decimal("120")
    assert contract.effective_client_rate(date(2026, 10, 1)) == Decimal("167.5")

    from app.core.database import AsyncSessionLocal

    async with AsyncSessionLocal() as db:
        order = await db.get(ClientOrder, ids["order_id"])
        assert order.rate_unit == RateUnit.daily
        assert (order.rate_candidate, order.rate_client) == (
            Decimal("960.000"),
            Decimal("1340.000"),
        )


async def test_contracts_export_has_no_daily_unit_after_conversion(
    app_client: AsyncClient, app_auth_headers: dict
):
    marker = f"eksport{uuid.uuid4().hex[:6]}"
    client_id, candidate_id = await _seed_people(marker)
    resp = await app_client.post(
        "/api/contracts",
        json={
            "candidate_id": candidate_id,
            "client_id": client_id,
            "start_date": "2026-09-01",
            "rate_unit": "daily",
            "rate_candidate": 960,
            "rate_client": 1340,
            "project_name": marker,
        },
        headers=app_auth_headers,
    )
    assert resp.status_code in (200, 201), resp.text

    export = await app_client.get(
        "/api/contracts/export",
        params={"format": "xlsx", "client_id": client_id},
        headers=app_auth_headers,
    )

    assert export.status_code == 200, export.text
    sheet = load_workbook(io.BytesIO(export.content)).active
    header = [cell.value for cell in sheet[1]]
    unit_col = header.index("Jednostka stawki")
    units = [row[unit_col] for row in sheet.iter_rows(min_row=2, values_only=True)]
    assert units == ["godzinowa"], units


# ── Jednorazowa korekta 0309 ────────────────────────────────────────────────


async def test_repair_converts_md_contracts_and_leaves_orders_and_others_untouched():
    from app.core.database import AsyncSessionLocal
    from app.models.app_setting import AppSetting
    from app.services.contract_hourly_rate_repair import (
        orders_checksum,
        run_contract_hourly_rate_repair,
        summarize_for_log,
    )

    marker = f"test_0309_{uuid.uuid4().hex[:8]}"
    details_key = f"repair_details_{marker}"
    client_id, candidate_id = await _seed_people()
    _, other_candidate = await _seed_people()
    _, monthly_candidate = await _seed_people()
    daily_id = await _seed_contract(
        client_id,
        candidate_id,
        rate_unit=RateUnit.daily,
        rate_candidate=Decimal("1001.550"),
        rate_client=Decimal("1340.000"),
        framework_rate=Decimal("1500.10"),
    )
    hourly_id = await _seed_contract(
        client_id,
        other_candidate,
        rate_unit=RateUnit.hourly,
        rate_candidate=Decimal("120"),
        rate_client=Decimal("150"),
    )
    monthly_id = await _seed_contract(
        client_id,
        monthly_candidate,
        rate_unit=RateUnit.monthly,
        rate_candidate=Decimal("17600"),
        rate_client=Decimal("22000"),
    )
    async with AsyncSessionLocal() as db:
        db.add_all(
            [
                ContractClientRate(
                    contract_id=daily_id,
                    rate=Decimal("1340.000"),
                    effective_from=date(2026, 1, 5),
                ),
                ClientOrder(
                    client_id=client_id,
                    contract_id=daily_id,
                    title="MD/ZAM",
                    status=ClientOrderStatus.active,
                    start_date=date(2026, 9, 1),
                    end_date=date(2026, 12, 31),
                    rate_unit=RateUnit.daily,
                    rate_client=Decimal("1340.000"),
                    rate_candidate=Decimal("1001.550"),
                    rate_client_currency="PLN",
                    rate_candidate_currency="PLN",
                    currency="PLN",
                ),
            ]
        )
        await db.commit()
    before = await _load(daily_id)
    monthly_before = before.monthly_rate(before.rate_client)
    async with AsyncSessionLocal() as db:
        orders_before = await orders_checksum(db)

    scope = {daily_id, hourly_id, monthly_id}
    try:
        async with AsyncSessionLocal() as db:
            summary = await run_contract_hourly_rate_repair(
                db, marker=marker, details_key=details_key, only_contract_ids=scope
            )
            await db.commit()
        assert summary is not None and "pending" not in summary, summary
        assert summary["converted_contract_ids"] == [daily_id]
        assert summary["client_orders_unchanged"] is True
        assert "converted 1/1" in summarize_for_log(summary)

        async with AsyncSessionLocal() as db:
            assert await orders_checksum(db) == orders_before
            assert (
                await run_contract_hourly_rate_repair(
                    db, marker=marker, details_key=details_key, only_contract_ids=scope
                )
                is None
            )
            activities = (
                await db.scalars(
                    select(Activity).where(
                        Activity.entity_type == "contract",
                        Activity.entity_id == daily_id,
                        Activity.action == "rate_unit_converted_to_hourly",
                    )
                )
            ).all()
            details = (await db.get(AppSetting, details_key)).value
            # Kontrakt zapisany w MD PO korekcie (stary kontener przy deployu).
            late = await db.get(Contract, hourly_id)
            late.rate_unit = RateUnit.daily
            late.rate_candidate = Decimal("960")
            late.rate_client = Decimal("1200")
            await db.commit()
        async with AsyncSessionLocal() as db:
            late_summary = await run_contract_hourly_rate_repair(
                db, marker=marker, details_key=details_key, only_contract_ids=scope
            )
            await db.commit()
            receipt = (await db.get(AppSetting, marker)).value
        assert late_summary == {"late_converted": [hourly_id], "skipped": []}
        assert receipt["late_converted_contract_ids"] == [hourly_id]
        relabelled = await _load(hourly_id)
        assert (relabelled.rate_unit, relabelled.rate_candidate) == (
            RateUnit.hourly,
            Decimal("120"),
        )
        # Przywróć stan „godzinowy od zawsze" dla asercji niżej.
        async with AsyncSessionLocal() as db:
            back = await db.get(Contract, hourly_id)
            back.rate_candidate = Decimal("120")
            back.rate_client = Decimal("150")
            back.billing_hours_per_month = 160
            await db.commit()
        assert len(activities) == 1
        assert details["converted"][0]["before"]["rate_candidate"] == "1001.550000"

        converted = await _load(daily_id)
        assert converted.rate_unit == RateUnit.hourly
        assert converted.billing_hours_per_month == 176
        assert converted.rate_candidate == Decimal("125.19375")
        assert converted.rate_client == Decimal("167.5")
        assert converted.framework_rate == Decimal("187.5125")
        assert converted.client_rate_schedule[0].rate == Decimal("167.5")
        assert converted.monthly_rate(converted.rate_client) == monthly_before

        untouched = await _load(hourly_id)
        assert (untouched.rate_unit, untouched.rate_candidate) == (
            RateUnit.hourly,
            Decimal("120"),
        )
        assert untouched.billing_hours_per_month == 160
        flat = await _load(monthly_id)
        assert (flat.rate_unit, flat.rate_client) == (RateUnit.monthly, Decimal("22000"))
    finally:
        async with AsyncSessionLocal() as db:
            await db.execute(
                delete(AppSetting).where(AppSetting.key.in_([marker, details_key]))
            )
            await db.commit()


async def test_repair_waits_when_columns_are_not_widened(monkeypatch):
    from app.core.database import AsyncSessionLocal
    from app.models.app_setting import AppSetting
    from app.services import contract_hourly_rate_repair as repair

    async def _narrow(_db) -> bool:
        return False

    monkeypatch.setattr(repair, "columns_widened", _narrow)
    marker = f"test_0309_wait_{uuid.uuid4().hex[:8]}"
    async with AsyncSessionLocal() as db:
        summary = await repair.run_contract_hourly_rate_repair(
            db, marker=marker, details_key=f"d_{marker}", only_contract_ids=set()
        )
        await db.commit()
        assert summary == {"pending": "columns_not_widened"}
        assert await db.get(AppSetting, marker) is None
