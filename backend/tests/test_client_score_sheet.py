"""0525: tabela punktowa klienta (Excel) — odczyt pól, punkty według zasad oceny
i zapis w układzie wzoru (bez bazy; trasy w ``test_client_forms_api.py``)."""

from __future__ import annotations

import io
import json
import zipfile

import pytest
from openpyxl import Workbook, load_workbook

from app.services import client_form_fill as fill
from app.services import client_form_render as render
from app.services import client_score_structure as score
from app.services.client_form_structure import FormRejected
from tests._client_score_xlsx import score_xlsx

CV = (
    "Anna Fikcyjna\n"
    "03.2019 – obecnie Tester automatyzujący / Bank Przykładowy S.A.\n"
    "Selenium WebDriver od 2021, SQL\n"
    "ISTQB Certified Tester FL (09.2022)"
)
RULES = (
    "Punkty za lata doświadczenia w automatyzacji: 0 pkt do 2 lat, 1 pkt więcej "
    "niż 2 lata, 2 pkt więcej niż 3 lata, 3 pkt więcej niż 4 lata. ISTQB FL: 3 pkt."
)


def _structure() -> dict:
    return score.read_structure(score_xlsx())


def _fields() -> dict[str, dict]:
    return {item["id"]: item for item in _structure()["fields"]}


def test_reads_person_lines_yes_no_results_and_points_with_maximum() -> None:
    fields = _fields()
    assert fields["s0.r2.c2"]["kind"] == "line"
    assert fields["s0.r2.c2"]["label"] == "Imię i nazwisko"
    assert fields["s0.r3.c2"]["label"] == "Firma"
    # „Tak/Nie” to metoda oceny — odpowiedź idzie do pustej kolumny „Wynik”.
    assert fields["s0.r7.c4"]["label"] == "Wynik (Tak/Nie)"
    assert fields["s0.r7.c4"]["row_label"].startswith("Minimum 3 lata")
    assert fields["s0.r7.c4"]["section"] == "Wymagania obligatoryjne"
    assert "s0.r7.c3" not in fields
    # Kryteria punktowane: maksimum z kolumny „Liczba punktów”.
    assert fields["s0.r12.c4"]["label"] == "Wynik (0–5 pkt)"
    assert fields["s0.r12.c4"]["max_points"] == 5
    assert fields["s0.r14.c4"]["max_points"] == 3
    # „Razem” z formułami zostaje nietknięte.
    assert not any(key.startswith("s0.r15.") for key in fields)
    assert score.max_points_total(_structure()) == 13


def test_summary_and_format() -> None:
    structure = _structure()
    assert structure["format"] == "xlsx"
    assert score.summary(structure).startswith("7 pól do wypełnienia · 2 pola nagłówka")


def test_filled_workbook_or_other_files_are_refused() -> None:
    book = Workbook()
    book.active["A1"] = "Stawki"
    book.active["A2"] = 100
    buffer = io.BytesIO()
    book.save(buffer)
    with pytest.raises(FormRejected) as no_form:
        score.read_structure(buffer.getvalue())
    assert str(no_form.value) == score.MSG_NO_FORM
    with pytest.raises(FormRejected) as not_excel:
        score.read_structure(b"%PDF-1.4")
    assert str(not_excel.value) == score.MSG_NOT_XLSX


def test_macros_are_refused() -> None:
    buffer = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(score_xlsx())) as source,
        zipfile.ZipFile(buffer, "w") as target,
    ):
        for entry in source.infolist():
            target.writestr(entry, source.read(entry))
        target.writestr("xl/vbaProject.bin", b"x")
    with pytest.raises(FormRejected) as refused:
        score.read_structure(buffer.getvalue())
    assert str(refused.value) == score.MSG_MACROS


def _parse(fields, rules: str = RULES) -> dict:
    return fill.check_parse(
        json.dumps({"fields": fields}),
        structure=_structure(),
        sources=fill.Sources(cv=CV, notes="", rules=rules),
        candidate_name="Anna Fikcyjna",
    )


def test_points_stay_within_the_row_maximum_and_need_rules() -> None:
    answer = [
        {"id": "s0.r12.c4", "value": "4 pkt", "quote": "03.2019 – obecnie"},
        {"id": "s0.r13.c4", "value": "6", "quote": "Selenium WebDriver od 2021"},
        {"id": "s0.r14.c4", "value": "3", "quote": "ISTQB Certified Tester FL"},
        {"id": "s0.r7.c4", "value": "tak", "quote": "03.2019 – obecnie"},
    ]
    result = _parse(answer)
    fields = result["fields"]
    assert fields["s0.r12.c4"]["value"] == "4"
    assert "s0.r13.c4" not in fields  # 6 > maksimum 5
    assert fields["s0.r14.c4"]["value"] == "3"
    assert fields["s0.r7.c4"]["value"] == "TAK"
    # „Firma” to wykonawca — my, nie pracodawca osoby; imię z profilu.
    assert fields["s0.r3.c2"] == {
        "value": "B2B.net S.A.",
        "source": "profile",
        "quote": None,
    }
    assert fields["s0.r2.c2"]["value"] == "Anna Fikcyjna"

    without_rules = _parse(answer, rules="")
    assert not any(key in without_rules["fields"] for key in ("s0.r12.c4", "s0.r14.c4"))
    assert without_rules["fields"]["s0.r7.c4"]["value"] == "TAK"


def test_prompt_carries_rules_and_maximum_only_when_rules_exist() -> None:
    with_rules = fill.build_prompt(
        _structure(), fill.Sources(cv=CV, notes="", rules=RULES)
    )
    assert "<scoring_rules>" in with_rules and '"max_points": 5' in with_rules
    without = fill.build_prompt(_structure(), fill.Sources(cv=CV, notes=""))
    assert '"s0.r12.c4"' not in without
    assert '"s0.r3.c2"' not in without  # „Firma” wpisuje kod


def test_render_writes_points_as_numbers_and_keeps_formulas() -> None:
    content = render.render_score(
        score_xlsx(),
        _structure(),
        {
            "s0.r2.c2": {"value": "Anna Fikcyjna"},
            "s0.r7.c4": {"value": "TAK"},
            "s0.r12.c4": {"value": "4"},
        },
    )
    sheet = load_workbook(io.BytesIO(content)).active
    assert sheet["B2"].value == "Anna Fikcyjna"
    assert sheet["D7"].value == "TAK"
    assert sheet["D12"].value == 4 and isinstance(sheet["D12"].value, int)
    assert sheet["C7"].value == "Tak/Nie"
    assert sheet["D15"].value == "=SUM(D12:D14)"
    assert sheet["D13"].value is None


def test_points_given_as_numbers_are_accepted() -> None:
    """Staging 27.10.2026: model oddał punkty liczbą i kod odrzucał wszystkie."""
    result = _parse(
        [
            {"id": "s0.r12.c4", "value": 2, "quote": "03.2019 – obecnie"},
            {"id": "s0.r14.c4", "value": 3.0, "quote": "ISTQB Certified Tester FL"},
            {"id": "s0.r13.c4", "value": True, "quote": "Selenium WebDriver od 2021"},
        ]
    )
    assert result["fields"]["s0.r12.c4"]["value"] == "2"
    assert result["fields"]["s0.r14.c4"]["value"] == "3"
    assert "s0.r13.c4" not in result["fields"]
