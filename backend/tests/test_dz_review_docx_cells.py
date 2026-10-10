"""QC CV czyta komórki tabeli Worda wprost z XML-a, nie przez ``row.cells``.

``row.cells`` z python-docx powtarza komórkę ``gridSpan`` razy, a tę liczbę
podaje plik. CV „…B2B…” przychodzi z formularza kariery i z importu, więc jest
niezaufane, a ``dz_review.docx_blocks`` czyta je w procesie aplikacji — ta sama
klasa błędu co w imporcie Worda do edytora (#2112).
"""

from __future__ import annotations

import io

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import _Row

from app.services import dz_review as svc


def _table_docx(*, grid_span: int | None = None, merged_down: bool = False) -> bytes:
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Java Developer"
    table.cell(0, 1).text = "Bank"
    table.cell(1, 1).text = "Kafka"
    if not merged_down:
        table.cell(1, 0).text = "Spring"
    # Komórki bierzemy przed zmianą XML-a: ``table.cell`` też rozwija gridSpan.
    top_left = table.cell(0, 0)._tc
    bottom_left = table.cell(1, 0)._tc
    if grid_span is not None:
        span = OxmlElement("w:gridSpan")
        span.set(qn("w:val"), str(grid_span))
        top_left.get_or_add_tcPr().append(span)
    if merged_down:
        start = OxmlElement("w:vMerge")
        start.set(qn("w:val"), "restart")
        top_left.get_or_add_tcPr().append(start)
        bottom_left.get_or_add_tcPr().append(OxmlElement("w:vMerge"))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _texts(data: bytes) -> list[str]:
    return ["".join(run["t"] for run in block.runs) for block in svc.docx_blocks(data)]


def test_reader_does_not_expand_cells_by_grid_span(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def _forbidden(self: _Row) -> tuple:
        raise AssertionError("row.cells powiela komórkę gridSpan razy — czytaj tc_lst")

    monkeypatch.setattr(_Row, "cells", property(_forbidden))

    texts = _texts(_table_docx())

    assert [t for t in texts if t in {"Java Developer", "Bank", "Spring", "Kafka"}] == [
        "Java Developer",
        "Bank",
        "Spring",
        "Kafka",
    ]


def test_cell_with_huge_grid_span_is_read_once() -> None:
    texts = _texts(_table_docx(grid_span=5_000_000))

    assert texts.count("Java Developer") == 1
    assert "Kafka" in texts


def test_vertically_merged_cell_is_not_repeated_in_every_row() -> None:
    texts = _texts(_table_docx(merged_down=True))

    assert texts.count("Java Developer") == 1
    assert "Kafka" in texts
