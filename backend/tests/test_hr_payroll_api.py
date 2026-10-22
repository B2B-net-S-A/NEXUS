"""Płace (E3, D2): stawki, premie, zakres widoczności, zamknięcie miesiąca, eksport.

Zakres: Płace widzą i zmieniają wszystko; osoba widzi swoje stawki i premie;
przełożony widzi godziny i premie zespołu, ale nie stawki ani kwot
wynagrodzenia; obcy dostaje 403/404. Zamknięty miesiąc zamraża linie,
blokuje timesheety (``payroll_locked``) i odpowiada 409 na zmiany.
Dane fikcyjne (repo jest publiczne).
"""

from __future__ import annotations

import io
import uuid
from datetime import date, timedelta
from decimal import Decimal

import pytest
import pytest_asyncio
from openpyxl import load_workbook
from sqlalchemy import delete, func, select, text

from app.core.database import AsyncSessionLocal
from app.core.scheduling import business_today, local_now
from app.core.security import create_access_token
from app.models.app_setting import AppSetting
from app.models.hr_ops import HrAuditEvent, HrJobRun
from app.models.hr_payroll import EmployeeRate, HrBonus, PayrollLine, PayrollPeriod
from app.models.hr_profile import EmployeeProfile
from app.models.notification import Notification, NotificationType
from app.models.section_permission import UserActionOverride
from app.models.user import User, UserRole
from app.services.hr.bonus_rules import SETTING_KEY as RULES_KEY
from app.services.hr.payroll_periods import add_months
from app.services.hr_foundation_schema import HR_MODULE_KEYS, HR_MODULES_SETTING
from tests.hr_settlements_seed import delete_settlement_rows

# ── Pomocnicze ────────────────────────────────────────────────────────────────


def _headers(user: User) -> dict[str, str]:
    token = create_access_token(subject=user.id, role=user.role.value)
    return {"Authorization": f"Bearer {token}"}


async def _user(db, role: str, *, profile: bool = False, **profile_values) -> User:
    marker = uuid.uuid4().hex[:10]
    user = User(
        email=f"payroll-{marker}@example.com",
        name=f"Osoba Przykładowa {marker}",
        role=UserRole(role),
        roles=[role],
        is_active=True,
    )
    db.add(user)
    await db.flush()
    if profile:
        db.add(EmployeeProfile(user_id=user.id, **profile_values))
        await db.flush()
    return user


async def _grant(db, user: User, action: str) -> None:
    db.add(UserActionOverride(user_id=user.id, action=action, access="manage"))
    await db.flush()


async def _set_setting(key: str, value) -> None:
    async with AsyncSessionLocal() as db:
        row = await db.get(AppSetting, key)
        if value is None:
            if row is not None:
                await db.delete(row)
        elif row is None:
            db.add(AppSetting(key=key, value=value))
        else:
            row.value = value
        await db.commit()


async def _wipe_payroll() -> None:
    async with AsyncSessionLocal() as db:
        await delete_settlement_rows(db)
        await db.execute(delete(PayrollLine))
        await db.execute(delete(PayrollPeriod))
        await db.execute(delete(HrBonus))
        await db.execute(delete(EmployeeRate))
        await db.commit()


@pytest_asyncio.fixture
async def payroll_on():
    async with AsyncSessionLocal() as db:
        previous = await db.scalar(
            select(AppSetting.value).where(AppSetting.key == HR_MODULES_SETTING)
        )
        rules = await db.scalar(
            select(AppSetting.value).where(AppSetting.key == RULES_KEY)
        )
    await _set_setting(HR_MODULES_SETTING, {key: True for key in HR_MODULE_KEYS})
    await _wipe_payroll()
    yield
    await _wipe_payroll()
    await _set_setting(RULES_KEY, rules)
    await _set_setting(
        HR_MODULES_SETTING, previous or {key: False for key in HR_MODULE_KEYS}
    )


def _months_ago(n: int) -> tuple[int, int]:
    today = business_today()
    return add_months(today.year, today.month, -n)


async def _timesheet(
    db, user_id: int, year: int, month: int, status: str, entries
) -> int:
    row = await db.execute(
        text(
            "INSERT INTO timesheets (user_id, year, month, status) "
            "VALUES (:u, :y, :m, :s) RETURNING id"
        ),
        {"u": user_id, "y": year, "m": month, "s": status},
    )
    timesheet_id = row.scalar_one()
    for day, hours, overtime in entries:
        await db.execute(
            text(
                "INSERT INTO timesheet_entries (timesheet_id, work_date, hours, "
                "description, is_overtime, overtime_reason) "
                "VALUES (:t, :d, :h, 'Praca w projekcie', :o, :r)"
            ),
            {
                "t": timesheet_id,
                "d": date(year, month, day),
                "h": Decimal(hours),
                "o": overtime,
                "r": "Wdrożenie u klienta" if overtime else None,
            },
        )
    return timesheet_id


async def _payroll_team():
    """Płace, przełożony, jego człowiek i osoba z innego zespołu."""

    async with AsyncSessionLocal() as db:
        finance = await _user(db, "recruiter")
        await _grant(db, finance, "payroll")
        boss = await _user(db, "recruiter", profile=True, employment_type="uop")
        worker = await _user(
            db, "employee", profile=True, manager_user_id=boss.id, employment_type="b2b"
        )
        other = await _user(db, "employee", profile=True, employment_type="uop")
        await db.commit()
    return finance, boss, worker, other


# ── Moduł i bramki ────────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_module_off_hides_payroll_except_for_admin(app_client) -> None:
    async with AsyncSessionLocal() as db:
        previous = await db.scalar(
            select(AppSetting.value).where(AppSetting.key == HR_MODULES_SETTING)
        )
        finance = await _user(db, "finance")
        admin = await _user(db, "admin")
        await db.commit()
    await _set_setting(HR_MODULES_SETTING, {key: False for key in HR_MODULE_KEYS})
    try:
        denied = await app_client.get(
            "/api/hr/payroll/months", headers=_headers(finance)
        )
        assert denied.status_code == 404
        allowed = await app_client.get(
            "/api/hr/payroll/months", headers=_headers(admin)
        )
        assert allowed.status_code == 200
    finally:
        await _set_setting(
            HR_MODULES_SETTING, previous or {key: False for key in HR_MODULE_KEYS}
        )


@pytest.mark.asyncio
async def test_finance_role_has_payroll_and_others_get_named_denial(
    app_client, payroll_on
) -> None:
    async with AsyncSessionLocal() as db:
        finance = await _user(db, "finance")
        recruiter = await _user(db, "recruiter", profile=True)
        outsider = await _user(db, "recruiter")
        await db.commit()

    assert (
        await app_client.get("/api/hr/payroll/months", headers=_headers(finance))
    ).status_code == 200
    refused = await app_client.get(
        "/api/hr/payroll/months", headers=_headers(recruiter)
    )
    assert refused.status_code == 403
    assert refused.json()["detail"]["permission"] == "payroll"
    # Bez kartoteki, bez podwładnych i bez Płac — nic z Płac.
    assert (
        await app_client.get("/api/hr/payroll/me", headers=_headers(outsider))
    ).status_code == 403


# ── Stawki ────────────────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_new_rate_starts_on_first_day_and_closes_previous(
    app_client, payroll_on
) -> None:
    finance, _boss, worker, _other = await _payroll_team()
    (y1, m1), (y2, m2) = _months_ago(3), _months_ago(1)
    url = f"/api/hr/payroll/people/{worker.id}/rates"

    first = await app_client.post(
        url,
        json={
            "hourly_rate": "100",
            "currency": "PLN",
            "effective_from": f"{y1}-{m1:02d}-01",
        },
        headers=_headers(finance),
    )
    assert first.status_code == 201, first.text
    second = await app_client.post(
        url,
        json={"hourly_rate": "120.50", "effective_from": f"{y2}-{m2:02d}-01"},
        headers=_headers(finance),
    )
    assert second.status_code == 201, second.text

    async with AsyncSessionLocal() as db:
        rows = list(
            await db.scalars(
                select(EmployeeRate)
                .where(EmployeeRate.user_id == worker.id)
                .order_by(EmployeeRate.effective_from)
            )
        )
    assert [str(r.hourly_rate) for r in rows] == ["100.00", "120.50"]
    assert rows[0].effective_to == date(y2, m2, 1) - timedelta(days=1)
    assert rows[1].effective_to is None

    # Krok wstawiony pomiędzy kończy się dzień przed kolejnym.
    (ym, mm) = _months_ago(2)
    middle = await app_client.post(
        url,
        json={"hourly_rate": "110", "effective_from": f"{ym}-{mm:02d}-01"},
        headers=_headers(finance),
    )
    assert middle.status_code == 201
    assert (
        middle.json()["effective_to"]
        == (date(y2, m2, 1) - timedelta(days=1)).isoformat()
    )

    not_first = await app_client.post(
        url,
        json={"hourly_rate": "130", "effective_from": f"{y2}-{m2:02d}-15"},
        headers=_headers(finance),
    )
    assert not_first.status_code == 422
    assert "1. dnia" in not_first.json()["detail"]

    ty, tm = _months_ago(13)
    too_old = await app_client.post(
        url,
        json={"hourly_rate": "130", "effective_from": f"{ty}-{tm:02d}-01"},
        headers=_headers(finance),
    )
    assert too_old.status_code == 422

    async with AsyncSessionLocal() as db:
        no_profile = await _user(db, "recruiter")
        await db.commit()
    missing = await app_client.post(
        f"/api/hr/payroll/people/{no_profile.id}/rates",
        json={"hourly_rate": "90", "effective_from": f"{y2}-{m2:02d}-01"},
        headers=_headers(finance),
    )
    assert missing.status_code == 404


@pytest.mark.asyncio
async def test_new_rate_after_gap_keeps_gap_open(app_client, payroll_on) -> None:
    """Luka z importu (stawka z końcem, bez następnej) zostaje luką (A2, 19.10)."""

    finance, _boss, worker, _other = await _payroll_team()
    (ya, ma), (ye, me), (yn, mn) = _months_ago(6), _months_ago(4), _months_ago(1)
    gap_end = date(*add_months(ye, me, 1), 1) - timedelta(days=1)
    async with AsyncSessionLocal() as db:
        db.add(
            EmployeeRate(
                user_id=worker.id,
                hourly_rate=Decimal("100.00"),
                currency="PLN",
                effective_from=date(ya, ma, 1),
                effective_to=gap_end,
                set_by=finance.id,
            )
        )
        await db.commit()

    response = await app_client.post(
        f"/api/hr/payroll/people/{worker.id}/rates",
        json={"hourly_rate": "120", "effective_from": f"{yn}-{mn:02d}-01"},
        headers=_headers(finance),
    )
    assert response.status_code == 201, response.text

    async with AsyncSessionLocal() as db:
        rows = list(
            await db.scalars(
                select(EmployeeRate)
                .where(EmployeeRate.user_id == worker.id)
                .order_by(EmployeeRate.effective_from)
            )
        )
    assert [r.effective_to for r in rows] == [gap_end, None]


@pytest.mark.asyncio
async def test_rate_change_notifies_without_amount_and_audits_as_payroll(
    app_client, payroll_on
) -> None:
    finance, _boss, worker, _other = await _payroll_team()
    y, m = _months_ago(1)
    response = await app_client.post(
        f"/api/hr/payroll/people/{worker.id}/rates",
        json={"hourly_rate": "987.65", "effective_from": f"{y}-{m:02d}-01"},
        headers=_headers(finance),
    )
    assert response.status_code == 201

    async with AsyncSessionLocal() as db:
        notes = list(
            await db.scalars(
                select(Notification).where(
                    Notification.user_id == worker.id,
                    Notification.notification_type == NotificationType.hr_rate_changed,
                )
            )
        )
        events = list(
            await db.scalars(
                select(HrAuditEvent).where(HrAuditEvent.subject_user_id == worker.id)
            )
        )
    assert len(notes) == 1
    assert "987" not in notes[0].title + notes[0].message
    with_amount = [e for e in events if "987.65" in str(e.details)]
    assert with_amount and all(e.sensitivity == "payroll" for e in with_amount)


@pytest.mark.asyncio
async def test_payroll_cannot_set_or_delete_own_rate_but_admin_can(
    app_client, payroll_on
) -> None:
    """Decyzja Artura 19.10.2026: własną stawkę zmienia druga osoba z Płac albo admin."""

    async with AsyncSessionLocal() as db:
        finance = await _user(db, "recruiter", profile=True, employment_type="b2b")
        await _grant(db, finance, "payroll")
        colleague = await _user(db, "recruiter")
        await _grant(db, colleague, "payroll")
        admin = await _user(db, "admin", profile=True, employment_type="b2b")
        await db.commit()
    y, m = _months_ago(1)
    body = {"hourly_rate": "101.00", "effective_from": f"{y}-{m:02d}-01"}

    own = await app_client.post(
        f"/api/hr/payroll/people/{finance.id}/rates",
        json=body,
        headers=_headers(finance),
    )
    assert own.status_code == 403
    assert own.json()["detail"] == (
        "Własnej stawki nie ustawiasz — zrobi to druga osoba z Płac albo administrator."
    )
    by_colleague = await app_client.post(
        f"/api/hr/payroll/people/{finance.id}/rates",
        json=body,
        headers=_headers(colleague),
    )
    assert by_colleague.status_code == 201, by_colleague.text
    own_delete = await app_client.delete(
        f"/api/hr/payroll/rates/{by_colleague.json()['id']}", headers=_headers(finance)
    )
    assert own_delete.status_code == 403
    assert own_delete.json()["detail"].startswith("Własnej stawki nie usuwasz")
    own_page = await app_client.get(
        f"/api/hr/payroll/people/{finance.id}", headers=_headers(finance)
    )
    assert own_page.status_code == 200
    assert own_page.json()["can_edit_rates"] is False

    admin_own = await app_client.post(
        f"/api/hr/payroll/people/{admin.id}/rates", json=body, headers=_headers(admin)
    )
    assert admin_own.status_code == 201, admin_own.text
    async with AsyncSessionLocal() as db:
        rates = list(
            await db.scalars(
                select(EmployeeRate.user_id).where(
                    EmployeeRate.user_id.in_([finance.id, admin.id])
                )
            )
        )
        trail = await db.scalar(
            select(HrAuditEvent.id).where(
                HrAuditEvent.subject_user_id == admin.id,
                HrAuditEvent.actor_user_id == admin.id,
                HrAuditEvent.action == "rate_set",
            )
        )
    assert sorted(rates) == sorted([finance.id, admin.id])
    assert trail is not None


# ── Zakres widoczności (D2) ───────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_manager_sees_team_hours_and_bonuses_but_not_rates(
    app_client, payroll_on
) -> None:
    finance, boss, worker, other = await _payroll_team()
    y, m = _months_ago(1)
    async with AsyncSessionLocal() as db:
        await _timesheet(db, worker.id, y, m, "approved", [(5, "8", False)])
        await db.commit()
    await app_client.post(
        f"/api/hr/payroll/people/{worker.id}/rates",
        json={"hourly_rate": "100", "effective_from": f"{y}-{m:02d}-01"},
        headers=_headers(finance),
    )
    bonus = await app_client.post(
        "/api/hr/payroll/bonuses",
        json={
            "recipient_user_id": worker.id,
            "amount": "300",
            "category": "custom",
            "period_year": y,
            "period_month": m,
        },
        headers=_headers(boss),
    )
    assert bonus.status_code == 201, bonus.text
    # Premia od przełożonego to propozycja — liczy się po zatwierdzeniu Płac.
    approved = await app_client.post(
        f"/api/hr/payroll/bonuses/{bonus.json()['id']}/approve",
        headers=_headers(finance),
    )
    assert approved.status_code == 200, approved.text

    month = await app_client.get(
        f"/api/hr/payroll/months/{y}/{m}", headers=_headers(boss)
    )
    assert month.status_code == 200
    body = month.json()
    assert body["scope"] == "team"
    assert "totals" not in body
    assert [p["user_id"] for p in body["people"]] == [worker.id]
    person = body["people"][0]
    assert person["hours"] == "8.00"
    assert "lines" not in person and "rate_missing" not in person
    assert person["bonuses"] == [{"currency": "PLN", "bonus_total": "300.00"}]
    assert not {"hourly_rate", "hours_amount", "total"} & set(person)

    card = await app_client.get(
        f"/api/hr/payroll/people/{worker.id}", headers=_headers(boss)
    )
    assert card.status_code == 200
    assert card.json()["scope"] == "team"
    assert "rates" not in card.json() and "current_rate" not in card.json()
    assert all("rule_snapshot" not in b for b in card.json()["bonuses"])

    own = await app_client.get(
        f"/api/hr/payroll/people/{worker.id}", headers=_headers(worker)
    )
    assert own.json()["scope"] == "full"
    assert own.json()["current_rate"]["hourly_rate"] == "100.00"

    me = await app_client.get("/api/hr/payroll/me", headers=_headers(worker))
    assert me.status_code == 200
    assert me.json()["rates"][0]["hourly_rate"] == "100.00"
    assert me.json()["bonuses"][0]["amount"] == "300.00"

    stranger = await app_client.get(
        f"/api/hr/payroll/people/{worker.id}", headers=_headers(other)
    )
    assert stranger.status_code == 404
    no_team = await app_client.get(
        f"/api/hr/payroll/months/{y}/{m}", headers=_headers(other)
    )
    assert no_team.status_code == 403

    full = await app_client.get(
        f"/api/hr/payroll/months/{y}/{m}", headers=_headers(finance)
    )
    line = next(p for p in full.json()["people"] if p["user_id"] == worker.id)["lines"]
    assert line == [
        {
            "currency": "PLN",
            "hourly_rate": "100.00",
            "hours_amount": "800.00",
            "bonus_total": "300.00",
            "total": "1100.00",
        }
    ]


# ── Premie ────────────────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_manual_bonus_scope_self_and_soft_cancel(app_client, payroll_on) -> None:
    finance, boss, worker, other = await _payroll_team()
    y, m = _months_ago(0)
    payload = {
        "amount": "450",
        "category": "custom",
        "period_year": y,
        "period_month": m,
        "description": "Za wdrożenie nowego klienta",
    }

    outside = await app_client.post(
        "/api/hr/payroll/bonuses",
        json={**payload, "recipient_user_id": other.id},
        headers=_headers(boss),
    )
    assert outside.status_code == 403
    selfish = await app_client.post(
        "/api/hr/payroll/bonuses",
        json={**payload, "recipient_user_id": boss.id},
        headers=_headers(boss),
    )
    assert selfish.status_code == 403
    ty, tm = _months_ago(13)
    too_old = await app_client.post(
        "/api/hr/payroll/bonuses",
        json={
            **payload,
            "recipient_user_id": worker.id,
            "period_year": ty,
            "period_month": tm,
        },
        headers=_headers(boss),
    )
    assert too_old.status_code == 422

    given = await app_client.post(
        "/api/hr/payroll/bonuses",
        json={**payload, "recipient_user_id": worker.id},
        headers=_headers(boss),
    )
    assert given.status_code == 201, given.text
    # Od przełożonego — propozycja dla Płac (19.10.2026); autor ją wycofuje.
    assert given.json()["status"] == "proposed"
    bonus_id = given.json()["id"]

    by_finance = await app_client.post(
        "/api/hr/payroll/bonuses",
        json={**payload, "recipient_user_id": worker.id, "amount": "100"},
        headers=_headers(finance),
    )
    assert by_finance.status_code == 201
    # Przełożony nie zmienia premii, której sam nie przyznał.
    foreign_edit = await app_client.patch(
        f"/api/hr/payroll/bonuses/{by_finance.json()['id']}",
        json={"amount": "999"},
        headers=_headers(boss),
    )
    assert foreign_edit.status_code == 403

    cancelled = await app_client.post(
        f"/api/hr/payroll/bonuses/{bonus_id}/cancel",
        json={"note": "Pomyłka w kwocie"},
        headers=_headers(boss),
    )
    assert cancelled.status_code == 200
    assert cancelled.json()["status"] == "cancelled"
    again = await app_client.post(
        f"/api/hr/payroll/bonuses/{bonus_id}/cancel", json={}, headers=_headers(boss)
    )
    assert again.status_code == 409

    async with AsyncSessionLocal() as db:
        row = await db.get(HrBonus, bonus_id)
        notes = list(
            await db.scalars(
                select(Notification).where(
                    Notification.user_id == worker.id,
                    Notification.notification_type
                    == NotificationType.hr_bonus_assigned,
                )
            )
        )
    assert row is not None and row.cancel_reason == "manual"
    assert notes and all("450" not in n.title + n.message for n in notes)


@pytest.mark.asyncio
async def test_proposal_waits_for_deadline_and_nobody_approves_own(
    app_client, payroll_on
) -> None:
    finance, boss, worker, _other = await _payroll_team()
    today = business_today()
    y, m = today.year, today.month
    async with AsyncSessionLocal() as db:
        early = HrBonus(
            recipient_user_id=worker.id,
            amount=Decimal("772.80"),
            category="delivery_lead",
            period_year=y,
            period_month=m,
            status="proposed",
            source="contract_placement",
            source_key=f"contract:test-{uuid.uuid4().hex}:dl",
            eligible_on=today + timedelta(days=5),
        )
        ready = HrBonus(
            recipient_user_id=finance.id,
            amount=Decimal("1500"),
            category="recruiter",
            period_year=y,
            period_month=m,
            status="proposed",
            source="contract_placement",
            source_key=f"contract:test-{uuid.uuid4().hex}:recruiter",
            eligible_on=today,
        )
        db.add_all([early, ready])
        await db.commit()

    waiting = await app_client.post(
        f"/api/hr/payroll/bonuses/{early.id}/approve", headers=_headers(boss)
    )
    assert waiting.status_code == 409
    assert "liczy się" in waiting.json()["detail"]

    own = await app_client.post(
        f"/api/hr/payroll/bonuses/{ready.id}/approve", headers=_headers(finance)
    )
    assert own.status_code in (403, 404)

    async with AsyncSessionLocal() as db:
        admin = await _user(db, "admin")
        await db.commit()
    approved = await app_client.post(
        f"/api/hr/payroll/bonuses/{ready.id}/approve", headers=_headers(admin)
    )
    assert approved.status_code == 200
    assert approved.json()["status"] == "assigned"


@pytest.mark.asyncio
async def test_manual_bonus_from_manager_waits_for_payroll(
    app_client, payroll_on
) -> None:
    """Decyzja Artura 19.10.2026: premia ręczna przełożonego = propozycja dla Płac."""

    finance, boss, worker, _other = await _payroll_team()
    y, m = _months_ago(0)
    proposed = await app_client.post(
        "/api/hr/payroll/bonuses",
        json={
            "recipient_user_id": worker.id,
            "amount": "250",
            "category": "custom",
            "period_year": y,
            "period_month": m,
            "description": "Za pomoc przy wdrożeniu",
        },
        headers=_headers(boss),
    )
    assert proposed.status_code == 201, proposed.text
    body = proposed.json()
    assert body["status"] == "proposed" and body["source"] == "manual"
    assert body["assigned_by"] is None and body["assigned_at"] is None
    assert body["can_approve"] is False
    bonus_id = body["id"]

    async def worker_bells() -> int:
        async with AsyncSessionLocal() as db:
            return len(
                list(
                    await db.scalars(
                        select(Notification.id).where(
                            Notification.user_id == worker.id,
                            Notification.related_entity_id == bonus_id,
                            Notification.notification_type
                            == NotificationType.hr_bonus_assigned,
                        )
                    )
                )
            )

    assert await worker_bells() == 0

    # Lista „Do decyzji”: przełożony nie ma „Zatwierdź”, Płace mają.
    listed = await app_client.get(
        "/api/hr/payroll/bonuses", params={"user_id": worker.id}, headers=_headers(boss)
    )
    row = next(item for item in listed.json()["items"] if item["id"] == bonus_id)
    assert row["can_approve"] is False
    listed = await app_client.get(
        "/api/hr/payroll/bonuses",
        params={"user_id": worker.id},
        headers=_headers(finance),
    )
    row = next(item for item in listed.json()["items"] if item["id"] == bonus_id)
    assert row["can_approve"] is True

    own = await app_client.post(
        f"/api/hr/payroll/bonuses/{bonus_id}/approve", headers=_headers(boss)
    )
    assert own.status_code == 403
    assert "Płace" in own.json()["detail"]

    # Propozycję przełożony może jeszcze poprawić.
    corrected = await app_client.patch(
        f"/api/hr/payroll/bonuses/{bonus_id}",
        json={"amount": "300"},
        headers=_headers(boss),
    )
    assert corrected.status_code == 200, corrected.text
    assert corrected.json()["status"] == "proposed"

    approved = await app_client.post(
        f"/api/hr/payroll/bonuses/{bonus_id}/approve", headers=_headers(finance)
    )
    assert approved.status_code == 200, approved.text
    assert approved.json()["status"] == "assigned"
    assert approved.json()["assigned_by"] == finance.id
    assert await worker_bells() == 1

    # Audyt 21.10.2026: po zatwierdzeniu przełożony nie podnosi kwoty bez Płac.
    raised = await app_client.patch(
        f"/api/hr/payroll/bonuses/{bonus_id}",
        json={"amount": "50000"},
        headers=_headers(boss),
    )
    assert raised.status_code == 403
    assert "Płace" in raised.json()["detail"]
    async with AsyncSessionLocal() as db:
        kept = await db.get(HrBonus, bonus_id)
    assert kept is not None and kept.amount == Decimal("300.00")
    assert await worker_bells() == 1
    by_payroll = await app_client.patch(
        f"/api/hr/payroll/bonuses/{bonus_id}",
        json={"amount": "320"},
        headers=_headers(finance),
    )
    assert by_payroll.status_code == 200, by_payroll.text

    # Decyzja Artura 22.10.2026: kwoty po zatwierdzeniu przełożony nie zmienia,
    # ale przyznaną przez siebie premię może anulować (osoba dostaje dzwonek).
    withdrawn = await app_client.post(
        f"/api/hr/payroll/bonuses/{bonus_id}/cancel", json={}, headers=_headers(boss)
    )
    assert withdrawn.status_code == 200, withdrawn.text
    assert withdrawn.json()["status"] == "cancelled"
    assert await worker_bells() == 2

    # Płace przyznają od razu, jak dotąd.
    direct = await app_client.post(
        "/api/hr/payroll/bonuses",
        json={
            "recipient_user_id": worker.id,
            "amount": "90",
            "period_year": y,
            "period_month": m,
        },
        headers=_headers(finance),
    )
    assert direct.status_code == 201
    assert direct.json()["status"] == "assigned"
    assert direct.json()["assigned_by"] == finance.id


@pytest.mark.asyncio
async def test_batch_approve_is_all_or_nothing_and_payroll_only(
    app_client, payroll_on
) -> None:
    """„Zatwierdź zaznaczone” (19.10.2026): jedna transakcja, tylko Płace."""

    finance, boss, worker, other = await _payroll_team()
    today = business_today()
    y, m = today.year, today.month

    def proposal(recipient: int, amount: str, **extra) -> HrBonus:
        return HrBonus(
            recipient_user_id=recipient,
            amount=Decimal(amount),
            category="recruiter",
            period_year=y,
            period_month=m,
            status="proposed",
            source="contract_placement",
            source_key=f"contract:test-{uuid.uuid4().hex}:recruiter",
            **extra,
        )

    async with AsyncSessionLocal() as db:
        first = proposal(worker.id, "1000", eligible_on=today)
        second = proposal(other.id, "529.20", eligible_on=today)
        early = proposal(other.id, "300", eligible_on=today + timedelta(days=5))
        db.add_all([first, second, early])
        await db.commit()

    # Przełożony (bez Płac) — odmowa bramki, nawet dla swojego człowieka.
    denied = await app_client.post(
        "/api/hr/payroll/bonuses/approve-batch",
        json={"ids": [first.id]},
        headers=_headers(boss),
    )
    assert denied.status_code == 403

    # Jedna premia przed terminem — nic nie wchodzi, komunikat mówi przy kim.
    refused = await app_client.post(
        "/api/hr/payroll/bonuses/approve-batch",
        json={"ids": [first.id, second.id, early.id]},
        headers=_headers(finance),
    )
    assert refused.status_code == 409
    assert refused.json()["detail"].startswith("Nic nie zatwierdzono")
    assert other.name in refused.json()["detail"]
    async with AsyncSessionLocal() as db:
        states = dict(
            (
                await db.execute(
                    select(HrBonus.id, HrBonus.status).where(
                        HrBonus.id.in_([first.id, second.id, early.id])
                    )
                )
            ).all()
        )
        bells = await db.scalar(
            select(func.count(Notification.id)).where(
                Notification.notification_type == NotificationType.hr_bonus_assigned,
                Notification.related_entity_id.in_([first.id, second.id]),
            )
        )
    assert set(states.values()) == {"proposed"}
    assert bells == 0

    approved = await app_client.post(
        "/api/hr/payroll/bonuses/approve-batch",
        json={"ids": [second.id, first.id, first.id]},
        headers=_headers(finance),
    )
    assert approved.status_code == 200, approved.text
    body = approved.json()
    assert body["approved"] == 2
    assert {item["id"] for item in body["items"]} == {first.id, second.id}
    assert {item["status"] for item in body["items"]} == {"assigned"}

    # Drugi raz te same — już nie czekają (409, bez zmian).
    again = await app_client.post(
        "/api/hr/payroll/bonuses/approve-batch",
        json={"ids": [first.id]},
        headers=_headers(finance),
    )
    assert again.status_code == 409

    empty = await app_client.post(
        "/api/hr/payroll/bonuses/approve-batch",
        json={"ids": []},
        headers=_headers(finance),
    )
    assert empty.status_code == 422


# ── Zamknięcie miesiąca ───────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_close_month_freezes_lines_locks_timesheets_and_is_idempotent(
    app_client, payroll_on
) -> None:
    finance, boss, worker, other = await _payroll_team()
    y, m = _months_ago(2)
    async with AsyncSessionLocal() as db:
        sheet = await _timesheet(
            db,
            worker.id,
            y,
            m,
            "approved",
            [(3, "8", False), (4, "5", False), (4, "5", True)],
        )
        admin = await _user(db, "admin")
        await db.commit()
    await app_client.post(
        f"/api/hr/payroll/people/{worker.id}/rates",
        json={"hourly_rate": "100", "effective_from": f"{y}-{m:02d}-01"},
        headers=_headers(finance),
    )
    for amount, currency in (("500", "PLN"), ("100", "EUR")):
        created = await app_client.post(
            "/api/hr/payroll/bonuses",
            json={
                "recipient_user_id": worker.id,
                "amount": amount,
                "currency": currency,
                "category": "sales",
                "period_year": y,
                "period_month": m,
            },
            headers=_headers(finance),
        )
        assert created.status_code == 201

    url = f"/api/hr/payroll/months/{y}/{m}"
    warned = await app_client.post(f"{url}/close", json={}, headers=_headers(finance))
    assert warned.status_code == 409
    assert warned.json()["detail"]["code"] == "close_warnings"

    closed = await app_client.post(
        f"{url}/close", json={"acknowledge_warnings": True}, headers=_headers(finance)
    )
    assert closed.status_code == 200, closed.text
    assert closed.json()["already_closed"] is False
    repeat = await app_client.post(
        f"{url}/close", json={"acknowledge_warnings": True}, headers=_headers(finance)
    )
    assert repeat.status_code == 200
    assert repeat.json()["already_closed"] is True

    async with AsyncSessionLocal() as db:
        lines = list(
            await db.scalars(
                select(PayrollLine)
                .where(PayrollLine.user_id == worker.id)
                .order_by(PayrollLine.currency)
            )
        )
        locked = await db.scalar(
            text("SELECT payroll_locked FROM timesheets WHERE id = :id"), {"id": sheet}
        )
        runs = list(
            await db.scalars(
                select(HrJobRun).where(
                    HrJobRun.kind == "payroll_close",
                    HrJobRun.period_key.like(f"{y}-{m:02d}#%"),
                )
            )
        )
        # Dane źródłowe zmieniają się po zamknięciu — linie nie.
        await db.execute(
            text("UPDATE timesheet_entries SET hours = 1 WHERE timesheet_id = :id"),
            {"id": sheet},
        )
        await db.commit()
    assert [(row.currency, str(row.hours), str(row.total)) for row in lines] == [
        ("EUR", "0.00", "100.00"),
        ("PLN", "18.00", "2300.00"),
    ]
    # 5 h + 5 h tego samego dnia = 2 h ponad normę (flaga wpisu to nie liczba nadgodzin).
    assert lines[1].overtime_hours == Decimal("2.00")
    assert locked is True
    assert len(runs) == 1 and "2300" not in str(runs[0].counts)

    frozen = await app_client.get(url, headers=_headers(finance))
    person = next(p for p in frozen.json()["people"] if p["user_id"] == worker.id)
    assert frozen.json()["status"] == "closed"
    assert person["hours"] == "18.00"

    # Zamknięty miesiąc: zmiany stawek i premii tego miesiąca → 409.
    rate = await app_client.post(
        f"/api/hr/payroll/people/{worker.id}/rates",
        json={"hourly_rate": "130", "effective_from": f"{y}-{m:02d}-01"},
        headers=_headers(finance),
    )
    assert rate.status_code == 409
    bonus = await app_client.post(
        "/api/hr/payroll/bonuses",
        json={
            "recipient_user_id": worker.id,
            "amount": "50",
            "category": "custom",
            "period_year": y,
            "period_month": m,
        },
        headers=_headers(boss),
    )
    assert bonus.status_code == 409
    ny, nm = add_months(y, m, 1)
    later = await app_client.post(
        f"/api/hr/payroll/people/{worker.id}/rates",
        json={"hourly_rate": "130", "effective_from": f"{ny}-{nm:02d}-01"},
        headers=_headers(finance),
    )
    assert later.status_code == 201

    # Ponowne otwarcie: tylko admin, z powodem i śladem.
    refused = await app_client.post(
        f"{url}/reopen", json={"reason": "Korekta godzin"}, headers=_headers(finance)
    )
    assert refused.status_code == 403
    reopened = await app_client.post(
        f"{url}/reopen", json={"reason": "Korekta godzin"}, headers=_headers(admin)
    )
    assert reopened.status_code == 200, reopened.text
    async with AsyncSessionLocal() as db:
        left = await db.scalar(
            select(PayrollLine.id).where(PayrollLine.user_id == worker.id).limit(1)
        )
        unlocked = await db.scalar(
            text("SELECT payroll_locked FROM timesheets WHERE id = :id"), {"id": sheet}
        )
        trail = list(
            await db.scalars(
                select(HrAuditEvent).where(
                    HrAuditEvent.entity_type == "payroll_period",
                    HrAuditEvent.action.in_(["month_closed", "month_reopened"]),
                )
            )
        )
    assert left is None and unlocked is False
    assert {e.action for e in trail} == {"month_closed", "month_reopened"}
    assert all(e.sensitivity == "payroll" for e in trail)


@pytest.mark.asyncio
async def test_pending_proposals_and_running_month_block_close(
    app_client, payroll_on
) -> None:
    finance, _boss, worker, _other = await _payroll_team()
    y, m = _months_ago(1)
    async with AsyncSessionLocal() as db:
        db.add(
            HrBonus(
                recipient_user_id=worker.id,
                amount=Decimal("5000"),
                category="league",
                period_year=y,
                period_month=m,
                status="proposed",
                source="league",
                source_key=f"league:test:{uuid.uuid4().hex}",
            )
        )
        await db.commit()
    blocked = await app_client.post(
        f"/api/hr/payroll/months/{y}/{m}/close",
        json={"acknowledge_warnings": True},
        headers=_headers(finance),
    )
    assert blocked.status_code == 409
    assert blocked.json()["detail"]["code"] == "pending_bonuses"

    today = business_today()
    running = await app_client.post(
        f"/api/hr/payroll/months/{today.year}/{today.month}/close",
        json={"acknowledge_warnings": True},
        headers=_headers(finance),
    )
    if today < date(*add_months(today.year, today.month, 1), 1) - timedelta(days=1):
        assert running.status_code == 422


# ── Reguły i eksport ──────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_bonus_rules_are_settings_of_payroll(app_client, payroll_on) -> None:
    finance, boss, worker, _other = await _payroll_team()
    defaults = await app_client.get("/api/hr/payroll/rules", headers=_headers(finance))
    assert defaults.status_code == 200
    assert defaults.json()["recruiter_tier_amounts"] == ["1000", "1500", "2000"]
    assert defaults.json()["extra_dl"] == []

    saved = await app_client.put(
        "/api/hr/payroll/rules",
        json={"extra_dl": [{"dl_user_id": boss.id, "recipient_user_id": worker.id}]},
        headers=_headers(finance),
    )
    assert saved.status_code == 200, saved.text
    assert saved.json()["extra_dl"][0]["recipient_user_id"] == worker.id

    async with AsyncSessionLocal() as db:
        no_profile = await _user(db, "recruiter")
        await db.commit()
    invalid = await app_client.put(
        "/api/hr/payroll/rules",
        json={
            "extra_dl": [{"dl_user_id": boss.id, "recipient_user_id": no_profile.id}]
        },
        headers=_headers(finance),
    )
    assert invalid.status_code == 422
    bad_bounds = await app_client.put(
        "/api/hr/payroll/rules",
        json={"recruiter_tier_bounds": [50, 40]},
        headers=_headers(finance),
    )
    assert bad_bounds.status_code == 422
    assert (
        await app_client.get("/api/hr/payroll/rules", headers=_headers(boss))
    ).status_code == 403


@pytest.mark.asyncio
async def test_automation_since_cannot_be_cleared_over_compass_imports(
    app_client, payroll_on
) -> None:
    # Audyt 19.10.2026: pusta albo wcześniejsza data startu automatu dałaby
    # drugie propozycje za placementy, których premie przyszły z Compassa.
    finance, boss, worker, _other = await _payroll_team()
    y, m = _months_ago(1)
    first = await app_client.put(
        "/api/hr/payroll/rules",
        json={"automation_since": "2026-10-01"},
        headers=_headers(finance),
    )
    assert first.status_code == 200, first.text
    async with AsyncSessionLocal() as db:
        db.add(
            HrBonus(
                recipient_user_id=worker.id,
                proposed_by=boss.id,
                amount=Decimal("100.00"),
                currency="PLN",
                category="recruiter",
                period_year=y,
                period_month=m,
                status="assigned",
                source="compass_import",
                source_key=f"compass:{uuid.uuid4()}",
                compass_id=uuid.uuid4(),
                assigned_by=boss.id,
                assigned_at=local_now(),
            )
        )
        await db.commit()
    for since in (None, "2026-09-01"):
        refused = await app_client.put(
            "/api/hr/payroll/rules",
            json={"automation_since": since},
            headers=_headers(finance),
        )
        assert refused.status_code == 422, refused.text
        assert "Compass" in refused.json()["detail"]
    later = await app_client.put(
        "/api/hr/payroll/rules",
        json={"automation_since": "2026-11-01"},
        headers=_headers(finance),
    )
    assert later.status_code == 200, later.text
    assert later.json()["automation_since"] == "2026-11-01"
    untouched = await app_client.put(
        "/api/hr/payroll/rules", json={"dl_percent": "10"}, headers=_headers(finance)
    )
    assert untouched.status_code == 200
    assert untouched.json()["automation_since"] == "2026-11-01"


@pytest.mark.asyncio
async def test_export_has_amounts_per_currency_and_no_absence_types(
    app_client, payroll_on
) -> None:
    finance, _boss, worker, _other = await _payroll_team()
    y, m = _months_ago(1)
    async with AsyncSessionLocal() as db:
        await _timesheet(db, worker.id, y, m, "approved", [(6, "8", False)])
        await db.commit()
    await app_client.post(
        f"/api/hr/payroll/people/{worker.id}/rates",
        json={"hourly_rate": "100", "effective_from": f"{y}-{m:02d}-01"},
        headers=_headers(finance),
    )
    response = await app_client.get(
        f"/api/hr/payroll/months/{y}/{m}/export.xlsx", headers=_headers(finance)
    )
    assert response.status_code == 200
    assert response.headers["content-disposition"].endswith(f'place-{y}-{m:02d}.xlsx"')
    book = load_workbook(io.BytesIO(response.content), read_only=True)
    rows = list(book["Płace"].iter_rows(values_only=True))
    header = rows[0]
    assert "Dni nieobecności z puli urlopowej" in header
    assert not any("chorob" in str(cell).lower() for row in rows for cell in row)
    mine = [row for row in rows[1:] if row[0] == worker.name]
    assert len(mine) == 1
    assert mine[0][2] == "PLN" and mine[0][3] == 8 and mine[0][8] == 800

    assert (
        await app_client.get(
            f"/api/hr/payroll/months/{y}/{m}/export.xlsx", headers=_headers(worker)
        )
    ).status_code == 403
