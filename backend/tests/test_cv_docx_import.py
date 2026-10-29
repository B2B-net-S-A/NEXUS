"""Plik Word z profilu kandydata → HTML edytora CV firmowego (bez bazy, bez AI).

Screening, 09.10.2026: rekruter wybiera gotowe CV z profilu i edytuje je.
34 z 50 osób w screeningu ma takie CV jako plik Word „…B2B…”, a edytor
przyjmował tylko CV z generatora. Testy pilnują dwóch rzeczy: tekst pliku
nigdy nie ginie, a plik w układzie szablonu firmowego wraca po zapisie
z edytora do tego samego układu.
"""

from __future__ import annotations

import copy
import re
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from fastapi import HTTPException, UploadFile
from lxml import html as lxml_html
from pydantic import ValidationError

from app.services.cv_approved_docx import render_approved_docx
from app.services.cv_document_assets import default_template
from app.services.cv_docx_import import (
    MAX_BYTES,
    MAX_LINES,
    CvImportError,
    docx_to_editor_html,
)
from app.services.cv_generator_b2b.docx_renderer import TRANSLATIONS
from app.services.html_sanitizer import sanitize_cv_html

RODO = "Wyrażam zgodę na przetwarzanie moich danych osobowych w tej rekrutacji."

FIRM_HTML = (
    "<h1>Tester QA – Anna Testowa</h1>"
    '<h2 data-cv-section="why_points">DLACZEGO NASZ KANDYDAT?</h2>'
    "<ul><li>5 lat w <strong>Python</strong></li><li>Automatyzacja testów</li></ul>"
    '<h2 data-cv-section="education">EDUKACJA</h2>'
    "<p>2015 - 2019 <strong>Politechnika Testowa</strong> — Informatyka · Wrocław</p>"
    '<h2 data-cv-section="skills">UMIEJĘTNOŚCI</h2>'
    "<p><strong>Języki: </strong>Python, SQL</p>"
    '<h2 data-cv-section="languages">JĘZYKI</h2><p>Polski - C2 · Angielski - B2</p>'
    '<h2 data-cv-section="experience">DOŚWIADCZENIE</h2>'
    '<p data-cv-section="role"><strong>Senior QA Engineer</strong> 12.2025 - obecnie</p>'
    '<p data-cv-section="employer">Acme · Technologie</p>'
    '<p data-cv-section="duties_label">Zakres zadań:</p>'
    "<ul><li>Rozwój testów w <strong>Python</strong></li><li>Utrzymanie CI</li></ul>"
    '<p data-cv-section="technologies">Technologie: <strong>Python</strong>, Jenkins</p>'
    '<p data-cv-section="role"><strong>QA Engineer</strong> 02.2023 - 11.2025</p>'
    '<p data-cv-section="employer">Beta</p>'
    '<p data-cv-section="duties_label">Zakres zadań:</p>'
    "<ul><li>Automatyzacja testów API</li></ul>"
    f'<p class="rodo">{RODO}</p>'
)


def _firm_docx() -> bytes:
    """Plik w układzie szablonu firmowego — taki, jaki dostaje klient."""
    return render_approved_docx(
        sanitize_cv_html(FIRM_HTML), default_template(), consent=None, language="pl"
    )


def _texts(data: bytes) -> list[str]:
    doc = Document(BytesIO(data))
    out = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            out.append(" | ".join(cell.text.strip() for cell in row.cells))
    return out


def _plain(html: str) -> str:
    return re.sub(r"\s+", " ", lxml_html.fromstring(html).text_content())


def _save(doc) -> bytes:
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def _hand_made_docx() -> bytes:
    """CV pisane ręcznie w Wordzie: inne etykiety niż szablon, tabela, lista."""
    doc = Document()
    doc.add_paragraph("Jan Fikcyjny")
    doc.add_heading("PROFILE", level=1)
    paragraph = doc.add_paragraph()
    paragraph.add_run("Senior engineer with ")
    paragraph.add_run("Kubernetes").bold = True
    paragraph.add_run(" and <b>cloud</b> background, ten years in delivery teams.")
    doc.add_heading("EXPERIENCE", level=1)
    doc.add_paragraph("ACME BANK")
    doc.add_paragraph("Platform Engineer, 2019 - 2024")
    doc.add_paragraph("Built the deployment platform", style="List Bullet")
    doc.add_paragraph("Ran the on-call rotation", style="List Bullet")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "2010 - 2014"
    table.rows[0].cells[1].text = "Example University"
    doc.add_paragraph(
        "Additional notes: open to hybrid work in Warsaw, available from November, "
        "holds a valid security clearance and a driving licence."
    )
    return _save(doc)


def test_firm_layout_survives_word_to_editor_to_word():
    """Plik z szablonu firmowego po wczytaniu i zapisie ma te same akapity."""
    first = _firm_docx()
    imported = docx_to_editor_html(first)
    second = render_approved_docx(
        imported.html, default_template(), consent=None, language=imported.language
    )
    assert _texts(second) == _texts(first)
    assert imported.language == "pl"


def test_firm_layout_comes_back_in_the_generator_shape():
    imported = docx_to_editor_html(_firm_docx())
    html = imported.html
    assert html.startswith('<article class="cv"><h1>Tester QA – Anna Testowa</h1>')
    for key in ("why_points", "education", "skills", "languages", "experience"):
        assert f'<h2 data-cv-section="{key}">' in html
    assert (
        '<p data-cv-section="role"><strong>Senior QA Engineer</strong> '
        "12.2025 - obecnie</p>"
        '<p data-cv-section="employer">Acme · Technologie</p>'
        '<p data-cv-section="duties_label">Zakres zadań:</p>'
        "<ul><li>Rozwój testów w <strong>Python</strong></li>"
    ) in html
    # Etykiety układu („Nazwa firmy:”, „Stanowisko:”) dokłada renderer — w
    # edytorze ich nie ma, jak w CV z generatora.
    assert "Nazwa firmy" not in html and "Stanowisko:" not in html
    assert (
        '<p data-cv-section="technologies">Technologie: <strong>Python</strong>, '
        "Jenkins</p>"
    ) in html
    assert (
        "<p>2015 - 2019 <strong>Politechnika Testowa</strong> — Informatyka · "
        "Wrocław</p>"
    ) in html
    # Nagłówek tabeli edukacji nie staje się wierszem.
    assert TRANSLATIONS["pl"]["education_header"] not in html
    assert imported.stats["roles"] == 2 and imported.stats["sections"] == 5


@pytest.mark.parametrize(
    "header",
    [
        ("", "Nazwa uczelni/kierunek"),
        ("Daty", ""),
        ("Daty", "Nazwa uczelni / kierunek"),
        ("Nazwa uczelni/kierunek", "Daty"),
    ],
)
def test_education_header_edited_in_word_does_not_become_a_row(header):
    """Zgłoszenie 19.10.2026: plik pobrany z NEXUSA, poprawiony w Wordzie
    i wgrany przez „Użyj innego CV” miał w tabeli edukacji wiersz
    „Nazwa uczelni/kierunek”, bo nagłówek bez „Daty” nie był rozpoznany."""
    doc = Document(BytesIO(_firm_docx()))
    cells = doc.tables[0].rows[0].cells
    for cell, text in zip(cells, header):
        cell.text = text
    html = docx_to_editor_html(_save(doc)).html
    assert "uczelni" not in html and "Daty" not in html
    assert (
        "<p>2015 - 2019 <strong>Politechnika Testowa</strong> — Informatyka · "
        "Wrocław</p>"
    ) in html


def test_consent_clause_is_read_from_the_pinned_text_box_once():
    imported = docx_to_editor_html(_firm_docx())
    assert imported.html.count('class="rodo"') == 1
    assert f'<p class="rodo">{RODO}</p>' in imported.html


def test_hand_made_file_keeps_every_paragraph_in_order():
    imported = docx_to_editor_html(_hand_made_docx())
    text = _plain(imported.html)
    expected = [
        "Jan Fikcyjny",
        "PROFILE",
        "Senior engineer with Kubernetes and <b>cloud</b> background",
        "EXPERIENCE",
        "ACME BANK",
        "Platform Engineer, 2019 - 2024",
        "Built the deployment platform",
        "Ran the on-call rotation",
        "2010 - 2014 Example University",
        "Additional notes: open to hybrid work in Warsaw",
    ]
    positions = [text.find(item) for item in expected]
    assert all(position >= 0 for position in positions), positions
    assert positions == sorted(positions)


def test_firm_layout_is_recognised_as_b2b_template():
    imported = docx_to_editor_html(_firm_docx())
    # „Dlaczego nasz kandydat?” + „Nazwa firmy:”/„Stanowisko:” dwóch stanowisk.
    assert imported.stats["template_markers"] >= 3
    assert imported.b2b_template is True


def test_candidate_cv_with_summary_and_experience_is_not_a_b2b_template():
    """Zgłoszenie 29.10.2026: CV kandydatki („Professional Summary”, „Core
    Expertise”, „Experience”) weszło jako CV do klienta 1:1. Te same nazwy
    sekcji ma szablon, ale nie ma go bez „Dlaczego nasz kandydat?” i etykiet
    bloku stanowiska — tak plik wygląda na CV kandydata."""
    doc = Document()
    doc.add_paragraph("Kandydatka Fikcyjna")
    for title, body in (
        ("SUMMARY", "Business analyst with eight years in payments and card schemes."),
        ("CORE EXPERTISE", "Requirements, BPMN, SQL, stakeholder workshops, UAT."),
        ("EXPERIENCE", "Senior Business Analyst, Example Bank, 2019 - 2026"),
        ("EDUCATION", "Example University, Economics"),
        ("SKILLS", "SQL, Jira, Confluence, Enterprise Architect"),
    ):
        doc.add_heading(title, level=1)
        doc.add_paragraph(body)
    buffer = BytesIO()
    doc.save(buffer)

    imported = docx_to_editor_html(buffer.getvalue())

    assert imported.stats["template_markers"] == 0
    assert imported.b2b_template is False


@pytest.mark.parametrize(
    "line",
    ["WHY OUR CANDIDATE?", "Nazwa firmy: Example Bank", "Company: Example Bank"],
)
def test_one_template_label_is_enough_to_skip_the_question(line):
    doc = Document()
    doc.add_paragraph("Kandydat Fikcyjny")
    doc.add_paragraph(line)
    doc.add_paragraph("Experience in payments, card schemes and SQL reporting " * 4)
    buffer = BytesIO()
    doc.save(buffer)

    assert docx_to_editor_html(buffer.getvalue()).b2b_template is True


def test_hand_made_file_keeps_structure_without_guessing_roles():
    imported = docx_to_editor_html(_hand_made_docx())
    html = imported.html
    assert "<h1>Jan Fikcyjny</h1>" in html
    assert "<strong>Kubernetes</strong>" in html
    # Tekst pliku nie staje się znacznikiem HTML.
    assert "&lt;b&gt;cloud&lt;/b&gt;" in html
    assert '<h2 data-cv-section="experience">EXPERIENCE</h2>' in html
    # Nazwa firmy wielkimi literami nie zamyka sekcji doświadczenia.
    assert "<h3>ACME BANK</h3>" in html
    assert "<h3>PROFILE</h3>" in html
    assert (
        "<ul><li>Built the deployment platform</li>"
        "<li>Ran the on-call rotation</li></ul>"
    ) in html
    assert 'data-cv-section="role"' not in html
    assert imported.stats["roles"] == 0
    assert imported.language == "en"


def test_missing_consent_clause_gets_the_standard_one_in_the_file_language():
    imported = docx_to_editor_html(_hand_made_docx())
    clause = TRANSLATIONS["en"]["rodo"]
    assert imported.html.count('class="rodo"') == 1
    assert clause[:60] in _plain(imported.html)


def test_text_in_links_and_form_fields_is_not_lost():
    doc = Document()
    doc.add_paragraph("Maria Przykładowa")
    linked = doc.add_paragraph("Portfolio: ")
    hyperlink = OxmlElement("w:hyperlink")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "projekty na GitHubie"
    run.append(text)
    hyperlink.append(run)
    linked._p.append(hyperlink)
    wrapped = doc.add_paragraph(
        "Akapit w polu formularza Worda — dziesięć lat pracy w zespołach wdrożeniowych, "
        "głównie w bankowości i telekomunikacji, ostatnio jako lider zespołu."
    )
    control = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    control.append(content)
    wrapped._p.addprevious(control)
    content.append(wrapped._p)
    doc.add_paragraph("Dodatkowy opis doświadczenia i kompetencji kandydatki. " * 3)
    imported = docx_to_editor_html(_save(doc))
    text_out = _plain(imported.html)
    assert "Portfolio: projekty na GitHubie" in text_out
    assert "Akapit w polu formularza Worda" in text_out


def test_text_box_content_does_not_leak_into_the_paragraph_it_is_anchored_in():
    data = _firm_docx()
    doc = Document(BytesIO(data))
    boxes = [
        element
        for element in doc.element.body.iter()
        if str(element.tag).endswith("}txbxContent")
    ]
    assert boxes, "szablon firmowy trzyma klauzulę zgody w polu tekstowym"
    imported = docx_to_editor_html(data)
    # Klauzula jest raz — jako klauzula, nie jako zwykły akapit treści.
    assert _plain(imported.html).count(RODO) == 1


@pytest.mark.parametrize(
    ("data", "fragment"),
    [
        (b"%PDF-1.7 fake pdf bytes", "tylko plik Word"),
        (b"PK\x03\x04 not really a zip", "Nie udało się odczytać"),
        (b"PK" + b"0" * (MAX_BYTES + 1), "za duży"),
    ],
    # Bez nazw pytest wpisuje w identyfikator testu całe 10 MB bajtów.
    ids=["pdf", "uszkodzony-zip", "za-duzy"],
)
def test_files_that_cannot_be_read_are_refused_in_polish(data, fragment):
    with pytest.raises(CvImportError) as error:
        docx_to_editor_html(data)
    assert fragment in str(error.value)


def _long_text() -> str:
    return (
        "Rozwój usług płatniczych w banku i integracje z systemami zewnętrznymi. " * 4
    )


def test_thousands_of_empty_paragraphs_are_refused_before_reading():
    """Pusty akapit nie daje wiersza, ale kosztuje: plik z dziesiątkami tysięcy
    pustych akapitów zajmował konwersję na wiele sekund (przegląd bezpieczeństwa
    13.10.2026). Limit liczy wszystkie akapity i wiersze tabel, także puste."""
    import time

    doc = Document()
    body = doc.element.body
    first = doc.add_paragraph("")._p
    for _ in range(12_000):
        body.insert(0, copy.deepcopy(first))
    doc.add_paragraph(_long_text())
    data = _save(doc)
    started = time.process_time()
    with pytest.raises(CvImportError) as refused:
        docx_to_editor_html(data)
    assert "zbyt długi" in str(refused.value)
    assert time.process_time() - started < 2.0


def test_thousands_of_styles_are_refused():
    """Każdy akapit szuka swojego stylu w `styles.xml`; tysiące stylów mnożyły
    koszt każdego akapitu."""
    doc = Document()
    styles = doc.styles.element
    base = styles.findall(qn("w:style"))[0]
    for index in range(1_200):
        clone = copy.deepcopy(base)
        clone.set(qn("w:styleId"), f"Zbedny{index}")
        styles.append(clone)
    doc.add_paragraph(_long_text())
    with pytest.raises(CvImportError) as refused:
        docx_to_editor_html(_save(doc))
    assert "stylów" in str(refused.value)


def test_paragraph_style_is_remembered_per_style_not_per_file():
    """Pogrubienie stylu czytamy raz na styl — ale każdy styl osobno."""
    from docx.enum.style import WD_STYLE_TYPE

    doc = Document()
    strong = doc.styles.add_style("Mocny akapit", WD_STYLE_TYPE.PARAGRAPH)
    strong.font.bold = True
    doc.styles.add_style("Zwykly akapit", WD_STYLE_TYPE.PARAGRAPH)
    # Pierwsza linia pliku to tytuł CV — porównujemy akapity treści.
    doc.add_paragraph("Senior Java Developer – Anna Przykładowa")
    doc.add_paragraph(
        "Pierwszy akapit o rozwoju usług płatniczych w banku.", style="Mocny akapit"
    )
    doc.add_paragraph(
        "Drugi akapit o integracjach z systemami zewnętrznymi.", style="Zwykly akapit"
    )
    doc.add_paragraph(
        "Trzeci akapit o testach jednostkowych i integracyjnych.", style="Mocny akapit"
    )
    doc.add_paragraph(_long_text())
    html = docx_to_editor_html(_save(doc)).html
    bold = " ".join(re.findall(r"<strong>(.*?)</strong>", html))
    assert "Pierwszy akapit" in bold and "Trzeci akapit" in bold
    assert "Drugi akapit" not in bold


def test_empty_document_is_refused():
    doc = Document()
    doc.add_paragraph("CV")
    with pytest.raises(CvImportError) as error:
        docx_to_editor_html(_save(doc))
    assert "nie ma tekstu" in str(error.value)


def test_role_block_without_position_stays_as_typed():
    """Daty i „Nazwa firmy:” bez stanowiska nie znikają ani nie udają roli."""
    doc = Document(BytesIO(_firm_docx()))
    body = doc.element.body
    target = next(p for p in doc.paragraphs if p.text.startswith("Stanowisko:"))
    filler = copy.deepcopy(target._p)
    for node in filler.iter(qn("w:t")):
        node.text = ""
    body.replace(target._p, filler)
    imported = docx_to_editor_html(_save(doc))
    text = _plain(imported.html)
    assert "12.2025 - obecnie" in text
    assert "Nazwa firmy: Acme · Technologie" in text
    assert imported.stats["roles"] == 1


# ── plik jest niezaufany (formularz kariery, import) ─────────────────────────


def test_huge_grid_span_does_not_multiply_cells():
    """``row.cells`` powtarza komórkę ``gridSpan`` razy; liczba idzie z pliku.

    Kilkukilobajtowy DOCX z ``gridSpan=2000000000`` budowałby w procesie
    aplikacji listę na gigabajty. Import czyta komórki wprost z XML-a.
    """
    doc = Document(BytesIO(_hand_made_docx()))
    cell = doc.tables[0].rows[0].cells[1]
    properties = cell._tc.get_or_add_tcPr()
    span = OxmlElement("w:gridSpan")
    span.set(qn("w:val"), "2000000000")
    properties.append(span)
    imported = docx_to_editor_html(_save(doc))
    assert _plain(imported.html).count("Example University") == 1


def test_vertically_merged_continuation_cell_is_not_read_twice():
    doc = Document()
    doc.add_paragraph("Opis kandydata i jego doświadczenia zawodowego. " * 6)
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Scalona komórka"
    table.rows[0].cells[1].text = "Pierwszy wiersz"
    table.rows[1].cells[1].text = "Drugi wiersz"
    for row, value in ((0, "restart"), (1, "continue")):
        merge = OxmlElement("w:vMerge")
        merge.set(qn("w:val"), value)
        table.rows[row].cells[0]._tc.get_or_add_tcPr().append(merge)
    text = _plain(docx_to_editor_html(_save(doc)).html)
    assert text.count("Scalona komórka") == 1
    assert "Pierwszy wiersz" in text and "Drugi wiersz" in text


def test_document_with_too_many_lines_is_refused():
    doc = Document()
    for index in range(MAX_LINES + 1):
        doc.add_paragraph(f"Linia {index}")
    with pytest.raises(CvImportError) as error:
        docx_to_editor_html(_save(doc))
    assert "zbyt długi" in str(error.value)


# ── Podgląd importu w edytorze CV: `…/import-preview` (Edytor CV 2, PR 5) ────
#
# Trasa zamienia plik Word na HTML edytora i NICZEGO nie zapisuje — treść
# przyjmuje edytor, a utrwala zwykły autozapis. Tu logika bez bazy (atrapa
# sesji); dostęp i brak zapisu na prawdziwej bazie: `test_stage_cv_copy.py`.


class _Upload(BytesIO):
    """Plik z formularza, który pamięta, o ile bajtów go poproszono."""

    def __init__(self, data: bytes) -> None:
        super().__init__(data)
        self.asked: list[int] = []

    def read(self, size: int = -1) -> bytes:
        self.asked.append(size)
        return super().read(size)


def _preview_db(job_id: int | None = 4) -> AsyncMock:
    db = AsyncMock()
    db.scalar.return_value = job_id
    db.add = Mock()
    return db


async def _preview(monkeypatch, data: bytes, name: str = "CV.docx", **overrides):
    from app.api import cv_editor_import as api

    scope = overrides.pop("scope", None) or AsyncMock()
    monkeypatch.setattr(api, "ensure_job_membership", scope)
    source = _Upload(data)
    call = {
        "request": SimpleNamespace(),
        "stage_id": 2,
        "current_user": SimpleNamespace(id=9),
        "file": UploadFile(file=source, filename=name),
        "db": _preview_db(),
        **overrides,
    }
    try:
        # Sam handler, bez dekoratora limitu (ten potrzebuje prawdziwego żądania).
        result = await api.preview_stage_cv_import.__wrapped__(**call)
    except HTTPException as refused:
        result = refused
    return result, call["db"], source, scope


async def test_import_preview_returns_editor_html_and_writes_nothing(monkeypatch):
    data = _hand_made_docx()

    result, db, source, scope = await _preview(
        monkeypatch, data, "C:\\Pulpit\\CV_od_klienta.docx"
    )

    # Ta sama konwersja co `…/upload` — z sanitizerem w środku.
    assert result.model_dump() == {
        "content_html": docx_to_editor_html(data).html,
        "filename": "CV_od_klienta.docx",
        # Zwykłe CV kandydata — okno „Użyj gotowego CV” zapyta o generację.
        "b2b_template": False,
    }
    assert "<b>cloud</b>" not in result.content_html
    # Zakres jak przy wgraniu pliku: zespół rekrutacji tego etapu.
    scope.assert_awaited_once_with(db, SimpleNamespace(id=9), 4)
    assert source.asked == [MAX_BYTES + 1]
    # Jedyny kontakt z bazą to odczyt rekrutacji etapu.
    assert db.scalar.await_count == 1
    assert not db.add.called
    for write in (db.execute, db.flush, db.commit, db.refresh, db.delete):
        write.assert_not_awaited()


@pytest.mark.parametrize(
    ("data", "fragment"),
    [
        (b"%PDF-1.7 to nie jest Word", "tylko plik Word"),
        (b"PK\x03\x04 to nie jest archiwum", "Nie udało się odczytać"),
        (b"", "pusty"),
    ],
    ids=["pdf", "uszkodzony-zip", "pusty"],
)
async def test_import_preview_refuses_unreadable_files_in_polish(
    monkeypatch, data, fragment
):
    refused, db, _, _ = await _preview(monkeypatch, data)

    assert isinstance(refused, HTTPException)
    assert refused.status_code == 422
    assert fragment in refused.detail
    # Komunikat jest dla rekrutera: bez nazw wyjątków i śladu stosu.
    assert "Traceback" not in refused.detail and "Error" not in refused.detail
    db.commit.assert_not_awaited()


async def test_import_preview_reads_the_upload_with_a_ceiling(monkeypatch):
    from app.services import cv_docx_import

    monkeypatch.setattr(cv_docx_import, "MAX_BYTES", 64)

    refused, _, source, _ = await _preview(monkeypatch, _hand_made_docx())

    assert isinstance(refused, HTTPException) and refused.status_code == 422
    assert "za duży" in refused.detail
    # Jeden odczyt o bajt ponad limit — reszta pliku nie trafia do pamięci.
    assert source.asked == [65]


async def test_import_preview_refuses_content_the_editor_could_not_save(monkeypatch):
    """Zapis z edytora (PATCH) ma limit znaków — dłuższego importu nie dałoby
    się utrwalić, więc podgląd odmawia od razu, zamiast dać treść do stracenia."""
    from app.api import cv_editor_import as api
    from app.schemas.candidate_stage_cv import CV_CONTENT_HTML_MAX, CVBrandedUpdate

    assert api.CV_CONTENT_HTML_MAX == CV_CONTENT_HTML_MAX == 500_000
    CVBrandedUpdate(expected_revision=0, content_html="x" * CV_CONTENT_HTML_MAX)
    with pytest.raises(ValidationError):
        CVBrandedUpdate(
            expected_revision=0, content_html="x" * (CV_CONTENT_HTML_MAX + 1)
        )

    data = _hand_made_docx()
    length = len(docx_to_editor_html(data).html)
    monkeypatch.setattr(api, "CV_CONTENT_HTML_MAX", length)
    fits, _, _, _ = await _preview(monkeypatch, data)
    monkeypatch.setattr(api, "CV_CONTENT_HTML_MAX", length - 1)
    refused, db, _, _ = await _preview(monkeypatch, data)

    assert len(fits.content_html) == length
    assert isinstance(refused, HTTPException) and refused.status_code == 422
    assert "za długa" in refused.detail
    db.commit.assert_not_awaited()


async def test_import_preview_checks_the_stage_before_touching_the_file(monkeypatch):
    missing, _, unread, scope = await _preview(
        monkeypatch, _hand_made_docx(), db=_preview_db(job_id=None)
    )
    outsider, _, untouched, _ = await _preview(
        monkeypatch,
        _hand_made_docx(),
        scope=AsyncMock(side_effect=HTTPException(403, "Brak dostępu.")),
    )

    assert isinstance(missing, HTTPException) and missing.status_code == 404
    scope.assert_not_awaited()
    assert isinstance(outsider, HTTPException) and outsider.status_code == 403
    assert unread.asked == [] and untouched.asked == []


async def test_import_preview_over_http_takes_a_form_file_and_keeps_its_limit(
    monkeypatch,
):
    """Cała trasa przez ASGI: formularz z plikiem, bramka roli i limit żądań.

    Testy wyżej wołają sam handler, więc nie widzą dekoratora limitu ani tego,
    jak FastAPI czyta parametry trasy pod tym dekoratorem (plik z formularza,
    bramka `Annotated`) — poprawne żądanie nie może skończyć się 422.
    """
    import httpx

    from app.api import cv_editor_import as api
    from app.api import deps
    from app.core.database import get_db
    from app.core.rate_limit import limiter
    from app.main import app
    from tests._authz_matrix import PERSONAS, build_user

    personas = {persona.key: persona for persona in PERSONAS}
    recruiter = build_user(personas["recruiter"], user_id=900_001)
    viewer = build_user(personas["user"], user_id=900_002)
    holder = {"user": recruiter}
    db = _preview_db()

    async def _current_user():
        return holder["user"]

    async def _database():
        return db

    monkeypatch.setattr(api, "ensure_job_membership", AsyncMock())
    monkeypatch.setattr(limiter, "enabled", True)
    monkeypatch.setitem(app.dependency_overrides, deps.get_current_user, _current_user)
    monkeypatch.setitem(app.dependency_overrides, get_db, _database)
    url = "/api/candidates/stages/2/cv/branded/import-preview"
    data = _hand_made_docx()
    mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    async def send(client: httpx.AsyncClient) -> httpx.Response:
        return await client.post(
            url, files={"file": ("CV po poprawkach.docx", data, mime)}
        )

    limiter.reset()
    try:
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=app), base_url="http://cv-import"
        ) as client:
            accepted = await send(client)
            without_file = await client.post(url, json={"file": "CV.docx"})
            holder["user"] = viewer
            refused = await send(client)
            holder["user"] = recruiter
            burst = [(await send(client)).status_code for _ in range(10)]
    finally:
        limiter.reset()

    assert accepted.status_code == 200, accepted.text
    assert accepted.json() == {
        "content_html": docx_to_editor_html(data).html,
        "filename": "CV po poprawkach.docx",
        "b2b_template": docx_to_editor_html(data).b2b_template,
    }
    assert without_file.status_code == 422
    # Konto tylko do odczytu odpada na bramce, zanim plik zostanie przeczytany.
    assert refused.status_code == 403
    # 10 na minutę na osobę; odmowy bramki i walidacji nie zużywają limitu.
    assert burst == [200] * 9 + [429]
    assert not db.add.called
    for write in (db.execute, db.flush, db.commit):
        write.assert_not_awaited()
