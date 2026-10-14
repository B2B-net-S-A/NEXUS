"""Plik CV dla klienta bez danych autora z szablonu (14.10.2026).

Szablon firmowy CV (a także zamrożony szablon etapu i szablon z reguły
klienta) niesie w ``docProps/core.xml`` login i imię i nazwisko osoby, która
go ostatnio zapisała. Oba renderery Worda zachowywały te pola, więc każdy
plik CV u klienta — i PDF zrobiony z tego pliku — pokazywał we
„Właściwościach” pracownika firmy. Renderery czyszczą teraz pola opisowe
dokumentu przed zapisem.
"""

from io import BytesIO

import pytest
from docx import Document

from app.services.cv_approved_docx import render_approved_docx
from app.services.cv_document_assets import default_template
from app.services.cv_generator_b2b.docx_renderer import render_cv_to_bytes
from app.services.html_sanitizer import sanitize_cv_html

# Pola tekstowe właściwości dokumentu; język i daty zostają.
FIELDS = (
    "author",
    "last_modified_by",
    "title",
    "subject",
    "comments",
    "keywords",
    "category",
    "content_status",
    "identifier",
    "version",
)

HTML = (
    "<h1>Tester QA – Anna Testowa</h1><hr>"
    '<h2 data-cv-section="why_points">DLACZEGO NASZ KANDYDAT?</h2>'
    "<ul><li><p>5 lat w <strong>Python</strong></p></li></ul>"
    '<h2 data-cv-section="experience">DOŚWIADCZENIE</h2>'
    '<p data-cv-section="role"><strong>QA Engineer</strong> 02.2023 - obecnie</p>'
    '<p data-cv-section="employer">Acme · Telekomunikacja</p>'
    '<p data-cv-section="duties_label">Zakres zadań:</p>'
    "<ul><li><p>Automatyzacja testów API</p></li></ul>"
)


def _payload() -> dict:
    return {
        "name": "Anna Testowa",
        "first_name": "Anna",
        "position": "QA Engineer",
        "language": "pl",
        "why_points": ["5 lat w Python"],
        "skills": [{"label": "Testy:", "content": "Pytest, Selenium"}],
        "languages": ["Polski – ojczysty"],
        "experience": [
            {
                "dates": "02.2023 – obecnie",
                "company": "Acme",
                "industry": "Telekomunikacja",
                "position": "QA Engineer",
                "responsibilities": ["Automatyzacja testów API"],
                "technologies": ["Python"],
            }
        ],
    }


def _company_template() -> bytes:
    return default_template()


def _client_template() -> bytes:
    """Szablon jak z Worda klienta: każde pole opisowe wypełnione (dane fikcyjne)."""
    doc = Document(BytesIO(default_template()))
    for name in FIELDS:
        setattr(doc.core_properties, name, f"Jan Przykładowy ({name})")
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


TEMPLATES = pytest.mark.parametrize(
    "template", [_company_template, _client_template], ids=["firmowy", "klienta"]
)


def _props(docx: bytes) -> dict[str, str]:
    properties = Document(BytesIO(docx)).core_properties
    return {name: getattr(properties, name) for name in FIELDS}


@TEMPLATES
def test_editor_docx_drops_template_author(template):
    docx = render_approved_docx(sanitize_cv_html(HTML), template())
    assert _props(docx) == dict.fromkeys(FIELDS, "")


@TEMPLATES
def test_generator_docx_drops_template_author(template):
    docx = render_cv_to_bytes(_payload(), BytesIO(template()))
    assert _props(docx) == dict.fromkeys(FIELDS, "")


def test_language_of_template_stays():
    """Czyszczenie nie rusza języka dokumentu (sprawdzanie pisowni w Wordzie)."""
    doc = Document(BytesIO(default_template()))
    doc.core_properties.language = "pl-PL"
    out = BytesIO()
    doc.save(out)
    docx = render_approved_docx(sanitize_cv_html(HTML), out.getvalue())
    assert Document(BytesIO(docx)).core_properties.language == "pl-PL"
