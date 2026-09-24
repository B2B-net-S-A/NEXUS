"""Synchronizacja kontrakt ↔ zamówienia (ticket 09.2026).

Zgłoszenie: kontrakt Bartosza Czapelki (Alior) stał jako „Szkic" ze stawką
kosztową 120 zł/h, bez stawki przychodowej i okresu zamówienia, choć jego
zamówienie OIT/0189/2026/ITVM miało okres 15.09–31.12.2026 i 1340 PLN/MD.

Testy opisują reguły z ticketu, każdą na liczbach z ticketu:

* zamówienie → kontrakt: okres zamówienia (osobne pole), stawka przychodowa
  od daty startu zamówienia, jednostka dopasowana do zamówienia (1 MD = 8 h),
  ale kontrakt nigdy nie przechodzi na MD — zamówienie w MD daje zł/h
  (ticket „Ujednolicenie stawek w module Kontrakty", 14.09.2026);
* kontrakt → zamówienie: stawka kosztowa z kontraktu, KAŻDA zaplanowana
  podwyżka w swoim dniu;
* jednorazowo: migawka raportu przed synchronizacją + żaden szkic.
"""

from __future__ import annotations

import io
import uuid
from datetime import date, timedelta
from decimal import Decimal

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import delete, select

from app.models.activity import Activity
from app.models.client_order import ClientOrder, ClientOrderStatus
from app.models.contract import Contract, ContractStatus, ContractType, RateUnit
from app.models.contract_candidate_rate import ContractCandidateRate
from app.models.contract_client_rate import ContractClientRate
from app.services.contract_order_sync import (
    apply_manual_client_rate,
    convert_rate_between,
    cost_reference_day,
    order_revenue_terms,
    sync_contract_from_orders,
    sync_orders_cost_from_contract,
)
from app.core.scheduling import business_today

pytestmark = pytest.mark.asyncio


class _FakeDb:
    """Wystarcza czystym przebiegom: aktywacja dopisuje wiersz audytu."""

    def __init__(self) -> None:
        self.added: list[object] = []

    def add(self, obj: object) -> None:
        self.added.append(obj)


def _contract(**overrides) -> Contract:
    """Kontrakt z generatora B2B: godzinowy, 120 zł/h, start 14.09.2026."""
    contract = Contract(
        id=overrides.pop("id", 1),
        client_id=overrides.pop("client_id", 10),
        candidate_id=5,
        contract_type=ContractType.b2b,
        status=overrides.pop("status", ContractStatus.draft),
        start_date=overrides.pop("start_date", date(2026, 9, 14)),
        end_date=None,
        rate_candidate=overrides.pop("rate_candidate", Decimal("120.000")),
        rate_client=overrides.pop("rate_client", None),
        rate_unit=overrides.pop("rate_unit", RateUnit.hourly),
        billing_hours_per_month=160,
        currency="PLN",
        rate_candidate_currency="PLN",
    )
    contract.candidate_rate_schedule = overrides.pop("candidate_rate_schedule", [])
    contract.client_rate_schedule = overrides.pop("client_rate_schedule", [])
    contract.framework_rate_schedule = []
    for key, value in overrides.items():
        setattr(contract, key, value)
    return contract


def _order(order_id: int = 100, **overrides) -> ClientOrder:
    """Zamówienie Aliora z ticketu: 15.09–31.12.2026, 1340 PLN/MD."""
    values = {
        "id": order_id,
        "client_id": 10,
        "contract_id": 1,
        "title": "OIT/0189/2026/ITVM",
        "status": ClientOrderStatus.active,
        "start_date": date(2026, 9, 15),
        "end_date": date(2026, 12, 31),
        "rate_client": Decimal("1340.000"),
        "rate_candidate": None,
        "rate_unit": RateUnit.daily,
        "billing_hours_per_month": 160,
        "rate_client_currency": "PLN",
        "rate_candidate_currency": "PLN",
        "currency": "PLN",
        "order_group_id": None,
    }
    values.update(overrides)
    return ClientOrder(**values)


# ── Przelicznik 1 MD = 8 h ───────────────────────────────────────────────────


def test_md_hour_conversion_uses_eight_hours_both_ways():
    assert convert_rate_between(
        Decimal("120"), RateUnit.hourly, RateUnit.daily
    ) == Decimal("960.000")
    assert convert_rate_between(
        Decimal("1340"), RateUnit.daily, RateUnit.hourly
    ) == Decimal("167.500")


def test_hour_to_month_uses_the_monthly_side_hours():
    """Liczba godzin miesiąca należy do strony MIESIĘCZNEJ, nie do wołającego."""
    assert convert_rate_between(
        Decimal("100"), RateUnit.hourly, RateUnit.monthly, to_hours=168
    ) == Decimal("16800.000")
    assert convert_rate_between(
        Decimal("16800"), RateUnit.monthly, RateUnit.hourly, from_hours=168
    ) == Decimal("100.000")


# ── Kiedy zamówienie jest „uzupełnione" ──────────────────────────────────────


def test_auto_draft_without_revenue_rate_does_not_pose_as_an_order_period():
    """Auto-szkic z podpisu ma start umowy i PUSTĄ stawkę klienta."""
    assert order_revenue_terms(_order(rate_client=None)) is None
    assert order_revenue_terms(_order(start_date=None)) is None
    assert order_revenue_terms(_order(status=ClientOrderStatus.cancelled)) is None
    terms = order_revenue_terms(_order())
    assert terms is not None and terms.unit == RateUnit.daily


def test_group_line_revenue_is_per_md():
    line = _order(order_group_id=7, rate_client=None, md_rate_revenue=Decimal("1200"))
    terms = order_revenue_terms(line)
    assert terms is not None
    assert (terms.rate, terms.unit, terms.currency) == (
        Decimal("1200"),
        RateUnit.daily,
        "PLN",
    )


# ── Zamówienie → kontrakt ────────────────────────────────────────────────────


async def test_czapelka_ticket_example_end_to_end_on_the_contract():
    """Szkic 120 zł/h + zamówienie 1340 PLN/MD → Aktywny, zł/h, okres zamówienia."""
    contract = _contract()
    db = _FakeDb()

    outcome = await sync_contract_from_orders(
        db, contract, [_order()], actor_id=None, today=date(2026, 9, 20)
    )

    assert contract.rate_unit == RateUnit.hourly, "kontrakt nigdy nie jest w MD"
    assert contract.rate_candidate == Decimal("120.000"), "koszt z umowy bez zmian"
    assert contract.effective_client_rate(date(2026, 9, 20)) == Decimal("167.5")
    assert contract.rate_client == Decimal("167.5"), "1340 zł/MD ÷ 8"
    # Miesiąc roboczy = 21 MD = 168 h (``app.core.work_time``): przychód
    # miesięczny kontraktu jest dokładnie tym, co zamówienie (1340 × 21).
    assert contract.billing_hours_per_month == 168
    assert contract.orders_in_md is True
    assert contract.monthly_rate(contract.rate_client) == Decimal("1340") * 21
    assert (contract.client_order_start_date, contract.client_order_end_date) == (
        date(2026, 9, 15),
        date(2026, 12, 31),
    )
    # Okres UMOWY nietknięty — okres zamówienia to osobne pole.
    assert (contract.start_date, contract.end_date) == (date(2026, 9, 14), None)
    assert contract.status == ContractStatus.active
    assert outcome.activated and outcome.unit_switched_to is None


async def test_hourly_order_keeps_the_hourly_contract_unit():
    contract = _contract()
    order = _order(rate_unit=RateUnit.hourly, rate_client=Decimal("167.500"))

    await sync_contract_from_orders(
        _FakeDb(), contract, [order], actor_id=None, today=date(2026, 9, 20)
    )

    assert contract.rate_unit == RateUnit.hourly
    assert contract.rate_candidate == Decimal("120.000")
    assert contract.rate_client == Decimal("167.500")


async def test_future_revenue_rate_applies_only_from_its_date():
    """Ticket: 120 zł obecnie, 130 zł od 01.10.2026 — w kontrakcie od 01.10."""
    contract = _contract(status=ContractStatus.active)
    current = _order(
        100,
        rate_unit=RateUnit.hourly,
        rate_client=Decimal("120"),
        end_date=date(2026, 9, 30),
    )
    future = _order(
        101,
        rate_unit=RateUnit.hourly,
        rate_client=Decimal("130"),
        start_date=date(2026, 10, 1),
        end_date=date(2026, 12, 31),
    )

    await sync_contract_from_orders(
        _FakeDb(), contract, [current, future], actor_id=None, today=date(2026, 9, 20)
    )

    assert contract.effective_client_rate(date(2026, 9, 30)) == Decimal("120.000")
    assert contract.effective_client_rate(date(2026, 10, 1)) == Decimal("130.000")
    assert contract.rate_client == Decimal("120.000"), "cache = stawka na dziś"


async def test_next_order_overwrites_the_order_period():
    contract = _contract(status=ContractStatus.active)
    first = _order(100)
    await sync_contract_from_orders(
        _FakeDb(), contract, [first], actor_id=None, today=date(2026, 9, 20)
    )
    extension = _order(101, start_date=date(2027, 1, 1), end_date=date(2027, 6, 30))

    await sync_contract_from_orders(
        _FakeDb(), contract, [first, extension], actor_id=None, today=date(2026, 9, 20)
    )

    assert (contract.client_order_start_date, contract.client_order_end_date) == (
        date(2027, 1, 1),
        date(2027, 6, 30),
    )


async def test_sync_is_idempotent_and_keeps_one_step_per_order():
    contract = _contract()
    orders = [_order()]
    await sync_contract_from_orders(
        _FakeDb(), contract, orders, actor_id=None, today=date(2026, 9, 20)
    )

    again = await sync_contract_from_orders(
        _FakeDb(), contract, orders, actor_id=None, today=date(2026, 9, 20)
    )

    assert not again.changed
    assert [s.source_order_id for s in contract.client_rate_schedule] == [100]


async def test_order_rate_correction_moves_the_same_step():
    contract = _contract()
    order = _order()
    await sync_contract_from_orders(
        _FakeDb(), contract, [order], actor_id=None, today=date(2026, 9, 20)
    )
    order.rate_client = Decimal("1400")

    await sync_contract_from_orders(
        _FakeDb(), contract, [order], actor_id=None, today=date(2026, 9, 20)
    )

    assert len(contract.client_rate_schedule) == 1
    assert contract.rate_client == Decimal("175"), "1400 zł/MD ÷ 8"


async def test_cancelled_order_takes_its_rate_and_period_back():
    contract = _contract()
    order = _order()
    await sync_contract_from_orders(
        _FakeDb(), contract, [order], actor_id=None, today=date(2026, 9, 20)
    )
    order.status = ClientOrderStatus.cancelled

    await sync_contract_from_orders(
        _FakeDb(), contract, [order], actor_id=None, today=date(2026, 9, 20)
    )

    assert contract.client_rate_schedule == []
    assert contract.client_order_start_date is None
    assert contract.client_order_end_date is None


async def test_md_then_hourly_orders_keep_the_contract_in_hours():
    """Kontrakt zostaje w zł/h przy zamówieniu w MD i przy kolejnym godzinowym."""
    contract = _contract()
    md_order = _order(100)
    await sync_contract_from_orders(
        _FakeDb(), contract, [md_order], actor_id=None, today=date(2026, 9, 20)
    )
    hourly_extension = _order(
        101,
        rate_unit=RateUnit.hourly,
        rate_client=Decimal("170"),
        start_date=date(2027, 1, 1),
        end_date=date(2027, 3, 31),
    )

    await sync_contract_from_orders(
        _FakeDb(),
        contract,
        [md_order, hourly_extension],
        actor_id=None,
        today=date(2026, 9, 20),
    )

    assert contract.rate_unit == RateUnit.hourly
    assert contract.rate_candidate == Decimal("120.000")
    # Stawka ze starszego zamówienia MD przeliczona na godziny.
    assert contract.effective_client_rate(date(2026, 10, 1)) == Decimal("167.500")
    assert contract.effective_client_rate(date(2027, 1, 1)) == Decimal("170.000")


async def test_explicit_contract_unit_wins_over_the_order_unit():
    contract = _contract(status=ContractStatus.active)

    await sync_contract_from_orders(
        _FakeDb(),
        contract,
        [_order()],
        actor_id=None,
        today=date(2026, 9, 20),
        follow_order_unit=False,
    )

    assert contract.rate_unit == RateUnit.hourly
    assert contract.rate_client == Decimal("167.500"), "1340 PLN/MD ÷ 8"


async def test_legacy_revenue_column_survives_as_history_before_first_order():
    contract = _contract(
        status=ContractStatus.active,
        rate_client=Decimal("150"),
        start_date=date(2026, 1, 1),
    )
    order = _order(rate_unit=RateUnit.hourly, rate_client=Decimal("170"))

    await sync_contract_from_orders(
        _FakeDb(), contract, [order], actor_id=None, today=date(2026, 9, 20)
    )

    assert contract.effective_client_rate(date(2026, 3, 1)) == Decimal("150")
    assert contract.effective_client_rate(date(2026, 9, 20)) == Decimal("170.000")


async def test_void_contract_is_never_touched():
    contract = _contract(status=ContractStatus.void)

    outcome = await sync_contract_from_orders(
        _FakeDb(), contract, [_order()], actor_id=None, today=date(2026, 9, 20)
    )

    assert not outcome.changed
    assert contract.rate_unit == RateUnit.hourly


def test_unchanged_manual_revenue_from_a_full_form_adds_no_step():
    """Formularz wysyła całą stawkę także wtedy, gdy nikt jej nie ruszał."""
    contract = _contract(
        client_rate_schedule=[
            ContractClientRate(
                rate=Decimal("1340"),
                effective_from=date(2026, 9, 15),
                source_order_id=100,
            )
        ],
        rate_unit=RateUnit.daily,
    )
    today = date(2026, 10, 5)

    assert not apply_manual_client_rate(
        contract, Decimal("1340"), actor_id=None, today=today
    )
    assert apply_manual_client_rate(
        contract, Decimal("1390"), actor_id=None, today=today
    )
    assert contract.effective_client_rate(today) == Decimal("1390")


# ── Kontrakt → zamówienie (stawka kosztowa) ──────────────────────────────────


def _raise_schedule() -> list[ContractCandidateRate]:
    """Ticket: 120 zł, od 01.10.2026 → 125 zł, od 01.12.2026 → 130 zł."""
    return [
        ContractCandidateRate(rate=Decimal("120"), effective_from=date(2026, 9, 14)),
        ContractCandidateRate(rate=Decimal("125"), effective_from=date(2026, 10, 1)),
        ContractCandidateRate(rate=Decimal("130"), effective_from=date(2026, 12, 1)),
    ]


@pytest.mark.parametrize(
    ("today", "expected"),
    [
        (date(2026, 9, 30), Decimal("120.000")),
        (date(2026, 10, 1), Decimal("125.000")),
        (date(2026, 11, 30), Decimal("125.000")),
        (date(2026, 12, 1), Decimal("130.000")),
    ],
)
async def test_every_planned_raise_reaches_the_order_on_its_day(today, expected):
    contract = _contract(
        status=ContractStatus.active, candidate_rate_schedule=_raise_schedule()
    )
    order = _order(rate_unit=RateUnit.hourly, rate_client=Decimal("167.5"))

    changed = await sync_orders_cost_from_contract(
        _FakeDb(), contract, [order], today=today
    )

    assert changed == [100]
    assert order.rate_candidate == expected


async def test_planned_raises_are_converted_to_the_md_order_unit():
    contract = _contract(
        status=ContractStatus.active, candidate_rate_schedule=_raise_schedule()
    )
    order = _order()  # PLN/MD

    await sync_orders_cost_from_contract(
        _FakeDb(), contract, [order], today=date(2026, 10, 1)
    )
    assert order.rate_candidate == Decimal("1000.000"), "125 zł/h × 8"
    await sync_orders_cost_from_contract(
        _FakeDb(), contract, [order], today=date(2026, 12, 1)
    )
    assert order.rate_candidate == Decimal("1040.000"), "130 zł/h × 8"


async def test_manual_cost_in_the_order_is_overwritten_by_the_contract():
    contract = _contract(status=ContractStatus.active)
    order = _order(rate_candidate=Decimal("999"))

    await sync_orders_cost_from_contract(
        _FakeDb(), contract, [order], today=date(2026, 9, 20)
    )

    assert order.rate_candidate == Decimal("960.000")


async def test_future_order_gets_the_cost_from_its_own_first_day():
    contract = _contract(
        status=ContractStatus.active, candidate_rate_schedule=_raise_schedule()
    )
    future = _order(start_date=date(2026, 12, 1), end_date=date(2027, 3, 31))

    assert cost_reference_day(future, date(2026, 9, 20)) == date(2026, 12, 1)
    await sync_orders_cost_from_contract(
        _FakeDb(), contract, [future], today=date(2026, 9, 20)
    )
    assert future.rate_candidate == Decimal("1040.000")


async def test_completed_and_cancelled_orders_keep_their_history():
    contract = _contract(status=ContractStatus.active)
    done = _order(100, status=ClientOrderStatus.completed, rate_candidate=Decimal("1"))
    gone = _order(101, status=ClientOrderStatus.cancelled, rate_candidate=Decimal("2"))

    changed = await sync_orders_cost_from_contract(
        _FakeDb(), contract, [done, gone], today=date(2026, 9, 20)
    )

    assert changed == []
    assert (done.rate_candidate, gone.rate_candidate) == (Decimal("1"), Decimal("2"))


async def test_contract_without_cost_does_not_blank_the_order():
    contract = _contract(status=ContractStatus.active, rate_candidate=None)
    order = _order(rate_candidate=Decimal("900"))

    changed = await sync_orders_cost_from_contract(
        _FakeDb(), contract, [order], today=date(2026, 9, 20)
    )

    assert changed == []
    assert order.rate_candidate == Decimal("900")


async def test_group_line_cost_stays_managed_per_line():
    """Linie zamówień zbiorczych MD/kosztowych prowadzi DL — kontrakt ich nie rusza."""
    contract = _contract(status=ContractStatus.active)
    line = _order(
        order_group_id=7, rate_candidate=Decimal("800"), md_rate_cost=Decimal("800")
    )

    changed = await sync_orders_cost_from_contract(
        _FakeDb(), contract, [line], today=date(2026, 9, 20)
    )

    assert changed == []
    assert (line.rate_candidate, line.md_rate_cost) == (Decimal("800"), Decimal("800"))


# ── Integracja: zapis zamówienia przez API ──────────────────────────────────


async def _seed_signed_contractor(
    *, contract_status: ContractStatus = ContractStatus.draft
) -> dict[str, int]:
    from app.core.database import AsyncSessionLocal
    from app.models.candidate import Candidate
    from app.models.client import Client

    suffix = uuid.uuid4().hex[:6]
    async with AsyncSessionLocal() as db:
        client = Client(name=f"Alior Sync {suffix}")
        candidate = Candidate(
            name="Bartosz",
            lastname=f"Czapelka-{suffix}",
            email=f"sync-{suffix}@example.com",
        )
        db.add_all([client, candidate])
        await db.flush()
        contract = Contract(
            candidate_id=candidate.id,
            client_id=client.id,
            contract_type=ContractType.b2b,
            status=contract_status,
            start_date=date(2026, 9, 14),
            rate_candidate=Decimal("120.000"),
            rate_unit=RateUnit.hourly,
            currency="PLN",
            rate_candidate_currency="PLN",
        )
        db.add(contract)
        await db.flush()
        order = ClientOrder(
            client_id=client.id,
            contract_id=contract.id,
            title="OIT/0189/2026/ITVM",
            status=ClientOrderStatus.draft,
            start_date=date(2026, 9, 14),
            rate_unit=RateUnit.hourly,
            rate_candidate=Decimal("120.000"),
            rate_client_currency="PLN",
            rate_candidate_currency="PLN",
            currency="PLN",
        )
        db.add(order)
        await db.commit()
        return {
            "client_id": client.id,
            "contract_id": contract.id,
            "order_id": order.id,
        }


async def _load_contract(contract_id: int) -> Contract:
    from app.core.database import AsyncSessionLocal
    from app.services.contract_rates import RATE_SCHEDULE_LOADS

    async with AsyncSessionLocal() as db:
        return await db.scalar(
            select(Contract)
            .where(Contract.id == contract_id)
            .options(*RATE_SCHEDULE_LOADS)
        )


async def test_filling_the_order_updates_the_contract_through_the_api(
    app_client: AsyncClient, app_auth_headers: dict, contract_order_sync_mode: str
):
    """Uzupełnienie zamówienia w UI = kontrakt Aktywny z okresem i przychodem.

    Macierz bramki (`contract_order_sync_mode`): bez markera korekty system
    zachowuje się jak przed wdrożeniem — zamówienie się zapisuje, kontrakt
    zostaje nietknięty. Inaczej padnięta korekta w entrypoincie + zapisy
    zamówień zatarłyby niezgodności, które migawka z następnego deployu miała
    pokazać. Zapis idzie przez `commit_order_write` (PATCH zamówienia).
    """
    ids = await _seed_signed_contractor()

    resp = await app_client.patch(
        f"/api/clients/{ids['client_id']}/orders/{ids['order_id']}",
        json={
            "start_date": "2026-09-15",
            "end_date": "2026-12-31",
            "rate_unit": "daily",
            "rate_client": "1340",
            # Ręczny koszt w zamówieniu przegrywa z kontraktem.
            "rate_candidate": "999",
        },
        headers=app_auth_headers,
    )

    assert resp.status_code == 200, resp.text
    if contract_order_sync_mode == "disabled":
        # Zamówienie zapisane dokładnie tak, jak je wysłano — łącznie z ręcznym
        # kosztem, którego bez synchronizacji nikt nie nadpisuje.
        body = resp.json()
        assert (body["start_date"], body["end_date"]) == ("2026-09-15", "2026-12-31")
        assert body["rate_unit"] == "daily"
        assert Decimal(str(body["rate_client"])) == Decimal("1340")
        assert Decimal(str(body["rate_candidate"])) == Decimal("999")
        contract = await _load_contract(ids["contract_id"])
        assert contract.status == ContractStatus.draft
        assert contract.rate_unit == RateUnit.hourly
        assert contract.billing_hours_per_month == 168, "domyślny miesiąc roboczy"
        assert not contract.orders_in_md
        assert (contract.client_order_start_date, contract.client_order_end_date) == (
            None,
            None,
        )
        assert contract.client_rate_schedule == []
        assert contract.rate_client is None
        assert contract.effective_candidate_rate(date(2026, 10, 1)) == Decimal("120")
        return
    assert Decimal(str(resp.json()["rate_candidate"])) == Decimal("960")
    contract = await _load_contract(ids["contract_id"])
    assert contract.status == ContractStatus.active
    assert contract.rate_unit == RateUnit.hourly
    assert contract.billing_hours_per_month == 168
    assert contract.orders_in_md is True
    assert contract.effective_client_rate(date(2026, 10, 1)) == Decimal("167.5")
    assert contract.effective_candidate_rate(date(2026, 10, 1)) == Decimal("120")
    assert (contract.client_order_start_date, contract.client_order_end_date) == (
        date(2026, 9, 15),
        date(2026, 12, 31),
    )
    assert contract.end_date is None

    # UAT M08-B03: odpowiedź kontraktu mówi, który krok pochodzi z zamówienia.
    detail = await app_client.get(
        f"/api/contracts/{ids['contract_id']}", headers=app_auth_headers
    )
    assert detail.status_code == 200, detail.text
    steps = detail.json()["client_rate_schedule"]
    assert [step["source_order_id"] for step in steps] == [ids["order_id"]]


async def test_contract_cost_change_reaches_the_order_through_the_api(
    app_client: AsyncClient, app_auth_headers: dict, contract_order_sync_mode: str
):
    """Kierunek kontrakt → zamówienie; bez markera zamówienie zostaje nietknięte."""
    ids = await _seed_signed_contractor(contract_status=ContractStatus.active)

    resp = await app_client.patch(
        f"/api/contracts/{ids['contract_id']}",
        json={"rate_candidate": 130},
        headers=app_auth_headers,
    )

    assert resp.status_code == 200, resp.text
    from app.core.database import AsyncSessionLocal

    contract = await _load_contract(ids["contract_id"])
    assert contract.effective_candidate_rate(business_today()) == Decimal("130")
    async with AsyncSessionLocal() as db:
        order = await db.get(ClientOrder, ids["order_id"])
        expected = "130.000" if contract_order_sync_mode == "enabled" else "120.000"
        assert order.rate_candidate == Decimal(expected)


async def test_sync_failure_never_blocks_the_order_write(
    app_client: AsyncClient, app_auth_headers: dict, monkeypatch
):
    """Synchronizacja jest projekcją — awaria ląduje w logu, zamówienie w bazie."""
    from app.services import contract_order_sync

    async def _boom(*args, **kwargs):
        raise RuntimeError("celowa awaria synchronizacji")

    monkeypatch.setattr(contract_order_sync, "resync_contract", _boom)
    ids = await _seed_signed_contractor()

    resp = await app_client.patch(
        f"/api/clients/{ids['client_id']}/orders/{ids['order_id']}",
        json={"end_date": "2026-12-31"},
        headers=app_auth_headers,
    )

    assert resp.status_code == 200, resp.text
    from app.core.database import AsyncSessionLocal

    async with AsyncSessionLocal() as db:
        order = await db.get(ClientOrder, ids["order_id"])
        assert order.end_date == date(2026, 12, 31)


# ── Jednorazowo: migawka raportu + korekta szkiców ──────────────────────────


async def test_repair_snapshots_the_report_first_then_activates_every_draft():
    from app.core.database import AsyncSessionLocal
    from app.models.app_setting import AppSetting
    from app.services.contract_order_sync import REPAIR_MARKER
    from app.services.contract_order_sync_repair import (
        build_reconciliation_rows,
        build_reconciliation_workbook,
        repair_contract_names,
        run_contract_order_sync_repair,
    )

    filled = await _seed_signed_contractor()
    empty = await _seed_signed_contractor()
    async with AsyncSessionLocal() as db:
        order = await db.get(ClientOrder, filled["order_id"])
        order.start_date = date(2026, 9, 15)
        order.end_date = date(2026, 12, 31)
        order.rate_unit = RateUnit.daily
        order.rate_client = Decimal("1340")
        order.rate_candidate = Decimal("960")
        empty_order = await db.get(ClientOrder, empty["order_id"])
        await db.delete(empty_order)
        await db.execute(delete(AppSetting).where(AppSetting.key == REPAIR_MARKER))
        await db.commit()

    today = date(2026, 9, 20)
    scope = {filled["contract_id"], empty["contract_id"]}
    async with AsyncSessionLocal() as db:
        summary = await run_contract_order_sync_repair(
            db, today=today, only_contract_ids=scope
        )
        await db.commit()
    assert summary is not None

    async with AsyncSessionLocal() as db:
        assert (
            await run_contract_order_sync_repair(
                db, today=today, only_contract_ids=scope
            )
            is None
        )
        receipt = (await db.get(AppSetting, REPAIR_MARKER)).value

    snapshot_row = next(
        row for row in receipt["snapshot"] if row["order_id"] == filled["order_id"]
    )
    # Migawka jest SPRZED korekty: kontrakt był jeszcze szkicem bez przychodu.
    assert snapshot_row["contract_status"] == "draft"
    assert snapshot_row["contract_revenue"] is None
    assert snapshot_row["cost_status"] == "zgodne", "120 zł/h = 960 zł/MD"
    assert snapshot_row["period_status"] == "brak w kontrakcie"

    filled_contract = await _load_contract(filled["contract_id"])
    assert filled_contract.status == ContractStatus.active
    assert filled_contract.rate_unit == RateUnit.hourly
    assert filled_contract.effective_client_rate(today) == Decimal("167.5")
    assert filled_contract.client_order_end_date == date(2026, 12, 31)
    empty_contract = await _load_contract(empty["contract_id"])
    assert empty_contract.status == ContractStatus.active
    assert empty_contract.rate_client is None
    assert empty_contract.client_order_start_date is None

    async with AsyncSessionLocal() as db:
        live = await build_reconciliation_rows(db, today=today)
        names = await repair_contract_names(
            db, [item["contract_id"] for item in receipt["repaired"]]
        )
    content = build_reconciliation_workbook(
        snapshot=receipt["snapshot"],
        repaired=receipt["repaired"],
        repair_names=names,
        live=live,
        summary=receipt,
    )
    workbook = load_workbook(io.BytesIO(content))
    assert workbook.sheetnames == [
        "Podsumowanie",
        "Przed wdrożeniem",
        "Poprawione szkice",
        "Stan bieżący",
    ]
    live_row = next(row for row in live if row["order_id"] == filled["order_id"])
    assert live_row["revenue_status"] == "zgodne"
    assert live_row["period_status"] == "zgodne"
    async with AsyncSessionLocal() as db:
        await db.execute(delete(AppSetting).where(AppSetting.key == REPAIR_MARKER))
        await db.commit()


async def test_daily_cost_sync_waits_for_the_repair_marker():
    from app.core.database import AsyncSessionLocal
    from app.models.app_setting import AppSetting
    from app.services.contract_order_sync import (
        REPAIR_MARKER,
        run_daily_order_cost_sync,
    )

    async with AsyncSessionLocal() as db:
        existing = await db.get(AppSetting, REPAIR_MARKER)
        saved = existing.value if existing else None
        await db.execute(delete(AppSetting).where(AppSetting.key == REPAIR_MARKER))
        await db.commit()
    try:
        async with AsyncSessionLocal() as db:
            assert await run_daily_order_cost_sync(db) == 0
        ids = await _seed_signed_contractor(contract_status=ContractStatus.active)
        async with AsyncSessionLocal() as db:
            db.add(AppSetting(key=REPAIR_MARKER, value=saved or {"test": True}))
            contract = await db.get(Contract, ids["contract_id"])
            db.add(
                ContractCandidateRate(
                    contract_id=contract.id,
                    rate=Decimal("125"),
                    effective_from=business_today() - timedelta(days=1),
                )
            )
            await db.commit()
        async with AsyncSessionLocal() as db:
            assert await run_daily_order_cost_sync(db) >= 1
            await db.commit()
        async with AsyncSessionLocal() as db:
            order = await db.get(ClientOrder, ids["order_id"])
            assert order.rate_candidate == Decimal("125.000")
    finally:
        async with AsyncSessionLocal() as db:
            await db.execute(delete(AppSetting).where(AppSetting.key == REPAIR_MARKER))
            if saved is not None:
                db.add(AppSetting(key=REPAIR_MARKER, value=saved))
            await db.commit()


async def test_daily_pass_backfills_the_order_period_without_touching_money():
    """UAT B-B02: kontrakt aktywny, którego zamówienia nikt nie zapisał po
    wdrożeniu synchronizacji, dostaje okres najnowszego uzupełnionego
    zamówienia — ale stawka, jednostka i harmonogram zostają nietknięte."""
    from app.core.database import AsyncSessionLocal
    from app.models.candidate import Candidate
    from app.models.client import Client
    from app.services.contract_order_sync import backfill_missing_order_periods

    suffix = uuid.uuid4().hex[:6]
    async with AsyncSessionLocal() as db:
        client = Client(name=f"Okres Backfill {suffix}")
        candidate = Candidate(
            name="Okres",
            lastname=f"Backfill-{suffix}",
            email=f"period-{suffix}@example.com",
        )
        db.add_all([client, candidate])
        await db.flush()
        contract = Contract(
            candidate_id=candidate.id,
            client_id=client.id,
            contract_type=ContractType.b2b,
            status=ContractStatus.active,
            start_date=date(2026, 1, 1),
            rate_candidate=Decimal("100.000"),
            rate_client=Decimal("150.000"),
            rate_unit=RateUnit.hourly,
            currency="PLN",
            rate_candidate_currency="PLN",
        )
        db.add(contract)
        await db.flush()
        common = {
            "client_id": client.id,
            "contract_id": contract.id,
            "rate_unit": RateUnit.daily,
            "rate_client": Decimal("1200.000"),
            "rate_client_currency": "PLN",
            "rate_candidate_currency": "PLN",
            "currency": "PLN",
        }
        older = ClientOrder(
            title="Z-1",
            status=ClientOrderStatus.completed,
            start_date=date(2026, 7, 1),
            end_date=date(2026, 8, 31),
            **common,
        )
        newest = ClientOrder(
            title="Z-2",
            status=ClientOrderStatus.active,
            start_date=date(2026, 9, 1),
            end_date=date(2026, 9, 30),
            **common,
        )
        cancelled = ClientOrder(
            title="Z-3",
            status=ClientOrderStatus.cancelled,
            start_date=date(2026, 10, 1),
            end_date=date(2026, 12, 31),
            **common,
        )
        db.add_all([older, newest, cancelled])
        await db.commit()
        contract_id = contract.id
        # Bez synchronizacji przy zapisie (bezpośredni zapis do bazy) — stan
        # kontraktów sprzed wdrożenia.

    async with AsyncSessionLocal() as db:
        assert await backfill_missing_order_periods(db) >= 1
        await db.commit()

    loaded = await _load_contract(contract_id)
    assert (loaded.client_order_start_date, loaded.client_order_end_date) == (
        date(2026, 9, 1),
        date(2026, 9, 30),
    )
    assert loaded.rate_unit == RateUnit.hourly
    assert loaded.rate_client == Decimal("150.000")
    assert list(loaded.client_rate_schedule) == []

    async with AsyncSessionLocal() as db:
        # Idempotentny: drugi przebieg nie wybiera już tego kontraktu.
        await backfill_missing_order_periods(db)
        await db.commit()
        activities = (
            await db.scalars(
                select(Activity).where(
                    Activity.entity_type == "contract",
                    Activity.entity_id == contract_id,
                    Activity.action == "synced_with_orders",
                )
            )
        ).all()
    assert len(activities) == 1
    assert activities[0].details["source"] == "daily_period_backfill"


async def _contract_with_order(
    *, order_status: ClientOrderStatus, manual_end: date | None = None
) -> tuple[int, int]:
    from app.core.database import AsyncSessionLocal
    from app.models.candidate import Candidate
    from app.models.client import Client

    suffix = uuid.uuid4().hex[:6]
    async with AsyncSessionLocal() as db:
        client = Client(name=f"Okres Ręczny {suffix}")
        candidate = Candidate(
            name="Okres",
            lastname=f"Reczny-{suffix}",
            email=f"period-manual-{suffix}@example.com",
        )
        db.add_all([client, candidate])
        await db.flush()
        contract = Contract(
            candidate_id=candidate.id,
            client_id=client.id,
            contract_type=ContractType.b2b,
            status=ContractStatus.active,
            start_date=date(2026, 1, 1),
            rate_candidate=Decimal("100.000"),
            rate_client=Decimal("150.000"),
            rate_unit=RateUnit.hourly,
            currency="PLN",
            rate_candidate_currency="PLN",
            client_order_end_date=manual_end,
        )
        db.add(contract)
        await db.flush()
        order = ClientOrder(
            title="Z-R",
            status=order_status,
            start_date=date(2026, 9, 1),
            end_date=date(2026, 9, 30),
            client_id=client.id,
            contract_id=contract.id,
            rate_unit=RateUnit.daily,
            rate_client=Decimal("1200.000"),
            rate_client_currency="PLN",
            rate_candidate_currency="PLN",
            currency="PLN",
        )
        db.add(order)
        await db.commit()
        return contract.id, order.id


async def test_daily_backfill_keeps_a_hand_entered_order_end_date():
    """Przebieg uzupełnia wyłącznie BRAK okresu — ręcznie wpisana data końca
    (sama, bez startu) nie może zostać nadpisana datą z zamówienia."""
    from app.core.database import AsyncSessionLocal
    from app.services.contract_order_sync import backfill_missing_order_periods

    contract_id, _ = await _contract_with_order(
        order_status=ClientOrderStatus.active, manual_end=date(2027, 3, 31)
    )
    async with AsyncSessionLocal() as db:
        await backfill_missing_order_periods(db)
        await db.commit()

    loaded = await _load_contract(contract_id)
    assert loaded.client_order_start_date is None
    assert loaded.client_order_end_date == date(2027, 3, 31)


async def test_cancelling_the_order_clears_a_backfilled_period():
    """Okres z nocnego przebiegu nie ma kroku stawki, więc anulowanie
    zamówienia musi go zdjąć inną drogą — inaczej kontrakt udaje zamówienie."""
    from app.core.database import AsyncSessionLocal
    from app.services.contract_order_sync import (
        backfill_missing_order_periods,
        resync_contract,
    )

    contract_id, order_id = await _contract_with_order(
        order_status=ClientOrderStatus.active
    )
    async with AsyncSessionLocal() as db:
        await backfill_missing_order_periods(db)
        await db.commit()
    assert (await _load_contract(contract_id)).client_order_start_date == date(
        2026, 9, 1
    )

    async with AsyncSessionLocal() as db:
        order = await db.get(ClientOrder, order_id)
        order.status = ClientOrderStatus.cancelled
        await db.flush()
        await resync_contract(db, contract_id, actor_id=None)
        await db.commit()

    loaded = await _load_contract(contract_id)
    assert (loaded.client_order_start_date, loaded.client_order_end_date) == (
        None,
        None,
    )


# ── Poprawki po przeglądzie adwersarialnym ──────────────────────────────────


async def test_foreign_order_does_not_relabel_the_contract_own_revenue():
    """125 PLN/h sprzed zmiany nie może zacząć się czytać jako 125 EUR/h."""
    contract = _contract(
        status=ContractStatus.active,
        rate_candidate=Decimal("100"),
        rate_client=Decimal("125"),
        start_date=date(2026, 1, 1),
    )
    eur_order = _order(
        rate_client=Decimal("250"), rate_client_currency="EUR", currency="EUR"
    )

    await sync_contract_from_orders(
        _FakeDb(), contract, [eur_order], actor_id=None, today=date(2026, 9, 20)
    )

    assert contract.resolved_rate_client_currency == "PLN"
    assert contract.effective_client_rate(date(2026, 3, 1)) == Decimal("125")
    assert all(step.source_order_id is None for step in contract.client_rate_schedule)


async def test_full_contract_form_with_a_stale_cache_does_not_lower_revenue(
    app_client: AsyncClient, app_auth_headers: dict
):
    """Edycja PM-a nie może dopisać kroku ze starą stawką przychodową.

    Kolumna ``rate_client`` była 167,5, a od wczoraj obowiązuje krok 175.
    Formularz pokazuje stawkę EFEKTYWNĄ (175, ``GET /contracts/{id}``) i odsyła
    ją razem z edycją nazwy projektu — to NIE jest zmiana stawki (audyt 22.09
    r2, FIN-03: porównanie idzie ze stawką efektywną, nie z kolumną).
    """
    from app.core.database import AsyncSessionLocal

    ids = await _seed_signed_contractor(contract_status=ContractStatus.active)
    today = business_today()
    async with AsyncSessionLocal() as db:
        contract = await db.get(Contract, ids["contract_id"])
        contract.rate_client = Decimal("167.5")
        db.add_all(
            [
                ContractClientRate(
                    contract_id=contract.id,
                    rate=Decimal("167.5"),
                    effective_from=today - timedelta(days=20),
                ),
                ContractClientRate(
                    contract_id=contract.id,
                    rate=Decimal("175"),
                    effective_from=today - timedelta(days=1),
                ),
            ]
        )
        await db.commit()

    resp = await app_client.patch(
        f"/api/contracts/{ids['contract_id']}",
        json={"project_name": "Nowa nazwa", "rate_client": 175},
        headers=app_auth_headers,
    )

    assert resp.status_code == 200, resp.text
    contract = await _load_contract(ids["contract_id"])
    assert contract.effective_client_rate(today) == Decimal("175")
    assert len(contract.client_rate_schedule) == 2


async def test_changing_revenue_to_the_stale_column_value_is_a_real_change(
    app_client: AsyncClient, app_auth_headers: dict
):
    """FIN-03: zmiana na wartość nieaktualnej kolumny była cichym no-opem."""
    from app.core.database import AsyncSessionLocal

    ids = await _seed_signed_contractor(contract_status=ContractStatus.active)
    # „Dziś” w czasie polskim, jak kod — `date.today()` (UTC na CI) myliło
    # się między 22:00 a 24:00 UTC.
    from app.core.scheduling import business_today

    today = business_today()
    async with AsyncSessionLocal() as db:
        contract = await db.get(Contract, ids["contract_id"])
        contract.rate_client = Decimal("167.5")
        db.add_all(
            [
                ContractClientRate(
                    contract_id=contract.id,
                    rate=Decimal("167.5"),
                    effective_from=today - timedelta(days=20),
                ),
                ContractClientRate(
                    contract_id=contract.id,
                    rate=Decimal("175"),
                    effective_from=today - timedelta(days=1),
                ),
            ]
        )
        await db.commit()

    resp = await app_client.patch(
        f"/api/contracts/{ids['contract_id']}",
        json={"rate_client": 167.5},
        headers=app_auth_headers,
    )

    assert resp.status_code == 200, resp.text
    contract = await _load_contract(ids["contract_id"])
    assert contract.effective_client_rate(today) == Decimal("167.5")
    assert len(contract.client_rate_schedule) == 3


async def test_daily_pass_refreshes_stale_cost_and_old_revenue_caches():
    """FIN-03: cache kosztu nie odświeżał się nigdy, a przychodu tylko w oknie
    31 dni — lista kontraktorów pokazywała stawki sprzed kroku."""
    from app.core.database import AsyncSessionLocal
    from app.services.contract_order_sync import _refresh_revenue_caches

    ids = await _seed_signed_contractor(contract_status=ContractStatus.active)
    today = business_today()
    async with AsyncSessionLocal() as db:
        contract = await db.get(Contract, ids["contract_id"])
        contract.rate_client = Decimal("220")
        contract.rate_candidate = Decimal("175")
        contract.margin = Decimal("45")
        db.add_all(
            [
                ContractClientRate(
                    contract_id=contract.id,
                    rate=Decimal("213"),
                    effective_from=today - timedelta(days=90),
                ),
                ContractCandidateRate(
                    contract_id=contract.id,
                    rate=Decimal("165"),
                    effective_from=today - timedelta(days=90),
                ),
            ]
        )
        await db.commit()

    async with AsyncSessionLocal() as db:
        await _refresh_revenue_caches(db, today=today)
        await db.commit()

    contract = await _load_contract(ids["contract_id"])
    assert contract.rate_client == Decimal("213")
    assert contract.rate_candidate == Decimal("165")
    assert contract.margin == Decimal("48")


async def _seed_contract(
    *,
    status: ContractStatus,
    end_date: date | None,
    candidate_id: int | None = None,
    client_id: int | None = None,
) -> dict[str, int]:
    from app.core.database import AsyncSessionLocal
    from app.models.candidate import Candidate
    from app.models.client import Client

    suffix = uuid.uuid4().hex[:6]
    async with AsyncSessionLocal() as db:
        if client_id is None:
            client = Client(name=f"Repair {suffix}")
            db.add(client)
            await db.flush()
            client_id = client.id
        if candidate_id is None:
            candidate = Candidate(name="Rep", lastname=f"Air-{suffix}")
            db.add(candidate)
            await db.flush()
            candidate_id = candidate.id
        contract = Contract(
            candidate_id=candidate_id,
            client_id=client_id,
            contract_type=ContractType.b2b,
            status=status,
            start_date=date(2026, 1, 5),
            end_date=end_date,
            rate_candidate=Decimal("120.000"),
            rate_unit=RateUnit.hourly,
            currency="PLN",
            rate_candidate_currency="PLN",
        )
        db.add(contract)
        await db.commit()
        return {
            "contract_id": contract.id,
            "candidate_id": candidate_id,
            "client_id": client_id,
        }


async def test_repair_never_blindly_activates_expired_or_duplicate_drafts():
    """Aktywny szkic z minioną datą końca zakończyłby cron razem z zamówieniami."""
    from app.core.database import AsyncSessionLocal
    from app.models.app_setting import AppSetting
    from app.services.contract_order_sync import REPAIR_MARKER
    from app.services.contract_order_sync_repair import (
        ACTION_DUPLICATE,
        ACTION_ENDED,
        ACTION_REVIVED,
        run_contract_order_sync_repair,
    )

    today = date(2026, 9, 20)
    expired_with_order = await _seed_contract(
        status=ContractStatus.draft, end_date=date(2026, 6, 30)
    )
    expired_without_order = await _seed_contract(
        status=ContractStatus.draft, end_date=date(2026, 6, 30)
    )
    live = await _seed_contract(status=ContractStatus.active, end_date=None)
    duplicate = await _seed_contract(
        status=ContractStatus.draft,
        end_date=None,
        candidate_id=live["candidate_id"],
        client_id=live["client_id"],
    )
    async with AsyncSessionLocal() as db:
        db.add(
            ClientOrder(
                client_id=expired_with_order["client_id"],
                contract_id=expired_with_order["contract_id"],
                title="ZAM/1",
                status=ClientOrderStatus.active,
                start_date=date(2026, 7, 1),
                end_date=date(2026, 12, 31),
                rate_unit=RateUnit.hourly,
                rate_client=Decimal("170"),
                rate_candidate=Decimal("120"),
                rate_client_currency="PLN",
                rate_candidate_currency="PLN",
                currency="PLN",
            )
        )
        await db.execute(delete(AppSetting).where(AppSetting.key == REPAIR_MARKER))
        await db.commit()

    scope = {
        expired_with_order["contract_id"],
        expired_without_order["contract_id"],
        duplicate["contract_id"],
    }
    try:
        async with AsyncSessionLocal() as db:
            await run_contract_order_sync_repair(
                db, today=today, only_contract_ids=scope
            )
            await db.commit()
        async with AsyncSessionLocal() as db:
            receipt = (await db.get(AppSetting, REPAIR_MARKER)).value
    finally:
        async with AsyncSessionLocal() as db:
            await db.execute(delete(AppSetting).where(AppSetting.key == REPAIR_MARKER))
            await db.commit()

    actions = {item["contract_id"]: item["action"] for item in receipt["repaired"]}
    assert actions[expired_with_order["contract_id"]] == ACTION_REVIVED
    assert actions[expired_without_order["contract_id"]] == ACTION_ENDED
    assert actions[duplicate["contract_id"]] == ACTION_DUPLICATE

    revived = await _load_contract(expired_with_order["contract_id"])
    assert (revived.status, revived.end_date) == (ContractStatus.active, None)
    assert revived.effective_client_rate(today) == Decimal("170.000")
    ended = await _load_contract(expired_without_order["contract_id"])
    assert ended.status == ContractStatus.ended
    assert ended.end_date == date(2026, 6, 30)
    left = await _load_contract(duplicate["contract_id"])
    assert left.status == ContractStatus.draft


# ── audyt 22.09 r2 (FIN-02): bez fantomowego przychodu ──────────────────────


async def test_cancelled_last_order_leaves_the_contract_without_revenue():
    """Decyzja właściciela: ostatni krok z zamówień znika → kontrakt bez przychodu."""
    contract = _contract()
    order = _order()
    await sync_contract_from_orders(
        _FakeDb(), contract, [order], actor_id=None, today=date(2026, 9, 20)
    )
    assert contract.rate_client == Decimal("167.5")
    order.status = ClientOrderStatus.cancelled

    await sync_contract_from_orders(
        _FakeDb(), contract, [order], actor_id=None, today=date(2026, 9, 20)
    )

    assert contract.client_rate_schedule == []
    assert contract.rate_client is None
    assert contract.margin is None
    assert contract.effective_client_rate(date(2026, 10, 1)) is None


async def test_deleting_the_only_order_takes_its_revenue_off_the_contract(
    app_client: AsyncClient, app_auth_headers: dict
):
    """Kaskada w bazie kasowała krok ZANIM synchronizacja go zobaczyła —
    kolumna ``rate_client`` zostawała ze stawką usuniętego zamówienia."""
    ids = await _seed_signed_contractor()
    resp = await app_client.patch(
        f"/api/clients/{ids['client_id']}/orders/{ids['order_id']}",
        json={
            "start_date": "2026-09-15",
            "end_date": "2026-12-31",
            "rate_unit": "daily",
            "rate_client": "1340",
        },
        headers=app_auth_headers,
    )
    assert resp.status_code == 200, resp.text
    contract = await _load_contract(ids["contract_id"])
    assert contract.rate_client == Decimal("167.5")

    resp = await app_client.delete(
        f"/api/clients/{ids['client_id']}/orders/{ids['order_id']}",
        headers=app_auth_headers,
    )
    assert resp.status_code == 204, resp.text

    contract = await _load_contract(ids["contract_id"])
    assert contract.client_rate_schedule == []
    assert contract.rate_client is None
    assert contract.margin is None
    assert contract.effective_client_rate(date(2026, 10, 1)) is None
    assert (contract.client_order_start_date, contract.client_order_end_date) == (
        None,
        None,
    )


# ── Audyt 24.09.2026 (blok C) ───────────────────────────────────────────────


async def test_order_in_a_new_currency_keeps_the_revenue_history_of_the_old_one():
    """W1: zamówienie PLN, potem EUR — historia przychodu nie może czytać się w EUR.

    Do poprawki synchronizacja przełączała walutę kontraktu na EUR i usuwała
    kroki zamówień w PLN; resolver brał wtedy najbliższy przyszły krok, więc
    wrzesień „kosztował" 250 EUR/MD zamiast 1340 PLN/MD.
    """
    contract = _contract(status=ContractStatus.active)
    pln = _order(100)
    await sync_contract_from_orders(
        _FakeDb(), contract, [pln], actor_id=None, today=date(2026, 9, 20)
    )
    eur = _order(
        101,
        start_date=date(2027, 1, 1),
        end_date=date(2027, 6, 30),
        rate_client=Decimal("250"),
        rate_client_currency="EUR",
        currency="EUR",
    )

    outcome = await sync_contract_from_orders(
        _FakeDb(), contract, [pln, eur], actor_id=None, today=date(2026, 9, 20)
    )

    assert contract.resolved_rate_client_currency == "PLN"
    assert [s.source_order_id for s in contract.client_rate_schedule] == [100]
    assert contract.effective_client_rate(date(2026, 10, 1)) == Decimal("167.500")
    assert outcome.currency_conflict == "EUR"
    assert outcome.as_details()["currency_conflict"] == "EUR"


async def test_single_currency_switch_without_history_still_follows_the_order():
    """Kontrakt bez przychodu w starej walucie przyjmuje walutę zamówienia."""
    contract = _contract(status=ContractStatus.active)
    eur = _order(
        101, rate_client=Decimal("250"), rate_client_currency="EUR", currency="EUR"
    )

    outcome = await sync_contract_from_orders(
        _FakeDb(), contract, [eur], actor_id=None, today=date(2026, 9, 20)
    )

    assert contract.resolved_rate_client_currency == "EUR"
    assert outcome.currency_conflict is None


def test_placeholder_order_is_not_a_revenue_source():
    """S4: zaślepka „(bez numeru)" ze stawką skopiowaną z kontraktu nie jest
    „uzupełnionym zamówieniem" — inaczej stawka wraca na kontrakt jako
    przychód z zamówienia."""
    assert order_revenue_terms(_order(title="(bez numeru)")) is None
    assert order_revenue_terms(_order(title=" (bez numeru) ")) is None
    assert order_revenue_terms(_order(title="OIT/0189/2026/ITVM")) is not None


async def _seed_revenue_case(
    *, contract_rate_client: Decimal | None, order_title: str = "Z-REV"
) -> tuple[int, int]:
    from app.core.database import AsyncSessionLocal
    from app.models.candidate import Candidate
    from app.models.client import Client

    suffix = uuid.uuid4().hex[:6]
    async with AsyncSessionLocal() as db:
        client = Client(name=f"Przychód Backfill {suffix}")
        candidate = Candidate(
            name="Przychód",
            lastname=f"Backfill-{suffix}",
            email=f"revenue-{suffix}@example.com",
        )
        db.add_all([client, candidate])
        await db.flush()
        contract = Contract(
            candidate_id=candidate.id,
            client_id=client.id,
            contract_type=ContractType.b2b,
            status=ContractStatus.active,
            start_date=date(2026, 1, 1),
            rate_candidate=Decimal("100.000"),
            rate_client=contract_rate_client,
            rate_unit=RateUnit.hourly,
            billing_hours_per_month=168,
            currency="PLN",
            rate_candidate_currency="PLN",
        )
        db.add(contract)
        await db.flush()
        db.add(
            ClientOrder(
                title=order_title,
                status=ClientOrderStatus.active,
                start_date=date(2026, 9, 1),
                end_date=date(2026, 12, 31),
                client_id=client.id,
                contract_id=contract.id,
                rate_unit=RateUnit.hourly,
                rate_client=Decimal("170.000"),
                rate_client_currency="PLN",
                rate_candidate_currency="PLN",
                currency="PLN",
            )
        )
        await db.commit()
        return contract.id, client.id


async def test_nightly_pass_fills_missing_revenue_and_only_reports_drift():
    """D1 (audyt 24.09.2026): noc UZUPEŁNIA brak przychodu z zamówienia,
    a przychód różny od zamówienia tylko raportuje — bez nadpisania."""
    from app.core.database import AsyncSessionLocal
    from app.models.app_setting import AppSetting
    from app.services.contract_order_sync import (
        REVENUE_DRIFT_REPORT_KEY,
        backfill_missing_order_revenue,
    )

    missing_id, _ = await _seed_revenue_case(contract_rate_client=None)
    drift_id, _ = await _seed_revenue_case(contract_rate_client=Decimal("150.000"))
    placeholder_id, _ = await _seed_revenue_case(
        contract_rate_client=None, order_title="(bez numeru)"
    )

    async with AsyncSessionLocal() as db:
        result = await backfill_missing_order_revenue(db, today=date(2026, 9, 20))
        await db.commit()

    assert result.filled >= 1
    assert drift_id in result.drift_contract_ids
    assert missing_id not in result.drift_contract_ids

    filled = await _load_contract(missing_id)
    assert [s.source_order_id is not None for s in filled.client_rate_schedule] == [
        True
    ]
    assert filled.effective_client_rate(date(2026, 10, 1)) == Decimal("170.000")
    assert filled.status == ContractStatus.active

    untouched = await _load_contract(drift_id)
    assert list(untouched.client_rate_schedule) == []
    assert untouched.rate_client == Decimal("150.000")

    placeholder = await _load_contract(placeholder_id)
    assert list(placeholder.client_rate_schedule) == []
    assert placeholder.rate_client is None

    async with AsyncSessionLocal() as db:
        report = await db.get(AppSetting, REVENUE_DRIFT_REPORT_KEY)
    assert report is not None and drift_id in report.value["drift_contract_ids"]


# ── Audyt 24.09.2026: przychód z zaślepki „(bez numeru)” nie znika ──────────


def _placeholder_contract(order_status=ClientOrderStatus.draft):
    """Kontrakt z „Dodaj kolejny projekt”: stawka 1000 zł/h skopiowana do
    zaślepki, a synchronizacja zrobiła z niej krok harmonogramu."""
    shell = _order(
        100,
        title="(bez numeru)",
        status=order_status,
        start_date=date(2026, 9, 14),
        end_date=None,
        rate_client=Decimal("1000.000"),
        rate_unit=RateUnit.hourly,
    )
    contract = _contract(
        status=ContractStatus.active,
        rate_client=Decimal("1000.000"),
        client_rate_schedule=[
            ContractClientRate(
                rate=Decimal("1000.000"),
                effective_from=date(2026, 9, 14),
                source_order_id=100,
                note="Z zamówienia klienta",
            )
        ],
    )
    return contract, shell


async def test_placeholder_step_survives_as_the_contract_own_revenue():
    """Zaślepka nie jest źródłem przychodu (S4), ale jej krok to kopia stawki
    umowy — skasowanie go zerowało przychód przy pierwszym przeliczeniu."""
    contract, shell = _placeholder_contract()

    await sync_contract_from_orders(
        _FakeDb(), contract, [shell], actor_id=None, today=date(2026, 9, 24)
    )

    assert [
        (s.rate, s.effective_from, s.source_order_id)
        for s in contract.client_rate_schedule
    ] == [(Decimal("1000.000"), date(2026, 9, 14), None)]
    assert contract.rate_client == Decimal("1000.000")
    # Idempotentne: drugi przebieg nic już nie zmienia.
    outcome = await sync_contract_from_orders(
        _FakeDb(), contract, [shell], actor_id=None, today=date(2026, 9, 24)
    )
    assert outcome.revenue_steps_changed == 0
    assert contract.rate_client == Decimal("1000.000")


async def test_cancelled_placeholder_still_takes_its_step_away():
    """Anulowane zamówienie (także zaślepka) — bez zmian: krok znika."""
    contract, shell = _placeholder_contract(ClientOrderStatus.cancelled)

    await sync_contract_from_orders(
        _FakeDb(), contract, [shell], actor_id=None, today=date(2026, 9, 24)
    )

    assert contract.client_rate_schedule == []
    assert contract.rate_client is None
