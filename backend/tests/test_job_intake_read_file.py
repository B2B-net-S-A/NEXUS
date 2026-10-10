"""Plik na /jobs/new (D4, 10.10.2026): wypełniony wzór Profilu Championa jest
czytany z tabel — bez modelu — a każdy inny plik idzie starą ścieżką.

Trasa: `POST /api/job-intake/read-file`. Reguły odczytu wzoru:
`app/services/job_intake_champion_file.py`; pomiar na plikach z produkcji:
`backend/scripts/eval_champion_file_read.py`.
"""

from __future__ import annotations

import json
import re
from datetime import date
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from app.api import job_request_intake as api
from app.services import champion_document
from app.services import job_intake_champion_file as champion_file
from app.services import job_request_intake as intake
from scripts import eval_champion_file_read as script
from tests.taxonomy_fixture import hydrated_taxonomy

BACKEND = Path(__file__).parents[1]
FILLED_V4 = BACKEND / "tests/fixtures/champion/v4-filled.docx"
ASSETS = BACKEND / "app/assets/champion"
BLANK_V5 = ASSETS / "Profil_Championa_v5.0.docx"
URL = "/api/job-intake/read-file"

REQUEST = (
    "Szukamy Senior Java Developera do projektu migracji płatności. "
    "Wymagania: Java 17+, Spring Boot, Kafka. Praca zdalna."
)

V5_VALUES = {
    "nazwa roli": "Tester automatyzujący",
    "maks. stawka": "120–140 zł/h",
    "tryb pracy": "hybrydowo",
    "dni w biurze": "2",
    "lokalizacja biura": "Gdansk or Warsaw",
    "start": "ASAP",
    "deadline": "2026-12-15",
    "must-have": (
        "Java 17, Kafka lub Docker\n"
        "Doświadczenie we wdrażaniu funkcjonalności w systemach bankowych\n"
        "brak"
    ),
    "nice-to-have": "Selenium",
    "dziedzina": "płatności kartowe (min. 2 lata), e-commerce (mile)",
    "certyfikaty": "ISTQB Foundation",
    "obowiązki na stanowisku": "Testy regresji i automatyzacja scenariuszy.",
    "od klienta": "Decyduje lead QA. " * 200,
    "insight od": "Zespół 6 osób, dużo spotkań z biznesem.",
}


def _fill(values: dict[str, str], questions: int = 2) -> bytes:
    """Pusty wzór v5 z repo wypełniony tak, jak zrobiłby to Delivery Lead."""
    doc = Document(BLANK_V5)
    for table in doc.tables:
        for row in table.rows:
            cells = row.cells
            for i in range(0, len(cells) - 1, 2):
                label = cells[i].text.strip().lower()
                for prefix, value in values.items():
                    if label.startswith(prefix):
                        cells[i + 1].text = value
    screening = next(
        t
        for t in doc.tables
        if t.rows[0].cells[0].text.startswith("Pytania od Delivery")
    )
    for i, row in enumerate(screening.rows[1 : 1 + questions], start=1):
        row.cells[0].text = f"Pytanie {i}: Jak testowałeś proces {i}?"
        row.cells[
            1
        ].text = "Idealna odpowiedź: konkretny przykład\n\nDeal breaker: sama teoria"
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def _plain_docx(text: str) -> bytes:
    doc = Document()
    doc.add_paragraph(text)
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


# ── reguły odczytu wzoru (bez bazy) ──────────────────────────────────────────


def test_champion_profile_maps_to_the_request_form_shape() -> None:
    with hydrated_taxonomy():
        profile = champion_file.read_champion_form(_fill(V5_VALUES), "profil.docx")
        form = champion_file.intake_from_champion(profile, today=date(2026, 10, 1))

    assert form.role_name == "Tester automatyzujący"
    # Technologia zostaje wierszem (bez wersji), „A lub B” to warianty jednego
    # wiersza, a zdanie klienta i wersja idą do opisu — jak w „Uprość do słów”.
    assert form.requirements == [
        {"words": ["Java"], "level": "must"},
        {"words": ["Kafka", "Docker"], "level": "must"},
        {"words": ["Selenium"], "level": "nice"},
    ]
    assert form.must == ["Java", "Kafka lub Docker"]
    assert form.nice == ["Selenium"]
    assert form.search_requirements == [["Java"], ["Kafka", "Docker"]]
    assert form.descriptive_requirements == [
        "Java 17",
        "Doświadczenie we wdrażaniu funkcjonalności w systemach bankowych",
    ]
    # Przedział stawki: budżetem jest góra, „od” zostaje do wyświetlenia.
    assert form.rate_budget_hourly == 140.0
    assert form.rate_budget_hourly_min == 120.0
    assert form.rate_quote == "120–140 zł/h"
    assert "górną granicę" in (form.rate_note or "")
    assert form.remote_policy == "hybrid"
    assert form.onsite_days_per_week == 2
    assert form.office_cities == ["Gdańsk", "Warszawa"]
    assert form.office_city == "Gdańsk, Warszawa"
    assert form.deadline == "2026-12-15"
    assert form.responsibilities == "Testy regresji i automatyzacja scenariuszy."
    assert [
        (q.question, q.ideal_answer, q.deal_breaker, q.from_request)
        for q in form.screening_questions
    ] == [
        ("Jak testowałeś proces 1?", "konkretny przykład", "sama teoria", True),
        ("Jak testowałeś proces 2?", "konkretny przykład", "sama teoria", True),
    ]
    assert form.experience["domains"] == [
        {"name": "płatności kartowe", "level": "must", "min_years": 2},
        {"name": "e-commerce", "level": "nice", "min_years": None},
    ]
    assert [c["name"] for c in form.experience["certifications"]] == [
        "ISTQB Foundation"
    ]
    # Ta sama reguła braków co po odczycie requestu: Java jest w słowniku,
    # więc Delivery Lead musi jeszcze zdecydować o krytycznych.
    assert form.missing == [intake.MISSING_CRITICAL]
    # To wpis ze wzoru, nie cytat z maila klienta.
    assert set(form.provenance.values()) == {"document"}
    assert {"role", "rate", "must", "nice", "questions", "deadline"} <= set(
        form.provenance
    )
    assert form.evidence == [] and form.ask_client == []
    # „ASAP” nie jest datą — pole zostaje puste, a wpis trafia do uwag.
    assert form.start_date is None
    assert form.advisories == [
        "Start: w pliku jest „ASAP” — nie udało się wpisać tego w pole. "
        "Uzupełnij je ręcznie."
    ]
    assert set(form.as_dict()) == set(intake.RequestIntake().as_dict())


def test_sections_without_a_form_field_come_back_as_extras() -> None:
    profile = champion_file.read_champion_form(_fill(V5_VALUES), "profil.docx")
    extras = champion_file.champion_extras(profile)

    assert extras["search"] == {}
    assert extras["client"] == {
        "consultant_insight": "Zespół 6 osób, dużo spotkań z biznesem."
    }
    [note] = extras["insights"]
    assert set(note) == {"id", "source", "topic", "audience", "text", "done", "origin"}
    assert (note["source"], note["audience"], note["origin"]) == (
        "client",
        "team",
        "document",
    )
    # Dłuższa notatka nie wywraca odczytu — limit jak przy zapisie profilu.
    assert len(note["text"]) == 2000

    v4 = champion_file.champion_extras(
        champion_file.read_champion_form(FILLED_V4.read_bytes(), "profil.docx")
    )
    assert set(v4["search"]) == {"keywords", "notes"}
    assert set(v4["client"]) == {"consultant_insight", "historical_questions"}
    assert v4["insights"] == []


def test_values_the_form_cannot_take_become_notes_not_guesses() -> None:
    values = {
        **V5_VALUES,
        "maks. stawka": "45 EUR/h",
        "tryb pracy": "zdalnie",
        "deadline": "2026-09-01",
    }
    profile = champion_file.read_champion_form(_fill(values), "profil.docx")
    form = champion_file.intake_from_champion(profile, today=date(2026, 10, 1))

    assert form.rate_budget_hourly is None
    assert "45 EUR/h" in (form.rate_note or "")
    assert intake.MISSING_BUDGET in form.missing
    # Praca zdalna: dni i miasto biura ze wzoru nie mają czego opisywać.
    assert form.remote_policy == "remote"
    assert form.onsite_days_per_week is None and form.office_cities == []
    # Termin, który minął, nie robi z nowej rekrutacji „po terminie”.
    assert form.deadline is None and "deadline" not in form.provenance
    assert any("01.09.2026" in note for note in form.advisories)


def test_only_a_docx_with_the_form_labels_is_a_champion_form() -> None:
    filled = FILLED_V4.read_bytes()
    assert champion_file.read_champion_form(filled, "profil.pdf") is None
    assert champion_file.read_champion_form(b"%PDF-1.4 fikcyjny", "x.docx") is None
    assert champion_file.read_champion_form(_plain_docx(REQUEST), "x.docx") is None
    with pytest.raises(champion_file.EmptyChampionForm):
        champion_file.read_champion_form(BLANK_V5.read_bytes(), "wzor.docx")
    assert champion_file.looks_like_champion_text(
        champion_document.document_text(filled)
    )
    assert not champion_file.looks_like_champion_text(REQUEST)


# ── trasa (z bazą, CI) ───────────────────────────────────────────────────────


async def _client_id(name: str) -> int:
    from app.core.database import AsyncSessionLocal
    from app.models.client import Client

    async with AsyncSessionLocal() as db:
        client = Client(name=name)
        db.add(client)
        await db.commit()
        return client.id


async def _upload(app_client, headers, client_id: int, name: str, data: bytes):
    return await app_client.post(
        URL,
        files={"file": (name, data, "application/octet-stream")},
        data={"client_id": str(client_id)},
        headers=headers,
    )


def _forbid_model(monkeypatch) -> None:
    """Wzór czytamy bez modelu i bez wpisu w telemetrii AI."""

    async def forbidden(*args, **kwargs):
        raise AssertionError("wypełniony wzór nie może wołać modelu")

    def no_telemetry(*args, **kwargs):
        raise AssertionError("odczyt wzoru nie otwiera ai_feature")

    monkeypatch.setattr(intake, "read_request", forbidden)
    monkeypatch.setattr(api, "ai_feature", no_telemetry)


def _model_reads(monkeypatch) -> dict:
    seen: dict = {}

    async def fake_read(db, *, client_id, request_text):
        seen["text"] = request_text
        return intake.normalize_model_output(
            {"role_name": "Senior Java Developer"}, request_text
        )

    monkeypatch.setattr(intake, "read_request", fake_read)
    return seen


@pytest.mark.asyncio
async def test_filled_champion_form_is_read_from_its_tables_without_the_model(
    app_client, app_auth_headers, monkeypatch
) -> None:
    _forbid_model(monkeypatch)
    client_id = await _client_id("Champion File Test Sp. z o.o.")

    resp = await _upload(
        app_client,
        app_auth_headers,
        client_id,
        "Profil Championa.docx",
        FILLED_V4.read_bytes(),
    )

    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["reader"] == "champion_form"
    form = body["intake"]
    assert form["role_name"] == "Senior Java Developer"
    assert [(row["words"][0], row["level"]) for row in form["requirements"]] == [
        ("Java", "must"),
        ("Spring Boot", "must"),
        ("Kafka", "must"),
        ("AWS", "nice"),
        ("Kubernetes", "nice"),
    ]
    assert form["must"] == ["Java", "Spring Boot", "Kafka"]
    # Wersja z rubryki MUST zostaje zdaniem opisowym, nie słowem kluczowym.
    assert form["descriptive_requirements"][0] == "Java 17"
    assert form["rate_budget_hourly"] == 150.0
    assert form["remote_policy"] == "hybrid"
    assert form["onsite_days_per_week"] == 2
    assert form["office_cities"] == ["Warszawa"]
    questions = form["screening_questions"]
    assert len(questions) == 3
    assert all(q["from_request"] for q in questions)
    assert [q["deal_breaker"] for q in questions] == [
        "Brak własnego udziału w opisanej migracji.",
        "Brak rozumienia problemu podwójnego wykonania.",
        "Brak gotowości do współpracy z biznesem.",
    ]
    assert form["evidence"] == []
    assert set(form["provenance"].values()) == {"document"}
    # „critical” zależy od wczytanego słownika umiejętności; innych braków nie ma.
    assert set(form["missing"]) <= {intake.MISSING_CRITICAL}
    # Termin ze wzoru (18.09.2026) już minął — pole puste, uwaga dla DL.
    assert form["deadline"] is None
    assert any("już minął" in note for note in form["advisories"])
    assert set(form) == set(intake.RequestIntake().as_dict())
    # Tekst obok formularza trzyma pary etykieta | wartość.
    assert "Senior Java Developer" in body["text"]
    extras = body["champion_extras"]
    assert set(extras) == {"search", "client", "insights"}
    assert set(extras["search"]) == {"keywords", "notes"}
    assert set(extras["client"]) == {"consultant_insight", "historical_questions"}


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "template", ["Profil_Championa_v4.0.docx", "Profil_Championa_v5.0.docx"]
)
async def test_blank_template_is_refused_with_a_plain_sentence(
    app_client, app_auth_headers, monkeypatch, template
) -> None:
    _forbid_model(monkeypatch)
    client_id = await _client_id("Champion Blank Test Sp. z o.o.")

    resp = await _upload(
        app_client,
        app_auth_headers,
        client_id,
        "wzor.docx",
        (ASSETS / template).read_bytes(),
    )

    assert resp.status_code == 422
    assert resp.json()["detail"] == (
        "To pusty wzór Profilu Championa — wypełnij go albo wgraj zapytanie klienta."
    )


@pytest.mark.asyncio
async def test_plain_docx_still_goes_to_the_request_reader(
    app_client, app_auth_headers, monkeypatch
) -> None:
    seen = _model_reads(monkeypatch)
    client_id = await _client_id("Champion Plain Test Sp. z o.o.")

    resp = await _upload(
        app_client, app_auth_headers, client_id, "zapytanie.docx", _plain_docx(REQUEST)
    )

    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["reader"] == "request"
    assert "champion_extras" not in body
    assert "Senior Java Developera" in seen["text"]
    assert body["intake"]["role_name"] == "Senior Java Developer"
    assert champion_file.PDF_NOTE not in body["intake"]["advisories"]


@pytest.mark.asyncio
async def test_form_whose_text_cannot_be_rendered_falls_back_to_the_request_reader(
    app_client, app_auth_headers, monkeypatch
) -> None:
    """Tekst obok formularza to drugi odczyt pliku — jego awaria nie może dać 500.

    ``document_text`` schodzi do tabel zagnieżdżonych, których ``table_profile``
    nie dotyka: zagnieżdżona tabela ze scaleniem w pionie w pierwszym wierszu
    kończy go wyjątkiem python-docx, choć sam wzór jest rozpoznany.
    """
    seen = _model_reads(monkeypatch)

    def broken(data: bytes) -> str:
        raise ValueError("no tr above topmost tr in w:tbl")

    monkeypatch.setattr(champion_document, "document_text", broken)
    client_id = await _client_id("Champion Nested Table Test Sp. z o.o.")

    resp = await _upload(
        app_client,
        app_auth_headers,
        client_id,
        "Profil Championa.docx",
        FILLED_V4.read_bytes(),
    )

    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["reader"] == "request"
    assert "champion_extras" not in body
    assert seen["text"]


@pytest.mark.asyncio
async def test_pdf_never_reaches_the_table_reader_and_gets_a_hint(
    app_client, app_auth_headers, monkeypatch
) -> None:
    table_reads: list[int] = []
    monkeypatch.setattr(
        champion_file, "table_profile", lambda data: table_reads.append(1)
    )
    # PDF wzoru: tekst z rubrykami, ale bez tabel, z których dałoby się czytać.
    champion_text = champion_document.document_text(FILLED_V4.read_bytes())
    monkeypatch.setattr(
        "app.services.cv_text_extractor.extract_text", lambda path, name: champion_text
    )
    _model_reads(monkeypatch)
    client_id = await _client_id("Champion PDF Test Sp. z o.o.")

    resp = await _upload(
        app_client, app_auth_headers, client_id, "profil.pdf", b"%PDF-1.4 fikcyjny"
    )

    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert table_reads == []
    assert body["reader"] == "request"
    assert "champion_extras" not in body
    assert body["intake"]["advisories"][-1] == (
        "Wygląda na Profil Championa w PDF — wgraj plik .docx, wtedy odczytamy "
        "tabelę bez AI."
    )


@pytest.mark.asyncio
async def test_upload_guards_run_before_any_reader(
    app_client, app_auth_headers, monkeypatch
) -> None:
    _forbid_model(monkeypatch)
    table_reads: list[int] = []
    monkeypatch.setattr(
        champion_file, "table_profile", lambda data: table_reads.append(1)
    )
    client_id = await _client_id("Champion Guards Test Sp. z o.o.")

    async def send(name: str, data: bytes):
        return await _upload(app_client, app_auth_headers, client_id, name, data)

    too_big = await send("duzy.docx", b"x" * (api.MAX_FILE_BYTES + 1))
    assert too_big.status_code == 413
    empty = await send("pusty.docx", b"")
    assert empty.status_code == 422
    assert empty.json()["detail"] == "Plik jest pusty."
    wrong_type = await send("arkusz.xlsx", b"PK fikcyjny arkusz")
    assert wrong_type.status_code == 422
    assert table_reads == []


@pytest.mark.asyncio
async def test_unreadable_docx_gets_the_ordinary_refusal(
    app_client, app_auth_headers, monkeypatch
) -> None:
    """Plik .docx, którego nie da się otworzyć, ocenia zwykła ścieżka."""
    _forbid_model(monkeypatch)
    client_id = await _client_id("Champion Broken Test Sp. z o.o.")

    resp = await _upload(
        app_client,
        app_auth_headers,
        client_id,
        "zepsuty.docx",
        b"to nie jest plik Word, tylko tekst",
    )

    assert resp.status_code == 422
    assert "nie ma tekstu" in resp.json()["detail"]


@pytest.mark.asyncio
async def test_pasted_text_reports_the_request_reader(
    app_client, app_auth_headers, monkeypatch
) -> None:
    _model_reads(monkeypatch)
    client_id = await _client_id("Champion Text Test Sp. z o.o.")

    resp = await app_client.post(
        "/api/job-intake/read",
        json={"client_id": client_id, "text": REQUEST},
        headers=app_auth_headers,
    )

    assert resp.status_code == 200, resp.text
    assert resp.json()["reader"] == "request"


# ── skrypt pomiaru: same liczby pod stałymi etykietami ───────────────────────


def _assert_numbers_only(rendered: str, *, forbidden: tuple[str, ...]) -> None:
    def walk(node) -> None:
        if isinstance(node, dict):
            for key, value in node.items():
                assert re.fullmatch(r"[A-Za-z_.]+", key), key
                walk(value)
        else:
            assert isinstance(node, (int, float)) and not isinstance(node, bool), node

    walk(json.loads(rendered))
    lowered = rendered.casefold()
    for word in forbidden:
        assert word.casefold() not in lowered, word


def test_eval_script_reports_repo_fixtures_as_numbers_only() -> None:
    report = script.fixtures_report(script.load_reader())

    # Dwa wypełnione wzory z `tests/fixtures/champion` i dwa puste z `app/assets`.
    files = report["files"]
    assert files >= 4
    assert report["by_extension"] == {"docx": files}
    assert report["champion"]["filled"] >= 2
    assert report["champion"]["blank"] >= 2
    assert "errors" not in report
    assert report["old_path_text"]["files"] == files
    _assert_numbers_only(
        script.render({"fixtures": report}),
        forbidden=("Java", "Warszawa", "Klient testowy", "v4-filled", "Profil_"),
    )


def test_eval_script_compares_fields_without_printing_values() -> None:
    assert script.verdict("Java ", "java") == "agree"
    assert script.verdict(["Kafka", "Java"], ["java", "kafka"]) == "agree"
    assert script.verdict(150.0, "151") == "differ"
    assert script.verdict(None, 0) == "reader_empty"
    assert script.verdict("x", [], "ai") == "ai_empty"
    assert script.verdict("", None) is None


def test_eval_script_bundle_carries_the_reader_module() -> None:
    bundled = script.bundle()
    scope: dict = {"__name__": "bundle_check"}
    exec(compile(bundled, "bundle", "exec"), scope)
    module = BACKEND / "app/services/job_intake_champion_file.py"
    assert scope["MODULE_SOURCE"] == module.read_text(encoding="utf-8")


@pytest.mark.asyncio
async def test_eval_script_reads_request_files_read_only(monkeypatch) -> None:
    from sqlalchemy import delete

    from app.core.database import AsyncSessionLocal
    from app.models.client import Client
    from app.models.job import Job, RemotePolicy
    from app.models.job_file import JobFile
    from app.services import storage_service

    async def no_refresh() -> int:
        return 0

    # Słownik umiejętności wczytany tu zostałby w procesie dla kolejnych testów.
    monkeypatch.setattr(
        "app.services.skill_taxonomy_loader.refresh_alias_map", no_refresh
    )
    rel, size = storage_service.save_job_file(
        "Fikcyjny Klient profil.docx", BytesIO(FILLED_V4.read_bytes())
    )
    async with AsyncSessionLocal() as db:
        client = Client(name="Eval Champion Test Sp. z o.o.")
        db.add(client)
        await db.flush()
        job = Job(
            title="Senior Java Developer",
            client_id=client.id,
            rate_budget_hourly=150,
            remote_policy=RemotePolicy.hybrid,
            onsite_days_per_week=2,
            location="Warszawa",
            champion_profile={
                "basics": {"role_name": "Senior Java Developer"},
                "stack": {"must": [{"name": "Java"}, {"name": "Spring Boot"}]},
            },
        )
        db.add(job)
        await db.flush()
        row = JobFile(
            job_id=job.id,
            source="request",
            filename="Fikcyjny Klient profil.docx",
            file_path=rel,
            size_bytes=size,
        )
        db.add(row)
        await db.commit()
        file_id = row.id
    try:
        # Najnowszy wiersz to ten sprzed chwili — limit 1 pomija resztę bazy.
        report = await script.production_report(
            script.load_reader(), limit=1, with_ai=False
        )
    finally:
        async with AsyncSessionLocal() as db:
            await db.execute(delete(JobFile).where(JobFile.id == file_id))
            await db.commit()
        storage_service.delete_job_file(rel)

    assert report["rows"] == 1 and report["attached_to_job"] == 1
    assert report["by_extension"] == {"docx": 1}
    assert report["champion"]["filled"] == 1
    assert "errors" not in report
    assert report["reader_vs_job.role"] == {"agree": 1}
    assert report["reader_vs_job.rate"] == {"agree": 1}
    assert report["reader_vs_job.work_mode"] == {"agree": 1}
    assert report["reader_vs_job.office_city"] == {"agree": 1}
    # Wzór ma trzy wymagania, rekrutacja dwa; pytań w rekrutacji nie zapisano.
    assert report["reader_vs_job.must"] == {"differ": 1}
    assert report["reader_vs_job.questions"] == {"job_empty": 1}
    _assert_numbers_only(
        script.render({"request_files": report}),
        forbidden=("Fikcyjny", "Eval Champion", "Java", "Warszawa"),
    )
