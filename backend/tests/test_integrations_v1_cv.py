"""API v1 — CV dla klienta (``integrations_v1/cv.py``, 15.10.2026).

Dwie części:

* bez bazy — decyzje scope'ów tras (status: ``cv:generate``, plik:
  ``cv:generate`` + ``candidates:cv``), ustawienia, które karta „CV do
  klienta” zaznacza sama (``services/integration_cv``), blokada od „CV
  wysłane” i funkcje wydzielone z tras ekranu (plik szkicu i zatwierdzonej
  wersji, numer projektu z polityki) — te same odmowy i nazwy plików;
* na bazie (CI) — trasy z prawdziwym kluczem, generator i model podstawione:
  zlecenie → 202 i stan ``processing``, konto serwisowe jako autor (bez osoby,
  kwota AI na „system”, wpis z ``details.via``), powtórka z paragonu, 403 bez
  scope'u, plik tylko z obydwoma scope'ami (ten sam plik i ślad co ekran),
  rekrutacja ukrytego klienta = 404, zgoda RODO jak na ekranie.

Dane są zmyślone (repo jest publiczne): „Anna Przykładowa”, „Analityk danych”.
"""

from __future__ import annotations

import ast
import inspect
import json
import uuid
from dataclasses import asdict
from datetime import datetime, timedelta, timezone
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock

import pytest
from docx import Document
from fastapi import HTTPException
from httpx import AsyncClient
from pydantic import ValidationError
from sqlalchemy import select

from app.api import candidate_stage_cv as stage_cv_api
from app.api import cv_generator_b2b as cv_api
from app.api.integrations_v1 import cv as v1_cv
from app.services import integration_cv
from app.services.cv_generator_b2b import durable_jobs

V1 = "/api/integrations/v1"


# ── Bez bazy: scope'y tras ───────────────────────────────────────────────────


def _required_scopes(path: str, method: str) -> frozenset[str]:
    for route in v1_cv.router.routes:
        if route.path != path or method not in route.methods:
            continue
        stack = list(route.dependant.dependencies)
        while stack:
            dep = stack.pop()
            if "require_service_scope" in getattr(dep.call, "__qualname__", ""):
                return inspect.getclosurevars(dep.call).nonlocals["required_scopes"]
            stack.extend(dep.dependencies)
    raise AssertionError(f"brak trasy {method} {path}")


PAIR = "/recruitments/{job_id}/candidates/{candidate_id}/cv"


def test_generation_and_status_need_only_cv_generate():
    # Stan nie niesie treści CV (same stany, liczby i daty).
    assert _required_scopes(PAIR, "POST") == {"cv:generate"}
    assert _required_scopes(PAIR, "GET") == {"cv:generate"}


def test_file_needs_cv_generate_and_candidates_cv():
    assert _required_scopes(PAIR + "/file", "GET") == {"cv:generate", "candidates:cv"}


def test_module_has_no_postponed_annotations():
    # slowapi #579: PEP 563 zamienia zależności w parametry query (422).
    tree = ast.parse(inspect.getsource(v1_cv))
    assert not any(
        isinstance(node, ast.ImportFrom) and node.module == "__future__"
        for node in ast.walk(tree)
    )


def test_generate_body_is_strict_and_short():
    assert v1_cv.CvGenerateIn().note_for_generator == ""
    with pytest.raises(ValidationError):
        v1_cv.CvGenerateIn(note="krócej")
    with pytest.raises(ValidationError):
        v1_cv.CvGenerateIn(note_for_generator="x" * 301)


# ── Bez bazy: reguły karty ───────────────────────────────────────────────────


@pytest.mark.parametrize(
    ("column", "locked"),
    [
        ("new", False),
        ("screening", False),
        ("verified", False),
        ("cv_qc", False),
        ("cv_sent", True),
        ("client_interview", True),
        ("contract", True),
        ("hired", True),
        # „Zamknięci”: z kolumny nie wynika, czy CV wyszło — jak `isAtClientColumn`.
        ("closed", False),
        (None, False),
    ],
)
def test_generation_is_locked_from_cv_sent(column, locked):
    assert integration_cv.locked_after_send(column) is locked


@pytest.mark.parametrize(
    ("forced", "both", "expected"),
    [
        (None, False, ("pl", "one")),
        ("en", False, ("en", "one")),
        ("pl", True, ("pl", "one")),  # wymuszony język wygrywa z „obie”
        (None, True, ("pl", "both")),
    ],
)
def test_language_is_the_form_default(forced, both, expected):
    assert integration_cv.language_request(forced, both) == expected


def test_content_mode_is_the_form_default():
    assert integration_cv.requested_content_mode(True, None) == "tailored"
    assert integration_cv.requested_content_mode(False, None) == "polished"
    assert integration_cv.requested_content_mode(True, "polished") == "tailored"
    assert integration_cv.requested_content_mode(True, "basic") == "basic"


@pytest.mark.parametrize(
    ("cv_state", "latest", "expected"),
    [
        ("none", None, "none"),
        ("none", "processing", "processing"),
        ("draft", "processing", "processing"),
        ("draft", "ready", "ready"),
        ("finalized", "failed", "ready"),
        ("none", "failed", "failed"),
        ("none", "ready", "none"),
    ],
)
def test_overall_status(cv_state, latest, expected):
    assert integration_cv.overall_status(cv_state, latest) == expected


def _settings_mocks(monkeypatch, *, documents, snapshot, champion, central=False):
    from app.services.cv_generator_b2b import central_policies, client_rules
    from app.services.cv_generator_b2b import standalone_service

    monkeypatch.setattr(
        standalone_service,
        "list_candidate_cv_sources",
        AsyncMock(return_value=documents),
    )
    monkeypatch.setattr(
        client_rules,
        "resolve_client_rule",
        AsyncMock(return_value=SimpleNamespace(managed_policy={"key": "x"})),
    )
    monkeypatch.setattr(client_rules, "snapshot_rule", Mock(return_value=snapshot))
    monkeypatch.setattr(
        standalone_service, "champion_present", Mock(return_value=champion)
    )
    monkeypatch.setattr(central_policies, "enabled", Mock(return_value=central))
    monkeypatch.setattr(
        central_policies,
        "automatic_mode_decision",
        Mock(return_value=("basic", None)),
    )
    monkeypatch.setattr(cv_api, "policy_project_ref", Mock(return_value="ZOB/7"))
    db = AsyncMock()
    db.scalar.return_value = None
    db.get.return_value = SimpleNamespace(cv_content_mode_cap="basic")
    return db


async def test_settings_follow_the_card_defaults(monkeypatch):
    db = _settings_mocks(
        monkeypatch,
        documents=[{"id": 9}, {"id": 10}],
        snapshot=SimpleNamespace(cv_language="en", requires_en_copy=True),
        champion=True,
    )
    job = SimpleNamespace(client_id=5, title="  Analityk danych  ")
    stage = SimpleNamespace(id=3, candidate_id=2)
    settings = await integration_cv.generation_settings(db, job=job, stage=stage)
    assert settings.cv_document_id == 9  # główny plik pierwszy
    assert (settings.language, settings.languages) == ("en", "one")
    assert settings.content_mode == "tailored"
    assert settings.position == "Analityk danych"
    # Bez polityki centralnej numer projektu wylicza wspólna ścieżka.
    assert settings.project_ref == ""
    assert settings.sources is None


async def test_settings_under_central_policy_take_basic_and_project_ref(monkeypatch):
    db = _settings_mocks(
        monkeypatch,
        documents=[{"id": 9}],
        snapshot=SimpleNamespace(cv_language=None, requires_en_copy=True),
        champion=True,
        central=True,
    )
    job = SimpleNamespace(client_id=5, title="Analityk danych")
    stage = SimpleNamespace(id=3, candidate_id=2)
    settings = await integration_cv.generation_settings(db, job=job, stage=stage)
    assert settings.content_mode == "basic"
    assert settings.project_ref == "ZOB/7"
    assert (settings.language, settings.languages) == ("pl", "both")


async def test_settings_without_a_cv_file_are_422(monkeypatch):
    db = _settings_mocks(
        monkeypatch,
        documents=[],
        snapshot=None,
        champion=False,
    )
    with pytest.raises(HTTPException) as exc:
        await integration_cv.generation_settings(
            db,
            job=SimpleNamespace(client_id=5, title="x"),
            stage=SimpleNamespace(id=3, candidate_id=2),
        )
    assert exc.value.status_code == 422
    assert exc.value.detail["error"] == "no_cv_file"


async def test_last_used_sources_come_from_the_current_version():
    db = AsyncMock()
    db.scalar.return_value = {"ai_sources": {"answers": False, "unknown": True}}
    assert await integration_cv.last_used_sources(db, 3) == {
        "answers": False,
        "notes": True,
        "transcripts": False,
        "prep": False,
    }
    db.scalar.return_value = {"content_mode": "tailored"}
    assert await integration_cv.last_used_sources(db, 3) is None


# ── Bez bazy: funkcje wydzielone z tras ekranu ──────────────────────────────


def _csv(**fields):
    base = dict(
        id=11,
        branded_status="draft",
        branded_version=2,
        branded_docx_filename="B2B_Analityk_Anna_Przykladowa.docx",
        branded_render_metadata={},
        branded_consent_content=None,
        generated_document_id=None,
    )
    return SimpleNamespace(**{**base, **fields})


def test_draft_file_name_is_the_screen_rule():
    assert (
        stage_cv_api.draft_docx_filename(_csv())
        == "SZKIC_B2B_Analityk_Anna_Przykladowa.docx"
    )
    assert stage_cv_api.draft_docx_filename(_csv(branded_docx_filename=None)) == (
        "SZKIC_CV_v2.docx"
    )


@pytest.mark.parametrize("status", ["none", "finalized"])
async def test_draft_file_refuses_what_is_not_a_draft(status):
    with pytest.raises(HTTPException) as exc:
        await stage_cv_api.draft_docx_for_download(
            AsyncMock(), _csv(branded_status=status), "<p>CV</p>"
        )
    assert exc.value.status_code == 409


async def test_draft_file_without_required_consent_is_refused(monkeypatch):
    render = AsyncMock(return_value=(b"PK", b""))
    monkeypatch.setattr(stage_cv_api, "_render_editor_docx", render)
    with pytest.raises(HTTPException) as exc:
        await stage_cv_api.draft_docx_for_download(
            AsyncMock(),
            _csv(branded_render_metadata={"consent_required": True}),
            "<p>CV</p>",
        )
    assert exc.value.status_code == 409
    assert exc.value.detail["code"] == "consent_required"
    render.assert_not_awaited()


async def test_draft_file_renders_the_given_editor_html(monkeypatch):
    render = AsyncMock(return_value=(b"PK-draft", b""))
    monkeypatch.setattr(stage_cv_api, "_render_editor_docx", render)
    csv = _csv()
    assert await stage_cv_api.draft_docx_for_download(AsyncMock(), csv, "<p>A</p>") == (
        b"PK-draft"
    )
    render.assert_awaited_once_with(csv, "<p>A</p>")


async def test_approved_version_refusals():
    db = AsyncMock()
    db.scalar.return_value = None
    with pytest.raises(HTTPException) as exc:
        await stage_cv_api.approved_version_for_download(db, _csv(), 2)
    assert exc.value.status_code == 404

    db.scalar.return_value = SimpleNamespace(
        generated_document_id=None, docx_content=None, consent_content=None
    )
    with pytest.raises(HTTPException) as exc:
        await stage_cv_api.approved_version_for_download(db, _csv(), 2)
    assert exc.value.status_code == 409

    db.scalar.return_value = SimpleNamespace(
        generated_document_id=None, docx_content=b"PK", consent_content=None
    )
    with pytest.raises(HTTPException) as exc:
        await stage_cv_api.approved_version_for_download(
            db, _csv(branded_render_metadata={"consent_required": True}), 2
        )
    assert exc.value.detail["code"] == "consent_required"

    version = SimpleNamespace(
        generated_document_id=None, docx_content=b"PK", consent_content=None
    )
    db.scalar.return_value = version
    assert await stage_cv_api.approved_version_for_download(db, _csv(), 2) is version


def test_approved_headers_carry_the_frozen_file_name():
    headers = stage_cv_api.approved_docx_headers(
        SimpleNamespace(
            docx_filename="B2B_Analityk.docx",
            docx_sha256="b" * 64,
            version=3,
            content_sha256="c" * 64,
        )
    )
    assert "B2B_Analityk.docx" in headers["Content-Disposition"]
    assert headers["ETag"] == '"' + "b" * 64 + '"'
    assert headers["X-CV-Version"] == "3"
    assert headers["Cache-Control"] == "private, no-store"


def test_policy_project_ref_only_for_rules_that_need_it(monkeypatch):
    from app.services import cv_packages

    monkeypatch.setattr(cv_packages, "pko_job_reference", Mock(return_value="ZOB/1"))
    plain = SimpleNamespace(requires_rodo_consent_block=False, managed_policy={})
    consent = SimpleNamespace(requires_rodo_consent_block=True, managed_policy={})
    numbered = SimpleNamespace(
        requires_rodo_consent_block=False,
        managed_policy={"require_project_ref": True},
    )
    assert cv_api.policy_project_ref(plain, object()) is None
    assert cv_api.policy_project_ref(consent, object()) == "ZOB/1"
    assert cv_api.policy_project_ref(numbered, object()) == "ZOB/1"


# ── Bez bazy: trasy na atrapie sesji ────────────────────────────────────────


def _wired_app(monkeypatch, db, scopes: set[str]):
    """Router v1 CV na atrapie sesji i z kluczem podstawionym w bramce."""
    from typing import get_args

    from fastapi import FastAPI

    from app.api.deps import Caller
    from app.api.integrations_v1.common import (
        PREFIX,
        IntegrationApiError,
        integration_error_handler,
    )
    from app.core.database import get_db
    from app.core.rate_limit import limiter
    from app.services.service_account_auth import ServicePrincipal

    monkeypatch.setattr(limiter, "enabled", False)
    app = FastAPI()
    app.include_router(v1_cv.router, prefix=PREFIX)
    app.add_exception_handler(IntegrationApiError, integration_error_handler)
    caller = Caller(
        service=ServicePrincipal(
            account_id=1, slug="rozliczenia", key_id="v2$00", scopes=frozenset(scopes)
        )
    )
    app.dependency_overrides[get_db] = lambda: db
    for alias in (v1_cv.CvGenerate, v1_cv.CvFile):
        app.dependency_overrides[get_args(alias)[1].dependency] = lambda: caller
    return app


def _pair_mocks(monkeypatch, *, column: str = "screening", in_progress=False):
    job = SimpleNamespace(id=4, client_id=5, title="Analityk danych")
    stage = SimpleNamespace(id=3, candidate_id=2, job_id=4)
    monkeypatch.setattr(v1_cv, "_job_or_404", AsyncMock(return_value=job))
    monkeypatch.setattr(integration_cv, "current_stage", AsyncMock(return_value=stage))
    monkeypatch.setattr(v1_cv, "stage_column", AsyncMock(return_value=column))
    monkeypatch.setattr(
        integration_cv, "generation_in_progress", AsyncMock(return_value=in_progress)
    )
    monkeypatch.setattr(
        integration_cv,
        "generation_settings",
        AsyncMock(
            return_value=integration_cv.GenerationSettings(
                cv_document_id=9,
                language="pl",
                languages="one",
                content_mode="tailored",
                project_ref="",
                position="Analityk danych",
                sources=None,
            )
        ),
    )
    return job, stage


def _receipts(monkeypatch):
    from app.services import integration_receipts

    finish = AsyncMock()
    monkeypatch.setattr(
        integration_receipts, "begin", AsyncMock(return_value=(Mock(), None))
    )
    monkeypatch.setattr(integration_receipts, "finish", finish)
    return finish


async def _call(app, method: str, url: str, **kwargs):
    from httpx import ASGITransport

    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        return await client.request(method, url, **kwargs)


async def test_generate_route_enqueues_as_the_service_account(monkeypatch):
    db = AsyncMock()
    db.add = Mock()
    db.get.return_value = SimpleNamespace(id=2, name="Anna", lastname="Przykładowa")
    _pair_mocks(monkeypatch)
    finish = _receipts(monkeypatch)
    enqueue = AsyncMock(return_value=(11, 21, "Anna Przykładowa"))
    monkeypatch.setattr(cv_api, "enqueue_candidate_generation", enqueue)
    execute = AsyncMock()
    monkeypatch.setattr(durable_jobs, "execute_job", execute)
    app = _wired_app(monkeypatch, db, {"cv:generate"})

    resp = await _call(
        app,
        "POST",
        V1 + "/recruitments/4/candidates/2/cv",
        headers={"Idempotency-Key": str(uuid.uuid4())},
        json={"note_for_generator": "  krócej  "},
    )
    assert resp.status_code == 202, resp.text
    assert resp.json() == {
        "job_id": 4,
        "candidate_id": 2,
        "stage_id": 3,
        "generation_id": 11,
        "status": "processing",
    }
    kwargs = enqueue.await_args.kwargs
    assert kwargs["user_id"] is None
    assert kwargs["attach_stage_draft"] is True
    assert kwargs["client_id"] == 5 and kwargs["cv_document_id"] == 9
    assert kwargs["blind_cv"] is False and kwargs["content_mode"] == "tailored"
    assert kwargs["instruction"] == "krócej"
    assert kwargs["stage"].id == 3
    [activity] = [c.args[0] for c in db.add.call_args_list]
    assert activity.action == v1_cv.REQUESTED_ACTION
    assert (activity.entity_type, activity.entity_id) == ("cv_generated_document", 11)
    assert activity.user_id is None
    finish.assert_awaited_once()
    db.commit.assert_awaited_once()
    execute.assert_awaited_once_with(21)


async def test_generate_route_without_body_uses_no_note(monkeypatch):
    db = AsyncMock()
    db.add = Mock()
    _pair_mocks(monkeypatch)
    _receipts(monkeypatch)
    enqueue = AsyncMock(return_value=(11, 21, "Anna Przykładowa"))
    monkeypatch.setattr(cv_api, "enqueue_candidate_generation", enqueue)
    monkeypatch.setattr(durable_jobs, "execute_job", AsyncMock())
    app = _wired_app(monkeypatch, db, {"cv:generate"})
    resp = await _call(
        app,
        "POST",
        V1 + "/recruitments/4/candidates/2/cv",
        headers={"Idempotency-Key": str(uuid.uuid4())},
    )
    assert resp.status_code == 202, resp.text
    assert enqueue.await_args.kwargs["instruction"] == ""


@pytest.mark.parametrize(
    ("column", "in_progress", "code"),
    [
        ("cv_sent", False, "cv_locked_after_send"),
        ("hired", False, "cv_locked_after_send"),
        ("screening", True, "cv_generation_in_progress"),
    ],
)
async def test_generate_route_refusals_charge_nothing(
    monkeypatch, column, in_progress, code
):
    db = AsyncMock()
    db.add = Mock()
    _pair_mocks(monkeypatch, column=column, in_progress=in_progress)
    finish = _receipts(monkeypatch)
    enqueue = AsyncMock()
    monkeypatch.setattr(cv_api, "enqueue_candidate_generation", enqueue)
    execute = AsyncMock()
    monkeypatch.setattr(durable_jobs, "execute_job", execute)
    app = _wired_app(monkeypatch, db, {"cv:generate"})
    resp = await _call(
        app,
        "POST",
        V1 + "/recruitments/4/candidates/2/cv",
        headers={"Idempotency-Key": str(uuid.uuid4())},
    )
    assert resp.status_code == 409, resp.text
    assert resp.json()["error"]["code"] == code
    enqueue.assert_not_awaited()
    finish.assert_not_awaited()
    execute.assert_not_awaited()
    db.commit.assert_not_awaited()


async def test_status_route_reports_counts_without_content(monkeypatch):
    from app.services import cv_qc

    db = AsyncMock()
    _pair_mocks(monkeypatch)
    csv = _csv(
        id=17,
        candidate_stage_id=2,
        branded_status="finalized",
        branded_version=3,
        branded_updated_at=datetime(2026, 10, 15, 9, 0, tzinfo=timezone.utc),
        branded_finalized_at=datetime(2026, 10, 15, 10, 0, tzinfo=timezone.utc),
        branded_render_metadata={"consent_required": True},
        branded_consent_content=b"png",
        branded_draft_html="<p>Treść CV, której stan nie oddaje</p>",
    )
    monkeypatch.setattr(integration_cv, "pair_stage_cv", AsyncMock(return_value=csv))
    monkeypatch.setattr(
        v1_cv,
        "stage_versions",
        AsyncMock(
            return_value=[
                {
                    "version": 3,
                    "approved_at": datetime(2026, 10, 15, 10, tzinfo=timezone.utc),
                    "approved_by_name": "Ola Testowa",
                    "language": "pl",
                    "docx_available": True,
                    "is_current": True,
                }
            ]
        ),
    )
    monkeypatch.setattr(
        integration_cv,
        "pair_generations",
        AsyncMock(
            return_value=[
                {
                    "id": 11,
                    "status": "failed",
                    "language": "pl",
                    "content_mode": "tailored",
                    "origin": "manual",
                    "created_at": None,
                    "attached": False,
                    "error_code": "ai_failed",
                }
            ]
        ),
    )
    monkeypatch.setattr(
        cv_qc,
        "pair_statuses",
        AsyncMock(
            return_value={
                (2, 4): {
                    "status": "failed",
                    "blocking_failed": 2,
                    "ok_at": None,
                    "dismissed_blocking": 0,
                    "warnings_count": 5,
                }
            }
        ),
    )
    app = _wired_app(monkeypatch, db, {"cv:generate"})
    resp = await _call(app, "GET", V1 + "/recruitments/4/candidates/2/cv")
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["status"] == "ready"  # CV jest, choć ostatnia generacja padła
    assert (body["cv_state"], body["finalized"], body["version"]) == (
        "finalized",
        True,
        3,
    )
    assert body["cv_stage_id"] == 2
    assert (body["consent_required"], body["consent_missing"]) == (True, False)
    assert body["versions"][0]["version"] == 3
    assert "approved_by_name" not in body["versions"][0]
    assert body["generations"][0]["error_code"] == "ai_failed"
    assert body["qc"] == {
        "status": "failed",
        "blocking_failed": 2,
        "warnings_count": 5,
        "ok_at": None,
    }
    assert "Treść CV" not in resp.text


async def test_file_route_serves_the_draft_and_leaves_both_trails(monkeypatch):
    db = AsyncMock()
    db.add = Mock()
    db.info = {}
    _pair_mocks(monkeypatch)
    csv = _csv(id=17, branded_version=1, branded_draft_html="<p>Anna</p>")
    monkeypatch.setattr(integration_cv, "pair_stage_cv", AsyncMock(return_value=csv))
    draft = AsyncMock(return_value=b"PK-draft")
    monkeypatch.setattr(stage_cv_api, "draft_docx_for_download", draft)
    log = AsyncMock()
    monkeypatch.setattr(v1_cv, "log_cv_download", log)
    app = _wired_app(monkeypatch, db, {"cv:generate", "candidates:cv"})

    resp = await _call(app, "GET", V1 + "/recruitments/4/candidates/2/cv/file")
    assert resp.status_code == 200, resp.text
    assert resp.content == b"PK-draft"
    assert (
        "SZKIC_B2B_Analityk_Anna_Przykladowa.docx"
        in resp.headers["content-disposition"]
    )
    assert resp.headers["x-cv-state"] == "draft"
    draft.assert_awaited_once_with(db, csv, "<p>Anna</p>")
    log_kwargs = log.await_args.kwargs
    assert log_kwargs["intent"] == "download"
    assert (log_kwargs["entity_type"], log_kwargs["entity_id"]) == (
        "candidate_stage_cv",
        17,
    )
    assert log_kwargs["user"].id is None
    assert (log_kwargs["file_format"], log_kwargs["state"]) == ("docx", "draft")
    [audit] = [c.args[0] for c in db.add.call_args_list]
    assert (audit.action, audit.entity_type, audit.entity_id) == (
        "cv_downloaded",
        "candidate",
        2,
    )
    assert audit.user_id is None
    db.commit.assert_awaited_once()


async def test_file_route_without_a_cv_is_404_without_trail(monkeypatch):
    db = AsyncMock()
    db.add = Mock()
    _pair_mocks(monkeypatch)
    monkeypatch.setattr(integration_cv, "pair_stage_cv", AsyncMock(return_value=None))
    log = AsyncMock()
    monkeypatch.setattr(v1_cv, "log_cv_download", log)
    app = _wired_app(monkeypatch, db, {"cv:generate", "candidates:cv"})
    resp = await _call(app, "GET", V1 + "/recruitments/4/candidates/2/cv/file")
    assert resp.status_code == 404, resp.text
    assert resp.json()["error"]["code"] == "cv_not_ready"
    log.assert_not_awaited()
    db.add.assert_not_called()


# ── Na bazie (CI) ────────────────────────────────────────────────────────────


def _h(api_key: str, idem: str | None = None) -> dict:
    out = {"X-API-Key": api_key}
    if idem is not None:
        out["Idempotency-Key"] = idem
    return out


async def _key(app_client: AsyncClient, headers: dict, scopes: list[str]):
    from tests.test_service_accounts import _account_with_key

    account, api_key, _ = await _account_with_key(app_client, headers, scopes)
    return account, api_key


def _docx() -> bytes:
    document = Document()
    document.add_paragraph("Anna Przykładowa. Analityk danych. Python.")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


async def _world(*, hidden: bool = False, stage: str = "screening") -> dict:
    from app.core.database import AsyncSessionLocal
    from app.models.candidate import Candidate
    from app.models.candidate_document import CandidateDocument, CandidateDocumentKind
    from app.models.client import Client
    from app.models.job import Job, JobStatus
    from app.models.recruitment_pipeline import CandidateStage, PipelineStage

    tag = uuid.uuid4().hex[:8]
    now = datetime.now(timezone.utc)
    async with AsyncSessionLocal() as db:
        client = Client(name=f"V1 CV {tag}", hidden=hidden)
        db.add(client)
        await db.flush()
        job = Job(
            title=f"Analityk danych {tag}",
            client_id=client.id,
            status=JobStatus.published,
        )
        anna = Candidate(name="Anna", lastname=f"Przykładowa{tag}")
        outsider = Candidate(name="Jan", lastname=f"Testowy{tag}")
        db.add_all([job, anna, outsider])
        await db.flush()
        document = CandidateDocument(
            candidate_id=anna.id,
            filename="cv-anna.docx",
            file_content=_docx(),
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            document_kind=CandidateDocumentKind.cv,
            is_primary=True,
            uploaded_at=now,
        )
        card = CandidateStage(
            candidate_id=anna.id,
            job_id=job.id,
            stage=PipelineStage(stage),
            moved_at=now - timedelta(hours=1),
        )
        db.add_all([document, card])
        await db.commit()
        return {
            "client_id": client.id,
            "job_id": job.id,
            "candidate_id": anna.id,
            "outsider_id": outsider.id,
            "document_id": document.id,
            "stage_id": card.id,
        }


def _stub_generation(monkeypatch, world: dict) -> dict:
    """Prawdziwa wspólna ścieżka; bez magazynu plików, modelu i workera.

    Bez polityki centralnej (domyślnie wyłączona): reguły klienta czyta wtedy
    ``ClientCvRule``, a świeży klient testu żadnej nie ma.
    """
    from app.core.config import settings
    from app.services.cv_generator_b2b.champion_builder import (
        ChampionProfileForPrompt,
    )
    from app.services.cv_generator_b2b.standalone_service import (
        CandidateGenerationSource,
    )

    seen: dict = {"loads": [], "charges": [], "persisted": [], "executed": []}
    cv_bytes = _docx()

    async def _load(db, *, candidate_id, stage_id, cv_document_id, **kwargs):
        seen["loads"].append({"cv_document_id": cv_document_id, **kwargs})
        return CandidateGenerationSource(
            cv_bytes=cv_bytes,
            cv_filename="cv-anna.docx",
            champion_json=json.dumps(asdict(ChampionProfileForPrompt())),
            has_champion=False,
            screening_notes_text="",
            source_warnings=(),
            fallback_name="Anna Przykładowa",
            job_id=world["job_id"],
            job_title="Analityk danych",
            client_content_mode_cap=None,
            candidate_id=candidate_id,
            stage_id=stage_id,
            cv_document_id=cv_document_id,
            client_id=world["client_id"],
        )

    async def _charge(db, user_id):
        seen["charges"].append(user_id)

    async def _persist(db, **kwargs):
        await kwargs["charge"]()
        seen["persisted"].append(kwargs)
        return 9000 + len(seen["persisted"])

    async def _execute(job_id):
        seen["executed"].append(job_id)

    monkeypatch.setattr(settings, "CV_CENTRAL_POLICIES_ENABLED", False)
    monkeypatch.setattr(cv_api, "load_candidate_generation_source", _load)
    monkeypatch.setattr(cv_api, "_charge_cv_generation_quota", _charge)
    monkeypatch.setattr(durable_jobs, "persist_job", _persist)
    monkeypatch.setattr(durable_jobs, "execute_job", _execute)
    return seen


def _pair(world: dict, candidate_id: int | None = None) -> str:
    return (
        f"{V1}/recruitments/{world['job_id']}/candidates/"
        f"{candidate_id or world['candidate_id']}/cv"
    )


async def _documents(world: dict) -> list:
    from app.core.database import AsyncSessionLocal
    from app.models.cv_generated_document import CvGeneratedDocument

    async with AsyncSessionLocal() as db:
        return list(
            (
                await db.scalars(
                    select(CvGeneratedDocument)
                    .where(CvGeneratedDocument.candidate_id == world["candidate_id"])
                    .order_by(CvGeneratedDocument.id)
                )
            ).all()
        )


async def _activities(action: str, entity_type: str, entity_id: int) -> list:
    from app.core.database import AsyncSessionLocal
    from app.models.activity import Activity

    async with AsyncSessionLocal() as db:
        return list(
            (
                await db.scalars(
                    select(Activity).where(
                        Activity.action == action,
                        Activity.entity_type == entity_type,
                        Activity.entity_id == entity_id,
                    )
                )
            ).all()
        )


async def _stage_cv(world: dict, **fields) -> int:
    from app.core.database import AsyncSessionLocal
    from app.models.candidate_stage_cv import CandidateStageCV

    values = dict(
        branded_status="draft",
        branded_draft_html="<p>Anna Przykładowa — Analityk danych</p>",
        branded_docx_filename="B2B_Analityk_Anna_Przykladowa.docx",
        branded_version=1,
        branded_updated_at=datetime.now(timezone.utc),
    )
    values.update(fields)
    async with AsyncSessionLocal() as db:
        csv = CandidateStageCV(
            candidate_stage_id=world["stage_id"],
            candidate_id=world["candidate_id"],
            job_id=world["job_id"],
            **values,
        )
        db.add(csv)
        await db.commit()
        return csv.id


class TestGenerate:
    async def test_start_is_202_processing_with_the_service_account_as_author(
        self, app_client: AsyncClient, app_auth_headers: dict, monkeypatch
    ):
        world = await _world()
        seen = _stub_generation(monkeypatch, world)
        account, api_key = await _key(app_client, app_auth_headers, ["cv:generate"])

        resp = await app_client.post(
            _pair(world),
            headers=_h(api_key, str(uuid.uuid4())),
            json={"note_for_generator": "krócej, więcej o SQL"},
        )
        assert resp.status_code == 202, resp.text
        body = resp.json()
        assert body["status"] == "processing"
        assert body["stage_id"] == world["stage_id"]
        assert (body["job_id"], body["candidate_id"]) == (
            world["job_id"],
            world["candidate_id"],
        )

        [row] = await _documents(world)
        assert row.id == body["generation_id"]
        assert (row.status, row.created_by, row.origin) == (
            "processing",
            None,
            "manual",
        )
        assert (row.job_id, row.stage_id, row.client_id) == (
            world["job_id"],
            world["stage_id"],
            world["client_id"],
        )
        # Kwota AI bez osoby („system”), zadanie bez autora, podpięcie szkicu
        # jak po kliknięciu, ustawienia karty.
        assert seen["charges"] == [None]
        [job] = seen["persisted"]
        assert job["user_id"] is None
        inputs = job["inputs"]
        assert inputs["user_id"] is None and inputs["quota_user_id"] is None
        assert inputs["attach_stage_draft"] is True
        assert inputs["stage_id"] == world["stage_id"]
        assert (inputs["language"], inputs["content_mode"]) == ("pl", "polished")
        [load] = seen["loads"]
        assert load["cv_document_id"] == world["document_id"]
        assert load["instruction"] == "krócej, więcej o SQL"
        assert load["position"].startswith("Analityk danych")
        # Zadanie w tle po commicie.
        assert seen["executed"] == [9001]

        [requested] = await _activities(
            v1_cv.REQUESTED_ACTION, "cv_generated_document", row.id
        )
        assert requested.user_id is None
        assert requested.details["via"] == f"integration:{account['slug']}"
        assert requested.details["stage_id"] == world["stage_id"]

        status = await app_client.get(_pair(world), headers=_h(api_key))
        assert status.status_code == 200, status.text
        state = status.json()
        assert state["status"] == "processing"
        assert state["cv_state"] == "none" and state["finalized"] is False
        assert state["generations"][0]["id"] == row.id
        assert state["generations"][0]["status"] == "processing"
        assert state["qc"]["status"] == "unchecked"

    async def test_replay_returns_the_receipt_without_a_second_generation(
        self, app_client: AsyncClient, app_auth_headers: dict, monkeypatch
    ):
        world = await _world()
        seen = _stub_generation(monkeypatch, world)
        _, api_key = await _key(app_client, app_auth_headers, ["cv:generate"])
        idem = str(uuid.uuid4())

        first = await app_client.post(_pair(world), headers=_h(api_key, idem))
        second = await app_client.post(_pair(world), headers=_h(api_key, idem))
        assert first.status_code == second.status_code == 202, second.text
        assert second.json() == first.json()
        assert second.headers.get("Idempotent-Replayed") == "true"
        assert len(await _documents(world)) == 1
        assert len(seen["persisted"]) == 1 and seen["executed"] == [9001]

    async def test_without_scope_is_403(
        self, app_client: AsyncClient, app_auth_headers: dict, monkeypatch
    ):
        world = await _world()
        seen = _stub_generation(monkeypatch, world)
        _, api_key = await _key(
            app_client, app_auth_headers, ["candidates:read", "candidates:cv"]
        )
        post = await app_client.post(_pair(world), headers=_h(api_key, "a" * 16))
        assert post.status_code == 403, post.text
        assert post.json()["error"]["code"] == "insufficient_scope"
        status = await app_client.get(_pair(world), headers=_h(api_key))
        assert status.status_code == 403
        assert seen["persisted"] == [] and await _documents(world) == []

    async def test_hidden_client_recruitment_is_404(
        self, app_client: AsyncClient, app_auth_headers: dict, monkeypatch
    ):
        world = await _world(hidden=True)
        seen = _stub_generation(monkeypatch, world)
        _, api_key = await _key(
            app_client, app_auth_headers, ["cv:generate", "candidates:cv"]
        )
        post = await app_client.post(
            _pair(world), headers=_h(api_key, str(uuid.uuid4()))
        )
        assert post.status_code == 404, post.text
        assert (
            await app_client.get(_pair(world), headers=_h(api_key))
        ).status_code == (404)
        file = await app_client.get(_pair(world) + "/file", headers=_h(api_key))
        assert file.status_code == 404
        assert seen["persisted"] == []

    async def test_person_outside_the_recruitment_is_404(
        self, app_client: AsyncClient, app_auth_headers: dict, monkeypatch
    ):
        world = await _world()
        _stub_generation(monkeypatch, world)
        _, api_key = await _key(app_client, app_auth_headers, ["cv:generate"])
        resp = await app_client.post(
            _pair(world, world["outsider_id"]), headers=_h(api_key, str(uuid.uuid4()))
        )
        assert resp.status_code == 404, resp.text

    async def test_card_at_cv_sent_is_locked(
        self, app_client: AsyncClient, app_auth_headers: dict, monkeypatch
    ):
        world = await _world(stage="cv_sent")
        seen = _stub_generation(monkeypatch, world)
        _, api_key = await _key(app_client, app_auth_headers, ["cv:generate"])
        resp = await app_client.post(
            _pair(world), headers=_h(api_key, str(uuid.uuid4()))
        )
        assert resp.status_code == 409, resp.text
        assert resp.json()["error"]["code"] == "cv_locked_after_send"
        assert seen["persisted"] == []

    async def test_running_generation_blocks_a_second_one(
        self, app_client: AsyncClient, app_auth_headers: dict, monkeypatch
    ):
        from app.core.database import AsyncSessionLocal
        from app.models.cv_generation_job import CvGenerationJob

        world = await _world()
        seen = _stub_generation(monkeypatch, world)
        _, api_key = await _key(app_client, app_auth_headers, ["cv:generate"])
        first = await app_client.post(
            _pair(world), headers=_h(api_key, str(uuid.uuid4()))
        )
        assert first.status_code == 202, first.text
        async with AsyncSessionLocal() as db:
            db.add(
                CvGenerationJob(
                    generated_id=first.json()["generation_id"],
                    kind="new",
                    status="queued",
                    input_storage_key=f"pytest/{uuid.uuid4().hex}",
                    input_sha256="0" * 64,
                )
            )
            await db.commit()

        second = await app_client.post(
            _pair(world), headers=_h(api_key, str(uuid.uuid4()))
        )
        assert second.status_code == 409, second.text
        assert second.json()["error"]["code"] == "cv_generation_in_progress"
        assert len(seen["persisted"]) == 1

    async def test_consent_rule_refuses_like_the_screen_before_any_charge(
        self, app_client: AsyncClient, app_auth_headers: dict, monkeypatch
    ):
        # Bez polityki centralnej zrzut zgody jest wymogiem GENERACJI (422),
        # a klucz go nie wgrywa — ta sama odmowa co formularz bez zrzutu.
        from app.core.database import AsyncSessionLocal
        from app.models.client_cv_rule import ClientCvRule

        world = await _world()
        async with AsyncSessionLocal() as db:
            db.add(
                ClientCvRule(
                    client_id=world["client_id"],
                    requires_rodo_consent_block=True,
                    confirmed_at=datetime.now(timezone.utc),
                )
            )
            await db.commit()
        seen = _stub_generation(monkeypatch, world)
        _, api_key = await _key(app_client, app_auth_headers, ["cv:generate"])
        resp = await app_client.post(
            _pair(world), headers=_h(api_key, str(uuid.uuid4()))
        )
        assert resp.status_code == 422, resp.text
        assert "zgodą kandydata" in resp.json()["error"]["message"]
        assert seen["charges"] == [] and seen["persisted"] == []
        assert await _documents(world) == []


class TestFile:
    async def test_file_needs_both_scopes(
        self, app_client: AsyncClient, app_auth_headers: dict
    ):
        world = await _world()
        await _stage_cv(world)
        for scopes in (["cv:generate"], ["candidates:cv"]):
            _, api_key = await _key(app_client, app_auth_headers, scopes)
            resp = await app_client.get(_pair(world) + "/file", headers=_h(api_key))
            assert resp.status_code == 403, (scopes, resp.text)
            assert resp.json()["error"]["code"] == "insufficient_scope"

    async def test_draft_is_the_screen_file_with_the_screen_trail(
        self, app_client: AsyncClient, app_auth_headers: dict, monkeypatch
    ):
        world = await _world()
        csv_id = await _stage_cv(world)
        render = AsyncMock(return_value=(b"PK-draft", b""))
        monkeypatch.setattr(stage_cv_api, "_render_editor_docx", render)
        account, api_key = await _key(
            app_client, app_auth_headers, ["cv:generate", "candidates:cv"]
        )

        resp = await app_client.get(_pair(world) + "/file", headers=_h(api_key))
        assert resp.status_code == 200, resp.text
        assert resp.content == b"PK-draft"
        assert (
            "SZKIC_B2B_Analityk_Anna_Przykladowa.docx"
            in (resp.headers["content-disposition"])
        )
        assert resp.headers["x-cv-state"] == "draft"
        render.assert_awaited_once()

        [download] = await _activities(
            "cv_file_downloaded", "candidate_stage_cv", csv_id
        )
        assert download.user_id is None
        assert download.details["format"] == "docx"
        assert download.details["state"] == "draft"
        assert download.details["via"] == f"integration:{account['slug']}"
        [audit] = await _activities("cv_downloaded", "candidate", world["candidate_id"])
        assert audit.user_id is None
        assert audit.details["endpoint"] == v1_cv.AUDIT_ENDPOINT

        status = await app_client.get(_pair(world), headers=_h(api_key))
        state = status.json()
        assert state["status"] == "ready"
        assert (state["cv_state"], state["finalized"]) == ("draft", False)
        assert state["cv_stage_id"] == world["stage_id"]

    async def test_draft_pdf_comes_from_the_same_word_file(
        self, app_client: AsyncClient, app_auth_headers: dict, monkeypatch
    ):
        world = await _world()
        await _stage_cv(world)
        monkeypatch.setattr(
            stage_cv_api,
            "_render_editor_docx",
            AsyncMock(return_value=(b"PK-draft", b"")),
        )
        to_pdf = AsyncMock(return_value=b"%PDF-1.4 fikcyjny")
        monkeypatch.setattr(v1_cv, "_pdf", to_pdf)
        _, api_key = await _key(
            app_client, app_auth_headers, ["cv:generate", "candidates:cv"]
        )
        resp = await app_client.get(
            _pair(world) + "/file", params={"format": "pdf"}, headers=_h(api_key)
        )
        assert resp.status_code == 200, resp.text
        assert resp.headers["content-type"] == "application/pdf"
        assert (
            "SZKIC_B2B_Analityk_Anna_Przykladowa.pdf"
            in (resp.headers["content-disposition"])
        )
        assert to_pdf.await_args.args[1] == b"PK-draft"

    async def test_finalized_gives_the_approved_version_file(
        self, app_client: AsyncClient, app_auth_headers: dict
    ):
        from app.core.database import AsyncSessionLocal
        from app.models.cv_document_version import CvDocumentVersion

        world = await _world()
        now = datetime.now(timezone.utc)
        csv_id = await _stage_cv(
            world,
            branded_status="finalized",
            branded_version=2,
            branded_finalized_at=now,
        )
        async with AsyncSessionLocal() as db:
            db.add(
                CvDocumentVersion(
                    candidate_stage_cv_id=csv_id,
                    version=2,
                    content_html="<p>Anna Przykładowa</p>",
                    content_sha256="a" * 64,
                    docx_content=b"PK-final",
                    docx_sha256="b" * 64,
                    docx_filename="B2B_Analityk_Anna_Przykladowa.docx",
                    approved_at=now,
                )
            )
            await db.commit()
        _, api_key = await _key(
            app_client, app_auth_headers, ["cv:generate", "candidates:cv"]
        )
        resp = await app_client.get(_pair(world) + "/file", headers=_h(api_key))
        assert resp.status_code == 200, resp.text
        assert resp.content == b"PK-final"
        assert "SZKIC_" not in resp.headers["content-disposition"]
        assert resp.headers["x-cv-version"] == "2"
        assert resp.headers["x-cv-state"] == "approved"

        state = (await app_client.get(_pair(world), headers=_h(api_key))).json()
        assert state["finalized"] is True and state["version"] == 2
        assert [v["version"] for v in state["versions"]] == [2]

    async def test_missing_consent_is_409_and_leaves_no_download_trail(
        self, app_client: AsyncClient, app_auth_headers: dict, monkeypatch
    ):
        world = await _world()
        csv_id = await _stage_cv(
            world, branded_render_metadata={"consent_required": True}
        )
        monkeypatch.setattr(
            stage_cv_api,
            "_render_editor_docx",
            AsyncMock(return_value=(b"PK-draft", b"")),
        )
        _, api_key = await _key(
            app_client, app_auth_headers, ["cv:generate", "candidates:cv"]
        )
        resp = await app_client.get(_pair(world) + "/file", headers=_h(api_key))
        assert resp.status_code == 409, resp.text
        assert resp.json()["error"]["code"] == "consent_required"
        assert (
            await _activities("cv_file_downloaded", "candidate_stage_cv", csv_id) == []
        )
        state = (await app_client.get(_pair(world), headers=_h(api_key))).json()
        assert state["consent_required"] is True and state["consent_missing"] is True

    async def test_pair_without_cv_is_404_not_ready(
        self, app_client: AsyncClient, app_auth_headers: dict
    ):
        world = await _world()
        _, api_key = await _key(
            app_client, app_auth_headers, ["cv:generate", "candidates:cv"]
        )
        resp = await app_client.get(_pair(world) + "/file", headers=_h(api_key))
        assert resp.status_code == 404, resp.text
        assert resp.json()["error"]["code"] == "cv_not_ready"
