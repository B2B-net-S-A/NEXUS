"""Przegląd DZ (0353): CV dla klienta obok oryginału i zapytania klienta.

Decyzja Artura 23.09.2026: Dominik przy DZ sprawdza, czy must-have są w CV,
czy są pogrubione i czy są wpisane w każdej roli, w której występują
w oryginale. Kod liczy to deterministycznie; Luna dokłada podpowiedzi.
"""

from __future__ import annotations

import json
import re
from pathlib import Path
from types import SimpleNamespace

import pytest
import pytest_asyncio
from httpx import AsyncClient
from sqlalchemy import delete, select

from app.core.database import AsyncSessionLocal
from app.models.candidate import Candidate
from app.models.candidate_stage_cv import CandidateStageCV
from app.models.dz_review_hint import DzReviewHint
from app.models.job import Job
from app.models.recruitment_pipeline import CandidateStage, PipelineStage
from app.models.user import UserRole
from app.services import dz_review as svc
from tests.test_board_tasks import _cleanup, _login, _seed_user, _seed_world


@pytest_asyncio.fixture
async def api_client():
    from httpx import ASGITransport

    from app.core.rate_limit import limiter as _limiter
    from app.main import app

    _limiter.enabled = False
    async with AsyncClient(
        transport=ASGITransport(app=app, raise_app_exceptions=False),
        base_url="http://testserver",
    ) as c:
        yield c


BRANDED_HTML = """
<h2 data-cv-section="summary">Podsumowanie</h2>
<p>Programista <b>Java</b> z doświadczeniem w chmurze.</p>
<h2 data-cv-section="experience">Doświadczenie</h2>
<p class="job-head" data-cv-section="role"><b>Senior Developer</b> 2021–obecnie</p>
<p class="job-co" data-cv-section="employer">Acme Bank · Bankowość</p>
<ul><li>Rozwój usług w <b>Java</b> i Spring.</li></ul>
<p class="job-head" data-cv-section="role"><b>Developer</b> 2018–2021</p>
<p class="job-co" data-cv-section="employer">Globex</p>
<ul><li>Aplikacje webowe, <b>React</b>.</li></ul>
<script>alert(1)</script>
"""

EXPERIENCE = [
    {
        "company": "Acme Bank",
        "role": "Senior Developer",
        "desc": "Mikroserwisy Java, Kubernetes, Kafka",
    },
    {"company": "Globex", "role": "Developer", "desc": "Java, React, Docker"},
    {"company": "Initech", "role": "Junior", "desc": "Java, SQL"},
]

ORIGINAL = (
    "Jan Kowalski\nAcme Bank 2021-obecnie Senior Developer Java Kubernetes Kafka\n"
    "Globex 2018-2021 Developer Java React Docker\nInitech 2016-2018 Junior Java SQL"
)


# ── Jednostkowe ──────────────────────────────────────────────────────────────


def test_html_blocks_keep_bold_runs_and_sections_but_drop_scripts() -> None:
    blocks = svc.html_blocks(BRANDED_HTML)
    text = svc.blocks_text(blocks)
    assert "alert" not in text
    assert [b.section for b in blocks if b.kind == "h"] == ["summary", "experience"]
    assert "Java" in svc.bold_texts(blocks)
    assert "**Java**" in svc.blocks_text(blocks, mark_bold=True)


def test_generated_roles_follow_generator_markers() -> None:
    roles = svc.generated_roles(svc.html_blocks(BRANDED_HTML))
    assert [r.company for r in roles] == ["Acme Bank", "Globex"]
    assert "Spring" in roles[0].text


def _reqs(*names: str) -> list[svc.Requirement]:
    return [svc.Requirement(label=n, alternatives=(n,)) for n in names]


def test_analyze_reports_the_three_dominik_checks() -> None:
    blocks = svc.html_blocks(BRANDED_HTML)
    result = svc.analyze(
        _reqs("Java", "Kubernetes", "Docker"),
        _reqs("Spring"),
        blocks,
        ORIGINAL,
        EXPERIENCE,
    )
    by = {c["label"]: c for c in result["checks"]}

    # Java: w CV i pogrubiona, ale w roli Globex zniknęła, a Initech wypadł z CV.
    assert by["Java"]["in_cv"] and by["Java"]["bolded"]
    assert by["Java"]["missing_in_roles"] == ["Developer · Globex"]
    assert by["Java"]["roles_absent"] == ["Junior · Initech"]
    # Kubernetes: jest w oryginale (Acme), nie ma go w CV dla klienta.
    assert not by["Kubernetes"]["in_cv"] and by["Kubernetes"]["in_original"]
    assert by["Kubernetes"]["missing_in_roles"] == ["Senior Developer · Acme Bank"]
    # Pogrubienie spoza wymagań (stanowiska z szablonu się nie liczą).
    assert result["extra_bold"] == ["React"]
    assert result["summary"]["must_in_cv"] == 1


@pytest.mark.parametrize(
    ("alternatives", "terms"),
    [
        # Prawdziwe wpisy DL-ów z produkcji (23.09.2026): nazwa + opis po myślniku.
        (
            ("IT Project Management – prowadzenie i koordynacja projektów",),
            ("IT Project Management",),
        ),
        (("Security frameworks – MITRE ATT&CK, NIST",), ("Security frameworks",)),
        # Kontrakt wymagań rozciął „lub” wewnątrz nawiasu — sklejamy i zdejmujemy nawias.
        (("react.js (v18", "higher)"), ("react.js",)),
        (("Java", "Kotlin"), ("Java", "Kotlin")),
        # Dywiz bez spacji to część nazwy, nie opis.
        (("CI-CD",), ("CI-CD",)),
        (("Docker (preferowany)",), ("Docker",)),
    ],
)
def test_requirement_terms_read_the_name_not_the_description(
    alternatives, terms
) -> None:
    assert svc.requirement_terms(alternatives) == terms


def test_prose_must_have_is_found_by_its_name() -> None:
    """Wymaganie opisane zdaniem nie może dawać „brak w CV”, gdy CV nazywa
    tę umiejętność — tak wyglądało 0 z 11 u kandydata na Case Managera."""
    alternatives = ("Kubernetes – utrzymanie klastrów produkcyjnych",)
    req = svc.Requirement(
        label=alternatives[0],
        alternatives=alternatives,
        terms=svc.requirement_terms(alternatives),
    )
    result = svc.analyze([req], [], svc.html_blocks(BRANDED_HTML), ORIGINAL, EXPERIENCE)
    check = result["checks"][0]
    assert check["in_original"] is True
    assert check["terms"] == ["Kubernetes"]
    assert check["label"] == alternatives[0]


def test_analyze_without_generator_markers_splits_by_company_names() -> None:
    blocks = svc.html_blocks(
        "<p>Acme Bank</p><p>Java i Kubernetes</p><p>Globex</p><p>React</p>"
    )
    result = svc.analyze(_reqs("Java"), [], blocks, ORIGINAL, EXPERIENCE)
    java = result["checks"][0]
    assert java["missing_in_roles"] == ["Developer · Globex"]
    assert not java["bolded"]


def test_blind_cv_pairs_roles_by_position_not_company() -> None:
    """Blind CV zamienia firmy na „Firma z branży …" — role idą po stanowisku
    i kolejności, a nie wszystkie lądują jako „pominięte"."""
    from app.services.cv_generator_b2b.html_export import render_interactive_html
    from app.services.cv_generator_b2b.public_view import build_public_payload

    payload = {
        "blind_cv": True,
        "full_name": "Jan Kowalski",
        "experience": [
            {
                "company": "PKO BP",
                "industry": "bankowej",
                "position": "Senior Developer",
                "dates": "2021 – obecnie",
                "responsibilities": ["Usługi w Java i Kubernetes"],
                "technologies": ["Java", "Kubernetes"],
            },
            {
                "company": "Globex",
                "industry": "IT",
                "position": "Developer",
                "dates": "2018 – 2021",
                "responsibilities": ["Aplikacje webowe"],
                "technologies": ["React"],
            },
        ],
    }
    html = render_interactive_html(
        build_public_payload(payload), [], document_only=True
    )
    assert "PKO BP" not in html
    experience = [
        {"company": "PKO BP", "role": "Senior Developer", "desc": "Java, Kubernetes"},
        {"company": "Globex", "role": "Developer", "desc": "Java, React"},
    ]
    result = svc.analyze(
        _reqs("Java"), [], svc.html_blocks(html), "PKO BP Java Globex Java", experience
    )
    java = result["checks"][0]
    assert java["roles_absent"] == []
    assert java["missing_in_roles"] == ["Developer · Globex"]


def _client_docx() -> bytes:
    """CV dla klienta w Wordzie, jak pliki „…B2B…" z Traffita."""
    import io

    from docx import Document

    doc = Document()
    doc.add_paragraph("PODSUMOWANIE")
    p = doc.add_paragraph("Programista ")
    p.add_run("Java").bold = True
    p.add_run(" w bankowości.")
    doc.add_paragraph("DOŚWIADCZENIE ZAWODOWE")
    role = doc.add_paragraph()
    role.add_run("Senior Developer | Acme Bank | 2021 – obecnie").bold = True
    item = doc.add_paragraph(style="List Bullet")
    item.add_run("Usługi w ")
    item.add_run("Java").bold = True
    item.add_run(", Kafka.")
    role2 = doc.add_paragraph()
    role2.add_run("Developer | Globex | 2018 – 2021").bold = True
    doc.add_paragraph("Aplikacje webowe, React.", style="List Bullet")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).paragraphs[0].add_run("Języki: polski, angielski")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_blocks_read_bold_bullets_headings_and_roles() -> None:
    blocks = svc.docx_blocks(_client_docx())
    heads = [b.text for b in blocks if b.kind == "h"]
    assert heads == ["PODSUMOWANIE", "DOŚWIADCZENIE ZAWODOWE"]
    assert [b.section for b in blocks if b.kind == "h"] == [None, "experience"]
    assert sum(1 for b in blocks if b.kind == "li") == 2
    assert [b.text for b in blocks if b.section == "role"] == [
        "Senior Developer | Acme Bank | 2021 – obecnie",
        "Developer | Globex | 2018 – 2021",
    ]
    # Nagłówki ról są z szablonu — do pogrubionych must-have liczy się treść.
    assert svc.bold_texts(blocks) == ["Java", "Java"]
    assert any("Języki" in b.text for b in blocks)


def test_docx_client_cv_pairs_roles_by_title_and_order() -> None:
    blocks = svc.docx_blocks(_client_docx())
    result = svc.analyze(_reqs("Java", "Kafka"), [], blocks, ORIGINAL, EXPERIENCE[:2])
    by = {c["label"]: c for c in result["checks"]}
    assert by["Java"]["bolded"] is True
    assert by["Java"]["missing_in_roles"] == ["Developer · Globex"]
    assert by["Kafka"]["bolded"] is False
    assert result["summary"]["roles_checked"] is True


def test_unknown_roles_are_not_reported_as_omitted() -> None:
    """CV, w którym nie da się rozpoznać żadnej roli (blind bez znaczników),
    nie ogłasza wszystkich ról pominiętymi — mówi „sprawdź ręcznie"."""
    blocks = svc.text_blocks("Programista Java\nFirma z branży bankowej: Java")
    result = svc.analyze(
        _reqs("Java"), [], blocks, ORIGINAL, EXPERIENCE, bold_known=False
    )
    java = result["checks"][0]
    assert java["roles_absent"] == [] and java["missing_in_roles"] == []
    assert java["original_roles"]
    assert java["bolded"] is None
    assert result["summary"] == {
        **result["summary"],
        "roles_checked": False,
        "bold_known": False,
    }


def test_word_boundaries_do_not_count_javascript_as_java() -> None:
    blocks = svc.html_blocks("<p>Senior <b>JavaScript</b> developer</p>")
    result = svc.analyze(_reqs("Java"), [], blocks, "JavaScript", [])
    assert result["checks"][0]["in_cv"] is False


def test_client_cv_file_is_picked_by_client_name_then_newest() -> None:
    """Runda 8 (R8-N8-1): jedna reguła wyboru pliku „…B2B…" dla QC i kolejki
    Cpro — plik z nazwą klienta wygrywa z nowszym plikiem innego klienta."""
    from datetime import datetime, timezone

    old = datetime(2026, 9, 1, tzinfo=timezone.utc)
    new = datetime(2026, 9, 20, tzinfo=timezone.utc)
    rows = [
        (100, "Ewa_B2B_Nordea.docx", old, old),
        (120, "Ewa_B2B_PKO.docx", new, new),
    ]
    assert svc.pick_document_cv(rows, "Nordea Bank Abp") == 100
    assert svc.pick_document_cv(rows, "Bank Pocztowy") == 120
    assert svc.pick_document_cv(rows, None) == 120
    # Bez daty wgrania liczy się data założenia, remis rozstrzyga id.
    assert (
        svc.pick_document_cv(
            [(7, "B2B.docx", None, old), (8, "B2B.docx", None, old)], None
        )
        == 8
    )
    assert svc.pick_document_cv([], "Nordea") is None


def test_parse_hints_drops_quotes_not_found_in_either_cv() -> None:
    material = {"generated": "Programista **Java**", "original": "Java Kubernetes"}
    raw = json.dumps(
        {
            "verdict": "fix",
            "hints": [
                {
                    "kind": "missing_must",
                    "severity": "high",
                    "must_have": "Kubernetes",
                    "message": "Dopisz Kubernetes w roli Acme.",
                    "quote": "Java Kubernetes",
                },
                {
                    "kind": "weird",
                    "severity": "urgent",
                    "message": "Coś",
                    "quote": "zmyślony cytat",
                },
                {"kind": "other", "message": ""},
            ],
        }
    )
    out = svc.parse_hints(f"```json\n{raw}\n```", material)
    assert out["verdict"] == "fix"
    assert out["hints"][0]["quote"] == "Java Kubernetes"
    assert out["hints"][1] == {
        "kind": "other",
        "severity": "medium",
        "must_have": None,
        "message": "Coś",
        "quote": None,
    }
    assert len(out["hints"]) == 2


def test_migration_is_mirrored_in_entrypoint() -> None:
    root = Path(__file__).resolve().parents[1]
    entry = (root / "entrypoint.sh").read_text()
    migration = (root / "alembic/versions/0353_dz_review_cpro_sender.py").read_text()
    assert "ALTER TYPE aifeaturekey ADD VALUE IF NOT EXISTS 'dz_review'" in entry
    assert "SELECT 'dz_review', TRUE, 0" in entry
    squash = lambda t: re.sub(r"\s+", " ", t)  # noqa: E731
    for fragment in (
        "ALTER TABLE jobs ADD COLUMN IF NOT EXISTS cpro_sender_id INTEGER NULL",
        "CREATE TABLE IF NOT EXISTS dz_review_hints (",
        "CONSTRAINT uq_dz_review_hints_stage_hash UNIQUE (candidate_stage_id, input_hash)",
    ):
        assert fragment in squash(migration)
        assert fragment in squash(entry)


# ── Integracyjne ─────────────────────────────────────────────────────────────


async def _seed_review(world: dict, hor: dict, api: AsyncClient) -> int:
    async with AsyncSessionLocal() as db:
        job = await db.get(Job, world["job_id"])
        job.must_skills = ["Java", "Kubernetes"]
        cand = await db.get(Candidate, world["candidate_id"])
        cand.experience = EXPERIENCE
        cand.raw_cv_text = ORIGINAL
        await db.commit()
    moved = await api.post(
        "/api/pipeline/move",
        headers=hor,
        json={
            "candidate_id": world["candidate_id"],
            "job_id": world["job_id"],
            "stage_def_id": world["defs"]["verified"],
        },
    )
    assert moved.status_code == 200, moved.text
    stage_id = moved.json()["id"]
    async with AsyncSessionLocal() as db:
        csv = await db.scalar(
            select(CandidateStageCV).where(
                CandidateStageCV.candidate_stage_id == stage_id
            )
        )
        if csv is None:
            csv = CandidateStageCV(
                candidate_stage_id=stage_id,
                candidate_id=world["candidate_id"],
                job_id=world["job_id"],
            )
            db.add(csv)
        csv.branded_status = "draft"
        csv.branded_draft_html = BRANDED_HTML
        await db.commit()
    return stage_id


async def _review(stage_id: int) -> dict:
    """Przegląd przez serwis — trasy `/api/board-tasks/dz/*` zniknęły w v5
    (QC CV, `/api/pipeline/stages/{id}/qc`), serwis zostaje dla QC."""
    async with AsyncSessionLocal() as db:
        stage = await db.get(CandidateStage, stage_id)
        review = await svc.build_review(db, stage)
    return json.loads(json.dumps(review, default=str))


async def _hints(stage_id: int, user_id: int) -> dict:
    async with AsyncSessionLocal() as db:
        stage = await db.get(CandidateStage, stage_id)
        review = await svc.build_review(db, stage)
        return await svc.generate_hints(db, stage, review, user_id=user_id)


async def _cleanup_review(world: dict, stage_id: int) -> None:
    async with AsyncSessionLocal() as db:
        await db.execute(
            delete(DzReviewHint).where(DzReviewHint.candidate_stage_id == stage_id)
        )
        await db.execute(
            delete(CandidateStageCV).where(
                CandidateStageCV.candidate_id == world["candidate_id"]
            )
        )
        await db.commit()


@pytest.mark.asyncio
async def test_review_shows_both_cvs_request_and_checks(
    api_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    world = await _seed_world()
    hor_id, hor_creds = await _seed_user(UserRole.head_of_recruitment)
    hor = await _login(api_client, hor_creds)
    stage_id = await _seed_review(world, hor, api_client)
    try:
        body = await _review(stage_id)
        assert body["client_request"]["must"] == ["Java", "Kubernetes"]
        assert body["generated_cv"]["source"] == "branded_draft"
        assert body["original_cv"]["source"] == "profile_text"
        assert "Kubernetes" in body["original_cv"]["text"]
        checks = {c["label"]: c for c in body["checks"]}
        assert checks["Java"]["bolded"] is True
        assert checks["Kubernetes"]["in_cv"] is False
        # Treść CV idzie jako bloki tekstu, nie HTML — skrypt z edytora odpada.
        flat = json.dumps(body["generated_cv"]["blocks"])
        assert "<" not in flat and "alert" not in flat
    finally:
        await _cleanup_review(world, stage_id)
        await _cleanup(world, [hor_id])


@pytest.mark.asyncio
async def test_original_comes_from_an_earlier_stage_snapshot_of_the_pair(
    api_client: AsyncClient,
) -> None:
    """Etap przeglądu bez snapshotu (import z Traffita) — oryginał bierzemy
    z wcześniejszego etapu TEJ SAMEJ pary, nie z profilu."""

    world = await _seed_world()
    hor_id, hor_creds = await _seed_user(UserRole.head_of_recruitment)
    hor = await _login(api_client, hor_creds)
    stage_id = await _seed_review(world, hor, api_client)
    try:
        async with AsyncSessionLocal() as db:
            earlier = CandidateStage(
                candidate_id=world["candidate_id"],
                job_id=world["job_id"],
                stage_def_id=world["defs"]["screening"],
                stage=PipelineStage.screening,
            )
            db.add(earlier)
            await db.flush()
            db.add(
                CandidateStageCV(
                    candidate_stage_id=earlier.id,
                    candidate_id=world["candidate_id"],
                    job_id=world["job_id"],
                    original_cv_filename="cv.txt",
                    original_cv_content="Snapshot: Java, Kubernetes, Kafka".encode(),
                )
            )
            await db.execute(
                delete(CandidateStageCV).where(
                    CandidateStageCV.candidate_stage_id == stage_id,
                    CandidateStageCV.original_cv_content.is_(None),
                    CandidateStageCV.original_cv_storage_key.is_(None),
                    CandidateStageCV.branded_status == "none",
                )
            )
            await db.commit()
        body = await _review(stage_id)
        assert body["original_cv"]["source"] == "snapshot", body["original_cv"]
        assert body["original_cv"]["filename"] == "cv.txt"
        assert "Snapshot: Java" in (body["original_cv"]["text"] or "")
    finally:
        await _cleanup_review(world, stage_id)
        await _cleanup(world, [hor_id])


@pytest.mark.asyncio
async def test_client_cv_made_outside_the_generator_comes_from_the_b2b_file(
    api_client: AsyncClient,
) -> None:
    """Nordea 23.09.2026: CV dla klienta to plik „…B2B…" z Traffita, nie
    dokument generatora — przegląd bierze go, z pogrubieniami z Worda."""
    from app.models.candidate_document import CandidateDocument

    world = await _seed_world()
    hor_id, hor_creds = await _seed_user(UserRole.head_of_recruitment)
    hor = await _login(api_client, hor_creds)
    stage_id = await _seed_review(world, hor, api_client)
    try:
        async with AsyncSessionLocal() as db:
            # Bez CV firmowego w NEXUSIE — zostaje tylko plik kandydata.
            await db.execute(
                delete(CandidateStageCV).where(
                    CandidateStageCV.candidate_id == world["candidate_id"]
                )
            )
            data = _client_docx()
            # Runda 8 (R8-N8-1): plik klienta jest STARSZY niż plik „Inny" —
            # kolejka Cpro brała najnowszy, QC plik z nazwą klienta.
            for name in (
                "Ewa_B2B_Nordea.docx",
                "Ewa_B2B_Inny.docx",
                "Ewa_oryginal.pdf",
            ):
                db.add(
                    CandidateDocument(
                        candidate_id=world["candidate_id"],
                        filename=name,
                        file_content=data,
                    )
                )
            await db.commit()
        body = await _review(stage_id)
        cv = body["generated_cv"]
        assert cv["source"] == "document"
        assert cv["filename"] == "Ewa_B2B_Nordea.docx"
        assert cv["bold_known"] is True
        checks = {c["label"]: c for c in body["checks"]}
        assert checks["Java"]["bolded"] is True
        from app.services.move_requirements import company_cv_refs

        pair = (world["candidate_id"], world["job_id"])
        async with AsyncSessionLocal() as db:
            refs = await company_cv_refs(db, [pair])
        assert refs[pair]["document_id"] == cv["document_id"]
    finally:
        async with AsyncSessionLocal() as db:
            await db.execute(
                delete(CandidateDocument).where(
                    CandidateDocument.candidate_id == world["candidate_id"]
                )
            )
            await db.commit()
        await _cleanup_review(world, stage_id)
        await _cleanup(world, [hor_id])


@pytest.mark.asyncio
async def test_hints_are_generated_once_then_read_from_memory(
    api_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    world = await _seed_world()
    hor_id, hor_creds = await _seed_user(UserRole.head_of_recruitment)
    hor = await _login(api_client, hor_creds)
    stage_id = await _seed_review(world, hor, api_client)
    calls: list[dict] = []

    def fake_call(**kwargs):
        calls.append(kwargs)
        payload = {
            "verdict": "fix",
            "hints": [
                {
                    "kind": "missing_must",
                    "severity": "high",
                    "must_have": "Kubernetes",
                    "message": "Dopisz Kubernetes w roli Acme Bank.",
                    "quote": "Java Kubernetes Kafka",
                }
            ],
        }
        return SimpleNamespace(
            content=[SimpleNamespace(type="text", text=json.dumps(payload))],
            model="gpt-6-luna",
        )

    monkeypatch.setattr("app.services.claude_client.call_claude", fake_call)
    monkeypatch.setattr(
        "app.services.llm_providers.api_key_configured", lambda _model: True
    )
    try:
        body = await _hints(stage_id, hor_id)
        assert body["status"] == "ok" and body["cached"] is False
        assert body["hints"][0]["must_have"] == "Kubernetes"
        assert body["hints"][0]["quote"] == "Java Kubernetes Kafka"
        prompt = calls[0]["messages"][0]["content"]
        assert "**Java**" in prompt and "Kubernetes" in prompt

        second = await _hints(stage_id, hor_id)
        assert second["cached"] is True
        assert len(calls) == 1
    finally:
        await _cleanup_review(world, stage_id)
        await _cleanup(world, [hor_id])


@pytest.mark.asyncio
async def test_hints_failure_is_advisory_not_an_error(
    api_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    world = await _seed_world()
    hor_id, hor_creds = await _seed_user(UserRole.head_of_recruitment)
    hor = await _login(api_client, hor_creds)
    stage_id = await _seed_review(world, hor, api_client)

    def boom(**_kwargs):
        raise RuntimeError("provider down")

    monkeypatch.setattr("app.services.claude_client.call_claude", boom)
    monkeypatch.setattr(
        "app.services.llm_providers.api_key_configured", lambda _model: True
    )
    try:
        result = await _hints(stage_id, hor_id)
        assert result["status"] == "unavailable"
        async with AsyncSessionLocal() as db:
            assert (
                await db.scalar(
                    select(DzReviewHint.id).where(
                        DzReviewHint.candidate_stage_id == stage_id
                    )
                )
                is None
            )
    finally:
        await _cleanup_review(world, stage_id)
        await _cleanup(world, [hor_id])


@pytest.mark.asyncio
async def test_stage_row_delete_removes_remembered_hints() -> None:
    """Podpowiedzi niosą cytaty z CV — odchodzą razem z wierszem etapu."""

    world = await _seed_world()
    try:
        async with AsyncSessionLocal() as db:
            stage = CandidateStage(
                candidate_id=world["candidate_id"],
                job_id=world["job_id"],
                stage_def_id=world["defs"]["verified"],
                stage=PipelineStage.verified,
            )
            db.add(stage)
            await db.flush()
            db.add(
                DzReviewHint(
                    candidate_stage_id=stage.id,
                    input_hash="x" * 64,
                    payload={"hints": []},
                )
            )
            await db.commit()
            stage_id = stage.id
        async with AsyncSessionLocal() as db:
            await db.execute(
                delete(CandidateStage).where(CandidateStage.id == stage_id)
            )
            await db.commit()
            assert (
                await db.scalar(
                    select(DzReviewHint.id).where(
                        DzReviewHint.candidate_stage_id == stage_id
                    )
                )
                is None
            )
    finally:
        await _cleanup(world, [])
