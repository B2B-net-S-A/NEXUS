"""Generator v3: ``POST /identify-upload`` — czy osoba z pliku jest już w bazie?

Te same darmowe sita co ``/api/candidates/from-cv`` (identyczny plik, e-mail
i telefon z nagłówka CV) i ZERO wywołań modelu — rozpoznanie ma być
natychmiastowe i bezpłatne.
"""

# ruff: noqa: F811  (fixture `pv_client` importowana z sąsiedniego pliku)
from __future__ import annotations

import hashlib
import uuid
from io import BytesIO

import pytest
from docx import Document
from httpx import AsyncClient

import app.models  # noqa: F401
from app.core.database import AsyncSessionLocal
from app.models.candidate import Candidate
from app.models.candidate_document import CandidateDocument
from tests.test_cv_auto_generate import _headers, _world
from tests.test_pending_gate_removed import pv_client  # noqa: F401

DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _cv(*lines: str) -> bytes:
    document = Document()
    for line in lines:
        document.add_paragraph(line)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def no_model(monkeypatch):
    """Każde wywołanie modelu kończy test."""
    from app.services import claude_client, cv_parser

    async def _forbidden(*args, **kwargs):
        raise AssertionError("identify-upload must not call a model")

    monkeypatch.setattr(cv_parser, "parse_cv", _forbidden)
    monkeypatch.setattr(claude_client, "call_claude", _forbidden)


async def _identify(client: AsyncClient, user_id: int, content: bytes, name="cv.docx"):
    return await client.post(
        "/api/cv-generator/identify-upload",
        headers=_headers(user_id),
        files={"cv_file": (name, content, DOCX_TYPE)},
    )


@pytest.mark.asyncio
async def test_email_in_the_header_finds_the_person(pv_client: AsyncClient, no_model):
    world = await _world()
    email = f"identify-{uuid.uuid4().hex[:8]}@example.com"
    async with AsyncSessionLocal() as db:
        candidate = Candidate(name="Anna", lastname="Wzorcowa", email=email)
        db.add(candidate)
        await db.commit()
        candidate_id = candidate.id

    response = await _identify(
        pv_client, world["user_id"], _cv("Anna Wzorcowa", email, "Python Developer")
    )
    assert response.status_code == 200, response.text
    [match] = [
        m for m in response.json()["matches"] if m["candidate_id"] == candidate_id
    ]
    assert match["full_name"] == "Anna Wzorcowa"
    assert match["match_reasons"]


@pytest.mark.asyncio
async def test_identical_file_is_recognised_by_its_bytes(
    pv_client: AsyncClient, no_model
):
    world = await _world()
    content = _cv("Jan Plik", f"Unikalny {uuid.uuid4().hex}")
    async with AsyncSessionLocal() as db:
        db.add(
            CandidateDocument(
                candidate_id=world["candidate_id"],
                filename="jan.docx",
                content_sha256=hashlib.sha256(content).hexdigest(),
            )
        )
        await db.commit()
    response = await _identify(pv_client, world["user_id"], content)
    assert response.status_code == 200, response.text
    [match] = response.json()["matches"]
    assert match["candidate_id"] == world["candidate_id"]
    assert match["match_reasons"] == ["identical_file"]


@pytest.mark.asyncio
async def test_unknown_person_is_an_empty_list(pv_client: AsyncClient, no_model):
    world = await _world()
    response = await _identify(
        pv_client,
        world["user_id"],
        _cv("Nowa Osoba", f"nowa-{uuid.uuid4().hex}@example.org"),
    )
    assert response.status_code == 200, response.text
    assert response.json() == {"matches": []}


@pytest.mark.asyncio
async def test_unreadable_file_is_refused(pv_client: AsyncClient, no_model):
    world = await _world()
    response = await _identify(pv_client, world["user_id"], b"not a document")
    assert response.status_code == 422, response.text


def test_route_is_mounted():
    from app.main import app

    assert "/api/cv-generator/identify-upload" in app.openapi()["paths"]
