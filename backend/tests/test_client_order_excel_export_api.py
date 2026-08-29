"""Legacy client-order Excel export: scope, order and financial columns."""

from __future__ import annotations

import io
import uuid
from datetime import date, timedelta
from decimal import Decimal

from httpx import AsyncClient
from openpyxl import load_workbook


async def _seed_order(label: str) -> tuple[int, int, int, str]:
    from app.core.database import AsyncSessionLocal
    from app.models.candidate import Candidate
    from app.models.client import Client
    from app.models.client_order import ClientOrder, ClientOrderStatus
    from app.models.contract import Contract, ContractStatus

    suffix = uuid.uuid4().hex[:8]
    async with AsyncSessionLocal() as db:
        candidate = Candidate(
            name="Anna",
            lastname=f"Eksportowa {suffix}",
            email=f"excel-{suffix}@example.com",
        )
        client = Client(name=f"{label} {suffix}")
        db.add_all([candidate, client])
        await db.commit()
        await db.refresh(candidate)
        await db.refresh(client)
        contract = Contract(
            candidate_id=candidate.id,
            client_id=client.id,
            status=ContractStatus.active,
            start_date=date.today() - timedelta(days=30),
            rate_candidate=Decimal("123.456"),
            rate_client=Decimal("180.000"),
        )
        db.add(contract)
        await db.commit()
        await db.refresh(contract)
        order = ClientOrder(
            client_id=client.id,
            contract_id=contract.id,
            title=f"PO-{suffix}",
            status=ClientOrderStatus.active,
            start_date=date(2026, 1, 1),
            end_date=date(2026, 12, 31),
            rate_client=Decimal("185.750"),
        )
        db.add(order)
        await db.commit()
        await db.refresh(order)
        return (
            client.id,
            order.id,
            contract.id,
            f"{candidate.name} {candidate.lastname}",
        )


async def _seed_md_group(client_id: int, order_number: str) -> int:
    from app.core.database import AsyncSessionLocal
    from app.models.client_order_group import ClientOrderGroup

    async with AsyncSessionLocal() as db:
        group = ClientOrderGroup(
            client_id=client_id,
            order_number=order_number,
            start_date=date(2026, 2, 1),
            end_date=date(2026, 11, 30),
            status="active",
            order_type="md",
            is_cost_based=False,
            is_md_budget_based=True,
            md_budget_total=Decimal("80"),
            md_budget_remaining=Decimal("80"),
            md_budget_manual_adjustment=Decimal("0"),
        )
        db.add(group)
        await db.commit()
        await db.refresh(group)
        return group.id


async def _finance_headers(app_client: AsyncClient) -> dict[str, str]:
    from app.core.database import AsyncSessionLocal
    from app.core.security import hash_password
    from app.models.user import User, UserRole

    suffix = uuid.uuid4().hex[:8]
    email = f"excel-finance-{suffix}@example.com"
    password = f"T3st_{suffix}!PassX"
    async with AsyncSessionLocal() as db:
        db.add(
            User(
                email=email,
                password_hash=hash_password(password),
                name=f"Excel Finance {suffix}",
                role=UserRole.finance,
                is_active=True,
                profile_completed=True,
            )
        )
        await db.commit()
    response = await app_client.post(
        "/api/auth/login", json={"email": email, "password": password}
    )
    assert response.status_code == 200, response.text
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


async def test_legacy_excel_export_has_visible_order_and_scoped_ids(
    app_client: AsyncClient, app_auth_headers: dict[str, str]
):
    client_id, order_id, _, consultant = await _seed_order("Klient Eksport")
    other_client_id, other_order_id, _, _ = await _seed_order("Inny Klient")

    response = await app_client.post(
        f"/api/clients/{client_id}/orders/export",
        headers=app_auth_headers,
        json={"order_ids": [order_id]},
    )
    assert response.status_code == 200, response.text
    assert "Zamowienia_Klient_Eksport" in response.headers["content-disposition"]
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert sheet.max_row == 2
    assert sheet["A2"].value == consultant
    assert sheet["C2"].value == 123.456
    assert sheet["D2"].value == 185.75
    assert sheet["E2"].value == "01.01.2026 – 31.12.2026"

    foreign = await app_client.post(
        f"/api/clients/{client_id}/orders/export",
        headers=app_auth_headers,
        json={"order_ids": [other_order_id]},
    )
    assert foreign.status_code == 404, foreign.text
    assert other_client_id != client_id


async def test_unified_excel_export_preserves_mixed_item_order_and_columns(
    app_client: AsyncClient, app_auth_headers: dict[str, str]
):
    client_id, order_id, _, consultant = await _seed_order("Eksport Łączony")
    group_id = await _seed_md_group(client_id, "MD-2026")

    response = await app_client.post(
        f"/api/clients/{client_id}/orders/export",
        headers=app_auth_headers,
        json={
            "items": [
                {"kind": "group", "id": group_id},
                {"kind": "order", "id": order_id},
            ]
        },
    )

    assert response.status_code == 200, response.text
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert [cell.value for cell in sheet[1]] == [
        "Imię i nazwisko",
        "Numer zamówienia",
        "Stawka kosztowa",
        "Stawka przychodowa",
        "Okres zamówienia",
        "Liczba MD / Kwota zamówienia",
        "Zużycie zamówienia",
        "Typ zamówienia",
    ]
    assert sheet.max_row == 3
    assert sheet["B2"].value == "MD-2026"
    assert sheet["F2"].value == 80
    assert sheet["H2"].value == "MD"
    assert sheet["A3"].value == consultant
    assert sheet["H3"].value == "Okresowe"


async def test_unified_excel_export_rejects_duplicate_and_foreign_items(
    app_client: AsyncClient, app_auth_headers: dict[str, str]
):
    client_id, order_id, _, _ = await _seed_order("Eksport Duplikat")
    other_client_id, _, _, _ = await _seed_order("Eksport Obcy")
    foreign_group_id = await _seed_md_group(other_client_id, "MD-FOREIGN")

    duplicate = await app_client.post(
        f"/api/clients/{client_id}/orders/export",
        headers=app_auth_headers,
        json={
            "items": [
                {"kind": "order", "id": order_id},
                {"kind": "order", "id": order_id},
            ]
        },
    )
    assert duplicate.status_code == 422, duplicate.text

    foreign = await app_client.post(
        f"/api/clients/{client_id}/orders/export",
        headers=app_auth_headers,
        json={"items": [{"kind": "group", "id": foreign_group_id}]},
    )
    assert foreign.status_code == 404, foreign.text


async def test_unified_export_keeps_group_and_standalone_permissions_separate(
    app_client: AsyncClient,
):
    client_id, order_id, _, _ = await _seed_order("Eksport Uprawnienia")
    group_id = await _seed_md_group(client_id, "MD-FINANCE")
    headers = await _finance_headers(app_client)

    group = await app_client.post(
        f"/api/clients/{client_id}/orders/export",
        headers=headers,
        json={"items": [{"kind": "group", "id": group_id}]},
    )
    assert group.status_code == 200, group.text

    standalone = await app_client.post(
        f"/api/clients/{client_id}/orders/export",
        headers=headers,
        json={"items": [{"kind": "order", "id": order_id}]},
    )
    assert standalone.status_code == 403, standalone.text
