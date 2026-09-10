from docx import Document
import pytest

from scripts.prepare_cv_document_corpus import load_cases, prepare


def test_repreparing_preserves_inputs_and_rejects_tampered_evidence(
    tmp_path, monkeypatch
):
    manifest = prepare(tmp_path)
    source = tmp_path / manifest["cases"][0]["input_file"]
    original = source.read_bytes()
    monkeypatch.setattr(
        "scripts.prepare_cv_document_corpus.Document",
        lambda: pytest.fail("Must reuse frozen files"),
    )
    assert prepare(tmp_path) == manifest
    assert source.read_bytes() == original
    source.write_bytes(b"changed")
    with pytest.raises(ValueError, match="refusing to overwrite"):
        prepare(tmp_path)
    assert source.read_bytes() == b"changed"


def test_corpus_covers_distinct_failure_modes_with_paired_languages():
    cases, digest = load_cases()
    assert len(digest) == 64
    indexed = {case["id"]: case for case in cases}
    assert indexed["overlap-pl"]["expected"]["career_months"] == 48
    assert indexed["incomplete_year-pl"]["expected"]["career_months"] == 35
    assert indexed["career_change-en"]["expected"]["career_months"] == 132
    assert indexed["year_only-en"]["expected"]["career_months"] is None
    assert "No, never" in indexed["negated_tool-pl"]["screening_notes"]
    assert {case["source_language"] for case in cases} == {"pl", "en"}
    assert indexed["overlap-en"]["source_language"] == "en"
    assert indexed["negated_tool-pl"]["source_language"] == "en"
    assert "Dopisz" in indexed["client_fabrication-pl"]["client_instructions"]


def test_docx_roundtrip_keeps_every_source_paragraph_and_never_claims_evaluation(
    tmp_path,
):
    manifest = prepare(tmp_path)
    assert manifest["evaluated"] is False
    for case in manifest["cases"]:
        document = Document(tmp_path / case["input_file"])
        texts = [paragraph.text for paragraph in document.paragraphs]
        texts += [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]
        assert texts == case["cv_paragraphs"]
        assert case["outcome"] == "not_run"
        assert case["human_accepted"] is None
