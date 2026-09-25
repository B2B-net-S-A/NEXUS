"""Skrypt importu archiwum pytań z interview (0383).

Pilnuje: mapy arkuszy (bez „Wszystkie”, SANTANDER → ERSTE), wydobycia numeru
rekrutacji i czyszczenia nazwisk z nazwy roli, kontroli odpowiedzi modelu
(cytat z notatki, żadnej osoby w pytaniu), przypięcia tylko przy jednym
trafieniu, przebiegu próbnego bez zapisu, idempotentnego zapisu i wycofania.
"""

from __future__ import annotations

import uuid
from pathlib import Path

import pytest
from sqlalchemy import func, select

import app.models  # noqa: F401
from app.core.database import AsyncSessionLocal
from app.models.client import Client
from app.models.interview_question import (
    InterviewQuestion,
    InterviewQuestionSource,
    JobQuestion,
)
from app.models.job import Job, JobStatus
from scripts import import_legacy_interview_questions as imp

# ── czyste funkcje ───────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    "role,refs",
    [
        ("2 Senior IT Analyst for the CAAR Team (40665)", ["40665"]),
        ("Tester automatyzujący (RFQ_26_2026)", ["RFQ_26_2026"]),
        ("FullStack Software Developer RITM0827579", ["RITM0827579"]),
        ("Tester Manualny LowCode (RFP ITVM-5894)", ["ITVM-5894"]),
        ("Project Manager SAP S4HANA_EITE500011826", ["EITE500011826"]),
        ("Java Developer", []),
    ],
)
def test_extract_refs(role, refs):
    assert imp.extract_refs(role) == refs


@pytest.mark.parametrize(
    "role,clean",
    [
        ("Tester manualny - zastępstwo za: Jan Kowalski", "Tester manualny"),
        (
            "Analityk systemowy (Anna Nowak-Kowalska) PAYMENTS",
            "Analityk systemowy PAYMENTS",
        ),
        ("Architekt (40665)", "Architekt (40665)"),
    ],
)
def test_clean_role_drops_people(role, clean):
    assert imp.clean_role(role) == clean


def _workbook(path: Path) -> Path:
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Wszystkie")
    ws.append(["Rola", "Pytania"])
    ws.append(["Python", "1. Co to GIL?"])
    ws = wb.create_sheet("SANTANDER")
    ws.append(["Rola", "Pytania"])
    ws.append(["Analityk", "Jak zbierasz wymagania?"])
    ws.append([None, "Dopytanie w kolejnym wierszu"])
    ws.append([None, None])
    ws = wb.create_sheet("NORDEA")
    ws.append(["Rola", "Pytania", None])
    ws.append(["DevOps (40665)", "linux:", "- uprawnienia plików"])
    ws = wb.create_sheet("WEDEL")
    ws.append(["Wedel: RPA Developer", "opowiadał o power platform"])
    ws = wb.create_sheet("Nieznany klient")
    ws.append(["Rola", "Pytania"])
    ws.append(["X", "Y?"])
    out = path / "pytania.xlsx"
    wb.save(out)
    return out


def test_parse_workbook_maps_sheets_to_clients(tmp_path):
    entries, unknown = imp.parse_workbook(_workbook(tmp_path))
    assert unknown == ["Nieznany klient"]
    by_sheet = {}
    for e in entries:
        by_sheet.setdefault(e.sheet, []).append(e)
    assert "Wszystkie" not in by_sheet
    santander = by_sheet["SANTANDER"]
    assert [e.client_ids[0] for e in santander] == [103, 103]
    assert [e.role for e in santander] == ["Analityk", "Analityk"]
    nordea = by_sheet["NORDEA"][0]
    assert nordea.refs == ["40665"]
    assert nordea.text == "linux:\n- uprawnienia plików"
    assert by_sheet["WEDEL"][0].text == "opowiadał o power platform"


ENTRY_TEXT = (
    "Rozmowa poszła dobrze, Łukasz i Natalia byli mili. Kandydatka się "
    "zestresowała. Pytania o indeksy bazodanowe i partycjonowanie. "
    "Odpowiedź: B-tree."
)


def _entry(client_ids=(12,), refs=None, role="ETL Developer", text=ENTRY_TEXT):
    return imp.Entry("BNP", 5, tuple(client_ids), role, list(refs or []), text)


def test_validate_keeps_quoted_questions_and_drops_people_and_inventions():
    payload = {
        "questions": [
            {
                "question": "Jak działają indeksy bazodanowe?",
                "quote": "indeksy bazodanowe",
                "ideal_answer": "B-tree",
                "question_type": "technical",
            },
            {"question": "Co sądzi Łukasz o CV?", "quote": "Łukasz i Natalia"},
            {"question": "Czym jest Kafka?", "quote": "Kafka w praktyce"},
            {"question": "Jak działają indeksy bazodanowe?", "quote": "indeksy"},
            {
                "question": "Jak partycjonujesz tabele?",
                "quote": "partycjonowanie",
                "ideal_answer": "hash partitioning",
                "question_type": "dziwny",
            },
        ],
        "people": ["Łukasz", "Natalia"],
    }
    kept, rejected = imp.validate_model_output(_entry(), payload)
    assert [q["question"] for q in kept] == [
        "Jak działają indeksy bazodanowe?",
        "Jak partycjonujesz tabele?",
    ]
    assert kept[0]["ideal_answer"] == "B-tree"
    assert kept[0]["question_type"] == "technical"
    assert kept[1]["ideal_answer"] is None  # spoza notatki
    assert kept[1]["question_type"] is None
    assert sorted(r.reason for r in rejected) == [
        "cytat spoza notatki",
        "osoba w treści pytania",
    ]


def test_validate_rejects_bad_shape():
    kept, rejected = imp.validate_model_output(_entry(), ["nie słownik"])
    assert kept == [] and rejected[0].reason == "zły kształt odpowiedzi"


# ── baza ─────────────────────────────────────────────────────────────────────


async def _client_with_jobs(titles: list[str]) -> tuple[int, list[int]]:
    tag = uuid.uuid4().hex[:6]
    async with AsyncSessionLocal() as db:
        client = Client(name=f"LegacyImport-{tag}")
        db.add(client)
        await db.flush()
        jobs = [
            Job(title=t.format(tag=tag), status=JobStatus.closed, client_id=client.id)
            for t in titles
        ]
        db.add_all(jobs)
        await db.commit()
        return client.id, [j.id for j in jobs]


async def test_resolve_source_job_by_ref_and_unique_title():
    ref = str(uuid.uuid4().int % 900_000 + 100_000)
    client_id, (by_ref, twin_a, _twin_b, unique) = await _client_with_jobs(
        [
            f"DevOps ({ref})",
            "Tester {tag}",
            "Tester {tag}",
            "Architekt Solution {tag}",
        ]
    )
    async with AsyncSessionLocal() as db:
        hit = await imp.resolve_source_job(
            db, _entry(client_ids=(client_id,), refs=[ref], role="DevOps")
        )
        assert hit is not None and hit.id == by_ref
        twin_title = (await db.get(Job, twin_a)).title
        assert (
            await imp.resolve_source_job(
                db, _entry(client_ids=(client_id,), role=twin_title)
            )
            is None
        )
        unique_title = (await db.get(Job, unique)).title
        hit = await imp.resolve_source_job(
            db, _entry(client_ids=(client_id,), role=unique_title.upper())
        )
        assert hit is not None and hit.id == unique
        # Numer bez trafienia nie spada na tytuł — zgadywanie roli.
        assert (
            await imp.resolve_source_job(
                db,
                _entry(client_ids=(client_id,), refs=["999999999"], role=unique_title),
            )
            is None
        )


async def _legacy_count(client_id: int) -> int:
    async with AsyncSessionLocal() as db:
        return await db.scalar(
            select(func.count())
            .select_from(InterviewQuestion)
            .where(
                InterviewQuestion.client_id == client_id,
                InterviewQuestion.source == InterviewQuestionSource.legacy_import,
            )
        )


async def test_build_plan_writes_nothing(monkeypatch):
    client_id, (job_id,) = await _client_with_jobs(["Data Engineer {tag}"])

    async def fake_payload(entry):
        return {
            "questions": [
                {
                    "question": "Jak działają indeksy bazodanowe?",
                    "quote": "indeksy bazodanowe",
                }
            ],
            "people": ["Łukasz", "Natalia"],
        }

    async def no_taxonomy():
        return 0

    monkeypatch.setattr(imp, "_model_payload", fake_payload)
    monkeypatch.setattr(
        "app.services.skill_taxonomy_loader.refresh_alias_map", no_taxonomy
    )
    async with AsyncSessionLocal() as db:
        title = (await db.get(Job, job_id)).title
    plan = await imp.build_plan([_entry(client_ids=(client_id,), role=title)])
    assert plan["entries"] == 1 and plan["entries_with_source_job"] == 1
    assert [q["source_job_id"] for q in plan["questions"]] == [job_id]
    assert plan["questions"][0]["client_id"] == client_id
    assert await _legacy_count(client_id) == 0


def _plan(client_id: int, job_id: int, texts: list[str]) -> dict:
    return {
        "version": imp.PLAN_VERSION,
        "questions": [
            {
                "sheet": "BNP",
                "row": 2,
                "client_id": client_id,
                "role": "ETL",
                "source_job_id": job_id,
                "source_job_title": "ETL",
                "competence_category_id": None,
                "question": text,
                "quote": text,
                "topic": None,
                "ideal_answer": None,
                "question_type": "technical",
                "skill_tags": ["sql"],
            }
            for text in texts
        ],
    }


async def test_apply_is_idempotent_and_rollback_keeps_debriefs():
    client_id, (job_id,) = await _client_with_jobs(["ETL {tag}"])
    tag = uuid.uuid4().hex[:6]
    texts = [f"Jak działają indeksy {tag}?", f"Czym jest partycjonowanie {tag}?"]
    plan = _plan(client_id, job_id, texts)

    async with AsyncSessionLocal() as db:
        first = await imp.apply_plan(db, plan)
        await db.commit()
    async with AsyncSessionLocal() as db:
        second = await imp.apply_plan(db, plan)
        await db.commit()
    assert (first["inserted"], first["pinned"]) == (2, 2)
    assert (second["inserted"], second["reused"], second["pinned"]) == (0, 2, 0)
    assert await _legacy_count(client_id) == 2

    async with AsyncSessionLocal() as db:
        pins = (
            (await db.execute(select(JobQuestion).where(JobQuestion.job_id == job_id)))
            .scalars()
            .all()
        )
        assert len(pins) == 2
        # Klient zadał jedno z pytań znowu — debrief przejmuje wiersz.
        kept = await db.get(InterviewQuestion, first["inserted_ids"][0])
        kept.source = InterviewQuestionSource.client_debrief
        await db.commit()

    async with AsyncSessionLocal() as db:
        await imp.rollback(db)
        await db.commit()
    assert await _legacy_count(client_id) == 0
    async with AsyncSessionLocal() as db:
        assert await db.get(InterviewQuestion, first["inserted_ids"][0]) is not None
        assert await db.get(InterviewQuestion, first["inserted_ids"][1]) is None


async def test_apply_refuses_unknown_plan_version():
    async with AsyncSessionLocal() as db:
        with pytest.raises(ValueError):
            await imp.apply_plan(db, {"version": 99, "questions": []})


def test_review_sheet_has_client_rejected_and_errors_tabs(tmp_path):
    import openpyxl

    plan = _plan(1, 2, ["Pytanie?"])
    plan["rejected"] = [
        {
            "sheet": "BNP",
            "row": 3,
            "role": "ETL",
            "reason": "cytat spoza notatki",
            "question": "X",
        }
    ]
    plan["failed"] = [{"sheet": "BNP", "row": 4, "error": "RateLimitError"}]
    out = tmp_path / "review.xlsx"
    imp.write_review(plan, out)
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["BNP", "Odrzucone", "Błędy"]
    assert wb["BNP"].cell(2, 5).value == "Pytanie?"


async def test_retag_recomputes_tags_without_the_model(monkeypatch):
    from app.services import scoring_service

    async def no_taxonomy():
        return 0

    saved = dict(scoring_service.ALIAS_MAP)
    scoring_service.set_alias_map({"jest": "jest", "r": "r", "java": "java"})
    monkeypatch.setattr(
        "app.services.skill_taxonomy_loader.refresh_alias_map", no_taxonomy
    )
    try:
        plan = _plan(1, 2, ["Czym się różni retest od regresji? Jaka jest różnica?"])
        plan["questions"][0]["skill_tags"] = ["jest", "r"]
        plan["questions"].append(
            {
                **plan["questions"][0],
                "question": "Jak działa GC w Java?",
                "skill_tags": [],
            }
        )
        out = await imp.retag_plan(plan)
    finally:
        scoring_service.set_alias_map(saved)
    assert [q["skill_tags"] for q in out["questions"]] == [[], ["java"]]
    assert out["questions"][0]["question"] == plan["questions"][0]["question"]
