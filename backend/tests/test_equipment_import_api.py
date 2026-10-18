"""Import pliku recepcji przez arkusz kontrolny: nic do rejestru przed zatwierdzeniem.

Fikcyjny skoroszyt: laptop osoby z NEXUSA (dopasowanie po imieniu i nazwisku),
laptop osoby spoza NEXUSA (zapisany na osobę spoza NEXUSA), laptop z opisem
zamiast osoby (do wyjaśnienia), wolny telefon, numer już w rejestrze
(konflikt) i wiersz bez numeru (pominięty).
"""

from __future__ import annotations

import io

import openpyxl
import pytest_asyncio
from sqlalchemy import select

from app.core.database import AsyncSessionLocal
from app.models.equipment import (
    EquipmentAssignment,
    EquipmentExternalPerson,
    EquipmentItem,
)
from tests.test_equipment_api import (
    _headers,
    _item,
    _model,
    _tag,
    make_people,
    modules_state,
)


@pytest_asyncio.fixture
async def modules_on():
    async with modules_state(True):
        yield


@pytest_asyncio.fixture
async def people():
    return await make_people()


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _workbook(worker_name: str, taken_serial: str, tag: str) -> bytes:
    workbook = openpyxl.Workbook()
    laptops = workbook.active
    laptops.title = "LAPTOPY"
    laptops.append(
        [
            "Nr\nLaptopa",
            "Model",
            "Numer seryjny / Service tag",
            "Użytkownik",
            "Klucz do Komputera",
        ]
    )
    laptops.append(
        [
            f"A{tag}",
            "Dell Vostro 3530",
            f"IMP1{tag}",
            worker_name,
            "111111-" * 7 + "111111",
        ]
    )
    laptops.append(
        [f"B{tag}", "Dell Vostro 3530", f"IMP2{tag}", "Osoba Spoza Nexusa", None]
    )
    laptops.append([f"C{tag}", "Dell Vostro 3530", taken_serial, None, None])
    laptops.append([None, "Dell Vostro 3530", None, "Ktoś", None])
    laptops.append([f"D{tag}", "Dell Vostro 3530", f"IMP4{tag}", "ktoś z biura", None])
    phones = workbook.create_sheet("TELEFONY")
    phones.append(
        ["Rodzaj urządzenia", "Model", "Numer seryjny / Service tag", "Użytkownik"]
    )
    phones.append(["Telefon", "Redmi A5", f"S/N: IMP3{tag}", "wolny-potwierdzony"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


async def test_import_waits_for_approval_then_creates_items(
    app_client, modules_on, people
) -> None:
    headers = _headers(people["reception"])
    tag = _tag()
    model = await _model(app_client, headers, "laptop", "2000.00")
    taken = f"IMP0{tag}"
    await _item(app_client, headers, model["id"], serial_number=taken)
    worker = people["worker"]

    refused = await app_client.post(
        "/api/equipment/imports",
        files={"file": ("sprzet.xlsx", b"nie excel", XLSX)},
        headers=headers,
    )
    assert refused.status_code == 422
    hr = await app_client.post(
        "/api/equipment/imports",
        files={"file": ("sprzet.xlsx", _workbook(worker.name, taken, tag), XLSX)},
        headers=_headers(people["hr"]),
    )
    assert hr.status_code == 403

    uploaded = await app_client.post(
        "/api/equipment/imports",
        files={"file": ("sprzet.xlsx", _workbook(worker.name, taken, tag), XLSX)},
        headers=headers,
    )
    assert uploaded.status_code == 201, uploaded.text
    batch = uploaded.json()
    assert batch["summary"]["counts"] == {
        "przyjmij": 3,
        "pomin": 1,
        "konflikt": 1,
        "do_wyjasnienia": 1,
    }
    assert batch["summary"]["removed_columns"] == [
        {
            "sheet": "LAPTOPY",
            "column": "Klucz do Komputera",
            "reason": "klucz BitLocker",
        }
    ]
    async with AsyncSessionLocal() as db:
        before = await db.scalar(
            select(EquipmentItem.id).where(EquipmentItem.serial_number == f"IMP1{tag}")
        )
    assert before is None  # nic w rejestrze przed zatwierdzeniem

    detail = (
        await app_client.get(f"/api/equipment/imports/{batch['id']}", headers=headers)
    ).json()
    rows = {
        row["data"].get("serial_number") or row["data"].get("model"): row
        for row in detail["rows"]
    }
    matched = rows[f"IMP1{tag}"]
    assert matched["decision"] == "przyjmij"
    assert matched["holder_user_id"] == worker.id
    assert "111111" not in str(detail)
    outsider = rows[f"IMP2{tag}"]
    assert outsider["decision"] == "przyjmij"
    assert outsider["holder_name"] == "Osoba Spoza Nexusa (spoza NEXUSA)"
    unknown = rows[f"IMP4{tag}"]
    assert unknown["decision"] == "do_wyjasnienia"
    assert rows[taken]["decision"] == "konflikt"
    assert rows[f"IMP3{tag}"]["data"]["item_status"] == "na_stanie"

    sheet = await app_client.get(
        f"/api/equipment/imports/{batch['id']}/control-sheet.xlsx", headers=headers
    )
    assert sheet.status_code == 200
    workbook = openpyxl.load_workbook(io.BytesIO(sheet.content))
    assert workbook.sheetnames == [
        "Podsumowanie",
        "Do przyjęcia",
        "Do wyjaśnienia",
        "Konflikty",
        "Pominięte",
    ]
    assert workbook["Do wyjaśnienia"].max_row == 2

    summary = (await app_client.get("/api/equipment/summary", headers=headers)).json()
    assert summary["import_unresolved"] >= 1

    applied = await app_client.post(
        f"/api/equipment/imports/{batch['id']}/apply", headers=headers
    )
    assert applied.json() == {"created": 3, "conflicts": 0}
    async with AsyncSessionLocal() as db:
        laptop = await db.scalar(
            select(EquipmentItem).where(EquipmentItem.serial_number == f"IMP1{tag}")
        )
        assert laptop.status == "wydany"
        holder = await db.scalar(
            select(EquipmentAssignment.user_id).where(
                EquipmentAssignment.item_id == laptop.id,
                EquipmentAssignment.returned_at.is_(None),
            )
        )
        assert holder == worker.id
        phone = await db.scalar(
            select(EquipmentItem).where(EquipmentItem.serial_number == f"IMP3{tag}")
        )
        assert phone.status == "na_stanie"
        outside = await db.scalar(
            select(EquipmentItem).where(EquipmentItem.serial_number == f"IMP2{tag}")
        )
        assert outside.status == "wydany"
        external = await db.scalar(
            select(EquipmentExternalPerson)
            .join(
                EquipmentAssignment,
                EquipmentAssignment.external_person_id == EquipmentExternalPerson.id,
            )
            .where(EquipmentAssignment.item_id == outside.id)
        )
        assert external.full_name == "Osoba Spoza Nexusa"

    # Recepcja wyjaśnia wiersz: laptop jest wolny w szafce → przyjmij bez osoby.
    decided = await app_client.patch(
        f"/api/equipment/imports/{batch['id']}/rows/{unknown['id']}",
        json={"decision": "przyjmij"},
        headers=headers,
    )
    assert decided.status_code == 200
    again = await app_client.post(
        f"/api/equipment/imports/{batch['id']}/apply", headers=headers
    )
    assert again.json() == {"created": 1, "conflicts": 0}
    third = await app_client.post(
        f"/api/equipment/imports/{batch['id']}/apply", headers=headers
    )
    assert third.json() == {"created": 0, "conflicts": 0}
    locked = await app_client.patch(
        f"/api/equipment/imports/{batch['id']}/rows/{matched['id']}",
        json={"decision": "pomin"},
        headers=headers,
    )
    assert locked.status_code == 409


def _cards_workbook(tag: str) -> bytes:
    """Plik kart jak u recepcji: stara kopia listy PRZED arkuszem „Aktualne”."""
    workbook = openpyxl.Workbook()
    old = workbook.active
    old.title = "Arkusz1"
    old.append(["NR. KARTY", "Karta przypisana", "Poziom dostępu"])
    old.append([f"K1{tag}", None, "ogólny"])
    old.append([f"K2{tag}", None, "ogólny"])
    old.append([f"K9{tag}", None, "ogólny"])
    current = workbook.create_sheet("Aktualne")
    current.append(["NR. KARTY", "Karta przypisana", "Poziom dostępu"])
    current.append([f"K1{tag}", None, "ogólny + biuro"])
    current.append([f"K2{tag}", None, "ogólny + biuro"])
    current.append([f"K3{tag}", None, "ogólny"])
    sales = workbook.create_sheet("Dział sprzedaży")
    sales.append(["NR. KARTY", "Karta przypisana", "Poziom dostępu"])
    sales.append([f"K9{tag}", None, "ogólny"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


async def test_current_sheet_wins_over_its_older_copy(
    app_client, modules_on, people
) -> None:
    """Arkusz „Aktualne” wygrywa z kopią listy stojącą przed nim w pliku.

    Na prawdziwym pliku kart (16.10, próba lokalna) stara kopia „Arkusz1” stała
    przed „Aktualne”, więc 66 aktualnych kart dostało „konflikt”. Kopia z innego
    arkusza ma być pominięta z powodem, a duplikat między dwoma arkuszami bez
    aktualnej listy zostaje konfliktem do decyzji człowieka.
    """
    headers = _headers(people["reception"])
    tag = _tag()
    uploaded = await app_client.post(
        "/api/equipment/imports",
        files={"file": ("karty.xlsx", _cards_workbook(tag), XLSX)},
        headers=headers,
    )
    assert uploaded.status_code == 201, uploaded.text
    detail = (
        await app_client.get(
            f"/api/equipment/imports/{uploaded.json()['id']}", headers=headers
        )
    ).json()
    rows = {
        (row["sheet"], row["data"].get("access_card_number")): row
        for row in detail["rows"]
    }
    for number in (f"K1{tag}", f"K2{tag}", f"K3{tag}"):
        assert rows[("Aktualne", number)]["decision"] == "przyjmij"
    for number in (f"K1{tag}", f"K2{tag}"):
        old = rows[("Arkusz1", number)]
        assert old["decision"] == "pomin"
        assert "Aktualne" in old["reason"]
    # K9 jest w dwóch arkuszach, żaden nie jest aktualną listą — decyduje człowiek.
    assert rows[("Arkusz1", f"K9{tag}")]["decision"] == "przyjmij"
    assert rows[("Dział sprzedaży", f"K9{tag}")]["decision"] == "konflikt"


async def test_control_sheet_neutralises_formulas(
    app_client, modules_on, people
) -> None:
    """Tekst z pliku zaczynający się od „=” nie staje się formułą w arkuszu kontrolnym."""
    headers = _headers(people["reception"])
    tag = _tag()
    workbook = openpyxl.Workbook()
    laptops = workbook.active
    laptops.title = "LAPTOPY"
    laptops.append(
        ["Nr\nLaptopa", "Model", "Numer seryjny / Service tag", "Użytkownik"]
    )
    laptops.append(
        [
            f"F{tag}",
            "Dell Vostro 3530",
            f"FX1{tag}",
            '=HYPERLINK("https://evil.example/","x")',
        ]
    )
    buffer = io.BytesIO()
    workbook.save(buffer)
    uploaded = await app_client.post(
        "/api/equipment/imports",
        files={"file": ("=formula.xlsx", buffer.getvalue(), XLSX)},
        headers=headers,
    )
    assert uploaded.status_code == 201, uploaded.text
    sheet = await app_client.get(
        f"/api/equipment/imports/{uploaded.json()['id']}/control-sheet.xlsx",
        headers=headers,
    )
    assert sheet.status_code == 200
    book = openpyxl.load_workbook(io.BytesIO(sheet.content))
    formulas = [
        cell.coordinate
        for ws in book.worksheets
        for row in ws.iter_rows()
        for cell in row
        if cell.data_type == "f"
    ]
    assert formulas == []


async def test_new_upload_of_same_file_supersedes_unapplied_batch(
    app_client, modules_on, people
) -> None:
    """Drugie wgranie pliku o tej samej nazwie zastępuje niezatwierdzoną partię.

    Wiersze „do wyjaśnienia” starej partii nie mogą dalej wisieć w liczniku
    ekranu (18.10.2026: pierwsza partia głównego pliku miała ich 149).
    """
    headers = _headers(people["reception"])
    tag = _tag()
    name = f"zastap-{tag}.xlsx"
    content = _workbook(people["worker"].name, f"IMPX{tag}", tag)
    first = await app_client.post(
        "/api/equipment/imports", files={"file": (name, content, XLSX)}, headers=headers
    )
    assert first.status_code == 201, first.text
    second = await app_client.post(
        "/api/equipment/imports", files={"file": (name, content, XLSX)}, headers=headers
    )
    assert second.status_code == 201, second.text
    old = (
        await app_client.get(
            f"/api/equipment/imports/{first.json()['id']}", headers=headers
        )
    ).json()
    assert {row["decision"] for row in old["rows"]} == {"pomin"}
    assert old["summary"]["superseded"] is True
    fresh = (
        await app_client.get(
            f"/api/equipment/imports/{second.json()['id']}", headers=headers
        )
    ).json()
    assert "przyjmij" in {row["decision"] for row in fresh["rows"]}
