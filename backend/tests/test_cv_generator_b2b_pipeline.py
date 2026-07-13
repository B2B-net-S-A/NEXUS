"""Unit tests for the CV Generator B2B pipeline helpers.

Covers the 2026-06-10 audit fixes:
  * champion keyword bolding — phrase-aware, word-boundary matching
    (multi-word MUST-HAVEs match, "Git" no longer bolds "digital", and a
    filler phrase like "Znajomość Java" still bolds the real term "Java"),
  * Claude response normalization (missing keys must not KeyError → 500),
  * the anti-fabrication seatbelt ("makijaż, nie inna osoba" — every
    technology/cert in the output must be traceable to CV or notes),
  * filename sanitization (ł → l, not dropped),
  * .doc upload rejection with a clear message.
"""

from __future__ import annotations

import io
import json

import docx as _docx_lib
import pytest

from app.services.cv_generator_b2b.champion_builder import (
    _split_skills,
    from_nexus_job,
    parse_champion_from_docx_bytes,
)
from app.services.cv_generator_b2b.docx_renderer import (
    add_text_with_highlights,
    compile_keyword_patterns,
    highlight_spans,
)
from app.services.skill_normalize import set_tech_taxonomy
from app.services.cv_generator_b2b.standalone_service import (
    StandaloneGenerationError,
    _build_download_filename,
    _close_truncated_json,
    _drop_empty_commas,
    _fabrication_warnings,
    _format_candidate_answers,
    _has_candidate_answers,
    _loads_cv_json,
    _normalize_candidate_data,
    _sanitize_filename_part,
    _sanitize_for_filename,
    _validate_upload,
    ascii_filename_fallback,
    rerender_docx_from_payload,
)


# ── Lenient JSON parsing (Claude 5 prose-wrapper resilience) ────────────────


def test_loads_cv_json_plain_object():
    assert _loads_cv_json('{"name": "Ada"}') == {"name": "Ada"}


def test_loads_cv_json_strips_prose_preamble():
    # Claude 5 sometimes prefixes the JSON with a sentence despite the
    # "return only JSON" instruction — the outermost {...} slice must win.
    wrapped = 'Oto dane kandydata:\n{"name": "Ada", "skills": ["Python"]}\nGotowe.'
    assert _loads_cv_json(wrapped) == {"name": "Ada", "skills": ["Python"]}


def test_loads_cv_json_raises_when_no_object():
    with pytest.raises(json.JSONDecodeError):
        _loads_cv_json("przepraszam, nie mogę tego zrobić")


# ── Empty-value comma repair (the production "Expecting value" failure) ──────


def test_loads_cv_json_trailing_comma_in_array():
    # Reproduces the reported bug: a trailing comma before ``]`` inside an
    # experience[].technologies line → json.loads "Expecting value", now
    # repaired instead of hard-failing the whole generation.
    raw = '{"experience": [{"technologies": ["Java", "Spring", "Docker",]}]}'
    assert _loads_cv_json(raw) == {
        "experience": [{"technologies": ["Java", "Spring", "Docker"]}]
    }


def test_loads_cv_json_trailing_comma_in_object():
    assert _loads_cv_json('{"name": "Ada", "position": "Dev",}') == {
        "name": "Ada",
        "position": "Dev",
    }


def test_loads_cv_json_doubled_and_leading_commas():
    raw = '{"why_points": [,"A",, "B"]}'
    assert _loads_cv_json(raw) == {"why_points": ["A", "B"]}


def test_loads_cv_json_repairs_prose_wrapped_trailing_comma():
    # Both defects at once: chatty preamble AND a trailing comma.
    raw = 'Oto JSON:\n{"skills": ["Python", "SQL",]}\nGotowe.'
    assert _loads_cv_json(raw) == {"skills": ["Python", "SQL"]}


def test_loads_cv_json_closes_truncated_tail():
    # Response cut mid-array (best-effort close). Keeps the complete items.
    raw = '{"experience": [{"company": "ACME", "responsibilities": ["Wdrożenie'
    parsed = _loads_cv_json(raw)
    assert parsed["experience"][0]["company"] == "ACME"
    assert parsed["experience"][0]["responsibilities"] == ["Wdrożenie"]


def test_drop_empty_commas_is_noop_on_valid_json():
    # A comma *between* values, and a comma inside a string literal, must both
    # survive untouched — repair only removes empty-value commas.
    valid = '{"a": [1, 2], "note": "Zrobił X, Y oraz Z"}'
    assert _drop_empty_commas(valid) == valid
    assert json.loads(_drop_empty_commas(valid)) == json.loads(valid)


def test_drop_empty_commas_keeps_comma_inside_string():
    # A trailing-comma pattern that lives INSIDE a string ("...,]") must NOT be
    # stripped — only structural commas count.
    raw = '{"t": "list ends with a comma here ,] literally"}'
    assert _drop_empty_commas(raw) == raw


def test_close_truncated_json_is_noop_when_balanced():
    balanced = '{"a": [1, 2], "b": {"c": "d"}}'
    assert _close_truncated_json(balanced) == balanced


# ── Keyword bolding ────────────────────────────────────────────────────────


def _matches(text: str, keywords: list[str]) -> list[str]:
    patterns = compile_keyword_patterns(keywords)
    return [text[s:e] for s, e in highlight_spans(text, patterns)]


def test_multiword_keyword_matches():
    assert _matches("Budowa pipeline'ów GitLab CI/CD od zera", ["GitLab CI/CD"]) == [
        "GitLab CI/CD"
    ]


def test_multiword_keyword_matches_with_spaced_slash():
    assert _matches("GitLab CI / CD w praktyce", ["GitLab CI/CD"]) == ["GitLab CI / CD"]


def test_substring_does_not_overbold():
    # Old token-based matching bolded "digital" for keyword "Git".
    assert _matches("digital transformation", ["Git"]) == []
    assert _matches("Praca z Git i GitLab", ["Git"]) == ["Git"]


def test_parenthesized_alias_matches_both_forms():
    found = _matches("klaster K8s, potem Kubernetes", ["Kubernetes (K8s)"])
    assert found == ["K8s", "Kubernetes"]


def test_special_char_keywords():
    assert _matches("programowanie w C++ oraz C", ["C++"]) == ["C++"]
    assert _matches("aplikacje .NET i C#", ["C#"]) == ["C#"]


def test_case_insensitive():
    assert _matches("doświadczenie z RABBITMQ", ["RabbitMQ"]) == ["RABBITMQ"]


def test_stopword_keywords_skipped():
    # Champion data sometimes contains noise like "i"/"w" — never bold those.
    assert _matches("praca w zespole i z klientem", ["w", "i"]) == []


def test_add_text_with_highlights_splits_runs():
    class FakeRun:
        def __init__(self) -> None:
            class Font:
                name = None
                color = type("C", (), {"rgb": None})()
                size = None
                bold = None

            self.font = Font()
            self.text = ""

    class FakePara:
        def __init__(self) -> None:
            self.runs: list[tuple[str, FakeRun]] = []

        def add_run(self, text: str) -> FakeRun:
            run = FakeRun()
            self.runs.append((text, run))
            return run

    para = FakePara()
    add_text_with_highlights(para, "Docker, Kubernetes i Bash", ["Kubernetes"])
    texts = [t for t, _ in para.runs]
    bolds = [r.font.bold for _, r in para.runs]
    assert texts == ["Docker, ", "Kubernetes", " i Bash"]
    assert bolds == [None, True, None]


def test_descriptive_parenthetical_bolds_only_base_term():
    # The screenshot bug: "Figma (zaawansowana znajomość)" used to split into a
    # bold keyword "zaawansowana znajomość" that then bolded those generic
    # words throughout the CV prose. Now only the real skill bolds.
    kw = ["Figma (zaawansowana znajomość)"]
    assert _matches("Projektowanie w Figma i Sketch", kw) == ["Figma"]
    # The qualifier must NEVER bold on its own.
    assert _matches("zaawansowana znajomość narzędzi UX", kw) == []


def test_acronym_parenthetical_still_splits():
    # Genuine aliases ("K8s", ".NET", "PL/SQL") must keep bolding both forms.
    assert _matches("klaster K8s, potem Kubernetes", ["Kubernetes (K8s)"]) == [
        "K8s",
        "Kubernetes",
    ]
    assert _matches("aplikacje .NET", ["Microsoft (.NET)"]) == [".NET"]


def test_generic_single_word_keyword_skipped():
    # Champion lists sometimes contain proficiency/filler entries on their own.
    assert _matches("zaawansowana znajomość Figma", ["znajomość"]) == []
    assert _matches("Docker mile widziane", ["mile widziane"]) == []
    assert _matches("bardzo dobra znajomość AWS", ["bardzo dobra znajomość"]) == []


def test_hyphen_space_spelling_drift_matches():
    # Champion list vs CV often disagree on hyphen vs space — match both ways
    # (multi-word tech "react native" ↔ "react-native").
    assert _matches("Aplikacje w react native", ["react-native"]) == ["react native"]
    assert _matches("Zaawansowany react-native", ["react native"]) == ["react-native"]


def test_filler_phrase_bolds_real_term():
    # "Bolds only part" bug: a champion entry written as natural language
    # ("Znajomość Java") never matched because the CV writes just "Java".
    # Now the filler is stripped and the real term still bolds.
    assert _matches("Programowanie w Java i Spring", ["Znajomość Java"]) == ["Java"]
    assert _matches("Backend w Spring Boot", ["Dobra znajomość Spring Boot"]) == [
        "Spring Boot"
    ]


def test_separate_skill_chips_each_bold():
    # Champion skills arrive as separate chips — each bolds where present.
    assert _matches("Stack: Java, Python i Go", ["Java", "Python"]) == [
        "Java",
        "Python",
    ]


def test_versioned_tech_bolds_bare_brand():
    # Reported bug: a champion skill carrying a version ("Java 17+") matched
    # only verbatim, so plain "Java" in the CV body never bolded. When the
    # compound is ONE real technology plus a version/qualifier token that
    # carries no letter ("17+", "17"), the bare brand is recovered and bolds.
    body = "Programowanie w Java i Spring Boot, mikroserwisy w Java."
    for kw in ("Java 17+", "Java 17"):
        assert _matches(body, [kw]) == ["Java", "Java"], kw
    # A bare digit/version token never bolds on its own (needs a letter).
    assert _matches("wersja 17 systemu", ["Java 17+"]) == []
    # ...but a two-real-word compound ("Java SE") is a phrase, not a lone
    # brand + junk, so it bolds only whole — bare "Java" does NOT leak out.
    assert _matches(body, ["Java SE 17"]) == []
    assert _matches("Migracja na Java SE 17 w projekcie", ["Java SE 17"]) == [
        "Java SE 17"
    ]


def test_multiword_product_name_bolds_only_whole():
    # A genuine multi-word product name whose words are EACH a technology
    # ("REST API") bolds ONLY as the whole phrase — never split into single
    # tokens that would bold unrelated occurrences. Recruiter directive: bold
    # the listed phrase, not words torn from it.
    assert _matches("Integracje przez REST oraz samodzielne API", ["REST API"]) == []
    assert _matches("Pełne REST API w projekcie", ["REST API"]) == ["REST API"]


def test_shared_vendor_prefix_never_leaks_to_other_products():
    # Reported bug: champion must-have "Apache Airflow" + "Apache Kafka" got
    # split into a standalone "Apache" pattern, which then bolded the shared
    # "Apache" prefix of NON-required products ("Apache NiFi", "Apache Spark").
    # A two-tech compound now bolds only whole, so the vendor prefix never
    # leaks — only the two listed phrases bold, nothing on Spark/NiFi.
    body = (
        "Delivered platforms using Apache Kafka, Apache Spark, Apache Airflow, "
        "and Apache NiFi for analytics."
    )
    assert _matches(body, ["Apache Airflow", "Apache Kafka"]) == [
        "Apache Kafka",
        "Apache Airflow",
    ]


def test_multiword_technology_stays_whole_concept_drops():
    # A genuine multi-word TECHNOLOGY bolds whole ("Spring Boot") — never its
    # bare common part ("boot"). A multi-word CONCEPT ("Design System") is not
    # a technology, so it must NOT bold at all (tech-only directive).
    assert _matches("Backend w Spring Boot dla klienta", ["Spring Boot"]) == [
        "Spring Boot"
    ]
    assert _matches("Tworzenie Design System dla klienta", ["Design System"]) == []
    assert _matches("migracja systemu do nowej wersji", ["Design System"]) == []


def test_verbose_requirement_never_bolds_generic_words():
    # The over-bolding recruiters reported: a wordy champion entry stays whole
    # and only matches verbatim, so its generic words never bold the CV prose.
    kw = ["Tworzenie i rozwój aplikacji webowych"]
    assert _matches("Tworzenie dokumentacji projektu", kw) == []
    assert _matches("dbam o rozwój zespołu i procesów", kw) == []


def test_requirement_prose_words_never_bold():
    # User-cited noise: these must never bold, even as single champion entries.
    assert (
        _matches("Posiada kluczowe technologie wymagane na stanowisku", ["Technologie"])
        == []
    )
    assert _matches("praca na stanowisku starszego developera", ["stanowisku"]) == []
    assert _matches("Tworzenie i utrzymanie usług", ["Tworzenie"]) == []
    # ...but a real skill wrapped in requirement prose still bolds.
    assert _matches("Posiada Java i Spring na stanowisku", ["Znajomość Java"]) == [
        "Java"
    ]


def test_verbose_requirement_bolds_only_buried_tech():
    # Recruiter directive: bold ONLY technologies. A wordy non-tech requirement
    # ("bazami danych SQL") never bolds the Polish prose — only the hard-tech
    # token buried inside it ("SQL") surfaces.
    kw = ["Doświadczenie z bazami danych SQL"]
    assert _matches("Pracował z bazami danych SQL i NoSQL", kw) == ["SQL"]
    # No tech token present → an unrelated prose line stays clean.
    assert _matches("migracja systemów do chmury", kw) == []


def test_strong_tech_tokens_pulled_from_requirement_prose():
    # Special-char / digit / acronym tokens are unmistakably tech, so they bold
    # even when the surrounding phrase is inflected differently in the CV.
    assert _matches("Aplikacje w C++ i Java", ["Programowanie w C++"]) == ["C++"]
    assert _matches("Backend w .NET", ["Tworzenie usług w .NET"]) == [".NET"]
    assert _matches("Praca z SQL Server", ["Microsoft SQL Server"]) == ["SQL"]


def test_non_tech_substantive_phrase_does_not_bold():
    # A substantive but non-technology requirement ("aplikacji webowych") must
    # NOT bold — it carries no tech token. Only technologies bold.
    kw = ["Tworzenie aplikacji webowych"]
    assert _matches("Projektowanie aplikacji webowych dla klienta", kw) == []
    assert _matches("Tworzenie dokumentacji technicznej", kw) == []


def test_champion_design_chips_only_hard_tech_bolds():
    # Real UI-Designer champion MUST-HAVE chips: "KONCEPT (pod-terminy)
    # (długie wyjaśnienie)". Recruiter directive — bold ONLY technologies:
    # platforms ("iOS/Android") and technical acronyms/standards ("RWD",
    # "WCAG"). Design methodologies / concepts ("User-Centered Design",
    # "Material Design", "Zasady Gestalt", "visual design", "design systemów")
    # must NOT bold.
    must = [
        "User-Centered Design (projektowanie w oparciu o potrzeby użytkownika, "
        "badania, feedback i dane, a nie wyłącznie wymagania biznesowe)",
        "Zasady projektowania (heurystyki, best practices UX) (np. spójność "
        "ekranów, przewidywalność zachowań aplikacji)",
        "Zasady Gestalt (hierarchia, percepcja, grupowanie) (umiejętność "
        "układania elementów na ekranie)",
        "Zaawansowany visual design (typografia, kolor, kompozycja, gridy) "
        "(projektowanie estetycznych interfejsów)",
        "Material Design oraz Human Interface Guidelines (znajomość standardów "
        "projektowania aplikacji dla Androida i iOS)",
        "Projektowanie aplikacji mobilnych (iOS/Android) (doświadczenie w "
        "projektowaniu ekranów)",
        "Responsywne projektowanie (RWD, multi-device) (tworzenie rozwiązań "
        "działających poprawnie)",
        "Zasady WCAG 2.1/2.2 (projektowanie zgodne z wymaganiami dostępności)",
        "Tworzenie i rozwój design systemów (komponenty, tokeny) "
        "(projektowanie i utrzymywanie bibliotek)",
    ]
    prose = (
        "Stosuje User-Centered Design, zasady projektowania oraz zasady Gestalt. "
        "Zaawansowany visual design; zna Material Design i Human Interface "
        "Guidelines. Projektowanie aplikacji mobilnych (iOS/Android), "
        "responsywne projektowanie (RWD). Zna zasady WCAG 2.1/2.2 oraz "
        "tworzenie i rozwój design systemów (komponenty, tokeny)."
    )
    found = [f.lower() for f in _matches(prose, must)]
    # Technologies / platforms / standards bold.
    assert "ios/android" in found, found
    assert "rwd" in found, found
    assert "wcag" in found, found
    # Methodologies / concepts / requirement prose must NEVER bold.
    for concept in [
        "user-centered",
        "user-centered design",
        "zasady projektowania",
        "gestalt",
        "visual design",
        "material design",
        "human interface",
        "design systemów",
        "komponenty",
        "tokeny",
        "tworzenie",
        "projektowanie",
    ]:
        assert all(concept not in f for f in found), (
            f"{concept!r} must not bold; found={found}"
        )


def test_only_listed_technologies_bold():
    # The headline requirement: in a mixed must-have/nice-to-have list, only
    # the concrete technologies bold; the design concepts alongside them do not.
    keywords = [
        "Figma (zaawansowana znajomość)",
        "React",
        "Doświadczenie w Python",
        "Docker oraz Kubernetes",
        "Znajomość SQL i PostgreSQL",
        "User-Centered Design",  # concept — must NOT bold
        "Zasady Gestalt",  # concept — must NOT bold
    ]
    prose = (
        "Projektant pracujący w Figma, buduje aplikacje w React i Python. "
        "Konteneryzacja: Docker, Kubernetes. Bazy: SQL, PostgreSQL. "
        "Stosuje User-Centered Design oraz zasady Gestalt."
    )
    found = [f.lower() for f in _matches(prose, keywords)]
    for tech in [
        "figma",
        "react",
        "python",
        "docker",
        "kubernetes",
        "sql",
        "postgresql",
    ]:
        assert tech in found, f"{tech!r} should bold; found={found}"
    joined = " ".join(found)
    assert "user-centered" not in joined
    assert "gestalt" not in joined


def test_multiword_microsoft_products_bold_whole():
    # Multi-word brand names (no per-token hard-tech signal) must bold WHOLE,
    # not partially. Before the allowlist extension "Microsoft Intune" and
    # "Microsoft 365" produced no span at all (recruiter: a key technology must
    # bold everywhere it appears).
    found = _matches(
        "Wdrożenie Microsoft Intune oraz migracja do Microsoft 365",
        ["Microsoft Intune", "Microsoft 365"],
    )
    assert found == ["Microsoft Intune", "Microsoft 365"]


def test_key_technology_bolds_consistently_across_sections():
    # The same champion technology must bold in EVERY place it appears — the
    # "why" headline, the skills list, a responsibility line and the per-role
    # "Technologie:" line all run through the same patterns.
    keywords = ["Microsoft Intune", "Active Directory"]
    sections = [
        "Posiada kluczowe technologie: Microsoft Intune, Active Directory",  # why
        "Zarządzanie urządzeniami: Microsoft Intune",  # skills
        "Konfiguracja zasad w Microsoft Intune i Active Directory",  # duty
        "Microsoft Intune, Active Directory, Azure AD",  # tech line
    ]
    for text in sections:
        found = {f.lower() for f in _matches(text, keywords)}
        assert "microsoft intune" in found, f"not bold in: {text!r}"
    # And it appears in every section that contains it.
    assert all("Microsoft Intune" in s for s in sections[:4])


def test_curated_multiword_suppresses_stray_token_bold():
    # A curated entry bolds the phrase WHOLE and skips token-splitting, so the
    # generic tail ("AD") no longer bolds on its own elsewhere in the CV.
    found = _matches("Integracja z Azure AD; osobny dział AD HR", ["Azure AD"])
    assert found == ["Azure AD"]


def test_allowlist_extension_does_not_bold_concepts():
    # Lock guard: adding product names must not regress the "only technologies"
    # rule — a design methodology alongside them still never bolds.
    keywords = ["Microsoft Intune", "User-Centered Design", "Material Design"]
    prose = "Wdraża Microsoft Intune; stosuje User-Centered Design i Material Design."
    found = [f.lower() for f in _matches(prose, keywords)]
    assert "microsoft intune" in found
    joined = " ".join(found)
    assert "user-centered" not in joined and "material design" not in joined


# ── Claude response normalization ──────────────────────────────────────────


def test_normalize_fills_missing_keys():
    data = _normalize_candidate_data(
        {
            "name": "Jan Kowalski",
            "experience": [{"company": "ACME"}],
            "education": [{"institution": "PW"}],
            "skills": [{"label": "DevOps:", "content": "Docker"}, "luźny string"],
        },
        fallback_name=None,
    )
    job = data["experience"][0]
    assert job["dates"] == ""
    assert job["responsibilities"] == []
    assert job["technologies"] == []
    edu = data["education"][0]
    assert edu["degree"] == ""
    assert data["skills"][1] == {"label": "", "content": "luźny string"}
    assert data["why_points"] == []
    assert data["languages"] == []


def test_normalize_uses_fallback_name():
    data = _normalize_candidate_data({}, fallback_name="Ada Nowak")
    assert data["name"] == "Ada Nowak"


def test_normalize_rejects_non_dict():
    with pytest.raises(StandaloneGenerationError) as exc:
        _normalize_candidate_data(["lista", "zamiast", "obiektu"], fallback_name=None)
    assert exc.value.code == "ai_failed"


def test_normalize_folds_typographic_dashes_to_hyphens():
    data = _normalize_candidate_data(
        {
            "name": "Jan Kowalski",
            "why_points": ["5 lat doświadczenia — w tym 3 w fintechu"],
            "languages": ["Polski – ojczysty", "Angielski – biegły"],
            "certifications": ["AWS SAA — Amazon (2022)"],
            "experience": [
                {
                    "company": "ACME",
                    "dates": "03.2020 – obecnie",
                    "responsibilities": ["Wdrożenie CI/CD — GitLab"],
                }
            ],
            "skills": [{"label": "DevOps:", "content": "Docker – zaawansowany"}],
        },
        fallback_name=None,
    )
    assert data["why_points"] == ["5 lat doświadczenia - w tym 3 w fintechu"]
    assert data["languages"] == ["Polski - ojczysty", "Angielski - biegły"]
    assert data["certifications"] == ["AWS SAA - Amazon (2022)"]
    job = data["experience"][0]
    assert job["dates"] == "03.2020 - obecnie"
    assert job["responsibilities"] == ["Wdrożenie CI/CD - GitLab"]
    assert data["skills"][0]["content"] == "Docker - zaawansowany"
    # No typographic dash survives anywhere in the output.
    blob = repr(data)
    for dash in ("–", "—", "‒", "―", "−"):
        assert dash not in blob


# ── Anti-fabrication guard ─────────────────────────────────────────────────

_SOURCE = """
Senior DevOps Engineer z doświadczeniem w Kubernetes, Docker oraz Terraform.
Certyfikat CKA (Certified Kubernetes Administrator) zdobyty w 2023.
Notatka rekrutera: kandydat potwierdził znajomość RabbitMQ na screeningu.
"""


def test_guard_passes_for_sourced_technologies():
    data = {
        "experience": [
            {
                "position": "DevOps",
                "technologies": ["Kubernetes", "Docker", "RabbitMQ"],
            }
        ],
        "certifications": ["CKA – Certified Kubernetes Administrator (2023)"],
    }
    assert _fabrication_warnings(data, _SOURCE, "pl") == []


def test_guard_flags_fabricated_technology():
    data = {
        "experience": [{"position": "DevOps", "technologies": ["Snowflake"]}],
        "certifications": [],
    }
    warnings = _fabrication_warnings(data, _SOURCE, "pl")
    assert len(warnings) == 1
    assert "Snowflake" in warnings[0]
    assert warnings[0].startswith("WERYFIKUJ")


def test_guard_flags_fabricated_certification():
    data = {
        "experience": [],
        "certifications": ["AWS Solutions Architect Professional"],
    }
    warnings = _fabrication_warnings(data, _SOURCE, "pl")
    assert len(warnings) == 1
    assert "AWS" in warnings[0]


def test_guard_handles_diacritics():
    data = {
        "experience": [
            {"position": "Dev", "technologies": ["zarządzanie łańcuchem CI"]}
        ],
        "certifications": [],
    }
    source = "Zarzadzanie lancuchem CI w GitLabie"
    assert _fabrication_warnings(data, source, "pl") == []


# ── Candidate screening answers (fed into the CV as notes context) ─────────

_QUESTIONS = [
    {"id": "q1", "question": "Ile lat pracowałeś z Kubernetes?"},
    {"id": "q2", "question": "Czy prowadziłeś migracje do chmury?"},
]


def test_candidate_answers_pairs_question_with_response():
    answers = {
        "answers": [
            {"question_id": "q1", "response": "5 lat, głównie na produkcji"},
            {"question_id": "q2", "response": "Tak, dwie migracje do AWS"},
        ]
    }
    text = _format_candidate_answers(answers, _QUESTIONS)
    assert "P: Ile lat pracowałeś z Kubernetes?" in text
    assert "O: 5 lat, głównie na produkcji" in text
    assert "P: Czy prowadziłeś migracje do chmury?" in text
    assert "O: Tak, dwie migracje do AWS" in text


def test_candidate_answers_skips_blank_responses():
    answers = {
        "answers": [
            {"question_id": "q1", "response": "  "},
            {"question_id": "q2", "response": "Tak, dwie migracje do AWS"},
        ]
    }
    text = _format_candidate_answers(answers, _QUESTIONS)
    assert "Kubernetes" not in text  # blank answer dropped along with its question
    assert text == "P: Czy prowadziłeś migracje do chmury?\nO: Tak, dwie migracje do AWS"


def test_candidate_answers_answer_only_when_question_missing():
    # Question id no longer present on the Champion Profile — keep the answer,
    # drop the dangling question label rather than emitting an empty "P:".
    answers = {"answers": [{"question_id": "gone", "response": "Node.js i Go"}]}
    text = _format_candidate_answers(answers, _QUESTIONS)
    assert text == "O: Node.js i Go"


def test_candidate_answers_excludes_recruiter_judgment():
    # overall_fit / deal_breaker_hit / recruiter notes must never reach the CV.
    answers = {
        "answers": [
            {"question_id": "q1", "response": "3 lata", "deal_breaker_hit": True}
        ],
        "overall_fit": "miss",
        "notes": "Sceptyczny, słaby angielski",
    }
    text = _format_candidate_answers(answers, _QUESTIONS)
    assert text == "P: Ile lat pracowałeś z Kubernetes?\nO: 3 lata"
    assert "miss" not in text
    assert "angielski" not in text


def test_candidate_answers_empty_and_malformed_inputs():
    assert _format_candidate_answers(None, _QUESTIONS) == ""
    assert _format_candidate_answers({}, _QUESTIONS) == ""
    assert _format_candidate_answers({"answers": []}, _QUESTIONS) == ""
    # A candidate answer with no question map at all still renders.
    assert _format_candidate_answers({"answers": [{"response": "Java"}]}, None) == (
        "O: Java"
    )


def test_has_candidate_answers_flag():
    assert _has_candidate_answers({"answers": [{"response": "cokolwiek"}]}) is True
    assert _has_candidate_answers({"answers": [{"response": "   "}]}) is False
    assert _has_candidate_answers({"answers": []}) is False
    assert _has_candidate_answers(None) is False
    assert _has_candidate_answers("not-a-dict") is False


# ── Filename sanitization ──────────────────────────────────────────────────


def test_sanitize_preserves_polish_letters():
    # The on-disk filename keeps the candidate's real spelling — Polish letters
    # (ł, ą, ż, ó…) survive; only whitespace is normalized to underscores.
    assert _sanitize_for_filename("Łukasz Kamecki") == "Łukasz_Kamecki"
    assert _sanitize_for_filename("Michał Bogdan") == "Michał_Bogdan"
    assert _sanitize_for_filename("Kamil Szukajło") == "Kamil_Szukajło"


def test_sanitize_replaces_unsafe_characters():
    assert _sanitize_for_filename("Jan/Kowalski") == "Jan_Kowalski"
    assert _sanitize_for_filename("   ") == "kandydat"


def test_ascii_filename_fallback_transliterates_polish():
    # The Content-Disposition ASCII fallback still folds Polish → ASCII so the
    # legacy ``filename="…"`` parameter stays latin-1 encodable.
    assert (
        ascii_filename_fallback("CV_B2B_Łukasz_Kamecki.docx")
        == "CV_B2B_Lukasz_Kamecki.docx"
    )
    assert (
        ascii_filename_fallback("CV_B2B_Kamil_Szukajło.docx")
        == "CV_B2B_Kamil_Szukajlo.docx"
    )


def test_sanitize_filename_part_preserves_spaces_and_diacritics():
    # The role/name component keeps spaces, Polish letters and tech tokens like
    # ``+``/``#`` — only runs of whitespace collapse to a single space.
    assert _sanitize_filename_part("IT Analyst") == "IT Analyst"
    assert _sanitize_filename_part("  IT   Analyst  ") == "IT Analyst"
    assert _sanitize_filename_part("C++ Developer") == "C++ Developer"
    assert _sanitize_filename_part("C# Developer") == "C# Developer"
    assert _sanitize_filename_part("Łukasz Kamecki") == "Łukasz Kamecki"


def test_sanitize_filename_part_strips_reserved_chars():
    # Filesystem-reserved characters are removed without leaving double spaces.
    assert _sanitize_filename_part("Data Engineer / ETL") == "Data Engineer ETL"
    assert _sanitize_filename_part('Role: "x"?') == "Role x"
    assert _sanitize_filename_part("   ") == ""


def test_build_download_filename_uses_role():
    # With a recruitment role the download name is "{role}_{name}.docx" and the
    # user-facing example spelling (spaces intact) is preserved.
    assert (
        _build_download_filename("IT Analyst", "Jan Kowalski")
        == "IT Analyst_Jan Kowalski.docx"
    )
    # Polish letters survive on disk (Content-Disposition folds ASCII separately).
    assert (
        _build_download_filename("Analityk IT", "Kamil Szukajło")
        == "Analityk IT_Kamil Szukajło.docx"
    )


def test_build_download_filename_falls_back_without_role():
    # No role (manual-upload mode / blank title) → legacy "CV_B2B_{name}.docx".
    assert _build_download_filename(None, "Jan Kowalski") == "CV_B2B_Jan_Kowalski.docx"
    assert _build_download_filename("   ", "Jan Kowalski") == "CV_B2B_Jan_Kowalski.docx"


# ── Upload validation ──────────────────────────────────────────────────────


def test_doc_upload_rejected_with_clear_message():
    with pytest.raises(StandaloneGenerationError) as exc:
        _validate_upload(
            b"\xd0\xcf\x11\xe0 fake doc",
            "cv_stare.doc",
            allowed_ext={".pdf", ".docx"},
            label="CV",
        )
    assert exc.value.code == "invalid_input"
    assert ".docx lub PDF" in exc.value.message


def test_pdf_upload_accepted():
    _validate_upload(
        b"%PDF-1.4 fake", "cv.pdf", allowed_ext={".pdf", ".docx"}, label="CV"
    )


# ── Per-role technology cap ────────────────────────────────────────────────


def test_cap_technologies_prioritizes_champion():
    from app.services.cv_generator_b2b.standalone_service import (
        _cap_role_technologies,
    )

    techs = [f"Tech{i}" for i in range(1, 12)] + ["Splunk", "QRadar", "Tech12"]
    data = {"experience": [{"position": "x", "technologies": list(techs)}]}
    _cap_role_technologies(data, ["Splunk", "QRadar"])
    kept = data["experience"][0]["technologies"]
    assert len(kept) == 12
    assert "Splunk" in kept and "QRadar" in kept
    # Original ordering preserved within the kept subset.
    assert kept.index("Tech1") < kept.index("Splunk")


def test_cap_technologies_noop_under_limit():
    from app.services.cv_generator_b2b.standalone_service import (
        _cap_role_technologies,
    )

    data = {"experience": [{"position": "x", "technologies": ["A", "B"]}]}
    _cap_role_technologies(data, ["A"])
    assert data["experience"][0]["technologies"] == ["A", "B"]


# ── Overlapping employment dates ───────────────────────────────────────────


def test_parse_date_range_variants():
    from app.services.cv_generator_b2b.standalone_service import _parse_date_range

    assert _parse_date_range("03.2020 – 11.2023") == (2020 * 12 + 2, 2023 * 12 + 10)
    assert _parse_date_range("2019") == (2019 * 12, 2019 * 12 + 11)
    ongoing = _parse_date_range("02.2024 – obecnie")
    assert ongoing is not None and ongoing[1] == 9999 * 12
    assert _parse_date_range("brak dat") is None
    assert _parse_date_range("") is None


def test_overlap_warning_flags_parallel_roles():
    from app.services.cv_generator_b2b.standalone_service import (
        _date_overlap_warnings,
    )

    data = {
        "experience": [
            {"company": "Contina", "dates": "02.2024 – obecnie"},
            {"company": "COI", "dates": "11.2023 – 02.2025"},
        ]
    }
    warnings = _date_overlap_warnings(data, "pl")
    assert len(warnings) == 1
    assert "Contina" in warnings[0] and "COI" in warnings[0]
    assert warnings[0].startswith("WERYFIKUJ")


def test_overlap_ignores_handover_month():
    from app.services.cv_generator_b2b.standalone_service import (
        _date_overlap_warnings,
    )

    data = {
        "experience": [
            {"company": "A", "dates": "01.2020 – 03.2022"},
            {"company": "B", "dates": "03.2022 – obecnie"},
        ]
    }
    assert _date_overlap_warnings(data, "pl") == []


def test_overlap_caps_warning_count():
    from app.services.cv_generator_b2b.standalone_service import (
        _date_overlap_warnings,
    )

    data = {
        "experience": [
            {"company": f"Firma{i}", "dates": "01.2020 – obecnie"} for i in range(5)
        ]
    }
    warnings = _date_overlap_warnings(data, "pl")
    assert len(warnings) == 4  # 3 pary + "… i N kolejnych"
    assert "kolejnych" in warnings[-1]


# ── Exact years of experience ──────────────────────────────────────────────


def test_total_experience_years_counts_full_tenure():
    from app.services.cv_generator_b2b.standalone_service import (
        _total_experience_years,
    )

    # 09.2018 – 09.2023 = 61 months ≈ 5.08 → 5 (the recruiter's "5 years").
    assert _total_experience_years([{"dates": "09.2018 – 09.2023"}]) == 5


def test_total_experience_years_merges_overlapping_roles():
    from app.services.cv_generator_b2b.standalone_service import (
        _total_experience_years,
    )

    # Parallel B2B contracts must be counted once, not summed to 10.
    years = _total_experience_years(
        [
            {"dates": "01.2019 – 12.2023"},
            {"dates": "01.2019 – 12.2023"},
        ]
    )
    assert years == 5


def test_total_experience_years_none_without_dates():
    from app.services.cv_generator_b2b.standalone_service import (
        _total_experience_years,
    )

    assert _total_experience_years([{"company": "X"}]) is None


def test_fix_experience_years_corrects_undercount():
    from app.services.cv_generator_b2b.standalone_service import (
        _fix_experience_years,
    )

    data = {
        "experience": [{"dates": "09.2018 – 09.2023"}],
        "why_points": [
            "Ponad 4 lata doświadczenia jako UX/UI Designer, w tym 3 lata w startupie",
            "Specjalizacja w technologiach: Figma, Sketch",
        ],
    }
    _fix_experience_years(data, "pl")
    # Headline total corrected to the exact figure; "w tym 3 lata" sub-figure
    # and the second point are left untouched.
    assert data["why_points"][0].startswith("5 lat doświadczenia jako UX/UI Designer")
    assert "w tym 3 lata w startupie" in data["why_points"][0]
    assert data["why_points"][1] == "Specjalizacja w technologiach: Figma, Sketch"


def test_fix_experience_years_polish_plural_unit():
    from app.services.cv_generator_b2b.standalone_service import (
        _fix_experience_years,
    )

    data = {
        "experience": [{"dates": "01.2022 – 12.2023"}],  # 24 months → 2 → "2 lata"
        "why_points": ["Ponad 1 rok doświadczenia jako Developer"],
    }
    _fix_experience_years(data, "pl")
    assert data["why_points"][0] == "2 lata doświadczenia jako Developer"


def test_fix_experience_years_skips_technology_specific_point():
    # A 6-year-total candidate who used Intune for only part of that time must
    # not have "2 lata z Microsoft Intune" rewritten to the total "6 lat z
    # Intune". The total belongs on the generic role headline instead.
    from app.services.cv_generator_b2b.standalone_service import (
        _fix_experience_years,
    )

    data = {
        "experience": [
            {  # 01.2018 – 12.2023 = 72 months → 6 years total
                "dates": "01.2018 – 12.2023",
                "technologies": ["Microsoft Intune", "Azure AD"],
            }
        ],
        "highlight_keywords": ["Microsoft Intune", "Microsoft Endpoint Manager"],
        "why_points": [
            "2 lata doświadczenia z Microsoft Intune i Microsoft Endpoint Manager",
            "Ponad 4 lata doświadczenia jako Specjalista MDM",
        ],
    }
    _fix_experience_years(data, "pl")
    # Tech-specific duration left untouched; total lands on the role headline.
    assert data["why_points"][0] == (
        "2 lata doświadczenia z Microsoft Intune i Microsoft Endpoint Manager"
    )
    assert data["why_points"][1] == "6 lat doświadczenia jako Specjalista MDM"


def test_fix_experience_years_never_inflates_lone_technology_point():
    # When the ONLY experience point is technology-specific, the recompute must
    # leave it alone rather than clobber the tech duration with the total.
    from app.services.cv_generator_b2b.standalone_service import (
        _fix_experience_years,
    )

    data = {
        "experience": [
            {"dates": "01.2019 – 12.2023", "technologies": ["Kubernetes"]}
        ],  # 5 years total
        "why_points": ["Ponad 3 lata doświadczenia z Kubernetes"],
    }
    _fix_experience_years(data, "pl")
    assert data["why_points"][0] == "Ponad 3 lata doświadczenia z Kubernetes"


def test_fix_experience_years_still_fixes_industry_duration():
    # An industry/domain duration ("w fintechu") is NOT a technology binding —
    # the total-tenure correction must still apply.
    from app.services.cv_generator_b2b.standalone_service import (
        _fix_experience_years,
    )

    data = {
        "experience": [
            {"dates": "01.2019 – 12.2023", "technologies": ["Java"]}
        ],  # 5 years total
        "why_points": ["Ponad 4 lata doświadczenia w fintechu"],
    }
    _fix_experience_years(data, "pl")
    assert data["why_points"][0] == "5 lat doświadczenia w fintechu"


def test_fix_experience_years_never_lands_on_company_subfigure():
    # Regression (Grzegorz Zieliński CV): "5+ years" is not parseable as the
    # headline figure, so the search used to drift into the "including 5 years
    # at Gigaset" sub-figure and stamp the 8-year career total there — a tenure
    # the candidate never had. The whole point must survive as written.
    from app.services.cv_generator_b2b.standalone_service import (
        _fix_experience_years,
    )

    data = {
        "experience": [
            {"company": "CGI / Fiserv", "dates": "06.2024 - currently"},
            {"company": "Gigaset", "dates": "03.2020 - 05.2024"},
            {"company": "Gigaset", "dates": "2019 - 2020"},
        ],  # merged career ≈ 8 years, Gigaset itself ≈ 5
        "why_points": [
            "5+ years of experience as a QA Automation Engineer, "
            "including 5 years at Gigaset"
        ],
    }
    _fix_experience_years(data, "en")
    assert data["why_points"][0] == (
        "5+ years of experience as a QA Automation Engineer, "
        "including 5 years at Gigaset"
    )


def test_fix_experience_years_fixes_headline_and_preserves_company_subfigure():
    from app.services.cv_generator_b2b.standalone_service import (
        _fix_experience_years,
    )

    data = {
        "experience": [
            {"company": "Gigaset", "dates": "01.2019 – 12.2023"}
        ],  # 5 years total
        "why_points": [
            "Ponad 3 lata doświadczenia jako inżynier QA, w tym 4 lata w Gigaset"
        ],
    }
    _fix_experience_years(data, "pl")
    assert data["why_points"][0] == (
        "5 lat doświadczenia jako inżynier QA, w tym 4 lata w Gigaset"
    )


def test_fix_experience_years_skips_company_bound_headline():
    # A duration tied to a company ("5 lat w Gigaset") is that company's
    # tenure, not the career total — inflating it would fabricate employment.
    from app.services.cv_generator_b2b.standalone_service import (
        _fix_experience_years,
    )

    data = {
        "experience": [
            {"company": "Gigaset", "dates": "01.2019 – 12.2023"},
            {"company": "Acme Corp", "dates": "01.2016 – 12.2018"},
        ],  # 8 years total
        "why_points": ["5 lat doświadczenia w Gigaset jako inżynier QA"],
    }
    _fix_experience_years(data, "pl")
    assert data["why_points"][0] == "5 lat doświadczenia w Gigaset jako inżynier QA"


def test_fix_experience_years_skips_company_bound_headline_en():
    from app.services.cv_generator_b2b.standalone_service import (
        _fix_experience_years,
    )

    data = {
        "experience": [
            {"company": "Gigaset", "dates": "01.2019 – 12.2023"},
            {"company": "Acme Corp", "dates": "01.2016 – 12.2018"},
        ],  # 8 years total
        "why_points": ["5 years of experience at Gigaset"],
    }
    _fix_experience_years(data, "en")
    assert data["why_points"][0] == "5 years of experience at Gigaset"


# ── Saved-CV re-render (panel list download/preview) ───────────────────────


def _sample_payload(blind: bool = False) -> dict:
    return {
        "name": "Jan Kowalski",
        "first_name": "Jan",
        "position": "Java Developer",
        "language": "pl",
        "blind_cv": blind,
        "why_points": ["8 lat doświadczenia jako Java Developer"],
        "skills": [{"label": "Backend:", "content": "Java, Spring Boot"}],
        "languages": ["Polski – ojczysty", "Angielski – biegły"],
        "experience": [
            {
                "dates": "01.2020 – obecnie",
                "company": "Acme",
                "industry": "IT",
                "position": "Java Developer",
                "responsibilities": ["Rozwój usług w Java"],
                "technologies": ["Java", "Spring Boot"],
            }
        ],
    }


def test_rerender_from_payload_produces_docx():
    # Saved-CV download/preview re-renders deterministically from the payload,
    # with no Claude call. DOCX is a zip → starts with the PK magic bytes.
    data = rerender_docx_from_payload(_sample_payload())
    assert data[:2] == b"PK"
    assert len(data) > 1000


def test_rerender_does_not_mutate_saved_payload():
    # render mutates candidate_data in place for blind anonymization; the helper
    # deep-copies so the STORED payload stays reusable for the next download.
    payload = _sample_payload(blind=True)
    rerender_docx_from_payload(payload)
    assert payload["name"] == "Jan Kowalski"
    assert payload["experience"][0]["company"] == "Acme"


def _full_page_payload() -> dict:
    # Enough experience to (over-)fill the page and spill across several pages —
    # the case where the old hybrid switched the RODO clause to an in-flow render
    # that dangled at the top of the last page in the paginated download.
    payload = _sample_payload()
    payload["experience"] = [
        {
            "dates": f"01.20{10 + i} – 12.20{11 + i}",
            "company": f"Firma {i}",
            "industry": "IT",
            "position": "Java Developer",
            "responsibilities": [
                f"Rozwój usług backendowych numer {j}" for j in range(4)
            ],
            "technologies": ["Java", "Spring Boot"],
        }
        for i in range(5)
    ]
    return payload


def test_rodo_clause_pinned_to_page_bottom_on_short_cv():
    # When the CV ends well short of the page, the RODO consent clause is pinned
    # to the foot of the last page — justified, once, with no divider above it
    # — so it reads as a footer instead of dangling mid-page below the last role.
    # It rides in a floating DrawingML text box that wraps top-and-bottom (so it
    # reserves its band and body text can never run underneath it), anchored to
    # the bottom margin, NOT a w:framePr (the frame approach spilled the clause
    # onto a blank second page on content-heavy CVs).
    import io
    import re
    import zipfile

    data = rerender_docx_from_payload(_sample_payload())
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        body = z.read("word/document.xml").decode("utf-8")

    assert "Wyrażam zgodę na przetwarzanie" in body, "RODO clause missing"
    # A single floating text box, pinned to the bottom page margin.
    assert body.count('name="RodoClause"') == 1, "RODO not in one floating box"
    # The box must RESERVE its band (top-and-bottom wrap), not float free
    # (wrapNone). wrapNone reserved no in-flow space, so a full-page body ran
    # under the pinned box and overlapped the clause — the reported bug.
    assert "wrapTopAndBottom" in body, "RODO box must reserve its band (top-and-bottom wrap)"
    assert "wrapNone" not in body, "wrapNone lets a full-page body overlap the clause"
    assert re.search(r'positionV[^>]*relativeFrom="margin"', body) and re.search(
        r"<[\w:]*align>bottom<", body
    ), "RODO not pinned to the bottom margin"
    # Justified, with no red top rule above the clause (removed on request). The
    # RODO paragraph was the only <w:pBdr> in the doc; section dividers use v:rect.
    assert 'w:val="both"' in body, "RODO clause not justified"
    assert "<w:pBdr" not in body, "RODO red top rule must be gone (removed on request)"
    assert "<w:framePr" not in body, "frame anchoring reintroduced (causes blank page)"


def test_rodo_clause_pinned_to_page_bottom_on_full_cv():
    # A content-heavy CV that spills across pages pins the RODO consent clause to
    # the foot of the LAST page via the SAME floating text box as a short CV —
    # never in the normal flow. The old hybrid switched a content-heavy CV to an
    # in-flow clause; in the paginated download that dangled at the TOP of the
    # last page whenever the body spilled just past a page boundary ("RODO na
    # górze strony"). A bottom-anchored float lands at the foot of whichever page
    # its anchor sits on, so the clause is always at the bottom of the last page.
    # The top-and-bottom wrap reserves its band, so on a body that fills the page
    # the last lines are pushed above the clause instead of overlapping it.
    import io
    import re
    import zipfile

    data = rerender_docx_from_payload(_full_page_payload())
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        body = z.read("word/document.xml").decode("utf-8")

    assert "Wyrażam zgodę na przetwarzanie" in body, "RODO clause missing"
    # Same single bottom-anchored floating box as the short-CV case: no in-flow
    # fallback, so the clause never dangles at the top of the last page. It wraps
    # top-and-bottom (reserves its band), so a full-page body is pushed above it
    # rather than overlapping it (the reported "tekst nachodzi na siebie").
    assert body.count('name="RodoClause"') == 1, "RODO not in one floating box"
    assert "wrapTopAndBottom" in body, "RODO box must reserve its band (top-and-bottom wrap)"
    assert "wrapNone" not in body, "wrapNone lets a full-page body overlap the clause"
    assert re.search(r'positionV[^>]*relativeFrom="margin"', body) and re.search(
        r"<[\w:]*align>bottom<", body
    ), "RODO not pinned to the bottom margin"
    assert 'w:val="both"' in body, "RODO clause not justified"
    assert "<w:framePr" not in body, "frame anchoring reintroduced (causes blank page)"


# ── PDF text extraction — OCR fallback for scanned CVs ─────────────────────
#
# A scanned / image-only PDF has no text layer, so pdftotext and pdfplumber
# both return nothing. Before the fallback the generator hard-failed with
# "Empty text extracted from <file>.pdf" even though the rest of the app
# (app.services.cv_text_extractor) already OCRs such files via tesseract.


def _patch_native_extractors(monkeypatch, pdftotext=None, pdfplumber=""):
    from app.services.cv_generator_b2b import text_extractor as te

    monkeypatch.setattr(te, "_extract_pdf_pdftotext", lambda data: pdftotext)
    if isinstance(pdfplumber, Exception):

        def _raise(data):
            raise pdfplumber

        monkeypatch.setattr(te, "_extract_pdf_pdfplumber", _raise)
    else:
        monkeypatch.setattr(te, "_extract_pdf_pdfplumber", lambda data: pdfplumber)
    return te


def test_scanned_pdf_falls_back_to_ocr(monkeypatch):
    te = _patch_native_extractors(monkeypatch, pdftotext=None, pdfplumber="")
    ocr_text = "Adam Drazkowski\nCheck Point CCSE / CCSA\n" + "x" * 200
    monkeypatch.setattr(te, "_extract_pdf_ocr", lambda data: ocr_text)

    assert te.extract_text_from_file(b"%PDF-fake", "skan.pdf") == ocr_text


def test_near_empty_native_text_prefers_ocr(monkeypatch):
    # Native extraction returning just page furniture (< threshold) must not
    # win over a real OCR read.
    te = _patch_native_extractors(monkeypatch, pdftotext="1/2", pdfplumber="")
    ocr_text = "Doświadczenie zawodowe\n" + "y" * 300
    monkeypatch.setattr(te, "_extract_pdf_ocr", lambda data: ocr_text)

    assert te.extract_text_from_file(b"%PDF-fake", "skan.pdf") == ocr_text


def test_good_native_text_skips_ocr(monkeypatch):
    native = "Solidne CV z pełną warstwą tekstową. " * 10
    te = _patch_native_extractors(monkeypatch, pdftotext=native)

    def _boom(data):
        raise AssertionError("OCR must not run when native text is good")

    monkeypatch.setattr(te, "_extract_pdf_ocr", _boom)

    assert te.extract_text_from_file(b"%PDF-fake", "cv.pdf") == native


def test_unreadable_pdf_raises_with_scan_hint(monkeypatch):
    from app.services.cv_generator_b2b.text_extractor import CVTextExtractionError

    te = _patch_native_extractors(monkeypatch, pdftotext=None, pdfplumber="")
    monkeypatch.setattr(te, "_extract_pdf_ocr", lambda data: None)

    with pytest.raises(CVTextExtractionError) as exc:
        te.extract_text_from_file(b"%PDF-fake", "skan.pdf")
    assert "Empty text extracted from skan.pdf" in str(exc.value)
    assert "skan" in str(exc.value)  # actionable hint for the recruiter


def test_corrupt_pdf_still_tries_ocr(monkeypatch):
    # pdfplumber blowing up on a malformed PDF must not crash the job —
    # poppler (pdf2image) can often still rasterize it for OCR.
    te = _patch_native_extractors(
        monkeypatch, pdftotext=None, pdfplumber=ValueError("No /Root object")
    )
    ocr_text = "Treść odzyskana przez OCR. " + "z" * 200
    monkeypatch.setattr(te, "_extract_pdf_ocr", lambda data: ocr_text)

    assert te.extract_text_from_file(b"%PDF-broken", "cv.pdf") == ocr_text


# ── Faza 3/5/6: taxonomy-backed bolding + legacy JSONB + manual parser ──────
#
# End-to-end tests that render a REAL DOCX and inspect the actual <w:b> runs
# (the contract the suite previously never asserted), gated on an injected
# technology taxonomy that mirrors the production skills/skill_aliases table.


@pytest.fixture
def tech_taxonomy():
    """Inject a small prod-like technology taxonomy for the duration of a test,
    then reset it so the taxonomy-free tests (and other modules) are unaffected."""
    tech = {
        "react", "kubernetes", "postgresql", "python", "docker", "jest",
        "selenium", "c", "r", "azure", "java", "spring boot", "sql", "git",
    }
    alias_to_canonical = {t: t for t in tech}
    alias_to_canonical.update(
        {
            "reactjs": "react", "react.js": "react",
            "k8s": "kubernetes", "postgres": "postgresql",
            "microsoft azure": "azure",
        }
    )
    canonical_to_aliases = {
        "react": ["reactjs", "react.js"],
        "kubernetes": ["k8s"],
        "postgresql": ["postgres"],
        "azure": ["microsoft azure"],
    }
    set_tech_taxonomy(
        tech_canonicals=tech,
        alias_to_canonical=alias_to_canonical,
        canonical_to_aliases=canonical_to_aliases,
    )
    yield
    set_tech_taxonomy(tech_canonicals=set(), alias_to_canonical={}, canonical_to_aliases={})


def _render_bold_texts(payload: dict) -> set[str]:
    """Render the payload to a real DOCX and return the set of bold run texts."""
    data = rerender_docx_from_payload(payload)
    doc = _docx_lib.Document(io.BytesIO(data))
    out: set[str] = set()

    def walk(paras):
        for p in paras:
            for run in p.runs:
                if run.font.bold and run.text.strip():
                    out.add(run.text.strip())

    walk(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                walk(cell.paragraphs)
    return out


def _bold_payload(keywords, *, why=None, resp=None) -> dict:
    p = _sample_payload()
    p["highlight_keywords"] = keywords
    p["skills"] = [{"label": "Technologie:", "content": ", ".join(keywords)}]
    p["why_points"] = why or ["Doświadczenie w " + ", ".join(keywords) + "."]
    p["experience"][0]["responsibilities"] = resp or [
        "Praca z " + ", ".join(keywords) + "."
    ]
    p["experience"][0]["technologies"] = keywords
    return p


def test_e2e_false_positives_never_bold(tech_taxonomy):
    bolds = _render_bold_texts(
        _bold_payload(
            ["Agile", "Scrum", "Leadership", "English", "UX"],
            why=["Pracował w Agile i Scrum; Leadership; English C1; obszar UX."],
        )
    )
    for token in ["Agile", "Scrum", "Leadership", "English", "UX"]:
        assert token not in bolds, f"{token!r} must not bold (not a technology)"


def test_e2e_taxonomy_alias_forms_bold(tech_taxonomy):
    # Chips written as ReactJS/K8s/Postgres must bold the canonical forms in prose.
    bolds = _render_bold_texts(
        _bold_payload(
            ["ReactJS", "K8s", "Postgres"],
            why=["Frontend w React, klaster Kubernetes, baza PostgreSQL."],
        )
    )
    for token in ["React", "Kubernetes", "PostgreSQL"]:
        assert token in bolds, f"{token!r} should bold via alias resolution"


def test_e2e_single_letter_and_stopword_tech_bold(tech_taxonomy):
    bolds = _render_bold_texts(
        _bold_payload(
            ["C", "R", "Jest"],
            why=["Testy w Jest; język C oraz R. Oto zdanie w którym jest czasownik."],
        )
    )
    assert "Jest" in bolds and "C" in bolds and "R" in bolds
    assert "jest" not in bolds  # the Polish verb (lowercase) must never bold


def test_e2e_polish_inflection_bolds(tech_taxonomy):
    bolds = _render_bold_texts(
        _bold_payload(
            ["Python", "Docker"],
            why=["Programował w Pythonie i konteneryzował w Dockerze."],
            resp=["Rozwój w Pythona, wdrożenia Dockera."],
        )
    )
    assert "Pythonie" in bolds  # inflected form bolds (was a false negative)
    assert "Dockera" in bolds


def test_e2e_multiword_product_bolds_whole(tech_taxonomy):
    bolds = _render_bold_texts(
        _bold_payload(
            ["Selenium WebDriver", "Microsoft Azure"],
            why=["Automatyzacja Selenium WebDriver na Microsoft Azure."],
        )
    )
    assert "Selenium WebDriver" in bolds  # not only "WebDriver"
    assert "Microsoft Azure" in bolds  # not nothing / not only "AD"


def test_from_nexus_job_handles_legacy_jsonb_shapes():
    # dict {"technologies": [...]} → the techs, NOT the literal key "technologies"
    assert from_nexus_job({"technologies": ["Java", "AWS"]}, None, None).must_have == [
        "Java",
        "AWS",
    ]
    # JSON-encoded string → decoded, NOT iterated character-by-character
    assert from_nexus_job('["Python", "Go"]', None, None).must_have == ["Python", "Go"]
    # list[dict] and list[str] mixed
    assert from_nexus_job([{"name": "React"}, "Docker"], None, None).must_have == [
        "React",
        "Docker",
    ]
    # comma-string
    assert from_nexus_job("Java, Spring", None, None).must_have == ["Java", "Spring"]


def test_split_skills_respects_parentheses():
    assert _split_skills("Java (Spring, Hibernate), Python") == [
        "Java (Spring, Hibernate)",
        "Python",
    ]
    assert _split_skills("WCAG 2.1/2.2 (AA, AAA)") == ["WCAG 2.1/2.2 (AA, AAA)"]
    assert _split_skills("Figma (badania, feedback)") == ["Figma (badania, feedback)"]
    # top-level comma / semicolon / newline still split
    assert _split_skills("Java, Python; Go\nRust") == ["Java", "Python", "Go", "Rust"]


def test_parse_champion_from_docx_keeps_parenthesized_skills():
    d = _docx_lib.Document()
    d.add_paragraph("2. PROFIL KANDYDATA")
    d.add_paragraph("MUST-HAVE:")
    d.add_paragraph("Java (Spring, Hibernate), Python")
    d.add_paragraph("NICE-TO-HAVE:")
    d.add_paragraph("Docker")
    d.add_paragraph("3. KONTEKST")
    buf = io.BytesIO()
    d.save(buf)
    cp = parse_champion_from_docx_bytes(buf.getvalue(), "champion.docx")
    assert "Java (Spring, Hibernate)" in cp.must_have
    assert "Python" in cp.must_have
    assert "Docker" in cp.nice_to_have
