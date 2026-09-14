import hashlib
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock
from zipfile import ZipFile

import pytest
from docx import Document
from fastapi import HTTPException
from app.services import object_storage
from app.services.cv_approval_provenance import capture_editor_origin
from app.services.cv_document_assets import default_template
from app.services.cv_standalone_approval import approve_unchanged_generation


@pytest.mark.parametrize(
    "case", ["verified", "blind", "unverified", "corrupt", "retry"]
)
async def test_standalone_approval_freezes_verified_artifact(case, pipeline_mode):
    """Zamrożenie zapisanego DOCX przy zatwierdzeniu — w OBU przepływach
    (`pipeline_mode`). Jedyna różnica trybów to wiersz BEZ kontroli treści:
    ścisłe dowody (v10) odmawiają 409, doradcze (legacy, domyślne na
    produkcji) zatwierdzają — pre-#1444 nie było tu żadnej bramki."""
    payload = {
        "name": "Jan Secret",
        "position": "Engineer",
        "language": "pl",
        "blind_cv": case == "blind",
    }
    digest = capture_editor_origin("", payload)["generated_factual_payload_sha256"]
    if case != "unverified":
        payload["factual_verification"] = {
            "status": "verified",
            "document_sha256": digest,
        }
    doc = SimpleNamespace(
        id=7,
        client_id=None,
        client_rule_version=None,
        status="ready",
        render_payload=payload,
        docx_content=b"exact docx",
        docx_sha256=hashlib.sha256(b"exact docx").hexdigest(),
        filename="cv.docx",
    )
    if case == "corrupt":
        doc.docx_content += b"changed"
    existing = object() if case == "retry" else None
    db = SimpleNamespace(
        scalar=AsyncMock(return_value=existing), add=Mock(), flush=AsyncMock()
    )
    evidence_enforced = pipeline_mode == "v10"
    if case == "corrupt" or (case == "unverified" and evidence_enforced):
        with pytest.raises(HTTPException) as exc:
            await approve_unchanged_generation(db, doc, 9)
        assert exc.value.status_code == 409
        db.add.assert_not_called()
        return
    version = await approve_unchanged_generation(db, doc, 9)
    if case == "unverified":
        # Dowody doradcze: zatwierdzenie bez kontroli AI, bez fałszywego
        # „verified" w metadanych.
        assert version.render_metadata["generation_review_available"] is False
        db.add.assert_called_once_with(version)
    if case == "retry":
        assert version is existing
        db.add.assert_not_called()
        return
    assert version.generated_owner_id == 7
    assert version.candidate_stage_cv_id is None
    assert version.docx_content == b"exact docx"
    assert (
        version.content_sha256
        == hashlib.sha256(version.content_html.encode()).hexdigest()
    )
    if case == "blind":
        assert "Secret" not in version.content_html
        assert "Jan" not in version.candidate_first_name


# ── CV sprzed zapisu DOCX (#1444, migracja 0292 bez backfillu) ─────────────


def _docx_text(content: bytes) -> str:
    document = Document(BytesIO(content))
    parts = [paragraph.text for paragraph in document.paragraphs]
    parts += [
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    ]
    return "\n".join(parts)


def _legacy_payload(**overrides):
    payload = {
        "name": "Jan Secret",
        "first_name": "Jan",
        "position": "Backend Engineer",
        "language": "pl",
        "blind_cv": False,
        "why_points": ["Pięć lat pracy z Pythonem"],
        "education": [],
        "skills": [{"label": "Backend", "content": "Python, PostgreSQL"}],
        "certifications": [],
        "languages": ["Polski – ojczysty"],
        "experience": [
            {
                "dates": "01.2020 – 12.2024",
                "company": "Acme Sp. z o.o.",
                "industry": "bankowa",
                "position": "Backend Engineer",
                "responsibilities": ["Utrzymanie API"],
                "technologies": ["Python"],
            }
        ],
    }
    payload.update(overrides)
    return payload


def _legacy_doc(**overrides):
    """A row generated before #1444 saved the file: both DOCX columns NULL."""
    fields = dict(
        id=7,
        client_id=None,
        client_rule_version=None,
        status="ready",
        blind=False,
        render_payload=_legacy_payload(),
        docx_content=None,
        docx_sha256=None,
        filename="cv.docx",
    )
    fields.update(overrides)
    return SimpleNamespace(**fields)


def _db():
    return SimpleNamespace(
        scalar=AsyncMock(return_value=None), add=Mock(), flush=AsyncMock()
    )


@pytest.fixture
def evidence_advisory(monkeypatch):
    """Production default — these rows never had an AI content review."""
    monkeypatch.delenv("CV_SOURCE_EVIDENCE_ENFORCED", raising=False)


async def test_legacy_generation_is_rendered_frozen_and_approved(evidence_advisory):
    doc = _legacy_doc()
    db = _db()
    version = await approve_unchanged_generation(db, doc, 9)
    assert doc.docx_content[:2] == b"PK"
    assert doc.docx_sha256 == hashlib.sha256(doc.docx_content).hexdigest()
    assert version.docx_content == doc.docx_content
    assert version.docx_sha256 == doc.docx_sha256
    assert version.template_content == default_template()
    assert version.render_metadata["docx_rendered_at_approval"] is True
    assert "Jan Secret" in _docx_text(doc.docx_content)
    db.add.assert_called_once_with(version)


async def test_legacy_blind_generation_renders_the_anonymized_docx(
    evidence_advisory,
):
    doc = _legacy_doc(blind=True, render_payload=_legacy_payload(blind_cv=True))
    version = await approve_unchanged_generation(_db(), doc, 9)
    text = _docx_text(doc.docx_content)
    for identity in ("Jan", "Secret", "Acme"):
        assert identity not in text
        assert identity not in version.content_html
    assert "Firma z branży bankowa" in text
    assert version.template == "blind"
    # The identity guard stays frozen for later edits of this approval.
    assert version.render_metadata["blind_identity_guard"] is True


async def test_blind_row_without_payload_anonymization_is_refused(
    evidence_advisory,
):
    doc = _legacy_doc(blind=True)  # the payload says blind_cv=False
    db = _db()
    with pytest.raises(HTTPException) as exc:
        await approve_unchanged_generation(db, doc, 9)
    assert exc.value.status_code == 409
    assert (doc.docx_content, doc.docx_sha256) == (None, None)
    db.add.assert_not_called()


async def test_legacy_consent_is_rendered_from_the_frozen_bytes(
    evidence_advisory, monkeypatch
):
    from PIL import Image

    image = BytesIO()
    Image.new("RGB", (8, 8), "white").save(image, "PNG")
    png = image.getvalue()
    download = Mock(return_value=png)
    monkeypatch.setattr(object_storage, "download_cv", download)
    doc = _legacy_doc(
        render_payload=_legacy_payload(
            consent_screenshot={"storage_key": "cv/consent.png", "filename": "z.png"}
        )
    )
    version = await approve_unchanged_generation(_db(), doc, 9)
    # One read from storage: the version and the DOCX carry the same image.
    download.assert_called_once_with("cv/consent.png")
    assert version.consent_content == png
    with ZipFile(BytesIO(doc.docx_content)) as archive:
        media = [
            archive.read(name)
            for name in archive.namelist()
            if name.startswith("word/media/")
        ]
    assert png in media


async def test_legacy_consent_that_cannot_be_read_blocks_approval(
    evidence_advisory, monkeypatch
):
    monkeypatch.setattr(
        object_storage, "download_cv", Mock(side_effect=OSError("storage down"))
    )
    doc = _legacy_doc(
        render_payload=_legacy_payload(
            consent_screenshot={"storage_key": "cv/consent.png"}
        )
    )
    db = _db()
    with pytest.raises(HTTPException) as exc:
        await approve_unchanged_generation(db, doc, 9)
    assert exc.value.status_code == 422
    assert (doc.docx_content, doc.docx_sha256) == (None, None)
    db.add.assert_not_called()


@pytest.mark.parametrize(
    "content,digest", [(None, "a" * 64), (b"docx", None), (b"", None)]
)
async def test_partial_docx_record_is_an_integrity_failure_not_legacy(
    evidence_advisory, content, digest
):
    doc = _legacy_doc(docx_content=content, docx_sha256=digest)
    with pytest.raises(HTTPException) as exc:
        await approve_unchanged_generation(_db(), doc, 9)
    assert exc.value.status_code == 409
    assert exc.value.detail == (
        "Brak poprawnego zapisanego dokumentu. Wygeneruj CV ponownie."
    )
    assert (doc.docx_content, doc.docx_sha256) == (content, digest)


async def test_legacy_generation_still_needs_review_when_evidence_is_enforced():
    # conftest pins CV_SOURCE_EVIDENCE_ENFORCED=true for this suite.
    doc = _legacy_doc()
    with pytest.raises(HTTPException) as exc:
        await approve_unchanged_generation(_db(), doc, 9)
    assert exc.value.status_code == 409
    assert "kontroli" in exc.value.detail
    assert doc.docx_content is None


async def test_unrenderable_legacy_payload_is_refused_without_a_write(
    evidence_advisory,
):
    payload = _legacy_payload()
    del payload["position"]  # the renderer indexes it directly
    doc = _legacy_doc(render_payload=payload)
    db = _db()
    with pytest.raises(HTTPException) as exc:
        await approve_unchanged_generation(db, doc, 9)
    assert exc.value.status_code == 409
    assert "odtworzyć" in exc.value.detail
    assert (doc.docx_content, doc.docx_sha256) == (None, None)
    db.add.assert_not_called()


# ── M05-B03: szkic założony samym otwarciem edytora nie blokuje zatwierdzenia ──


def _draft_row(**overrides):
    fields = dict(id=31, edit_revision=0, branded_status="draft")
    fields.update(overrides)
    return SimpleNamespace(**fields)


def _db_with_draft(draft, job_id=None):
    return SimpleNamespace(
        # kolejno: istniejąca wersja 1, szkic, kontrola w toku dla szkicu
        scalar=AsyncMock(side_effect=[None, draft, job_id]),
        add=Mock(),
        flush=AsyncMock(),
        delete=AsyncMock(),
    )


async def test_untouched_editor_draft_is_dropped_and_generation_approved(
    evidence_advisory,
):
    draft = _draft_row()
    db = _db_with_draft(draft)
    version = await approve_unchanged_generation(db, _legacy_doc(), 9)
    db.delete.assert_awaited_once_with(draft)
    db.add.assert_called_once_with(version)
    assert version.version == 1


@pytest.mark.parametrize(
    "draft,job_id",
    [
        (_draft_row(edit_revision=1), None),  # zapisane zmiany w edytorze
        (_draft_row(branded_status="finalized", edit_revision=0), None),
        (_draft_row(), 77),  # kontrola treści z edytora w toku
    ],
)
async def test_edited_or_reviewed_draft_still_blocks_approval(
    evidence_advisory, draft, job_id
):
    db = _db_with_draft(draft, job_id)
    doc = _legacy_doc()
    with pytest.raises(HTTPException) as exc:
        await approve_unchanged_generation(db, doc, 9)
    assert exc.value.status_code == 409
    assert "szkic" in exc.value.detail
    db.delete.assert_not_awaited()
    db.add.assert_not_called()
    assert doc.docx_content is None
