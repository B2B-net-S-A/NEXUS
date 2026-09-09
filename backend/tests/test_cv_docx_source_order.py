"""Real DOCX layout must keep each role's dates and duties together."""

from io import BytesIO

from docx import Document
import pytest

from app.services.cv_generator_b2b.text_extractor import (
    CVTextExtractionError,
    extract_text_from_file,
)


def test_interleaved_tables_nested_and_merged_cells_keep_source_order():
    doc = Document()
    doc.add_paragraph("Doświadczenie zawodowe")
    first = doc.add_table(rows=1, cols=2)
    first.cell(0, 0).text = "2010–2017 | Firma A | Magazynier"
    first.cell(0, 0).merge(first.cell(0, 1))
    doc.add_paragraph("Kompletowanie zamówień magazynowych")
    second = doc.add_table(rows=1, cols=1)
    cell = second.cell(0, 0)
    cell.text = "2018–2020 | Firma B | Analityk"
    nested = cell.add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Analiza sprzedaży SQL"
    cell.add_paragraph("Raportowanie wyników")
    doc.add_paragraph("Wykształcenie")
    data = BytesIO()
    doc.save(data)
    assert extract_text_from_file(data.getvalue(), "cv.docx").splitlines() == [
        "Doświadczenie zawodowe",
        "2010–2017 | Firma A | Magazynier",
        "Kompletowanie zamówień magazynowych",
        "2018–2020 | Firma B | Analityk",
        "Analiza sprzedaży SQL",
        "Raportowanie wyników",
        "Wykształcenie",
    ]


def test_corrupt_docx_raises_the_preflight_error_contract():
    with pytest.raises(CVTextExtractionError, match="DOCX"):
        extract_text_from_file(b"not a zip archive", "cv.docx")
