"""Nagłówek stanowiska w Wordzie nie zostaje sam na dole strony.

Zgłoszenie 22.10.2026: w CV na dole 1. strony stała sama data stanowiska
(„02.2019 - 03.2025”), a firma, stanowisko i zadania zaczynały się na 2.
Od #503 akapity CV nie są sklejane, żeby Enter w Wordzie przesuwał tekst
o jeden wiersz. Wyjątek: daty, firma, stanowisko i „Zakres zadań:” trzymają
się razem i z pierwszym zadaniem. Same zadania dalej płyną wiersz po wierszu.
"""

from io import BytesIO

from docx import Document

from app.services.cv_approved_docx import render_approved_docx
from app.services.cv_document_assets import default_template
from app.services.cv_generator_b2b.docx_renderer import render_cv_to_bytes
from app.services.html_sanitizer import sanitize_cv_html

HEADER = [
    "02.2019 - 03.2025",
    "Nazwa firmy: Acme · Bankowość",
    "Stanowisko: Network Engineer",
    "Zakres zadań:",
]
DUTIES = ["Utrzymanie zapór Check Point", "Administracja Cisco AnyConnect"]

EDITOR_HTML = (
    "<h1>Network Engineer – Jan Testowy</h1><hr>"
    '<h2 data-cv-section="experience">DOŚWIADCZENIE</h2>'
    '<p data-cv-section="role"><strong>Network Engineer</strong> 02.2019 - 03.2025</p>'
    '<p data-cv-section="employer">Acme · Bankowość</p>'
    '<p data-cv-section="duties_label">Zakres zadań:</p><ul>'
    + "".join(f"<li><p>{duty}</p></li>" for duty in DUTIES)
    + "</ul>"
)


def _payload() -> dict:
    return {
        "name": "Jan Testowy",
        "first_name": "Jan",
        "position": "Network Engineer",
        "language": "pl",
        "why_points": ["Sieci w bankowości"],
        "experience": [
            {
                "dates": "02.2019 - 03.2025",
                "company": "Acme · Bankowość",
                "industry": "",
                "position": "Network Engineer",
                "responsibilities": DUTIES,
                "technologies": [],
            }
        ],
    }


def _cohesion(doc: Document) -> list[tuple[str, bool]]:
    paragraphs = doc.paragraphs
    texts = [p.text.strip() for p in paragraphs]
    start = texts.index(HEADER[0])
    block = paragraphs[start : start + len(HEADER) + len(DUTIES)]
    # Generator dopisuje interpunkcję na końcu zadań („,” / „.”).
    return [
        (p.text.strip().rstrip(",.;"), bool(p.paragraph_format.keep_with_next))
        for p in block
    ]


EXPECTED = [(line, True) for line in HEADER] + [(duty, False) for duty in DUTIES]


def test_generator_docx_keeps_role_header_with_first_task():
    docx = render_cv_to_bytes(_payload(), BytesIO(default_template()))
    assert _cohesion(Document(BytesIO(docx))) == EXPECTED


def test_editor_docx_keeps_role_header_with_first_task():
    docx = render_approved_docx(sanitize_cv_html(EDITOR_HTML), default_template())
    assert _cohesion(Document(BytesIO(docx))) == EXPECTED
