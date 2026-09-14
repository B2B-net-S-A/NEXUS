from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock
import pytest
from fastapi import HTTPException
from app.services import cv_generated_editor as editor


def draft():
    return SimpleNamespace(
        id=5,
        generated_document_id=7,
        edit_revision=2,
        branded_version=1,
        branded_status="draft",
        branded_draft_html="<p>Original</p>",
        branded_template_content=b"template",
        branded_consent_content=None,
        branded_render_metadata={},
        branded_language="pl",
        branded_template="standard",
        branded_docx_filename="cv.docx",
    )


@pytest.mark.parametrize("revision", [1, 2])
async def test_save_checks_revision_and_preserves_approved_content(revision):
    item = draft()
    if revision != 2:
        with pytest.raises(HTTPException) as exc:
            editor.save(item, revision, "<p>Edited</p>")
        assert exc.value.status_code == 409
        assert item.branded_draft_html == "<p>Original</p>"
    else:
        editor.save(item, 2, "<p>Edited</p><script>alert(1)</script>")
        assert item.edit_revision == 3
        assert "<script>" not in item.branded_draft_html
        item.branded_status = "finalized"
        with pytest.raises(HTTPException):
            editor.save(item, 3, "<p>Other</p>")
        editor.new_draft(item, 3)
        assert item.branded_version == 2
        assert item.edit_revision == 4
        assert item.branded_status == "draft"


@pytest.mark.parametrize("verified", [True, False])
async def test_finalize_stores_exact_submitted_content_only_after_review(
    monkeypatch, verified
):
    item = draft()
    db = SimpleNamespace(add=Mock(), flush=AsyncMock())
    monkeypatch.setattr(
        editor, "render", AsyncMock(return_value=b"rendered edited docx")
    )
    review = AsyncMock(return_value={"status": "verified"})
    if not verified:
        review.side_effect = HTTPException(422, "Unsupported claim")
    monkeypatch.setattr(editor, "review_for_approval", review)
    if verified:
        version = await editor.finalize(db, item, 2, "<p>Submitted</p>", 9)
        assert version.content_html == "<p>Submitted</p>"
        assert version.docx_content == b"rendered edited docx"
        assert version.generated_owner_id == 7
        assert version.render_metadata["content_review"]["status"] == "verified"
        assert item.branded_status == "finalized"
        assert item.edit_revision == 3
    else:
        with pytest.raises(HTTPException):
            await editor.finalize(db, item, 2, "<p>Submitted</p>", 9)
        db.add.assert_not_called()
        assert item.branded_status == "draft"
        assert item.edit_revision == 2
        assert item.branded_draft_html == "<p>Original</p>"
    review.assert_awaited_once_with(db, item, "<p>Submitted</p>", 9)


async def test_real_editor_docx_keeps_submitted_text_and_bold(monkeypatch):
    from io import BytesIO
    from docx import Document
    from app.services.cv_document_assets import default_template

    item = draft()
    item.branded_template_content = default_template()
    db = SimpleNamespace(add=Mock(), flush=AsyncMock())
    # Exercise the actual DOCX renderer; this test does not claim model acceptance.
    monkeypatch.setattr(
        editor, "review_for_approval", AsyncMock(return_value={"status": "verified"})
    )
    version = await editor.finalize(
        db,
        item,
        2,
        "<p>Nie pracował z <strong>Kubernetes</strong>. Utrzymywał <strong>Python</strong> API.</p>",
        9,
    )
    document = Document(BytesIO(version.docx_content))
    paragraph = next(p for p in document.paragraphs if "Kubernetes" in p.text)
    assert paragraph.text == "Nie pracował z Kubernetes. Utrzymywał Python API."
    bold_text = "".join(run.text for run in paragraph.runs if run.bold)
    assert bold_text == "KubernetesPython"
    assert "Original" not in "\n".join(p.text for p in document.paragraphs)
    assert version.content_html == item.branded_draft_html
    assert version.render_metadata["content_review"]["status"] == "verified"


async def test_missing_assets_returns_recovery_code_without_creating_draft(monkeypatch):
    db = SimpleNamespace(
        scalar=AsyncMock(side_effect=[None, None]), add=Mock(), flush=AsyncMock()
    )
    generated = SimpleNamespace(
        id=7, status="ready", render_payload={"language": "pl", "name": "Synthetic"}
    )
    monkeypatch.setattr(
        editor,
        "generated_assets",
        Mock(side_effect=editor.CvAssetsError("Missing original template")),
    )
    with pytest.raises(HTTPException) as exc:
        await editor.load_draft(db, generated)
    assert exc.value.status_code == 422
    assert exc.value.detail["code"] == "cv_editor_assets_unavailable"
    db.add.assert_not_called()
    db.flush.assert_not_awaited()


@pytest.mark.parametrize("persist", [False, True])
async def test_reading_the_editor_does_not_create_a_draft_row(monkeypatch, persist):
    """M05-B03: GET edytora nie zapisuje szkicu; mutacja zakłada ten sam stan."""
    db = SimpleNamespace(
        scalar=AsyncMock(side_effect=[None, None]), add=Mock(), flush=AsyncMock()
    )
    generated = SimpleNamespace(
        id=7,
        status="ready",
        filename="cv.docx",
        render_payload={"language": "pl", "name": "Synthetic", "position": "Dev"},
    )
    monkeypatch.setattr(
        editor, "generated_assets", Mock(return_value=(b"template", None, {}))
    )
    monkeypatch.setattr(editor, "capture_editor_origin", Mock(return_value={}))
    item = await editor.load_draft(db, generated, persist=persist)
    assert item.edit_revision == 0
    assert item.branded_status == "draft"
    assert item.branded_version == 1
    if persist:
        db.add.assert_called_once_with(item)
        db.flush.assert_awaited_once()
    else:
        db.add.assert_not_called()
        db.flush.assert_not_awaited()


def test_editor_read_endpoints_do_not_persist_a_draft():
    """Kontrakt tras: odczyty edytora wołają ``persist=False``."""
    import ast
    import inspect

    from app.api import cv_generator_b2b as api

    for fn in (
        api.get_generated_editor,
        api.preview_generated_editor,
        api.print_generated_editor,
    ):
        tree = ast.parse(inspect.getsource(fn).lstrip())
        calls = [
            node
            for node in ast.walk(tree)
            if isinstance(node, ast.Call)
            and getattr(node.func, "id", None) == "_load_generated_editor"
        ]
        assert calls, fn.__name__
        assert all(
            any(
                kw.arg == "persist"
                and isinstance(kw.value, ast.Constant)
                and kw.value.value is False
                for kw in call.keywords
            )
            for call in calls
        ), fn.__name__
