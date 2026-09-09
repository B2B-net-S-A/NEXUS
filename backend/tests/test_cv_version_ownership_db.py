"""Hosted PostgreSQL coverage for standalone/pipeline approval ownership."""

import hashlib
from datetime import datetime, timezone
from uuid import uuid4

import pytest
from sqlalchemy import delete, select
from sqlalchemy.exc import IntegrityError

from app.core.database import AsyncSessionLocal
from app.models.candidate import Candidate
from app.models.client import Client
from app.models.job import Job
from app.models.recruitment_pipeline import CandidateStage
from app.models.candidate_stage_cv import CandidateStageCV
from app.models.cv_document_version import CvDocumentVersion
from app.models.cv_generated_document import CvGeneratedDocument
from app.models.cv_generated_share import CvGeneratedShareToken


@pytest.mark.parametrize("standalone", [True, False])
async def test_generation_deletion_respects_approval_owner(standalone):
    async with AsyncSessionLocal() as db:
        # Keep the entire synthetic fixture in a rolled-back transaction.
        try:
            candidate = Candidate(name="Approval", lastname="Fixture")
            client = Client(name=f"approval-{uuid4().hex}")
            db.add_all([candidate, client])
            await db.flush()
            job = Job(title="Approval fixture", client_id=client.id)
            db.add(job)
            await db.flush()
            stage = CandidateStage(candidate_id=candidate.id, job_id=job.id)
            db.add(stage)
            await db.flush()
            stage_cv = CandidateStageCV(
                candidate_stage_id=stage.id, candidate_id=candidate.id, job_id=job.id
            )
            generated = CvGeneratedDocument(
                candidate_name="Fixture",
                filename="fixture.docx",
                mode="upload",
                status="ready",
            )
            db.add_all([stage_cv, generated])
            await db.flush()
            version = CvDocumentVersion(
                candidate_stage_cv_id=None if standalone else stage_cv.id,
                generated_owner_id=generated.id if standalone else None,
                generated_document_id=generated.id,
                version=1,
                content_html="<p>Fixture</p>",
                content_sha256=hashlib.sha256(b"<p>Fixture</p>").hexdigest(),
                approved_at=datetime.now(timezone.utc),
            )
            db.add(version)
            await db.flush()
            version_id = version.id
            token_key = uuid4().hex
            db.add(
                CvGeneratedShareToken(
                    token=token_key,
                    token_sha256=uuid4().hex * 2,
                    generated_document_id=generated.id,
                    document_version_id=version_id,
                )
            )
            await db.flush()
            await db.execute(
                delete(CvGeneratedDocument).where(
                    CvGeneratedDocument.id == generated.id
                )
            )
            # Query columns rather than identity-map objects after DB-side cascades.
            retained = (
                await db.execute(
                    select(
                        CvDocumentVersion.id, CvDocumentVersion.generated_document_id
                    ).where(CvDocumentVersion.id == version_id)
                )
            ).first()
            if standalone:
                assert retained is None
            else:
                assert retained.id == version_id
                assert retained.generated_document_id is None
            assert (
                await db.scalar(
                    select(CvGeneratedShareToken.token).where(
                        CvGeneratedShareToken.token == token_key
                    )
                )
                is None
            )
        finally:
            await db.rollback()


async def test_approval_without_an_owner_is_rejected():
    async with AsyncSessionLocal() as db:
        try:
            db.add(
                CvDocumentVersion(
                    version=1,
                    content_html="Fixture",
                    content_sha256="a" * 64,
                    approved_at=datetime.now(timezone.utc),
                )
            )
            with pytest.raises(IntegrityError):
                await db.flush()
        finally:
            await db.rollback()


async def test_standalone_approval_share_roundtrip(
    app_client, app_auth_headers, monkeypatch
):
    from app.services.cv_approval_provenance import capture_editor_origin

    payload = {
        "name": "Private Identity",
        "position": "Engineer",
        "language": "pl",
        "blind_cv": True,
    }
    digest = capture_editor_origin("", payload)["generated_factual_payload_sha256"]
    payload["factual_verification"] = {"status": "verified", "document_sha256": digest}
    async with AsyncSessionLocal() as db:
        generated = CvGeneratedDocument(
            candidate_name="Private Identity",
            filename="fixture.docx",
            mode="upload",
            status="ready",
            render_payload=payload,
            docx_content=b"frozen fixture",
            docx_sha256=hashlib.sha256(b"frozen fixture").hexdigest(),
        )
        db.add(generated)
        await db.commit()
        generated_id = generated.id
    try:
        url = f"/api/cv-generator/generated/{generated_id}"
        unapproved = await app_client.post(
            f"{url}/share-token", headers=app_auth_headers
        )
        assert unapproved.status_code == 409, unapproved.text
        first = await app_client.post(f"{url}/approve", headers=app_auth_headers)
        assert first.status_code == 200, first.text
        version_id = first.json()["document_version_id"]
        again = await app_client.post(f"{url}/approve", headers=app_auth_headers)
        assert again.status_code == 200, again.text
        assert again.json()["document_version_id"] == version_id
        versions = await app_client.get(
            f"{url}/approved-versions", headers=app_auth_headers
        )
        assert versions.status_code == 200, versions.text
        assert [row["id"] for row in versions.json()] == [version_id]
        created = await app_client.post(
            f"{url}/share-token",
            params={"document_version_id": version_id},
            headers=app_auth_headers,
        )
        assert created.status_code == 201, created.text
        token = created.json()["token"]
        revoke_key = created.json()["revoke_key"]
        assert token != revoke_key
        assert revoke_key.startswith("v2$")
        async with AsyncSessionLocal() as db:
            stored_token = await db.get(CvGeneratedShareToken, revoke_key)
            assert stored_token.document_version_id == version_id
            assert (
                stored_token.token_sha256 == hashlib.sha256(token.encode()).hexdigest()
            )
        assert (
            await app_client.get(f"/api/public/cv-i/{revoke_key}")
        ).status_code == 404
        public = await app_client.get(f"/api/public/cv-i/{token}")
        assert public.status_code == 200, public.text
        assert public.json()["document_version_id"] == version_id
        assert "Private Identity" not in public.text
        assert public.json()["cv_html"]
        assert public.json()["chat_enabled"] is False
        async with AsyncSessionLocal() as db:
            version = await db.get(CvDocumentVersion, version_id)
            assert version.docx_content == b"frozen fixture"
            assert public.json()["cv_html"] == version.content_html
        chat = await app_client.post(
            f"/api/public/cv-i/{token}/chat", json={"question": "What experience?"}
        )
        assert chat.status_code == 404, chat.text

        from unittest.mock import AsyncMock
        from app.services import cv_generated_editor

        review = AsyncMock(return_value={"status": "verified"})
        monkeypatch.setattr(cv_generated_editor, "review_for_approval", review)
        loaded = await app_client.get(f"{url}/editor", headers=app_auth_headers)
        assert loaded.status_code == 200, loaded.text
        assert loaded.json()["status"] == "finalized"
        revision = loaded.json()["edit_revision"]
        opened = await app_client.post(
            f"{url}/editor/new-draft",
            headers=app_auth_headers,
            json={"expected_revision": revision},
        )
        assert opened.status_code == 200, opened.text
        revision = opened.json()["edit_revision"]
        submitted = "<p>Reviewed <strong>Python</strong> API maintenance.</p>"
        stale = await app_client.patch(
            f"{url}/editor",
            headers=app_auth_headers,
            json={"expected_revision": revision - 1, "content_html": "<p>Stale</p>"},
        )
        assert stale.status_code == 409, stale.text
        finalized = await app_client.post(
            f"{url}/editor/finalize",
            headers=app_auth_headers,
            json={"expected_revision": revision, "content_html": submitted},
        )
        assert finalized.status_code == 200, finalized.text
        assert finalized.json()["version"] == 2
        assert finalized.json()["document_version_id"] != version_id
        assert review.await_count == 1
        assert review.call_args.args[2] == submitted
        still_original = await app_client.get(f"/api/public/cv-i/{token}")
        assert still_original.status_code == 200, still_original.text
        assert still_original.json()["cv_html"] == public.json()["cv_html"]
        old_download = await app_client.get(
            f"{url}/editor/versions/1/docx", headers=app_auth_headers
        )
        assert old_download.status_code == 200, old_download.text
        assert old_download.content == b"frozen fixture"
        new_download = await app_client.get(
            f"{url}/editor/versions/2/docx", headers=app_auth_headers
        )
        assert new_download.status_code == 200
        from docx import Document
        from io import BytesIO

        rendered = Document(BytesIO(new_download.content))
        assert any(
            "Reviewed Python API maintenance." == p.text for p in rendered.paragraphs
        )
    finally:
        async with AsyncSessionLocal() as db:
            await db.execute(
                delete(CvGeneratedDocument).where(
                    CvGeneratedDocument.id == generated_id
                )
            )
            await db.commit()
