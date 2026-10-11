"""Żaden czytnik wgranego pliku Worda nie rozwija komórek przez ``row.cells``.

``row.cells`` z python-docx powtarza komórkę ``gridSpan`` razy, a tę liczbę
podaje plik — bez górnej granicy. CV z formularza kariery, plik requestu
i wzór Profilu Championa są niezaufane i czytane w procesie aplikacji, więc
jedna komórka z ``gridSpan="5000000"`` zajmowała proces na minuty. #2112
i #2119 naprawiły import do edytora i QC CV; tu pozostałe pięć czytników.

Dla prawdziwych dokumentów wynik ma zostać taki sam jak dotąd — dlatego testy
porównują nowe czytniki z dawną regułą na dokumencie ze scaleniami.

Druga liczba z pliku to długość scalenia w pionie: ``row.cells`` wspina się do
początku scalenia osobno z każdego wiersza (koszt kwadratowy), a tekst scalonej
komórki wraca raz na wiersz — stąd testy długiego scalenia i limitu znaków.
"""

from __future__ import annotations

import ast
import copy
import io
import time
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import _Row

from app.core import docx_cells
from app.services import champion_document
from app.services import cv_text_extractor
from app.services.cv_generator_b2b import text_extractor as generator_text
from app.services.cv_generator_b2b.legacy_v7 import text_extractor as legacy_text

HUGE_SPAN = 5_000_000


def _set(tc, tag: str, value: str | None = None) -> None:
    element = OxmlElement(tag)
    if value is not None:
        element.set(qn("w:val"), value)
    tc.get_or_add_tcPr().append(element)


def _docx(*, grid_span: int | None = None, merged: bool = False) -> bytes:
    """Tabela 3 × 3; ``merged`` scala komórkę w poziomie i kolumnę w pionie."""
    doc = Document()
    doc.add_paragraph("Doświadczenie")
    table = doc.add_table(rows=3, cols=3)
    texts = [
        ["2021–2024", "Java Developer", "Bank"],
        ["2019–2021", "Spring", "Kafka"],
        ["2017–2019", "SQL", "Oracle"],
    ]
    # Komórki bierzemy przed zmianą XML-a: ``table.cell`` też rozwija gridSpan.
    tcs = [[table.cell(r, c)._tc for c in range(3)] for r in range(3)]
    for r, row in enumerate(texts):
        for c, text in enumerate(row):
            table.cell(r, c).text = text
    if merged:
        # Wiersz 1: środkowa komórka zajmuje dwie kolumny (trzecia znika).
        _set(tcs[0][1], "w:gridSpan", "2")
        tcs[0][2].getparent().remove(tcs[0][2])
        # Pierwsza kolumna wierszy 2–3 scalona w pionie.
        _set(tcs[1][0], "w:vMerge", "restart")
        _set(tcs[2][0], "w:vMerge")
        for paragraph in list(tcs[2][0].p_lst):
            for run in list(paragraph.r_lst):
                paragraph.remove(run)
    if grid_span is not None:
        _set(tcs[0][0], "w:gridSpan", str(grid_span))
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _chain_docx(rows: int, text: str = "Okres", *, root: bool = True) -> bytes:
    """Jedna kolumna scalona w pionie przez ``rows`` wierszy.

    ``root=False`` — spreparowany plik: dalszy ciąg scalenia już w pierwszym
    wierszu, bez początku nad sobą.
    """
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = text
    first = table.rows[0]._tr
    _set(first.tc_lst[0], "w:vMerge", "restart" if root else None)
    ghost = copy.deepcopy(first)
    ghost_tc = ghost.tc_lst[0]
    ghost_tc.tcPr.find(qn("w:vMerge")).attrib.pop(qn("w:val"), None)
    if root:
        for paragraph in ghost_tc.p_lst:
            for run in list(paragraph.r_lst):
                paragraph.remove(run)
    for _ in range(rows - 1):
        first.addnext(copy.deepcopy(ghost))
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def forbid_row_cells(monkeypatch):
    def boom(self):
        raise AssertionError("czytnik wgranego pliku użył row.cells")

    monkeypatch.setattr(_Row, "cells", property(boom))


def _table(data: bytes):
    return Document(io.BytesIO(data)).tables[0]


def _rows(data: bytes):
    return _table(data).rows


# ── pomocnicy ────────────────────────────────────────────────────────────────


def test_row_cells_reads_a_spanned_cell_once_and_skips_vertical_continuations():
    rows = _rows(_docx(merged=True))
    assert [cell.text for cell in docx_cells.row_cells(rows[0])] == [
        "2021–2024",
        "Java Developer",
    ]
    # Dalszy ciąg scalenia w pionie nie ma własnej treści.
    assert [cell.text for cell in docx_cells.row_cells(rows[2])] == ["SQL", "Oracle"]


def test_row_texts_keeps_column_positions_of_vertical_continuations():
    rows = _rows(_docx(merged=True))
    assert docx_cells.row_texts(rows[1]) == ["2019–2021", "Spring", "Kafka"]
    assert docx_cells.row_texts(rows[2]) == ["", "SQL", "Oracle"]


def test_grid_rows_match_row_cells_of_python_docx_for_real_documents():
    for data in (_docx(), _docx(merged=True), _chain_docx(40)):
        table = _table(data)
        for cells, row in zip(docx_cells.grid_rows(table), table.rows, strict=True):
            assert [c._tc for c in cells] == [c._tc for c in row.cells]


def test_grid_rows_cap_the_span_taken_from_the_file():
    table = _table(_docx(grid_span=HUGE_SPAN))
    started = time.perf_counter()
    cells = next(docx_cells.grid_rows(table))
    assert time.perf_counter() - started < 1
    assert len(cells) == docx_cells.MAX_GRID_SPAN + 2
    assert len(docx_cells.row_cells(table.rows[0])) == 3


def test_grid_rows_read_a_merge_continuation_without_a_root_as_its_own_cell():
    # python-docx rzuca tu ValueError („no tr above topmost tr”).
    table = _table(_chain_docx(3, root=False))
    assert [[c.text for c in cells] for cells in docx_cells.grid_rows(table)] == [
        ["Okres"],
        ["Okres"],
        ["Okres"],
    ]


def test_grid_cell_texts_stop_at_the_character_limit(monkeypatch):
    monkeypatch.setattr(docx_cells, "MAX_TABLE_TEXT_CHARS", 1_000)
    table = _table(_chain_docx(200, "x" * 300))
    texts = docx_cells.grid_cell_texts([table])
    # Bez limitu: 200 wierszy × 300 znaków = 60 000 znaków.
    assert len(texts) == 4
    assert sum(len(text) for text in texts) == 1_200


# ── czytniki: plik z ogromnym gridSpan ───────────────────────────────────────


def _read_cv_text(data: bytes, tmp_path: Path) -> str:
    path = tmp_path / "cv.docx"
    path.write_bytes(data)
    text = cv_text_extractor._extract_docx(str(path))
    assert text is not None
    return text


READERS = {
    "champion_document.document_text": lambda data, tmp: (
        champion_document.document_text(data)
    ),
    "cv_text_extractor._extract_docx": _read_cv_text,
    "generator _extract_docx": lambda data, tmp: generator_text._extract_docx(data),
    "legacy_v7 _extract_docx": lambda data, tmp: legacy_text._extract_docx(data),
}


@pytest.mark.parametrize("name", sorted(READERS))
def test_reader_does_not_expand_a_huge_grid_span(name, tmp_path, forbid_row_cells):
    started = time.perf_counter()
    text = READERS[name](_docx(grid_span=HUGE_SPAN), tmp_path)
    assert time.perf_counter() - started < 5
    # Treść jest odczytana, a tekst scalonej komórki nie wraca milion razy.
    assert "Java Developer" in text and "Oracle" in text
    assert text.count("2021–2024") <= docx_cells.MAX_GRID_SPAN


LONG_MERGE = 6_000


@pytest.mark.parametrize("name", sorted(READERS))
def test_reader_reads_a_long_vertical_merge_in_linear_time(
    name, tmp_path, forbid_row_cells
):
    # Wspinanie się do początku scalenia z każdego wiersza: 600 wierszy = 18 s,
    # 2000 wierszy = 2 minuty (pomiar 11.10.2026); jednym przejściem — ułamek sekundy.
    started = time.perf_counter()
    text = READERS[name](_chain_docx(LONG_MERGE), tmp_path)
    assert time.perf_counter() - started < 10
    assert "Okres" in text


@pytest.mark.parametrize("name", sorted(READERS))
def test_reader_survives_a_merge_continuation_without_a_root(name, tmp_path):
    # Na `row.cells` taki plik kończył się ValueError. Czytniki, które biorą
    # komórkę raz, pomijają dalszy ciąg scalenia; pozostałe czytają go jak komórkę.
    text = READERS[name](_chain_docx(3, root=False), tmp_path)
    assert ("Okres" in text) == (
        name in {"cv_text_extractor._extract_docx", "legacy_v7 _extract_docx"}
    )


def test_table_profile_does_not_expand_a_huge_grid_span(forbid_row_cells):
    started = time.perf_counter()
    # To nie jest wzór Profilu Championa — liczy się, że odczyt się kończy.
    assert champion_document.table_profile(_docx(grid_span=HUGE_SPAN)) is None
    assert time.perf_counter() - started < 5


# ── czytniki: prawdziwy dokument czyta się jak dotąd ─────────────────────────


def _old_grid_text(data: bytes) -> str:
    """Dawna reguła `cv_text_extractor` i `legacy_v7`: akapity, potem każda pozycja siatki."""
    doc = Document(io.BytesIO(data))
    lines = [p.text for p in doc.paragraphs if p.text]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    lines.append(cell.text)
    return "\n".join(lines)


@pytest.mark.parametrize("merged", [False, True])
def test_cv_text_and_legacy_generator_text_are_unchanged(merged, tmp_path):
    data = _docx(merged=merged)
    expected = _old_grid_text(data)
    if merged:
        # Dawna reguła powtarza scaloną komórkę — i tak ma zostać.
        assert expected.count("Java Developer") == 2
        assert expected.count("2019–2021") == 2
    assert _read_cv_text(data, tmp_path) == expected
    assert legacy_text._extract_docx(data) == expected


def test_cv_text_of_a_vertical_merge_is_unchanged(tmp_path):
    data = _chain_docx(40)
    expected = _old_grid_text(data)
    assert expected.count("Okres") == 40
    assert _read_cv_text(data, tmp_path) == expected
    assert legacy_text._extract_docx(data) == expected


def test_deduplicating_readers_read_a_merged_cell_once():
    data = _docx(merged=True)
    for text in (
        champion_document.document_text(data),
        generator_text._extract_docx(data),
    ):
        assert text.count("Java Developer") == 1
        assert text.count("2019–2021") == 1
        assert "Oracle" in text


# ── strażnik: nowy czytnik nie wraca do row.cells ────────────────────────────

# Pliki, które chodzą po tabelach dokumentów zbudowanych PRZEZ NAS (szablony
# umów, wygenerowane CV) — tam gridSpan pochodzi z naszego kodu.
TRUSTED_DOCUMENTS = {
    "app/services/b2b_contract_generator/company_variant.py",
    "app/services/cv_approved_docx.py",
    "app/services/cv_generator_b2b/docx_renderer.py",
    # `row.cells` to tam słownik kolumn odczytu PDF-a, nie tabela Worda.
    "app/services/order_policies/pko_bp_layout.py",
}


def test_no_reader_of_uploaded_word_files_uses_row_cells():
    backend = Path(__file__).resolve().parents[1]
    offenders = []
    for path in sorted((backend / "app").rglob("*.py")):
        relative = path.relative_to(backend).as_posix()
        if relative in TRUSTED_DOCUMENTS:
            continue
        tree = ast.parse(path.read_text(encoding="utf-8"))
        for node in ast.walk(tree):
            if isinstance(node, ast.Attribute) and node.attr == "cells":
                offenders.append(f"{relative}:{node.lineno}")
    assert offenders == [], (
        "Komórki tabeli Worda czytaj przez app.core.docx_cells (row_cells / row_texts / "
        "grid_cell_texts), nie przez row.cells — gridSpan pochodzi z pliku: "
        + ", ".join(offenders)
    )
