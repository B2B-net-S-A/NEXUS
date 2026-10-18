"""Wspólne narzędzia testów pakietu klienta, wzorów protokołów i maila powitalnego.

Dane wyłącznie fikcyjne (repo jest publiczne): „Przykładowi” ludzie, adresy
``@example.com``, kwoty zmyślone.
"""

from __future__ import annotations

import io
import uuid
from contextlib import asynccontextmanager
from datetime import date, timedelta
from decimal import Decimal
from typing import Optional

import docx
from sqlalchemy import select, text

import app.models  # noqa: F401  (rejestracja wszystkich mapperów)
from app.core.database import AsyncSessionLocal
from app.core.scheduling import business_today
from app.core.security import create_access_token
from app.models.app_setting import AppSetting
from app.models.candidate import Candidate
from app.models.client import Client, ClientStatus
from app.models.contact import Contact
from app.models.contract import Contract, ContractStatus, RateUnit
from app.models.section_permission import UserActionOverride
from app.models.team_structure import DeliveryLeadClientAssignment
from app.models.user import User, UserRole
from app.services.hr_foundation_schema import HR_MODULE_KEYS, HR_MODULES_SETTING

PREFIX = "/api/client-onboarding"


def tag() -> str:
    return uuid.uuid4().hex[:8]


def headers(user: User) -> dict[str, str]:
    token = create_access_token(subject=user.id, role=user.role.value)
    return {"Authorization": f"Bearer {token}"}


def ensure_router_mounted() -> None:
    """Zamontuj router, jeśli koordynator jeszcze go nie wpiął w ``main.py``.

    Warunek staje się no-opem, gdy montaż wejdzie do ``main.py`` — testy
    sprawdzają realny router pod docelowym adresem.
    """
    from app.main import app

    path = f"{PREFIX}/clients/{{client_id}}/pack"
    if any(getattr(route, "path", None) == path for route in app.routes):
        return
    from app.api import client_onboarding

    app.include_router(client_onboarding.router, prefix=PREFIX, tags=["hr"])


async def make_user(
    role: UserRole, *extra: UserRole, grants: tuple[str, ...] = ()
) -> User:
    marker = tag()
    async with AsyncSessionLocal() as db:
        user = User(
            email=f"onb-{marker}@example.com",
            name=f"Osoba Przykładowa {marker}",
            role=role,
            roles=[role.value, *(r.value for r in extra)],
            is_active=True,
            profile_completed=True,
        )
        db.add(user)
        await db.flush()
        for action in grants:
            db.add(UserActionOverride(user_id=user.id, action=action, access="manage"))
        await db.commit()
        return user


async def make_client(*, dl: Optional[User] = None, provider: str = "nieznany") -> int:
    async with AsyncSessionLocal() as db:
        client = Client(name=f"Bank Przykładowy {tag()}", status=ClientStatus.active)
        db.add(client)
        await db.flush()
        await db.execute(
            text("UPDATE clients SET equipment_provider = :p WHERE id = :id"),
            {"p": provider, "id": client.id},
        )
        if dl is not None:
            db.add(
                DeliveryLeadClientAssignment(
                    delivery_lead_user_id=dl.id, client_id=client.id, is_head=True
                )
            )
        await db.commit()
        return client.id


async def make_contact(client_id: int, name: str = "Piotr Przykładowy") -> int:
    async with AsyncSessionLocal() as db:
        contact = Contact(client_id=client_id, name=f"{name} {tag()}")
        db.add(contact)
        await db.commit()
        return contact.id


async def make_contract(
    client_id: int,
    *,
    start: Optional[date] = None,
    email: Optional[str] = "auto",
    manager_contact_id: Optional[int] = None,
    status: ContractStatus = ContractStatus.active,
) -> int:
    marker = tag()
    async with AsyncSessionLocal() as db:
        candidate = Candidate(
            name="Marek",
            lastname=f"Przykładowy-{marker}",
            email=f"kons-{marker}@example.com" if email == "auto" else email,
        )
        db.add(candidate)
        await db.flush()
        contract = Contract(
            candidate_id=candidate.id,
            client_id=client_id,
            status=status,
            start_date=start or business_today() + timedelta(days=14),
            rate_candidate=Decimal("100.000"),
            rate_client=Decimal("130.000"),
            rate_unit=RateUnit.hourly,
            client_pm_contact_id=manager_contact_id,
        )
        db.add(contract)
        await db.commit()
        return contract.id


async def set_setting(key: str, value: Optional[dict]) -> None:
    async with AsyncSessionLocal() as db:
        row = await db.get(AppSetting, key)
        if value is None:
            if row is not None:
                await db.delete(row)
        elif row is None:
            db.add(AppSetting(key=key, value=value))
        else:
            row.value = value
        await db.commit()


async def get_setting(key: str) -> Optional[dict]:
    async with AsyncSessionLocal() as db:
        return await db.scalar(select(AppSetting.value).where(AppSetting.key == key))


@asynccontextmanager
async def modules_state(on: bool):
    """Przełączniki modułów „Ludzi” na czas testu (stan sprzed testu wraca)."""
    previous = await get_setting(HR_MODULES_SETTING)
    await set_setting(HR_MODULES_SETTING, {key: on for key in HR_MODULE_KEYS})
    try:
        yield
    finally:
        await set_setting(
            HR_MODULES_SETTING, previous or {key: False for key in HR_MODULE_KEYS}
        )


def docx_bytes(*paragraphs: str) -> bytes:
    """DOCX z podanymi akapitami (pola jinja w tekście akapitów)."""
    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


#: Wzór wydania z kaucją: wszystkie wymagane pola i pętla pozycji akapitami.
DEPOSIT_PARAGRAPHS = (
    "PROTOKÓŁ Z KAUCJĄ nr {{ number }}",
    "Data: {{ protocol_date }}",
    "Odbiera: {{ holder_name }}, wydaje: {{ issuer_name }}",
    "{%p for line in lines %}",
    "{{ line.position }}. {{ line.name }} ({{ line.identifier }}) {{ line.value }}",
    "{%p endfor %}",
    "Razem {{ total }} zł ({{ total_words }}). Kaucja według umowy.",
)


def deposit_template() -> bytes:
    return docx_bytes(*DEPOSIT_PARAGRAPHS)


def document_text(content: bytes) -> str:
    return "\n".join(p.text for p in docx.Document(io.BytesIO(content)).paragraphs)
