"""Runda 9 (R9-X1-10): eksport powiadomień DL bez ORM i openpyxl na pętli."""

from __future__ import annotations

import io
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock

import pytest
from openpyxl import load_workbook

from app.api import dl_alerts as mod
from app.models.dl_alert import DL_ALERT_STATUS_HANDLED, DL_ALERT_STATUS_NEW


@pytest.mark.asyncio
async def test_export_builds_the_sheet_in_a_thread_from_plain_rows(monkeypatch) -> None:
    created = datetime(2026, 9, 1, 8, 0, tzinfo=timezone.utc)
    rows = [
        (
            5,
            "Anna",
            "a@example.com",
            "",
            "=HYPERLINK(1)",
            "order_ending",
            "Treść",
            created,
            created + timedelta(hours=2, minutes=5),
            DL_ALERT_STATUS_HANDLED,
        ),
        (
            6,
            "",
            "b@example.com",
            None,
            "Klient B",
            "x",
            "M",
            created,
            None,
            DL_ALERT_STATUS_NEW,
        ),
    ]
    result = MagicMock()
    result.all.return_value = rows
    db = AsyncMock()
    db.execute.return_value = result

    threaded: list[object] = []
    real_to_thread = mod.asyncio.to_thread

    async def spy(func, *args, **kwargs):
        threaded.append(func)
        return await real_to_thread(func, *args, **kwargs)

    monkeypatch.setattr(mod.asyncio, "to_thread", spy)

    resp = await mod.export_alerts(
        user=SimpleNamespace(id=5), db=db, scope="mine", limit=10
    )

    assert threaded == [mod._build_export_xlsx]
    body = b"".join([chunk async for chunk in resp.body_iterator])
    sheet = load_workbook(io.BytesIO(body)).active
    values = [[c.value for c in r] for r in sheet.iter_rows()]
    assert values[0] == mod._EXPORT_HEADER
    assert values[1][0] == "Anna"
    assert values[1][1] == "'=HYPERLINK(1)"  # formuła zneutralizowana
    assert values[1][6] == "2 h 5 min"
    assert values[2][0] == "b@example.com"
    assert values[2][5] in ("", None)
    assert values[2][6] == "—"
