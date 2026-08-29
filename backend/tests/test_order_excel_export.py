from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from app.services.order_excel_export import (
    OrderExportRow,
    build_orders_workbook,
    orders_export_filename,
)


def test_multi_order_workbook_has_typed_values_and_readable_headers():
    content = build_orders_workbook(
        [
            OrderExportRow(
                consultant_name="Jan Nowak",
                order_number="274607",
                cost_rate=Decimal("100.25"),
                revenue_rate=Decimal("150.50"),
                start_date=date(2026, 1, 1),
                end_date=date(2026, 12, 31),
                allocation=Decimal("60"),
                consumption=Decimal("50"),
            ),
            OrderExportRow(
                consultant_name="Jan Kowalski",
                order_number="274607",
                cost_rate=Decimal("110"),
                revenue_rate=Decimal("160"),
                start_date=date(2026, 1, 1),
                end_date=date(2026, 12, 31),
                allocation=Decimal("50"),
                consumption=Decimal("45"),
            ),
        ],
        include_model_columns=True,
    )
    sheet = load_workbook(io.BytesIO(content)).active
    assert [cell.value for cell in sheet[1]] == [
        "Imię i nazwisko",
        "Numer zamówienia",
        "Stawka kosztowa",
        "Stawka przychodowa",
        "Okres zamówienia",
        "Liczba MD / Kwota zamówienia",
        "Zużycie zamówienia",
    ]
    assert sheet["B2"].value == "274607"
    assert sheet["C2"].value == 100.25
    assert sheet["F2"].value == 60
    assert sheet["G3"].value == 45
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:G3"


def test_export_filename_contains_folded_client_name_and_date():
    assert (
        orders_export_filename("Nordea Bank ABP", date(2026, 8, 21))
        == "Zamowienia_Nordea_Bank_ABP_21.08.2026.xlsx"
    )


def test_user_text_cannot_become_an_excel_formula():
    content = build_orders_workbook(
        [
            OrderExportRow(
                consultant_name='=HYPERLINK("bad")',
                order_number="+CMD",
                cost_rate=None,
                revenue_rate=None,
                start_date=None,
                end_date=None,
            )
        ],
        include_model_columns=False,
    )
    sheet = load_workbook(io.BytesIO(content), data_only=False).active
    assert sheet["A2"].data_type == "s"
    assert sheet["A2"].value.startswith("'=")
    assert sheet["B2"].value.startswith("'+")


def test_unified_workbook_adds_order_type_after_full_combined_columns():
    content = build_orders_workbook(
        [
            OrderExportRow(
                consultant_name="Anna Nowak",
                order_number="PO-42",
                cost_rate=Decimal("100"),
                revenue_rate=Decimal("150"),
                start_date=date(2026, 1, 1),
                end_date=None,
                allocation=Decimal("80"),
                consumption=Decimal("12"),
                order_type="MD",
            )
        ],
        include_model_columns=True,
        include_order_type=True,
    )

    sheet = load_workbook(io.BytesIO(content)).active
    assert [cell.value for cell in sheet[1]] == [
        "Imię i nazwisko",
        "Numer zamówienia",
        "Stawka kosztowa",
        "Stawka przychodowa",
        "Okres zamówienia",
        "Liczba MD / Kwota zamówienia",
        "Zużycie zamówienia",
        "Typ zamówienia",
    ]
    assert sheet["H2"].value == "MD"
    assert sheet.auto_filter.ref == "A1:H2"
