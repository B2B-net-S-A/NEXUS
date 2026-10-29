"""Fikcyjna „Karta przed CV” w układzie wzoru Credit Agricole (29.10.2026).

Prawdziwy wzór ma telefony opiekunów i nie trafia do publicznego repo —
ten plik ma ten sam układ akapitów i przebiegów, ale zmyślone osoby.
"""

from __future__ import annotations

import io

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

GUARDIAN_LINE = "Anna Przykładowa, tel. +48 000 000 000, e-mail: "
GUARDIAN_EMAIL = "anna.przykladowa@example.com"


def _runs(paragraph, *parts: tuple[str, bool]) -> None:
    for text, bold in parts:
        run = paragraph.add_run(text)
        run.bold = bold or None


def _hyperlink(paragraph, text: str) -> None:
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), "kontakt")
    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    link.append(run)
    paragraph._p.append(link)


def cover_docx(*, labels: bool = True) -> bytes:
    document = docx.Document()
    document.add_paragraph().add_run("B2B.NET S.A. ")
    document.add_paragraph()
    if labels:
        _runs(document.add_paragraph(), ("Imie", True), (" i nazwisko", True))
        _runs(document.add_paragraph(), ("Rola:", True), (" ", False), ("xxx", False))
        _runs(
            document.add_paragraph(),
            ("Poziom doświadczenia:", True),
            (" ", False),
            ("Junior", False),
            ("/Specjalista/", False),
            ("Expert", False),
        )
        _runs(
            document.add_paragraph(),
            ("Staż:", True),
            (" ", False),
            ("x", False),
            (" doświadczenia komercyjnego", False),
        )
        _runs(
            document.add_paragraph(),
            ("Profil kompetencji:", True),
            (" ", False),
            ("xxx", False),
        )
        availability = document.add_paragraph()
        _runs(availability, ("Dostępność:", True), (" ", False), ("xx.xxxx", False))
        _runs(availability, (" ", False))
        availability.add_run().add_break()
        _runs(availability, ("Lokalizacja:", True), (" ", False), ("xxxxxx", False))
        _runs(
            document.add_paragraph(),
            ("Wynagrodzenie:", True),
            (" ", False),
            ("xxxx", False),
            (",00 PLN netto za 1 MD", False),
        )
    else:
        document.add_paragraph("Oferta współpracy — opis zespołu i zasad rozliczeń.")
        _runs(document.add_paragraph(), ("Rola:", True), (" ", False), ("xxx", False))
    _runs(document.add_paragraph(), ("Opiekun po stronie dostawcy:", True))
    guardian = document.add_paragraph()
    _runs(guardian, (GUARDIAN_LINE, False))
    _hyperlink(guardian, GUARDIAN_EMAIL)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()
