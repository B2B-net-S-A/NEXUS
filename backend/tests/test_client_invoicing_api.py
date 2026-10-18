"""Faktury do klientów (E8) — trasy: dostęp, zakładka klienta, Finanse, wyciąg, eksport.

Zakres: rekruter, TCM i DL cudzego klienta nie widzą kwot (403), DL widzi
swojego klienta, Finanse wszystkich; wyłączony moduł = 404 (admin przechodzi);
„do faktury” obok faktur z KSeF ze stanem; przypisanie, pominięcie, uznanie
różnicy i ręczne „opłacono” tylko Finanse; uznania z wyciągu opłacają faktury
(także z importu Płac, bez wpływu na obciążenia). Dane fikcyjne.
"""

from __future__ import annotations

import io
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path

import pytest
import pytest_asyncio
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.api import client_sales_invoices as api
from app.api import hr_payments
from app.core.config import settings
from app.core.database import AsyncSessionLocal
from app.models.client_sales_invoice import ClientSalesInvoicePayment
from app.models.section_permission import UserActionOverride
from app.services.client_invoicing import sync
from app.services.client_invoicing.views import ROW_STATUS_LABELS
from app.services.hr.bank import statements
from app.services.hr_foundation_schema import HR_MODULE_KEYS, HR_MODULES_SETTING
from tests.client_invoicing import helpers as fx

pytestmark = pytest.mark.asyncio

CLIENTS = "/api/clients"
FINANCE = "/api/finance/sales-invoices"
STATEMENT = (
    Path(__file__).parent / "data" / "bank" / "sales" / "camt053_wplaty_klientow.xml"
)
TODAY = date(2026, 11, 20)
OCTOBER = {"year": 2026, "month": 10}


@pytest_asyncio.fixture
async def client():
    from app.main import app

    mounted = any(getattr(r, "path", "").startswith(FINANCE) for r in app.routes)
    target = app
    if not mounted:
        target = FastAPI()
        target.include_router(api.router, prefix=CLIENTS)
        target.include_router(api.finance_router, prefix=FINANCE)
    async with AsyncClient(
        transport=ASGITransport(app=target), base_url="http://testserver"
    ) as http:
        yield http


@pytest_asyncio.fixture
async def payroll_client():
    """Trasy faktur do klientów i wyciągu Płac (``/api/hr/payments``) razem."""

    target = FastAPI()
    target.include_router(api.router, prefix=CLIENTS)
    target.include_router(api.finance_router, prefix=FINANCE)
    target.include_router(hr_payments.router, prefix="/api/hr/payments")
    async with AsyncClient(
        transport=ASGITransport(app=target), base_url="http://testserver"
    ) as http:
        yield http


@pytest_asyncio.fixture
async def seed(monkeypatch):
    monkeypatch.setattr(api, "business_today", lambda: TODAY)
    monkeypatch.setattr(settings, "KSEF_MODE", "simulate")
    monkeypatch.setattr(settings, "KSEF_CONTEXT_NIP", fx.COMPANY_NIP)
    monkeypatch.setattr(settings, "KSEF_SIMULATE_DIR", str(fx.DATA_DIR.parent))
    data = await fx.seed_month()
    await fx.set_module(True)
    async with AsyncSessionLocal() as db:
        result = await sync.run_once(
            db, now=datetime(2026, 11, 20, 12, 0, tzinfo=timezone.utc)
        )
    assert result["saved"] == 5, result
    yield data
    await fx.wipe()
    await fx.set_module(None)


async def _tab(client, user, seed, **params):
    return await client.get(
        f"{CLIENTS}/{seed.client_id}/invoicing",
        params={**OCTOBER, **params},
        headers=fx.headers(user),
    )


async def _upload(client, user, raw: bytes):
    return await client.post(
        f"{FINANCE}/statements",
        files={"file": ("wyciag.xml", io.BytesIO(raw), "application/xml")},
        headers=fx.headers(user),
    )


# ── Dostęp ────────────────────────────────────────────────────────────────────


async def test_amounts_only_for_finance_and_the_clients_delivery_lead(client, seed):
    users = seed.users
    assert (await _tab(client, users["finance"], seed)).status_code == 200
    assert (await _tab(client, users["dl"], seed)).status_code == 200
    assert (await _tab(client, users["admin"], seed)).status_code == 200

    foreign = await _tab(client, users["other_dl"], seed)
    assert foreign.status_code == 403
    assert "poza Twoim portfelem" in foreign.text
    for role in ("recruiter", "tcm"):
        denied = await _tab(client, users[role], seed)
        assert denied.status_code == 403, role
        assert "25200" not in denied.text

    # Delivery Lead nie wchodzi do ekranu Finansów ani nie przypisuje faktur.
    invoice = await fx.invoice(fx.KSEF_UNKNOWN)
    listing = await client.get(FINANCE, params=OCTOBER, headers=fx.headers(users["dl"]))
    assert listing.status_code == 403
    assign = await client.post(
        f"{FINANCE}/{invoice.id}/ignore",
        json={"reason": "sprzedaż sprzętu"},
        headers=fx.headers(users["dl"]),
    )
    assert assign.status_code == 403


async def test_module_off_hides_screens_except_for_admin(client, seed):
    await fx.set_module(False)
    off = await _tab(client, seed.users["finance"], seed)
    assert off.status_code == 404
    assert off.json()["detail"] == "Moduł faktur do klientów jest wyłączony."
    admin = await _tab(client, seed.users["admin"], seed)
    assert admin.status_code == 200
    assert admin.json()["module_enabled"] is False
    # Przełącznik zmienia tylko admin.
    finance_toggle = await client.put(
        f"{FINANCE}/module",
        json={"enabled": True},
        headers=fx.headers(seed.users["finance"]),
    )
    assert finance_toggle.status_code == 403
    toggle = await client.put(
        f"{FINANCE}/module",
        json={"enabled": True},
        headers=fx.headers(seed.users["admin"]),
    )
    assert toggle.status_code == 200 and toggle.json() == {"enabled": True}


async def test_unknown_client_is_404(client, seed):
    response = await client.get(
        f"{CLIENTS}/999999999/invoicing",
        params=OCTOBER,
        headers=fx.headers(seed.users["finance"]),
    )
    assert response.status_code == 404


# ── Zakładka klienta ──────────────────────────────────────────────────────────


async def test_client_tab_compares_orders_with_ksef_invoices(client, seed):
    response = await _tab(client, seed.users["dl"], seed)
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["period"] == "2026-10"
    assert body["can_manage"] is False
    rows = fx.as_dict(body["rows"], "order_number")
    assert rows["ZAM/2026/118"]["status"] == "ok"
    assert rows["ZAM/2026/118"]["expected_net"] == "25200.00"
    assert rows["ZAM/2026/121"]["status"] == "ok"
    assert rows["ZAM/2026/121"]["lines"][0]["unit"] == "h"
    # 22 500 z faktury − 75 z korekty = 22 425 do faktury.
    assert rows["ZAM/2026/124"]["status"] == "ok"
    assert rows["ZAM/2026/124"]["invoiced_net"] == "22425.00"
    assert rows["ZAM/2026/130"]["status"] == "not_invoiced"
    assert rows["ZAM/2026/130"]["status_label"] == ROW_STATUS_LABELS["not_invoiced"]
    assert body["summary"]["rows"] == 4
    assert body["summary"]["not_invoiced"] == 1
    invoices = fx.as_dict(body["invoices"], "invoice_number")
    corrected = invoices["FS 10/124/2026"]
    assert corrected["net"] == "22425.00"
    assert corrected["corrections"][0]["net"] == "-75.00"
    assert corrected["check"] == "ok"
    # Nieopłacone, termin 10.11 minął.
    assert invoices["FS 10/121/2026"]["overdue_days"] == 10
    assert {item["invoice_number"] for item in body["overdue"]} >= {"FS 10/121/2026"}
    # Faktura dla nabywcy spoza NEXUSA nie trafia do zakładki klienta.
    assert "FS 11/7/2026" not in invoices

    finance = (await _tab(client, seed.users["finance"], seed)).json()
    assert finance["can_manage"] is True


async def test_export_has_every_position_with_polish_decimals(client, seed):
    csv_response = await client.get(
        f"{CLIENTS}/{seed.client_id}/invoicing/export",
        params={**OCTOBER, "format": "csv"},
        headers=fx.headers(seed.users["dl"]),
    )
    assert csv_response.status_code == 200
    assert "do-faktury-klient" in csv_response.headers["content-disposition"]
    text_ = csv_response.content.decode("utf-8-sig")
    assert "ZAM/2026/130" in text_ and "Tomasz Wzorcowy" in text_
    assert "24000,00" in text_
    assert fx.CLIENT_NIP in text_

    xlsx = await client.get(
        f"{CLIENTS}/{seed.client_id}/invoicing/export",
        params=OCTOBER,
        headers=fx.headers(seed.users["finance"]),
    )
    sheet = load_workbook(io.BytesIO(xlsx.content)).active
    assert sheet.max_row == 5
    nets = sorted(row[10] for row in sheet.iter_rows(min_row=2, values_only=True))
    assert nets == [22425.0, 24000.0, 25200.0, 31920.0]

    denied = await client.get(
        f"{CLIENTS}/{seed.client_id}/invoicing/export",
        params=OCTOBER,
        headers=fx.headers(seed.users["other_dl"]),
    )
    assert denied.status_code == 403


# ── Finanse: przypisanie, pominięcie, różnica ─────────────────────────────────


async def test_finance_resolves_unmatched_and_accepts_a_difference(client, seed):
    finance = fx.headers(seed.users["finance"])
    november = await client.get(
        FINANCE, params={"year": 2026, "month": 11}, headers=finance
    )
    assert november.status_code == 200
    unknown = fx.as_dict(november.json()["invoices"], "invoice_number")["FS 11/7/2026"]
    assert unknown["match_status"] == "unmatched"
    assert unknown["match_reason_text"].startswith("NIP nabywcy")

    short = await client.post(
        f"{FINANCE}/{unknown['id']}/ignore", json={"reason": "x"}, headers=finance
    )
    assert short.status_code == 422
    ignored = await client.post(
        f"{FINANCE}/{unknown['id']}/ignore",
        json={"reason": "Sprzedaż sprzętu spoza zamówień"},
        headers=finance,
    )
    assert ignored.json()["match_status"] == "ignored"
    reopened = await client.post(f"{FINANCE}/{unknown['id']}/reopen", headers=finance)
    assert reopened.json()["match_status"] == "unmatched"

    # Faktura 118 przypięta do 130: różnica 1 200 do sprawdzenia, potem uznana.
    first = await fx.invoice(fx.KSEF_118)
    moved = await client.post(
        f"{FINANCE}/{first.id}/assign",
        json={"targets": [{"group_id": seed.groups["130"]}], **OCTOBER},
        headers=finance,
    )
    assert moved.status_code == 200, moved.text
    assert (moved.json()["check"], moved.json()["difference"]) == (
        "difference",
        "1200.00",
    )
    options = await client.get(f"{FINANCE}/{first.id}/options", headers=finance)
    labels = {o["label"] for o in options.json()["options"]}
    assert "ZAM/2026/130 · Tomasz Wzorcowy" in labels

    accepted = await client.post(
        f"{FINANCE}/{first.id}/accept-difference",
        json={"reason": "Klient zapłaci za dodatkowy dzień"},
        headers=finance,
    )
    assert accepted.json()["check"] == "accepted"
    tab = (await _tab(client, seed.users["finance"], seed)).json()
    rows = fx.as_dict(tab["rows"], "order_number")
    assert rows["ZAM/2026/130"]["status"] == "accepted"
    assert rows["ZAM/2026/118"]["status"] == "not_invoiced"

    other = await client.post(
        f"{FINANCE}/{first.id}/assign",
        json={"targets": [{"order_id": 999999999}], **OCTOBER},
        headers=finance,
    )
    assert other.status_code == 404
    correction = await fx.invoice(fx.KSEF_CORRECTION)
    refused = await client.post(
        f"{FINANCE}/{correction.id}/assign",
        json={"targets": [{"group_id": seed.groups["124"]}], **OCTOBER},
        headers=finance,
    )
    assert refused.status_code == 409
    no_diff = await client.post(
        f"{FINANCE}/{(await fx.invoice(fx.KSEF_121)).id}/accept-difference",
        json={"reason": "bez różnicy"},
        headers=finance,
    )
    assert no_diff.status_code == 409


async def test_assign_to_another_clients_order_is_refused(client, seed):
    async with AsyncSessionLocal() as db:
        foreign = await fx.periodic_order(
            db, seed.other_client_id, "PO-INNY", rate="100", unit=fx.RateUnit.monthly
        )
        await db.commit()
    first = await fx.invoice(fx.KSEF_118)
    response = await client.post(
        f"{FINANCE}/{first.id}/assign",
        json={"targets": [{"order_id": foreign}], **OCTOBER},
        headers=fx.headers(seed.users["finance"]),
    )
    assert response.status_code == 422
    assert "innego klienta" in response.text


# ── Zapłata ───────────────────────────────────────────────────────────────────


async def test_statement_credits_pay_invoices_and_reimport_changes_nothing(
    client, seed
):
    finance = seed.users["finance"]
    raw = STATEMENT.read_bytes()
    response = await _upload(client, finance, raw)
    assert response.status_code == 200, response.text
    result = response.json()
    assert (
        result["credits"],
        result["matched"],
        result["unmatched"],
        result["skipped"],
    ) == (
        5,
        3,
        1,
        1,
    )
    # Wynik bez kwot i nazw klientów.
    assert "30996" not in response.text and "Bank Testowy" not in response.text

    again = (await _upload(client, finance, raw)).json()
    assert (again["matched"], again["known"]) == (0, 4)

    tab = (await _tab(client, finance, seed)).json()
    invoices = fx.as_dict(tab["invoices"], "invoice_number")
    assert invoices["FS 10/118/2026"]["payment_status"] == "paid"
    assert (
        invoices["FS 10/124/2026"]["payment_status"] == "paid"
    )  # MPP: VAT na rachunku VAT
    partial = invoices["FS 10/121/2026"]
    assert (partial["payment_status"], partial["to_pay"]) == ("partial", "19261.60")
    assert partial["overdue_days"] == 10

    late = await client.get(
        f"{CLIENTS}/sales-invoices/overdue", headers=fx.headers(seed.users["dl"])
    )
    assert [item["invoice_number"] for item in late.json()["invoices"]] == [
        "FS 10/121/2026"
    ]
    nobody = await client.get(
        f"{CLIENTS}/sales-invoices/overdue", headers=fx.headers(seed.users["other_dl"])
    )
    assert nobody.json()["invoices"] == []
    recruiter = await client.get(
        f"{CLIENTS}/sales-invoices/overdue", headers=fx.headers(seed.users["recruiter"])
    )
    assert recruiter.status_code == 403

    credits = (
        await client.get(f"{FINANCE}/credits", headers=fx.headers(finance))
    ).json()
    (pending,) = credits["credits"]
    assert pending["hints"][0]["invoice_number"] == "FS 11/7/2026"
    assigned = await client.post(
        f"{FINANCE}/credits/{pending['id']}/assign",
        json={"invoice_id": pending["hints"][0]["invoice_id"]},
        headers=fx.headers(finance),
    )
    assert assigned.status_code == 200
    detail = (
        await client.get(
            f"{FINANCE}/{pending['hints'][0]['invoice_id']}",
            headers=fx.headers(finance),
        )
    ).json()
    assert detail["payment_status"] == "paid"
    twice = await client.post(
        f"{FINANCE}/credits/{pending['id']}/assign",
        json={"invoice_id": pending["hints"][0]["invoice_id"]},
        headers=fx.headers(finance),
    )
    assert twice.status_code == 409


async def _payroll_upload(client, user, raw: bytes):
    """Wyciąg przez Płace (``/api/hr/payments/statements``) przy włączonych modułach HR."""

    async with AsyncSessionLocal() as db:
        row = await db.get(fx.AppSetting, HR_MODULES_SETTING)
        before = None if row is None else dict(row.value)
        if row is None:
            db.add(
                fx.AppSetting(
                    key=HR_MODULES_SETTING, value={k: True for k in HR_MODULE_KEYS}
                )
            )
        else:
            row.value = {k: True for k in HR_MODULE_KEYS}
        await db.commit()
    try:
        return await client.post(
            "/api/hr/payments/statements",
            files={"file": ("wyciag.xml", io.BytesIO(raw), "application/xml")},
            headers=fx.headers(user),
        )
    finally:
        async with AsyncSessionLocal() as db:
            row = await db.get(fx.AppSetting, HR_MODULES_SETTING)
            if before is None:
                await db.delete(row)
            else:
                row.value = before
            await db.commit()


async def test_payroll_without_finance_only_counts_client_credits(payroll_client, seed):
    """Płace bez modułu Finanse nie opłacają faktur klientów i nie widzą, co pasuje."""

    async with AsyncSessionLocal() as db:
        payroll = await fx.make_user(db, "recruiter", name="Paula Płacowa")
        db.add(
            UserActionOverride(user_id=payroll.id, action="payroll", access="manage")
        )
        await db.commit()
    raw = STATEMENT.read_bytes()
    response = await _payroll_upload(payroll_client, payroll, raw)
    assert response.status_code == 200, response.text
    async with AsyncSessionLocal() as db:
        stored = await db.scalar(
            select(func.count()).select_from(ClientSalesInvoicePayment)
        )
    assert stored == 0
    tab = (await _tab(payroll_client, seed.users["finance"], seed)).json()
    invoices = fx.as_dict(tab["invoices"], "invoice_number")
    assert invoices["FS 10/118/2026"]["payment_status"] == "unpaid"
    credits = response.json()["credits"]
    # Odpowiedź nie mówi, czy numer faktury z tytułu istnieje — tylko ile uznań.
    assert (
        credits["matched"],
        credits["unmatched"],
        credits["skipped"],
        credits["known"],
    ) == (0, 0, 0, 0)
    assert credits["credits"] == 5
    assert credits["restricted"] is True

    # Ten sam plik wgrany potem przez Finanse dopasowuje uznania normalnie.
    finance = (await _upload(payroll_client, seed.users["finance"], raw)).json()
    assert (finance["matched"], finance["unmatched"]) == (3, 1)


async def test_payroll_with_finance_books_client_credits_through_the_payroll_route(
    payroll_client, seed
):
    response = await _payroll_upload(
        payroll_client, seed.users["finance"], STATEMENT.read_bytes()
    )
    assert response.status_code == 200, response.text
    credits = response.json()["credits"]
    assert credits["restricted"] is False
    assert (credits["matched"], credits["unmatched"]) == (3, 1)


async def test_payroll_statement_import_also_books_client_credits(seed, monkeypatch):
    raw = STATEMENT.read_bytes()
    async with AsyncSessionLocal() as db:
        result = await statements.import_statement(
            db, raw, actor_id=seed.users["finance"].id, match_credits=True
        )
        await db.commit()
    # Obciążenia: czynsz nie dotyczy zespołu — jak dotąd tylko licznik.
    assert result.debits == 1
    assert result.matched == 0
    assert result.credits["matched"] == 3 and result.credits["unmatched"] == 1
    async with AsyncSessionLocal() as db:
        stored = list(
            await db.scalars(
                select(ClientSalesInvoicePayment).where(
                    ClientSalesInvoicePayment.match_status == "matched"
                )
            )
        )
    assert sum((p.amount for p in stored), Decimal("0")) == Decimal("78578.75")


async def test_payroll_import_survives_a_failure_on_the_client_side(seed, monkeypatch):
    async def broken(*_args, **_kwargs):
        raise RuntimeError("awaria testowa")

    monkeypatch.setattr(statements.client_payments, "import_credits", broken)
    async with AsyncSessionLocal() as db:
        result = await statements.import_statement(
            db,
            STATEMENT.read_bytes(),
            actor_id=seed.users["finance"].id,
            match_credits=True,
        )
        await db.commit()
    assert result.credits["error"] is True
    assert result.debits == 1


async def test_payroll_import_skips_credits_when_module_is_off(seed):
    await fx.set_module(False)
    async with AsyncSessionLocal() as db:
        result = await statements.import_statement(
            db,
            STATEMENT.read_bytes(),
            actor_id=seed.users["finance"].id,
            match_credits=True,
        )
        await db.commit()
    assert result.credits["disabled"] is True
    assert result.credits["matched"] == 0


async def test_manual_paid_needs_reason_and_a_past_date(client, seed):
    finance = fx.headers(seed.users["finance"])
    invoice = await fx.invoice(fx.KSEF_121)
    future = await client.post(
        f"{FINANCE}/{invoice.id}/mark-paid",
        json={"amount": "100.00", "paid_on": "2026-12-01", "reason": "kompensata"},
        headers=finance,
    )
    assert future.status_code == 422
    paid = await client.post(
        f"{FINANCE}/{invoice.id}/mark-paid",
        json={
            "amount": "39261.60",
            "paid_on": "2026-11-19",
            "reason": "Kompensata z notą",
        },
        headers=finance,
    )
    assert paid.status_code == 200
    assert paid.json()["payment_status"] == "paid"
    again = await client.post(
        f"{FINANCE}/{invoice.id}/mark-paid",
        json={"amount": "1.00", "paid_on": "2026-11-19", "reason": "Kompensata z notą"},
        headers=finance,
    )
    assert again.status_code == 409


async def test_module_state_tells_the_profile_whether_to_show_the_tab(client, seed):
    path = f"{CLIENTS}/sales-invoices/module"
    assert (await client.get(path, headers=fx.headers(seed.users["dl"]))).json() == {
        "enabled": True,
        "open": True,
    }
    await fx.set_module(False)
    assert (await client.get(path, headers=fx.headers(seed.users["dl"]))).json() == {
        "enabled": False,
        "open": False,
    }
    admin = await client.get(path, headers=fx.headers(seed.users["admin"]))
    assert admin.json() == {"enabled": False, "open": True}
    denied = await client.get(path, headers=fx.headers(seed.users["recruiter"]))
    assert denied.status_code == 403
