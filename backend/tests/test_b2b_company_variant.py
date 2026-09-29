"""Generator Umów B2B — wariant umowy dla spółki (ticket 8, 28.09.2026).

Do tej zmiany umowa dla spółki powstawała ze wzoru JDG: komparycja mówiła
„Panem … prowadzącym działalność gospodarczą pod firmą …”, a wzór nie miał
zapisów o osobie, która faktycznie świadczy usługi. Testy pilnują czterech
rzeczy: umowa JDG jest bez zmian, spółka dostaje komparycję z KRS, nowy § 12
przesuwa kolejne paragrafy (także odwołania w klauzulach Klienta) i Załącznik
nr 3 wskazuje osobę z rekrutacji — nie spółkę i nie jej reprezentanta.
"""

from __future__ import annotations

import io
import json
import re
import uuid
from datetime import date
from pathlib import Path

import docx as _docx
import pytest

from app.api.contract_templates import render_contract_template
from app.schemas.b2b_contract_generator import B2BRenderRequest
from app.services.b2b_contract_generator import registry_lookup
from app.services.b2b_contract_generator.clause_overrides import (
    apply_ops_html_counted,
    resolve_override,
)
from app.services.b2b_contract_generator.company_variant import (
    CompanyVariantError,
    apply_company_variant_docx,
    apply_company_variant_html,
    format_share_capital,
    representation_en,
)
from app.services.b2b_contract_generator.docx_renderer import render_from_context
from app.services.b2b_contract_generator.render_context import build_render_context
from app.services.b2b_documents.contract_versions import (
    COMPANY_VERSION,
    CURRENT_VERSION,
    refs_for,
    version_for_variant,
)

TEMPLATE_DIR = Path(__file__).resolve().parents[1] / "app" / "templates" / "contract"
FIXTURES = Path(__file__).resolve().parent / "fixtures" / "krs"

COURT = (
    "Sąd Rejonowy dla m.st. Warszawy w Warszawie, XII Wydział Gospodarczy "
    "Krajowego Rejestru Sądowego"
)

_BASE = dict(
    gender="m",
    partner_name="Jan Kowalski",
    partner_instrumental="Janem Kowalskim",
    partner_legal_name="JK Soft Sp. z o.o.",
    partner_business_address="ul. Prosta 1, 00-001 Warszawa",
    partner_nip="1234567890",
    partner_regon="123456789",
    partner_email="biuro@jksoft.pl",
    partner_phone="+48 600 100 200",
    client_name="Acme",
    project_city="Warszawa",
    project_description="Rozwój platformy.",
    contract_number="1600/2026",
    signing_date=date(2026, 9, 28),
    start_date=date(2026, 10, 1),
    rate_candidate=150,
)

_COMPANY = dict(
    contract_variant="company",
    # W wariancie spółki płeć dotyczy osoby reprezentującej (tu: Anna Nowak).
    gender="k",
    partner_krs="0000123456",
    partner_seat="Warszawa",
    partner_seat_locative="w Warszawie",
    partner_registry_court=COURT,
    partner_share_capital="5000",
    partner_representative_name="Anna Nowak",
    partner_representative_function="Prezes Zarządu",
    partner_representation="Panią Annę Nowak – Prezes Zarządu",
    assigned_person_name="Piotr Wiśniewski",
)


def _req(lang: str = "pl", **extra) -> B2BRenderRequest:
    return B2BRenderRequest(**{**_BASE, "language": lang, **extra})


def _texts(data: bytes) -> list[str]:
    doc = _docx.Document(io.BytesIO(data))
    lines = [p.text for p in doc.paragraphs]
    lines += ["|".join(c.text for c in row.cells) for t in doc.tables for row in t.rows]
    return lines


def _render(req: B2BRenderRequest) -> list[str]:
    lang = req.language
    return _texts(render_from_context(build_render_context(req, None), language=lang))


def _headings(lines: list[str]) -> list[str]:
    return [t.strip() for t in lines if re.fullmatch(r"§\s*\d+A?\s*\.?", t.strip())]


def _after(lines: list[str], heading: str) -> str:
    """Tytuł paragrafu — pierwszy niepusty akapit po nagłówku."""
    idx = next(i for i, t in enumerate(lines) if t.strip() == heading)
    return next(t.strip() for t in lines[idx + 1 :] if t.strip())


def _html(req: B2BRenderRequest) -> str:
    lang = req.language
    ctx = build_render_context(req, None)
    out = render_contract_template(
        (TEMPLATE_DIR / f"umowa_b2b_{lang}.html").read_text(encoding="utf-8"), ctx
    )
    _key, ops = resolve_override(req.client_name, lang)
    if ops:
        out, _ = apply_ops_html_counted(out, ops)
    if ctx["b2b"]["is_company"]:
        out = apply_company_variant_html(
            out, language=lang, assigned_person=ctx["b2b"]["assigned_person"]
        )
    return out


# ── JDG zostaje bez zmian ───────────────────────────────────────────────────


@pytest.mark.parametrize("lang", ["pl", "en"])
def test_default_variant_is_sole_trader_and_old_payloads_render_as_before(lang):
    """Payload zapisany przed ticketem nie ma pola wariantu → umowa JDG."""
    payload = {k: v for k, v in _BASE.items()}
    payload["signing_date"] = "2026-09-28"
    payload["start_date"] = "2026-10-01"
    req = B2BRenderRequest(**payload, language=lang)
    assert req.contract_variant == "jdg"

    lines = _render(req)
    text = "\n".join(lines)
    marker = "pod firmą" if lang == "pl" else "under the name"
    assert marker in text
    assert "Osoby skierowane" not in text and "Designated Person" not in text
    heads = _headings(lines)
    assert heads[-2:] == (["§ 12", "§ 13"])
    assert _after(lines, "§ 12").lower().startswith(("rozwiązanie", "termination"))


@pytest.mark.parametrize("lang", ["pl", "en"])
def test_sole_trader_ignores_company_fields(lang):
    """Dane spółki wpisane w formularzu nie wyciekają do umowy JDG."""
    text = "\n".join(_render(_req(lang, **{**_COMPANY, "contract_variant": "jdg"})))
    assert "0000123456" not in text
    assert "Piotr Wiśniewski" not in text
    assert "Nowak" not in text


# ── Komparycja spółki ───────────────────────────────────────────────────────


def test_company_komparycja_uses_krs_data_pl():
    lines = _render(_req("pl", **_COMPANY))
    komparycja = next(t for t in lines if t.startswith("JK Soft Sp. z o.o. z siedzibą"))
    assert komparycja.startswith(
        "JK Soft Sp. z o.o. z siedzibą w Warszawie, pod adresem ul. Prosta 1, "
        "00-001 Warszawa, wpisaną do Rejestru Przedsiębiorców Krajowego Rejestru "
        f"Sądowego, prowadzonego przez {COURT}, pod numerem KRS: 0000123456, "
        "NIP: 1234567890, REGON: 123456789, o kapitale zakładowym 5.000,00 zł, "
        "reprezentowaną przez Panią Annę Nowak – Prezes Zarządu,"
    )
    text = "\n".join(lines)
    assert "prowadzącym działalność" not in text
    assert "prowadzącą działalność" not in text
    # Partnerem jest spółka — rodzaj żeński w umowie i w umowie powierzenia.
    assert "zwaną w dalszej części umowy „Partnerem”" in text
    assert "zwaną dalej „Podmiotem Przetwarzającym”" in text
    # Umowa powierzenia (Załącznik nr 2) też ma strony w wariancie spółki.
    assert text.count("JK Soft Sp. z o.o. z siedzibą w Warszawie") == 2


def test_company_komparycja_en_translates_representation():
    text = "\n".join(_render(_req("en", **_COMPANY)))
    assert (
        "JK Soft Sp. z o.o. with its registered office in Warszawa, at the "
        "address: ul. Prosta 1, 00-001 Warszawa" in text
    )
    assert "with share capital of PLN 5,000.00" in text
    assert "represented by Ms Anna Nowak – President of the Management Board" in text
    assert "conducting business activity under the name" not in text


def test_representative_is_not_the_assigned_person():
    """Ticket: reprezentant (zarząd) i osoba skierowana to różne osoby."""
    lines = _render(_req("pl", **_COMPANY))
    row = next(t for t in lines if t.startswith("Osoba skierowana:"))
    assert row == "Osoba skierowana:|Piotr Wiśniewski"
    komparycja = next(t for t in lines if t.startswith("JK Soft Sp. z o.o. z siedzibą"))
    assert "Piotr Wiśniewski" not in komparycja


# ── § 12 i przenumerowanie ──────────────────────────────────────────────────


@pytest.mark.parametrize("lang", ["pl", "en"])
def test_company_inserts_section_12_and_shifts_the_rest(lang):
    lines = _render(_req(lang, **_COMPANY))
    heads = _headings(lines)
    assert heads[-3:] == ["§ 12", "§ 13", "§ 14"]
    titles = {
        "pl": ("Osoby skierowane do realizacji Usług", "rozwiązanie", "postanowienia"),
        "en": ("Persons designated to provide the Services", "termination", "final"),
    }[lang]
    assert _after(lines, "§ 12") == titles[0]
    assert _after(lines, "§ 13").lower().startswith(titles[1])
    assert _after(lines, "§ 14").lower().startswith(titles[2])
    text = "\n".join(lines)
    point_8 = (
        "8. Zmiana Osoby Skierowanej wymaga uprzedniej zgody B2BNET lub Klienta B2BNET."
        if lang == "pl"
        else "8. A change of a Designated Person requires the prior consent of B2BNET"
    )
    assert point_8 in text
    # Umowa powierzenia ma własną numerację § 1–5 — bez przesunięcia.
    assert any(t.startswith("§ 5.") for t in lines)


@pytest.mark.parametrize(
    "lang, ref_before, ref_after",
    [
        ("pl", "§ 12 ust. 3 Umowy Głównej", "§ 13 ust. 3 Umowy Głównej"),
        (
            "pl",
            "§ 12 ust. 4 lit. a) Umowy Głównej",
            "§ 13 ust. 4 lit. a) Umowy Głównej",
        ),
        (
            "en",
            "§ 12 section 3 of the Main Agreement",
            "§ 13 section 3 of the Main Agreement",
        ),
    ],
)
def test_client_appendix_references_follow_the_shift(lang, ref_before, ref_after):
    """Załącznik Credit Agricole cytuje rozwiązanie umowy („§ 12 … Umowy
    Głównej”); w umowie spółki ten paragraf ma numer 13. Własna numeracja
    załącznika („§ 12. Sankcje międzynarodowe”) zostaje."""
    client = "Credit Agricole Bank Polska"
    jdg = "\n".join(_render(_req(lang, client_name=client)))
    company = "\n".join(_render(_req(lang, client_name=client, **_COMPANY)))
    assert ref_before in jdg
    assert ref_after in company and ref_before not in company
    own = (
        "§ 12. Sankcje międzynarodowe"
        if lang == "pl"
        else "§ 12. International sanctions"
    )
    assert own in jdg and own in company
    # Odwołania do paragrafów < 12 bez zmian.
    same = "§ 8 Umowy Głównej" if lang == "pl" else "§ 8 of the Main Agreement"
    assert same in company


def test_statute_paragraphs_are_not_shifted():
    """„art. 22 § 1 Kodeksu Pracy” to przepis ustawy, nie paragraf umowy."""
    text = "\n".join(_render(_req("pl", **_COMPANY)))
    assert "art. 22.§ 1 Kodeksu Pracy" in text
    assert "art. 78¹ §1 Kodeksu cywilnego" in text


def test_every_client_override_still_renders_as_company():
    from app.services.b2b_contract_generator.clause_override_content import (
        CLIENT_OVERRIDES,
    )

    for needles, _builder in CLIENT_OVERRIDES:
        for lang in ("pl", "en"):
            lines = _render(_req(lang, client_name=needles[0], **_COMPANY))
            assert _headings(lines)[-3:] == ["§ 12", "§ 13", "§ 14"], needles[0]
            assert _html(_req(lang, client_name=needles[0], **_COMPANY))


# ── Podgląd HTML = DOCX ─────────────────────────────────────────────────────


@pytest.mark.parametrize("lang", ["pl", "en"])
def test_html_preview_mirrors_the_company_docx(lang):
    out = _html(_req(lang, **_COMPANY))
    assert "{{" not in out and "{%" not in out
    assert "<h2>§ 14</h2>" in out
    label = "Osoba skierowana:" if lang == "pl" else "Designated Person:"
    assert f"<tr><td>{label}</td><td>Piotr Wiśniewski</td></tr>" in out
    assert out.index("<h2>§ 12</h2>") < out.index("<h2>§ 13</h2>")


def test_html_preview_escapes_the_assigned_person():
    out = _html(_req("pl", **{**_COMPANY, "assigned_person_name": "<b>X</b> & Y"}))
    assert "&lt;b&gt;X&lt;/b&gt; &amp; Y" in out


def test_html_sole_trader_has_no_company_rows():
    out = _html(_req("pl"))
    assert "Osoba skierowana" not in out
    assert "<h2>§ 14</h2>" not in out


def test_missing_anchor_fails_loudly_instead_of_issuing_a_jdg_contract():
    with pytest.raises(CompanyVariantError):
        apply_company_variant_html(
            "<h1>Umowa</h1><p>bez paragrafów</p>", language="pl", assigned_person="X"
        )
    doc = _docx.Document()
    doc.add_paragraph("Umowa bez paragrafów")
    with pytest.raises(CompanyVariantError):
        apply_company_variant_docx(doc, language="pl", assigned_person="X")


# ── Formatowanie danych ─────────────────────────────────────────────────────


@pytest.mark.parametrize(
    "raw, pl, en",
    [
        ("5000", "5.000,00 zł", "PLN 5,000.00"),
        ("5 000,00", "5.000,00 zł", "PLN 5,000.00"),
        ("1.360.000,00", "1.360.000,00 zł", "PLN 1,360,000.00"),
        ("1360000,5 zł", "1.360.000,50 zł", "PLN 1,360,000.50"),
        ("10 000 EUR", "10 000 EUR", "10 000 EUR"),
        ("", None, None),
    ],
)
def test_share_capital_format(raw, pl, en):
    assert format_share_capital(raw, "pl") == pl
    assert format_share_capital(raw, "en") == en


def test_representation_en():
    assert (
        representation_en("Jan Kowalski", "Członek Zarządu", "m")
        == "Mr Jan Kowalski – Member of the Management Board"
    )
    assert representation_en("Jan Kowalski", "Dyrektor", "m") == (
        "Mr Jan Kowalski – Dyrektor"
    )
    assert representation_en("", "Prezes Zarządu", "m") is None


# ── Numeracja w dokumentach pochodnych ──────────────────────────────────────


def test_derived_documents_cite_shifted_paragraphs_for_company_contracts():
    """Aneks daty startu i wypowiedzenie cytują paragrafy umowy bazowej —
    w umowie spółki rozwiązanie to § 13, a data startu § 14."""
    assert version_for_variant("jdg") == CURRENT_VERSION
    assert version_for_variant("company") == COMPANY_VERSION
    assert len(COMPANY_VERSION) <= 16  # kolumna VARCHAR(16)
    base, company = refs_for(CURRENT_VERSION), refs_for(COMPANY_VERSION)
    assert (base.start_paragraph, company.start_paragraph) == (
        "§ 13 ust. 2",
        "§ 14 ust. 2",
    )
    assert company.notice_paragraph == "§ 13 ust. 2 pkt 2"
    assert company.non_compete_paragraph == base.non_compete_paragraph == "§ 10 ust. 1"


# ── Odpis KRS ───────────────────────────────────────────────────────────────


def test_krs_details_parse_seat_capital_and_board():
    data = json.loads((FIXTURES / "odpis_aktualny.json").read_text(encoding="utf-8"))
    details = registry_lookup._parse_krs_company_details(data)
    assert details["seat"] == "Warszawa"
    assert details["share_capital"] == "1.360.000,00"
    # Publiczne API maskuje imiona i nazwiska — nie wymyślamy ich.
    assert [r["name"] for r in details["representatives"]] == [None, None]
    assert [r["function"] for r in details["representatives"]] == [
        "Prezes Zarządu",
        "Członek Zarządu",
    ]
    assert details["representation_method"].startswith("W przypadku zarządu")


def test_registry_court_comes_from_the_last_court_entry_not_system():
    data = json.loads((FIXTURES / "odpis_pelny.json").read_text(encoding="utf-8"))
    assert registry_lookup._registry_court_from_full(data) == (
        "Sąd Rejonowy dla m.st. Warszawy w Warszawie, XIV Wydział Gospodarczy "
        "Krajowego Rejestru Sądowego"
    )


@pytest.mark.parametrize(
    "raw, expected",
    [
        (
            "SĄD  REJONOWY  GDAŃSK-PÓŁNOC W GDAŃSKU, VII  WYDZIAŁ GOSPODARCZY",
            "Sąd Rejonowy Gdańsk-Północ w Gdańsku, VII Wydział Gospodarczy",
        ),
        (
            "SĄD REJONOWY DLA M. ST. WARSZAWY W WARSZAWIE, XIII WYDZIAŁ",
            "Sąd Rejonowy dla m.st. Warszawy w Warszawie, XIII Wydział",
        ),
        ("BIELSKO-BIAŁA", "Bielsko-Biała"),
    ],
)
def test_title_case_pl(raw, expected):
    assert registry_lookup.title_case_pl(raw) == expected


async def test_lookup_adds_krs_details_only_for_companies(monkeypatch):
    async def fake_base(nip=None, krs=None):
        return {"name": "X", "krs": "0000123456", "entity_type": "company"}

    async def fake_details(krs):
        return {"seat": "Kraków", "registry_court": "Sąd Rejonowy"}

    monkeypatch.setattr(registry_lookup, "_lookup_company_base", fake_base)
    monkeypatch.setattr(registry_lookup, "lookup_krs_company_details", fake_details)
    result = await registry_lookup.lookup_company(nip="1234567890")
    assert result["seat"] == "Kraków" and result["registry_court"] == "Sąd Rejonowy"

    async def fake_jdg(nip=None, krs=None):
        return {"name": "Y", "krs": None, "entity_type": "sole_trader"}

    monkeypatch.setattr(registry_lookup, "_lookup_company_base", fake_jdg)
    assert "seat" not in await registry_lookup.lookup_company(nip="1234567890")


# ── Endpoint: osoba skierowana z rekrutacji ─────────────────────────────────


async def _seed_candidate_in_job(created_by: int) -> tuple[int, int, str]:
    from app.core.database import AsyncSessionLocal
    from app.models.candidate import Candidate
    from app.models.recruitment_pipeline import CandidateStage, PipelineStage
    from app.models.client import Client
    from app.models.job import Job, JobStatus

    unique = uuid.uuid4().hex[:8]
    async with AsyncSessionLocal() as db:
        client = Client(name=f"Company Variant Client {unique}")
        candidate = Candidate(
            name="Piotr",
            lastname=f"Skierowany-{unique}",
            email=f"company-variant-{unique}@example.com",
            created_by=created_by,
        )
        db.add_all([client, candidate])
        await db.flush()
        job = Job(
            title=f"Tester {unique}", client_id=client.id, status=JobStatus.published
        )
        db.add(job)
        await db.flush()
        db.add(
            CandidateStage(
                candidate_id=candidate.id,
                job_id=job.id,
                stage=PipelineStage.onboarding,
                moved_by=created_by,
            )
        )
        await db.commit()
        return candidate.id, job.id, f"Piotr Skierowany-{unique}"


async def test_render_takes_assigned_person_from_the_recruitment(
    app_client, app_auth_headers
):
    from sqlalchemy import select

    from app.core.database import AsyncSessionLocal
    from app.models.b2b_generated_contract import B2BGeneratedContract
    from tests.test_b2b_generated_contract_status import _admin_user_id

    admin_id = await _admin_user_id(app_client)
    candidate_id, job_id, full_name = await _seed_candidate_in_job(admin_id)
    seq = 300000 + (uuid.uuid4().int % 500000)
    number = f"{seq}/2026"
    payload = {
        **{k: v for k, v in _BASE.items() if k not in ("signing_date", "start_date")},
        **_COMPANY,
        "language": "pl",
        "candidate_id": candidate_id,
        "job_id": job_id,
        "contract_number": number,
        "signing_date": "2026-09-28",
        "start_date": "2026-10-01",
        # Formularz podał kogoś innego — dokument i tak wskazuje kandydata.
        "assigned_person_name": "Ktoś Inny",
    }
    resp = await app_client.post(
        "/api/b2b-generator/render?format=docx", headers=app_auth_headers, json=payload
    )
    assert resp.status_code == 200, resp.text
    lines = _texts(resp.content)
    assert f"Osoba skierowana:|{full_name}" in lines
    assert "Ktoś Inny" not in "\n".join(lines)

    async with AsyncSessionLocal() as db:
        row = await db.scalar(
            select(B2BGeneratedContract).where(
                B2BGeneratedContract.contract_number == number
            )
        )
        assert row.template_version == COMPANY_VERSION
        assert row.render_payload["assigned_person_name"] == full_name
        assert row.render_payload["contract_variant"] == "company"
        row_id = row.id

    # Ponowne pobranie z zapisanego payloadu = ten sam dokument.
    again = await app_client.get(
        f"/api/b2b-generator/generated/{row_id}/docx", headers=app_auth_headers
    )
    assert again.status_code == 200, again.text
    assert f"Osoba skierowana:|{full_name}" in _texts(again.content)

    # Poprawka na JDG pod tym samym numerem cofa wersję wzoru.
    fixed = await app_client.post(
        f"/api/b2b-generator/generated/{row_id}/rerender",
        headers=app_auth_headers,
        json={**payload, "contract_variant": "jdg"},
    )
    assert fixed.status_code == 200, fixed.text
    assert "Osoba skierowana" not in "\n".join(_texts(fixed.content))
    async with AsyncSessionLocal() as db:
        row = await db.get(B2BGeneratedContract, row_id)
        assert row.template_version == CURRENT_VERSION


async def test_html_preview_endpoint_renders_company_variant(
    app_client, app_auth_headers
):
    payload = {
        **{k: v for k, v in _BASE.items() if k not in ("signing_date", "start_date")},
        **_COMPANY,
        "language": "pl",
        "signing_date": "2026-09-28",
        "start_date": "2026-10-01",
    }
    resp = await app_client.post(
        "/api/b2b-generator/render?format=html", headers=app_auth_headers, json=payload
    )
    assert resp.status_code == 200, resp.text
    html = resp.json()["html"]
    assert "Osoby skierowane do realizacji Usług" in html
    assert "<tr><td>Osoba skierowana:</td><td>Piotr Wiśniewski</td></tr>" in html
