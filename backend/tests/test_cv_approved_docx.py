import hashlib
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import Mock

from docx import Document
import pytest

from app.services.cv_approved_docx import ApprovedDocxError, render_approved_docx
from app.services.cv_document_assets import (
    default_template,
    generated_assets,
    CvAssetsError,
)
from app.services.html_sanitizer import sanitize_cv_html
from app.schemas.candidate_stage_cv import CVBrandedFinalize
from app.models.cv_document_version import CvDocumentVersion
from tests.test_cv_document_versions_unit import context
from app.api import candidate_stage_cv as api


def test_editor_text_bolding_and_nested_order_are_preserved():
    content = (
        "<h1>Audyt Testowy</h1><div><span>Alpha</span> <span>Beta</span></div>"
        "<p>Python <strong>FastAPI</strong> oraz <em>pytest</em>. Bez AWS.</p>"
        "<ul><li><p>Przed</p><ul><li><p>Zagnieżdżone</p></li></ul><p>Po</p></li></ul>"
    )
    doc = Document(BytesIO(render_approved_docx(content, default_template())))
    paragraphs = [p.text for p in doc.paragraphs]
    assert "Alpha Beta" in paragraphs
    assert (
        paragraphs.index("• Przed")
        < paragraphs.index("• Zagnieżdżone")
        < paragraphs.index("Po")
    )
    narrative = next(p for p in doc.paragraphs if p.text.startswith("Python"))
    assert narrative.text == "Python FastAPI oraz pytest. Bez AWS."
    assert [r.text for r in narrative.runs if r.bold] == ["FastAPI"]
    assert [r.text for r in narrative.runs if r.italic] == ["pytest"]


def test_full_document_import_does_not_print_css_or_scripts():
    clean = sanitize_cv_html(
        "<!doctype html><html><head><style>SECRET CSS</style><title>Title</title></head>"
        "<body><h1>CV</h1><script>BAD CODE</script><p>A &amp; B <strong>Python</strong></p></body></html>"
    )
    assert "SECRET" not in clean and "BAD" not in clean and "Title" not in clean
    doc = Document(BytesIO(render_approved_docx(clean, default_template())))
    assert [p.text for p in doc.paragraphs] == ["CV", "A & B Python"]


def test_table_text_and_cell_merges_are_preserved():
    doc = Document(
        BytesIO(
            render_approved_docx(
                '<table><tr><th colspan="2">Kontakt</th></tr>'
                "<tr><td>Telefon</td><td>123</td></tr></table>",
                default_template(),
            )
        )
    )
    table = doc.tables[0]
    assert table.cell(0, 0)._tc is table.cell(0, 1)._tc
    assert table.cell(1, 0).text == "Telefon"
    assert table.cell(1, 1).text == "123"


def test_consent_style_survives_editor_serialization_without_rewriting_text():
    clean = sanitize_cv_html(
        '<h1>Audyt Testowy</h1><p data-cv-section="rodo">Moja treść zgody.</p>'
    )
    doc = Document(BytesIO(render_approved_docx(clean, default_template())))
    assert doc.paragraphs[0].text == "Audyt Testowy"
    assert doc.paragraphs[0].style.font.all_caps is False
    paragraph = doc.paragraphs[1]
    assert paragraph.text == "Moja treść zgody."
    assert paragraph.style.font.size.pt == 7
    assert paragraph.style.paragraph_format.keep_together


def test_overlapping_table_spans_are_rejected_instead_of_losing_text():
    with pytest.raises(ApprovedDocxError, match="Nakładające"):
        render_approved_docx(
            '<table><tr><td>A</td><td rowspan="2">B</td></tr>'
            '<tr><td colspan="2">C</td></tr></table>',
            default_template(),
        )


def test_remote_image_is_rejected_without_any_network_fetch():
    with pytest.raises(ApprovedDocxError, match="Nieobsługiwany obraz"):
        render_approved_docx(
            '<p><img src="https://example.com/private.png"></p>', default_template()
        )


def test_missing_selected_consent_blocks_selection(monkeypatch):
    from app.services import object_storage

    download = Mock(side_effect=OSError("PRIVATE DETAILS"))
    monkeypatch.setattr(object_storage, "download_cv", download)
    source = SimpleNamespace(
        render_payload={"consent_screenshot": {"storage_key": "synthetic/key"}}
    )
    with pytest.raises(CvAssetsError) as exc:
        generated_assets(source)
    assert "PRIVATE" not in str(exc.value)


async def test_approval_freezes_actual_edits_and_download_never_rerenders(monkeypatch):
    from app.services import cv_approval_review
    from unittest.mock import AsyncMock

    monkeypatch.setattr(
        cv_approval_review,
        "review_for_approval",
        AsyncMock(return_value={"status": "verified"}),
    )
    csv, db, user, _, loader = context(monkeypatch)
    result = await api.finalize_branded_cv(
        2,
        CVBrandedFinalize(
            expected_revision=5,
            content_html="<p>Poprawione <b>Python</b>, bez dopisanej kompetencji.</p>",
        ),
        user,
        db,
    )
    version = next(
        c.args[0]
        for c in db.add.call_args_list
        if isinstance(c.args[0], CvDocumentVersion)
    )
    assert result.document_version_id == version.id
    assert hashlib.sha256(version.docx_content).hexdigest() == version.docx_sha256
    doc = Document(BytesIO(version.docx_content))
    assert [p.text for p in doc.paragraphs] == [
        "Poprawione Python, bez dopisanej kompetencji."
    ]
    assert version.render_metadata["renderer_version"] == "approved-html-1"
    csv.branded_draft_html = "<p>Późniejsze poprawki</p>"
    csv.branded_version = 2
    csv.branded_status = "draft"
    db.scalar.return_value = version
    monkeypatch.setattr(
        api, "render_approved_docx", Mock(side_effect=AssertionError("must not render"))
    )
    first = await api.download_approved_docx(2, 1, user, db)
    second = await api.download_approved_docx(2, 1, user, db)
    assert first.body == second.body == version.docx_content
    assert first.headers["etag"] == '"' + version.docx_sha256 + '"'
    loader.assert_awaited_with(db, 2, user, read_access=True)


async def test_invalid_asset_prevents_approval_commit(monkeypatch):
    csv, db, user, _, _ = context(monkeypatch)
    from fastapi import HTTPException

    with pytest.raises(HTTPException) as exc:
        await api.finalize_branded_cv(
            2,
            CVBrandedFinalize(
                expected_revision=5,
                content_html='<p><img src="https://example.com/private.png"></p>',
            ),
            user,
            db,
        )
    assert exc.value.status_code == 422
    assert csv.branded_status == "draft"
    db.commit.assert_not_awaited()


async def test_preview_uses_unsaved_editor_text_without_saving_or_approving(
    monkeypatch,
):
    csv, db, user, _, _ = context(monkeypatch)
    previous_html = csv.branded_draft_html
    result = await api.preview_branded_docx(
        2,
        CVBrandedFinalize(
            expected_revision=5, content_html="<p>Jeszcze niezapisana poprawka</p>"
        ),
        user,
        db,
    )
    assert (
        Document(BytesIO(result.body)).paragraphs[0].text
        == "Jeszcze niezapisana poprawka"
    )
    assert csv.branded_draft_html == previous_html
    assert csv.branded_status == "draft" and csv.edit_revision == 5
    db.commit.assert_not_awaited()
