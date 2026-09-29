"""Generator Umów B2B — § 13 ust. 6 w umowie dla BIK (ticket 8, 09.2026).

Umowa BIK od zawsze dostawała Załącznik nr 4 („Szczególne wymagania…”), ale
lista załączników w § 13 ust. 6 kończyła się na Załączniku nr 3 — umowa
wymieniała trzy integralne załączniki, a miała cztery. Pozostali Klienci
zostają przy treści szablonu.
"""

from __future__ import annotations

import io
from datetime import date
from pathlib import Path

import docx as _docx
import pytest

from app.api.contract_templates import render_contract_template
from app.schemas.b2b_contract_generator import B2BRenderRequest
from app.services.b2b_contract_generator.clause_overrides import (
    apply_ops_html_counted,
    resolve_override,
)
from app.services.b2b_contract_generator.company_variant import (
    apply_company_variant_html,
)
from app.services.b2b_contract_generator.docx_renderer import render_from_context
from app.services.b2b_contract_generator.render_context import build_render_context

TEMPLATE_DIR = Path(__file__).resolve().parents[1] / "app" / "templates" / "contract"

BIK = "Biuro Informacji Kredytowej S.A."

EXPECTED = {
    "pl": (
        "Integralną część niniejszej Umowy stanowią następujące załączniki: "
        "Załącznik nr 1 – Deklaracja Poufności, Załącznik nr 2 – Umowa "
        "Powierzenia Przetwarzania Danych Osobowych (DPA), Załącznik nr 3 – "
        "Klient Projektu, Załącznik nr 4 – Szczególne wymagania dotyczące "
        "realizacji usług na rzecz Klienta Projektu – Biuro Informacji "
        "Kredytowej S.A."
    ),
    "en": (
        "The following appendices form an integral part of this Agreement: "
        "Appendix 1 – Confidentiality Agreement, Appendix No. 2 – Data "
        "Processing Agreement (DPA), Appendix No. 3 – Project Customer, "
        "Appendix No. 4 – Special requirements for the provision of services "
        "to the Project Customer – Biuro Informacji Kredytowej S.A."
    ),
}

TEMPLATE_TEXT = {
    "pl": (
        "Integralną część niniejszej Umowy stanowią następujące załączniki: "
        "Załącznik nr 1 – Deklaracja Poufności, Załącznik nr 2 – Umowa "
        "Powierzenia Przetwarzania Danych Osobowych (DPA), Załącznik nr 3 – "
        "Klient Projektu."
    ),
    "en": (
        "The following appendices form an integral part of this Agreement: "
        "Appendix 1 – Confidentiality Agreement, Appendix No. 2 – Data "
        "Processing Agreement (DPA), Appendix No. 3 – Project Customer."
    ),
}

_BASE = dict(
    gender="m",
    partner_name="Jan Kowalski",
    partner_instrumental="Janem Kowalskim",
    partner_legal_name="JK Soft",
    partner_business_address="ul. Prosta 1, 00-001 Warszawa",
    partner_nip="1234567890",
    partner_regon="123456789",
    partner_email="biuro@jksoft.pl",
    partner_phone="+48 600 100 200",
    project_city="Warszawa",
    project_description="Rozwój platformy.",
    contract_number="1600/2026",
    signing_date=date(2026, 9, 28),
    start_date=date(2026, 10, 1),
    rate_candidate=150,
)

_COMPANY = dict(
    contract_variant="company",
    gender="k",
    partner_legal_name="JK Soft Sp. z o.o.",
    partner_krs="0000123456",
    partner_seat="Warszawa",
    partner_seat_locative="w Warszawie",
    partner_registry_court=(
        "Sąd Rejonowy dla m.st. Warszawy w Warszawie, XII Wydział Gospodarczy "
        "Krajowego Rejestru Sądowego"
    ),
    partner_share_capital="5000",
    partner_representative_name="Anna Nowak",
    partner_representative_function="Prezes Zarządu",
    partner_representation="Panią Annę Nowak – Prezes Zarządu",
    assigned_person_name="Piotr Wiśniewski",
)


def _req(lang: str, client_name: str, **extra) -> B2BRenderRequest:
    return B2BRenderRequest(
        **{**_BASE, "language": lang, "client_name": client_name, **extra}
    )


def _docx_paragraphs(req: B2BRenderRequest) -> list[str]:
    data = render_from_context(build_render_context(req, None), language=req.language)
    return [p.text for p in _docx.Document(io.BytesIO(data)).paragraphs]


def _html(req: B2BRenderRequest) -> str:
    lang = req.language
    ctx = build_render_context(req, None)
    out = render_contract_template(
        (TEMPLATE_DIR / f"umowa_b2b_{lang}.html").read_text(encoding="utf-8"), ctx
    )
    _key, ops = resolve_override(req.client_name, lang)
    if ops:
        out, applied = apply_ops_html_counted(out, ops)
        assert applied == len(ops)
    if ctx["b2b"]["is_company"]:
        out = apply_company_variant_html(
            out, language=lang, assigned_person=ctx["b2b"]["assigned_person"]
        )
    return out


@pytest.mark.parametrize("lang", ["pl", "en"])
@pytest.mark.parametrize("variant", [{}, _COMPANY], ids=["jdg", "company"])
def test_bik_contract_lists_appendix_4(lang, variant):
    req = _req(lang, BIK, **variant)
    paragraphs = _docx_paragraphs(req)
    assert EXPECTED[lang] in paragraphs
    assert TEMPLATE_TEXT[lang] not in paragraphs
    # Załącznik, do którego odsyła nowa treść, nadal jest w dokumencie.
    appendix = (
        "Załącznik nr 4 - Szczególne" if lang == "pl" else "Appendix No. 4 - Special"
    )
    assert any(p.startswith(appendix) for p in paragraphs)
    html = _html(req)
    assert f"<p>{EXPECTED[lang]}</p>" in html
    assert TEMPLATE_TEXT[lang] not in html


@pytest.mark.parametrize("lang", ["pl", "en"])
@pytest.mark.parametrize(
    "client", ["Acme", "Credit Agricole Bank Polska S.A.", "Alior Bank S.A."]
)
def test_other_clients_keep_template_attachment_list(lang, client):
    req = _req(lang, client)
    assert TEMPLATE_TEXT[lang] in _docx_paragraphs(req)
    assert EXPECTED[lang] not in "\n".join(_docx_paragraphs(req))
    assert TEMPLATE_TEXT[lang] in _html(req)


@pytest.mark.parametrize("lang", ["pl", "en"])
def test_bik_paragraph_keeps_template_formatting(lang):
    """Podmieniona zostaje treść akapitu, nie jego styl z szablonu."""
    req = _req(lang, BIK)
    data = render_from_context(build_render_context(req, None), language=lang)
    doc = _docx.Document(io.BytesIO(data))
    para = next(p for p in doc.paragraphs if p.text == EXPECTED[lang])
    tpl = _docx.Document(str(TEMPLATE_DIR / f"umowa_b2b_{lang}.docx"))
    orig = next(p for p in tpl.paragraphs if p.text == TEMPLATE_TEXT[lang])
    assert para.style.name == orig.style.name
    assert para.alignment == orig.alignment
    assert para.runs[0].font.size == orig.runs[0].font.size
