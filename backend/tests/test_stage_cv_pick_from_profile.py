"""Screening: wybór gotowego CV z profilu kandydata i edycja (09.10.2026).

Do tej daty CV firmowe etapu dało się zacząć wyłącznie od CV z generatora dla
TEJ rekrutacji. Pomiar na produkcji: 34 z 50 osób w screeningu ma gotowe CV
jako plik Word „…B2B…”, 9 — CV z generatora z innej rekrutacji. Testy pilnują
obu dróg, ich odmów oraz tego, że auto-CV po ruchu na „Zweryfikowany” nie
przykrywa CV wybranego wcześniej.
"""

from __future__ import annotations

import uuid
from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document
from httpx import AsyncClient
from sqlalchemy import select

import app.models  # noqa: F401
from app.core.database import AsyncSessionLocal
from app.models.activity import Activity
from app.models.candidate import Candidate
from app.models.candidate_document import CandidateDocument, CandidateDocumentKind
from app.models.candidate_stage_cv import CandidateStageCV
from app.models.client import Client
from app.models.cv_generated_document import CvGeneratedDocument
from app.models.job import Job, JobStatus
from app.models.recruitment_pipeline import CandidateStage, PipelineStage
from app.models.user import User, UserRole
from app.services import cv_auto_generate as auto
from app.services.candidate_stage_cv_service import create_original_cv_snapshot

MARKER = "Zbudowała platformę wdrożeniową dla zespołów produktowych"


def _docx(marker: str = MARKER) -> bytes:
    doc = Document()
    doc.add_paragraph("Inżynier platformy – Ewa Fikcyjna")
    doc.add_heading("DOŚWIADCZENIE", level=1)
    doc.add_paragraph("01.2020 - 03.2024")
    doc.add_paragraph("Nazwa firmy: Firma Przykładowa")
    doc.add_paragraph("Stanowisko: Platform Engineer")
    doc.add_paragraph(marker, style="List Bullet")
    doc.add_paragraph(
        "Dziesięć lat w zespołach wdrożeniowych, ostatnio jako liderka zespołu "
        "platformy w bankowości; prowadziła migrację usług do chmury i dyżury."
    )
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


async def _seed(*, with_cv_row: bool = True) -> dict:
    tag = uuid.uuid4().hex[:8]
    async with AsyncSessionLocal() as db:
        client = Client(name=f"Klient wybór CV {tag}")
        other_client = Client(name=f"Inny klient wybór CV {tag}")
        db.add_all([client, other_client])
        await db.flush()
        candidate = Candidate(name="Ewa", lastname=f"Fikcyjna{tag}")
        job = Job(title=f"Platform Engineer {tag}", client_id=client.id)
        other_job = Job(title=f"DevOps {tag}", client_id=other_client.id)
        db.add_all([candidate, job, other_job])
        await db.flush()
        stage = CandidateStage(
            candidate_id=candidate.id, job_id=job.id, stage=PipelineStage.screening
        )
        content = _docx()
        document = CandidateDocument(
            candidate_id=candidate.id,
            filename=f"CV_B2B_Ewa_Fikcyjna_{tag}.docx",
            file_content=content,
            document_kind=CandidateDocumentKind.cv,
        )
        db.add_all([stage, document])
        await db.flush()
        if with_cv_row:
            await create_original_cv_snapshot(db, stage)
        await db.commit()
        return {
            "stage_id": stage.id,
            "candidate_id": candidate.id,
            "job_id": job.id,
            "other_job_id": other_job.id,
            "other_client_id": other_client.id,
            "document_id": document.id,
            "document_bytes": content,
            "tag": tag,
        }


async def _generated(candidate_id: int, job_id: int | None, **extra) -> int:
    async with AsyncSessionLocal() as db:
        row = CvGeneratedDocument(
            candidate_id=candidate_id,
            job_id=job_id,
            candidate_name="Ewa Fikcyjna",
            filename="ZOB-1_CV_dla_innego_klienta.docx",
            status="ready",
            render_payload={
                "name": "Ewa Fikcyjna",
                "why_points": ["Prowadziła migrację do chmury"],
            },
            **extra,
        )
        db.add(row)
        await db.commit()
        return row.id


async def _csv(stage_id: int) -> CandidateStageCV | None:
    async with AsyncSessionLocal() as db:
        return await db.scalar(
            select(CandidateStageCV).where(
                CandidateStageCV.candidate_stage_id == stage_id
            )
        )


def _base(stage_id: int) -> str:
    return f"/api/candidates/stages/{stage_id}/cv/branded"


async def _revision(client: AsyncClient, headers: dict, stage_id: int) -> int:
    response = await client.get(_base(stage_id), headers=headers)
    return response.json()["edit_revision"] if response.status_code == 200 else 0


@pytest.mark.asyncio
async def test_word_file_from_the_profile_becomes_an_editable_stage_cv(
    app_client: AsyncClient, app_auth_headers: dict, monkeypatch
):
    from unittest.mock import AsyncMock

    from app.services import cv_approval_review

    monkeypatch.setattr(
        cv_approval_review,
        "review_for_approval",
        AsyncMock(return_value={"status": "unverified", "method": "test"}),
    )
    world = await _seed()
    base = _base(world["stage_id"])
    chosen = await app_client.post(
        base + "/select-document",
        headers=app_auth_headers,
        json={
            "expected_revision": await _revision(
                app_client, app_auth_headers, world["stage_id"]
            ),
            "document_id": world["document_id"],
        },
    )
    assert chosen.status_code == 200, chosen.text
    body = chosen.json()
    assert (body["status"], body["source"]) == ("draft", "document")
    assert body["generated_document_id"] is None and body["from_generator"] is False
    assert MARKER in body["content_html"]
    assert 'data-cv-section="role"' in body["content_html"]

    edited = body["content_html"].replace("Zbudowała", "Zaprojektowała")
    saved = await app_client.patch(
        base,
        headers=app_auth_headers,
        json={"expected_revision": body["edit_revision"], "content_html": edited},
    )
    assert saved.status_code == 200, saved.text
    approved = await app_client.post(
        base + "/finalize",
        headers=app_auth_headers,
        json={
            "expected_revision": saved.json()["edit_revision"],
            "content_html": edited,
        },
    )
    assert approved.status_code == 200, approved.text
    docx = await app_client.get(base + "/versions/1/docx", headers=app_auth_headers)
    assert docx.status_code == 200, docx.text
    text = " ".join(p.text for p in Document(BytesIO(docx.content)).paragraphs)
    assert "Zaprojektowała platformę wdrożeniową" in text

    csv = await _csv(world["stage_id"])
    assert csv.branded_render_metadata["source"] == "document"
    assert csv.branded_render_metadata["source_document_id"] == world["document_id"]
    # Nazwa pliku liczy się z tej rekrutacji, nie z nazwy pliku w profilu.
    assert world["tag"] in csv.branded_docx_filename
    async with AsyncSessionLocal() as db:
        document = await db.get(CandidateDocument, world["document_id"])
        await db.refresh(document, attribute_names=["file_content"])
        assert bytes(document.file_content) == world["document_bytes"]
        activity = await db.scalar(
            select(Activity).where(
                Activity.entity_type == "candidate_stage_cv",
                Activity.entity_id == csv.id,
                Activity.action == "branded_cv_imported_from_document",
            )
        )
        assert activity.details["source_document_id"] == world["document_id"]


@pytest.mark.asyncio
async def test_stage_without_a_cv_row_gets_one_when_a_cv_is_chosen(
    app_client: AsyncClient, app_auth_headers: dict
):
    """Etapy z importu Traffita nie mają wiersza CV — wybór go zakłada."""
    world = await _seed(with_cv_row=False)
    assert await _csv(world["stage_id"]) is None
    chosen = await app_client.post(
        _base(world["stage_id"]) + "/select-document",
        headers=app_auth_headers,
        json={"expected_revision": 0, "document_id": world["document_id"]},
    )
    assert chosen.status_code == 200, chosen.text
    csv = await _csv(world["stage_id"])
    assert csv is not None and csv.branded_status == "draft"
    assert (csv.candidate_id, csv.job_id) == (world["candidate_id"], world["job_id"])


@pytest.mark.asyncio
@pytest.mark.parametrize("source_job", ["other", None])
async def test_cv_generated_for_another_recruitment_is_a_detached_copy(
    app_client: AsyncClient, app_auth_headers: dict, source_job
):
    world = await _seed()
    generated_id = await _generated(
        world["candidate_id"],
        world["other_job_id"] if source_job else None,
        client_id=world["other_client_id"],
        consent_content=b"zrzut zgody dla innego klienta",
        central_policy={"requires_rodo_consent_block": True},
    )
    chosen = await app_client.post(
        _base(world["stage_id"]) + "/select-generated",
        headers=app_auth_headers,
        json={
            "expected_revision": await _revision(
                app_client, app_auth_headers, world["stage_id"]
            ),
            "generated_document_id": generated_id,
        },
    )
    assert chosen.status_code == 200, chosen.text
    body = chosen.json()
    assert (body["status"], body["source"]) == ("draft", "generator")
    assert body["generated_document_id"] is None
    assert "Prowadziła migrację do chmury" in body["content_html"]

    csv = await _csv(world["stage_id"])
    metadata = csv.branded_render_metadata
    assert metadata["source"] == "generated_other_job"
    assert metadata["source_generated_id"] == generated_id
    # Zgoda i jej wymóg należą do klienta źródłowego — nie jadą z kopią.
    assert csv.branded_consent_content is None
    assert "consent_required" not in metadata
    assert csv.branded_docx_filename != "ZOB-1_CV_dla_innego_klienta.docx"
    assert world["tag"] in csv.branded_docx_filename


@pytest.mark.asyncio
async def test_cv_generated_for_this_recruitment_keeps_the_generator_link(
    app_client: AsyncClient, app_auth_headers: dict
):
    world = await _seed()
    generated_id = await _generated(world["candidate_id"], world["job_id"])
    chosen = await app_client.post(
        _base(world["stage_id"]) + "/select-generated",
        headers=app_auth_headers,
        json={
            "expected_revision": await _revision(
                app_client, app_auth_headers, world["stage_id"]
            ),
            "generated_document_id": generated_id,
        },
    )
    assert chosen.status_code == 200, chosen.text
    assert chosen.json()["generated_document_id"] == generated_id
    assert chosen.json()["source"] == "generator"
    csv = await _csv(world["stage_id"])
    assert csv.branded_docx_filename == "ZOB-1_CV_dla_innego_klienta.docx"


@pytest.mark.asyncio
async def test_another_candidates_cv_cannot_be_chosen(
    app_client: AsyncClient, app_auth_headers: dict
):
    world = await _seed()
    stranger = await _seed()
    generated_id = await _generated(stranger["candidate_id"], world["job_id"])
    revision = await _revision(app_client, app_auth_headers, world["stage_id"])
    by_generated = await app_client.post(
        _base(world["stage_id"]) + "/select-generated",
        headers=app_auth_headers,
        json={"expected_revision": revision, "generated_document_id": generated_id},
    )
    by_document = await app_client.post(
        _base(world["stage_id"]) + "/select-document",
        headers=app_auth_headers,
        json={"expected_revision": revision, "document_id": stranger["document_id"]},
    )
    assert (by_generated.status_code, by_document.status_code) == (404, 404)
    assert (await _csv(world["stage_id"])).branded_status == "none"


@pytest.mark.asyncio
async def test_file_that_is_not_word_is_refused_with_a_polish_sentence(
    app_client: AsyncClient, app_auth_headers: dict
):
    world = await _seed()
    async with AsyncSessionLocal() as db:
        pdf = CandidateDocument(
            candidate_id=world["candidate_id"],
            filename="CV_B2B_skan.pdf",
            file_content=b"%PDF-1.7 fake",
            document_kind=CandidateDocumentKind.cv,
        )
        db.add(pdf)
        await db.commit()
        pdf_id = pdf.id
    response = await app_client.post(
        _base(world["stage_id"]) + "/select-document",
        headers=app_auth_headers,
        json={
            "expected_revision": await _revision(
                app_client, app_auth_headers, world["stage_id"]
            ),
            "document_id": pdf_id,
        },
    )
    assert response.status_code == 422, response.text
    assert "tylko plik Word" in response.json()["detail"]
    assert (await _csv(world["stage_id"])).branded_status == "none"


@pytest.mark.asyncio
async def test_client_requiring_consent_screenshot_sends_the_recruiter_to_the_generator(
    app_client: AsyncClient, app_auth_headers: dict, monkeypatch
):
    """Kopii z profilu nie da się dołączyć zgody — pobranie byłoby zablokowane."""
    from app.services.cv_generator_b2b import client_rules

    async def _rule(db, client_id):
        return object()

    monkeypatch.setattr(client_rules, "resolve_client_rule", _rule)
    monkeypatch.setattr(
        client_rules,
        "snapshot_rule",
        lambda rule: SimpleNamespace(requires_rodo_consent_block=True),
    )
    world = await _seed()
    generated_id = await _generated(world["candidate_id"], world["other_job_id"])
    revision = await _revision(app_client, app_auth_headers, world["stage_id"])
    by_document = await app_client.post(
        _base(world["stage_id"]) + "/select-document",
        headers=app_auth_headers,
        json={"expected_revision": revision, "document_id": world["document_id"]},
    )
    by_generated = await app_client.post(
        _base(world["stage_id"]) + "/select-generated",
        headers=app_auth_headers,
        json={"expected_revision": revision, "generated_document_id": generated_id},
    )
    for response in (by_document, by_generated):
        assert response.status_code == 422, response.text
        assert response.json()["detail"]["code"] == "consent_client_needs_generator"
    assert (await _csv(world["stage_id"])).branded_status == "none"


@pytest.mark.asyncio
async def test_stale_revision_does_not_replace_a_draft(
    app_client: AsyncClient, app_auth_headers: dict
):
    world = await _seed()
    base = _base(world["stage_id"])
    revision = await _revision(app_client, app_auth_headers, world["stage_id"])
    payload = {"expected_revision": revision, "document_id": world["document_id"]}
    first = await app_client.post(
        base + "/select-document", headers=app_auth_headers, json=payload
    )
    second = await app_client.post(
        base + "/select-document", headers=app_auth_headers, json=payload
    )
    assert (first.status_code, second.status_code) == (200, 409)


@pytest.mark.asyncio
async def test_qc_reads_the_cv_chosen_from_the_profile(
    app_client: AsyncClient, app_auth_headers: dict
):
    """QC CV i przegląd DL czytają szkic pary — także ten z pliku Word."""
    from app.services import dz_review

    world = await _seed()
    chosen = await app_client.post(
        _base(world["stage_id"]) + "/select-document",
        headers=app_auth_headers,
        json={
            "expected_revision": await _revision(
                app_client, app_auth_headers, world["stage_id"]
            ),
            "document_id": world["document_id"],
        },
    )
    assert chosen.status_code == 200, chosen.text
    async with AsyncSessionLocal() as db:
        source = await dz_review._generated_cv(
            db, world["candidate_id"], world["job_id"]
        )
    assert source["source"] == "branded_draft"
    assert source["stage_id"] == world["stage_id"]
    assert MARKER in source["html"]


@pytest.mark.asyncio
async def test_automation_does_not_generate_over_a_cv_chosen_in_screening(
    monkeypatch,
):
    """Ruch na „Zweryfikowany” zakłada nowy wiersz etapu; auto-CV podpięte do
    niego przykryłoby CV wybrane w screeningu i kosztowało drugą generację."""
    from app.api import cv_generator_b2b as api

    tag = uuid.uuid4().hex[:8]
    async with AsyncSessionLocal() as db:
        user = User(
            email=f"cv-pick-{tag}@example.com",
            name="CV Pick",
            role=UserRole.recruiter,
            is_active=True,
        )
        client = Client(name=f"Klient auto pick {tag}")
        db.add_all([user, client])
        await db.flush()
        job = Job(
            title=f"Auto pick {tag}",
            client_id=client.id,
            status=JobStatus.published,
            recruiter_id=user.id,
        )
        candidate = Candidate(name="Auto", lastname=f"Pick{tag}")
        db.add_all([job, candidate])
        await db.flush()
        db.add(
            CandidateDocument(
                candidate_id=candidate.id, filename="cv.pdf", is_primary=True
            )
        )
        screening = CandidateStage(
            candidate_id=candidate.id, job_id=job.id, stage=PipelineStage.screening
        )
        verified = CandidateStage(
            candidate_id=candidate.id,
            job_id=job.id,
            stage=PipelineStage.verified,
            moved_by=user.id,
        )
        db.add_all([screening, verified])
        await db.flush()
        db.add(
            CandidateStageCV(
                candidate_stage_id=screening.id,
                candidate_id=candidate.id,
                job_id=job.id,
                branded_status="draft",
                branded_draft_html="<p>CV wybrane w screeningu</p>",
            )
        )
        await db.commit()
        ids = (job.id, verified.id, user.id)

    async def _must_not_run(db, **kwargs):
        raise AssertionError("auto-CV wygenerowało dokument mimo CV pary")

    monkeypatch.setattr(api, "enqueue_candidate_generation", _must_not_run)
    async with AsyncSessionLocal() as db:
        result = await auto._enqueue(db, stage_id=ids[1], user_id=ids[2])
    assert result is None
    async with AsyncSessionLocal() as db:
        events = (
            await db.scalars(
                select(Activity).where(
                    Activity.entity_type == auto.ACTIVITY_ENTITY,
                    Activity.entity_id == ids[0],
                )
            )
        ).all()
    assert [event.details["reason"] for event in events] == ["pair_cv_exists"]
