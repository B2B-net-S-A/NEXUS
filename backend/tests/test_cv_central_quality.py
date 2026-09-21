"""Quality fixes for central CV policies measured in production (21.09.2026).

* Content mode: every client defaults to "Pod rekrutację" (tailored), the
  recruiter may change it, "tailored" without a Champion falls back to
  polished with a notice (never a 422), the client cap wins.
* Upload: an uploaded Champion DOCX is used again (tailored share fell from
  49% to 12% because central mode discarded it).
* No forced ``why_points_max`` rendered as a client instruction.
* Startup never crashes on a catalogue/client mismatch.
* Legacy "Considered for" line is not printed when it repeats the header.
* The system prompt stays byte-identical across clients and positions.
"""

import ast
import json
import time
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from typing import get_args
from unittest.mock import AsyncMock

import pytest
from docx import Document
from fastapi import FastAPI, HTTPException
from httpx import ASGITransport, AsyncClient

from app.models.client_cv_rule import ClientCvRule
from app.services.cv_generator_b2b import central_policies as policies
from app.services.cv_generator_b2b.client_rules import (
    build_client_presentation_rules_block,
    snapshot_rule,
)
from app.services.cv_generator_b2b.presentation_title import considered_for_line

NORDEA = 11
ALIOR = 39


def _central_rule(client_id):
    policy = policies.policy_for(client_id)
    return ClientCvRule(
        client_id=client_id,
        version=1,
        managed_policy=policies.metadata(policy),
        **policies.recipe_for(policy),
    )


@pytest.fixture
def champion(monkeypatch):
    """``job_supports_tailored`` is True for jobs flagged ``has_champion``."""
    import app.services.champion_intake as intake
    from app.services.cv_generator_b2b import standalone_service

    monkeypatch.setattr(
        standalone_service,
        "champion_present",
        lambda job: bool(getattr(job, "has_champion", False)),
    )
    monkeypatch.setattr(intake, "enforce_operation", lambda *a, **kw: None)


def _job(has_champion=True):
    return SimpleNamespace(
        champion_profile={"any": "profile"} if has_champion else None,
        has_champion=has_champion,
    )


# ── Catalogue: everyone defaults to tailored, versions bumped ────────────


def test_every_catalogue_policy_and_the_standard_default_is_tailored():
    for entry in policies.catalog():
        assert entry["content_mode"] == "tailored", entry["key"]
        assert entry["version"] >= 2, entry["key"]  # republished on startup
    standard = policies.policy_for(None)
    assert standard["content_mode"] == "tailored"
    assert standard["version"] >= 2
    assert policies.metadata(standard)["content_mode"] == "tailored"


@pytest.mark.parametrize("client_id", [NORDEA, ALIOR, None, 999999])
def test_default_mode_is_tailored_with_a_champion(champion, client_id):
    policy = policies.metadata(policies.policy_for(client_id))
    assert policies.automatic_mode_decision(_job(), policy=policy) == (
        "tailored",
        None,
    )


def test_without_a_champion_the_default_falls_back_with_a_notice(champion):
    policy = policies.metadata(policies.policy_for(ALIOR))
    mode, notice = policies.automatic_mode_decision(_job(False), policy=policy)
    assert mode == "polished"
    assert notice == policies.MISSING_CHAMPION_NOTICE
    # An uploaded Champion is enough, with no recruitment at all.
    assert policies.automatic_mode_decision(
        None, uploaded_champion=True, policy=policy
    ) == ("tailored", None)


def test_client_cap_still_wins(champion):
    assert policies.automatic_mode_decision(_job(), "basic")[0] == "basic"
    assert policies.resolve_mode("tailored", _job(), "polished")[0] == "polished"


def test_requested_mode_is_honoured_not_locked(champion):
    assert policies.resolve_mode("polished", _job()) == ("polished", None)
    assert policies.resolve_mode("basic", _job()) == ("basic", None)
    assert policies.resolve_mode("tailored", _job()) == ("tailored", None)
    assert policies.resolve_mode("tailored", _job(False)) == (
        "polished",
        policies.MISSING_CHAMPION_NOTICE,
    )


def test_a_single_client_can_still_be_switched_by_the_catalogue(champion):
    assert policies.automatic_mode(_job(), policy={"content_mode": "polished"}) == (
        "polished"
    )


def test_job_predicate_rejects_incomplete_champions(monkeypatch, champion):
    import app.services.champion_intake as intake

    assert policies.job_supports_tailored(None) is False
    assert policies.job_supports_tailored(_job(False)) is False

    def refuse(*args, **kwargs):
        assert kwargs["force"] is True
        raise HTTPException(422, "incomplete")

    monkeypatch.setattr(intake, "enforce_operation", refuse)
    assert policies.job_supports_tailored(_job()) is False


def test_central_rule_does_not_lock_the_mode():
    rule = snapshot_rule(_central_rule(ALIOR))
    assert not rule.content_mode_locked


# ── No forced why_points_max rendered as a client instruction ────────────


@pytest.mark.parametrize("client_id", [NORDEA, ALIOR, 26, None])
def test_central_policy_renders_no_client_instruction(client_id):
    rule = snapshot_rule(_central_rule(client_id))
    assert rule.why_points_max is None
    assert build_client_presentation_rules_block(rule, "pl") == ""
    assert build_client_presentation_rules_block(rule, "en") == ""


# ── Startup resilience ───────────────────────────────────────────────────


def _catalogue_clients(**overrides):
    entries = iter(sorted(policies.catalog(), key=lambda p: p["client_id"]))

    async def scalar(statement):
        entry = next(entries)
        return SimpleNamespace(
            **{
                "id": entry["client_id"],
                "external_source": entry["external_source"],
                "external_id": entry["external_id"],
                "hidden": False,
                "merged_into_client_id": None,
                **overrides,
            }
        )

    return scalar


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "overrides",
    [
        {"external_id": "not-the-catalogue-id"},
        {"hidden": True},
        {"merged_into_client_id": 5},
    ],
)
async def test_synchronize_skips_non_matching_clients_instead_of_raising(overrides):
    db = AsyncMock()
    db.scalar.side_effect = _catalogue_clients(**overrides)
    assert await policies.synchronize(db) == 0
    db.add.assert_not_called()
    db.commit.assert_awaited_once()


def test_startup_synchronization_is_guarded():
    source = (Path(__file__).parents[1] / "app/main.py").read_text()
    tree = ast.parse(source)
    guarded = [
        node
        for node in ast.walk(tree)
        if isinstance(node, ast.Try)
        and "central_policies.synchronize" in ast.unparse(node)
    ]
    assert guarded, "central_policies.synchronize must run inside try/except"


# ── Legacy "Considered for" duplicate ────────────────────────────────────


def test_considered_for_line_skips_a_repeat_of_the_header():
    assert considered_for_line("Java Developer", "java  developer ") == ""
    assert considered_for_line("  ", "Java Developer") == ""
    assert considered_for_line("Tech Lead", "Java Developer") == "Tech Lead"


def _legacy_payload(considered_for):
    return {
        "name": "Jan Kowalski",
        "first_name": "Jan",
        "position": "Backend Developer",
        "considered_for": considered_for,
        "why_points": ["Python"],
        "education": [],
        "skills": [{"label": "Backend", "content": "Python"}],
        "certifications": [],
        "languages": ["Polski"],
        "experience": [
            {
                "dates": "01.2015 – 07.2019",
                "company": "Acme",
                "position": "Backend Developer",
                "responsibilities": ["API"],
                "technologies": ["Python"],
            }
        ],
        "language": "pl",
    }


@pytest.mark.parametrize(
    "considered_for,printed",
    [("backend developer", False), ("Tech Lead", True)],
)
def test_legacy_document_prints_considered_for_only_when_different(
    considered_for, printed
):
    from app.services.cv_generator_b2b import standalone_service as svc
    from app.services.cv_generator_b2b.html_export import render_interactive_html
    from app.services.cv_generator_b2b.public_view import build_public_payload

    payload = _legacy_payload(considered_for)
    document = Document(BytesIO(svc.rerender_docx_from_payload(payload)))
    text = "\n".join(p.text for p in document.paragraphs)
    assert ("Rozważany na stanowisko" in text) is printed
    public = build_public_payload(payload)
    assert (public["considered_for"] is not None) is printed
    html = render_interactive_html(public, [], document_only=True)
    assert ('class="considered"' in html) is printed


# ── System prompt stays cacheable ────────────────────────────────────────


def test_system_prompt_is_identical_across_clients_and_positions(monkeypatch):
    from app.services.cv_generator_b2b import standalone_service as svc
    from app.services.cv_generator_b2b.legacy_v7 import pipeline as legacy

    monkeypatch.delenv("CV_GENERATION_PIPELINE", raising=False)
    monkeypatch.setenv("CV_FINAL_REVIEW_ENABLED", "false")
    systems: dict[tuple, list[str]] = {}
    current: list[tuple] = []

    def analyze(user_content, request_id, system=None, **kwargs):
        systems.setdefault(current[0], []).append(system)
        return json.dumps(
            {
                "name": "Jan Kowalski",
                "first_name": "Jan",
                "position": "Backend Developer",
                "presentation_position": "Programista",
                "why_points": ["Python"],
                "education": [],
                "skills": [{"label": "Backend", "content": "Python"}],
                "certifications": [],
                "languages": ["Polski"],
                "experience": [],
            }
        )

    monkeypatch.setattr(
        legacy, "extract_text_from_file", lambda *a: "Jan Kowalski\nPython"
    )
    monkeypatch.setattr(legacy, "analyze_with_ai", analyze)
    monkeypatch.setattr(legacy, "render_cv_to_bytes", lambda *a, **k: b"DOCX")
    for language in ("pl", "en"):
        for client_id, title in ((NORDEA, "Java Developer"), (ALIOR, "QA"), (None, "")):
            current[:] = [(language,)]
            svc._run_generation_pipeline(
                cv_bytes=b"cv",
                cv_filename="cv.pdf",
                champion_dto=None,
                screening_notes_text="",
                language=language,
                blind_cv=False,
                request_id="cache-test",
                fallback_name="Jan Kowalski",
                started_at=time.time(),
                job_id=1,
                job_title=title,
                content_mode="polished",
                client_rule=snapshot_rule(_central_rule(client_id)),
            )
    for language, prompts in systems.items():
        assert len(prompts) == 3
        assert len(set(prompts)) == 1, language


# ── Upload API under central policies ────────────────────────────────────


def _docx(text):
    buffer = BytesIO()
    document = Document()
    document.add_paragraph(text)
    document.save(buffer)
    return buffer.getvalue()


_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


async def _upload(
    monkeypatch, *, content_mode, with_champion, cap=None, stage=False, extra=None
):
    from app.api import cv_generator_b2b as api
    from app.api import champion_intake as champion_api
    from app.models.user import User, UserRole
    from app.services.cv_generator_b2b import durable_jobs

    monkeypatch.setattr(policies.settings, "CV_CENTRAL_POLICIES_ENABLED", True)
    app = FastAPI()
    app.include_router(api.router)
    app.state.limiter = api.limiter
    db = AsyncMock()

    async def get(model, ident):
        if model is api.CandidateStage:
            return SimpleNamespace(id=3, candidate_id=2, job_id=4)
        if model is api.Job:
            return SimpleNamespace(
                id=4, client_id=ALIOR, title="QA", champion_profile=None
            )
        if model is api.Client:
            return SimpleNamespace(id=ALIOR, cv_content_mode_cap=cap)
        return SimpleNamespace(id=ident)

    db.get.side_effect = get
    app.dependency_overrides[api.get_db] = lambda: db
    app.dependency_overrides[get_args(api.CandidateWriteAccess)[1].dependency] = (
        lambda: User(id=7, role=UserRole.admin, roles=["admin"])
    )

    async def rule(_db, client_id):
        return _central_rule(client_id)

    monkeypatch.setattr(api, "resolve_client_rule", rule)

    async def no_preview(*args, **kwargs):
        raise HTTPException(503, "model unavailable")

    monkeypatch.setattr(champion_api, "read_preview", no_preview)
    pending = AsyncMock(return_value=11)
    monkeypatch.setattr(api, "_create_pending_row", pending)
    monkeypatch.setattr(api, "_charge_cv_generation_quota", AsyncMock())
    persisted = AsyncMock(return_value=21)
    monkeypatch.setattr(durable_jobs, "persist_job", persisted)
    monkeypatch.setattr(durable_jobs, "execute_job", AsyncMock())
    files = {"cv_file": ("cv.docx", _docx("Jan Kowalski. Python."), _DOCX)}
    if with_champion:
        files["champion_file"] = ("Champion.docx", _docx("MUST: Python"), _DOCX)
    data = {"content_mode": content_mode}
    if stage:
        data.update(candidate_id="2", stage_id="3")
    data.update(extra or {})
    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        response = await client.post(
            "/cv-generator/generate-upload", data=data, files=files
        )
    assert response.status_code == 202, response.text
    return persisted.call_args.kwargs["inputs"]["payload"], pending.call_args.kwargs


@pytest.mark.asyncio
async def test_upload_uses_the_uploaded_champion_for_tailored(monkeypatch):
    payload, row = await _upload(
        monkeypatch, content_mode="tailored", with_champion=True
    )
    assert payload.content_mode == row["content_mode"] == "tailored"
    assert payload.champion_bytes is not None
    assert payload.source_warnings == ()


@pytest.mark.asyncio
async def test_upload_without_champion_falls_back_with_a_notice(monkeypatch):
    payload, row = await _upload(
        monkeypatch, content_mode="tailored", with_champion=False
    )
    assert payload.content_mode == row["content_mode"] == "polished"
    assert payload.source_warnings == (policies.MISSING_CHAMPION_NOTICE,)


@pytest.mark.asyncio
async def test_upload_honours_an_explicit_polished_request(monkeypatch):
    payload, _ = await _upload(monkeypatch, content_mode="polished", with_champion=True)
    assert payload.content_mode == "polished"
    assert payload.source_warnings == ()


@pytest.mark.asyncio
async def test_upload_client_cap_wins_over_the_uploaded_champion(monkeypatch):
    payload, _ = await _upload(
        monkeypatch,
        content_mode="tailored",
        with_champion=True,
        cap="basic",
        stage=True,
    )
    assert payload.content_mode == "basic"


@pytest.mark.asyncio
async def test_upload_manual_requirements_are_not_a_champion(monkeypatch):
    payload, _ = await _upload(
        monkeypatch,
        content_mode="tailored",
        with_champion=False,
        extra={"must_requirements": "Python", "nice_requirements": "AWS"},
    )
    assert payload.champion_bytes is None
    assert payload.content_mode == "polished"
