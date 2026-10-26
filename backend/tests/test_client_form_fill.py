"""0522: wypełnianie formularza CV klienta — kontrola odpowiedzi modelu przez kod
i plik Word w układzie wzoru (bez bazy; trasy w ``test_client_forms_api.py``)."""

from __future__ import annotations

import io
import json
from contextlib import asynccontextmanager
from unittest.mock import AsyncMock

import docx
import pytest

from app.services import client_form_fill as fill
from app.services import client_form_render as render
from app.services import client_form_structure as form
from tests._client_form_docx import form_docx

CV = (
    "Jan Fikcyjny\nTester automatyzujący\n"
    "07.2023 – obecnie Automation Test Engineer / Bank Przykładowy S.A. / Warszawa\n"
    "- tworzenie frameworku do testów automatycznych w Rest Assured\n"
    "06.2020 – 01.2023 Specjalista / Firma Testowa Sp. z o.o.\n"
    "Umiejętności: SQL, Selenium WebDriver, Java\n"
    "Certyfikaty: ISTQB Certified Tester FL (09.2022)\n"
    "Kontakt: jan.fikcyjny@example.com, +48 600 100 200"
)
NOTES = "Kandydat nie pracował przy projektach związanych ze służbą zdrowia."
SOURCES = fill.Sources(cv=CV, notes=NOTES)


def _structure() -> dict:
    return form.read_structure(form_docx()).to_json()


def _parse(fields=(), extra_rows=()) -> dict:
    raw = json.dumps({"fields": list(fields), "extra_rows": list(extra_rows)})
    return fill.check_parse(
        raw, structure=_structure(), sources=SOURCES, candidate_name="Jan Fikcyjny"
    )


def test_value_needs_a_quote_from_cv_or_notes() -> None:
    result = _parse(
        [
            {
                "id": "t1.r1.c1",
                "value": "Bank Przykładowy S.A.",
                "quote": "Bank Przykładowy S.A.",
            },
            {
                "id": "t2.r6.c2",
                "value": "0",
                "quote": "nie pracował przy projektach związanych ze służbą zdrowia",
            },
            {
                "id": "t1.r1.c3",
                "value": "Architekt",
                "quote": "Architekt rozwiązań w chmurze",
            },
            {"id": "t1.r1.c4", "value": "Opis", "quote": None},
            {"id": "nie-ma-takiego", "value": "x", "quote": "SQL"},
        ]
    )
    fields = result["fields"]
    assert fields["t1.r1.c1"] == {
        "value": "Bank Przykładowy S.A.",
        "source": "cv",
        "quote": "Bank Przykładowy S.A.",
    }
    assert fields["t2.r6.c2"]["source"] == "notes"
    assert "t1.r1.c3" not in fields and "t1.r1.c4" not in fields
    assert result["dropped"] == 3


def test_name_comes_from_the_profile_not_the_model() -> None:
    result = _parse([{"id": "p2", "value": "Ktoś Inny", "quote": "Jan Fikcyjny"}])
    assert result["fields"]["p2"] == {
        "value": "Jan Fikcyjny",
        "source": "profile",
        "quote": None,
    }


def test_yes_no_scale_and_contact_rules() -> None:
    result = _parse(
        [
            {"id": "t2.r2.c2", "value": "tak", "quote": "SQL"},
            # Skala 1–10: liczba musi stać w cytacie — inaczej to zgadywanie.
            {"id": "t2.r5.c3", "value": "8", "quote": "Selenium WebDriver"},
            {"id": "t2.r6.c3", "value": "11", "quote": "Selenium WebDriver"},
            {
                "id": "t0.r0.c0",
                "value": "Kontakt: jan.fikcyjny@example.com",
                "quote": "Kontakt",
            },
            {"id": "t6.r0.c0", "value": "Tel. +48 600 100 200", "quote": "Kontakt"},
            {"id": "t1.r1.c2", "value": "2019 - 2023", "quote": "07.2023 – obecnie"},
        ]
    )
    fields = result["fields"]
    assert fields["t2.r2.c2"]["value"] == "TAK"
    assert "t2.r5.c3" not in fields and "t2.r6.c3" not in fields
    assert "t0.r0.c0" not in fields and "t6.r0.c0" not in fields
    # Okres to nie telefon.
    assert fields["t1.r1.c2"]["value"] == "2019 - 2023"


def test_extra_rows_fill_empty_template_rows_first_then_overflow() -> None:
    row = {
        "table": 1,
        "cells": {
            "c1": {
                "value": "Firma Testowa Sp. z o.o.",
                "quote": "Firma Testowa Sp. z o.o.",
            },
            "c2": {"value": "06.2020 – 01.2023", "quote": "06.2020 – 01.2023"},
        },
    }
    result = _parse(
        [
            {
                "id": "t1.r1.c1",
                "value": "Bank Przykładowy S.A.",
                "quote": "Bank Przykładowy S.A.",
            }
        ],
        [
            row,
            row,
            row,
            {"table": 2, "cells": {"c1": {"value": "SQL", "quote": "SQL"}}},
        ],
    )
    fields = result["fields"]
    # Wiersze 2 i 3 wzoru zajęte, trzeci dopisany ponad wzór.
    assert fields["t1.r2.c1"]["value"] == "Firma Testowa Sp. z o.o."
    assert fields["t1.r3.c2"]["value"] == "06.2020 – 01.2023"
    assert [item["table"] for item in result["extra_rows"]] == [1]
    # Tabela kompetencji fikcyjnego wzoru nie ma „…” — dopisać do niej nie wolno.
    assert result["dropped"] == 1


def test_extra_rows_are_capped_per_table(monkeypatch) -> None:
    monkeypatch.setattr(fill, "EXTRA_ROWS_MAX", 1)
    row = {
        "table": 1,
        "cells": {"c3": {"value": "Specjalista", "quote": "Specjalista"}},
    }
    result = _parse([], [row] * 6)
    assert len(result["extra_rows"]) == 1
    assert result["dropped"] == 2  # 3 w wiersze wzoru, 1 dopisany, 2 ponad limit


def test_prompt_fences_sources_and_skips_the_name_field() -> None:
    prompt = fill.build_prompt(_structure(), SOURCES)
    assert "<cv>" in prompt and "<screening_notes>" in prompt
    assert '"p2"' not in prompt
    assert '"t1.r1.c1"' in prompt
    assert '"columns": {"c1": "Miejsce pracy"' in prompt


@pytest.mark.asyncio
async def test_without_model_the_form_still_has_the_name(monkeypatch) -> None:
    monkeypatch.setattr(fill, "api_key_configured", lambda _model: False)
    result = await fill.fill(
        AsyncMock(),
        structure=_structure(),
        sources=SOURCES,
        candidate_name="Jan Fikcyjny",
        user_id=1,
    )
    assert result["status"] == "unavailable"
    assert result["message"] == fill.MSG_UNAVAILABLE
    assert result["fields"] == {
        "p2": {"value": "Jan Fikcyjny", "source": "profile", "quote": None}
    }


@asynccontextmanager
async def _no_quota(*_args, **_kwargs):
    yield


@pytest.mark.asyncio
async def test_model_error_is_a_hint_not_a_failure(monkeypatch) -> None:
    monkeypatch.setattr(fill, "api_key_configured", lambda _model: True)
    monkeypatch.setattr(fill, "ai_feature", _no_quota)

    def _boom(*_args):
        raise RuntimeError("przeciążenie")

    monkeypatch.setattr(fill, "_call_model", _boom)
    result = await fill.fill(
        AsyncMock(),
        structure=_structure(),
        sources=SOURCES,
        candidate_name="Jan Fikcyjny",
        user_id=1,
    )
    assert result["status"] == "unavailable" and result["model"] is None


@pytest.mark.asyncio
async def test_model_answer_goes_through_the_code_checks(monkeypatch) -> None:
    monkeypatch.setattr(fill, "api_key_configured", lambda _model: True)
    monkeypatch.setattr(fill, "ai_feature", _no_quota)
    answer = {
        "fields": [
            {"id": "p4", "value": "Warszawa", "quote": "Warszawa"},
            {"id": "t2.r2.c2", "value": "TAK", "quote": "SQL, Selenium"},
            {"id": "t2.r5.c3", "value": "7", "quote": "Selenium WebDriver"},
        ]
    }
    monkeypatch.setattr(fill, "_call_model", lambda *_a: json.dumps(answer))
    result = await fill.fill(
        AsyncMock(),
        structure=_structure(),
        sources=SOURCES,
        candidate_name="Jan Fikcyjny",
        user_id=1,
    )
    assert result["status"] == "ok"
    assert set(result["fields"]) == {"p2", "p4", "t2.r2.c2"}
    assert result["dropped"] == 1


# ── plik Word ────────────────────────────────────────────────────────────────


def _cell(document, table: int, row: int, grid: int) -> str:
    return document.tables[table].rows[row].cells[grid].text


def test_values_land_in_place_and_the_template_text_stays() -> None:
    structure = _structure()
    values = {
        "p2": {"value": "Jan Fikcyjny"},
        "t0.r0.c0": {"value": "Tester automatyzujący.\nTesty API i GUI."},
        "t1.r1.c1": {"value": "Bank Przykładowy S.A."},
        "t2.r2.c2": {"value": "TAK"},
        "t3.r1.c3": {"value": "06.2020 – 01.2023"},
    }
    content = render.render(form_docx(), structure, values, [])
    document = docx.Document(io.BytesIO(content))
    assert document.paragraphs[2].text == "Imię i nazwisko:  Jan Fikcyjny"
    assert _cell(document, 0, 0, 0) == "Tester automatyzujący.\nTesty API i GUI."
    assert _cell(document, 1, 1, 1) == "Bank Przykładowy S.A."
    assert _cell(document, 2, 2, 1) == "Znajomość języka SQL;"  # tekst wzoru
    assert _cell(document, 2, 2, 2) == "TAK"
    assert _cell(document, 3, 1, 3) == "06.2020 – 01.2023"
    # Bez dopisanych wierszy „…” znika, a wiersz zostaje pusty.
    assert len(document.tables[1].rows) == 6
    assert _cell(document, 1, 4, 0) == "4" and _cell(document, 1, 4, 1) == ""


def test_extra_rows_continue_numbering_and_replace_the_ellipsis_row() -> None:
    structure = _structure()
    rows = [
        {"table": 1, "cells": {"c1": {"value": "Firma A"}, "c2": {"value": "2019"}}},
        {"table": 1, "cells": {"c1": {"value": "Firma B"}}},
    ]
    content = render.render(form_docx(), structure, {}, rows)
    document = docx.Document(io.BytesIO(content))
    experience = document.tables[1]
    assert [row.cells[0].text for row in experience.rows] == [
        "Lp.",
        "1",
        "2",
        "3",
        "4",
        "5",
        "",
    ]
    assert experience.rows[4].cells[1].text == "Firma A"
    assert experience.rows[4].cells[2].text == "2019"
    assert experience.rows[5].cells[1].text == "Firma B"
    assert not any("…" in cell.text for row in experience.rows for cell in row.cells)


def test_quote_may_be_fragments_each_in_the_source() -> None:
    result = _parse(
        [
            {
                "id": "t1.r1.c4",
                "value": "Framework w Rest Assured; testy w Java.",
                "quote": [
                    "tworzenie frameworku do testów automatycznych w Rest Assured",
                    "SQL, Selenium WebDriver, Java",
                ],
            },
            {
                "id": "t1.r2.c4",
                "value": "Coś więcej",
                "quote": ["tworzenie frameworku", "fragment, którego nie ma"],
            },
        ]
    )
    assert result["fields"]["t1.r1.c4"]["source"] == "cv"
    assert " … " in result["fields"]["t1.r1.c4"]["quote"]
    assert "t1.r2.c4" not in result["fields"]


def test_city_comes_from_the_profile_when_known() -> None:
    raw = json.dumps({"fields": [{"id": "p4", "value": "Kraków", "quote": "Warszawa"}]})
    with_city = fill.check_parse(
        raw,
        structure=_structure(),
        sources=SOURCES,
        candidate_name="Jan Fikcyjny",
        candidate_city="Warszawa",
    )
    assert with_city["fields"]["p4"] == {
        "value": "Warszawa",
        "source": "profile",
        "quote": None,
    }
    without_city = fill.check_parse(
        raw, structure=_structure(), sources=SOURCES, candidate_name="Jan Fikcyjny"
    )
    assert without_city["fields"]["p4"]["source"] == "cv"
