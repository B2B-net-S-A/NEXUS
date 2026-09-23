"""CI PostgreSQL: three normal flows must keep a stored Champion budget.

Profiles stored before the v4 intake — ~949 of them from the 08.2026 bulk
import, in the FLAT shape — carry `rate_value` 140 next to the document text
in `rate_raw` ("140 zł netto/h", "do 140 zł netto/h", "120/140 zł/h" …).
After #1477 three flows re-derived that stored number from its text, and the
first PLN/h grammar could not read most of these spellings, so the budget —
and, with the sync box ticked, `jobs.rate_budget_hourly` — became empty:

1. "Uzgodnij profil i pola rekrutacji" (the reconcile dialog over the draft),
2. "Importuj Word/PDF" onto a job that already has a rate (rate not taken
   from the document),
3. `POST /api/jobs` with `from_job_id` (template copy).

The dialog is driven the way `ChampionImportReview` drives it: the fields as
strings (`readValues`/`withValues`) → `/api/champion/validate` →
`apply-import`. Every flow runs with the current client (no `rate_raw` for a
rate the document did not bring) AND with a client that sends the stored text
back — the server must hold the budget either way. Nothing is mocked but the
AI parser of the preview and the re-embed.
"""

import uuid
from contextlib import asynccontextmanager
from copy import deepcopy
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from sqlalchemy import select

from app.core.database import AsyncSessionLocal
from app.models.client import Client
from app.models.job import Job, JobStatus

RATE = "basics.rate_value"
FIXTURES = Path(__file__).parent / "fixtures" / "champion"

# The spellings seen next to a 140 PLN/h budget in the 08.2026 import.
STORED_TEXTS = [
    "140 zł netto/h",
    "140 PLN netto/h",
    "do 140 zł netto/h",
    "120-140 zł netto/h",
    "140 zł/h (netto, B2B)",
    "Stawka: 140 zł/h",
    "120/140 zł/h",
]

# The fields of `ChampionImportReview` (frontend/src/components/ChampionIntake.tsx).
DIALOG_FIELDS = [
    "basics.role_name",
    "basics.seniority_min_years",
    "basics.rate_value",
    "basics.work_mode",
    "basics.onsite_days_per_week",
    "basics.candidate_location_pref",
    "basics.language",
    "basics.start_date",
    "basics.deadline",
    "basics.contract_length",
    "search.keywords",
    "search.target_companies",
    "search.disqualifiers",
    "search.notes",
    "stack.must",
    "stack.nice",
    "stack.notes",
    "experience.domains",
    "experience.certifications",
    "experience.regulations",
    "project.about",
    "project.responsibilities",
    "client.selling_points",
    "client.consultant_insight",
    "client.historical_questions",
    "client.priority_rules",
]


def flat_profile(rate_raw: str) -> dict:
    """The flat pre-09.2026 shape of the August bulk import (invented data)."""
    return {
        "role_name": "Senior Java Developer",
        "seniority_min_years": 5,
        "rate_value": 140.0,
        "rate_raw": rate_raw,
        "work_mode": "hybrydowo",
        "start_date": "01.10.2026",
        "contract_length": "12 miesięcy",
        "project_context": {
            "about": "Platforma płatności dla klienta z sektora finansowego.",
            "responsibilities": "Rozwój usług REST i integracji.",
            "selling_points": "Stabilny projekt, nowoczesny stack.",
        },
        "sourcing": {"keywords": "Java Spring Kafka", "notes": ""},
        "disqualifiers": ["Brak doświadczenia komercyjnego"],
        "must_skills": [{"name": "Java 17"}, {"name": "Spring Boot"}],
        "screening_questions": [
            {
                "id": "q1",
                "question": "Jak projektujesz idempotentne API?",
                "ideal_answer": "Klucze idempotencji i ponowienia.",
                "deal_breaker": "",
            },
            {
                "id": "q2",
                "question": "Jak planujesz migrację schematu?",
                "ideal_answer": "Etapy, ryzyka i plan wycofania.",
                "deal_breaker": "",
            },
        ],
        "_source": "champion_upload",
        "_parsed_at": "2026-08-14T10:00:00+00:00",
        "_parser": "champion_parse:v3:haiku-4.5",
    }


# ── What the dialog sends (a port of `readValues` / `withValues`) ────────────


def _js_string(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _experience_line(item: dict) -> str:
    """Lustro `experienceToText` (frontend/src/lib/champion-experience.ts)."""
    line = item["name"]
    if item.get("min_years"):
        line += f" (min. {item['min_years']} lat)"
    if item.get("level") == "nice":
        line += " (mile)"
    return line


def read_values(profile: dict) -> dict:
    notes = (profile.get("intake") or {}).get("unresolved") or {}
    values = {}
    for path in DIALOG_FIELDS:
        section, key = path.split(".")
        value = (profile.get(section) or {}).get(key)
        if isinstance(value, list) and section == "experience":
            canonical = "\n".join(_experience_line(item) for item in value)
        elif isinstance(value, list):
            canonical = "\n".join(
                item["name"] if isinstance(item, dict) else str(item) for item in value
            )
        else:
            canonical = _js_string(value)
        note = notes.get(path)
        if isinstance(value, list) and note:
            values[path] = "\n".join(x for x in (canonical, note) if x)
        else:
            values[path] = note if note is not None else canonical
    return values


def dialog_payload(source: dict, *, current: dict | None, rate_raw) -> dict:
    """`ChampionImportReview.apply` with the default selection.

    A field is taken from the document when the current profile has no value
    for it (or there is no current profile — the reconcile dialog).
    """
    source_values = read_values(source)
    old = read_values(current) if current else {}
    base = deepcopy(current if current else source)
    for path in DIALOG_FIELDS:
        section, key = path.split(".")
        value = source_values[path] if not old.get(path) else old[path]
        base.setdefault(section, {})
        base[section] = {
            **base[section],
            key: (value or None) if section == "basics" else value,
        }
    base["basics"]["rate_raw"] = rate_raw
    source_meta = source.get("intake") or {}
    intake = {
        **(base.get("intake") or {}),
        "policy_version": 1,
        "unresolved": {},
        "template_version": source_meta.get("template_version"),
    }
    if source_meta.get("document_context") is not None:
        intake["document_context"] = source_meta["document_context"]
    base["intake"] = intake
    if not (current and current.get("screening_questions")):
        base["screening_questions"] = source.get("screening_questions") or []
    return base


# ── Fixtures and HTTP helpers ────────────────────────────────────────────────


@pytest.fixture
def no_refresh(monkeypatch):
    async def noop(*args, **kwargs):
        return None

    monkeypatch.setattr("app.services.job_matching_refresh.refresh_job_matching", noop)


@pytest.fixture
def ai_document(monkeypatch, no_refresh):
    """The AI parser answers with the given rate; quota stubbed."""
    answer: dict = {}

    async def fake_parse(text):
        return {
            "basics": {
                "role_name": "Senior Java Developer",
                "rate_value": answer["number"],
                "rate_raw": answer["raw"],
                "work_mode": "zdalnie",
            },
            "stack": {"must": [{"name": "Java 17"}], "nice": []},
            "project": {"about": "Nowy projekt.", "responsibilities": "API"},
            "screening_questions": [],
        }

    @asynccontextmanager
    async def no_quota(*args, **kwargs):
        yield None

    monkeypatch.setattr(
        "app.services.champion_profile_ingest.parse_champion_document", fake_parse
    )
    monkeypatch.setattr("app.services.ai_quota.ai_feature", no_quota)
    return answer


async def _seed_job(profile: dict | None, *, budget: float | None = 140.0) -> int:
    async with AsyncSessionLocal() as db:
        client = Client(name=f"Synthetic budget retention {uuid.uuid4().hex[:8]}")
        db.add(client)
        await db.flush()
        job = Job(
            title="Synthetic Java role",
            status=JobStatus.draft,
            client_id=client.id,
            champion_profile=profile,
            rate_budget_hourly=budget,
        )
        db.add(job)
        await db.commit()
        return job.id


async def _job(job_id: int) -> Job:
    async with AsyncSessionLocal() as db:
        return await db.scalar(select(Job).where(Job.id == job_id))


def _word_document() -> bytes:
    return (FIXTURES / "v4-two-questions.docx").read_bytes()


def _unreadable_document() -> bytes:
    """No v4 tables: the preview goes to the (mocked) AI parser."""
    document = Document()
    document.add_paragraph("Profil Championa — rola Java, opis w treści.")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


async def _preview(app_client, headers, data: bytes) -> dict:
    response = await app_client.post(
        "/api/champion/preview",
        files={
            "file": (
                "profil.docx",
                data,
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document",
            )
        },
        headers=headers,
    )
    assert response.status_code == 200, response.text
    return response.json()


async def _validate_and_apply(
    app_client, headers, job_id: int, payload: dict, sync_fields: list[str]
) -> dict:
    route = f"/api/jobs/{job_id}/champion-profile"
    before = (await app_client.get(route, headers=headers)).json()
    checked = await app_client.post(
        "/api/champion/validate", json={"profile": payload}, headers=headers
    )
    assert checked.status_code == 200, checked.text
    applied = await app_client.post(
        route + "/apply-import",
        json={
            "profile": checked.json()["champion_profile"],
            "expected_fingerprint": before["fingerprint"],
            "sync_fields": sync_fields,
        },
        headers=headers,
    )
    assert applied.status_code == 200, applied.text
    return applied.json()


async def _assert_budget_kept(job_id: int, text: str) -> None:
    job = await _job(job_id)
    basics = job.champion_profile["basics"]
    assert basics["rate_value"] == 140.0, f"{text!r}: the stored budget was wiped"
    assert basics["rate_raw"] == text
    assert RATE not in (job.champion_profile.get("intake") or {}).get("unresolved", {})
    assert float(job.rate_budget_hourly) == 140.0


# ── 1. The reconcile dialog over the stored draft ────────────────────────────


@pytest.mark.parametrize("resend_text", [False, True], ids=["no-text", "old-client"])
@pytest.mark.parametrize("text", STORED_TEXTS)
async def test_reconcile_dialog_keeps_the_stored_budget(
    app_client, app_auth_headers, no_refresh, text, resend_text
):
    job_id = await _seed_job(flat_profile(text))
    draft = (
        await app_client.get(
            f"/api/jobs/{job_id}/champion-profile", headers=app_auth_headers
        )
    ).json()["champion_profile"]
    assert draft["basics"]["rate_value"] == 140.0

    payload = dialog_payload(
        draft,
        current=None,
        rate_raw=draft["basics"]["rate_raw"] if resend_text else None,
    )
    applied = await _validate_and_apply(
        app_client, app_auth_headers, job_id, payload, ["rate_value"]
    )

    assert applied["champion_profile"]["basics"]["rate_value"] == 140.0
    assert applied["job_values"]["rate_value"] == 140.0
    codes = {i["code"] for i in applied["validation"]["issues"]}
    assert "missing_budget" not in codes
    await _assert_budget_kept(job_id, text)


# ── 2. A document imported onto a job that already has a rate ───────────────


@pytest.mark.parametrize("resend_text", [False, True], ids=["no-text", "old-client"])
@pytest.mark.parametrize("text", STORED_TEXTS)
async def test_import_that_keeps_the_current_rate_keeps_the_budget(
    app_client, app_auth_headers, ai_document, text, resend_text
):
    ai_document.update(raw="150 zł/h", number=150.0)
    job_id = await _seed_job(flat_profile(text))
    current = (
        await app_client.get(
            f"/api/jobs/{job_id}/champion-profile", headers=app_auth_headers
        )
    ).json()["champion_profile"]
    preview = await _preview(app_client, app_auth_headers, _unreadable_document())
    assert preview["champion_profile"]["basics"]["rate_value"] == 150.0

    # The rate is not taken from the document (the current one is non-empty).
    payload = dialog_payload(
        preview["champion_profile"],
        current=current,
        rate_raw=current["basics"]["rate_raw"] if resend_text else None,
    )
    assert payload["basics"]["rate_value"] == "140"
    await _validate_and_apply(
        app_client, app_auth_headers, job_id, payload, ["rate_value"]
    )

    await _assert_budget_kept(job_id, text)


# ── 3. Template copy ─────────────────────────────────────────────────────────


@pytest.mark.parametrize("text", STORED_TEXTS)
async def test_template_copy_copies_the_budget_as_stored(
    app_client, app_auth_headers, text
):
    source_id = await _seed_job(flat_profile(text))
    source = await _job(source_id)

    response = await app_client.post(
        "/api/jobs",
        json={
            "title": "Synthetic Java role (kopia)",
            "client_id": source.client_id,
            "from_job_id": source_id,
        },
        headers=app_auth_headers,
    )
    assert response.status_code == 201, response.text

    copy = (await _job(response.json()["id"])).champion_profile
    assert copy["basics"]["rate_value"] == 140.0, f"{text!r}: the copy lost the budget"
    assert copy["basics"]["rate_raw"] == text
    assert RATE not in (copy.get("intake") or {}).get("unresolved", {})
    # Nothing else was re-normalised either: the date keeps its stored form.
    assert copy["basics"]["start_date"] == "01.10.2026"


# ── Fresh documents: the text is the source of truth ─────────────────────────


@pytest.mark.parametrize(
    "raw,number,budget,warned",
    [
        ("120–140 zł/h", 130.0, 140.0, True),
        ("do 140 zł netto/h", None, 140.0, True),
        ("140 zł/h", 150.0, 140.0, False),
        ("45 EUR/h", 45.0, None, False),
    ],
)
async def test_a_fresh_document_rate_through_preview_and_apply(
    app_client, app_auth_headers, ai_document, raw, number, budget, warned
):
    ai_document.update(raw=raw, number=number)
    job_id = await _seed_job(None, budget=None)
    preview = await _preview(app_client, app_auth_headers, _unreadable_document())
    source = preview["champion_profile"]

    # The rate is taken from the document untouched: the dialog sends its text.
    payload = dialog_payload(
        source, current=None, rate_raw=source["basics"]["rate_raw"]
    )
    applied = await _validate_and_apply(
        app_client, app_auth_headers, job_id, payload, []
    )

    profile = applied["champion_profile"]
    assert profile["basics"]["rate_value"] == budget
    assert profile["basics"]["rate_raw"] == raw
    assert (RATE in profile["intake"]["advisory"]) is warned
    assert (RATE in profile["intake"]["unresolved"]) is (budget is None)
    job = await _job(job_id)
    if budget is None:
        assert job.rate_budget_hourly is None
    else:
        assert float(job.rate_budget_hourly) == budget


async def test_a_word_form_rate_cell_with_qualifiers_is_a_budget(
    app_client, app_auth_headers, no_refresh
):
    document = Document(BytesIO(_word_document()))
    row = next(
        row
        for table in document.tables
        for row in table.rows
        if any("maks. stawka kandydata" in cell.text.lower() for cell in row.cells)
    )
    labels = [cell.text.lower() for cell in row.cells]
    rate_cell = next(i for i, label in enumerate(labels) if "maks. stawka" in label) + 1
    row.cells[rate_cell].text = "140 zł/h netto"
    stream = BytesIO()
    document.save(stream)
    job_id = await _seed_job(None, budget=None)

    preview = await _preview(app_client, app_auth_headers, stream.getvalue())
    assert preview["summary"]["rate_value"] == 140.0
    source = preview["champion_profile"]
    assert source["basics"]["rate_raw"] == "140 zł/h netto"

    payload = dialog_payload(
        source, current=None, rate_raw=source["basics"]["rate_raw"]
    )
    applied = await _validate_and_apply(
        app_client, app_auth_headers, job_id, payload, []
    )

    assert applied["champion_profile"]["basics"]["rate_value"] == 140.0
    assert float((await _job(job_id)).rate_budget_hourly) == 140.0
