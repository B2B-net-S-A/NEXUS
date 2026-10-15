"""Nowe CV z rekrutacji — `POST /api/candidates/stages/{id}/cv/original/upload`
i `…/cv/original/refresh` na etapie bez wiersza CV (13.10.2026).

Kafelek „Wgraj nowe CV” w panelu osoby i okno „Dodaj z CV” robią jedno: plik
staje się głównym CV profilu i od razu kopią CV etapu („Oryginał CV”). Do tej
daty nowe CV wgrane w profilu nie zmieniało kopii w rekrutacji, a ręczne
„Aktualizuj snapshot” dawało 404 na etapach z importu Traffita (bez wiersza
CV), czyli na większości kart w „Nowych” i „Screeningu”.
"""

from __future__ import annotations

import hashlib
import json
import uuid

import pytest
from httpx import AsyncClient
from sqlalchemy import func, select

from app.core.database import AsyncSessionLocal
from app.models.activity import Activity
from app.models.candidate_document import CandidateDocument
from app.models.candidate_stage_cv import CandidateStageCV
from app.models.recruitment_pipeline import CandidateStage
from app.models.user import UserRole
from app.services.candidate_stage_cv_service import create_original_cv_snapshot
from tests.test_candidate_stage_cv_snapshot import (
    _seed_candidate_with_cv,
    _seed_job,
    _seed_stage,
)
from tests.test_job_proposals import _user

pytestmark = pytest.mark.asyncio

NEW_PDF = (
    b"%PDF-1.4\n1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
    b"2 0 obj<</Type/Pages/Kids[]/Count 0>>endobj\ntrailer<</Root 1 0 R>>\n%%EOF\n"
)


@pytest.fixture(autouse=True)
def _no_cv_reading(monkeypatch: pytest.MonkeyPatch) -> None:
    # Odczyt AI biegnie w tle po odpowiedzi; pusty tekst kończy go od razu.
    monkeypatch.setattr(
        "app.services.cv_text_extractor.extract_text",
        lambda *_args, **_kwargs: "",
    )


async def _stage_with_snapshot(*, snapshot: bool = True) -> tuple[int, int]:
    cid = await _seed_candidate_with_cv(
        cv_bytes=b"%PDF-1.4 OLD", cv_filename="stare.pdf"
    )
    sid = await _seed_stage(cid, await _seed_job())
    if snapshot:
        async with AsyncSessionLocal() as db:
            stage = await db.scalar(
                select(CandidateStage).where(CandidateStage.id == sid)
            )
            await create_original_cv_snapshot(db, stage)
            await db.commit()
    return cid, sid


async def _upload(
    client: AsyncClient,
    headers: dict[str, str],
    stage_id: int,
    *,
    filename: str = "nowe-cv.pdf",
    content: bytes = NEW_PDF,
):
    return await client.post(
        f"/api/candidates/stages/{stage_id}/cv/original/upload",
        headers=headers,
        files={"file": (filename, content, "application/pdf")},
    )


async def _documents(candidate_id: int) -> list[CandidateDocument]:
    async with AsyncSessionLocal() as db:
        return list(
            (
                await db.scalars(
                    select(CandidateDocument)
                    .where(CandidateDocument.candidate_id == candidate_id)
                    .order_by(CandidateDocument.id)
                )
            ).all()
        )


async def test_refresh_creates_missing_stage_cv_row(
    app_client: AsyncClient, app_auth_headers: dict[str, str]
) -> None:
    cid, sid = await _stage_with_snapshot(snapshot=False)

    res = await app_client.post(
        f"/api/candidates/stages/{sid}/cv/original/refresh", headers=app_auth_headers
    )

    assert res.status_code == 200, res.text
    body = res.json()
    assert body["has_snapshot"] is True
    assert body["original_snapshot_source"] == "manual_refresh"
    assert body["candidate_id"] == cid


async def test_upload_makes_primary_cv_and_replaces_stage_copy(
    app_client: AsyncClient, app_auth_headers: dict[str, str]
) -> None:
    cid, sid = await _stage_with_snapshot()

    res = await _upload(app_client, app_auth_headers, sid)

    assert res.status_code == 200, res.text
    body = res.json()
    assert body["original_cv_filename"] == "nowe-cv.pdf"
    assert body["original_snapshot_source"] == "manual_upload"
    assert body["has_snapshot"] is True

    docs = await _documents(cid)
    primary = [d for d in docs if d.is_primary]
    assert [d.filename for d in primary] == ["nowe-cv.pdf"]

    download = await app_client.get(
        f"/api/candidates/stages/{sid}/cv/original/download", headers=app_auth_headers
    )
    assert download.status_code == 200, download.text
    assert download.content == NEW_PDF

    async with AsyncSessionLocal() as db:
        csv_id = await db.scalar(
            select(CandidateStageCV.id).where(
                CandidateStageCV.candidate_stage_id == sid
            )
        )
        details = await db.scalar(
            select(Activity.details)
            .where(
                Activity.entity_type == "candidate_stage_cv",
                Activity.entity_id == csv_id,
                Activity.action == "snapshot_refreshed",
            )
            .order_by(Activity.id.desc())
            .limit(1)
        )
    assert details["snapshot_source"] == "manual_upload"
    assert details["old_filename"] == "stare.pdf"
    assert details["new_filename"] == "nowe-cv.pdf"


async def test_upload_on_stage_without_cv_row_creates_it(
    app_client: AsyncClient, app_auth_headers: dict[str, str]
) -> None:
    _, sid = await _stage_with_snapshot(snapshot=False)

    res = await _upload(app_client, app_auth_headers, sid)

    assert res.status_code == 200, res.text
    assert res.json()["original_snapshot_source"] == "manual_upload"


async def test_upload_refuses_formats_the_stage_copy_cannot_use(
    app_client: AsyncClient, app_auth_headers: dict[str, str]
) -> None:
    cid, sid = await _stage_with_snapshot()

    res = await _upload(
        app_client, app_auth_headers, sid, filename="cv.odt", content=b"odt"
    )

    assert res.status_code == 415, res.text
    assert "PDF, DOC albo DOCX" in res.json()["detail"]
    assert await _documents(cid) == []
    original = await app_client.get(
        f"/api/candidates/stages/{sid}/cv/original", headers=app_auth_headers
    )
    assert original.json()["original_cv_filename"] == "stare.pdf"


async def test_same_file_already_in_profile_becomes_primary(
    app_client: AsyncClient, app_auth_headers: dict[str, str]
) -> None:
    cid, sid = await _stage_with_snapshot()
    for filename, content, primary in (
        ("glowne.pdf", b"%PDF-1.4 glowne " + uuid.uuid4().hex.encode(), "true"),
        ("angielskie.pdf", NEW_PDF, "false"),
    ):
        res = await app_client.post(
            f"/api/candidates/{cid}/documents",
            headers=app_auth_headers,
            files={"file": (filename, content, "application/pdf")},
            data={"document_kind": "cv", "is_primary": primary},
        )
        assert res.status_code == 201, res.text

    res = await _upload(app_client, app_auth_headers, sid, filename="angielskie.pdf")

    assert res.status_code == 200, res.text
    docs = await _documents(cid)
    assert len(docs) == 2
    assert [d.filename for d in docs if d.is_primary] == ["angielskie.pdf"]


async def test_role_without_candidate_write_is_refused_before_any_write(
    app_client: AsyncClient,
) -> None:
    # Bramka zespołu rekrutacji jest otwarta dla każdej roli wewnętrznej
    # (decyzja z 23.09.2026), więc odmowa idzie z prawa zapisu kandydata.
    cid, sid = await _stage_with_snapshot()
    _, viewer_headers = await _user(UserRole.user)

    res = await _upload(app_client, viewer_headers, sid)

    assert res.status_code == 403, res.text
    async with AsyncSessionLocal() as db:
        count = await db.scalar(
            select(func.count())
            .select_from(CandidateDocument)
            .where(CandidateDocument.candidate_id == cid)
        )
    assert count == 0


async def _stage_copy_of(
    content: bytes, *, candidate_cv: bytes | None
) -> tuple[int, int]:
    """Etap z kopią CV `content` u osoby, której bieżące CV to `candidate_cv`."""
    cid = await _seed_candidate_with_cv(cv_bytes=candidate_cv, cv_filename="wlasne.pdf")
    sid = await _seed_stage(cid, await _seed_job())
    async with AsyncSessionLocal() as db:
        db.add(
            CandidateStageCV(
                candidate_stage_id=sid,
                candidate_id=cid,
                job_id=(
                    await db.scalar(
                        select(CandidateStage.job_id).where(CandidateStage.id == sid)
                    )
                ),
                original_cv_filename="cudze.pdf",
                original_cv_content=content,
                original_cv_sha256=hashlib.sha256(content).hexdigest(),
                original_snapshot_source="manual_upload",
            )
        )
        await db.commit()
    return cid, sid


async def test_quarantine_puts_the_persons_own_cv_back_into_stage_copies() -> None:
    from app.models.candidate import Candidate
    from app.services.candidate_stage_cv_service import (
        restore_stage_copies_after_quarantine,
    )

    wrong = b"%PDF-1.4 CV innej osoby"
    cid, sid = await _stage_copy_of(wrong, candidate_cv=b"%PDF-1.4 wlasne CV")
    async with AsyncSessionLocal() as db:
        candidate = await db.scalar(select(Candidate).where(Candidate.id == cid))
        fixed = await restore_stage_copies_after_quarantine(
            db,
            candidate=candidate,
            quarantined_sha256=hashlib.sha256(wrong).hexdigest(),
        )
        await db.commit()
        row = await db.scalar(
            select(CandidateStageCV).where(CandidateStageCV.candidate_stage_id == sid)
        )
    assert fixed == 1
    assert row.original_snapshot_source == "quarantine_restore"
    assert row.original_cv_sha256 == hashlib.sha256(b"%PDF-1.4 wlasne CV").hexdigest()
    assert row.original_cv_filename == "wlasne.pdf"


async def test_quarantine_without_own_cv_leaves_stage_copy_empty() -> None:
    from app.models.candidate import Candidate
    from app.services.candidate_stage_cv_service import (
        restore_stage_copies_after_quarantine,
        snapshot_exists,
    )

    wrong = b"%PDF-1.4 CV innej osoby " + uuid.uuid4().hex.encode()
    cid, sid = await _stage_copy_of(wrong, candidate_cv=None)
    async with AsyncSessionLocal() as db:
        candidate = await db.scalar(select(Candidate).where(Candidate.id == cid))
        await restore_stage_copies_after_quarantine(
            db,
            candidate=candidate,
            quarantined_sha256=hashlib.sha256(wrong).hexdigest(),
        )
        await db.commit()
        row = await db.scalar(
            select(CandidateStageCV).where(CandidateStageCV.candidate_stage_id == sid)
        )
    assert not snapshot_exists(row)
    assert row.original_cv_sha256 is None


async def test_candidate_pipelines_say_which_board_column_the_card_is_in(
    app_client: AsyncClient, app_auth_headers: dict[str, str]
) -> None:
    # Okno „Dodaj z CV” decyduje o podmianie CV tą samą regułą co Tablica
    # i `proposals/bulk` (zatrudniony ≠ zamknięty).
    from app.models.recruitment_pipeline import PipelineStage

    cid = await _seed_candidate_with_cv()
    in_progress_job, hired_job = await _seed_job(), await _seed_job()
    await _seed_stage(cid, in_progress_job)
    async with AsyncSessionLocal() as db:
        db.add(
            CandidateStage(
                candidate_id=cid, job_id=hired_job, stage=PipelineStage.hired
            )
        )
        await db.commit()

    res = await app_client.get(
        f"/api/candidates/{cid}/pipelines", headers=app_auth_headers
    )

    assert res.status_code == 200, res.text
    columns = {row["job_id"]: row["board_column"] for row in res.json()["pipelines"]}
    assert columns == {in_progress_job: "new", hired_job: "hired"}


# ── Gotowe CV do klienta z kafelka (15.10.2026, D1–D3) ───────────────────────
#
# Zgłoszenie rekrutera: plik „B2B_…docx” wgrany kafelkiem stawał się oryginałem
# kandydata i głównym CV profilu, a CV do klienta trzeba było jeszcze podmienić
# paskiem „Ta osoba ma już gotowe CV”. Teraz plik nazwany jak CV do klienta
# staje się nim od razu; CV kandydata i profil zostają bez zmian.

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
CLIENT_CV_NAME = "B2B_Analityk_Ewa_Fikcyjna.docx"


def _client_docx(header: str = "Analityk – Ewa Fikcyjna", marker: str = "") -> bytes:
    from io import BytesIO

    from docx import Document

    doc = Document()
    doc.add_paragraph(header)
    doc.add_heading("DOŚWIADCZENIE", level=1)
    doc.add_paragraph("01.2020 - 03.2024")
    doc.add_paragraph("Nazwa firmy: Firma Przykładowa")
    doc.add_paragraph("Stanowisko: Analityk systemowy")
    doc.add_paragraph(
        "Prowadziła analizę wymagań dla systemu rozliczeń i odbiory z klientem "
        f"biznesowym; spisywała procesy i scenariusze testów akceptacyjnych. {marker}",
        style="List Bullet",
    )
    doc.add_paragraph(
        "Osiem lat w zespołach analitycznych, ostatnio w bankowości; prowadziła "
        "warsztaty z biznesem i dokumentację wymagań w Confluence."
    )
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


async def _client_world(
    *,
    stage: str = "screening",
    primary_cv: bool = True,
) -> dict:
    """Osoba „Ewa Fikcyjna” z głównym CV w PDF i kartą w rekrutacji."""
    from app.models.candidate import Candidate
    from app.models.candidate_document import CandidateDocumentKind
    from app.models.client import Client
    from app.models.job import Job
    from app.models.recruitment_pipeline import PipelineStage

    tag = uuid.uuid4().hex[:8]
    async with AsyncSessionLocal() as db:
        client = Client(name=f"Klient CV do klienta {tag}")
        db.add(client)
        await db.flush()
        job = Job(title=f"Analityk {tag}", client_id=client.id, description="x")
        candidate = Candidate(
            name="Ewa", lastname="Fikcyjna", email=f"ewa-{tag}@example.com"
        )
        db.add_all([job, candidate])
        await db.flush()
        if primary_cv:
            db.add(
                CandidateDocument(
                    candidate_id=candidate.id,
                    filename="ewa-cv.pdf",
                    file_content=b"%PDF-1.4 wlasne CV " + tag.encode(),
                    document_kind=CandidateDocumentKind.cv,
                    is_primary=True,
                )
            )
        row = CandidateStage(
            candidate_id=candidate.id, job_id=job.id, stage=PipelineStage(stage)
        )
        db.add(row)
        await db.flush()
        await create_original_cv_snapshot(db, row)
        await db.commit()
        return {"candidate_id": candidate.id, "job_id": job.id, "stage_id": row.id}


async def _upload_docx(
    client: AsyncClient,
    headers: dict[str, str],
    stage_id: int,
    content: bytes,
    *,
    filename: str = CLIENT_CV_NAME,
    data: dict | None = None,
):
    return await client.post(
        f"/api/candidates/stages/{stage_id}/cv/original/upload",
        headers=headers,
        files={"file": (filename, content, DOCX_MIME)},
        data=data or {},
    )


async def _stage_cv(stage_id: int) -> CandidateStageCV:
    async with AsyncSessionLocal() as db:
        return await db.scalar(
            select(CandidateStageCV).where(
                CandidateStageCV.candidate_stage_id == stage_id
            )
        )


async def _auto_cv(world: dict, stage_id: int | None = None) -> int:
    """Szkic od automatu po „Zweryfikowany”, którego nikt nie ruszał."""
    from app.api.candidate_stage_cv import apply_generated_to_stage_cv
    from app.models.cv_generated_document import CvGeneratedDocument
    from app.services.cv_auto_generate import ATTACH_ACTION

    stage_id = stage_id or world["stage_id"]
    async with AsyncSessionLocal() as db:
        generated = CvGeneratedDocument(
            candidate_id=world["candidate_id"],
            job_id=world["job_id"],
            candidate_name="Ewa Fikcyjna",
            filename="auto.docx",
            status="ready",
            origin="auto",
            stage_id=stage_id,
            source_cv_revision=uuid.uuid4().hex,
            render_payload={
                "name": "Ewa Fikcyjna",
                "position": "Analityk",
                "why_points": ["Prowadziła analizę wymagań"],
            },
        )
        db.add(generated)
        await db.flush()
        csv = await db.scalar(
            select(CandidateStageCV).where(
                CandidateStageCV.candidate_stage_id == stage_id
            )
        )
        await apply_generated_to_stage_cv(
            db, csv, generated, user_id=None, activity_action=ATTACH_ACTION
        )
        await db.commit()
        return generated.id


@pytest.fixture
def no_profile_reading(monkeypatch: pytest.MonkeyPatch) -> list:
    """Zadanie AI czytające profil — CV do klienta nie może go uruchomić."""
    calls: list = []

    async def _record(*args, **kwargs):
        calls.append((args, kwargs))

    monkeypatch.setattr(
        "app.api.candidates._enrich_candidate_from_document_task", _record
    )
    return calls


async def test_client_cv_file_becomes_the_client_cv_not_the_original(
    app_client: AsyncClient, app_auth_headers: dict[str, str], no_profile_reading
) -> None:
    world = await _client_world()
    before = await _stage_cv(world["stage_id"])

    res = await _upload_docx(
        app_client, app_auth_headers, world["stage_id"], _client_docx()
    )

    assert res.status_code == 200, res.text
    body = res.json()
    assert body["outcome"] == "client_cv"
    assert (body["replaced"], body["undo_available"], body["stored_in_profile"]) == (
        "none",
        True,
        True,
    )
    assert body["client_cv_stage_id"] == world["stage_id"]
    assert body["client_cv"]["status"] == "draft"
    assert body["client_cv"]["origin"]["kind"] == "upload"
    assert "analizę wymagań" in body["client_cv"]["content_html"]
    # Oryginał kandydata i główne CV profilu bez zmian, bez odczytu AI.
    assert body["original_cv_filename"] == before.original_cv_filename == "ewa-cv.pdf"
    assert body["original_snapshot_source"] == before.original_snapshot_source
    docs = await _documents(world["candidate_id"])
    assert [(d.filename, d.is_primary) for d in docs] == [
        ("ewa-cv.pdf", True),
        (CLIENT_CV_NAME, False),
    ]
    assert no_profile_reading == []
    async with AsyncSessionLocal() as db:
        details = await db.scalar(
            select(Activity.details).where(
                Activity.entity_type == "candidate_stage_cv",
                Activity.entity_id == before.id,
                Activity.action == "branded_cv_uploaded_in_recruitment",
            )
        )
    assert details["prior"] == {"status": "none", "generated_document_id": None}
    assert details["source_document_id"] == docs[1].id


async def test_client_filename_pattern_of_the_recruitment_client_is_recognised(
    app_client: AsyncClient,
    app_auth_headers: dict[str, str],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    async def _pattern(db, client_id):
        return "ZOB-{PROJEKT}_{STANOWISKO}_{IMIE_NAZWISKO}"

    monkeypatch.setattr(
        "app.services.client_cv_upload.client_filename_pattern", _pattern
    )
    world = await _client_world()

    res = await _upload_docx(
        app_client,
        app_auth_headers,
        world["stage_id"],
        _client_docx(),
        filename="ZOB-2962_Analityk_Ewa_Fikcyjna.docx",
    )

    assert res.status_code == 200, res.text
    assert res.json()["outcome"] == "client_cv"


@pytest.mark.parametrize(
    ("filename", "content", "mime"),
    [
        ("B2B_Analityk_Ewa_Fikcyjna.pdf", NEW_PDF, "application/pdf"),
        ("Ewa_Fikcyjna_CV.docx", None, DOCX_MIME),
    ],
)
async def test_candidate_cv_files_still_become_the_original(
    app_client: AsyncClient, app_auth_headers: dict[str, str], filename, content, mime
) -> None:
    world = await _client_world()

    res = await app_client.post(
        f"/api/candidates/stages/{world['stage_id']}/cv/original/upload",
        headers=app_auth_headers,
        files={"file": (filename, content or _client_docx(), mime)},
    )

    assert res.status_code == 200, res.text
    body = res.json()
    assert body["outcome"] == "original"
    assert body["original_cv_filename"] == filename
    assert body["original_snapshot_source"] == "manual_upload"
    assert (await _stage_cv(world["stage_id"])).branded_status == "none"


async def test_untouched_auto_cv_is_replaced_without_asking_and_undo_restores_it(
    app_client: AsyncClient, app_auth_headers: dict[str, str]
) -> None:
    from app.models.cv_stage_draft_snapshot import CvStageDraftSnapshot

    world = await _client_world(stage="verified")
    auto_id = await _auto_cv(world)
    content = _client_docx(marker=uuid.uuid4().hex)

    res = await _upload_docx(app_client, app_auth_headers, world["stage_id"], content)

    assert res.status_code == 200, res.text
    body = res.json()
    assert (body["replaced"], body["undo_available"]) == ("auto_cv", True)
    csv = await _stage_cv(world["stage_id"])
    assert csv.generated_document_id is None
    async with AsyncSessionLocal() as db:
        snapshots = (
            await db.scalars(
                select(CvStageDraftSnapshot.reason).where(
                    CvStageDraftSnapshot.candidate_stage_cv_id == csv.id
                )
            )
        ).all()
    assert snapshots == ["replaced"]

    # „Cofnij”: ten sam plik jako CV kandydata, auto-CV wraca.
    undo = await _upload_docx(
        app_client,
        app_auth_headers,
        world["stage_id"],
        content,
        data={"mode": "original"},
    )

    assert undo.status_code == 200, undo.text
    assert undo.json()["outcome"] == "original"
    csv = await _stage_cv(world["stage_id"])
    assert (csv.branded_status, csv.generated_document_id) == ("draft", auto_id)
    assert csv.original_cv_filename == CLIENT_CV_NAME
    docs = await _documents(world["candidate_id"])
    assert [d.filename for d in docs if d.is_primary] == [CLIENT_CV_NAME]


async def test_undo_on_a_card_without_client_cv_empties_it_again(
    app_client: AsyncClient, app_auth_headers: dict[str, str]
) -> None:
    world = await _client_world()
    content = _client_docx(marker=uuid.uuid4().hex)
    first = await _upload_docx(app_client, app_auth_headers, world["stage_id"], content)
    assert first.status_code == 200, first.text

    undo = await _upload_docx(
        app_client,
        app_auth_headers,
        world["stage_id"],
        content,
        data={"mode": "original"},
    )

    assert undo.status_code == 200, undo.text
    csv = await _stage_cv(world["stage_id"])
    assert (csv.branded_status, csv.branded_draft_html) == ("none", None)
    assert csv.original_snapshot_source == "manual_upload"


async def test_edited_client_cv_is_replaced_only_after_confirmation(
    app_client: AsyncClient, app_auth_headers: dict[str, str]
) -> None:
    world = await _client_world()
    async with AsyncSessionLocal() as db:
        csv = await db.scalar(
            select(CandidateStageCV).where(
                CandidateStageCV.candidate_stage_id == world["stage_id"]
            )
        )
        csv.branded_status = "draft"
        csv.branded_draft_html = "<p>CV poprawione przez rekrutera</p>"
        csv.edit_revision = 4
        await db.commit()

    asked = await _upload_docx(
        app_client, app_auth_headers, world["stage_id"], _client_docx()
    )

    assert asked.status_code == 409, asked.text
    detail = asked.json()["detail"]
    assert (detail["code"], detail["status"]) == ("client_cv_replace_confirm", "draft")
    assert "Ta rekrutacja ma już CV do klienta" in detail["message"]
    assert [d.filename for d in await _documents(world["candidate_id"])] == [
        "ewa-cv.pdf"
    ]

    confirmed = await _upload_docx(
        app_client,
        app_auth_headers,
        world["stage_id"],
        _client_docx(),
        data={"replace_client_cv": "true"},
    )

    assert confirmed.status_code == 200, confirmed.text
    body = confirmed.json()
    assert (body["replaced"], body["undo_available"]) == ("confirmed", False)
    assert "poprawione przez rekrutera" not in body["client_cv"]["content_html"]


async def test_approved_client_cv_asks_before_replacing(
    app_client: AsyncClient, app_auth_headers: dict[str, str]
) -> None:
    from datetime import datetime, timezone

    world = await _client_world()
    async with AsyncSessionLocal() as db:
        csv = await db.scalar(
            select(CandidateStageCV).where(
                CandidateStageCV.candidate_stage_id == world["stage_id"]
            )
        )
        csv.branded_status = "finalized"
        csv.branded_draft_html = "<p>Zatwierdzone CV</p>"
        csv.branded_finalized_at = datetime(2026, 10, 15, 12, 20, tzinfo=timezone.utc)
        await db.commit()

    asked = await _upload_docx(
        app_client, app_auth_headers, world["stage_id"], _client_docx()
    )

    assert asked.status_code == 409, asked.text
    detail = asked.json()["detail"]
    assert detail["status"] == "finalized"
    # Czas firmy (Europe/Warsaw), nie UTC.
    assert "15.10.2026, 14:20" in detail["message"]


async def test_after_cv_sent_the_file_only_goes_to_the_profile(
    app_client: AsyncClient, app_auth_headers: dict[str, str]
) -> None:
    world = await _client_world(stage="cv_sent")

    res = await _upload_docx(
        app_client, app_auth_headers, world["stage_id"], _client_docx()
    )

    assert res.status_code == 200, res.text
    body = res.json()
    assert (body["outcome"], body["stored_in_profile"], body["client_cv"]) == (
        "client_cv_kept",
        True,
        None,
    )
    assert (await _stage_cv(world["stage_id"])).branded_status == "none"
    docs = await _documents(world["candidate_id"])
    assert [(d.filename, d.is_primary) for d in docs] == [
        ("ewa-cv.pdf", True),
        (CLIENT_CV_NAME, False),
    ]


async def test_after_cv_sent_without_primary_cv_nothing_is_saved(
    app_client: AsyncClient, app_auth_headers: dict[str, str]
) -> None:
    world = await _client_world(stage="cv_sent", primary_cv=False)

    res = await _upload_docx(
        app_client, app_auth_headers, world["stage_id"], _client_docx()
    )

    assert res.status_code == 409, res.text
    assert res.json()["detail"]["code"] == "client_cv_locked_after_send"
    assert await _documents(world["candidate_id"]) == []


async def test_person_without_primary_cv_keeps_the_file_out_of_the_profile(
    app_client: AsyncClient, app_auth_headers: dict[str, str]
) -> None:
    world = await _client_world(primary_cv=False)

    res = await _upload_docx(
        app_client, app_auth_headers, world["stage_id"], _client_docx()
    )

    assert res.status_code == 200, res.text
    body = res.json()
    assert (body["outcome"], body["stored_in_profile"]) == ("client_cv", False)
    assert await _documents(world["candidate_id"]) == []


async def test_cv_of_another_person_is_refused_before_any_write(
    app_client: AsyncClient, app_auth_headers: dict[str, str]
) -> None:
    world = await _client_world()

    res = await _upload_docx(
        app_client,
        app_auth_headers,
        world["stage_id"],
        _client_docx(header="Analityk – Jan Przykładowy"),
        filename="B2B_Analityk_Jan_Przykładowy.docx",
    )

    assert res.status_code == 422, res.text
    detail = res.json()["detail"]
    assert detail["code"] == "client_cv_other_person"
    assert "Jan Przykładowy" in detail["message"]
    assert [d.filename for d in await _documents(world["candidate_id"])] == [
        "ewa-cv.pdf"
    ]
    assert (await _stage_cv(world["stage_id"])).branded_status == "none"


async def test_client_cv_lands_on_the_stage_row_that_holds_the_pairs_cv(
    app_client: AsyncClient, app_auth_headers: dict[str, str]
) -> None:
    from app.models.recruitment_pipeline import PipelineStage

    world = await _client_world(stage="verified")
    await _auto_cv(world)
    async with AsyncSessionLocal() as db:
        later = CandidateStage(
            candidate_id=world["candidate_id"],
            job_id=world["job_id"],
            stage=PipelineStage.interview,
        )
        db.add(later)
        await db.commit()
        later_id = later.id

    res = await _upload_docx(app_client, app_auth_headers, later_id, _client_docx())

    assert res.status_code == 200, res.text
    assert res.json()["client_cv_stage_id"] == world["stage_id"]
    assert (await _stage_cv(world["stage_id"])).branded_render_metadata["origin"][
        "kind"
    ] == "upload"


async def test_use_other_cv_routes_refuse_after_cv_sent(
    app_client: AsyncClient, app_auth_headers: dict[str, str]
) -> None:
    world = await _client_world(stage="cv_sent")
    base = f"/api/candidates/stages/{world['stage_id']}/cv/branded"
    docs = await _documents(world["candidate_id"])

    responses = [
        await app_client.post(
            f"{base}/upload",
            headers=app_auth_headers,
            files={"file": (CLIENT_CV_NAME, _client_docx(), DOCX_MIME)},
            data={"expected_revision": "0"},
        ),
        await app_client.post(
            f"{base}/select-document",
            headers=app_auth_headers,
            json={"document_id": docs[0].id, "expected_revision": 0},
        ),
        await app_client.post(
            f"{base}/select-stage-cv",
            headers=app_auth_headers,
            json={"source_job_id": await _seed_job(), "expected_revision": 0},
        ),
    ]

    for res in responses:
        assert res.status_code == 409, res.text
        assert res.json()["detail"]["code"] == "client_cv_locked_after_send"
    assert (await _stage_cv(world["stage_id"])).branded_status == "none"


# ── Naprawa profili z CV do klienta jako głównym CV (D3) ─────────────────────


async def _profile_with_client_cv_as_primary() -> tuple[int, int, int, str]:
    """Osoba po starym kafelku: „B2B_…docx” główne, własne PDF obok, kopia etapu z B2B."""
    from datetime import timedelta

    from app.core.scheduling import local_now
    from app.models.candidate import Candidate
    from app.models.candidate_document import CandidateDocumentKind
    from app.models.recruitment_pipeline import PipelineStage

    tag = uuid.uuid4().hex[:8]
    client_bytes = _client_docx(marker=tag)
    sha = hashlib.sha256(client_bytes).hexdigest()
    job_id = await _seed_job()
    async with AsyncSessionLocal() as db:
        candidate = Candidate(
            name="Ewa", lastname="Fikcyjna", cv_filename=CLIENT_CV_NAME
        )
        db.add(candidate)
        await db.flush()
        own = CandidateDocument(
            candidate_id=candidate.id,
            filename="ewa-cv.pdf",
            file_content=b"%PDF-1.4 wlasne CV " + tag.encode(),
            document_kind=CandidateDocumentKind.other,
            uploaded_at=local_now() - timedelta(days=30),
        )
        client_cv = CandidateDocument(
            candidate_id=candidate.id,
            filename=CLIENT_CV_NAME,
            file_content=client_bytes,
            content_sha256=sha,
            document_kind=CandidateDocumentKind.cv,
            is_primary=True,
            uploaded_at=local_now(),
        )
        stage = CandidateStage(
            candidate_id=candidate.id, job_id=job_id, stage=PipelineStage.screening
        )
        db.add_all([own, client_cv, stage])
        await db.flush()
        db.add(
            CandidateStageCV(
                candidate_stage_id=stage.id,
                candidate_id=candidate.id,
                job_id=job_id,
                original_cv_filename=CLIENT_CV_NAME,
                original_cv_content=client_bytes,
                original_cv_sha256=sha,
                original_snapshot_source="manual_upload",
            )
        )
        await db.commit()
        return candidate.id, own.id, stage.id, sha


async def test_repair_gives_the_primary_cv_back_to_the_persons_own_cv(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture
) -> None:
    from app.cli import client_cv_primary_repair as repair_cli
    from app.core.scheduling import local_now
    from app.models.candidate import Candidate
    from app.models.candidate_document import CandidateDocumentKind

    monkeypatch.setattr(
        repair_cli, "RECEIPT_KEY", f"test_repair_{uuid.uuid4().hex[:8]}"
    )
    cid, own_id, sid, sha = await _profile_with_client_cv_as_primary()
    lonely, _ = await _stage_with_snapshot(snapshot=False)
    async with AsyncSessionLocal() as db:
        db.add(
            CandidateDocument(
                candidate_id=lonely,
                filename="B2B_Analityk_Bez_Innego_CV.docx",
                file_content=b"docx " + uuid.uuid4().hex.encode(),
                document_kind=CandidateDocumentKind.cv,
                is_primary=True,
                uploaded_at=local_now(),
            )
        )
        await db.commit()
    scope = [cid, lonely]

    # Na sucho: liczby i identyfikatory, nic nie zmienia.
    assert await repair_cli.main([], candidate_ids=scope, reread=False) == 0
    plan = json.loads(capsys.readouterr().out.splitlines()[0])["plan"]
    assert (plan["with_successor"], plan["without_successor"]) == (1, 1)
    assert (plan["candidate_ids"], plan["skipped_candidate_ids"]) == ([cid], [lonely])
    assert plan["stage_copies"] == 1
    assert plan["successor_kinds"] == {"other.pdf": 1}
    assert [d.is_primary for d in await _documents(cid)] == [False, True]

    # Liczba z próby musi się zgadzać.
    assert (
        await repair_cli.main(
            ["--apply", "--expected", "2"], candidate_ids=scope, reread=False
        )
        == 2
    )
    assert (
        await repair_cli.main(
            ["--apply", "--expected", "1"], candidate_ids=scope, reread=False
        )
        == 0
    )
    docs = {d.id: d for d in await _documents(cid)}
    assert [d.filename for d in docs.values() if d.is_primary] == ["ewa-cv.pdf"]
    assert docs[own_id].document_kind.value == "cv"
    csv = await _stage_cv(sid)
    assert csv.original_snapshot_source == "client_cv_repair"
    assert csv.original_cv_filename == "ewa-cv.pdf"
    assert csv.original_cv_sha256 != sha
    async with AsyncSessionLocal() as db:
        candidate = await db.scalar(select(Candidate).where(Candidate.id == cid))
    assert candidate.cv_filename == "ewa-cv.pdf"
    # Osoba bez innego CV zostaje bez zmian.
    assert [d.is_primary for d in await _documents(lonely)] == [True]

    # Paragon blokuje drugi bieg.
    capsys.readouterr()
    assert (
        await repair_cli.main(
            ["--apply", "--expected", "0"], candidate_ids=scope, reread=False
        )
        == 2
    )


async def test_repair_plan_follows_current_cv_order_and_skips_pdfs() -> None:
    """Następca jak w `get_current_cv` (rodzaj „cv” przed resztą, bez daty na
    końcu); PDF „B2B” jest CV kandydata, a pliki po wdrożeniu odcina `until`."""
    from datetime import timedelta

    from app.cli import client_cv_primary_repair as repair_cli
    from app.core.scheduling import local_now
    from app.models.candidate import Candidate
    from app.models.candidate_document import CandidateDocumentKind

    now = local_now()
    async with AsyncSessionLocal() as db:
        people = [Candidate(name="Ewa", lastname="Fikcyjna") for _ in range(2)]
        db.add_all(people)
        await db.flush()
        with_docx, with_pdf = (person.id for person in people)
        tag = uuid.uuid4().hex.encode()
        db.add_all(
            [
                CandidateDocument(
                    candidate_id=with_docx,
                    filename=CLIENT_CV_NAME,
                    file_content=b"docx " + tag,
                    document_kind=CandidateDocumentKind.cv,
                    is_primary=True,
                    uploaded_at=now,
                ),
                # Starszy plik rodzaju „cv” wygrywa z nowszym załącznikiem „other”
                # i z plikiem bez daty.
                CandidateDocument(
                    candidate_id=with_docx,
                    filename="ewa-cv.pdf",
                    file_content=b"%PDF cv " + tag,
                    document_kind=CandidateDocumentKind.cv,
                    uploaded_at=now - timedelta(days=14),
                ),
                CandidateDocument(
                    candidate_id=with_docx,
                    filename="zaswiadczenie.pdf",
                    file_content=b"%PDF other " + tag,
                    document_kind=CandidateDocumentKind.other,
                    uploaded_at=now - timedelta(days=2),
                ),
                CandidateDocument(
                    candidate_id=with_docx,
                    filename="bez-daty.pdf",
                    file_content=b"%PDF none " + tag,
                    document_kind=CandidateDocumentKind.cv,
                ),
                CandidateDocument(
                    candidate_id=with_pdf,
                    filename="B2B_Analityk_Ewa_Fikcyjna.pdf",
                    file_content=b"%PDF b2b " + tag,
                    document_kind=CandidateDocumentKind.cv,
                    is_primary=True,
                    uploaded_at=now,
                ),
                CandidateDocument(
                    candidate_id=with_pdf,
                    filename="ewa-cv.pdf",
                    file_content=b"%PDF own " + tag,
                    document_kind=CandidateDocumentKind.cv,
                    uploaded_at=now - timedelta(days=14),
                ),
            ]
        )
        await db.commit()

    async with AsyncSessionLocal() as db:
        planned = await repair_cli.plan(db, candidate_ids=[with_docx, with_pdf])
        before_upload = await repair_cli.plan(
            db, candidate_ids=[with_docx], until=now - timedelta(hours=1)
        )

    assert [item.candidate_id for item in planned] == [with_docx]
    successor = next(
        d for d in await _documents(with_docx) if d.id == planned[0].successor_id
    )
    assert successor.filename == "ewa-cv.pdf"
    assert before_upload == []
