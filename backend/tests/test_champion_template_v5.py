"""Wzór Word v5 (09.2026): sekcje „Doświadczenie poza stackiem” i „Wiedza z rozmów”.

Formularz czyta ścieżka tabelowa (bez AI). Test wypełnia PUSTY wzór z repo tak,
jak zrobiłby to Delivery Lead, i sprawdza, że nowe wiersze trafiają do
`experience` i `insights`, a v4 wciąż się wczytuje (etykiety, nie numery).
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from app.services.champion_document import experience_from_text, table_profile
from app.services.champion_intake import prepare_profile, preview_document, user_edit

ASSETS = Path(__file__).parents[1] / "app/assets/champion"
BLANK_V5 = ASSETS / "Profil_Championa_v5.0.docx"


def _fill(values: dict[str, str]) -> bytes:
    doc = Document(BLANK_V5)
    for table in doc.tables:
        for row in table.rows:
            cells = row.cells
            label = cells[0].text.strip().lower()
            for prefix, value in values.items():
                if label.startswith(prefix) and len(cells) > 1:
                    cells[1].text = value
            if label.startswith("pytania od delivery"):
                continue
    screening = next(
        t
        for t in doc.tables
        if t.rows[0].cells[0].text.startswith("Pytania od Delivery")
    )
    for i, row in enumerate(screening.rows[1:3], start=1):
        row.cells[0].text = f"Pytanie {i}: Jak testowałeś proces {i}?"
        row.cells[
            1
        ].text = "Idealna odpowiedź: konkretny przykład\n\nDeal breaker: brak"
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


FILLED = {
    "must-have": "Selenium, SQL",
    "nice-to-have": "Postman",
    "dziedzina": "płatności kartowe (min. 2 lata), e-commerce (mile)",
    "certyfikaty": "ISTQB Foundation",
    "regulacje": "PSD2 (mile), PCI DSS",
    "od klienta": "Decyduje lead QA; odrzucali za brak znajomości chargebacku.",
    "insight od": "Zespół 6 osób, dużo spotkań z biznesem.",
}


def test_blank_v5_keeps_every_v4_label_and_adds_the_new_rows() -> None:
    texts = [
        cell.text.strip().lower()
        for table in Document(BLANK_V5).tables
        for row in table.rows
        for cell in row.cells
    ]
    for label in (
        "must-have",
        "nice-to-have",
        "dziedzina",
        "certyfikaty",
        "regulacje",
        "od klienta",
        "insight od",
        "co przekona",
        "historyczne pytania",
    ):
        assert any(t.startswith(label) for t in texts), label


def test_filled_v5_reads_experience_and_notes_without_ai() -> None:
    parsed = table_profile(_fill(FILLED))
    assert parsed is not None
    assert parsed["template_version"] == "5.0"
    profile = prepare_profile(
        parsed["profile"],
        raw_fields=parsed["raw_fields"],
        template_version=parsed["template_version"],
    )
    domains = {d["name"]: d for d in profile["experience"]["domains"]}
    assert domains["płatności kartowe"]["min_years"] == 2
    assert domains["płatności kartowe"]["level"] == "must"
    assert domains["e-commerce"]["level"] == "nice"
    assert [c["name"] for c in profile["experience"]["certifications"]] == [
        "ISTQB Foundation"
    ]
    assert {r["name"]: r["level"] for r in profile["experience"]["regulations"]} == {
        "PSD2": "nice",
        "PCI DSS": "must",
    }
    assert profile["client"]["consultant_insight"].startswith("Zespół 6 osób")
    [note] = profile["insights"]
    assert note["source"] == "client"
    assert note["audience"] == "team"
    assert "chargebacku" in note["text"]
    # Projekt nadal rozpoznany po nazwie nagłówka („5. O projekcie”).
    assert profile["stack"]["must"] == [{"name": "Selenium"}, {"name": "SQL"}]


def test_imported_note_gets_a_server_id_and_the_importer_as_author() -> None:
    parsed = table_profile(_fill(FILLED))
    saved = user_edit({}, parsed["profile"], 42, imported=True, actor_name="DL Import")
    [note] = saved["insights"]
    assert note["id"].startswith("n-")
    assert note["origin"] == "document"
    assert note["author_name"] == "DL Import"


@pytest.mark.asyncio
async def test_preview_of_v5_does_not_call_ai(monkeypatch) -> None:
    async def forbidden(*args, **kwargs):
        raise AssertionError("formularz v5 nie może wołać AI")

    monkeypatch.setattr(
        "app.services.champion_profile_ingest.parse_champion_document", forbidden
    )
    result = await preview_document(_fill(FILLED), "profil.docx")
    assert result["champion_profile"]["experience"]["domains"]


def test_experience_cell_grammar() -> None:
    assert experience_from_text(
        "ubezpieczenia (mile); karty min. 3 lat", with_years=True
    ) == [
        {"name": "ubezpieczenia", "level": "nice", "min_years": None},
        {"name": "karty", "level": "must", "min_years": 3},
    ]
    assert experience_from_text("ISTQB min. 3 lata", with_years=False) == [
        {"name": "ISTQB min. 3 lata", "level": "must", "min_years": None},
    ]
