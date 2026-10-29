"""„Karta przed CV” (Credit Agricole, 29.10.2026): odczyt wzoru, zapis pliku
i wartości — bez bazy. Wzór fikcyjny (``tests/_cover_sheet_docx.py``)."""

from __future__ import annotations

import io
import zipfile
from datetime import date
from decimal import Decimal

import docx
import pytest

from app.services import cover_sheet as cs
from tests._cover_sheet_docx import GUARDIAN_EMAIL, GUARDIAN_LINE, cover_docx

VALUES = {
    cs.NAME: "Anna Przykładowa",
    cs.ROLE: "Inżynier danych",
    cs.LEVEL: "Specjalista",
    cs.EXPERIENCE: "7 lat",
    cs.PROFILE: "Oracle, PL/SQL",
    cs.AVAILABILITY: "12.2026",
    cs.LOCATION: "Kraków",
    cs.RATE: "1 298,64",
}


def _paragraphs(content: bytes) -> list:
    return docx.Document(io.BytesIO(content)).paragraphs


def _text(content: bytes) -> list[str]:
    return [p.text for p in _paragraphs(content)]


# ── Odczyt wzoru ─────────────────────────────────────────────────────────────


def test_reader_finds_every_label_of_the_template() -> None:
    structure = cs.read_template(cover_docx())
    assert [slot.field for slot in structure.slots] == list(cs.FIELD_LABELS)
    experience = structure.slot(cs.EXPERIENCE)
    assert experience is not None and experience.suffix == " doświadczenia komercyjnego"
    rate = structure.slot(cs.RATE)
    assert rate is not None and rate.suffix == " PLN netto za 1 MD"
    assert structure.fields()[0] == {"id": "name", "label": "Imię i nazwisko"}
    texts = [item["text"] for item in structure.layout if item["kind"] == "text"]
    # Nagłówek i blok opiekuna (z łączem) idą do podglądu jako tekst.
    assert texts[0] == "B2B.NET S.A."
    assert texts[-1] == (GUARDIAN_LINE + GUARDIAN_EMAIL).strip()


def test_labels_match_without_case_and_polish_letters() -> None:
    assert cs.fold("Doświadczenie ŁĄKA") == "doswiadczenie laka"
    assert len(cs.fold("Staż: Łódź")) == len("Staż: Łódź")


def test_document_without_the_labels_is_rejected() -> None:
    with pytest.raises(cs.SheetRejected) as exc:
        cs.read_template(cover_docx(labels=False))
    assert "nie rozpoznał" in str(exc.value)


def test_non_docx_and_macro_files_are_rejected_in_polish() -> None:
    with pytest.raises(cs.SheetRejected) as exc:
        cs.read_template(b"%PDF-1.7 to nie jest Word")
    assert str(exc.value) == cs.MSG_NOT_DOCX

    original = zipfile.ZipFile(io.BytesIO(cover_docx()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as package:
        for item in original.infolist():
            package.writestr(item, original.read(item.filename))
        package.writestr("word/vbaProject.bin", b"makro")
    with pytest.raises(cs.SheetRejected) as exc:
        cs.read_template(out.getvalue())
    assert str(exc.value) == cs.MSG_MACROS


def test_too_big_file_is_rejected() -> None:
    with pytest.raises(cs.SheetRejected) as exc:
        cs.read_template(b"0" * (cs.MAX_TEMPLATE_BYTES + 1))
    assert str(exc.value) == cs.MSG_TOO_BIG


# ── Zapis pliku ──────────────────────────────────────────────────────────────


def test_render_replaces_placeholders_and_keeps_the_rest() -> None:
    template = cover_docx()
    rendered = cs.render(template, VALUES)
    texts = _text(rendered)
    assert "Anna Przykładowa" in texts
    assert "Rola: Inżynier danych" in texts
    assert "Poziom doświadczenia: Specjalista" in texts
    assert "Staż: 7 lat doświadczenia komercyjnego" in texts
    assert "Profil kompetencji: Oracle, PL/SQL" in texts
    assert "Dostępność: 12.2026 \nLokalizacja: Kraków" in texts
    assert "Wynagrodzenie: 1 298,64 PLN netto za 1 MD" in texts
    assert not any("xx" in text for text in texts)
    # Blok opiekuna bajt w bajt (z łączem).
    before = [p._p.xml for p in _paragraphs(template)[-2:]]
    after = [p._p.xml for p in _paragraphs(rendered)[-2:]]
    assert before == after


def test_render_keeps_run_formatting() -> None:
    paragraphs = _paragraphs(cs.render(cover_docx(), VALUES))
    name = next(p for p in paragraphs if p.text == "Anna Przykładowa")
    assert all(run.bold for run in name.runs if run.text)
    role = next(p for p in paragraphs if p.text.startswith("Rola:"))
    assert [(run.text, bool(run.bold)) for run in role.runs] == [
        ("Rola:", True),
        (" ", False),
        ("Inżynier danych", False),
    ]


def _two_fields_in_one_paragraph(*, break_in_label_run: bool) -> bytes:
    document = docx.Document()
    document.add_paragraph().add_run("Imie i nazwisko").bold = True
    document.add_paragraph().add_run("Rola:").bold = True
    paragraph = document.add_paragraph()
    paragraph.add_run("Poziom doświadczenia:").bold = True
    if break_in_label_run:
        # Word po Shift+Enter: złamanie i etykieta w jednym przebiegu.
        paragraph.add_run("\nStaż:").bold = True
    else:
        paragraph.add_run().add_break()
        paragraph.add_run("Staż:").bold = True
    paragraph.add_run(" x doświadczenia komercyjnego")
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def test_label_after_a_line_break_inside_the_same_run_is_read() -> None:
    template = _two_fields_in_one_paragraph(break_in_label_run=True)
    fields = [slot.field for slot in cs.read_template(template).slots]
    assert cs.EXPERIENCE in fields
    texts = _text(cs.render(template, {**VALUES, cs.EXPERIENCE: "9 lat"}))
    assert "Poziom doświadczenia: Specjalista\nStaż: 9 lat doświadczenia komercyjnego" in texts


def test_value_added_after_a_bold_label_does_not_shift_the_next_line() -> None:
    # Pusty „Poziom” dostaje nowy przebieg za etykietą — „Staż” w następnym
    # wierszu tego samego akapitu ma nadal trafić we własne przebiegi.
    template = _two_fields_in_one_paragraph(break_in_label_run=False)
    texts = _text(cs.render(template, {**VALUES, cs.EXPERIENCE: "9 lat"}))
    assert "Poziom doświadczenia: Specjalista\nStaż: 9 lat doświadczenia komercyjnego" in texts


@pytest.mark.parametrize("dirty",["a\x01b", "a\x0bb", "a￾b", "a\ud800b", "a\x00b"])
def test_characters_word_cannot_store_are_dropped(dirty: str) -> None:
    # Tytuł albo imię wyciągnięte z PDF-a bywa ze znakiem sterującym — lxml
    # odrzuca go przy zapisie, a DL nie poprawi imienia w oknie karty (500).
    texts = _text(cs.render(cover_docx(), {**VALUES, cs.ROLE: dirty}))
    assert "Rola: ab" in texts
    assert cs.clean_value(f" {dirty} ") == "ab"


def test_empty_value_drops_the_placeholder_and_its_words() -> None:
    texts = _text(cs.render(cover_docx(), {**VALUES, cs.EXPERIENCE: ""}))
    assert "Staż:" in [text.strip() for text in texts]
    assert not any("komercyjnego" in text for text in texts)


def test_whole_amount_keeps_two_decimals() -> None:
    texts = _text(cs.render(cover_docx(), {**VALUES, cs.RATE: "1 300,00"}))
    assert "Wynagrodzenie: 1 300,00 PLN netto za 1 MD" in texts


# ── Wartości ─────────────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    ("years", "level"),
    [
        (None, None),
        (0, "Junior"),
        (2.9, "Junior"),
        (3, "Specjalista"),
        (7, "Specjalista"),
        (7.9, "Specjalista"),
        (8, "Expert"),
        (20, "Expert"),
    ],
)
def test_level_thresholds(years, level) -> None:
    assert cs.level_for_years(years) == level


@pytest.mark.parametrize(
    ("count", "text"),
    [
        (1, "1 rok"),
        (2, "2 lata"),
        (4, "4 lata"),
        (5, "5 lat"),
        (12, "12 lat"),
        (14, "14 lat"),
        (21, "21 lat"),
        (22, "22 lata"),
        (24, "24 lata"),
        (25, "25 lat"),
        (112, "112 lat"),
    ],
)
def test_polish_plural_of_years(count, text) -> None:
    assert cs.experience_text(count) == text


def test_experience_without_years_is_empty() -> None:
    assert cs.experience_text(None) is None


TODAY = date(2026, 10, 29)


def test_availability_from_date_and_from_the_past() -> None:
    kwargs = {"notice_period": 3, "notice_unit": "months", "today": TODAY}
    assert cs.availability_text(available_from=date(2026, 12, 1), **kwargs) == "12.2026"
    # Termin, który minął, to „od teraz”, nie miesiąc z przeszłości.
    assert cs.availability_text(available_from=date(2025, 3, 1), **kwargs) == "10.2026"


def test_availability_from_notice_period() -> None:
    def text(count, unit):
        return cs.availability_text(
            available_from=None, notice_period=count, notice_unit=unit, today=TODAY
        )

    assert text(3, "months") == "01.2027 (3 mies. wypowiedzenia)"
    assert text(2, "weeks") == "11.2026 (2 tyg. wypowiedzenia)"
    assert text(30, None) == "11.2026 (30 dni wypowiedzenia)"


def test_availability_unknown_is_empty() -> None:
    assert (
        cs.availability_text(
            available_from=None, notice_period=None, notice_unit=None, today=TODAY
        )
        is None
    )


def test_md_rate_from_hourly_and_daily() -> None:
    assert cs.md_rate(162.33, "hourly") == Decimal("1298.64")
    assert cs.md_rate(1300, "daily") == Decimal("1300.00")
    assert cs.md_rate(20000, "monthly") is None
    assert cs.md_rate(None, "hourly") is None
    assert cs.md_rate(0, "hourly") is None
    assert cs.amount_text(Decimal("1298.64")) == "1 298,64"
    assert cs.amount_text(Decimal("950")) == "950,00"
    assert cs.amount_text(Decimal("12345.5")) == "12 345,50"


def test_profile_from_bold_terms_of_why_section() -> None:
    html = (
        "<article><h1>Inżynier danych – Anna Przykładowa</h1>"
        '<h2 data-cv-section="why_points">DLACZEGO NASZ KANDYDAT?</h2>'
        "<ul><li>7 lat z <b>Oracle</b> i <strong>PL/SQL</strong></li>"
        "<li>Hurtownie w <b>oracle</b>, procesy <b>ETL</b>.</li></ul>"
        '<h2 data-cv-section="skills">UMIEJĘTNOŚCI</h2>'
        "<p><b>Bazy danych: </b><b>MS SQL</b></p></article>"
    )
    assert cs.profile_from_cv_html(html) == ["Oracle", "PL/SQL", "ETL"]


def test_profile_from_why_heading_without_section_marker() -> None:
    html = "<h2>Dlaczego nasz kandydat?</h2><p><b>Kafka</b></p><h2>Doświadczenie</h2><b>Java</b>"
    assert cs.profile_from_cv_html(html) == ["Kafka"]
    assert cs.profile_from_cv_html(None) == []


def test_profile_fallback_is_must_skills_the_person_has() -> None:
    must = [{"name": "Oracle"}, {"name": "Ab Initio"}, "Python"]
    have = [{"name": "oracle"}, {"name": "Python"}, {"name": "Go"}]
    assert cs.profile_from_skills(must, have) == ["Oracle", "Python"]
    assert cs.profile_from_skills(None, have) == []
