"""Draft validation and gradual adoption contracts; no database or AI needed."""

from copy import deepcopy
from datetime import date
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
import pytest
from docx import Document
from fastapi import HTTPException
from app.services.champion_document import document_text, table_profile
from app.services.champion_intake import (
    prepare_profile,
    validation,
    user_edit,
    fingerprint,
    enforce_operation,
    preview_document,
)

FIXTURES = Path(__file__).parent / "fixtures" / "champion"
BLANK = Path(__file__).parents[1] / "app/assets/champion/Profil_Championa_v4.0.docx"


@pytest.fixture(autouse=True)
def _enable_champion_gate(monkeypatch):
    # The gate defaults OFF in production (advisory only); these tests cover
    # the blocking contract, so turn it on. Individual tests may delenv to
    # exercise the default advisory behavior.
    monkeypatch.setenv("CHAMPION_INTAKE_GATE_ENABLED", "true")


def filled():
    d = table_profile((FIXTURES / "v4-two-questions.docx").read_bytes())
    return prepare_profile(
        d["profile"], raw_fields=d["raw_fields"], template_version=d["template_version"]
    )


def job(cp):
    return SimpleNamespace(
        champion_profile=cp,
        title="Java Developer",
        client_id=1,
        rate_budget_hourly=None,
        onsite_days_per_week=None,
        location=None,
        remote_policy=None,
        matching_requirements=None,
        requirements_reviewed=False,
    )


@pytest.mark.asyncio
async def test_blank_and_two_question_template_without_ai(monkeypatch):
    async def forbidden(*args, **kwargs):
        raise AssertionError("v4 must not call AI")

    monkeypatch.setattr(
        "app.services.champion_profile_ingest.parse_champion_document", forbidden
    )
    blank = await preview_document(BLANK.read_bytes(), BLANK.name)
    assert blank["champion_profile"]["screening_questions"] == []
    assert blank["champion_profile"]["stack"] == {"must": [], "nice": [], "notes": ""}
    assert blank["validation"]["blocked_operations"] == ["cv", "handoff", "search"]
    for name, count in [("v4-two-questions.docx", 2), ("v4-filled.docx", 3)]:
        valid = await preview_document((FIXTURES / name).read_bytes(), "input.docx")
        assert len(valid["champion_profile"]["screening_questions"]) == count
        assert valid["summary"]["onsite_days_per_week"] == 2
        assert valid["validation"]["issues"] == []
        assert valid["champion_profile"]["intake"]["template_version"] == "4.0"
        assert (
            valid["champion_profile"]["intake"]["document_context"]["client_name"]
            == "Klient testowy"
        )
        assert (
            "Rekruter"
            in valid["champion_profile"]["intake"]["document_context"]["prep_owner"]
        )
        assert valid["champion_profile"]["_parser"].startswith("champion_parse:v7:")


def test_order_merged_cells_and_empty_rows():
    d = Document()
    d.add_paragraph("before")
    t = d.add_table(rows=3, cols=2)
    t.cell(0, 0).merge(t.cell(0, 1)).text = "merged"
    t.cell(2, 0).text = "label"
    t.cell(2, 1).text = "answer"
    d.add_paragraph("after")
    stream = BytesIO()
    d.save(stream)
    assert document_text(stream.getvalue()) == "before\nmerged\nlabel | answer\nafter"


@pytest.mark.parametrize(
    "field,value",
    [
        ("rate_value", "120–150 PLN/h"),
        ("rate_value", "150 EUR/h"),
        ("rate_value", "1200 PLN/dzień"),
        ("rate_value", 0),
        ("rate_value", 2001),
        ("onsite_days_per_week", "2 lub 3"),
        ("onsite_days_per_week", 2.5),
        ("onsite_days_per_week", 8),
        ("work_mode", "remote / hybrydowo"),
        ("start_date", "2026-02-30"),
        ("deadline", "za tydzień"),
    ],
)
def test_ambiguous_values_stay_visible_and_null(field, value):
    cp = filled()
    cp["basics"][field] = value
    cp["basics"]["rate_raw"] = None
    cp = prepare_profile(cp)
    assert cp["basics"][field] is None
    assert cp["intake"]["unresolved"][f"basics.{field}"] == str(value)
    assert "search" in validation(cp)["blocked_operations"]


def test_legacy_dates_zero_and_missing_are_distinct():
    cp = filled()
    cp["basics"].update(
        start_date="01.10.2026", deadline="18/09/2026", onsite_days_per_week=0
    )
    cp = prepare_profile(cp)
    assert cp["basics"]["start_date"] == "2026-10-01"
    assert cp["basics"]["onsite_days_per_week"] == 0
    assert "missing_office_days" in {i["code"] for i in validation(cp)["issues"]}
    cp["basics"]["onsite_days_per_week"] = None
    assert prepare_profile(cp)["basics"]["onsite_days_per_week"] is None


def test_skills_alternatives_placeholders_and_versions():
    cp = filled()
    cp["stack"] = {
        "must": "Java 17; JAVA 17\nSpring (MVC, WebFlux), Kafka lub RabbitMQ; do ustalenia",
        "nice": "brak",
        "notes": "",
    }
    cp = prepare_profile(cp)
    assert cp["stack"]["must"] == [
        {"name": "Java 17"},
        {"name": "Spring (MVC, WebFlux)"},
        {"name": "Kafka lub RabbitMQ"},
    ]
    assert cp["stack"]["nice"] == []
    assert cp["intake"]["unresolved"]["stack.must"] == "do ustalenia"
    from app.services.requirement_contract import explicit_contract

    contract = explicit_contract(["Kafka lub RabbitMQ"], [])
    assert len(contract.all_of[0].any_of) == 2


def test_grandfathering_noop_workflow_changes_and_server_stamps():
    legacy = filled()
    legacy.pop("intake")
    legacy["screening_questions"] = []
    assert validation(legacy)["blocked_operations"] == []
    assert user_edit(legacy, deepcopy(legacy), 99)["intake"] is None
    assert (
        user_edit(
            legacy,
            {"intake": {"policy_version": 1}, "verification": {"fake": True}},
            99,
        )["intake"]
        is None
    )
    edited = user_edit(legacy, {"project": {"about": "Nowy projekt"}}, 99)
    assert edited["intake"]["applied_by"] == 99
    assert "cv" in validation(edited)["blocked_operations"]
    assert user_edit(legacy, legacy, 100, imported=True)["intake"]["applied_by"] == 100
    assert user_edit({}, {}, 100)["intake"]["policy_version"] == 1
    assert user_edit(edited, {"intake": None}, 101)["intake"] == edited["intake"]


def test_gate_off_by_default_is_advisory_and_never_blocks(monkeypatch):
    monkeypatch.delenv("CHAMPION_INTAKE_GATE_ENABLED", raising=False)
    cp = filled()
    cp["basics"].update(rate_value=None, rate_raw=None)
    draft = job(cp)
    result = validation(cp, draft)
    # Issues are still surfaced (advisory), but nothing is blocked and
    # enforce_operation is a no-op — pre-#1477 generator behavior.
    assert result["issues"]
    assert result["blocked_operations"] == []
    assert all(i["blocked_operations"] == [] for i in result["issues"])
    enforce_operation(draft, "cv")
    enforce_operation(draft, "search")
    enforce_operation(draft, "handoff")
    enforce_operation(draft, "search", force=True)


def test_draft_operation_specific_gates_and_revalidation():
    cp = filled()
    cp["basics"].update(rate_value=None, rate_raw=None)
    draft = job(cp)
    enforce_operation(draft, "cv")
    with pytest.raises(HTTPException) as blocked:
        enforce_operation(draft, "search")
    assert blocked.value.status_code == 422
    assert (
        blocked.value.detail["validation"]["issues"][0]["path"] == "basics.rate_value"
    )
    draft.rate_budget_hourly = 150
    assert validation(cp, draft)["blocked_operations"] == []
    draft.location = "Kraków"
    assert "column_conflict" in {i["code"] for i in validation(cp, draft)["issues"]}
    cp["intake"] = None
    enforce_operation(draft, "search")
    with pytest.raises(HTTPException):
        enforce_operation(draft, "search", force=True)


def test_fingerprint_changes_on_profile_and_recruitment_edits():
    j = job(filled())
    before = fingerprint(j)
    j.champion_profile["stack"]["notes"] = "Nowe kryteria"
    assert before != fingerprint(j)
    before = fingerprint(j)
    j.rate_budget_hourly = 200
    assert before != fingerprint(j)
    before = fingerprint(j)
    j.deadline = date(2026, 12, 1)
    assert before != fingerprint(j)


def test_job_values_carry_role_deadline_and_office_location():
    from app.services.champion_intake import response_context

    j = job(filled())
    j.title = "Senior Go Developer"
    j.deadline = date(2026, 11, 30)
    j.office_location = "Kraków"
    j.location = "Warszawa"

    job_values = response_context(j)["job_values"]

    assert job_values["role_name"] == "Senior Go Developer"
    assert job_values["deadline"] == "2026-11-30"
    # `office_location` wins over the legacy `location` column, like the
    # `basics.candidate_location_pref` comparison in `validation()`.
    assert job_values["candidate_location_pref"] == "Kraków"


@pytest.mark.asyncio
async def test_oversize_is_not_truncated():
    d = Document()
    d.add_paragraph("x" * 14001)
    stream = BytesIO()
    d.save(stream)
    with pytest.raises(ValueError, match="limit"):
        await preview_document(stream.getvalue(), "large.docx")


def test_cv_context_parity_and_no_fake_screening():
    from app.services.cv_generator_b2b.champion_builder import (
        from_nexus_job,
        parse_champion_from_docx_bytes,
        build_champion_section,
    )

    cp = filled()
    uploaded = parse_champion_from_docx_bytes(
        (FIXTURES / "v4-two-questions.docx").read_bytes(), "input.docx"
    )
    stored = from_nexus_job(cp["stack"]["must"], cp["stack"]["nice"], cp)
    assert build_champion_section(uploaded, "pl") == build_champion_section(
        stored, "pl"
    )
    assert cp["stack"]["notes"] in build_champion_section(uploaded, "pl")
    assert cp["client"]["selling_points"] in build_champion_section(uploaded, "pl")
    assert parse_champion_from_docx_bytes(BLANK.read_bytes(), "blank.docx").is_empty()


def test_legacy_malformed_fields_warn_without_adopting():
    cp = filled()
    cp.pop("intake")
    cp["basics"]["rate_value"] = "100–150 EUR/h"
    assert validation(cp)["blocked_operations"] == []


def test_invalid_edit_is_retained_even_when_canonical_value_was_empty():
    cp = filled()
    cp["basics"]["start_date"] = None
    updated = user_edit(cp, {"basics": {"start_date": "za miesiąc"}}, 9)
    assert updated["basics"]["start_date"] is None
    assert updated["intake"]["unresolved"]["basics.start_date"] == "za miesiąc"


def test_conflicting_skill_columns_block_cv_and_search_until_reconciled():
    from app.services.champion_intake import sync_selected_rubrics

    cp = filled()
    j = job(cp)
    j.must_skills = [{"name": "Python", "level": None}]
    j.nice_skills = []
    assert "cv" in validation(cp, j)["blocked_operations"]
    sync_selected_rubrics(j, cp, ["must"])
    assert validation(cp, j)["blocked_operations"] == []
    assert j.matching_requirements is None
    assert not j.requirements_reviewed


@pytest.mark.asyncio
async def test_parser_does_not_repair_a_truncated_response_into_success(monkeypatch):
    from app.services import champion_profile_ingest as ingest

    async def truncated(*args, **kwargs):
        return SimpleNamespace(stop_reason="max_tokens", content=[])

    monkeypatch.setattr(ingest, "run_in_threadpool", truncated)
    with pytest.raises(ValueError, match="całego dokumentu"):
        await ingest.parse_champion_document("synthetic legacy profile")


def test_ai_legacy_envelope_preserves_zero_and_splits_string_skills():
    from app.services.champion_profile_ingest import build_champion_dict

    cp = prepare_profile(
        build_champion_dict(
            {
                "basics": {"rate_value": 0, "seniority_min_years": 0},
                "stack": {"must": "Python; Java"},
            },
            None,
        )
    )
    assert cp["basics"]["seniority_min_years"] == 0
    assert cp["intake"]["unresolved"]["basics.rate_value"] == "0"
    assert cp["stack"]["must"] == [{"name": "Python"}, {"name": "Java"}]


def test_reviewed_contract_cannot_mask_a_different_profile():
    from app.services.champion_intake import sync_selected_rubrics
    from app.services.requirement_contract import explicit_contract

    cp = filled()
    j = job(cp)
    j.must_skills = cp["stack"]["must"]
    j.matching_requirements = explicit_contract(
        ["Python"], [], reviewed=True
    ).model_dump()
    j.requirements_reviewed = True
    assert "cv" in validation(cp, j)["blocked_operations"]
    sync_selected_rubrics(j, cp, ["must"])
    assert j.matching_requirements is None
    assert validation(cp, j)["blocked_operations"] == []


def test_preview_job_rate_is_numeric_even_when_database_uses_decimal():
    from decimal import Decimal
    from app.services.champion_intake import response_context

    j = job(filled())
    j.rate_budget_hourly = Decimal("170.00")
    value = response_context(j)["job_values"]["rate_value"]
    assert isinstance(value, float)
    assert value == 170.0


def test_vertical_merge_is_read_once_and_does_not_multiply_questions():
    d = Document(FIXTURES / "v4-two-questions.docx")
    questions = next(t for t in d.tables if "Pytania od Delivery" in t.cell(0, 0).text)
    for col in (0, 1):
        first = questions.cell(1, col).text
        questions.cell(1, col).merge(questions.cell(2, col)).text = first
    first_question = questions.cell(1, 0).text
    stream = BytesIO()
    d.save(stream)
    data = stream.getvalue()
    assert document_text(data).count(first_question) == 1
    assert len(table_profile(data)["profile"]["screening_questions"]) == 1
