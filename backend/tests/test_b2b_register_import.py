"""Powtarzalny import rejestru umów z Excela działu — prawdziwa baza, in-process API.

Baza jest wspólna i nieczyszczona: każdy test ma własny zakres numerów
(losowa baza ≥ 20 000 000) i własne, losowe nazwiska; asercje wyłącznie na
własnych wierszach. Dane w plikach są fikcyjne.
"""

from __future__ import annotations

import io
import uuid
from datetime import date

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

import app.models.b2b_register_import  # noqa: F401 — rejestracja tabel w metadanych
from app.core.database import AsyncSessionLocal
from app.models.b2b_generated_contract import B2BGeneratedContract
from app.models.b2b_register_import import B2BRegisterImportRun
from app.models.candidate import Candidate
from app.models.client import Client
from tests.test_b2b_register_import_parser import (
    build_register,
    contract_row,
    no_business_row,
)

BASE = "/api/b2b-generator"
IMPORT = f"{BASE}/register-import"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture(autouse=True)
def _router_mounted():
    """Router rejestruje ``main.py`` — do tego czasu test montuje go sam."""
    from app.api import b2b_register_import
    from app.main import app

    if not any(getattr(r, "path", "") == IMPORT for r in app.routes):
        app.include_router(b2b_register_import.router, prefix=BASE)
    yield


def _base() -> int:
    return 20_000_000 + (uuid.uuid4().int % 900_000_000)


def _tag() -> str:
    """Losowy człon nazwiska — bez kolizji z danymi innych testów."""
    return "Tst" + uuid.uuid4().hex[:10]


async def _upload(
    app_client, headers, payload: bytes, *, dry_run: bool, name="rejestr.xlsx"
):
    return await app_client.post(
        f"{IMPORT}?dry_run={'true' if dry_run else 'false'}",
        headers=headers,
        files={"file": (name, payload, XLSX)},
    )


async def _preview_and_apply(app_client, headers, payload: bytes) -> dict:
    preview = await _upload(app_client, headers, payload, dry_run=True)
    assert preview.status_code == 200, preview.text
    applied = await _upload(app_client, headers, payload, dry_run=False)
    assert applied.status_code == 200, applied.text
    return applied.json()


async def _rows(numbers: list[str]) -> dict[str, B2BGeneratedContract]:
    async with AsyncSessionLocal() as db:
        result = await db.execute(
            select(B2BGeneratedContract).where(
                B2BGeneratedContract.raw_contract_number.in_(numbers)
            )
        )
        return {row.raw_contract_number: row for row in result.scalars().all()}


def _file(base: int, tag: str, **overrides) -> bytes:
    rows = [
        contract_row(
            f"{tag}a Anna", base, signing=date(2025, 3, 1), start=date(2025, 3, 10)
        ),
        contract_row(
            f"{tag}b Jan",
            base + 1,
            signing=date(2025, 4, 1),
            start="nie później niż 01.05.2025",
        ),
        contract_row(
            f"{tag}c Ewa",
            base + 2,
            signing="umowa nie doszła do skutku",
        ),
        contract_row(
            f"{tag}d Karol",
            base + 3,
            signing=date(2023, 1, 5),
            notes="porozumienie stron z dniem 31.03.2024",
            red=True,
        ),
        contract_row(f"{tag}e Leon", base + 4, signing=date(2023, 1, 6), red=True),
    ]
    rows.extend(overrides.get("extra", []))
    return build_register(rows, no_business=overrides.get("no_business"))


async def test_dry_run_writes_no_register_rows(app_client, app_auth_headers):
    base, tag = _base(), _tag()
    payload = _file(base, tag)
    resp = await _upload(app_client, app_auth_headers, payload, dry_run=True)
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["mode"] == "dry_run"
    assert body["counters"]["created"] == 5
    assert body["counters"]["cancelled"] == 1
    assert body["counters"]["closed"] == 1
    assert body["counters"]["likely_ended"] == 1
    assert await _rows([str(base + i) for i in range(5)]) == {}
    async with AsyncSessionLocal() as db:
        run = await db.get(B2BRegisterImportRun, body["run_id"])
    assert run is not None and run.mode == "dry_run"


async def test_apply_requires_preview_of_the_same_file(app_client, app_auth_headers):
    payload = _file(_base(), _tag())
    resp = await _upload(app_client, app_auth_headers, payload, dry_run=False)
    assert resp.status_code == 409, resp.text
    assert "podgląd" in resp.json()["detail"]


async def test_apply_creates_rows_and_reimport_is_idempotent(
    app_client, app_auth_headers
):
    base, tag = _base(), _tag()
    payload = _file(base, tag)
    body = await _preview_and_apply(app_client, app_auth_headers, payload)
    assert body["counters"]["created"] == 5

    rows = await _rows([str(base + i) for i in range(5)])
    anna = rows[str(base)]
    assert anna.source == "excel"
    assert (
        anna.contract_number == f"{base}/2025"
        and anna.seq == base
        and anna.year == 2025
    )
    assert anna.contract_status == "active" and anna.signature_status == "signed_both"
    assert anna.partner_name == f"Anna {tag}a"
    assert anna.render_payload is None
    assert rows[str(base + 1)].start_date_mode == "not_later"
    ewa = rows[str(base + 2)]
    assert ewa.contract_status == "cancelled" and ewa.year is None and ewa.seq is None
    karol = rows[str(base + 3)]
    assert karol.contract_status == "closed"
    assert karol.closure_date == date(2024, 3, 31) and karol.closure_reason == "other"
    leon = rows[str(base + 4)]
    assert leon.contract_status == "active"
    assert "likely_ended" in leon.legacy_data["flags"]

    again = await _preview_and_apply(app_client, app_auth_headers, payload)
    assert again["counters"]["created"] == 0
    assert again["counters"]["updated"] == 0
    assert again["counters"]["unchanged"] == 5


async def test_changed_file_updates_and_flags_missing_rows(
    app_client, app_auth_headers
):
    base, tag = _base(), _tag()
    await _preview_and_apply(app_client, app_auth_headers, _file(base, tag))

    changed = build_register(
        [
            contract_row(
                f"{tag}a Anna",
                base,
                signing=date(2025, 3, 1),
                start=date(2025, 3, 10),
                position="Java developer",
            ),
            contract_row(
                f"{tag}b Jan",
                base + 1,
                signing=date(2025, 4, 1),
                start="nie później niż 01.05.2025",
            ),
        ]
    )
    body = await _preview_and_apply(app_client, app_auth_headers, changed)
    assert body["counters"]["updated"] >= 1
    rows = await _rows([str(base + i) for i in range(5)])
    assert rows[str(base)].position == "Java developer"
    assert rows[str(base + 2)].excel_missing_since is not None
    assert rows[str(base)].excel_missing_since is None


async def test_manual_status_change_in_nexus_survives_reimport(
    app_client, app_auth_headers
):
    base, tag = _base(), _tag()
    payload = _file(base, tag)
    await _preview_and_apply(app_client, app_auth_headers, payload)
    anna = (await _rows([str(base)]))[str(base)]
    patched = await app_client.patch(
        f"{BASE}/generated/{anna.id}",
        headers=app_auth_headers,
        json={
            "contract_status": "closed",
            "closure_reason": "project_completed",
            "closure_date": "2026-01-31",
        },
    )
    assert patched.status_code == 200, patched.text
    await _preview_and_apply(app_client, app_auth_headers, payload)
    anna = (await _rows([str(base)]))[str(base)]
    assert anna.contract_status == "closed"
    assert anna.closure_reason == "project_completed"


async def test_generator_row_with_the_same_number_is_untouched(
    app_client, app_auth_headers
):
    base, tag = _base(), _tag()
    async with AsyncSessionLocal() as db:
        generator = B2BGeneratedContract(
            year=2026,
            seq=base,
            contract_number=f"{base}/2026",
            partner_name="Inna Osoba",
            client_name="Inny Klient",
            language="pl",
            signature_status="unsigned",
            contract_status="in_progress",
        )
        db.add(generator)
        await db.commit()
        await db.refresh(generator)
    payload = build_register(
        [
            contract_row(
                f"{tag} Adam", base, signing=date(2026, 5, 1), client="Klient Testowy"
            )
        ]
    )
    body = await _preview_and_apply(app_client, app_auth_headers, payload)
    assert body["counters"]["created"] == 0
    assert body["counters"]["generator_discrepancies"] == 1
    discrepancy = body["generator_discrepancies"][0]
    assert set(discrepancy["differences"]) == {"partner", "client"}
    async with AsyncSessionLocal() as db:
        fresh = await db.get(B2BGeneratedContract, generator.id)
        count = await db.scalar(
            select(func.count(B2BGeneratedContract.id)).where(
                B2BGeneratedContract.seq == base
            )
        )
    assert fresh.partner_name == "Inna Osoba" and fresh.source == "generator"
    assert count == 1


async def test_rollback_deletes_created_and_restores_updated(
    app_client, app_auth_headers
):
    base, tag = _base(), _tag()
    await _preview_and_apply(app_client, app_auth_headers, _file(base, tag))
    extra = contract_row(f"{tag}f Nowa", base + 5, signing=date(2025, 6, 1))
    changed = build_register(
        [
            contract_row(
                f"{tag}a Anna",
                base,
                signing=date(2025, 3, 1),
                start=date(2025, 3, 10),
                position="PM",
            ),
            extra,
        ]
    )
    body = await _preview_and_apply(app_client, app_auth_headers, changed)
    run_id = body["run_id"]

    runs = await app_client.get(f"{IMPORT}/runs", headers=app_auth_headers)
    assert runs.status_code == 200
    mine = next(r for r in runs.json() if r["id"] == run_id)
    assert mine["can_rollback"] is True

    resp = await app_client.post(
        f"{IMPORT}/runs/{run_id}/rollback", headers=app_auth_headers
    )
    assert resp.status_code == 200, resp.text
    rows = await _rows([str(base), str(base + 2), str(base + 5)])
    assert str(base + 5) not in rows
    assert rows[str(base)].position == "tester manualny"
    assert rows[str(base + 2)].excel_missing_since is None

    again = await app_client.post(
        f"{IMPORT}/runs/{run_id}/rollback", headers=app_auth_headers
    )
    assert again.status_code == 409


async def test_rollback_keeps_a_status_changed_by_hand_after_import(
    app_client, app_auth_headers
):
    """Ręczna zmiana statusu w NEXUSIE wygrywa z cofnięciem, jak z plikiem."""
    base, tag = _base(), _tag()
    body = await _preview_and_apply(app_client, app_auth_headers, _file(base, tag))
    anna = (await _rows([str(base)]))[str(base)]
    patched = await app_client.patch(
        f"{BASE}/generated/{anna.id}",
        headers=app_auth_headers,
        json={
            "contract_status": "closed",
            "closure_reason": "project_completed",
            "closure_date": "2026-01-31",
        },
    )
    assert patched.status_code == 200, patched.text
    resp = await app_client.post(
        f"{IMPORT}/runs/{body['run_id']}/rollback", headers=app_auth_headers
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["kept_created"] >= 1
    rows = await _rows([str(base), str(base + 1)])
    assert rows[str(base)].contract_status == "closed"
    assert str(base + 1) not in rows


async def test_row_without_number_keeps_its_key_when_dates_are_filled_in(
    app_client, app_auth_headers
):
    """Dział uzupełnia datę podpisania później — to nadal ten sam wiersz."""
    tag = _tag()
    first = build_register(
        [contract_row(f"{tag} Ola", "bez numeru", signing=None, start=None)]
    )
    await _preview_and_apply(app_client, app_auth_headers, first)
    second = build_register(
        [
            contract_row(
                f"{tag} Ola",
                "bez numeru",
                signing=date(2026, 2, 1),
                start=date(2026, 2, 15),
            )
        ]
    )
    body = await _preview_and_apply(app_client, app_auth_headers, second)
    assert body["counters"]["created"] == 0
    async with AsyncSessionLocal() as db:
        count = await db.scalar(
            select(func.count(B2BGeneratedContract.id)).where(
                B2BGeneratedContract.partner_name.ilike(f"%{tag}%")
            )
        )
    assert count == 1


async def test_generator_refuses_a_number_already_in_the_excel_register(
    app_client, app_auth_headers
):
    """Excel bez daty podpisania ma numer bez roku — ten sam numer porządkowy
    w dowolnym roku jest już wydany (numeracja ciągła między latami)."""
    base, tag = _base(), _tag()
    payload = build_register(
        [contract_row(f"{tag} Zenon", base, signing=None, start=None)]
    )
    await _preview_and_apply(app_client, app_auth_headers, payload)
    resp = await app_client.post(
        "/api/b2b-generator/render?format=docx",
        headers=app_auth_headers,
        json={
            "language": "pl",
            "contract_number": f"{base}/2026",
            "partner_name": "Inna Osoba",
            "client_name": "Nordea Bank",
            "signing_date": "2026-08-01",
        },
    )
    assert resp.status_code == 409, resp.text
    assert "Excela" in resp.json()["detail"]


async def test_next_seq_counts_excel_numbers(app_client, app_auth_headers):
    from app.api.b2b_contract_generator import _next_seq

    base, tag = _base() + 900_000_000, _tag()
    payload = build_register(
        [contract_row(f"{tag} Max", base, signing=None, start=None)]
    )
    await _preview_and_apply(app_client, app_auth_headers, payload)
    async with AsyncSessionLocal() as db:
        suggested = await _next_seq(db)
    assert suggested > base


async def test_excel_row_is_read_only_in_the_register(app_client, app_auth_headers):
    base, tag = _base(), _tag()
    await _preview_and_apply(app_client, app_auth_headers, _file(base, tag))
    rows = await _rows([str(base), str(base + 4)])
    anna = rows[str(base)]

    listed = await app_client.get(
        f"{BASE}/generated",
        headers=app_auth_headers,
        params={"q": f"{base}/2025", "source": "excel"},
    )
    assert listed.status_code == 200, listed.text
    item = next(i for i in listed.json() if i["id"] == anna.id)
    assert item["source"] == "excel"
    assert item["raw_contract_number"] == str(base)
    assert item["can_edit"] is False and item["can_delete"] is False
    assert item["can_download"] is False and item["can_confirm_signed"] is False
    assert item["can_change_status"] is True

    leon = await app_client.get(
        f"{BASE}/generated", headers=app_auth_headers, params={"q": f"{base + 4}/2023"}
    )
    assert "likely_ended" in leon.json()[0]["legacy_flags"]

    deleted = await app_client.delete(
        f"{BASE}/generated/{anna.id}", headers=app_auth_headers
    )
    assert deleted.status_code == 409
    form = await app_client.get(
        f"{BASE}/generated/{anna.id}/form", headers=app_auth_headers
    )
    assert form.status_code == 409
    rename = await app_client.patch(
        f"{BASE}/generated/{anna.id}",
        headers=app_auth_headers,
        json={"client_name": "X"},
    )
    assert rename.status_code == 409


async def test_candidate_client_and_annex_matching(app_client, app_auth_headers):
    base, tag = _base(), _tag()
    async with AsyncSessionLocal() as db:
        client = Client(name=f"Klient {tag}")
        db.add(client)
        single = Candidate(
            name="Jedna", lastname=f"{tag}x", email=f"{tag}x@example.com"
        )
        twin_a = Candidate(
            name="Bliźniak", lastname=f"{tag}y", email=f"{tag}ya@example.com"
        )
        twin_b = Candidate(
            name="Bliźniak", lastname=f"{tag}y", email=f"{tag}yb@example.com"
        )
        db.add_all([single, twin_a, twin_b])
        await db.commit()
        for obj in (client, single):
            await db.refresh(obj)
    payload = build_register(
        [
            contract_row(
                f"{tag}x Jedna",
                base,
                client=f"Klient {tag}",
                signing=date(2025, 1, 2),
                start=date(2025, 1, 6),
            ),
            contract_row(
                f"{tag}y Bliźniak",
                base + 1,
                client=f"Klient {tag}",
                signing=date(2025, 1, 3),
            ),
            contract_row(
                f"{tag}z Nikt",
                base + 2,
                client=f"Nieznany {tag}",
                signing=date(2025, 1, 4),
            ),
        ],
        no_business=[
            no_business_row(
                f"Jedna {tag}x",
                date(2025, 1, 6),
                client=f"Klient {tag}",
                annex="aneks zrobiony 03.03.2025",
            )
        ],
    )
    body = await _preview_and_apply(app_client, app_auth_headers, payload)
    rows = await _rows([str(base), str(base + 1), str(base + 2)])
    assert rows[str(base)].candidate_id == single.id
    assert rows[str(base)].client_id == client.id
    assert rows[str(base)].needs_business_data_annex is True
    assert rows[str(base)].business_data_annex_done_at == date(2025, 3, 3)
    assert rows[str(base + 1)].candidate_id is None
    assert any(e["number"] == str(base + 1) for e in body["ambiguous_candidates"])
    assert any(e["text"] == f"Nieznany {tag}" for e in body["unknown_clients"])
    assert body["counters"]["annex_matched"] == 1


async def test_export_has_register_headers(app_client, app_auth_headers):
    resp = await app_client.get(
        f"{BASE}/generated/export.xlsx", headers=app_auth_headers
    )
    assert resp.status_code == 200, resp.text
    sheet = load_workbook(io.BytesIO(resp.content), read_only=True).active
    header = [c.value for c in next(sheet.iter_rows(max_row=1))]
    assert header[:2] == ["NAZWISKO, PÓŹNIEJ IMIE", "Numer umowy"]
    assert header[13] == "ZMIANY W UMOWIE" and len(header) == 14


async def test_non_xlsx_upload_is_refused(app_client, app_auth_headers):
    resp = await app_client.post(
        f"{IMPORT}?dry_run=true",
        headers=app_auth_headers,
        files={"file": ("rejestr.csv", b"a;b", "text/csv")},
    )
    assert resp.status_code == 422
