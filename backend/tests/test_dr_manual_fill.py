"""Ręczne uzupełnienie statystyk sprzed 12.10.2026 — bez bazy.

* `scripts.import_dr_manual_fill.parse_payload`: granica 12.10.2026 (od niej
  liczy NEXUS), październik DL = 1–11.10, przed 2026 klienci tylko z Rady,
  tydzień rekrutera (dzień roboczy, jeden wpis na tydzień i miesiąc),
  odmowa całości przy błędnym wierszu;
* `scripts.dr_manual_fill_sheet`: szablon z zakładką „Braki” i arkusz → JSON
  z samymi id kont.

Na bazie (CI): `test_dr_manual_fill_db.py`.
"""

from __future__ import annotations

from datetime import date

import openpyxl
import pytest

from scripts.dr_manual_fill_sheet import (
    SheetError,
    build_template,
    sheet_to_payload,
)
from scripts.import_dr_manual_fill import InputError, parse_payload


def _week(**overrides):
    return {
        "user_id": 86,
        "report_date": "2025-03-03",
        "verifications": 6,
        "recommendations": 4,
        "interviews": 2,
        "placements": 1,
        "days_worked": 5,
        **overrides,
    }


def _refused(payload, match: str) -> None:
    with pytest.raises(InputError, match=match):
        parse_payload(payload)


def test_parse_accepts_every_section_before_the_cutover() -> None:
    parsed = parse_payload(
        {
            "weeks": [_week(), _week(report_date="2026-10-05", interviews=None)],
            "placements": [
                {"user_id": 86, "placement_date": "2026-03-05", "client": " Nordea "}
            ],
            "dl_months": [
                {"user_id": 14, "report_month": "2026-03", "placements": 2},
                # Październik 2026 = dni 1–11.10 — wolno.
                {"user_id": 14, "report_month": "2026-10-01", "placements": 1},
            ],
            "board_months": [{"report_month": "2025-03", "placements": 12}],
            "board_clients": [
                {
                    "report_month": "2025-03",
                    "client_name": "Nordea",
                    "placement_count": 4,
                }
            ],
        }
    )
    assert parsed["weeks"][0]["week_number"] == 10
    # Puste pole zostaje `None` — import zostawi zapisaną wartość.
    assert parsed["weeks"][1]["interviews"] is None
    assert parsed["placements"] == [
        {
            "user_id": 86,
            "placement_date": date(2026, 3, 5),
            "client": "Nordea",
            "count": 1,
        }
    ]
    assert [d["report_month"] for d in parsed["dl_months"]] == [
        date(2026, 3, 1),
        date(2026, 10, 1),
    ]
    assert parsed["dl_months"][0]["requests"] is None


@pytest.mark.parametrize(
    "payload",
    [
        {"weeks": [_week(report_date="2026-10-12")]},
        {
            "placements": [
                {"user_id": 86, "placement_date": "2026-10-13", "client": "Nordea"}
            ]
        },
        {"dl_months": [{"user_id": 14, "report_month": "2026-11", "placements": 1}]},
    ],
)
def test_parse_refuses_days_counted_by_nexus(payload) -> None:
    _refused(payload, "od 2026-10-12 liczy NEXUS")


def test_parse_keeps_clients_before_2026_in_the_board_report() -> None:
    _refused(
        {
            "placements": [
                {"user_id": 86, "placement_date": "2025-12-15", "client": "Nordea"}
            ]
        },
        "przed 2026 klientów liczy raport Rady",
    )
    _refused(
        {"board_months": [{"report_month": "2026-01", "placements": 13}]},
        "od 2026 placementy firmy",
    )
    _refused(
        {
            "board_clients": [
                {"report_month": "2026-02", "client_name": "X", "placement_count": 1}
            ]
        },
        "od 2026 placementy firmy",
    )


def test_parse_checks_the_shape_of_a_week() -> None:
    _refused({"weeks": [_week(report_date="2025-03-01")]}, "dzień wolny")
    _refused({"weeks": [_week(), _week()]}, r"ten sam wpis co weeks\[0\]")
    _refused(
        {"weeks": [_week(), _week(report_date="2025-03-04")]},
        "drugi wpis tego samego tygodnia",
    )
    _refused({"weeks": [_week(days_worked=8)]}, "najwyżej 7 dni")
    _refused(
        {
            "weeks": [
                _week(
                    verifications=None,
                    recommendations=None,
                    interviews=None,
                    placements=None,
                    days_worked=None,
                )
            ]
        },
        "bez żadnej liczby",
    )
    # Tydzień przez koniec miesiąca to dwa wpisy — jak w DynaReporterze.
    parsed = parse_payload(
        {"weeks": [_week(report_date="2025-09-29"), _week(report_date="2025-10-01")]}
    )
    assert len(parsed["weeks"]) == 2


def test_parse_refuses_the_whole_input_and_lists_every_error() -> None:
    with pytest.raises(InputError) as exc:
        parse_payload(
            {
                "weeks": [
                    _week(user_id=True),
                    _week(placements=-1),
                    _week(report_date="03.03.2025"),
                ],
                "dl_months": [{"user_id": 14, "report_month": "2026-03-02"}],
            }
        )
    message = str(exc.value)
    assert "weeks[0]: brak id konta" in message
    assert "weeks[1]: placements" in message
    assert "weeks[2]: report_date" in message
    assert "dl_months[0]: report_month" in message


def test_parse_refuses_unknown_sections_and_empty_input() -> None:
    _refused({"tygodnie": []}, "nieznane sekcje")
    _refused({"weeks": []}, "brak wierszy")
    _refused([], "obiekt JSON")


# ── Arkusz ──────────────────────────────────────────────────────────────


GAPS = {
    "cutover": "2026-10-12",
    "months": [
        {"month": "2025-01", "weeks": 0, "dl_entries": 0, "board_placements": 16},
        {"month": "2025-08", "weeks": 0, "dl_entries": 6, "board_placements": 11},
        {"month": "2026-01", "weeks": 89, "dl_entries": 7, "board_placements": 13},
    ],
    "dl_missing": [
        {"user_id": 2, "month": "2026-03"},
        {"user_id": 2, "month": "2026-10"},
    ],
    "recruiter_october_missing": [{"user_id": 1, "report_date": "2026-10-05"}],
}
LISTS = {
    "people": [
        {"id": 1, "name": "Rekruter Testowy", "role": "recruiter", "active": True},
        {"id": 2, "name": "Lead Testowy", "role": "delivery_lead", "active": True},
        # Ta sama osoba z importu Traffita — zostaje konto aktywne.
        {"id": 140, "name": "Lead Testowy", "role": "recruiter", "active": False},
        {"id": 150, "name": "Była Osoba", "role": "recruiter", "active": False},
    ],
    "clients": ["Klient Testowy"],
}


@pytest.fixture
def template(tmp_path):
    path = tmp_path / "uzupelnienie.xlsx"
    build_template(GAPS, LISTS, path)
    return path


def test_template_lists_gaps_and_prepares_rows(template) -> None:
    wb = openpyxl.load_workbook(template)
    gaps = list(wb["Braki"].iter_rows(min_row=2, values_only=True))
    kinds = [g[0] for g in gaps]
    # 2025-01 i 2025-08 bez tygodni rekruterów; wpisy DL brakują tylko w 2025-01.
    assert kinds.count("Tygodnie rekruterów (KPI)") == 2
    assert gaps[2][:3] == ("Wpisy Delivery Leadów", "wszyscy DL", "2025-01")
    assert kinds.count("Wpisy Delivery Leadów") == 1
    assert "Wpis Delivery Leada 1–11.10.2026" in kinds
    assert kinds[-1] == "Bez tabeli w DynaReporterze"
    dl = list(wb["Miesiące DL"].iter_rows(min_row=2, max_row=3, values_only=True))
    assert [(r[0], r[1]) for r in dl] == [
        ("Lead Testowy", "2026-03"),
        ("Lead Testowy", "2026-10"),
    ]
    people = [r[:2] for r in wb["Osoby"].iter_rows(min_row=2, values_only=True)]
    assert people == [
        ("Lead Testowy", 2),
        ("Rekruter Testowy", 1),
        ("Była Osoba (nieaktywne)", 150),
    ]


def test_sheet_to_payload_skips_prepared_rows_left_blank(template) -> None:
    wb = openpyxl.load_workbook(template)
    wb["Miesiące DL"]["D2"] = 3  # placementy za 2026-03
    weeks = wb["Tygodnie rekruterów"]
    weeks["A3"], weeks["B3"], weeks["F3"] = "Była Osoba (nieaktywne)", "2025-03-03", 2
    weeks["A4"], weeks["B4"], weeks["C4"] = 1, date(2025, 3, 10), 4.0
    board = wb["Rada — klienci"]
    board["A2"], board["B2"], board["C2"] = date(2025, 3, 1), "Klient Testowy", 4
    wb.save(template)

    payload, skipped = sheet_to_payload(template)
    assert skipped["dl_months"] == 1 and skipped["weeks"] == 1
    assert payload["dl_months"] == [
        {
            "user_id": 2,
            "report_month": "2026-03",
            "requests": None,
            "placements": 3,
            "vacancies": None,
            "open_requests": None,
            "open_vacancies": None,
        }
    ]
    assert [(w["user_id"], w["report_date"]) for w in payload["weeks"]] == [
        (150, "2025-03-03"),
        (1, "2025-03-10"),
    ]
    assert payload["weeks"][1]["verifications"] == 4
    assert payload["board_clients"] == [
        {
            "report_month": "2025-03",
            "client_name": "Klient Testowy",
            "placement_count": 4,
        }
    ]
    # Wynik przechodzi walidację importu — i nie niesie nazwisk.
    parse_payload(payload)
    assert "Testowy" not in str(payload["weeks"])


def test_sheet_to_payload_names_rows_with_errors(template) -> None:
    wb = openpyxl.load_workbook(template)
    weeks = wb["Tygodnie rekruterów"]
    weeks["A3"], weeks["B3"], weeks["F3"] = "Ktoś Spoza Listy", "2025-03-03", 1
    weeks["A4"], weeks["F4"] = "Rekruter Testowy", 1
    weeks["A5"], weeks["B5"], weeks["F5"] = "Rekruter Testowy", "2025-03-17", "dwa"
    wb.save(template)
    with pytest.raises(SheetError) as exc:
        sheet_to_payload(template)
    message = str(exc.value)
    assert "Tygodnie rekruterów, wiersz 3: osoby nie ma" in message
    assert "wiersz 4: brak — Pierwszy dzień wpisu" in message
    assert "wiersz 5: Placementy — liczba" in message
    # Komunikaty nie powtarzają nazwisk z arkusza.
    assert "Spoza" not in message
