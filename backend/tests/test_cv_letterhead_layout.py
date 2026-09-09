"""Brand artwork must not wrap body words or lose its page origin."""

from pathlib import Path
from lxml import etree
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml.ns import qn
from docx.shared import Cm
from app.services.cv_generator_b2b.docx_renderer import normalize_letterhead_layout

TEMPLATE = Path(__file__).parents[1] / "app/templates/cv/szablon_firmowy.docx"


def test_real_template_reserves_artwork_and_preserves_images_and_position():
    doc = Document(TEMPLATE)
    anchor = doc.sections[0].header._element.xpath(".//wp:anchor")[0]
    horizontal = etree.tostring(anchor.find(qn("wp:positionH")))
    images = [
        p.blob for p in doc.part.package.parts if p.content_type.startswith("image/")
    ]
    text = [p.text for p in doc.paragraphs]
    normalize_letterhead_layout(doc)
    assert doc.sections[0].top_margin.cm > 4.6
    assert anchor.find(qn("wp:wrapTight")) is None
    assert anchor.find(qn("wp:wrapNone")) is not None
    assert etree.tostring(anchor.find(qn("wp:positionH"))) == horizontal
    assert [
        p.blob for p in doc.part.package.parts if p.content_type.startswith("image/")
    ] == images
    assert [p.text for p in doc.paragraphs] == text
    first = doc._element.xml, doc.sections[0].header._element.xml
    normalize_letterhead_layout(doc)
    assert (doc._element.xml, doc.sections[0].header._element.xml) == first


def test_inherited_header_keeps_page_origin_across_different_margins():
    doc = Document(TEMPLATE)
    doc.add_section(WD_SECTION_START.NEW_PAGE).left_margin = Cm(4)
    normalize_letterhead_layout(doc)
    for section in doc.sections:
        assert section.top_margin.cm > 4.6
        h = section.header._element.xpath(".//wp:positionH")[0]
        assert h.get("relativeFrom") == "page"
        assert h.find(qn("wp:posOffset")).text == "-46990"


def test_plain_document_and_foreground_logo_are_unchanged():
    plain = Document()
    original = plain._element.xml
    normalize_letterhead_layout(plain)
    assert plain._element.xml == original
    doc = Document(TEMPLATE)
    anchor = doc.sections[0].header._element.xpath(".//wp:anchor")[0]
    anchor.set("behindDoc", "0")
    original = doc._element.xml, doc.sections[0].header._element.xml
    normalize_letterhead_layout(doc)
    assert (doc._element.xml, doc.sections[0].header._element.xml) == original


def test_missing_page_size_does_not_crash_or_reposition_artwork():
    doc = Document(TEMPLATE)
    section = doc.sections[0]
    section._sectPr.remove(section._sectPr.find(qn("w:pgSz")))
    original = doc._element.xml, section.header._element.xml
    normalize_letterhead_layout(doc)
    assert (doc._element.xml, section.header._element.xml) == original
