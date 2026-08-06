"""Integration contract for the scope-shaped Clients directory."""

from __future__ import annotations

import uuid
from datetime import date, datetime, timedelta, timezone
from io import BytesIO

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import delete, select

from app.core.database import AsyncSessionLocal
from app.core.security import create_access_token, hash_password
from app.models.activity import Activity
from app.models.candidate import Candidate
from app.models.client import Client, ClientStatus
from app.models.client_directory import (
    ClientAlias,
    ClientPortfolioScope,
    PortfolioCategory,
)
from app.models.client_framework_contract import (
    ClientFrameworkContract,
    FrameworkContractStatus,
)
from app.models.contract import Contract, ContractStatus
from app.models.user import User, UserRole

pytestmark = [pytest.mark.asyncio, pytest.mark.integration]


async def _seed_directory() -> dict[str, object]:
    suffix = uuid.uuid4().hex[:10]
    today = date.today()
    async with AsyncSessionLocal() as db:
        alpha = Client(
            name=f"Source Alpha {suffix}",
            display_name=f"Alpha {suffix}",
            legal_name=f"Alpha Legal {suffix} S.A.",
            industry="Banking",
            status=ClientStatus.active,
        )
        zulu = Client(
            name=f"Source Zulu {suffix}",
            display_name=f"Zulu {suffix}",
            legal_name=f"Zulu Legal {suffix} S.A.",
            industry="Engineering",
            status=ClientStatus.prospect,
        )
        hidden = Client(
            name=f"Hidden {suffix}",
            hidden=True,
            status=ClientStatus.active,
        )
        db.add_all((alpha, zulu, hidden))
        await db.flush()

        msa = ClientFrameworkContract(
            client_id=zulu.id,
            name=f"MSA {suffix}",
            status=FrameworkContractStatus.active,
            effective_date=today - timedelta(days=30),
            expiry_date=None,
        )
        db.add(msa)
        await db.flush()

        scopes = [
            ClientPortfolioScope(
                client_id=alpha.id,
                category=PortfolioCategory.active,
                label=None,
                source_system="test",
                source_key=f"{suffix}:alpha",
            ),
            ClientPortfolioScope(
                client_id=zulu.id,
                framework_contract_id=msa.id,
                category=PortfolioCategory.active,
                label="B scope",
                source_system="test",
                source_key=f"{suffix}:zulu-b",
            ),
            ClientPortfolioScope(
                client_id=zulu.id,
                category=PortfolioCategory.active,
                label="A scope",
                source_system="test",
                source_key=f"{suffix}:zulu-a",
            ),
            ClientPortfolioScope(
                client_id=zulu.id,
                category=PortfolioCategory.relationship,
                label="Pentesty",
                source_system="test",
                source_key=f"{suffix}:zulu-rel",
            ),
            ClientPortfolioScope(
                client_id=hidden.id,
                category=PortfolioCategory.active,
                source_system="test",
                source_key=f"{suffix}:hidden",
            ),
        ]
        db.add_all(scopes)
        archived_alias = f"Archived Alias {suffix}"
        db.add_all(
            [
                ClientAlias(
                    client_id=zulu.id,
                    alias=f"Legacy Alias {suffix}",
                    normalized_alias=f"legacy alias {suffix}",
                    source_system="test",
                    source_key=f"{suffix}:alias",
                ),
                ClientAlias(
                    client_id=zulu.id,
                    alias=archived_alias,
                    normalized_alias=archived_alias.lower(),
                    source_system="test",
                    source_key=f"{suffix}:archived-alias",
                    archived_at=datetime.now(timezone.utc),
                ),
            ]
        )

        candidate_one = Candidate(name="Anna", lastname=f"One-{suffix}")
        candidate_two = Candidate(name="Jan", lastname=f"Two-{suffix}")
        candidate_future = Candidate(name="Future", lastname=f"Three-{suffix}")
        candidate_ended = Candidate(name="Ended", lastname=f"Four-{suffix}")
        db.add_all((candidate_one, candidate_two, candidate_future, candidate_ended))
        await db.flush()
        contracts = [
            Contract(
                candidate_id=candidate_one.id,
                client_id=zulu.id,
                status=ContractStatus.active,
                start_date=today - timedelta(days=10),
                end_date=None,
            ),
            # Same candidate twice must still count once.
            Contract(
                candidate_id=candidate_one.id,
                client_id=zulu.id,
                status=ContractStatus.ending,
                start_date=today - timedelta(days=5),
                end_date=today + timedelta(days=5),
            ),
            Contract(
                candidate_id=candidate_two.id,
                client_id=zulu.id,
                status=ContractStatus.ending,
                start_date=today,
                end_date=today,
            ),
            Contract(
                candidate_id=candidate_future.id,
                client_id=zulu.id,
                status=ContractStatus.active,
                start_date=today + timedelta(days=1),
                end_date=None,
            ),
            Contract(
                candidate_id=candidate_ended.id,
                client_id=zulu.id,
                status=ContractStatus.active,
                start_date=today - timedelta(days=20),
                end_date=today - timedelta(days=1),
            ),
        ]
        db.add_all(contracts)
        await db.commit()
        return {
            "suffix": suffix,
            "archived_alias": archived_alias,
            "client_ids": [alpha.id, zulu.id, hidden.id],
            "candidate_ids": [
                candidate_one.id,
                candidate_two.id,
                candidate_future.id,
                candidate_ended.id,
            ],
            "msa_id": msa.id,
        }


async def _cleanup_directory(seed: dict[str, object]) -> None:
    client_ids = seed["client_ids"]
    candidate_ids = seed["candidate_ids"]
    async with AsyncSessionLocal() as db:
        await db.execute(delete(Contract).where(Contract.client_id.in_(client_ids)))
        await db.execute(
            delete(ClientAlias).where(ClientAlias.client_id.in_(client_ids))
        )
        await db.execute(
            delete(ClientPortfolioScope).where(
                ClientPortfolioScope.client_id.in_(client_ids)
            )
        )
        await db.execute(
            delete(ClientFrameworkContract).where(
                ClientFrameworkContract.client_id.in_(client_ids)
            )
        )
        await db.execute(delete(Candidate).where(Candidate.id.in_(candidate_ids)))
        await db.execute(delete(Client).where(Client.id.in_(client_ids)))
        await db.commit()


@pytest.mark.parametrize(
    ("portfolio_category", "expected_category"),
    [
        (None, PortfolioCategory.active),
        (PortfolioCategory.relationship.value, PortfolioCategory.relationship),
    ],
)
async def test_create_client_atomically_creates_initial_portfolio_scope(
    app_client: AsyncClient,
    app_auth_headers: dict[str, str],
    portfolio_category: str | None,
    expected_category: PortfolioCategory,
) -> None:
    name = f"Atomic client {uuid.uuid4().hex[:10]}"
    client_id: int | None = None
    params = (
        {"portfolio_category": portfolio_category}
        if portfolio_category is not None
        else None
    )

    try:
        response = await app_client.post(
            "/api/clients",
            params=params,
            json={
                "name": name,
                "industry": "Test",
                "status": "inactive",
            },
            headers=app_auth_headers,
        )
        assert response.status_code == 201, response.text
        body = response.json()
        client_id = body["id"]
        # Portfolio category is independent of the legacy client status.
        assert body["status"] == "inactive"

        async with AsyncSessionLocal() as db:
            scopes = list(
                (
                    await db.scalars(
                        select(ClientPortfolioScope).where(
                            ClientPortfolioScope.client_id == client_id
                        )
                    )
                ).all()
            )
            assert len(scopes) == 1
            assert scopes[0].category == expected_category
            assert scopes[0].source_system == "manual"
            assert scopes[0].framework_contract_id is None
    finally:
        if client_id is not None:
            async with AsyncSessionLocal() as db:
                await db.execute(
                    delete(Activity).where(
                        Activity.entity_type == "client",
                        Activity.entity_id == client_id,
                    )
                )
                await db.execute(
                    delete(ClientPortfolioScope).where(
                        ClientPortfolioScope.client_id == client_id
                    )
                )
                await db.execute(delete(Client).where(Client.id == client_id))
                await db.commit()


async def test_create_client_rejects_unknown_portfolio_category_without_writes(
    app_client: AsyncClient,
    app_auth_headers: dict[str, str],
) -> None:
    name = f"Invalid portfolio category {uuid.uuid4().hex[:10]}"

    response = await app_client.post(
        "/api/clients",
        params={"portfolio_category": "not-a-category"},
        json={"name": name},
        headers=app_auth_headers,
    )

    assert response.status_code == 422, response.text
    async with AsyncSessionLocal() as db:
        assert await db.scalar(select(Client.id).where(Client.name == name)) is None


async def test_directory_requires_auth(app_client: AsyncClient) -> None:
    response = await app_client.get("/api/clients/directory")
    assert response.status_code == 401


async def test_directory_counts_scopes_sorts_and_counts_consultants(
    app_client: AsyncClient,
    app_auth_headers: dict[str, str],
) -> None:
    seed = await _seed_directory()
    try:
        response = await app_client.get(
            "/api/clients/directory",
            params={
                "category": "active",
                "q": seed["suffix"],
                "page": 1,
                "page_size": 50,
            },
            headers=app_auth_headers,
        )
        assert response.status_code == 200, response.text
        body = response.json()
        assert body["total_rows"] == 3
        assert body["total_clients"] == 2
        assert body["page_size"] == 50
        assert body["category_counts"]["active"] >= 2
        assert body["category_counts"]["relationship"] >= 1

        items = body["items"]
        assert [item["display_name"].split()[0] for item in items] == [
            "Alpha",
            "Zulu",
            "Zulu",
        ]
        assert [item["scope_label"] for item in items[1:]] == [
            "A scope",
            "B scope",
        ]
        assert {item["active_consultants_count"] for item in items[1:]} == {2}
        assert items[1]["client_status"] == "prospect"
        assert items[2]["effective_date"] is not None
        assert items[2]["expiry_date"] is None
        assert items[2]["msa_id"] == seed["msa_id"]
    finally:
        await _cleanup_directory(seed)


@pytest.mark.parametrize("needle_kind", ["alias", "legal", "industry", "scope"])
async def test_directory_search_is_scoped_to_selected_category(
    app_client: AsyncClient,
    app_auth_headers: dict[str, str],
    needle_kind: str,
) -> None:
    seed = await _seed_directory()
    suffix = seed["suffix"]
    needles = {
        "alias": f"Legacy Alias {suffix}",
        "legal": f"Zulu Legal {suffix}",
        "industry": "Engineering",
        "scope": "A scope",
    }
    try:
        active = await app_client.get(
            "/api/clients/directory",
            params={"category": "active", "q": needles[needle_kind]},
            headers=app_auth_headers,
        )
        assert active.status_code == 200, active.text
        assert active.json()["total_clients"] == 1
        assert all(
            item["display_name"].startswith("Zulu") for item in active.json()["items"]
        )

        relationship = await app_client.get(
            "/api/clients/directory",
            params={"category": "relationship", "q": needles[needle_kind]},
            headers=app_auth_headers,
        )
        assert relationship.status_code == 200, relationship.text
        if needle_kind in {"alias", "legal", "industry"}:
            assert relationship.json()["total_rows"] == 1
            assert relationship.json()["items"][0]["scope_label"] == "Pentesty"
        else:
            assert relationship.json()["total_rows"] == 0
    finally:
        await _cleanup_directory(seed)


async def test_directory_search_ignores_archived_aliases(
    app_client: AsyncClient,
    app_auth_headers: dict[str, str],
) -> None:
    seed = await _seed_directory()
    try:
        response = await app_client.get(
            "/api/clients/directory",
            params={"category": "active", "q": seed["archived_alias"]},
            headers=app_auth_headers,
        )
        assert response.status_code == 200, response.text
        assert response.json()["total_rows"] == 0
        assert response.json()["total_clients"] == 0
    finally:
        await _cleanup_directory(seed)


async def test_scope_crud_archives_instead_of_deleting(
    app_client: AsyncClient,
    app_auth_headers: dict[str, str],
) -> None:
    seed = await _seed_directory()
    client_id = seed["client_ids"][0]
    try:
        created = await app_client.post(
            f"/api/clients/{client_id}/portfolio-scopes",
            json={"category": "inactive", "label": "Manual"},
            headers=app_auth_headers,
        )
        assert created.status_code == 201, created.text
        scope_id = created.json()["id"]

        updated = await app_client.patch(
            f"/api/clients/{client_id}/portfolio-scopes/{scope_id}",
            json={"category": "relationship", "label": "Relacja"},
            headers=app_auth_headers,
        )
        assert updated.status_code == 200, updated.text
        assert updated.json()["category"] == "relationship"

        archived = await app_client.delete(
            f"/api/clients/{client_id}/portfolio-scopes/{scope_id}",
            headers=app_auth_headers,
        )
        assert archived.status_code == 204, archived.text
        async with AsyncSessionLocal() as db:
            scope = await db.get(ClientPortfolioScope, scope_id)
            assert scope is not None
            assert scope.archived_at is not None
    finally:
        await _cleanup_directory(seed)


async def test_scope_crud_rejects_null_category_foreign_and_occupied_msa(
    app_client: AsyncClient,
    app_auth_headers: dict[str, str],
) -> None:
    seed = await _seed_directory()
    alpha_id, zulu_id, _hidden_id = seed["client_ids"]
    msa_id = seed["msa_id"]
    try:
        foreign_msa = await app_client.post(
            f"/api/clients/{alpha_id}/portfolio-scopes",
            json={
                "category": "active",
                "framework_contract_id": msa_id,
            },
            headers=app_auth_headers,
        )
        assert foreign_msa.status_code == 422, foreign_msa.text

        occupied_msa = await app_client.post(
            f"/api/clients/{zulu_id}/portfolio-scopes",
            json={
                "category": "active",
                "framework_contract_id": msa_id,
            },
            headers=app_auth_headers,
        )
        assert occupied_msa.status_code == 409, occupied_msa.text

        created = await app_client.post(
            f"/api/clients/{alpha_id}/portfolio-scopes",
            json={"category": "inactive", "label": "Validate null"},
            headers=app_auth_headers,
        )
        assert created.status_code == 201, created.text
        scope_id = created.json()["id"]

        null_category = await app_client.patch(
            f"/api/clients/{alpha_id}/portfolio-scopes/{scope_id}",
            json={"category": None},
            headers=app_auth_headers,
        )
        assert null_category.status_code == 422, null_category.text
    finally:
        await _cleanup_directory(seed)


# ── Placement overrides (manifest-invisible curation) ─────────────────────────


async def _scope_id_by_source_key(source_key: str) -> int:
    async with AsyncSessionLocal() as db:
        scope = (
            await db.execute(
                select(ClientPortfolioScope).where(
                    ClientPortfolioScope.source_key == source_key
                )
            )
        ).scalar_one()
        return scope.id


async def test_placement_override_moves_tab_without_touching_base(
    app_client: AsyncClient,
    app_auth_headers: dict[str, str],
) -> None:
    seed = await _seed_directory()
    suffix = seed["suffix"]
    alpha_id = seed["client_ids"][0]
    scope_id = await _scope_id_by_source_key(f"{suffix}:alpha")
    try:
        moved = await app_client.patch(
            f"/api/clients/{alpha_id}/portfolio-scopes/{scope_id}/placement",
            json={"category": "inactive"},
            headers=app_auth_headers,
        )
        assert moved.status_code == 200, moved.text
        assert moved.json()["category_override"] == "inactive"
        # The manifest base column is untouched — the invariant still sees active.
        assert moved.json()["category"] == "active"

        async with AsyncSessionLocal() as db:
            scope = await db.get(ClientPortfolioScope, scope_id)
            assert scope.category == PortfolioCategory.active
            assert scope.category_override == PortfolioCategory.inactive

        # The directory now lists alpha under Nieaktywni, not Aktywni.
        active = await app_client.get(
            "/api/clients/directory",
            params={"category": "active", "q": f"Alpha {suffix}"},
            headers=app_auth_headers,
        )
        assert active.status_code == 200
        assert all(item["client_id"] != alpha_id for item in active.json()["items"])

        inactive = await app_client.get(
            "/api/clients/directory",
            params={"category": "inactive", "q": f"Alpha {suffix}"},
            headers=app_auth_headers,
        )
        assert inactive.status_code == 200
        rows = [i for i in inactive.json()["items"] if i["client_id"] == alpha_id]
        assert len(rows) == 1
        assert rows[0]["category"] == "inactive"
        assert rows[0]["category_override"] == "inactive"
        # The manifest/base category is exposed unchanged for the UI.
        assert rows[0]["category_base"] == "active"
        # The badge (Client.status) is a different field — still active.
        assert rows[0]["client_status"] == "active"

        # Clearing the override returns the row to the manifest tab.
        cleared = await app_client.patch(
            f"/api/clients/{alpha_id}/portfolio-scopes/{scope_id}/placement",
            json={"category": None},
            headers=app_auth_headers,
        )
        assert cleared.status_code == 200, cleared.text
        assert cleared.json()["category_override"] is None
        back = await app_client.get(
            "/api/clients/directory",
            params={"category": "active", "q": f"Alpha {suffix}"},
            headers=app_auth_headers,
        )
        assert any(item["client_id"] == alpha_id for item in back.json()["items"])
    finally:
        await _cleanup_directory(seed)


async def test_placement_override_pins_and_validates_contract_dates(
    app_client: AsyncClient,
    app_auth_headers: dict[str, str],
) -> None:
    seed = await _seed_directory()
    suffix = seed["suffix"]
    alpha_id = seed["client_ids"][0]
    scope_id = await _scope_id_by_source_key(f"{suffix}:alpha")
    try:
        # Alpha's scope has no linked MSA — dates are empty until pinned.
        pinned = await app_client.patch(
            f"/api/clients/{alpha_id}/portfolio-scopes/{scope_id}/placement",
            json={"contract_start": "2026-05-11", "contract_end": "2027-12-31"},
            headers=app_auth_headers,
        )
        assert pinned.status_code == 200, pinned.text
        assert pinned.json()["contract_start_override"] == "2026-05-11"

        listing = await app_client.get(
            "/api/clients/directory",
            params={"category": "active", "q": f"Alpha {suffix}"},
            headers=app_auth_headers,
        )
        row = next(i for i in listing.json()["items"] if i["client_id"] == alpha_id)
        assert row["effective_date"] == "2026-05-11"
        assert row["expiry_date"] == "2027-12-31"
        assert row["msa_id"] is None  # dates come from the override, not an MSA

        # End before start is rejected without persisting.
        bad = await app_client.patch(
            f"/api/clients/{alpha_id}/portfolio-scopes/{scope_id}/placement",
            json={"contract_start": "2027-01-01", "contract_end": "2026-01-01"},
            headers=app_auth_headers,
        )
        assert bad.status_code == 422, bad.text

        # Partial update: only the end date is sent, but it precedes the ALREADY
        # pinned start (2026-05-11). The prospective pair is validated against the
        # stored start, so this is 422 too — and the existing override is intact.
        partial_bad = await app_client.patch(
            f"/api/clients/{alpha_id}/portfolio-scopes/{scope_id}/placement",
            json={"contract_end": "2026-01-01"},
            headers=app_auth_headers,
        )
        assert partial_bad.status_code == 422, partial_bad.text
        async with AsyncSessionLocal() as db:
            scope = await db.get(ClientPortfolioScope, scope_id)
            assert scope.contract_start_override == date(2026, 5, 11)
            assert scope.contract_end_override == date(2027, 12, 31)
    finally:
        await _cleanup_directory(seed)


async def test_base_scope_patch_guards_manifest_owned_scope(
    app_client: AsyncClient,
    app_auth_headers: dict[str, str],
) -> None:
    seed = await _seed_directory()
    alpha_id = seed["client_ids"][0]
    suffix = seed["suffix"]
    manifest_scope_id = await _scope_id_by_source_key(f"{suffix}:alpha")
    try:
        # A manifest-owned (non-manual) scope refuses in-place category edits …
        blocked = await app_client.patch(
            f"/api/clients/{alpha_id}/portfolio-scopes/{manifest_scope_id}",
            json={"category": "inactive"},
            headers=app_auth_headers,
        )
        assert blocked.status_code == 409, blocked.text

        # … but the label stays freely editable on the same scope.
        relabel = await app_client.patch(
            f"/api/clients/{alpha_id}/portfolio-scopes/{manifest_scope_id}",
            json={"label": "Etykieta"},
            headers=app_auth_headers,
        )
        assert relabel.status_code == 200, relabel.text

        # A genuinely manual scope may still change category in place.
        created = await app_client.post(
            f"/api/clients/{alpha_id}/portfolio-scopes",
            json={"category": "inactive", "label": "Manual"},
            headers=app_auth_headers,
        )
        assert created.status_code == 201, created.text
        manual_scope_id = created.json()["id"]
        moved = await app_client.patch(
            f"/api/clients/{alpha_id}/portfolio-scopes/{manual_scope_id}",
            json={"category": "relationship"},
            headers=app_auth_headers,
        )
        assert moved.status_code == 200, moved.text
        assert moved.json()["category"] == "relationship"
    finally:
        await _cleanup_directory(seed)


async def test_archive_guards_manifest_owned_scope(
    app_client: AsyncClient,
    app_auth_headers: dict[str, str],
) -> None:
    """Archiving a manifest scope would drop it from the invariant join (the
    2026-08-03 outage). Only manual scopes may be archived from the API."""
    seed = await _seed_directory()
    alpha_id = seed["client_ids"][0]
    suffix = seed["suffix"]
    manifest_scope_id = await _scope_id_by_source_key(f"{suffix}:alpha")
    try:
        blocked = await app_client.delete(
            f"/api/clients/{alpha_id}/portfolio-scopes/{manifest_scope_id}",
            headers=app_auth_headers,
        )
        assert blocked.status_code == 409, blocked.text
        async with AsyncSessionLocal() as db:
            scope = await db.get(ClientPortfolioScope, manifest_scope_id)
            assert scope.archived_at is None  # untouched

        created = await app_client.post(
            f"/api/clients/{alpha_id}/portfolio-scopes",
            json={"category": "inactive", "label": "Manual"},
            headers=app_auth_headers,
        )
        assert created.status_code == 201, created.text
        manual_scope_id = created.json()["id"]
        archived = await app_client.delete(
            f"/api/clients/{alpha_id}/portfolio-scopes/{manual_scope_id}",
            headers=app_auth_headers,
        )
        assert archived.status_code == 204, archived.text
    finally:
        await _cleanup_directory(seed)


# ── Directory export (CSV / XLSX) ─────────────────────────────────────────────

_XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


async def _seed_role_user(role: UserRole) -> tuple[int, dict[str, str]]:
    """Seed a user with ``role`` and return (id, bearer headers)."""
    unique = uuid.uuid4().hex[:8]
    async with AsyncSessionLocal() as db:
        user = User(
            email=f"dir-export-{role.value}-{unique}@example.com",
            password_hash=hash_password(f"T3st_{unique}!Export"),
            name=f"DirExport {role.value}",
            role=role,
            is_active=True,
        )
        db.add(user)
        await db.commit()
        await db.refresh(user)
        user_id = user.id
    token = create_access_token(user_id, role.value)
    return user_id, {"Authorization": f"Bearer {token}"}


async def _delete_user(user_id: int) -> None:
    async with AsyncSessionLocal() as db:
        await db.execute(delete(User).where(User.id == user_id))
        await db.commit()


def _read_sheet(content: bytes) -> tuple[list[str], list[dict]]:
    """Return (header, records) from the first worksheet of an xlsx payload."""
    workbook = load_workbook(BytesIO(content))
    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    header = [str(cell) if cell is not None else "" for cell in rows[0]]
    records = [dict(zip(header, row)) for row in rows[1:]]
    return header, records


async def test_directory_export_requires_auth(app_client: AsyncClient) -> None:
    response = await app_client.get("/api/clients/directory/export")
    assert response.status_code == 401


async def test_directory_export_xlsx_mirrors_list_and_includes_legal(
    app_client: AsyncClient,
    app_auth_headers: dict[str, str],
) -> None:
    seed = await _seed_directory()
    alpha_id = seed["client_ids"][0]
    try:
        # NIP/REGON round-trip through the export for a privileged role.
        async with AsyncSessionLocal() as db:
            alpha = await db.get(Client, alpha_id)
            alpha.nip = "1234567890"
            alpha.regon = "998877665"
            await db.commit()

        response = await app_client.get(
            "/api/clients/directory/export",
            params={"category": "active", "q": seed["suffix"]},
            headers=app_auth_headers,
        )
        assert response.status_code == 200, response.text
        assert response.headers["content-type"].startswith(_XLSX_MEDIA_TYPE)
        assert "attachment" in response.headers["content-disposition"]
        assert ".xlsx" in response.headers["content-disposition"]

        header, records = _read_sheet(response.content)
        # Admin (legal-privileged) sees the base + legal columns.
        assert header[:9] == [
            "ID klienta",
            "Klient",
            "Zakres",
            "Kategoria",
            "Branża",
            "Status klienta",
            "Aktywni konsultanci",
            "Start umowy ramowej",
            "Koniec umowy ramowej",
        ]
        assert header[9:] == ["Nazwa prawna", "NIP", "REGON"]

        # "Export what I see": same three active scopes the list returns.
        assert len(records) == 3

        alpha_row = next(r for r in records if r["ID klienta"] == alpha_id)
        assert alpha_row["Klient"].startswith("Alpha")
        assert alpha_row["Kategoria"] == "Aktywny"
        assert alpha_row["Status klienta"] == "Aktywny"
        assert alpha_row["Branża"] == "Banking"
        assert alpha_row["Nazwa prawna"].startswith("Alpha Legal")
        assert alpha_row["NIP"] == "1234567890"
        assert alpha_row["REGON"] == "998877665"

        # The Zulu "B scope" carries an MSA with no expiry → open-ended.
        b_scope = next(
            r
            for r in records
            if r["Klient"].startswith("Zulu") and r["Zakres"] == "B scope"
        )
        assert b_scope["Status klienta"] == "Prospekt"
        assert b_scope["Koniec umowy ramowej"] == "Bezterminowa"
        assert b_scope["Start umowy ramowej"]  # non-empty ISO date
        assert b_scope["Aktywni konsultanci"] == 2
    finally:
        await _cleanup_directory(seed)


async def test_directory_export_csv_has_bom_and_header(
    app_client: AsyncClient,
    app_auth_headers: dict[str, str],
) -> None:
    seed = await _seed_directory()
    try:
        response = await app_client.get(
            "/api/clients/directory/export",
            params={"category": "active", "q": seed["suffix"], "format": "csv"},
            headers=app_auth_headers,
        )
        assert response.status_code == 200, response.text
        assert response.headers["content-type"].startswith("text/csv")
        assert ".csv" in response.headers["content-disposition"]
        text = response.content.decode("utf-8-sig")
        assert response.content.startswith(b"\xef\xbb\xbf")  # UTF-8 BOM for Excel
        lines = [line for line in text.splitlines() if line]
        assert lines[0].split(",")[0] == "ID klienta"
        assert "Nazwa prawna" in lines[0]
        # header + three active scopes
        assert len(lines) == 4
    finally:
        await _cleanup_directory(seed)


async def test_directory_export_hides_legal_columns_for_non_privileged_role(
    app_client: AsyncClient,
) -> None:
    seed = await _seed_directory()
    user_id, headers = await _seed_role_user(UserRole.recruiter)
    try:
        response = await app_client.get(
            "/api/clients/directory/export",
            params={"category": "active", "q": seed["suffix"]},
            headers=headers,
        )
        assert response.status_code == 200, response.text
        header, records = _read_sheet(response.content)
        assert header == [
            "ID klienta",
            "Klient",
            "Zakres",
            "Kategoria",
            "Branża",
            "Status klienta",
            "Aktywni konsultanci",
            "Start umowy ramowej",
            "Koniec umowy ramowej",
        ]
        assert "Nazwa prawna" not in header
        assert "NIP" not in header
        assert "REGON" not in header
        # The rows themselves are still visible — only legal columns are gated.
        assert len(records) == 3
    finally:
        await _cleanup_directory(seed)
        await _delete_user(user_id)


async def test_directory_export_neutralises_formula_injection(
    app_client: AsyncClient,
    app_auth_headers: dict[str, str],
) -> None:
    suffix = uuid.uuid4().hex[:10]
    payload = f"=SUM(1+9)*cmd|{suffix}"
    client_ids: list[int] = []
    async with AsyncSessionLocal() as db:
        evil = Client(
            name=f"Evil {suffix}",
            display_name=payload,
            legal_name=f"+ATTACK {suffix}",
            industry="@formula",
            status=ClientStatus.active,
        )
        db.add(evil)
        await db.flush()
        db.add(
            ClientPortfolioScope(
                client_id=evil.id,
                category=PortfolioCategory.active,
                label="-danger",
                source_system="test",
                source_key=f"{suffix}:evil",
            )
        )
        client_ids.append(evil.id)
        await db.commit()
    try:
        response = await app_client.get(
            "/api/clients/directory/export",
            params={"category": "active", "q": suffix},
            headers=app_auth_headers,
        )
        assert response.status_code == 200, response.text
        _, records = _read_sheet(response.content)
        row = next(r for r in records if r["ID klienta"] == client_ids[0])
        # Every free-text cell that would start a formula is defused with a
        # leading apostrophe; the raw value is otherwise preserved.
        assert row["Klient"] == "'" + payload
        assert row["Zakres"] == "'-danger"
        assert row["Branża"] == "'@formula"
        assert row["Nazwa prawna"] == "'+ATTACK " + suffix
    finally:
        async with AsyncSessionLocal() as db:
            await db.execute(
                delete(ClientPortfolioScope).where(
                    ClientPortfolioScope.client_id.in_(client_ids)
                )
            )
            await db.execute(delete(Client).where(Client.id.in_(client_ids)))
            await db.commit()


async def test_directory_export_flags_truncation_at_limit(
    app_client: AsyncClient,
    app_auth_headers: dict[str, str],
) -> None:
    seed = await _seed_directory()  # 3 active scopes match the suffix
    try:
        capped = await app_client.get(
            "/api/clients/directory/export",
            params={"category": "active", "q": seed["suffix"], "limit": 2},
            headers=app_auth_headers,
        )
        assert capped.status_code == 200, capped.text
        assert capped.headers.get("x-export-truncated") == "true"

        full = await app_client.get(
            "/api/clients/directory/export",
            params={"category": "active", "q": seed["suffix"], "limit": 100},
            headers=app_auth_headers,
        )
        assert full.status_code == 200, full.text
        assert "x-export-truncated" not in full.headers
    finally:
        await _cleanup_directory(seed)
