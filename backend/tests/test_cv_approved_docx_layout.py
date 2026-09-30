"""CV z rekrutacji (szkic, zatwierdzone, DL, Cpro) w układzie szablonu firmowego.

Zgłoszenie 30.09.2026: „szkic” z edytora CV etapu był „rozjechany” w Wordzie —
kropki wpisane jako tekst, bez linii między sekcjami i bez tabeli edukacji —
a to samo CV pobrane z profilu (generator) wyglądało poprawnie. Ten sam
konwerter robił zatwierdzone CV wysyłane klientowi. Testy pilnują, że Word
z HTML-a edytora ma układ generatora i że żadne słowo z edytora nie ginie.
"""

import re
from collections import Counter
from io import BytesIO

import pytest
from docx import Document
from lxml import html

from app.services.cv_approved_docx import render_approved_docx
from app.services.cv_document_assets import default_template
from app.services.html_sanitizer import sanitize_cv_html

RODO = "Wyrażam zgodę na przetwarzanie moich danych osobowych."
JOBS = [
    (
        "Senior QA Engineer",
        "12.2025 - obecnie",
        "Acme",
        "Technologie",
        ["Rozwój testów w <strong>Python</strong>", "Utrzymanie CI w Jenkins"],
        "<strong>Python</strong>, Jenkins",
    ),
    (
        "QA Engineer",
        "02.2023 - 11.2025",
        "Beta",
        "Telekomunikacja",
        ["Automatyzacja testów API"],
        "Pytest",
    ),
]


def _editor_html() -> str:
    """Kształt po edytorze (Tiptap): akapity w listach, bez div i span."""
    html_ = (
        "<h1>Tester QA – Anna Testowa</h1><hr>"
        '<h2 data-cv-section="why_points">DLACZEGO NASZ KANDYDAT?</h2>'
        "<ul><li><p>5 lat w <strong>Python</strong></p></li><li><p>Automatyzacja</p></li></ul>"
        '<h2 data-cv-section="education">EDUKACJA</h2>'
        "<p>2015 - 2019 <strong>Politechnika Testowa</strong> — Informatyka · Wrocław</p>"
        '<h2 data-cv-section="skills">UMIEJĘTNOŚCI</h2>'
        "<p><strong>Języki: </strong><strong>Python</strong>, SQL</p>"
        '<h2 data-cv-section="languages">JĘZYKI</h2><p>Polski - C2 · Angielski - B2</p>'
        '<h2 data-cv-section="experience">DOŚWIADCZENIE</h2>'
    )
    for position, dates, company, industry, duties, tech in JOBS:
        html_ += (
            f'<p data-cv-section="role"><strong>{position}</strong> {dates}</p>'
            f'<p data-cv-section="employer">{company} · {industry}</p>'
            '<p data-cv-section="duties_label">Zakres zadań:</p><ul>'
            + "".join(f"<li><p>{duty}</p></li>" for duty in duties)
            + f'</ul><p data-cv-section="technologies">Technologie: {tech}</p>'
        )
    return html_ + f'<p data-cv-section="rodo">{RODO}</p>'


def _generator_html() -> str:
    """Kształt z generatora (zapisany szkic, który pobierają DL i Cpro)."""
    html_ = (
        '<article class="cv"><h1>Tester QA – Anna Testowa</h1><hr>'
        '<h2 data-cv-section="why_points">DLACZEGO NASZ KANDYDAT?</h2>'
        "<ul><li>5 lat w <b>Python</b></li><li>Automatyzacja</li></ul>"
        '<h2 data-cv-section="education">EDUKACJA</h2>'
        '<div class="edu"><span class="edu-dates">2015 - 2019</span> '
        "<span><b>Politechnika Testowa</b> — Informatyka · Wrocław</span></div>"
        '<h2 data-cv-section="skills">UMIEJĘTNOŚCI</h2>'
        '<p class="skill"><b>Języki: </b><b>Python</b>, SQL</p>'
        '<h2 data-cv-section="languages">JĘZYKI</h2><p>Polski - C2 · Angielski - B2</p>'
        '<h2 data-cv-section="experience">DOŚWIADCZENIE</h2>'
    )
    for index, (position, dates, company, industry, duties, tech) in enumerate(JOBS):
        html_ += (
            f'<div class="job" id="exp-{index}">'
            f'<p class="job-head" data-cv-section="role"><b>{position}</b> '
            f'<span class="job-dates">{dates}</span></p>'
            f'<p class="job-co" data-cv-section="employer">{company} · {industry}</p>'
            '<p class="lbl" data-cv-section="duties_label">Zakres zadań:</p><ul>'
            + "".join(f"<li>{duty}</li>" for duty in duties)
            + '</ul><p class="lbl" data-cv-section="technologies">Technologie: '
            f'<span class="tech">{tech}</span></p></div>'
        )
    return html_ + f'<p class="rodo">{RODO}</p></article>'


def _unmarked_html() -> str:
    """Starsze zatwierdzone CV: bez znaczników sekcji i bez linii pod tytułem."""
    return re.sub(r' data-cv-section="[a-z_]+"', "", _editor_html()).replace("<hr>", "")


def _render(content: str) -> Document:
    return Document(
        BytesIO(render_approved_docx(sanitize_cv_html(content), default_template()))
    )


def _is_rule(paragraph) -> bool:
    return bool(paragraph._p.xpath(".//*[local-name()='rect']"))


def _bullet_id(paragraph):
    num_pr = paragraph._p.pPr.numPr if paragraph._p.pPr is not None else None
    return None if num_pr is None else num_pr.numId.val


def _words(text: str) -> Counter:
    return Counter(re.findall(r"[0-9A-Za-zÀ-ž+#]+", text))


def _texts(element) -> list[str]:
    return element.xpath(".//*[local-name()='t']/text()")


def _document_text(doc: Document) -> str:
    return " ".join(_texts(doc.element.body))


SHAPES = [_editor_html, _generator_html, _unmarked_html]


@pytest.mark.parametrize("shape", SHAPES, ids=["editor", "generator", "unmarked"])
def test_recruitment_cv_gets_the_corporate_layout(shape):
    doc = _render(shape())
    paragraphs = doc.paragraphs
    texts = [p.text for p in paragraphs]

    # Każdy tytuł sekcji stoi pod czerwoną linią szablonu.
    for heading in ("DLACZEGO NASZ KANDYDAT?", "UMIEJĘTNOŚCI", "DOŚWIADCZENIE"):
        position = texts.index(heading)
        assert _is_rule(paragraphs[position - 1]), heading

    # Punktory to lista Worda z szablonu (czerwony kwadrat), nie wpisane „• ".
    for bullet in ("5 lat w Python", "Języki: Python, SQL", "Automatyzacja testów API"):
        assert _bullet_id(paragraphs[texts.index(bullet)]) == 1, bullet
    assert not any(text.startswith("• ") for text in texts)
    # Języki: jeden punktor na język, jak w generatorze.
    assert _bullet_id(paragraphs[texts.index("Polski - C2")]) == 1
    assert _bullet_id(paragraphs[texts.index("Angielski - B2")]) == 1

    # Edukacja w tabeli generatora: daty | uczelnia (pogrubiona) + kierunek.
    table = doc.tables[0]
    assert [c.text for c in table.rows[0].cells] == ["Daty", "Nazwa uczelni/kierunek"]
    dates, school = table.rows[1].cells
    assert dates.text == "2015 - 2019"
    assert school.text == "Politechnika Testowa\nInformatyka · Wrocław"
    assert [r.text for r in school.paragraphs[0].runs if r.bold] == [
        "Politechnika Testowa"
    ]

    # Blok stanowiska w kolejności generatora: daty, firma, stanowisko, zadania.
    start = texts.index("12.2025 - obecnie")
    assert texts[start : start + 4] == [
        "12.2025 - obecnie",
        "Nazwa firmy: Acme · Technologie",
        "Stanowisko: Senior QA Engineer",
        "Zakres zadań:",
    ]
    assert all(run.bold for run in paragraphs[start].runs)
    assert all(run.underline for run in paragraphs[start + 3].runs)
    tech = paragraphs[texts.index("Technologie: Python, Jenkins")]
    assert [r.text for r in tech.runs if r.bold] == ["Technologie:", "Python"]
    # Linia oddziela kolejne stanowiska (nie pierwsze).
    second = texts.index("02.2023 - 11.2025")
    assert _is_rule(paragraphs[second - 1])
    assert not _is_rule(paragraphs[start - 1])

    # Klauzula RODO przypięta do dołu ostatniej strony, jak w generatorze.
    boxes = doc.element.body.xpath(".//*[local-name()='txbxContent']")
    assert ["".join(_texts(box)) for box in boxes] == [RODO]


@pytest.mark.parametrize("shape", SHAPES, ids=["editor", "generator", "unmarked"])
def test_no_word_from_the_editor_is_lost(shape):
    source = sanitize_cv_html(shape())
    expected = _words(
        " ".join(html.fragment_fromstring(source, create_parent="div").itertext())
    )
    rendered = _words(_document_text(_render(shape())))
    assert expected - rendered == Counter()
    # Dochodzą wyłącznie etykiety układu szablonu.
    assert set(rendered - expected) <= {
        "Daty",
        "Nazwa",
        "uczelni",
        "kierunek",
        "firmy",
        "Stanowisko",
    }


def test_edited_role_without_bold_position_keeps_its_text():
    # Rekruter zdjął pogrubienie stanowiska — układu nie da się ustalić,
    # więc akapit zostaje taki, jaki jest, zamiast zgubić stanowisko.
    content = (
        '<h2 data-cv-section="experience">DOŚWIADCZENIE</h2>'
        '<p data-cv-section="role">Tester 2020 - 2021</p>'
        '<p data-cv-section="employer">Acme</p>'
    )
    texts = [p.text for p in _render(content).paragraphs]
    assert "Tester 2020 - 2021" in texts
    assert "Nazwa firmy: Acme" in texts


def test_emphasis_set_in_the_editor_survives_the_layout():
    # Pogrubienia z QC i ręczne podkreślenia przechodzą 1:1 do punktorów.
    doc = _render(
        '<h2 data-cv-section="why_points">DLACZEGO NASZ KANDYDAT?</h2>'
        "<ul><li><p>Zna <strong>Kafka</strong> i <em>Spark</em></p></li></ul>"
    )
    bullet = next(p for p in doc.paragraphs if p.text == "Zna Kafka i Spark")
    assert [r.text for r in bullet.runs if r.bold] == ["Kafka"]
    assert [r.text for r in bullet.runs if r.italic] == ["Spark"]


def _boxes(doc: Document) -> list:
    return doc.element.body.xpath(".//*[local-name()='txbxContent']")


def test_text_typed_below_the_clause_stays_in_the_body_with_one_box():
    # Starszy edytor kopiował znacznik RODO do akapitu utworzonego Enterem;
    # dwie przypięte ramki leżałyby jedna na drugiej.
    doc = _render(
        f'<p>Treść</p><p data-cv-section="rodo">{RODO}</p>'
        '<p data-cv-section="rodo">Dopisek pod klauzulą</p>'
    )
    assert ["".join(_texts(box)) for box in _boxes(doc)] == [RODO]
    assert "Dopisek pod klauzulą" in [p.text for p in doc.paragraphs]


def test_clause_keeps_line_breaks_and_emphasis():
    doc = _render(
        '<p data-cv-section="rodo">Wyrażam zgodę<br>na przetwarzanie '
        "<strong>moich</strong> danych.</p>"
    )
    (box,) = _boxes(doc)
    assert "".join(_texts(box)) == "Wyrażam zgodę na przetwarzanie moich danych."
    bold = box.xpath(".//*[local-name()='r'][.//*[local-name()='b']]")
    assert ["".join(_texts(run)) for run in bold] == ["moich"]


def test_only_the_line_under_the_position_is_labelled_as_employer():
    texts = [
        p.text
        for p in _render(
            '<h2 data-cv-section="experience">DOŚWIADCZENIE</h2>'
            '<p data-cv-section="role"><strong>Dev</strong> 2020 - 2022</p>'
            '<p data-cv-section="employer">Firma X</p>'
            '<p data-cv-section="employer">Projekt dla banku, 5 osób</p>'
            '<p data-cv-section="duties_label">Dopisek pod etykietą</p>'
        ).paragraphs
    ]
    assert "Nazwa firmy: Firma X" in texts
    assert "Projekt dla banku, 5 osób" in texts
    assert "Nazwa firmy: Projekt dla banku, 5 osób" not in texts


def test_bold_prose_in_unmarked_experience_is_not_taken_for_a_role():
    doc = _render(
        "<h2>DOŚWIADCZENIE</h2>"
        "<p><strong>Projekt:</strong> nowy system płatności dla 2000 użytkowników</p>"
        "<p>Zespół 5 osób, Scrum.</p>"
    )
    texts = [p.text for p in doc.paragraphs]
    assert texts[-2:] == [
        "Projekt: nowy system płatności dla 2000 użytkowników",
        "Zespół 5 osób, Scrum.",
    ]
