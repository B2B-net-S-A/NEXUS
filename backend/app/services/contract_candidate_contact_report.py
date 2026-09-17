"""Arkusz braków kontaktu do konsultanta — po jednorazowej korekcie 0320.

Ticket: „Po zakończeniu backfillu system generuje plik Excel z listą
kontraktów, dla których e-mail i/lub telefon pozostały puste (brak danych
w obu źródłach) — do ręcznego uzupełnienia przez zespół."

„Brak" liczy się **tą samą regułą, co ekran** (``contract_candidate_contact``),
a nie zapytaniem o pustą kolumnę: kolumna na umowie jest nadpisaniem, więc
pusta kolumna przy wypełnionym profilu to NIE jest brak — wypisanie takich
wierszy wysyłałoby zespół do ręcznego przepisywania danych, które i tak widać.

Arkusz **nie zawiera wartości kontaktu** — z definicji ich nie ma. Niesie
nazwisko konsultanta i nazwę klienta, więc jest za bramką administratora,
tak jak ``order-sync-report``.
"""

from __future__ import annotations

import io
from datetime import date
from typing import Any, Iterable, Optional, Sequence

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.services.order_excel_export import _safe_text

GAP_HEADERS = (
    "ID kontraktu",
    "Konsultant",
    "Klient",
    "Rekrutacja",
    "Status",
    "Typ kontraktu",
    "Okres umowy",
    "Brak e-maila",
    "Brak telefonu",
)

_COLUMN_WIDTHS = (14, 30, 30, 30, 16, 16, 26, 14, 14)


def _period(start: Optional[date], end: Optional[date]) -> str:
    left = start.strftime("%d.%m.%Y") if start else "—"
    right = end.strftime("%d.%m.%Y") if end else "bezterminowo"
    return f"{left} – {right}"


def _cell(value: Any) -> Any:
    """Tekst przepuszczony przez ochronę przed wstrzyknięciem formuły."""

    if value is None:
        return "—"
    if isinstance(value, str):
        return _safe_text(value)
    return value


def gap_row(
    *,
    contract_id: int,
    candidate_name: Optional[str],
    client_name: Optional[str],
    job_title: Optional[str],
    status: str,
    contract_type: str,
    start_date: Optional[date],
    end_date: Optional[date],
    missing_email: bool,
    missing_phone: bool,
) -> list[Any]:
    return [
        contract_id,
        _cell(candidate_name),
        _cell(client_name),
        _cell(job_title),
        _cell(status),
        _cell(contract_type),
        _period(start_date, end_date),
        "tak" if missing_email else "nie",
        "tak" if missing_phone else "nie",
    ]


def _write_summary(workbook: Workbook, receipt: dict[str, Any], gaps: int) -> None:
    sheet = workbook.create_sheet(title="Podsumowanie")
    sheet.append(["Pozycja", "Wartość"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    skipped: dict[str, int] = {}
    for item in receipt.get("skipped") or []:
        reason = str(item.get("reason", "?"))
        skipped[reason] = skipped.get(reason, 0) + 1
    rows: Sequence[tuple[str, Any]] = (
        ("Korekta wykonana", receipt.get("executed_at", "—")),
        ("Umowy z dokumentem generatora", receipt.get("linked_documents_found", 0)),
        ("Uzupełniony e-mail", receipt.get("contracts_email_filled", 0)),
        ("Uzupełniony telefon", receipt.get("contracts_phone_filled", 0)),
        ("Kontrakty z brakami (arkusz „Braki”)", gaps),
        *(
            (f"Pominięte — {reason}", count)
            for reason, count in sorted(skipped.items())
        ),
    )
    for label, value in rows:
        sheet.append([_cell(label), value])
    sheet.column_dimensions["A"].width = 40
    sheet.column_dimensions["B"].width = 30
    sheet.sheet_view.showGridLines = False


def build_gap_workbook(*, receipt: dict[str, Any], rows: Iterable[list[Any]]) -> bytes:
    """Zbuduj skoroszyt (CPU-bound — wołaj poza pętlą zdarzeń)."""

    materialized = list(rows)
    workbook = Workbook()
    # Domyślny arkusz zostaje jako „Braki" — to on jest treścią raportu.
    gaps = workbook.active
    gaps.title = "Braki"
    gaps.append(list(GAP_HEADERS))
    for cell in gaps[1]:
        cell.font = Font(bold=True)
    for row in materialized:
        gaps.append(row)
    gaps.freeze_panes = "A2"
    if materialized:
        gaps.auto_filter.ref = gaps.dimensions
    for index, width in enumerate(_COLUMN_WIDTHS, start=1):
        gaps.column_dimensions[get_column_letter(index)].width = width
    gaps.sheet_view.showGridLines = False

    _write_summary(workbook, receipt, len(materialized))

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
