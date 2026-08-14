"""Parser miesięcznego arkusza wyników finansowych kontraktorów.

Pierwszy odczyt XLSX w tym backendzie — dotąd `openpyxl` służył wyłącznie do
EKSPORTU. Stąd komplet zabezpieczeń wejścia w jednym miejscu, zamiast rozsypania
ich po handlerze.

DWIE KLASY BŁĘDU, CELOWO ROZDZIELONE:

1. Nagłówki — arkusz bez wymaganych kolumn nie jest „arkuszem z brakami", tylko
   innym plikiem. Import odrzucany w całości z listą braków i nadmiarów.

2. Wiersz — brak lub niepoprawna wartość LICZBOWA **nie pomija wiersza**
   (wymóg pkt 4.1 ticketu). Wiersz wjeżdża z polem `None`, liczony do
   `needs_completion`, a operator uzupełnia go ręcznie w tabeli. Pomijanie
   takich wierszy było najgorszym z możliwych zachowań: suma w kaflu cicho
   przestawała obejmować kogoś, kogo w arkuszu widać.

   Odrzucany jest wyłącznie wiersz, którego nie da się w ogóle zinterpretować
   jako rekordu — bez nazwiska nie ma czego uzupełniać ani do czego wrócić.

WARTOŚCI Z ARKUSZA BYWAJĄ CZYM CHCĄ. Ta sama kolumna potrafi zwrócić `float`
(komórka liczbowa), `str` („1 234,50 zł", „12%", „b/d"), `datetime` (komórka
sformatowana jako data) albo `None`. Koercja jest więc jawna i wybaczająca —
z jednym wyjątkiem: nie zgadujemy. Czego nie da się odczytać jako liczby,
zostaje `None` i trafia do ręcznego uzupełnienia.
"""

from __future__ import annotations

import math
import re
from dataclasses import dataclass, field
from decimal import Decimal, DecimalException
from io import BytesIO
from typing import Any, Optional

# Nazwy DOKŁADNIE jak w arkuszu (pkt 4.1 ticketu). Sześć z nich nie ma
# odpowiednika w tabeli wynikowej i nie jest nigdzie zapisywanych — są tu,
# bo ich BRAK oznacza, że wgrano inny plik.
REQUIRED_HEADERS: tuple[str, ...] = (
    "Imię i nazwisko",
    "Średnia Stawka MD",
    "Ilość MD",
    "Wynagrodzenie",
    "Klient",
    "Uwagi",
    "Projekt",
    "Stawka z VD",
    "Stawka MD",
    "Faktura",
    "Marża PLN",
    "Marża %",
    "Czy wystawiono fakturę",
    "Data wysłania",
    "Płatny urlop",
)

# Nagłówek arkusza → kolumna modelu. Sześć pól spoza tego mapowania jest
# świadomie porzucanych po walidacji — żyją wyłącznie w oryginalnym pliku.
# Dwa nagłówki zmieniają nazwę w UI („Średnia Stawka MD" → „Stawka kosztowa
# MD", „Stawka MD" → „Stawka przychodowa MD"); nazwa w pliku zostaje.
HEADER_TO_FIELD: dict[str, str] = {
    "Imię i nazwisko": "consultant_name",
    "Klient": "client_name",
    "Średnia Stawka MD": "cost_rate_md",
    "Ilość MD": "md_count",
    "Wynagrodzenie": "compensation",
    "Stawka MD": "revenue_rate_md",
    "Faktura": "invoice_amount",
    "Marża PLN": "margin_pln",
    "Marża %": "margin_pct",
}

NUMERIC_FIELDS: tuple[str, ...] = (
    "cost_rate_md",
    "md_count",
    "compensation",
    "revenue_rate_md",
    "invoice_amount",
    "margin_pln",
    "margin_pct",
)

# Sufit wierszy. Arkusz miesięczny to dziesiątki pozycji; cokolwiek w tej skali
# to pomyłka albo próba wysycenia pamięci procesu (uvicorn ma jednego workera).
MAX_ROWS = 5000

# Znaki, które w kwotach z Excela pojawiają się regularnie i nie niosą wartości:
# spacja zwykła i niełamliwa (separator tysięcy), „zł", „%", apostrof.
_STRIP_RE = re.compile(r"[\s  ']|zł|PLN|%", re.IGNORECASE)
_NUMERIC_RE = re.compile(r"^-?\d+(?:[.,]\d+)?$")


class FinanceHeaderError(Exception):
    """Arkusz nie ma wymaganych kolumn — import odrzucany w całości."""

    def __init__(self, missing: list[str], unexpected: list[str]) -> None:
        self.missing = missing
        self.unexpected = unexpected
        super().__init__(f"missing={missing} unexpected={unexpected}")


class FinanceWorkbookError(Exception):
    """Plik nie jest czytelnym arkuszem XLSX (uszkodzony / zły format)."""


@dataclass(frozen=True)
class FinanceRowError:
    row_number: int
    reason: str


@dataclass(frozen=True)
class FinanceRowDraft:
    row_number: int
    consultant_name: str
    client_name: Optional[str]
    cost_rate_md: Optional[Decimal]
    md_count: Optional[Decimal]
    compensation: Optional[Decimal]
    revenue_rate_md: Optional[Decimal]
    invoice_amount: Optional[Decimal]
    margin_pln: Optional[Decimal]
    margin_pct: Optional[Decimal]
    """Pola, których import nie wypełnił — w UI dostają ramkę „Uzupełnij"."""
    missing_fields: tuple[str, ...] = ()


@dataclass(frozen=True)
class FinanceParseResult:
    rows: list[FinanceRowDraft] = field(default_factory=list)
    critical_errors: list[FinanceRowError] = field(default_factory=list)

    @property
    def needs_completion(self) -> int:
        """Ile wierszy ma choć jedno pole do ręcznego uzupełnienia."""
        return sum(1 for row in self.rows if row.missing_fields)


def _normalize_header(value: Any) -> str:
    """Nagłówki z Excela bywają z ogonem spacji albo twardą spacją w środku."""
    if value is None:
        return ""
    return re.sub(r"[\s ]+", " ", str(value)).strip()


def _coerce_decimal(value: Any) -> Optional[Decimal]:
    """Wartość z komórki → Decimal albo None. Nigdy nie rzuca.

    None zwracamy zarówno dla pustej komórki, jak i dla treści, której nie da
    się odczytać jako liczby („b/d", „-", „do ustalenia"). Rozróżnienie tych
    dwóch przypadków nic by nie dało: w obu wierszu brakuje liczby i w obu
    człowiek musi ją wpisać.
    """
    if value is None:
        return None
    if isinstance(value, bool):
        # bool jest podklasą int — bez tej gałęzi TRUE wjechałoby jako 1.
        return None
    if isinstance(value, Decimal):
        return value if value.is_finite() else None
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float):
        if math.isnan(value) or math.isinf(value):
            return None
        return Decimal(str(value))
    if not isinstance(value, str):
        # datetime i spółka — kolumna liczbowa sformatowana jako data to
        # pomyłka w arkuszu, nie liczba do odgadnięcia.
        return None

    cleaned = _STRIP_RE.sub("", value).strip()
    if not cleaned:
        return None
    # Polski zapis dziesiętny: „1234,50". Kropka jako separator tysięcy jest
    # niejednoznaczna („1.234" = tysiąc czy jeden i 234 tysięczne?), więc
    # akceptujemy wyłącznie zapis z jednym separatorem dziesiętnym.
    if not _NUMERIC_RE.match(cleaned):
        return None
    cleaned = cleaned.replace(",", ".")
    try:
        parsed = Decimal(cleaned)
    except (DecimalException, ValueError):
        return None
    return parsed if parsed.is_finite() else None


def _coerce_text(value: Any) -> Optional[str]:
    if value is None:
        return None
    text = re.sub(r"[\s ]+", " ", str(value)).strip()
    return text[:255] or None


def parse_finance_workbook(data: bytes) -> FinanceParseResult:
    """Odczytaj arkusz wyników. Wołać przez ``run_in_threadpool``.

    openpyxl jest synchroniczny i CPU-bound, a backend chodzi na jednym
    workerze uvicorna — wywołanie wprost z handlera zablokowałoby pętlę
    zdarzeń całej aplikacji na czas parsowania.
    """

    from openpyxl import load_workbook

    try:
        # read_only=True → strumieniowanie zamiast budowy pełnego drzewa
        # komórek. data_only=True → wartości, nie formuły (arkusz liczy marże
        # formułami; bez tego dostalibyśmy „=G2-D2").
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl rzuca kilkoma różnymi typami
        raise FinanceWorkbookError(str(exc)) from exc

    try:
        sheet = workbook.worksheets[0]
        rows_iter = sheet.iter_rows(values_only=True)

        try:
            header_row = next(rows_iter)
        except StopIteration:
            raise FinanceHeaderError(list(REQUIRED_HEADERS), []) from None

        headers = [_normalize_header(cell) for cell in header_row]
        present = {h for h in headers if h}
        missing = [h for h in REQUIRED_HEADERS if h not in present]
        if missing:
            unexpected = sorted(present - set(REQUIRED_HEADERS))
            raise FinanceHeaderError(missing, unexpected)

        # Pierwsze wystąpienie każdego nagłówka. Zduplikowana kolumna
        # w arkuszu nie jest błędem krytycznym — bierzemy tę po lewej.
        index_of: dict[str, int] = {}
        for position, name in enumerate(headers):
            if name and name not in index_of:
                index_of[name] = position

        rows: list[FinanceRowDraft] = []
        errors: list[FinanceRowError] = []

        for offset, raw in enumerate(rows_iter):
            row_number = offset + 2  # +1 za nagłówek, +1 bo Excel liczy od 1
            if len(rows) >= MAX_ROWS:
                errors.append(
                    FinanceRowError(
                        row_number=row_number,
                        reason=f"Przekroczono limit {MAX_ROWS} wierszy — reszta pominięta.",
                    )
                )
                break

            if raw is None or all(cell is None for cell in raw):
                continue  # pusty wiersz separujący — nie błąd

            def cell(header: str) -> Any:
                position = index_of[header]
                return raw[position] if position < len(raw) else None

            consultant = _coerce_text(cell("Imię i nazwisko"))
            if not consultant:
                # Jedyny warunek odrzucenia. Wiersz bez osoby nie ma tożsamości:
                # nie da się go ani uzupełnić, ani odnaleźć w tabeli.
                errors.append(
                    FinanceRowError(
                        row_number=row_number,
                        reason="Brak wartości w kolumnie „Imię i nazwisko”.",
                    )
                )
                continue

            values: dict[str, Any] = {
                "consultant_name": consultant,
                "client_name": _coerce_text(cell("Klient")),
            }
            missing_fields: list[str] = []
            for header, target in HEADER_TO_FIELD.items():
                if target not in NUMERIC_FIELDS:
                    continue
                parsed = _coerce_decimal(cell(header))
                values[target] = parsed
                if parsed is None:
                    missing_fields.append(target)

            rows.append(
                FinanceRowDraft(
                    row_number=row_number,
                    missing_fields=tuple(missing_fields),
                    **values,
                )
            )

        return FinanceParseResult(rows=rows, critical_errors=errors)
    finally:
        workbook.close()
