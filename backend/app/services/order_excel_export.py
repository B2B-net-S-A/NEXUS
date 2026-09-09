"""Excel export shared by legacy and multi-consultant client orders."""

from __future__ import annotations

import io
import re
import unicodedata
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from typing import TYPE_CHECKING, Callable, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.shared_md_orders import uses_shared_md_pool

if TYPE_CHECKING:
    from app.schemas.client_order_group import OrderGroupRead


COST_GROUP_TOTAL_LABEL = "Całe zamówienie (kwota łączna)"
MD_GROUP_TOTAL_LABEL = "Całe zamówienie (wspólna pula MD)"


@dataclass(frozen=True)
class OrderExportRow:
    consultant_name: str
    order_number: str
    cost_rate: Optional[Decimal]
    revenue_rate: Optional[Decimal]
    start_date: Optional[date]
    end_date: Optional[date]
    allocation: Optional[Decimal] = None
    consumption: Optional[Decimal] = None
    order_type: Optional[str] = None
    remaining_md: Optional[Decimal] = None


BASE_HEADERS = (
    "Imię i nazwisko",
    "Numer zamówienia",
    "Stawka kosztowa",
    "Stawka przychodowa",
    "Okres zamówienia",
)
MODEL_HEADERS = ("Liczba MD / Kwota zamówienia", "Zużycie zamówienia")
ORDER_TYPE_HEADER = "Typ zamówienia"

ORDER_TYPE_LABELS = {
    "md": "MD",
    "cost": "Kosztowe",
    "periodic": "Okresowe",
}


def is_current_order_period(
    start_date: Optional[date], end_date: Optional[date], today: date
) -> bool:
    """Czy okres zamówienia OBOWIĄZUJE dziś — jedna reguła dla całego eksportu.

    Eksport odpowiada na pytanie „kto i na jakich warunkach pracuje u tego
    klienta DZIŚ". Zamówienia zakończone i jeszcze nierozpoczęte wchodziły do
    arkusza razem z bieżącym, więc ten sam konsultant pojawiał się w pliku
    tyle razy, ile zamówień przewinęło się przez jego kontrakt.

    Brak daty końca = „bezterminowo" (sięga w prawo bez granicy); brak daty
    startu = zamówienie już obowiązuje — wiersz bez daty rozpoczęcia jest
    w tym module normą historyczną, a wykluczenie go zabierałoby z arkusza
    czyjąś realną, trwającą współpracę. Zamówienie „kończące się" mieści się
    w tej regule i ma się w eksporcie znaleźć.
    """

    if start_date is not None and start_date > today:
        return False
    return end_date is None or end_date >= today


def order_type_export_label(value: object) -> str:
    """Polish workbook label for an explicit/effective order type."""

    raw = getattr(value, "value", value)
    return ORDER_TYPE_LABELS.get(str(raw), str(raw))


def export_rows_for_group(
    group: OrderGroupRead,
    *,
    include_line: Optional[Callable[[object], bool]] = None,
) -> list[OrderExportRow]:
    """Build workbook rows at the correct group/consultant granularity.

    ``include_line`` zawęża obsadę do wierszy, które mają trafić do
    arkusza (eksport bierze wyłącznie aktualnych konsultantów). Wiersz
    zbiorczy grupy zostaje niezależnie od filtra: opisuje CAŁE zamówienie,
    a nie osobę.
    """

    rows: list[OrderExportRow] = []
    shared_md = uses_shared_md_pool(group)
    if group.is_cost_based and group.lines:
        # A cost group's amount is one shared pool.  Repeating it per person
        # makes spreadsheet sums grow with the number of consultants.
        rows.append(
            OrderExportRow(
                consultant_name=COST_GROUP_TOTAL_LABEL,
                order_number=group.order_number,
                cost_rate=None,
                revenue_rate=None,
                start_date=group.start_date,
                end_date=group.end_date,
                allocation=group.budget_amount,
                consumption=None,
            )
        )
    if shared_md and group.lines:
        rows.append(
            OrderExportRow(
                consultant_name=MD_GROUP_TOTAL_LABEL,
                order_number=group.order_number,
                cost_rate=None,
                revenue_rate=None,
                start_date=group.start_date,
                end_date=group.end_date,
                allocation=group.md_budget_total,
                consumption=group.md_budget_used,
                remaining_md=group.md_budget_remaining
                if group.md_budget_mode is not None
                else None,
            )
        )
    if not group.lines:
        rows.append(
            OrderExportRow(
                consultant_name="",
                order_number=group.order_number,
                cost_rate=None,
                revenue_rate=None,
                start_date=group.start_date,
                end_date=group.end_date,
                allocation=(
                    group.budget_amount
                    if group.is_cost_based
                    else group.md_budget_total
                    if shared_md
                    else None
                ),
                consumption=(group.md_budget_used if shared_md else None),
                remaining_md=group.md_budget_remaining
                if group.md_budget_mode is not None and shared_md
                else None,
            )
        )
        return rows
    for line in group.lines:
        if include_line is not None and not include_line(line):
            continue
        consumption: Optional[Decimal]
        if group.is_cost_based:
            consumption = line.invoiced_total
        elif shared_md:
            consumption = None
        elif line.md_total is None:
            consumption = None
        else:
            consumption = (
                line.md_total
                + (line.md_manual_adjustment or Decimal("0"))
                - (line.md_remaining or Decimal("0"))
            )
        rows.append(
            OrderExportRow(
                consultant_name=line.consultant_name,
                order_number=group.order_number,
                cost_rate=line.rate_cost,
                revenue_rate=line.rate_revenue,
                start_date=group.start_date,
                end_date=group.end_date,
                allocation=(
                    None if group.is_cost_based or shared_md else line.md_total
                ),
                consumption=consumption,
                remaining_md=line.md_remaining
                if group.md_budget_mode == "per_person"
                else None,
            )
        )
    return rows


def _safe_text(value: str) -> str:
    """Prevent user-controlled text from becoming an Excel formula."""

    return f"'{value}" if value.startswith(("=", "+", "-", "@")) else value


def _period(row: OrderExportRow) -> str:
    start = row.start_date.strftime("%d.%m.%Y") if row.start_date else "—"
    end = row.end_date.strftime("%d.%m.%Y") if row.end_date else "bezterminowo"
    return f"{start} – {end}"


def build_orders_workbook(
    rows: list[OrderExportRow],
    *,
    include_model_columns: bool,
    include_order_type: bool = False,
) -> bytes:
    """Build a readable, typed .xlsx workbook entirely in memory."""

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Zamówienia"
    include_remaining = include_model_columns and any(
        row.remaining_md is not None for row in rows
    )
    headers = [
        *BASE_HEADERS,
        *(MODEL_HEADERS if include_model_columns else ()),
        *((ORDER_TYPE_HEADER,) if include_order_type else ()),
        *(("Pozostały budżet MD",) if include_remaining else ()),
    ]
    sheet.append(headers)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    for item in rows:
        values: list[object] = [
            _safe_text(item.consultant_name),
            _safe_text(item.order_number),
            item.cost_rate,
            item.revenue_rate,
            _period(item),
        ]
        if include_model_columns:
            values.extend([item.allocation, item.consumption])
        if include_order_type:
            values.append(_safe_text(item.order_type or ""))
        if include_remaining:
            values.append(item.remaining_md)
        sheet.append(values)

    for row in sheet.iter_rows(min_row=2):
        for column in (3, 4):
            if row[column - 1].value is not None:
                row[column - 1].number_format = "#,##0.###;[Red]-#,##0.###"
        for column in (6, 7):
            if column <= len(row) and row[column - 1].value is not None:
                row[column - 1].number_format = "#,##0.######;[Red]-#,##0.######"

    if include_remaining:
        for cell in sheet.iter_rows(
            min_row=2, min_col=len(headers), max_col=len(headers)
        ):
            cell[0].number_format = "#,##0.######;[Red]-#,##0.######"
    widths = [
        30,
        20,
        19,
        21,
        27,
        *([31, 24] if include_model_columns else []),
        *([18] if include_order_type else []),
        *([24] if include_remaining else []),
    ]
    for index, width in enumerate(widths[: len(headers)], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def orders_export_filename(
    client_name: str, generated_on: Optional[date] = None
) -> str:
    """Return the stable ASCII filename required by the product ticket."""

    normalized = unicodedata.normalize("NFKD", client_name)
    ascii_name = normalized.encode("ascii", "ignore").decode("ascii")
    safe_name = re.sub(r"[^A-Za-z0-9]+", "_", ascii_name).strip("_") or "Klient"
    stamp = (generated_on or date.today()).strftime("%d.%m.%Y")
    return f"Zamowienia_{safe_name}_{stamp}.xlsx"
