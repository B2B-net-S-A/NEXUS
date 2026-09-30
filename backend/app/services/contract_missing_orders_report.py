"""Raport kontraktów bez podpiętego zamówienia (ticket 10, 30.09.2026).

„Podpięte zamówienie" liczy się tak, jak pokazuje je zakładka „Dokumenty"
kontraktu — i tylko tak. Zakładka składa dwie listy:

* dokumenty kontraktu typu „Zamówienie" (``contract_documents.doc_type =
  'order'``): ręczny upload albo automatyczna kopia PDF-u zamówienia
  MD/kosztowego (``source_order_group_id``);
* „Dokumenty zamówień": zamówienia kontraktu z wgranym plikiem PO
  (``client_orders.file_path``).

Kontrakt bez żadnego z nich trafia do raportu. Plik z SharePointa, który
klasyfikator zapisał jako „Inny", się NIE liczy — w zakładce też nie widać go
jako zamówienia, a raport ma odpowiadać temu, co widzi człowiek.

Poza raportem są wyłącznie umowy unieważnione (``void``): takiej umowy nikt nie
uzupełnia. Status jest kolumną, więc resztę zawęża filtr w Excelu.

Dwie kolumny pomocnicze odróżniają przyczyny braku, bo każda ma inną naprawę:
zamówienie jest w systemie, ale bez PDF-u (wgrać plik), albo PDF leży na
zamówieniu MD/kosztowym, a kopia nie trafiła do dokumentów kontraktu (wgrać
PDF zamówienia ponownie — kopia powstaje przy wgraniu).

Trzecia kolumna wskazuje dokumenty typu „Inny”, których nazwa wygląda na
zamówienie (ta sama reguła co klasyfikator plików z SharePointa,
``looks_like_order``). Typ zmienia człowiek — nazwa pliku to za słaby dowód,
żeby przestawiać go automatycznie.
"""

from __future__ import annotations

import io
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any, Iterable, Optional, Sequence

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import exists, func, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.core.export_safety import safe_cell
from app.models.client_order import ClientOrder, ClientOrderStatus
from app.models.client_order_group import ClientOrderGroup
from app.models.contract import Contract, ContractStatus
from app.models.contract_document import ContractDocument, ContractDocumentType
from app.services.client_identity import client_display_name
from app.services.contract_folder_docs.classify import looks_like_order

SHEET_MISSING = "Bez zamówienia"
SHEET_SUMMARY = "Podsumowanie"

HEADERS = (
    "ID kontraktu",
    "Kontrakt w NEXUSIE",
    "Konsultant",
    "Klient",
    "Rekrutacja",
    "Numer umowy B2B",
    "Typ kontraktu",
    "Status",
    "Start umowy",
    "Koniec umowy",
    "Zamówienia w systemie bez PDF",
    "Numery tych zamówień",
    "Plik typu „Inny” wyglądający na zamówienie",
    "Uwaga",
)

_COLUMN_WIDTHS = (13, 40, 28, 28, 30, 18, 14, 16, 13, 14, 16, 32, 36, 60)

RULE_DESCRIPTION = (
    "Kontrakt ma podpięte zamówienie, gdy w zakładce „Dokumenty” jest dokument "
    "typu „Zamówienie” albo zamówienie kontraktu ma wgrany plik PDF. "
    "Umowy anulowane są pominięte."
)

STATUS_LABELS = {
    "draft": "Szkic",
    "ready_for_signature": "Do podpisu",
    "active": "Aktywny",
    "ending": "Kończący się",
    "ended": "Zakończony",
    "void": "Anulowany",
}
_STATUS_ORDER = ("active", "ending", "ready_for_signature", "draft", "ended")

TYPE_LABELS = {"b2b": "B2B", "uop": "Umowa o pracę", "uzlecenie": "Umowa zlecenie"}

GROUP_PDF_NOTE = (
    "PDF zamówienia MD/kosztowego jest w systemie, ale nie ma kopii "
    "w Dokumentach — wgraj PDF zamówienia ponownie."
)
POSSIBLE_ORDER_NOTE = (
    "Plik wygląda na zamówienie — zmień jego typ na „Zamówienie” w zakładce Dokumenty."
)


@dataclass(frozen=True)
class MissingOrderRow:
    contract_id: int
    candidate_name: Optional[str]
    client_name: Optional[str]
    job_title: Optional[str]
    contract_number: Optional[str]
    contract_type: str
    status: str
    start_date: Optional[date]
    end_date: Optional[date]
    orders_without_pdf: int = 0
    order_numbers: tuple[str, ...] = ()
    group_pdf_without_copy: bool = False
    possible_order_files: tuple[str, ...] = ()


@dataclass(frozen=True)
class MissingOrdersReport:
    rows: list[MissingOrderRow]
    total_by_status: dict[str, int] = field(default_factory=dict)


def _enum_value(value: Any) -> str:
    return str(getattr(value, "value", value))


def has_order_document_clause():
    """SQL: kontrakt ma zamówienie widoczne w zakładce „Dokumenty"."""

    document = exists().where(
        ContractDocument.contract_id == Contract.id,
        ContractDocument.doc_type == ContractDocumentType.order,
    )
    order_file = exists().where(
        ClientOrder.contract_id == Contract.id,
        ClientOrder.file_path.is_not(None),
    )
    return document | order_file


_CHUNK = 5000


async def load_missing_orders_report(db: AsyncSession) -> MissingOrdersReport:
    """Kontrakty (poza anulowanymi) bez podpiętego zamówienia + liczby."""

    totals = {
        _enum_value(status_): count
        for status_, count in (
            await db.execute(
                select(Contract.status, func.count())
                .where(Contract.status != ContractStatus.void)
                .group_by(Contract.status)
            )
        ).all()
    }

    contracts = (
        (
            await db.execute(
                select(Contract)
                .where(
                    Contract.status != ContractStatus.void,
                    ~has_order_document_clause(),
                )
                .options(
                    selectinload(Contract.candidate),
                    selectinload(Contract.client),
                    selectinload(Contract.job),
                    selectinload(Contract.b2b_detail),
                )
                .order_by(Contract.id)
            )
        )
        .scalars()
        .all()
    )

    numbers: dict[int, list[str]] = defaultdict(list)
    counts: Counter[int] = Counter()
    group_pdf: set[int] = set()
    possible: dict[int, list[str]] = defaultdict(list)
    ids = [contract.id for contract in contracts]
    for start in range(0, len(ids), _CHUNK):
        chunk = ids[start : start + _CHUNK]
        order_rows = (
            await db.execute(
                select(
                    ClientOrder.contract_id,
                    ClientOrder.title,
                    ClientOrderGroup.order_number,
                    ClientOrderGroup.file_path,
                )
                .outerjoin(
                    ClientOrderGroup,
                    ClientOrderGroup.id == ClientOrder.order_group_id,
                )
                .where(
                    ClientOrder.contract_id.in_(chunk),
                    ClientOrder.status != ClientOrderStatus.cancelled,
                )
                .order_by(ClientOrder.contract_id, ClientOrder.id)
            )
        ).all()
        for contract_id, title, group_number, group_file in order_rows:
            counts[contract_id] += 1
            label = (group_number or title or "").strip()
            if label and label not in numbers[contract_id]:
                numbers[contract_id].append(label)
            if group_file:
                group_pdf.add(contract_id)
        other_documents = (
            await db.execute(
                select(ContractDocument.contract_id, ContractDocument.filename)
                .where(
                    ContractDocument.contract_id.in_(chunk),
                    ContractDocument.doc_type == ContractDocumentType.other,
                )
                .order_by(ContractDocument.contract_id, ContractDocument.id)
            )
        ).all()
        for contract_id, filename in other_documents:
            if looks_like_order(filename):
                possible[contract_id].append(filename)

    rows = []
    for contract in contracts:
        candidate = contract.candidate
        detail = contract.b2b_detail
        rows.append(
            MissingOrderRow(
                contract_id=contract.id,
                candidate_name=(
                    f"{candidate.name} {candidate.lastname}".strip()
                    if candidate
                    else None
                ),
                client_name=(
                    client_display_name(contract.client) if contract.client else None
                ),
                job_title=contract.job.title if contract.job else None,
                contract_number=detail.contract_number if detail else None,
                contract_type=_enum_value(contract.contract_type),
                status=_enum_value(contract.status),
                start_date=contract.start_date,
                end_date=contract.end_date,
                orders_without_pdf=counts.get(contract.id, 0),
                order_numbers=tuple(numbers.get(contract.id, ())),
                group_pdf_without_copy=contract.id in group_pdf,
                possible_order_files=tuple(possible.get(contract.id, ())),
            )
        )
    return MissingOrdersReport(rows=rows, total_by_status=totals)


def _cell(value: Any) -> Any:
    if value is None:
        return "—"
    if isinstance(value, str):
        return safe_cell(value) if value else "—"
    return value


def _date(value: Optional[date], empty: str = "—") -> str:
    return value.strftime("%d.%m.%Y") if value else empty


def _note(row: MissingOrderRow) -> str:
    notes = []
    if row.possible_order_files:
        notes.append(POSSIBLE_ORDER_NOTE)
    if row.group_pdf_without_copy:
        notes.append(GROUP_PDF_NOTE)
    return " ".join(notes)


def report_row(row: MissingOrderRow, base_url: str) -> list[Any]:
    return [
        row.contract_id,
        f"{base_url.rstrip('/')}/contracts/{row.contract_id}",
        _cell(row.candidate_name),
        _cell(row.client_name),
        _cell(row.job_title),
        _cell(row.contract_number),
        TYPE_LABELS.get(row.contract_type, row.contract_type),
        STATUS_LABELS.get(row.status, row.status),
        _date(row.start_date),
        _date(row.end_date, "bezterminowo"),
        row.orders_without_pdf,
        _cell(", ".join(row.order_numbers)) if row.order_numbers else "—",
        _cell(", ".join(row.possible_order_files)) if row.possible_order_files else "",
        _note(row),
    ]


def _bold_header(sheet) -> None:
    for cell in sheet[1]:
        cell.font = Font(bold=True)


def _write_summary(
    workbook: Workbook,
    report: MissingOrdersReport,
    generated_at: datetime,
) -> None:
    sheet = workbook.create_sheet(title=SHEET_SUMMARY)
    missing_by_status = Counter(row.status for row in report.rows)
    total = sum(report.total_by_status.values())
    missing = len(report.rows)
    sheet.append(["Pozycja", "Wartość"])
    _bold_header(sheet)
    lines: Sequence[tuple[str, Any]] = (
        ("Wygenerowano", generated_at.strftime("%d.%m.%Y %H:%M")),
        ("Reguła", RULE_DESCRIPTION),
        ("Kontrakty (bez anulowanych)", total),
        ("Z podpiętym zamówieniem", total - missing),
        (f"Bez zamówienia (arkusz „{SHEET_MISSING}”)", missing),
        (
            "— w tym z zamówieniem w systemie bez PDF",
            sum(1 for row in report.rows if row.orders_without_pdf),
        ),
        (
            "— w tym z PDF-em zamówienia MD/kosztowego bez kopii",
            sum(1 for row in report.rows if row.group_pdf_without_copy),
        ),
        (
            "— w tym z plikiem typu „Inny” wyglądającym na zamówienie",
            sum(1 for row in report.rows if row.possible_order_files),
        ),
    )
    for label, value in lines:
        sheet.append([_cell(label), value])
    sheet.append([])
    sheet.append(["Status", "Kontrakty", "Bez zamówienia"])
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
    statuses = [s for s in _STATUS_ORDER if s in report.total_by_status] + sorted(
        s for s in report.total_by_status if s not in _STATUS_ORDER
    )
    for status_ in statuses:
        sheet.append(
            [
                STATUS_LABELS.get(status_, status_),
                report.total_by_status[status_],
                missing_by_status.get(status_, 0),
            ]
        )
    sheet.column_dimensions["A"].width = 52
    sheet.column_dimensions["B"].width = 60
    sheet.column_dimensions["C"].width = 16
    sheet.sheet_view.showGridLines = False


def build_missing_orders_workbook(
    report: MissingOrdersReport,
    *,
    base_url: str,
    generated_at: datetime,
) -> bytes:
    """Zbuduj skoroszyt (CPU-bound — wołaj poza pętlą zdarzeń)."""

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_MISSING
    sheet.append(list(HEADERS))
    _bold_header(sheet)
    rows: Iterable[MissingOrderRow] = report.rows
    for row in rows:
        values = report_row(row, base_url)
        sheet.append(values)
        link_cell = sheet.cell(row=sheet.max_row, column=2)
        link_cell.hyperlink = values[1]
        link_cell.style = "Hyperlink"
    sheet.freeze_panes = "A2"
    if report.rows:
        sheet.auto_filter.ref = sheet.dimensions
    for index, width in enumerate(_COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.sheet_view.showGridLines = False

    _write_summary(workbook, report, generated_at)

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
