"""Render approved editor content without extracting, rewriting or inferring facts."""

from __future__ import annotations

import base64
from io import BytesIO
import re

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from lxml import html

from app.services.cv_generator_b2b.docx_renderer import (
    add_consent_screenshot,
    add_horizontal_line,
    normalize_letterhead_layout,
    TRANSLATIONS,
)

RENDERER_VERSION = "approved-html-1"
BLOCKS = {
    "p",
    "div",
    "section",
    "article",
    "main",
    "header",
    "footer",
    "blockquote",
    "h1",
    "h2",
    "h3",
    "h4",
    "h5",
    "h6",
    "ul",
    "ol",
    "li",
    "table",
    "hr",
    "pre",
}


class ApprovedDocxError(ValueError):
    pass


def _image(value: str) -> bytes:
    # Never fetch arbitrary URLs from editable HTML. Tiptap currently emits no
    # remote images; fail explicitly if an API client submits one.
    match = re.fullmatch(r"data:image/(?:png|jpeg);base64,([A-Za-z0-9+/=\s]+)", value)
    if not match or len(match[1]) > 8_000_000:
        raise ApprovedDocxError(
            "Nieobsługiwany obraz w CV. Użyj osadzonego PNG lub JPEG."
        )
    try:
        data = base64.b64decode(match[1], validate=True)
        from PIL import Image

        with Image.open(BytesIO(data)) as im:
            im.verify()
    except Exception as error:
        raise ApprovedDocxError("Nie można odczytać obrazu w CV.") from error
    return data


def render_approved_docx(
    content_html: str,
    template: bytes,
    *,
    consent: bytes | None = None,
    language: str = "pl",
) -> bytes:
    """The HTML is already sanitized. Preserve text, emphasis and order exactly.

    Uses the corporate letterhead but never the original generated narrative.
    Text formatting comes from the editor, not a fresh keyword/model pass.
    """
    doc = Document(BytesIO(template))
    normalize_letterhead_layout(doc)
    for child in list(doc.element.body):
        if not child.tag.endswith("}sectPr"):
            doc.element.body.remove(child)
    # Use private body styles: changing Normal also changes the letterhead,
    # while template Heading styles may contain all-caps and border residue.
    normal = doc.styles.add_style("Approved CV Body", WD_STYLE_TYPE.PARAGRAPH)
    normal.base_style = doc.styles["Normal"]
    normal.font.name = "Montserrat"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string("373535")
    normal.paragraph_format.space_after = Pt(5)
    for level, size in [(1, 22), (2, 14), (3, 11)]:
        heading = doc.styles.add_style(
            f"Approved CV Heading {level}", WD_STYLE_TYPE.PARAGRAPH
        )
        heading.base_style = normal
        heading.font.name = "Montserrat SemiBold"
        heading.font.size = Pt(size)
        heading.font.color.rgb = RGBColor.from_string("E14F4F")
        heading.font.all_caps = False
        heading.paragraph_format.keep_with_next = True
        heading.paragraph_format.space_before = Pt(10)
    rodo = doc.styles.add_style("Approved CV Consent Text", WD_STYLE_TYPE.PARAGRAPH)
    rodo.base_style = normal
    rodo.font.size = Pt(7)
    rodo.paragraph_format.keep_together = True
    rodo.paragraph_format.space_before = Pt(10)
    root = html.fragment_fromstring(content_html or "", create_parent="div")

    def text_run(paragraph, text, state):
        if not text:
            return
        run = paragraph.add_run(text)
        run.bold = state.get("bold", False)
        run.italic = state.get("italic", False)
        run.underline = state.get("underline", False)
        run.font.strike = state.get("strike", False)
        run.font.subscript = state.get("sub", False)
        run.font.superscript = state.get("sup", False)
        if state.get("code"):
            run.font.name = "Courier New"

    def inline(node, paragraph, state=None):
        state = dict(state or {})
        tag = str(node.tag).lower()
        for tags, key in [
            ({"b", "strong"}, "bold"),
            ({"em", "i"}, "italic"),
            ({"u"}, "underline"),
            ({"s", "del"}, "strike"),
            ({"sub"}, "sub"),
            ({"sup"}, "sup"),
            ({"code"}, "code"),
        ]:
            if tag in tags:
                state[key] = True
        style = node.get("style", "").lower()
        if re.search(r"font-weight\s*:\s*(bold|[7-9]00)", style):
            state["bold"] = True
        if "font-style: italic" in style or "font-style:italic" in style:
            state["italic"] = True
        if tag == "br":
            paragraph.add_run().add_break()
        elif tag == "img":
            try:
                paragraph.add_run().add_picture(
                    BytesIO(_image(node.get("src", ""))), width=Inches(5.5)
                )
            except ApprovedDocxError:
                raise
            except Exception as error:
                raise ApprovedDocxError("Nie można wstawić obrazu do DOCX.") from error
        else:
            text_run(paragraph, node.text, state)
            for child in node:
                inline(child, paragraph, state)
                text_run(paragraph, child.tail, state)

    def paragraph_for(container, tag="p", style="", prefix=""):
        paragraph = container.add_paragraph(style=normal)
        if prefix:
            paragraph.add_run(prefix)
        if tag.startswith("h") and len(tag) == 2 and tag[1].isdigit():
            paragraph.style = "Approved CV Heading " + str(min(int(tag[1]), 3))
        align = re.search(r"text-align\s*:\s*(left|right|center|justify)", style)
        if align:
            paragraph.alignment = {
                "left": WD_ALIGN_PARAGRAPH.LEFT,
                "right": WD_ALIGN_PARAGRAPH.RIGHT,
                "center": WD_ALIGN_PARAGRAPH.CENTER,
                "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
            }[align[1]]
        return paragraph

    def render_table(node, container):
        rows = node.xpath("./tr|./thead/tr|./tbody/tr|./tfoot/tr")
        if not rows:
            return
        # Respect merged cells; reject pathological dimensions before allocating.
        occupied = {}
        cells = []
        max_col = 0
        for row_index, row in enumerate(rows):
            col = 0
            for cell in row.xpath("./td|./th"):
                while (row_index, col) in occupied:
                    col += 1
                try:
                    width = int(cell.get("colspan", "1"))
                    height = int(cell.get("rowspan", "1"))
                except ValueError as error:
                    raise ApprovedDocxError(
                        "Nieprawidłowe scalanie komórek tabeli."
                    ) from error
                if (
                    not (1 <= width <= 20 and 1 <= height <= len(rows) - row_index)
                    or col + width > 20
                ):
                    raise ApprovedDocxError("Nieobsługiwany rozmiar tabeli w CV.")
                for ri in range(row_index, row_index + height):
                    for ci in range(col, col + width):
                        if (ri, ci) in occupied:
                            raise ApprovedDocxError(
                                "Nakładające się komórki tabeli w CV."
                            )
                        occupied[ri, ci] = True
                cells.append((row_index, col, height, width, cell))
                max_col = max(max_col, col + width)
                col += width
        if not max_col or len(rows) > 1000:
            raise ApprovedDocxError("Nieobsługiwany rozmiar tabeli w CV.")
        table = container.add_table(rows=len(rows), cols=max_col)
        for ri, ci, height, width, cell_node in cells:
            target = table.cell(ri, ci)
            if height > 1 or width > 1:
                target = target.merge(table.cell(ri + height - 1, ci + width - 1))
            target.text = ""
            render_container(cell_node, target)
            if (
                target.paragraphs
                and not target.paragraphs[0].text
                and len(target.paragraphs) > 1
            ):
                first = target.paragraphs[0]._element
                first.getparent().remove(first)

    def render_container(node, container, depth=0):
        # Mixed inline/block contents remain in document order. Unknown cosmetic
        # wrappers keep all their text; no source selection occurs in this layer.
        paragraph = None
        if node.text and node.text.strip():
            paragraph = paragraph_for(container)
            text_run(paragraph, node.text, {})
        for child in node:
            tag = str(child.tag).lower()
            if tag in {"ul", "ol"}:
                paragraph = None
                for number, li in enumerate(child.xpath("./li"), 1):
                    p = paragraph_for(
                        container, prefix=f"{number}. " if tag == "ol" else "• "
                    )
                    p.paragraph_format.left_indent = Inches(0.15 * (depth + 1))
                    text_run(p, li.text, {})
                    for part in list(li):
                        if str(part.tag).lower() in {"ul", "ol"}:
                            wrapper = html.Element("div")
                            wrapper.append(part)
                            render_container(wrapper, container, depth + 1)
                            p = None
                        elif str(part.tag).lower() == "p":
                            if p is None:
                                p = paragraph_for(container)
                            # Tiptap lists wrap item text in paragraphs.
                            if p.text and p.text not in {"• ", f"{number}. "}:
                                p.add_run().add_break()
                            inline(part, p)
                        else:
                            if p is None:
                                p = paragraph_for(container)
                            inline(part, p)
                        if part.tail and (part.tail.strip() or p is not None):
                            if p is None:
                                p = paragraph_for(container)
                            text_run(p, part.tail, {})
            elif tag == "table":
                paragraph = None
                render_table(child, container)
            elif tag == "hr":
                paragraph = None
                if container is doc:
                    add_horizontal_line(doc)
            elif tag in {"p", "pre", "blockquote", "h1", "h2", "h3", "h4", "h5", "h6"}:
                paragraph = None
                p = paragraph_for(container, tag, child.get("style", ""))
                if (
                    child.get("data-cv-section") == "rodo"
                    or "rodo" in child.get("class", "").split()
                ):
                    p.style = rodo
                inline(child, p)
            elif tag in BLOCKS:
                paragraph = None
                render_container(child, container, depth)
            else:
                if paragraph is None:
                    paragraph = paragraph_for(container)
                inline(child, paragraph)
            if child.tail and (child.tail.strip() or paragraph is not None):
                if paragraph is None:
                    paragraph = paragraph_for(container)
                text_run(paragraph, child.tail, {})

    render_container(root, doc)
    if consent is not None:
        if not add_consent_screenshot(
            doc,
            consent,
            TRANSLATIONS.get(language, TRANSLATIONS["pl"])["consent_heading"],
        ):
            raise ApprovedDocxError(
                "Nie można dołączyć zrzutu zgody. Zatwierdzenie zatrzymane."
            )
    output = BytesIO()
    doc.save(output)
    return output.getvalue()
