"""Parser rejestru umów z Excela działu — bajty pliku → wiersze, bez bazy.

Arkusz „Umowy B2B” (nagłówki szukane po synonimach, wzór
``md_import_parser.py``) i arkusz „Bez dzialalności ” (nazwa ze spacją na
końcu — porównujemy po znormalizowanym początku nazwy).

Kolory komórek są DANYMI: czerwone tło kolumny A to w praktyce „współpraca
zakończona”, pomarańczowy wiersz — „nie doszła do skutku / bez projektu”,
zielony wiersz w arkuszu „Bez działalności” — aneks zrobiony. Dlatego plik
czytamy tak, żeby widzieć style (``ReadOnlyCell.fill`` działa też w trybie
``read_only`` — pełne wczytanie arkusza, który deklaruje 1 048 576 wierszy
formatowania, trwałoby minuty).

Moduł nie dotyka bazy — testy parsera nie potrzebują Postgresa.
"""

from __future__ import annotations

import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any, Optional

from openpyxl import load_workbook

from app.services.candidate_identity_quarantine import normalize_person_name_part

_MAX_HEADER_SCAN_ROWS = 10
# Arkusz ma formatowanie do końca (1 048 576 wierszy), więc koniec danych to
# dopiero długa seria pustych wierszy.
_MAX_BLANK_RUN = 40

MAIN_SHEET = "umowy b2b"
NO_BUSINESS_SHEET_PREFIX = "bez dzialal"


class RegisterParseError(ValueError):
    """Plik nie wygląda na rejestr umów działu — komunikat po polsku."""


# ── Normalizacja ─────────────────────────────────────────────────────────────


def fold(value: Any) -> str:
    """Małe litery bez polskich znaków, DŁUGOŚĆ ZACHOWANA (pozycje zgodne
    z oryginałem — potrzebne do szukania daty obok słowa kluczowego)."""
    text = str(value or "").replace("Ł", "L").replace("ł", "l")
    decomposed = unicodedata.normalize("NFKD", text)
    return "".join(c for c in decomposed if not unicodedata.combining(c)).lower()


def norm(value: Any) -> str:
    """Klucz porównania nagłówków i słów: bez znaków, pojedyncze spacje."""
    return re.sub(r"[^a-z0-9]+", " ", fold(value)).strip()


def name_tokens(value: Any) -> frozenset[str]:
    return frozenset(
        token
        for part in str(value or "").split()
        if (token := normalize_person_name_part(part))
    )


# ── Nagłówki ─────────────────────────────────────────────────────────────────

_MAIN_HEADERS: dict[str, tuple[str, ...]] = {
    "name": ("nazwisko pozniej imie", "nazwisko imie", "nazwisko i imie", "nazwisko"),
    "number": ("numer umowy", "nr umowy", "numer"),
    "client": ("klient",),
    "position": ("stanowisko",),
    "kind": ("rodzaj umowy", "rodzaj"),
    "signing": ("data podpisania",),
    "start": ("data startu pracy", "data startu", "data rozpoczecia"),
    "end": ("data zakonczenia umowy", "data zakonczenia"),
    "loyalty": ("okres lojalnosci",),
    "recruiter": ("rekruter",),
    "settlement": ("czy wyslano informacje o rozliczeniach", "czy wyslano"),
    "welcome": ("mail powitalny",),
    "notes": ("uwagi",),
    "changes": ("zmiany w umowie",),
}
_MAIN_REQUIRED = ("name", "number")

_NO_BUSINESS_HEADERS: dict[str, tuple[str, ...]] = {
    "name": ("imie i nazwisko", "nazwisko i imie", "imie nazwisko"),
    "start": ("data startu pracy", "data startu"),
    "recruiter": ("rekruter",),
    "annex": ("aneksy do umow kiedy zrobione", "aneksy"),
    "client": ("klient",),
    "annex_notes": ("uwagi do wpisania w aneksie", "uwagi"),
}


def _map_headers(
    values: list[Any], synonyms: dict[str, tuple[str, ...]]
) -> dict[str, int]:
    normalized = [norm(v) for v in values]
    mapping: dict[str, int] = {}
    # Dokładne dopasowania przed prefiksami: „Uwagi” nie może zabrać kolumny
    # „UWAGI DO WPISANIA W ANEKSIE”, gdy obie istnieją.
    for exact in (True, False):
        for fieldname, options in synonyms.items():
            if fieldname in mapping:
                continue
            for index, header in enumerate(normalized):
                if not header or index in mapping.values():
                    continue
                if any(
                    header == option if exact else header.startswith(option)
                    for option in options
                ):
                    mapping[fieldname] = index
                    break
    return mapping


# ── Wartości komórek ─────────────────────────────────────────────────────────


def _cell_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str):
        stripped = value.strip()
        return stripped or None
    return value


def _raw(value: Any) -> Any:
    """Wartość do JSON-a (`raw` wiersza stagingu)."""
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, str)) or value is None:
        return value
    return str(value)


def _fill_descriptor(cell: Any) -> Optional[str]:
    fill = getattr(cell, "fill", None)
    if fill is None or not getattr(fill, "fill_type", None):
        return None
    color = getattr(fill, "fgColor", None)
    if color is None:
        return None
    kind = getattr(color, "type", None)
    if kind == "rgb" and isinstance(color.rgb, str):
        rgb = color.rgb.upper()[-6:]
        return None if rgb == "000000" and fill.fill_type != "solid" else rgb
    if kind == "theme":
        return f"theme:{color.theme}:{round(float(color.tint or 0), 2)}"
    if kind == "indexed":
        return f"indexed:{color.indexed}"
    return None


def _rgb(descriptor: Optional[str]) -> Optional[tuple[int, int, int]]:
    if not descriptor or not re.fullmatch(r"[0-9A-F]{6}", descriptor):
        return None
    return tuple(int(descriptor[i : i + 2], 16) for i in (0, 2, 4))  # type: ignore[return-value]


def is_red(descriptor: Optional[str]) -> bool:
    if descriptor in ("indexed:2", "indexed:10"):
        return True
    rgb = _rgb(descriptor)
    return bool(rgb and rgb[0] >= 0xCC and rgb[1] <= 0x40 and rgb[2] <= 0x40)


def is_orange(descriptor: Optional[str]) -> bool:
    if descriptor and descriptor.startswith("theme:5:"):
        # accent2 motywu Office z rozjaśnieniem (legenda: „nie doszła do
        # skutku” / „bez projektu”).
        return float(descriptor.split(":")[2]) > 0
    rgb = _rgb(descriptor)
    return bool(
        rgb
        and rgb[0] >= 0xF0
        and 0x80 <= rgb[1] <= 0xE8
        and rgb[2] <= 0xC0
        and rgb[0] - rgb[2] >= 0x30
    )


def is_green(descriptor: Optional[str]) -> bool:
    rgb = _rgb(descriptor)
    return bool(
        rgb and rgb[1] >= 0xC0 and rgb[1] - rgb[0] >= 0x20 and rgb[1] - rgb[2] >= 0x20
    )


# ── Daty ─────────────────────────────────────────────────────────────────────

_DMY_RE = re.compile(
    r"(?<!\d)(\d{1,2})\s*[.\-/]\s*(\d{1,2})\s*[.\-/]\s*(\d{2,4})(?!\d)"
)
_ISO_RE = re.compile(r"(?<!\d)(\d{4})-(\d{1,2})-(\d{1,2})(?!\d)")


def _year(raw: str) -> int:
    value = int(raw)
    if len(raw) <= 3 and value < 100:
        # „30.07.026”, „30.07.26” — literówki w roku.
        return 2000 + value
    return value


def dates_in(text: str) -> list[tuple[int, date]]:
    """Wszystkie daty w tekście z pozycją początku."""
    found: list[tuple[int, date]] = []
    for match in _DMY_RE.finditer(text):
        try:
            found.append(
                (
                    match.start(),
                    date(
                        _year(match.group(3)), int(match.group(2)), int(match.group(1))
                    ),
                )
            )
        except ValueError:
            continue
    for match in _ISO_RE.finditer(text):
        try:
            found.append(
                (
                    match.start(),
                    date(int(match.group(1)), int(match.group(2)), int(match.group(3))),
                )
            )
        except ValueError:
            continue
    return sorted(found)


def first_date(text: str) -> Optional[date]:
    found = dates_in(text)
    return found[0][1] if found else None


_NOT_CONCLUDED_RE = re.compile(r"nie\s+(doszl|dojdzie)\w*\s+do\s+skutku")
_CLOSURE_WORDS = ("porozumien", "wypowiedzen", "rozwiazan", "rozwiazal")
_CLOSURE_WINDOW = 60


def closure_date_from(*texts: Optional[str]) -> Optional[date]:
    """Najpóźniejsza data stojąca obok „porozumienie/wypowiedzenie/rozwiązanie”."""
    best: Optional[date] = None
    for text in texts:
        if not text:
            continue
        folded = fold(text)
        keyword_positions = [
            m.start() for word in _CLOSURE_WORDS for m in re.finditer(word, folded)
        ]
        if not keyword_positions:
            continue
        for position, found in dates_in(text):
            if any(abs(position - k) <= _CLOSURE_WINDOW for k in keyword_positions):
                if best is None or found > best:
                    best = found
    return best


# ── Wiersze ──────────────────────────────────────────────────────────────────

_LEGEND_MARKERS = (
    "umow",
    "numer",
    "legenda",
    "projektu",
    "zlecenie",
    "cos nie tak",
    "numeracja",
    "skutku",
)


def _looks_like_legend(text: Any) -> bool:
    if not isinstance(text, str):
        return False
    folded = norm(text)
    return (
        any(marker in folded for marker in _LEGEND_MARKERS) or len(folded.split()) > 5
    )


@dataclass
class ContractRow:
    row_number: int
    raw: dict[str, Any]
    name_raw: str
    partner_name: str
    tokens: frozenset[str]
    number_raw: Optional[str]
    number_int: Optional[int]
    client_raw: Optional[str]
    position: Optional[str]
    kind: Optional[str]
    kind_raw: Optional[str]
    signing_date: Optional[date]
    signing_raw: Optional[str]
    cancelled: bool
    start_date: Optional[date]
    start_date_mode: Optional[str]
    start_raw: Optional[str]
    end_date: Optional[date]
    end_raw: Optional[str]
    loyalty_raw: Optional[str]
    recruiter_raw: Optional[str]
    settlement_raw: Optional[str]
    welcome_raw: Optional[str]
    notes: Optional[str]
    changes: Optional[str]
    name_fill: Optional[str]
    row_fill: Optional[str]
    red: bool
    orange: bool
    closure_date: Optional[date]
    flags: list[str] = field(default_factory=list)


@dataclass
class NoBusinessRow:
    row_number: int
    raw: dict[str, Any]
    name_raw: str
    tokens: frozenset[str]
    start_date: Optional[date]
    client_raw: Optional[str]
    recruiter_raw: Optional[str]
    annex_raw: Optional[str]
    annex_notes: Optional[str]
    done: bool
    done_date: Optional[date]
    row_fill: Optional[str]


@dataclass
class ParsedRegister:
    contracts: list[ContractRow]
    no_business: list[NoBusinessRow]
    skipped: list[dict[str, Any]]
    no_business_sheet_found: bool


def _text(value: Any) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    return str(value).strip() or None


def _kind(raw: Optional[str]) -> Optional[str]:
    folded = norm(raw)
    if not folded:
        return None
    if "b2b" in folded:
        return "b2b"
    if "zlecen" in folded:
        return "mandate"
    if "dziel" in folded:
        return "work"
    if "uop" in folded or "o prace" in folded:
        return "employment"
    return None


def _partner_name(name_raw: str) -> str:
    """„Nazwisko Imię” (kolejność w Excelu) → „Imię Nazwisko”."""
    parts = name_raw.split()
    if len(parts) < 2:
        return name_raw
    return " ".join(parts[1:] + parts[:1])


def _parse_contract_row(
    row_number: int, values: dict[str, Any], fills: dict[str, Optional[str]]
) -> ContractRow:
    name_raw = _text(values.get("name")) or ""
    flags: list[str] = []

    number = values.get("number")
    number_int: Optional[int] = None
    number_raw: Optional[str] = None
    if isinstance(number, int) and not isinstance(number, bool):
        number_int, number_raw = number, str(number)
    elif number is not None:
        number_raw = _text(number)
        if number_raw and re.fullmatch(r"\d+", number_raw):
            number_int = int(number_raw)

    position = _text(values.get("position"))
    cancelled = False
    if position:
        folded_position = fold(position)
        if "bez projektu" in folded_position:
            flags.append("without_project")
        if _NOT_CONCLUDED_RE.search(folded_position):
            cancelled = True

    signing = values.get("signing")
    signing_date: Optional[date] = None
    signing_raw: Optional[str] = None
    if isinstance(signing, date):
        signing_date = signing
    elif signing is not None:
        signing_raw = _text(signing)
        if signing_raw and _NOT_CONCLUDED_RE.search(fold(signing_raw)):
            cancelled = True
        elif signing_raw:
            signing_date = first_date(signing_raw)
            if signing_date is None and norm(signing_raw):
                flags.append("signing_date_unparsed")

    start = values.get("start")
    start_date: Optional[date] = None
    start_mode: Optional[str] = None
    start_raw: Optional[str] = None
    if isinstance(start, date):
        start_date, start_mode = start, "exact"
    elif start is not None:
        start_raw = _text(start)
        folded_start = fold(start_raw)
        start_date = first_date(start_raw or "")
        if start_date is not None:
            if "nie pozniej" in folded_start:
                start_mode = "not_later"
            elif "nie wczesniej" in folded_start:
                start_mode = "not_earlier"
            else:
                start_mode = "exact"
            if "prawdopodob" in folded_start:
                flags.append("start_date_uncertain")
        elif norm(start_raw):
            flags.append("start_date_unknown")

    end = values.get("end")
    end_date: Optional[date] = None
    end_raw: Optional[str] = None
    if isinstance(end, date):
        end_date = end
    elif end is not None:
        end_raw = _text(end)
        folded_end = norm(end_raw)
        # „czas nieokreślony” i jego literówki („nieokeślony”, „nie określony”).
        if not (
            "nieok" in folded_end or "okres" in folded_end or "nieokr" in folded_end
        ):
            end_date = first_date(end_raw or "")

    kind_raw = _text(values.get("kind"))
    notes = _text(values.get("notes"))
    changes = _text(values.get("changes"))

    name_fill = fills.get("name")
    row_fill = next(
        (
            fills[key]
            for key in ("number", "client", "position", "kind", "signing", "start")
            if fills.get(key)
        ),
        None,
    )
    red = is_red(name_fill)
    orange = is_orange(row_fill) or is_orange(name_fill)
    if orange:
        flags.append("row_highlight_orange")
    closure = closure_date_from(notes, changes) if red else None
    if red and closure is None and not cancelled:
        flags.append("likely_ended")

    return ContractRow(
        row_number=row_number,
        raw={key: _raw(value) for key, value in values.items()},
        name_raw=name_raw,
        partner_name=_partner_name(name_raw),
        tokens=name_tokens(name_raw),
        number_raw=number_raw,
        number_int=number_int,
        client_raw=_text(values.get("client")),
        position=position,
        kind=_kind(kind_raw),
        kind_raw=kind_raw,
        signing_date=signing_date,
        signing_raw=signing_raw,
        cancelled=cancelled,
        start_date=start_date,
        start_date_mode=start_mode,
        start_raw=start_raw,
        end_date=end_date,
        end_raw=end_raw,
        loyalty_raw=_text(values.get("loyalty")),
        recruiter_raw=_text(values.get("recruiter")),
        settlement_raw=_text(values.get("settlement")),
        welcome_raw=_text(values.get("welcome")),
        notes=notes,
        changes=changes,
        name_fill=name_fill,
        row_fill=row_fill,
        red=red,
        orange=orange,
        closure_date=closure,
        flags=flags,
    )


_ANNEX_DONE_RE = re.compile(r"aneks\w*\s+(zrobion|podpisan|wyslan)")


def _parse_no_business_row(
    row_number: int, values: dict[str, Any], fills: dict[str, Optional[str]]
) -> NoBusinessRow:
    name_raw = _text(values.get("name")) or ""
    annex_raw = _text(values.get("annex"))
    annex_notes = _text(values.get("annex_notes"))
    done_date: Optional[date] = None
    done = False
    for text in (annex_raw, annex_notes):
        if text and _ANNEX_DONE_RE.search(fold(text)):
            done = True
            done_date = done_date or first_date(text)
    row_fill = next(
        (fills[k] for k in ("name", "start", "annex") if fills.get(k)), None
    )
    if is_green(row_fill):
        done = True
    start = values.get("start")
    start_date = start if isinstance(start, date) else first_date(_text(start) or "")
    return NoBusinessRow(
        row_number=row_number,
        raw={key: _raw(value) for key, value in values.items()},
        name_raw=name_raw,
        tokens=name_tokens(name_raw),
        start_date=start_date,
        client_raw=_text(values.get("client")),
        recruiter_raw=_text(values.get("recruiter")),
        annex_raw=annex_raw,
        annex_notes=annex_notes,
        done=done,
        done_date=done_date,
        row_fill=row_fill,
    )


def _iter_table(ws, synonyms: dict[str, tuple[str, ...]], required: tuple[str, ...]):
    """(nr wiersza, wartości, kolory) wierszy pod znalezionym nagłówkiem."""
    rows = ws.iter_rows()
    mapping: Optional[dict[str, int]] = None
    header_row = 0
    for index, cells in enumerate(rows, start=1):
        if index > _MAX_HEADER_SCAN_ROWS:
            break
        candidate = _map_headers([getattr(c, "value", None) for c in cells], synonyms)
        if all(key in candidate for key in required):
            mapping, header_row = candidate, index
            break
    if mapping is None:
        return None
    return header_row, mapping, rows


def _read_rows(header_row: int, mapping: dict[str, int], rows):
    blank_run = 0
    for index, cells in enumerate(rows, start=header_row + 1):
        values: dict[str, Any] = {}
        fills: dict[str, Optional[str]] = {}
        for fieldname, column in mapping.items():
            cell = cells[column] if column < len(cells) else None
            values[fieldname] = _cell_value(getattr(cell, "value", None))
            fills[fieldname] = _fill_descriptor(cell) if cell is not None else None
        if all(value is None for value in values.values()):
            blank_run += 1
            if blank_run >= _MAX_BLANK_RUN:
                return
            continue
        blank_run = 0
        yield index, values, fills


# Rejestr działu to ~1200 wierszy (kilkaset KB po rozpakowaniu). Sufit chroni
# proces web przed spreparowanym plikiem o wysokim współczynniku kompresji.
_MAX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024


def _assert_reasonable_archive(payload: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            total = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile as exc:
        raise RegisterParseError(
            "Nie udało się otworzyć pliku — wymagany jest skoroszyt Excel (.xlsx)."
        ) from exc
    if total > _MAX_UNCOMPRESSED_BYTES:
        raise RegisterParseError(
            "Plik po rozpakowaniu jest zbyt duży jak na rejestr umów."
        )


def parse_register(payload: bytes) -> ParsedRegister:
    _assert_reasonable_archive(payload)
    try:
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — zły plik to 422, nie 500
        raise RegisterParseError(
            "Nie udało się otworzyć pliku — wymagany jest skoroszyt Excel (.xlsx)."
        ) from exc
    try:
        main = next(
            (ws for ws in workbook.worksheets if norm(ws.title) == MAIN_SHEET), None
        )
        candidates = [main] if main is not None else list(workbook.worksheets)
        table = None
        for ws in candidates:
            table = _iter_table(ws, _MAIN_HEADERS, _MAIN_REQUIRED)
            if table is not None:
                break
        if table is None:
            raise RegisterParseError(
                "Nie znaleziono arkusza „Umowy B2B” z nagłówkami "
                "„NAZWISKO, PÓŹNIEJ IMIE” i „Numer umowy”."
            )
        header_row, mapping, rows = table
        contracts: list[ContractRow] = []
        skipped: list[dict[str, Any]] = []
        for row_number, values, fills in _read_rows(header_row, mapping, rows):
            name = values.get("name")
            number = values.get("number")
            if _looks_like_legend(name) or (
                name is None and _looks_like_legend(number)
            ):
                skipped.append({"row": row_number, "reason": "legend"})
                continue
            if name is None and number is None:
                skipped.append({"row": row_number, "reason": "no_name_no_number"})
                continue
            if name is None:
                skipped.append({"row": row_number, "reason": "no_name"})
                continue
            contracts.append(_parse_contract_row(row_number, values, fills))

        no_business: list[NoBusinessRow] = []
        sheet = next(
            (
                ws
                for ws in workbook.worksheets
                if norm(ws.title).startswith(NO_BUSINESS_SHEET_PREFIX)
            ),
            None,
        )
        found = False
        if sheet is not None:
            nb_table = _iter_table(sheet, _NO_BUSINESS_HEADERS, ("name",))
            if nb_table is not None:
                found = True
                nb_header, nb_mapping, nb_rows = nb_table
                for row_number, values, fills in _read_rows(
                    nb_header, nb_mapping, nb_rows
                ):
                    if values.get("name") is None or _looks_like_legend(
                        values.get("name")
                    ):
                        continue
                    no_business.append(
                        _parse_no_business_row(row_number, values, fills)
                    )
        return ParsedRegister(
            contracts=contracts,
            no_business=no_business,
            skipped=skipped,
            no_business_sheet_found=found,
        )
    finally:
        workbook.close()
