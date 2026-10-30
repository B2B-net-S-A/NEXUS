"""Rozliczenie premii miesiąca z arkusza Delivery: podgląd, potem przyznanie.

Arkusz ma kolumny: osoba (wypełniona tylko w pierwszym wierszu grupy), kandydat,
marża, kwota, uwagi. Plik nie jest zapisywany. Podgląd niczego nie zmienia,
a zapis biegnie z wierszy zatwierdzonych przez Płace na ekranie.

Premia przypisana do kontraktu dostaje klucz automatu
(``contract:{id}:{rola}``), więc automat nie zaproponuje jej drugi raz. Istniejąca
propozycja automatu jest przejmowana: osoba, kwota i miesiąc przychodzą z arkusza.
Wiersz z kwotą 0 odrzuca propozycję (powód poza ``AUTO_CANCEL_REASONS``, więc
automat jej nie przywraca). Uwagi z arkusza są widoczne tylko w podglądzie: do
opisu premii, który czyta osoba, trafia zdanie zbudowane z nazw.
"""

from __future__ import annotations

import io
import unicodedata
from dataclasses import dataclass
from datetime import date
from decimal import Decimal, InvalidOperation
from typing import Any, Optional
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.scheduling import business_today, local_now
from app.models.candidate import Candidate
from app.models.contract import Contract, ContractStatus
from app.models.hr_payroll import HrBonus
from app.models.hr_profile import EmployeeProfile
from app.models.user import User
from app.services.contract_rates import RATE_SCHEDULE_LOADS
from app.services.hr import placement_bonuses
from app.services.hr.bonus_rules import load_rules, money
from app.services.hr.payroll_access import PayrollViewer
from app.services.hr.payroll_bonuses import (
    _audit,
    _notify,
    clean_amount,
    clean_category,
)
from app.services.hr.payroll_errors import (
    PayrollError,
    PayrollForbidden,
    PayrollNotFound,
)
from app.services.hr.payroll_periods import (
    assert_open,
    lock_period_shared,
    validate_period,
)

#: Powód anulowania propozycji wierszem z kwotą 0 (automat jej nie przywraca).
CANCEL_REASON = "settlement_import"
MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 2000
NOT_A_PLACEMENT = "premia uznaniowa"
SKIPPED_MARKERS = ("gwarancje",)
ROLE_CATEGORIES = {
    "recruiter": "recruiter",
    "dl": "delivery_lead",
    "dl_extra": "delivery_lead",
}


class SettlementFileError(PayrollError):
    """Plik nie daje się odczytać (komunikat po polsku, dla użytkownika)."""


@dataclass(frozen=True)
class SheetRow:
    index: int
    person: str
    candidate: str
    amount: Decimal
    note: str


def name_key(value: Optional[str]) -> str:
    """Klucz nazwiska niezależny od kolejności słów, wielkości liter i ogonków."""

    text = (value or "").replace("ł", "l").replace("Ł", "L")
    plain = "".join(
        char
        for char in unicodedata.normalize("NFKD", text)
        if not unicodedata.combining(char)
    )
    return " ".join(sorted(plain.lower().split()))


def _amount(value: object) -> Optional[Decimal]:
    if value is None or isinstance(value, bool):
        return None
    text = str(value).replace("\xa0", "").replace(" ", "").replace(",", ".")
    try:
        number = Decimal(text)
    except InvalidOperation:
        return None
    return money(number) if number.is_finite() else None


def parse_workbook(content: bytes) -> list[SheetRow]:
    """Wiersze z kwotą z pierwszego arkusza; błąd pliku to komunikat dla Płac."""

    if len(content) > MAX_FILE_BYTES:
        raise SettlementFileError("Plik jest za duży (limit 5 MB).")
    try:
        book = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except (
        BadZipFile,
        InvalidFileException,
        KeyError,
        OSError,
        ParseError,
        ValueError,
    ) as exc:
        raise SettlementFileError("To nie jest plik Excela (.xlsx).") from exc
    sheet = book.worksheets[0]
    rows: list[SheetRow] = []
    person = ""
    for number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
        if number == 1:
            continue
        if number > MAX_ROWS:
            raise SettlementFileError(f"Plik ma więcej niż {MAX_ROWS} wierszy.")
        cells = (tuple(values) + (None,) * 5)[:5]
        if cells[0]:
            person = str(cells[0]).strip()
        candidate = str(cells[1] or "").strip()
        amount = _amount(cells[3])
        if (
            not person
            or not candidate
            or candidate == "-"
            or candidate.lower().startswith(SKIPPED_MARKERS)
            or amount is None
        ):
            continue
        rows.append(
            SheetRow(number, person, candidate, amount, str(cells[4] or "").strip())
        )
    book.close()
    if not rows:
        raise SettlementFileError(
            "Nie ma wierszy z kwotą. Kolumny: osoba, kandydat, marża, kwota, uwagi."
        )
    return rows


def resolve_role(
    note: str,
    *,
    recipient_id: Optional[int],
    automat_dl: Optional[int],
    automat_recruiter: Optional[int],
    is_dl: bool,
    is_recruiter: bool,
) -> str:
    """Rola wiersza: uwaga wygrywa, potem wskazanie automatu, potem rola konta."""

    lowered = note.lower()
    if "premia dl" in lowered:
        return "dl"
    if "premia rekrutersk" in lowered:
        return "recruiter"
    if recipient_id is not None and recipient_id == automat_recruiter:
        return "recruiter"
    if recipient_id is not None and recipient_id == automat_dl:
        return "dl"
    if is_dl and not is_recruiter:
        return "dl"
    return "recruiter"


async def _contracts_by_candidate(
    db: AsyncSession, names: set[str]
) -> dict[str, Contract]:
    """Najnowszy kontrakt każdego kandydata z arkusza, także szkic.

    Szkic liczy się, bo arkusz płaci już po podpisie: kontrakt bez klucza dostałby
    po aktywacji drugą premię od automatu.
    """

    rows = await db.execute(
        select(Contract, Candidate.name, Candidate.lastname)
        .join(Candidate, Candidate.id == Contract.candidate_id)
        .options(*RATE_SCHEDULE_LOADS)
        .where(Contract.status != ContractStatus.void)
        .order_by(func.coalesce(Contract.start_date, date.min), Contract.id)
    )
    found: dict[str, Contract] = {}
    for contract, first, last in rows.all():
        key = name_key(" ".join(part for part in (first, last) if part))
        if key in names:
            found[key] = contract
    return found


async def build_preview(
    db: AsyncSession, viewer: PayrollViewer, content: bytes, *, year: int, month: int
) -> dict[str, Any]:
    if not viewer.payroll:
        raise PayrollForbidden("Rozliczenie z arkusza wczytują Płace.")
    await assert_open(db, year, month)
    rows = parse_workbook(content)
    users = (
        await db.execute(
            select(User.id, User.name, User.role, User.roles, User.is_active).where(
                User.id.in_(select(EmployeeProfile.user_id))
            )
        )
    ).all()
    people = {name_key(row.name): row for row in users}
    contracts = await _contracts_by_candidate(db, {name_key(r.candidate) for r in rows})
    context = await placement_bonuses.load_context(db, list(contracts.values()))
    rules = await load_rules(db)
    existing = {
        bonus.source_key: bonus
        for bonus in await db.scalars(
            select(HrBonus).where(
                HrBonus.source == placement_bonuses.SOURCE,
                HrBonus.contract_id.in_([c.id for c in contracts.values()] or [0]),
            )
        )
    }
    names = dict((await db.execute(select(User.id, User.name))).all())
    taken: set[str] = set()
    result: list[dict[str, Any]] = []
    # Najpierw wiersze zgodne z automatem: to one dostają klucz roli, gdy kilka
    # osób ma pozycję DL za ten sam placement.
    ordered = sorted(
        rows,
        key=lambda row: 0 if _matches_automat(row, people, contracts, context) else 1,
    )
    for row in ordered:
        person = people.get(name_key(row.person))
        contract = contracts.get(name_key(row.candidate))
        result.append(
            await _preview_row(
                db,
                row,
                person=person,
                contract=contract,
                context=context,
                rules=rules,
                existing=existing,
                names=names,
                taken=taken,
                viewer=viewer,
            )
        )
    result.sort(key=lambda item: item["index"])
    counts: dict[str, int] = {}
    for item in result:
        counts[item["status"]] = counts.get(item["status"], 0) + 1
    return {
        "year": year,
        "month": month,
        "rows": result,
        "counts": counts,
        "total": str(sum((Decimal(item["amount"]) for item in result), Decimal(0))),
    }


def _matches_automat(row, people, contracts, context) -> bool:
    person = people.get(name_key(row.person))
    contract = contracts.get(name_key(row.candidate))
    if person is None or contract is None:
        return False
    dl_id, recruiter_id, _source = context.people(contract)
    return person.id in (dl_id, recruiter_id)


async def _preview_row(
    db: AsyncSession,
    row: SheetRow,
    *,
    person,
    contract: Optional[Contract],
    context,
    rules,
    existing: dict[str, HrBonus],
    names: dict[int, str],
    taken: set[str],
    viewer: PayrollViewer,
) -> dict[str, Any]:
    item: dict[str, Any] = {
        "index": row.index,
        "person": row.person,
        "candidate": row.candidate,
        "amount": str(row.amount),
        "note": row.note,
        "recipient_user_id": person.id if person else None,
        "recipient_name": person.name if person else None,
        "recipient_active": bool(person.is_active) if person else None,
        "contract_id": contract.id if contract else None,
        "contract_label": None,
        "start_date": None,
        "role": None,
        "keyed": False,
        "category": "custom",
        "automat_name": None,
        "automat_amount": None,
        "selected": True,
    }
    if person is None:
        item.update(status="no_person", selected=False)
        return item
    if person.id == viewer.user_id:
        item.update(status="self", selected=False)
        return item
    inactive = not person.is_active
    if row.candidate.lower().startswith(NOT_A_PLACEMENT) or contract is None:
        item["status"] = (
            "no_placement"
            if row.candidate.lower().startswith(NOT_A_PLACEMENT)
            else "no_contract"
        )
        return item
    candidate_name, client_name = context.names(contract)
    dl_id, recruiter_id, _source = context.people(contract)
    role = resolve_role(
        row.note,
        recipient_id=person.id,
        automat_dl=dl_id,
        automat_recruiter=recruiter_id,
        is_dl=placement_bonuses.holds_dl_role(person.role, person.roles),
        is_recruiter=placement_bonuses.holds_recruiter_role(person.role, person.roles),
    )
    keyed_role: Optional[str] = role
    if placement_bonuses.source_key(contract.id, role) in taken:
        keyed_role = (
            "dl_extra"
            if role == "dl"
            and placement_bonuses.source_key(contract.id, "dl_extra") not in taken
            else None
        )
    if keyed_role is not None:
        taken.add(placement_bonuses.source_key(contract.id, keyed_role))
    bonus = (
        existing.get(placement_bonuses.source_key(contract.id, keyed_role))
        if keyed_role
        else None
    )
    margin = await placement_bonuses.contract_margin(
        db, contract, contract.start_date or business_today()
    )
    target = placement_bonuses.placement_targets(
        margin.per_hour, dl_id, recruiter_id, rules
    ).get(role)
    automat_id = target.recipient_user_id if target else None
    item.update(
        contract_label=f"{candidate_name} ({client_name})",
        start_date=contract.start_date.isoformat() if contract.start_date else None,
        role=keyed_role or role,
        keyed=keyed_role is not None,
        category=ROLE_CATEGORIES[role],
        automat_name=names.get(automat_id) if automat_id else None,
        automat_amount=str(target.amount)
        if target and target.amount is not None
        else None,
    )
    if bonus is not None and bonus.status == "assigned":
        item.update(status="already_assigned", selected=False)
    elif inactive:
        # Decyzja Artura 30.10.2026: osobie po odejściu domyślnie nie płacimy.
        item.update(status="inactive", selected=False)
    elif row.amount == 0:
        item["status"] = "zero"
    elif automat_id == person.id:
        item["status"] = "matches"
    else:
        item["status"] = "differs"
    return item


async def _bonus_for_row(
    db: AsyncSession, row: dict[str, Any], contract: Optional[Contract]
) -> Optional[HrBonus]:
    if contract is None or not row.get("keyed"):
        return None
    return await db.scalar(
        select(HrBonus)
        .where(
            HrBonus.source_key == placement_bonuses.source_key(contract.id, row["role"])
        )
        .with_for_update()
        .execution_options(populate_existing=True)
    )


def _description(row: dict[str, Any], contract: Optional[Contract]) -> str:
    if contract is None or not row.get("contract_label"):
        return str(row["candidate"])[:1000]
    label = placement_bonuses.ROLE_LABELS.get(row.get("role"), "Premia")
    return f"{label} — placement: {row['contract_label']}"[:1000]


async def apply_rows(
    db: AsyncSession,
    viewer: PayrollViewer,
    rows: list[dict[str, Any]],
    *,
    year: int,
    month: int,
    as_proposal: bool = False,
) -> dict[str, int]:
    """Przyznaje zatwierdzone wiersze w jednej transakcji (bez commita).

    Powtórzenie jest bezpieczne: wiersz z kluczem przyznany wcześniej i wiersz bez
    klucza z tym samym numerem z arkusza są pomijane. ``as_proposal`` wpisuje je
    jako propozycje do potwierdzenia przez Płace: bez przyznania i bez dzwonka
    (decyzja Artura 30.10.2026). Automat nie nadpisuje ich osoby, kwoty ani miesiąca.
    """

    if not viewer.payroll:
        raise PayrollForbidden("Rozliczenie z arkusza przyznają Płace.")
    await lock_period_shared(db, year, month)
    await assert_open(db, year, month)
    validate_period(year, month)
    imported = set(
        await db.scalars(
            select(HrBonus.source_key).where(
                HrBonus.source_key.like(f"settlement:{year}-{month:02d}:%")
            )
        )
    )
    counts = {"assigned": 0, "proposed": 0, "rejected": 0, "skipped": 0}
    for row in rows:
        contract = (
            await db.get(Contract, row["contract_id"])
            if row.get("contract_id")
            else None
        )
        if row.get("contract_id") and contract is None:
            raise PayrollNotFound(f"Wiersz {row['index']}: nie ma takiego kontraktu.")
        if row.get("keyed") and row.get("role") not in placement_bonuses.ROLES:
            raise PayrollError(f"Wiersz {row['index']}: nieznana rola premii.")
        amount = money(Decimal(str(row["amount"])))
        bonus = await _bonus_for_row(db, row, contract)
        if amount == 0:
            if bonus is not None and bonus.status == "proposed":
                bonus.status = "cancelled"
                bonus.cancelled_by = viewer.user_id
                bonus.cancelled_at = local_now()
                bonus.cancel_reason = CANCEL_REASON
                _audit(db, bonus, actor_id=viewer.user_id, action="bonus_cancelled")
                counts["rejected"] += 1
            else:
                counts["skipped"] += 1
            continue
        amount = clean_amount(amount)
        category = clean_category(row.get("category") or "custom")
        if not row.get("recipient_user_id"):
            raise PayrollError(f"Wiersz {row['index']}: wybierz osobę.")
        recipient = int(row["recipient_user_id"])
        if recipient == viewer.user_id:
            raise PayrollForbidden("Premii nie przyznajesz sobie.")
        has_profile = await db.scalar(
            select(EmployeeProfile.id).where(EmployeeProfile.user_id == recipient)
        )
        if has_profile is None:
            raise PayrollNotFound(f"Wiersz {row['index']}: osoba nie ma kartoteki.")
        if bonus is not None and bonus.status == "assigned":
            counts["skipped"] += 1
            continue
        keyless = f"settlement:{year}-{month:02d}:{row['index']}"
        if bonus is None and not row.get("keyed") and keyless in imported:
            counts["skipped"] += 1
            continue
        if bonus is None:
            keyed = contract is not None and row.get("keyed")
            bonus = HrBonus(
                recipient_user_id=recipient,
                proposed_by=None if keyed else viewer.user_id,
                source=placement_bonuses.SOURCE if keyed else "manual",
                source_key=(
                    placement_bonuses.source_key(contract.id, row["role"])
                    if keyed
                    else keyless
                ),
                contract_id=contract.id if contract else None,
                currency="PLN",
                category=category,
                period_year=year,
                period_month=month,
                amount=amount,
                status="proposed" if as_proposal else "assigned",
            )
            db.add(bonus)
        bonus.recipient_user_id = recipient
        bonus.amount = amount
        bonus.category = category
        bonus.period_year, bonus.period_month = year, month
        bonus.status = "proposed" if as_proposal else "assigned"
        bonus.cancel_reason = None
        bonus.cancelled_at = None
        bonus.cancelled_by = None
        bonus.assigned_by = None if as_proposal else viewer.user_id
        bonus.assigned_at = None if as_proposal else local_now()
        bonus.description = _description(row, contract)
        bonus.rule_snapshot = {
            **(bonus.rule_snapshot or {}),
            "settlement_import": True,
            "edited_amount": True,
            "edited_recipient": True,
            "edited_period": True,
        }
        await db.flush()
        if as_proposal:
            _audit(db, bonus, actor_id=viewer.user_id, action="bonus_proposed")
            counts["proposed"] += 1
            continue
        _audit(db, bonus, actor_id=viewer.user_id, action="bonus_assigned")
        await _notify(db, bonus, event="new")
        counts["assigned"] += 1
    await db.flush()
    return counts


__all__ = [
    "SettlementFileError",
    "SheetRow",
    "apply_rows",
    "build_preview",
    "name_key",
    "parse_workbook",
    "resolve_role",
]
