"""Formularz CV klienta: wpisanie wartości do kopii wzoru DOCX (0522, PR 2).

Wzór zostaje taki, jaki dał klient — logo, stopka, etykiety, klauzula RODO.
Kod wpisuje tekst wyłącznie w miejsca odczytane przez
``client_form_structure`` (za etykietą, do pustej ramki, do pustej komórki)
i dokłada wiersze tylko w tabelach z „…”: kopia wiersza wzorcowego
z kolejnym numerem, wstawiona za nim; wiersz z „…” wtedy znika (jego numer
by się powtórzył), a bez dopisanych wierszy zostaje pusty. Puste pole zostaje puste: plik idzie do
klienta, więc żadnych „UZUPEŁNIJ” w treści.
"""

from __future__ import annotations

import copy
import io
import re
from typing import Any, Optional

from docx.oxml.ns import qn
from docx.table import _Cell

from app.core import docx_cells
from app.services import client_form_structure

_CELL_ID = re.compile(r"^t(\d+)\.r(\d+)\.c(\d+)$")
_ORDINAL = re.compile(r"^\s*(\d{1,3})(\s*\.?)\s*$")


def _clear(paragraph: Any) -> None:
    for run in list(paragraph._p.findall(qn("w:r"))):
        paragraph._p.remove(run)


def _write(cell: Any, value: str) -> None:
    """Tekst w komórce: pierwszy akapit z formatowaniem wzoru, linie jako złamania."""
    paragraphs = cell.paragraphs
    first = paragraphs[0]
    for extra in paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    _clear(first)
    lines = value.split("\n")
    run = first.add_run(lines[0])
    for line in lines[1:]:
        run.add_break()
        run = first.add_run(line)


def _tr_cells(tr: Any, table: Any) -> list[Any]:
    """Komórki jednego wiersza na każdą pozycję siatki (sufit na ``gridSpan``)."""
    cells: list[Any] = []
    for tc in tr.tc_lst:
        cells.extend([_Cell(tc, table)] * min(tc.grid_span, docx_cells.MAX_GRID_SPAN))
    return cells


def _is_ellipsis(cell: Any) -> bool:
    return bool(client_form_structure._ELLIPSIS.match(cell.text or ""))


def _drop_ellipsis(table: Any, *, remove_row: bool) -> None:
    """Wiersz z „…”: znika, gdy dochodzą wiersze (jego numer by się powtórzył);
    inaczej zostaje pusty jak inne niewypełnione wiersze wzoru."""
    rows = list(zip(list(table._tbl.tr_lst), docx_cells.grid_rows(table)))
    for tr, cells in rows:
        if not any(_is_ellipsis(cell) for cell in cells):
            continue
        if remove_row:
            tr.getparent().remove(tr)
            continue
        for cell in cells:
            if _is_ellipsis(cell):
                for paragraph in cell.paragraphs:
                    _clear(paragraph)


def _append_rows(table: Any, template_row: int, rows: list[dict[str, Any]]) -> None:
    """Kopie wiersza wzorcowego zaraz za nim, z kolejnym numerem w „Lp.”."""
    template = table._tbl.tr_lst[template_row]
    first = _tr_cells(template, table)
    ordinal = _ORDINAL.match(first[0].text or "") if first else None
    number = int(ordinal.group(1)) if ordinal else 0
    suffix = ordinal.group(2).strip() if ordinal else ""
    anchor = template
    for row in rows:
        new_tr = copy.deepcopy(template)
        anchor.addnext(new_tr)
        anchor = new_tr
        cells = _tr_cells(new_tr, table)
        for cell in cells:
            for paragraph in cell.paragraphs:
                _clear(paragraph)
        number += 1
        if ordinal and cells:
            _write(cells[0], f"{number}{suffix}")
        for column, cell in row["cells"].items():
            grid = int(column[1:])
            if grid < len(cells) and cell.get("value"):
                _write(cells[grid], cell["value"])


def render(
    template: bytes,
    structure: dict[str, Any],
    fields: dict[str, Any],
    extra_rows: Optional[list[dict[str, Any]]] = None,
) -> bytes:
    """Wzór z wpisanymi wartościami; odmowa wzoru = ``FormRejected``."""
    document = client_form_structure.open_document(template)
    paragraphs = document.paragraphs
    tables = document.tables
    grids: dict[int, list[list[Any]]] = {}
    for item in structure.get("fields") or []:
        value = (fields.get(item["id"]) or {}).get("value")
        if not value:
            continue
        if item["kind"] == "line":
            run = paragraphs[int(item["id"][1:])].add_run(f" {value}")
            run.bold = False
            continue
        match = _CELL_ID.match(item["id"])
        if match is None:
            continue
        t_index, r_index, grid = (int(part) for part in match.groups())
        if t_index not in grids:
            grids[t_index] = list(docx_cells.grid_rows(tables[t_index]))
        _write(grids[t_index][r_index][grid], value)

    by_table: dict[int, list[dict[str, Any]]] = {}
    for row in extra_rows or []:
        by_table.setdefault(int(row["table"]), []).append(row)
    for table_meta in structure.get("tables") or []:
        if not table_meta.get("extendable"):
            continue
        index = int(table_meta["index"])
        table = tables[index]
        template_row = table_meta.get("template_row")
        appending = template_row is not None and bool(by_table.get(index))
        _drop_ellipsis(table, remove_row=appending)
        if appending:
            _append_rows(table, int(template_row), by_table[index])

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ── 0525: tabela punktowa klienta (Excel) ───────────────────────────────────


def render_score(
    template: bytes, structure: dict[str, Any], fields: dict[str, Any]
) -> bytes:
    """Wzór Excela z wpisanymi wartościami; odmowa wzoru = ``FormRejected``.

    Wartość trafia tylko w komórkę odczytaną przez ``client_score_structure``.
    Punkty idą jako liczby, żeby formuły sum we wzorze („Razem”) policzyły się
    same po otwarciu; formuł, formatów i walidacji („Tak/Nie”) kod nie rusza.
    """
    from app.services import client_score_structure  # noqa: PLC0415

    workbook = client_score_structure.open_workbook(template)
    sheets = workbook.worksheets
    for item in structure.get("fields") or []:
        value = (fields.get(item["id"]) or {}).get("value")
        match = _SCORE_ID.match(item["id"])
        if not value or match is None:
            continue
        s_index, row, column = (int(part) for part in match.groups())
        if s_index >= len(sheets):
            continue
        cell = sheets[s_index].cell(row=row, column=column)
        if item.get("max_points") is not None and str(value).isdigit():
            cell.value = int(value)
        else:
            cell.value = value
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


_SCORE_ID = re.compile(r"^s(\d+)\.r(\d+)\.c(\d+)$")
