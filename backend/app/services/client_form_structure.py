"""Formularz CV klienta: odczyt pól do wypełnienia z wzoru DOCX. Bez AI, bez bazy.

Część klientów nie przyjmuje CV firmowego, tylko własny formularz. Centrum
e-Zdrowia ma go w każdym przetargu osobno („Załącznik nr 4a/4b do SWZ —
Formularz CV”): ten sam szkielet, inne wiersze kompetencji. Delivery Lead
wgrywa ten plik do Plików rekrutacji, a NEXUS wypełnia go danymi kandydata
(decyzje Artura 26.10.2026, ``docs/claude/generator-cv/formularz-klienta-zamiast-cv-firmowego.md``).

Ten moduł raz odczytuje strukturę wzoru — gdzie są puste miejsca na dane:

* ``line`` — akapit-etykieta zakończona dwukropkiem, bez treści po nim
  („Imię i nazwisko:”); wartość dopisuje się za etykietą,
* ``box`` — pusta tabela 1×1 pod etykietą („Cel zawodowy:” + ramka),
* ``cell`` — pusta komórka tabeli: nagłówek kolumny daje etykietę, a wypełniona
  komórka wiersza (np. nazwa kompetencji ze wzoru) jego nazwę.

Wiersz z samym „…” oznacza, że tabelę można wydłużać — liczy się jako
miejsce na kolejne wiersze, nie jako pole. Etykiet i tekstu wzoru kod nigdy
nie zmienia: SWZ zabrania modyfikowania pól formularza.

Identyfikator pola wskazuje miejsce w dokumencie (``p3`` — akapit, ``t2.r7.c3``
— tabela, wiersz, komórka po scaleniu), więc wypełnianie trafia w to samo
miejsce bez ponownego zgadywania. Kolejność elementów liczy się po treści
dokumentu; tabele zagnieżdżone i nagłówki stron są pomijane.
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import asdict, dataclass
from typing import Any, Optional

import docx
from docx.oxml.ns import qn

from app.core import docx_cells

# 2 (PR 2, 26.10.2026): tabela z „…” pamięta wiersz wzorcowy i jego kolumny
# — dopisywanie wierszy. Struktura starszej wersji czyta się ponownie z pliku.
STRUCTURE_VERSION = 2

MAX_ZIP_ENTRIES = 1000
MAX_UNPACKED_BYTES = 60 * 1024 * 1024
# Mniej pól to nie formularz, tylko dokument z tabelką (np. opis zamówienia).
MIN_FIELDS = 5
MAX_LABEL_CHARS = 120

MSG_NOT_DOCX = "Formularz musi być plikiem Word (.docx)."
MSG_BROKEN = "Nie da się odczytać pliku Word — zapisz go w Wordzie od nowa."
MSG_TOO_BIG = "Plik Word po rozpakowaniu jest za duży."
MSG_MACROS = "Plik z makrami nie może być wzorem — zapisz go jako zwykły .docx."
MSG_NO_FORM = (
    "W tym pliku nie ma pustych pól do wypełnienia — NEXUS nie rozpoznał w nim "
    "formularza CV."
)

_ELLIPSIS = re.compile(r"^\s*(…|\.{3,})\s*$")
_ORDINAL = re.compile(r"^\s*\d{1,3}\s*\.?\s*$")
_LP = re.compile(r"^\s*l\.?\s*p\.?\s*$", re.IGNORECASE)


class FormRejected(Exception):
    """Plik nie nadaje się na wzór formularza — zdanie dla użytkownika."""


@dataclass(frozen=True)
class FormField:
    id: str
    kind: str  # "line" | "box" | "cell"
    label: str
    section: Optional[str] = None
    row_label: Optional[str] = None
    table: Optional[str] = None


@dataclass(frozen=True)
class FormTable:
    index: int
    title: Optional[str]
    rows: int
    extendable: bool
    # Wiersz, którego kopie dokłada wypełnianie: ostatni wiersz danych przed
    # „…”, i wszystkie jego kolumny poza „Lp.” (``c<kolumna siatki>`` → etykieta).
    template_row: Optional[int] = None
    columns: Optional[dict[str, str]] = None


@dataclass(frozen=True)
class FormStructure:
    fields: tuple[FormField, ...]
    tables: tuple[FormTable, ...]
    version: int = STRUCTURE_VERSION

    def to_json(self) -> dict[str, Any]:
        return {
            "version": self.version,
            "fields": [asdict(item) for item in self.fields],
            "tables": [asdict(item) for item in self.tables],
        }


def _clean(text: str) -> str:
    return " ".join((text or "").split())


def _label(text: str) -> str:
    value = _clean(text).rstrip(":").rstrip("*").strip()
    return value[:MAX_LABEL_CHARS]


def _check_package(content: bytes) -> None:
    try:
        package = zipfile.ZipFile(io.BytesIO(content))
    except zipfile.BadZipFile as exc:
        raise FormRejected(MSG_NOT_DOCX) from exc
    with package:
        entries = package.infolist()
        if len(entries) > MAX_ZIP_ENTRIES:
            raise FormRejected(MSG_BROKEN)
        if sum(entry.file_size for entry in entries) > MAX_UNPACKED_BYTES:
            raise FormRejected(MSG_TOO_BIG)
        names = {entry.filename.lower() for entry in entries}
        if "word/document.xml" not in names:
            raise FormRejected(MSG_NOT_DOCX)
        if any(name.endswith("vbaproject.bin") for name in names):
            raise FormRejected(MSG_MACROS)


def open_document(content: bytes) -> Any:
    """Dokument python-docx po kontroli paczki; odmowa = ``FormRejected``."""
    _check_package(content)
    try:
        return docx.Document(io.BytesIO(content))
    except Exception as exc:  # uszkodzony DOCX (python-docx, lxml)
        raise FormRejected(MSG_BROKEN) from exc


def grid_cells(row: list[Any]) -> list[tuple[int, int, Any]]:
    """(pierwsza kolumna siatki, szerokość, komórka) po scaleniu.

    ``row`` to wiersz z ``docx_cells.grid_rows`` — komórka na każdą pozycję
    siatki, scalona powtórzona (z sufitem na ``gridSpan``, który podaje plik:
    ``row.cells`` python-docx rozwijałby go bez granicy). Liczy się kolumna
    siatki, nie kolejność komórek: nagłówek i wiersze danych bywają scalone
    inaczej (w formularzu „Analityk” CEZ „Lp.” zajmuje dwie kolumny nagłówka,
    a pod nim pusta komórka wiersza obejmuje drugą z nich i „Nazwę uczelni”).
    """
    cells: list[tuple[int, int, Any]] = []
    previous = None
    for grid, cell in enumerate(row):
        if previous is not None and cell._tc is previous:
            start, span, kept = cells[-1]
            cells[-1] = (start, span + 1, kept)
            continue
        cells.append((grid, 1, cell))
        previous = cell._tc
    return cells


def _named_headers(raw_rows: list[Any], headers_row: Optional[int]) -> list[str]:
    """Nazwane nagłówki wiersza nagłówkowego po kolei, bez „Lp.” i bez powtórzeń scalenia."""
    if headers_row is None:
        return []
    names: list[str] = []
    for _, _, cell in grid_cells(raw_rows[headers_row]):
        label = _label(cell.text)
        if label and not _LP.match(label) and (not names or names[-1] != label):
            names.append(label)
    return names


def _header_for(headers: dict[int, str], start: int, span: int) -> Optional[str]:
    """Nagłówek nad komórką: pierwszy spoza „Lp.” w jej kolumnach.

    ``""`` — komórka leży wyłącznie pod „Lp.” (numer wiersza, nie pole);
    ``None`` — nad tymi kolumnami nie ma nagłówka.
    """
    labels = [headers[grid] for grid in range(start, start + span) if grid in headers]
    if not labels:
        return None
    named = next((label for label in labels if not _LP.match(label)), None)
    return named if named is not None else ""


def body_items(document: Any) -> list[tuple[str, int]]:
    """(``"p"`` | ``"t"``, indeks) w kolejności treści dokumentu."""
    items: list[tuple[str, int]] = []
    p_index = t_index = 0
    for element in document.element.body.iterchildren():
        if element.tag == qn("w:p"):
            items.append(("p", p_index))
            p_index += 1
        elif element.tag == qn("w:tbl"):
            items.append(("t", t_index))
            t_index += 1
    return items


def _is_label(text: str) -> bool:
    value = _clean(text)
    return value.endswith(":") and 1 < len(value) <= MAX_LABEL_CHARS


def _table_fields(
    table: Any, t_index: int, title: Optional[str]
) -> tuple[list[FormField], FormTable]:
    fields: list[FormField] = []
    raw_rows = list(docx_cells.grid_rows(table))
    rows = [grid_cells(row) for row in raw_rows]

    # Ramka na tekst: tabela z jedną komórką pod etykietą („Cel zawodowy:”).
    if len(rows) == 1 and len(rows[0]) == 1:
        if not _clean(rows[0][0][2].text):
            fields.append(
                FormField(id=f"t{t_index}.r0.c0", kind="box", label=title or "Opis")
            )
        return fields, FormTable(t_index, title, 0, False)

    numbered = any(cells and _ORDINAL.match(_clean(cells[0][2].text)) for cells in rows)
    headers: dict[int, str] = {}
    headers_row: Optional[int] = None
    section: Optional[str] = None
    extendable = False
    data_rows = 0
    last_row: Optional[tuple[int, dict[str, str]]] = None
    template: Optional[tuple[int, dict[str, str]]] = None
    for r_index, cells in enumerate(rows):
        texts = [(start, span, _clean(cell.text)) for start, span, cell in cells]
        filled = [text for _, _, text in texts if text]
        if filled and (
            len(cells) == 1 or len(set(filled)) == 1 and len(filled) == len(texts)
        ):
            # Wiersz-sekcja przez całą szerokość („Kompetencje obligatoryjne”).
            section = _label(filled[0])
            continue
        if not filled and (numbered or len(cells) == 1):
            continue  # pusty wiersz bez numeru w tabeli numerowanej
        if any(_ELLIPSIS.match(text) for _, _, text in texts):
            extendable = True
            if template is None and last_row is not None:
                template = last_row
            continue
        if len(filled) == len(texts):
            if not headers or any(_LP.match(text) for text in filled):
                headers = {
                    grid: _label(cell.text)
                    for grid, cell in enumerate(raw_rows[r_index])
                }
                headers_row = r_index
            continue
        # Komórki wiersza poza kolumną „Lp.”. Gdy jest ich tyle, ile nazwanych
        # nagłówków, etykiety idą po kolei — siatka nagłówka i wierszy bywa
        # przesunięta; dopasowanie po kolumnach siatki zostaje na resztę.
        own = [item for item in texts if _header_for(headers, item[0], item[1]) != ""]
        names = _named_headers(raw_rows, headers_row)
        by_order = len(own) == len(names)
        row_label = next(
            (text for _, _, text in own if text and not _ORDINAL.match(text)), None
        )
        ordinal = next((text for _, _, text in texts[:1] if _ORDINAL.match(text)), None)
        added = 0
        for position, (start, span, text) in enumerate(own):
            if text:
                continue
            header = names[position] if by_order else _header_for(headers, start, span)
            fields.append(
                FormField(
                    id=f"t{t_index}.r{r_index}.c{start}",
                    kind="cell",
                    label=header or f"Kolumna {start + 1}",
                    section=section,
                    row_label=_label(row_label)
                    if row_label
                    else (f"Wiersz {ordinal.rstrip('.')}" if ordinal else None),
                    table=title,
                )
            )
            added += 1
        if added:
            data_rows += 1
            last_row = (
                r_index,
                {
                    f"c{start}": (names[position] if by_order else None)
                    or _header_for(headers, start, span)
                    or f"Kolumna {start + 1}"
                    for position, (start, span, _) in enumerate(own)
                },
            )
    if extendable and template is None:
        template = last_row
    return fields, FormTable(
        t_index,
        title,
        data_rows,
        extendable,
        template_row=template[0] if template else None,
        columns=template[1] if template else None,
    )


def read_structure(content: bytes) -> FormStructure:
    """Pola formularza albo ``FormRejected`` ze zdaniem dla użytkownika."""
    document = open_document(content)
    paragraphs = document.paragraphs
    tables = document.tables
    items = body_items(document)

    fields: list[FormField] = []
    table_meta: list[FormTable] = []
    for position, (kind, index) in enumerate(items):
        if kind == "p":
            text = paragraphs[index].text
            if not _is_label(text):
                continue
            following = items[position + 1] if position + 1 < len(items) else None
            if following is not None and following[0] == "t":
                continue  # etykieta tabeli — tytuł, nie pole
            fields.append(FormField(id=f"p{index}", kind="line", label=_label(text)))
            continue
        title = None
        if position > 0 and items[position - 1][0] == "p":
            previous = _clean(paragraphs[items[position - 1][1]].text)
            if 1 < len(previous) <= MAX_LABEL_CHARS:
                title = _label(previous)
        table_fields, meta = _table_fields(tables[index], index, title)
        fields.extend(table_fields)
        if meta.rows:
            table_meta.append(meta)

    if len(fields) < MIN_FIELDS or not any(item.kind == "cell" for item in fields):
        raise FormRejected(MSG_NO_FORM)
    return FormStructure(fields=tuple(fields), tables=tuple(table_meta))


def summary(structure: dict[str, Any]) -> str:
    """Jedno zdanie do Plików rekrutacji: co NEXUS rozpoznał we wzorze."""
    fields = structure.get("fields") or []
    lines = sum(1 for item in fields if item.get("kind") == "line")
    boxes = sum(1 for item in fields if item.get("kind") == "box")
    parts: list[str] = []
    if lines:
        parts.append(f"{lines} {_plural(lines, 'pole', 'pola', 'pól')} nagłówka")
    if boxes:
        parts.append(f"{boxes} {_plural(boxes, 'ramka', 'ramki', 'ramek')} na opis")
    for table in structure.get("tables") or []:
        rows = int(table.get("rows") or 0)
        name = table.get("title") or f"tabela {int(table.get('index', 0)) + 1}"
        text = f"„{name}”: {rows} {_plural(rows, 'wiersz', 'wiersze', 'wierszy')}"
        if table.get("extendable"):
            text += " (można dokładać)"
        parts.append(text)
    total = len(fields)
    head = f"{total} {_plural(total, 'pole', 'pola', 'pól')} do wypełnienia"
    return head + (" · " + " · ".join(parts) if parts else "")


def _plural(count: int, one: str, few: str, many: str) -> str:
    if count == 1:
        return one
    if count % 10 in (2, 3, 4) and count % 100 not in (12, 13, 14):
        return few
    return many
