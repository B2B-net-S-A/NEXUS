"""Szablony protokołów sprzętu (docxtpl) — budowane kodem ze wzoru recepcji.

Wzór wydania to protokół zdawczo-odbiorczy recepcji (16.10.2026) słowo
w słowo: ten sam układ akapitów, ta sama tabela (Lp., Nazwa sprzętu,
Nr seryjny, Ilość, Cena, Wartość), te same klauzule i podpisy. Dane osoby,
numery i kwoty są polami Jinja; jedyny dopisek to numer protokołu pod
tytułem (makieta „Sprzęt → Wydaj sprzęt”), bez którego nie da się go
odszukać ani podpisać linkiem.

Wzoru zwrotu recepcja nie miała — protokół zwrotu ma ten sam nagłówek
i strony zamienione miejscami, a w tabeli stan przy zwrocie zamiast ceny.

Szablony w ``app/templates/equipment`` są WYNIKIEM tego modułu (jak
``scripts/build_academy_templates.py`` dla Akademii)::

    python -m app.services.equipment.protocol_template

Zmiana treści = zmiana tutaj, ponowne uruchomienie i nowa
``TEMPLATE_VERSION`` (protokół zapisuje wersję wzoru, z której powstał).

28.10.2026 (zgłoszenie recepcji): siatka tabeli niesie szerokości kolumn
(LibreOffice robi z niej PDF), podpisy stoją w tabelce bez ramek z nazwiskami
stron, a pod linią podpisu pola ``holder_signature`` / ``issuer_signature``
— dopiski z NEXUSA po potwierdzeniu linkiem i zatwierdzeniu recepcji.

Wzory wgrane przez recepcję (etap 3, O3)
----------------------------------------

Recepcja i admin wgrywają własne wzory (np. „z kaucją”, „zwrot u klienta”)
jako DOCX z polami ``{{ … }}`` — te same pola, które wypełnia standard
(``FIELDS``, wiersze tabeli ``LINE_FIELDS``). Plik przychodzi od człowieka,
więc:

* przed zapisem przechodzi ``check_upload``: tylko ``.docx`` (zip bez makr,
  limit rozmiaru spakowanego i rozpakowanego), wymagane pola rodzaju, żadnych
  pól spoza listy i próbne wypełnienie fikcyjnym protokołem,
* wypełnia się WYŁĄCZNIE w ``SandboxedEnvironment`` jinja2 (docxtpl
  ``render(…, jinja_env=…)``): bez atrybutów Pythona (``__class__``,
  ``__globals__``…), bez makr, dołączania plików, funkcji globalnych
  (``range``, ``lipsum``…), potęgowania, mnożenia i formatowania napisów
  (pamięć), z krótką listą filtrów i ``StrictUndefined`` — literówka w polu
  pada przy wgraniu, a nie w protokole do podpisu.
"""

from __future__ import annotations

import hashlib
import io
import re
import zipfile
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable, Mapping, Optional

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt, Twips
from docx.table import Table
from docxtpl import DocxTemplate
from jinja2 import StrictUndefined, TemplateError, TemplateSyntaxError, nodes
from jinja2.exceptions import SecurityError
from jinja2.sandbox import SandboxedEnvironment
from sqlalchemy import func, or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.docx_cells import row_cells
from app.models.equipment import (
    ClientOnboardingPack,
    EquipmentProtocol,
    EquipmentProtocolTemplate,
)
from app.models.user import User
from app.services.equipment.errors import (
    EquipmentConflict,
    EquipmentError,
    EquipmentNotFound,
)

TEMPLATE_DIR = Path(__file__).resolve().parents[2] / "templates" / "equipment"
ISSUE_TEMPLATE = "protokol_wydania.docx"
RETURN_TEMPLATE = "protokol_zwrotu.docx"
TEMPLATE_VERSION = "2026-10-28"

COMPANY_LINE = (
    "Spółka B2B.net SA, z siedzibą w Warszawie 02-486, Al. Jerozolimskie 180, "
    "z upoważnienia "
)
LIABILITY_TAIL = (
    ", iż przyjmuje na siebie odpowiedzialność materialną za mienie Spółki "
    "B2B.net SA powierzone na podstawie niniejszego protokołu z obowiązkiem "
    "jego zwrotu w stanie niepogorszonym. "
)
CONDITION_TAIL = (
    "się ze stanem technicznym powierzonego mienia i nie wnosi zastrzeżeń. "
)
CLAUSE_1 = (
    "Sprzęt będzie oddany w dokładnie takim samym stanie, w jakim został przekazany."
)
CLAUSE_A = (
    " Wszelkie uszkodzenia lub zmiany w stanie sprzętu, które nie wynikają "
    "z normalnego użytkowania, zostaną udokumentowane i mogą skutkować "
    "dodatkowymi opłatami."
)
CLAUSE_B = (
    " Osoba odpowiedzialna za sprzęt będzie ponosić wszelkie opłaty związane "
    "z przywróceniem sprzętu do pierwotnego stanu, jeśli zostanie on oddany "
    "w stanie odmiennym od tego, w jakim został przekazany."
)
CLAUSE_SIGNING = (
    "Podpisując niniejszy protokół, osoba odbierająca sprzęt potwierdza, że "
    "zapoznała się z jego aktualnym stanem i akceptuje odpowiedzialność za "
    "utrzymanie go w tym samym stanie podczas użytkowania."
)
CLAUSE_C = (
    " Strony ustalają, że niniejszy protokół może być zawarty i podpisany "
    "w formie elektronicznej z wykorzystaniem podpisu zaufanego, o którym mowa "
    "w art. 3 pkt 14 ustawy z dnia 17 lutego 2005 r. o informatyzacji "
    "działalności podmiotów realizujących zadania publiczne. Strony uznają, "
    "że podpis zaufany użyty przez każdą ze stron jest równoważny z podpisem "
    "własnoręcznym dla celów prawnych i dowodowych."
)
CLAUSE_2 = (
    "Osoba kończąca świadczenie usług zobowiązana jest do zwrotu powierzonego "
    "sprzętu osobiście w siedzibie biura, znajdującej się przy Alejach "
    "Jerozolimskich 180 w Warszawie. W przypadku braku możliwości osobistego "
    "zwrotu, dopuszcza się przesłanie sprzętu na własny koszt za pośrednictwem "
    "firmy kurierskiej, na adres siedziby: B2B.NET S.A. Aleje Jerozolimskie "
    "180, 02-486 Warszawa"
)
CLAUSE_3 = (
    "3. Zabrania się dokonywania jakichkolwiek zmian lub ingerencji "
    "w przekazany sprzęt; naruszenie zakazu skutkuje odpowiedzialnością "
    "materialną oraz obowiązkiem pokrycia kosztów przywrócenia sprzętu do "
    "stanu pierwotnego."
)
COPIES = "Powyższy protokół został sporządzony w 2 jednobrzmiących egzemplarzach. "
RETURN_RECEIPT = (
    "Strona przyjmująca potwierdza odbiór sprzętu wymienionego w tabeli "
    "w stanie opisanym w kolumnie „Stan przy zwrocie”."
)

_BODY = "Cambria"
_DATA = "Calibri"
_LINE = Emu(190500)  # 15 pt — interlinia akapitów wzoru

_ISSUE_COLUMNS = ("Lp.", "Nazwa sprzętu", "Nr seryjny", "Ilość", "Cena", "Wartość")
_RETURN_COLUMNS = (
    "Lp.",
    "Nazwa sprzętu",
    "Nr seryjny",
    "Ilość",
    "Stan przy zwrocie",
    "Uwagi",
)
#: Twipy jak we wzorze, tylko „Lp.” szersze (418 łamało nagłówek w PDF).
_COLUMN_WIDTHS = (567, 2811, 2643, 992, 1843, 1275)


@dataclass(frozen=True)
class Run:
    text: str
    bold: bool = False
    size: float | None = 9.0
    font: str = _BODY


def _new_document() -> docx.document.Document:
    document = docx.Document()
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = 1.0
    section = document.sections[0]
    section.page_width = Emu(7560310)
    section.page_height = Emu(10692130)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Emu(457200))
    props = document.core_properties
    props.author = "NEXUS"
    props.last_modified_by = "NEXUS"
    props.title = "Protokół zdawczo-odbiorczy"
    props.comments = ""
    return document


def _paragraph(
    document: docx.document.Document,
    *runs: Run,
    center: bool = False,
    spaced: bool = True,
) -> None:
    paragraph = document.add_paragraph()
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if spaced:
        paragraph.paragraph_format.line_spacing = _LINE
    for spec in runs:
        run = paragraph.add_run(spec.text)
        run.bold = spec.bold
        run.font.name = spec.font
        if spec.size is not None:
            run.font.size = Pt(spec.size)


def _empty(document: docx.document.Document, *, center: bool = False) -> None:
    _paragraph(document, center=center)


def _cell(cell, *runs: Run, center: bool = False) -> None:
    paragraph = cell.paragraphs[0]
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for spec in runs:
        run = paragraph.add_run(spec.text)
        run.bold = spec.bold
        run.font.name = spec.font
        if spec.size is not None:
            run.font.size = Pt(spec.size)


def _table(
    document: docx.document.Document,
    columns: tuple[str, ...],
    data: tuple[str, ...],
    total: tuple[Run | None, Run | None] | None,
) -> None:
    rows = 4 + (1 if total else 0)
    table = document.add_table(rows=rows, cols=len(columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Komórki przez docx_cells (strażnik `test_docx_table_cells.py`); tabela
    # jest świeża, bez scaleń, więc to te same komórki co w `row.cells`.
    cells = [row_cells(row) for row in table.rows]
    _fix_widths(table, _COLUMN_WIDTHS)
    for row in cells:
        for index, cell in enumerate(row):
            cell.width = Twips(_COLUMN_WIDTHS[index])
    for index, title in enumerate(columns):
        _cell(cells[0][index], Run(title, bold=True), center=True)
    # docxtpl zamienia CAŁY wiersz z „{%tr …%}” na znacznik pętli.
    _cell(cells[1][0], Run("{%tr for line in lines %}"))
    for index, field in enumerate(data):
        size, font = (9.0, _BODY) if index in (0, 3) else (11.0, _DATA)
        _cell(
            cells[2][index],
            Run("{{ line.%s }}" % field, size=size, font=font),
            center=index >= 2,
        )
    _cell(cells[3][0], Run("{%tr endfor %}"))
    if total:
        label, value = total
        if label:
            _cell(
                cells[4][4],
                label,
                Run("{{ total }} zł", size=11.0, font=_DATA),
            )
        if value:
            _cell(cells[4][5], value)


def _fix_widths(table: Table, widths: tuple[int, ...]) -> None:
    """Siatka tabeli = szerokości kolumn. LibreOffice (PDF) czyta ``gridCol``,
    nie ``tcW`` komórek — bez tego kolumny wychodzą równe i IMEI łamie się
    na trzy wiersze."""
    table.autofit = False
    grid = table._tbl.tblGrid
    for col, width in zip(grid.findall(qn("w:gridCol")), widths):
        col.set(qn("w:w"), str(width))
    props = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    props.append(layout)


def _signatures(document: docx.document.Document, kind: str) -> None:
    """Dwie kolumny bez ramek: strona, imię i nazwisko, linia podpisu
    i — po potwierdzeniu w NEXUSIE — dopisek, kto i kiedy potwierdził."""
    _empty(document)
    _empty(document)
    left, right = ("issuer", "holder") if kind == "wydanie" else ("holder", "issuer")
    table = document.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _fix_widths(table, (5233, 5233))
    rows = [row_cells(row) for row in table.rows]
    labels = ("Strona przekazująca", "Strona przyjmująca")
    for index, side in enumerate((left, right)):
        _cell(rows[0][index], Run(labels[index], bold=True, size=11.0), center=True)
        _cell(
            rows[1][index],
            Run("{{ %s_name }}" % side, size=10.0, font=_DATA),
            center=True,
        )
        _cell(rows[2][index], Run("." * 60, size=9.0), center=True)
        _cell(rows[3][index], Run("(podpis)", size=9.0), center=True)
        paragraph = rows[3][index].add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run("{{ %s_signature }}" % side)
        run.font.size = Pt(8)
        run.font.name = _DATA


def _header(document: docx.document.Document) -> None:
    _paragraph(document, Run("PROTOKÓŁ", bold=True, size=None), center=True)
    _paragraph(document, Run("ZDAWCZO-ODBIORCZY", bold=True, size=None), center=True)
    _paragraph(document, Run("nr {{ number }}", size=None), center=True)
    _empty(document, center=True)
    _paragraph(
        document,
        Run("Sporządzony w dniu {{ protocol_date }} r. roku w Warszawie,", size=None),
    )
    _empty(document)


def build_issue_template() -> docx.document.Document:
    """Wzór recepcji: protokół wydania sprzętu."""
    document = _new_document()
    _header(document)
    _paragraph(document, Run("Strona przekazująca: ", bold=True, size=None))
    _paragraph(
        document,
        Run(COMPANY_LINE, size=None),
        Run("{{ issuer_name }}", bold=True, size=None, font=_DATA),
    )
    _empty(document)
    _paragraph(
        document,
        Run("Strona przyjmująca ", bold=True, size=None),
        Run("{{ holder_name }}", size=11.0, font=_DATA),
        spaced=False,
    )
    _paragraph(document, Run("PRZEDMIOT PRZEKAZANIA", bold=True, size=None))
    _empty(document)
    _table(
        document,
        _ISSUE_COLUMNS,
        ("position", "name", "identifier", "quantity", "price", "value"),
        (Run("Razem ", bold=True), Run("{{ total }} zł", size=11.0, font=_DATA)),
    )
    _empty(document)
    _paragraph(
        document,
        Run("Wartość przedmiotu przekazania słownie", size=None),
        Run(":", size=None),
        Run(" ", size=None),
        Run("{{ total_words }}.", size=None),
    )
    _empty(document)
    _paragraph(
        document,
        Run("Pan"),
        Run("/Pani"),
        Run(" ", bold=True, size=None, font=_DATA),
        Run("{{ holder_name }}", bold=True, size=None, font=_DATA),
        Run(" "),
        Run("oświadcza"),
        Run(LIABILITY_TAIL),
        spaced=False,
    )
    _empty(document)
    _paragraph(
        document,
        Run("Pan"),
        Run("/Pani"),
        Run(" "),
        Run("{{ holder_name }}", bold=True, size=None, font=_DATA),
        Run(" "),
        Run("zapoznał"),
        Run("/-a"),
        Run(" "),
        Run(CONDITION_TAIL),
        spaced=False,
    )
    _empty(document, center=False)
    _paragraph(document, Run("1.", bold=True), Run(CLAUSE_1), spaced=False)
    _paragraph(document, spaced=False)
    _paragraph(document, Run("a.", bold=True), Run(CLAUSE_A), spaced=False)
    _paragraph(document, Run("b.", bold=True), Run(CLAUSE_B), spaced=False)
    _paragraph(document, Run(CLAUSE_SIGNING))
    _paragraph(document, Run("c.", bold=True), Run(CLAUSE_C))
    _empty(document)
    _paragraph(document, Run("2.", bold=True), Run(CLAUSE_2))
    _empty(document)
    _paragraph(document, Run(CLAUSE_3))
    _empty(document)
    _paragraph(document, Run(COPIES))
    _signatures(document, "wydanie")
    return document


def build_return_template() -> docx.document.Document:
    """Zwrot: nagłówek wzoru, strony zamienione, stan przy zwrocie w tabeli."""
    document = _new_document()
    _header(document)
    _paragraph(document, Run("Strona przekazująca: ", bold=True, size=None))
    _paragraph(document, Run("{{ holder_name }}", bold=True, size=None, font=_DATA))
    _empty(document)
    _paragraph(document, Run("Strona przyjmująca ", bold=True, size=None))
    _paragraph(
        document,
        Run(COMPANY_LINE, size=None),
        Run("{{ issuer_name }}", bold=True, size=None, font=_DATA),
    )
    _empty(document)
    _paragraph(document, Run("PRZEDMIOT ZWROTU", bold=True, size=None))
    _empty(document)
    _table(
        document,
        _RETURN_COLUMNS,
        ("position", "name", "identifier", "quantity", "condition", "notes"),
        None,
    )
    _empty(document)
    _paragraph(document, Run(RETURN_RECEIPT))
    _empty(document)
    _paragraph(document, Run(COPIES))
    _signatures(document, "zwrot")
    return document


def write_templates(target: Path = TEMPLATE_DIR) -> list[Path]:
    target.mkdir(parents=True, exist_ok=True)
    written = []
    for name, builder in (
        (ISSUE_TEMPLATE, build_issue_template),
        (RETURN_TEMPLATE, build_return_template),
    ):
        path = target / name
        builder().save(str(path))
        written.append(path)
    return written


# ── Wzory wgrane przez recepcję (etap 3, O3) ──────────────────────────────────

#: Limit pliku wzoru (spakowanego) i jego treści po rozpakowaniu (zip bomb).
MAX_UPLOAD_BYTES = 2 * 1024 * 1024
MAX_UNPACKED_BYTES = 40 * 1024 * 1024
MAX_ZIP_ENTRIES = 400
DOCX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)

#: Pola protokołu (kontekst ``protocol_document.context_for``) z opisem dla
#: recepcji — lista na ekranie „Wzory protokołów”.
FIELDS: dict[str, str] = {
    "number": "numer protokołu (PZO/rok/miesiąc/kolejny)",
    "protocol_date": "data protokołu (dd.mm.rrrr)",
    "holder_name": "imię i nazwisko osoby, która odbiera albo zwraca sprzęt",
    "issuer_name": "kto z recepcji wydaje albo przyjmuje",
    "lines": "tabela pozycji: wiersz {%tr for line in lines %} … {%tr endfor %}",
    "total": "suma wartości (np. 3 903,73)",
    "total_words": "suma słownie",
    "holder_signature": "dopisek pod podpisem osoby po potwierdzeniu linkiem (puste przed)",
    "issuer_signature": "dopisek pod podpisem recepcji po potwierdzeniu w NEXUSIE (puste przed)",
}
#: Pola wiersza tabeli (``line.<pole>``).
LINE_FIELDS: dict[str, str] = {
    "position": "Lp.",
    "name": "nazwa sprzętu",
    "identifier": "numer seryjny, IMEI albo numer karty",
    "quantity": "ilość",
    "price": "cena",
    "value": "wartość",
    "condition": "stan przy zwrocie",
    "notes": "uwagi",
}
#: Bez tych pól protokołu nie da się odszukać ani przypisać osobie.
REQUIRED_FIELDS: dict[str, tuple[str, ...]] = {
    "wydanie": ("number", "protocol_date", "holder_name", "lines"),
    "zwrot": ("number", "protocol_date", "holder_name", "lines"),
}
_MACRO_PARTS = ("vbaproject.bin", "vbadata.xml")


class TemplateRejected(ValueError):
    """Wgrany wzór nie nadaje się do użycia — zdanie dla recepcji."""


#: Filtry, które nie budują dużych napisów z małego wzoru (bez ``center``,
#: ``indent``, ``format``, ``replace``…).
_SAFE_FILTERS = frozenset(
    {
        "abs",
        "capitalize",
        "count",
        "d",
        "default",
        "e",
        "escape",
        "first",
        "float",
        "int",
        "join",
        "last",
        "length",
        "lower",
        "round",
        "string",
        "title",
        "trim",
        "truncate",
        "upper",
    }
)
#: Węzły, z których może się składać wzór z pliku (lista dozwolonych, nie
#: zakazanych): pola, tekst, warunki, porównania, filtry z ``_SAFE_FILTERS``
#: i JEDNA pętla po ``lines``. Bez wywołań metod (``ljust``), przypisań
#: (``set s = s ~ s``), sklejania ``~``, arytmetyki (mnożenie wielkich liczb
#: zajmuje procesor; kwoty przychodzą już jako tekst) i zagnieżdżonych pętli —
#: każda z tych konstrukcji pozwala małym wzorem zająć pamięć albo procesor.
_ALLOWED_NODES = (
    nodes.Template,
    nodes.Output,
    nodes.TemplateData,
    nodes.Name,
    nodes.Const,
    nodes.Getattr,
    nodes.Getitem,
    nodes.Slice,
    nodes.Filter,
    nodes.Test,
    nodes.Keyword,
    nodes.If,
    nodes.CondExpr,
    nodes.Compare,
    nodes.Operand,
    nodes.And,
    nodes.Or,
    nodes.Not,
    nodes.List,
    nodes.Tuple,
    nodes.For,
)
#: Górna granica liczby węzłów — standardowe wzory mają ich ok. 45.
MAX_TEMPLATE_NODES = 2000
#: Jedyna lista, po której wzór może iterować (pozycje protokołu).
LOOP_FIELD = "lines"
_SEQUENCES = (str, bytes, list, tuple)


def _reject_node(message: str, node: nodes.Node, name: Any, filename: Any) -> None:
    raise TemplateSyntaxError(message, node.lineno, name, filename)


def _check_tree(tree: nodes.Template, name: Any, filename: Any) -> None:
    count = 0
    for node in tree.find_all(nodes.Node):
        count += 1
        if count > MAX_TEMPLATE_NODES:
            _reject_node("Wzór ma za dużo pól i warunków.", node, name, filename)
        if not isinstance(node, _ALLOWED_NODES):
            _reject_node(
                "We wzorze są dostępne tylko pola, warunki i jedna pętla po "
                "pozycjach protokołu.",
                node,
                name,
                filename,
            )
        if isinstance(node, nodes.For):
            if (
                node.recursive
                or not isinstance(node.iter, nodes.Name)
                or node.iter.name != LOOP_FIELD
                or not isinstance(node.target, nodes.Name)
            ):
                _reject_node(
                    "Pętla we wzorze może iść tylko po pozycjach protokołu "
                    "(for line in lines).",
                    node,
                    name,
                    filename,
                )
            if next(iter(node.find_all(nodes.For)), None) is not None:
                _reject_node(
                    "Pętla w pętli nie jest dostępna we wzorze.", node, name, filename
                )


class _ProtocolSandbox(SandboxedEnvironment):
    """Sandbox wzorów z pliku: atrybuty Pythona blokuje jinja2, a ta klasa
    dokłada listę dozwolonych węzłów (``_check_tree``), brak funkcji
    globalnych, krótką listę filtrów oraz brak potęgowania, mnożenia,
    dodawania i formatowania napisów (pamięć)."""

    intercepted_binops = frozenset({"*", "**", "%", "+"})

    def __init__(self, **options: Any) -> None:
        super().__init__(**options)
        self.globals.clear()
        self.filters = {
            name: value for name, value in self.filters.items() if name in _SAFE_FILTERS
        }

    def _parse(self, source: str, name: Any, filename: Any) -> nodes.Template:
        tree = super()._parse(source, name, filename)
        _check_tree(tree, name, filename)
        return tree

    def call_binop(self, context: Any, operator: str, left: Any, right: Any) -> Any:
        if operator == "**":
            raise SecurityError("Potęgowanie nie jest dostępne we wzorze.")
        if isinstance(left, _SEQUENCES) or isinstance(right, _SEQUENCES):
            raise SecurityError(
                "Mnożenie, dodawanie i formatowanie tekstu nie są dostępne we wzorze."
            )
        return super().call_binop(context, operator, left, right)


def sandboxed_env() -> SandboxedEnvironment:
    """Środowisko jinja2 dla wzorów z pliku (każde wypełnienie — nowe)."""
    return _ProtocolSandbox(autoescape=True, undefined=StrictUndefined)


#: Osadzone obiekty (OLE, ActiveX) — nie we wzorze protokołu.
_EMBEDDED_PREFIXES = ("word/embeddings/", "word/activex/")
_RELATION = re.compile(r"<Relationship\b[^>]*>", re.IGNORECASE)


def _external_relation(rels_xml: str) -> bool:
    """Relacja do zasobu spoza paczki (dołączony szablon, obraz z sieci, OLE).

    Zwykłe hiperłącza zostają — LibreOffice ich nie pobiera przy konwersji.
    """

    for match in _RELATION.finditer(rels_xml):
        tag = match.group(0)
        if not re.search(r"TargetMode\s*=\s*[\"']External[\"']", tag):
            continue
        if not re.search(r"Type\s*=\s*[\"'][^\"']*/hyperlink[\"']", tag):
            return True
    return False


def _check_package(content: bytes) -> None:
    try:
        package = zipfile.ZipFile(io.BytesIO(content))
    except zipfile.BadZipFile as exc:
        raise TemplateRejected("To nie jest plik Word (.docx).") from exc
    try:
        with package:
            _check_parts(package)
    except TemplateRejected:
        raise
    except (KeyError, RuntimeError, NotImplementedError, zipfile.BadZipFile) as exc:
        # Część pod inną wielkością liter, zaszyfrowana albo uszkodzona.
        raise TemplateRejected(
            "Nie da się odczytać pliku Word — zapisz go w Wordzie od nowa."
        ) from exc


def _check_parts(package: zipfile.ZipFile) -> None:
    entries = package.infolist()
    if len(entries) > MAX_ZIP_ENTRIES:
        raise TemplateRejected("Plik ma za dużo części — zapisz go w Wordzie od nowa.")
    if sum(entry.file_size for entry in entries) > MAX_UNPACKED_BYTES:
        raise TemplateRejected("Plik po rozpakowaniu jest za duży.")
    names = {entry.filename.lower() for entry in entries}
    if "word/document.xml" not in names or "[content_types].xml" not in names:
        raise TemplateRejected("To nie jest plik Word (.docx).")
    if any(name.rsplit("/", 1)[-1] in _MACRO_PARTS for name in names):
        raise TemplateRejected(
            "Wzór z makrami nie jest przyjmowany — zapisz go jako .docx."
        )
    types_name = next(
        entry.filename
        for entry in entries
        if entry.filename.lower() == "[content_types].xml"
    )
    types = package.read(types_name).decode("utf-8", "replace")
    if "macroEnabled" in types:
        raise TemplateRejected(
            "Wzór z makrami nie jest przyjmowany — zapisz go jako .docx."
        )
    if any(name.startswith(_EMBEDDED_PREFIXES) for name in names):
        raise TemplateRejected(
            "Wzór z osadzonymi obiektami nie jest przyjmowany — "
            "zapisz go jako zwykły .docx."
        )
    for entry in entries:
        if entry.filename.lower().endswith(".rels") and _external_relation(
            package.read(entry).decode("utf-8", "replace")
        ):
            raise TemplateRejected(
                "Wzór odwołuje się do plików spoza dokumentu (szablon, "
                "obraz z sieci) — zapisz go bez łączy zewnętrznych."
            )


def _sample_context(kind: str) -> dict[str, Any]:
    """Fikcyjny protokół do próbnego wypełnienia (te same pola co prawdziwy)."""
    from app.services.equipment import protocol_document

    view = protocol_document.ProtocolView(
        kind=kind,
        number="PZO/2026/01/001",
        protocol_date=date(2026, 1, 15),
        holder_name="Anna Przykładowa",
        issuer_name="Recepcja Przykładowa",
        lines=(
            protocol_document.LineView(
                position=1,
                name="Laptop przykładowy 14 cali",
                identifier="TEST0001",
                quantity=1,
                unit_value=Decimal("1234.56"),
                condition="sprawny" if kind == "zwrot" else None,
            ),
        ),
    )
    return protocol_document.context_for(view)


def _short(error: Exception) -> str:
    text_value = " ".join(str(error).split())
    return text_value[:200] or error.__class__.__name__


def template_fields(content: bytes) -> set[str]:
    """Pola najwyższego poziomu użyte we wzorze (bez zmiennych pętli)."""
    try:
        return set(
            DocxTemplate(io.BytesIO(content)).get_undeclared_template_variables(
                jinja_env=sandboxed_env()
            )
        )
    except TemplateError as exc:
        raise TemplateRejected(f"Błąd w polach wzoru: {_short(exc)}") from exc
    except Exception as exc:  # uszkodzony DOCX (python-docx, lxml)
        raise TemplateRejected("Nie da się odczytać pliku Word.") from exc


def render_uploaded(content: bytes, context: dict[str, Any]) -> bytes:
    """Wzór z pliku wypełniony w sandboxie (``SecurityError`` / ``TemplateError`` lecą dalej)."""
    template = DocxTemplate(io.BytesIO(content))
    template.render(context, jinja_env=sandboxed_env())
    buffer = io.BytesIO()
    template.save(buffer)
    return buffer.getvalue()


def check_upload(filename: str, content: bytes, kind: str) -> list[str]:
    """Walidacja wgranego wzoru; zwraca użyte pola albo ``TemplateRejected``."""
    if kind not in REQUIRED_FIELDS:
        raise TemplateRejected("Wybierz rodzaj wzoru: wydanie albo zwrot.")
    if not (filename or "").lower().endswith(".docx"):
        raise TemplateRejected("Wgraj plik Word w formacie .docx.")
    if not content or len(content) > MAX_UPLOAD_BYTES:
        raise TemplateRejected("Plik jest pusty albo większy niż 2 MB.")
    _check_package(content)
    fields = template_fields(content)
    unknown = sorted(fields - set(FIELDS))
    if unknown:
        raise TemplateRejected(
            "Wzór ma pola, których protokół nie zna: "
            + ", ".join(unknown)
            + ". Dostępne pola: "
            + ", ".join(FIELDS)
            + "."
        )
    missing = [field for field in REQUIRED_FIELDS[kind] if field not in fields]
    if missing:
        raise TemplateRejected(
            "Wzór nie ma wymaganych pól: " + ", ".join(missing) + "."
        )
    try:
        render_uploaded(content, _sample_context(kind))
    except SecurityError as exc:
        raise TemplateRejected(
            "Wzór próbuje sięgnąć poza dane protokołu — takiego wzoru nie da się użyć."
        ) from exc
    except TemplateError as exc:
        raise TemplateRejected(f"Wzór nie daje się wypełnić: {_short(exc)}") from exc
    except Exception as exc:
        raise TemplateRejected("Wzór nie daje się wypełnić.") from exc
    return sorted(fields)


def standard_template_bytes(kind: str) -> bytes:
    """Wbudowany standard — punkt wyjścia do własnego wzoru recepcji."""
    name = ISSUE_TEMPLATE if kind == "wydanie" else RETURN_TEMPLATE
    return (TEMPLATE_DIR / name).read_bytes()


def field_rows() -> Iterable[dict[str, str]]:
    """Pola do listy na ekranie „Wzory protokołów” (pole, opis, gdzie)."""
    for key, label in FIELDS.items():
        yield {"field": key, "label": label, "scope": "protocol"}
    for key, label in LINE_FIELDS.items():
        yield {"field": f"line.{key}", "label": label, "scope": "line"}


# ── Rejestr wzorów (ekran „Sprzęt → Wzory protokołów”) ───────────────────────

KIND_LABELS = {"wydanie": "wydania", "zwrot": "zwrotu"}


def template_payload(
    template: EquipmentProtocolTemplate, *, packs: int = 0, protocols: int = 0
) -> dict[str, Any]:
    return {
        "id": template.id,
        "kind": template.kind,
        "name": template.name,
        "description": template.description,
        "file_name": template.file_name,
        "file_size": template.file_size,
        "fields": list(template.fields or []),
        "is_active": template.is_active,
        "created_at": template.created_at.isoformat() if template.created_at else None,
        "packs": packs,
        "protocols": protocols,
    }


async def _usage(
    db: AsyncSession, ids: list[int]
) -> tuple[dict[int, int], dict[int, int]]:
    if not ids:
        return {}, {}
    packs: dict[int, int] = {}
    for column in (
        ClientOnboardingPack.issue_template_id,
        ClientOnboardingPack.return_template_id,
    ):
        for template_id, count in (
            await db.execute(
                select(column, func.count()).where(column.in_(ids)).group_by(column)
            )
        ).all():
            packs[template_id] = packs.get(template_id, 0) + int(count)
    protocols = {
        template_id: int(count)
        for template_id, count in (
            await db.execute(
                select(EquipmentProtocol.template_id, func.count())
                .where(EquipmentProtocol.template_id.in_(ids))
                .group_by(EquipmentProtocol.template_id)
            )
        ).all()
    }
    return packs, protocols


async def list_templates(db: AsyncSession) -> list[dict[str, Any]]:
    """Wgrane wzory: najpierw w użyciu, potem wycofane (z liczbą pakietów i protokołów)."""
    rows = (
        await db.scalars(
            select(EquipmentProtocolTemplate).order_by(
                EquipmentProtocolTemplate.is_active.desc(),
                EquipmentProtocolTemplate.kind,
                EquipmentProtocolTemplate.name,
            )
        )
    ).all()
    packs, protocols = await _usage(db, [row.id for row in rows])
    return [
        template_payload(
            row, packs=packs.get(row.id, 0), protocols=protocols.get(row.id, 0)
        )
        for row in rows
    ]


def _clean_name(value: Any) -> str:
    name = " ".join(str(value or "").split())
    if not name:
        raise EquipmentError("Nazwij wzór, np. „Standard z kaucją”.")
    if len(name) > 120:
        raise EquipmentError("Nazwa wzoru może mieć najwyżej 120 znaków.")
    return name


def _clean_description(value: Any) -> Optional[str]:
    text_value = " ".join(str(value or "").split())
    if len(text_value) > 300:
        raise EquipmentError("Opis wzoru może mieć najwyżej 300 znaków.")
    return text_value or None


#: Unikat aktywnej nazwy w bazie (``equipment_schema``) — ostatnia linia obrony.
_ACTIVE_NAME_INDEX = "ux_equipment_protocol_templates_active_name"


def _duplicate_name(kind: str, name: str) -> EquipmentConflict:
    return EquipmentConflict(
        f"Wzór {KIND_LABELS[kind]} o nazwie „{name}” już jest — wybierz inną nazwę."
    )


async def _ensure_unique_name(
    db: AsyncSession, kind: str, name: str, *, exclude_id: Optional[int] = None
) -> None:
    query = select(EquipmentProtocolTemplate.id).where(
        EquipmentProtocolTemplate.kind == kind,
        EquipmentProtocolTemplate.is_active.is_(True),
        func.lower(EquipmentProtocolTemplate.name) == name.lower(),
    )
    if exclude_id is not None:
        query = query.where(EquipmentProtocolTemplate.id != exclude_id)
    if await db.scalar(query.limit(1)) is not None:
        raise _duplicate_name(kind, name)


async def _flush_unique(
    db: AsyncSession, template: EquipmentProtocolTemplate, *, add: bool = False
) -> None:
    """Zapis wzoru; dwa równoczesne wgrania tej samej nazwy przechodzą
    sprawdzenie obok siebie — drugie zatrzymuje unikat bazy, a osoba dostaje
    409 z powodem zamiast 500."""

    # Przed zapisem: po wycofaniu punktu zapisu atrybuty wzoru wygasają.
    kind, name = template.kind, template.name
    try:
        async with db.begin_nested():
            if add:
                db.add(template)
            await db.flush()
    except IntegrityError as exc:
        if _ACTIVE_NAME_INDEX not in str(exc.orig):
            raise
        raise _duplicate_name(kind, name) from exc


async def create_template(
    db: AsyncSession,
    *,
    kind: str,
    name: Any,
    description: Any,
    filename: str,
    content: bytes,
    actor: User,
) -> EquipmentProtocolTemplate:
    """Nowy wzór z pliku po ``check_upload`` (odmowa = ``EquipmentError`` 422)."""
    from app.services.hr import audit

    clean = _clean_name(name)
    try:
        fields = check_upload(filename, content, kind)
    except TemplateRejected as exc:
        raise EquipmentError(str(exc)) from exc
    await _ensure_unique_name(db, kind, clean)
    template = EquipmentProtocolTemplate(
        kind=kind,
        name=clean,
        description=_clean_description(description),
        file_name=Path(filename).name[:255] or "wzor.docx",
        content=content,
        file_sha256=hashlib.sha256(content).hexdigest(),
        file_size=len(content),
        fields=fields,
        created_by=actor.id,
        updated_by=actor.id,
    )
    await _flush_unique(db, template, add=True)
    audit.record(
        db,
        actor_user_id=actor.id,
        subject_user_id=None,
        entity_type="equipment_protocol_template",
        entity_id=template.id,
        action="template_uploaded",
        details={"kind": kind, "fields": fields},
    )
    return template


async def get_template(db: AsyncSession, template_id: int) -> EquipmentProtocolTemplate:
    template = await db.get(EquipmentProtocolTemplate, template_id)
    if template is None:
        raise EquipmentNotFound("Nie ma takiego wzoru protokołu.")
    return template


async def update_template(
    db: AsyncSession, template_id: int, values: Mapping[str, Any], *, actor: User
) -> EquipmentProtocolTemplate:
    """Nazwa, opis i wycofanie. Treści pliku się nie zmienia — nowy plik = nowy wzór."""
    from app.services.hr import audit

    template = await get_template(db, template_id)
    if "name" in values and values["name"] is not None:
        name = _clean_name(values["name"])
        if template.is_active:
            await _ensure_unique_name(db, template.kind, name, exclude_id=template.id)
        template.name = name
    if "description" in values:
        template.description = _clean_description(values["description"])
    if values.get("is_active") is False and template.is_active:
        clients = (
            await db.scalars(
                select(ClientOnboardingPack.client_id).where(
                    or_(
                        ClientOnboardingPack.issue_template_id == template.id,
                        ClientOnboardingPack.return_template_id == template.id,
                    )
                )
            )
        ).all()
        if clients:
            raise EquipmentConflict(
                f"Wzór jest wybrany w pakiecie {len(clients)} "
                + ("klienta" if len(clients) == 1 else "klientów")
                + " — najpierw wybierz tam inny wzór."
            )
        template.is_active = False
    elif values.get("is_active") is True and not template.is_active:
        await _ensure_unique_name(
            db, template.kind, template.name, exclude_id=template.id
        )
        template.is_active = True
    template.updated_by = actor.id
    await _flush_unique(db, template)
    audit.record(
        db,
        actor_user_id=actor.id,
        subject_user_id=None,
        entity_type="equipment_protocol_template",
        entity_id=template.id,
        action="template_updated",
        details={"fields": sorted(values), "is_active": template.is_active},
    )
    return template


async def template_content(
    db: AsyncSession, template_id: int
) -> tuple[EquipmentProtocolTemplate, bytes]:
    template = await get_template(db, template_id)
    content = await db.scalar(
        select(EquipmentProtocolTemplate.content).where(
            EquipmentProtocolTemplate.id == template_id
        )
    )
    return template, content or b""


def sample_document(content: bytes, kind: str) -> bytes:
    """Wzór wypełniony fikcyjnym protokołem — podgląd dla recepcji."""
    try:
        return render_uploaded(content, _sample_context(kind))
    except (SecurityError, TemplateError) as exc:
        raise EquipmentError(
            "Tego wzoru nie da się wypełnić — wgraj poprawiony plik."
        ) from exc


if __name__ == "__main__":
    for written_path in write_templates():
        print(f"Zapisano {written_path}")
