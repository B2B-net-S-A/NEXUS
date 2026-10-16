"""Import plików recepcji (Excel) przez arkusz kontrolny (krok S2).

Przebieg: wgranie pliku → partia z wierszami i decyzjami (nic nie trafia do
rejestru) → arkusz kontrolny XLSX dla recepcji → poprawki decyzji →
zatwierdzenie partii (zapis). Ponowne zatwierdzenie zapisuje tylko wiersze,
które od tamtej pory dostały decyzję „przyjmij”.

Parser czyta kolumny po nagłówkach (synonimy z plików recepcji: „Numer seryjny
/ Service tag”, „Nr Laptopa”, „Numer usługi”, „NR. KARTY”…). Kolumna z sekretem
w nagłówku (klucz BitLocker, PUK, PIN, PESEL, hasło) nie jest czytana wcale;
wartość wyglądająca na sekret w czytanej kolumnie znika, a wiersz dostaje
powód. Kolumny bez znanego nagłówka są pomijane (w raporcie tylko nazwa).
Plik nie jest przechowywany — partia ma jego SHA-256 i oczyszczone wiersze.

Decyzje: ``przyjmij`` (egzemplarz powstanie, z posiadaczem albo na stanie),
``pomin`` (brak numeru egzemplarza), ``konflikt`` (numer już w rejestrze
albo dwa razy w pliku), ``do_wyjasnienia`` (osoby z pliku nie ma w NEXUSIE;
laptop bez posiadacza i bez stanu).
"""

from __future__ import annotations

import asyncio
import hashlib
import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from typing import Any, Iterable, Optional

import openpyxl
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.candidate import Candidate
from app.models.contract import Contract
from app.models.equipment import (
    EquipmentAssignment,
    EquipmentImportBatch,
    EquipmentImportRow,
    EquipmentItem,
)
from app.models.user import User
from app.core.export_safety import safe_row
from app.core.zip_guard import assert_safe_ooxml
from app.services import equipment_schema as schema
from app.services.equipment import identifiers, registry, sanitize
from app.services.equipment.registry import (
    ISSUABLE_CONTRACT_STATUSES,
    EquipmentConflict,
    EquipmentError,
    EquipmentNotFound,
)
from app.services.hr import audit

MAX_FILE_BYTES = 10 * 1024 * 1024
HEADER_SCAN_ROWS = 6
MAX_ROWS_PER_SHEET = 5000

IDENTIFIER_FIELDS = schema.ITEM_UNIQUE_COLUMNS

#: Synonimy nagłówków (po ``sanitize.fold``) — kolejność = pierwszeństwo, gdy
#: dwie kolumny pasują do jednego pola (np. w pliku kart „Użytkownik” to firma,
#: a osoba stoi w „Karta przypisana”).
FIELD_SYNONYMS: dict[str, tuple[str, ...]] = {
    "inventory_number": (
        "nr laptopa",
        "nr ewidencyjny",
        "numer ewidencyjny",
        "nr inwentarzowy",
        "numer inwentarzowy",
        "nr sprzetu",
        "numer sprzetu",
    ),
    "serial_number": (
        "numer seryjny / service tag",
        "numer seryjny/service tag",
        "numer seryjny",
        "nr seryjny",
        "service tag",
        "s/n",
        "sn",
        "serial",
    ),
    "imei": ("imei",),
    "iccid": ("iccid", "numer karty sim", "nr karty sim", "nr sim"),
    "phone_number": ("numer uslugi", "numer telefonu", "nr telefonu", "msisdn"),
    "access_card_number": ("nr. karty", "nr karty", "numer karty", "nr karty dostepu"),
    "category": ("rodzaj urzadzenia", "typ urzadzenia", "rodzaj", "kategoria", "typ"),
    "model": ("model", "model urzadzenia", "nazwa", "urzadzenie"),
    "holder": (
        "karta przypisana",
        "uzytkownik",
        "posiadacz",
        "pracownik",
        "osoba",
        "kto",
    ),
    "project": ("projekt", "klient", "jaki dzial", "dzial", "firma"),
    "status": ("status laptopa", "status", "stan", "informacja"),
    "purchase_date": ("data zakupu",),
    "price": ("cena", "cena zakupu", "wartosc"),
    "notes": (
        "uwagi",
        "po kim",
        "po kim telefon",
        "dostep",
        "taryfa",
        "ram",
        "procesor",
    ),
}
_SINGLE_FIELDS = tuple(f for f in FIELD_SYNONYMS if f != "notes")

_SHEET_HINTS: tuple[tuple[str, str], ...] = (
    ("laptop", "laptop"),
    ("telefon", "telefon"),
    ("monitor", "monitor"),
    ("sluchawk", "akcesorium"),
    ("airpods", "akcesorium"),
    ("akcesori", "akcesorium"),
    ("sim", "karta_sim"),
    ("kart", "karta_dostepu"),
)
_VALUE_HINTS: tuple[tuple[str, str], ...] = (
    ("laptop", "laptop"),
    ("notebook", "laptop"),
    ("telefon", "telefon"),
    ("smartfon", "telefon"),
    ("monitor", "monitor"),
    ("sim", "karta_sim"),
    ("karta", "karta_dostepu"),
    ("sluchawk", "akcesorium"),
    ("airpods", "akcesorium"),
    ("mysz", "akcesorium"),
    ("stacja", "akcesorium"),
    ("dock", "akcesorium"),
)
_FREE_PREFIXES = ("woln", "na stanie", "biuro", "szafk", "magazyn", "brak", "-")
_STATUS_WORDS: tuple[tuple[str, str], ...] = (
    ("utyliz", "wycofany"),
    ("zlomow", "wycofany"),
    ("wycofan", "wycofany"),
    ("serwis", "serwis"),
    ("naprawa", "serwis"),
    ("zgubion", "zgubiony"),
    ("skradzion", "zgubiony"),
    ("kradziez", "zgubiony"),
)
_FREE_STATUS = (
    "woln",
    "na stanie",
    "do czyszczenia",
    "do przygotowania",
    "gotowy",
    "nowy",
)
_PREFIX = re.compile(r"^(?:s/?n|st|service tag)\s*[:.]?\s*", re.IGNORECASE)
_IMEI_IN_TEXT = re.compile(r"imei\s*:?\s*(\d{14,16})", re.IGNORECASE)


@dataclass
class SheetReport:
    name: str
    header_row: Optional[int] = None
    rows: int = 0
    skipped: Optional[str] = None

    def as_dict(self) -> dict[str, Any]:
        return {
            "name": self.name,
            "header_row": self.header_row,
            "rows": self.rows,
            "skipped": self.skipped,
        }


@dataclass
class ParsedRow:
    sheet: str
    row_number: int
    data: dict[str, Any]
    reasons: list[str] = field(default_factory=list)
    #: Powody, przez które wiersz idzie „do wyjaśnienia” zamiast do rejestru
    #: (np. numer dłuższy niż kolumna — wiersz opisuje prawdziwy egzemplarz).
    needs_review: list[str] = field(default_factory=list)


@dataclass
class ParseResult:
    rows: list[ParsedRow]
    sheets: list[SheetReport]
    removed_columns: list[dict[str, str]]
    ignored_columns: list[dict[str, str]]
    removed_values: dict[str, int]


def _cell_text(value: Any) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    text = " ".join(str(value).split())
    return text or None


def match_header(header: str) -> tuple[Optional[str], int]:
    """(pole, pierwszeństwo) dla nagłówka albo (None, 0)."""
    folded = sanitize.fold(header).strip(" :.")
    for field_name, synonyms in FIELD_SYNONYMS.items():
        for index, synonym in enumerate(synonyms):
            if folded == synonym:
                return field_name, index
    for field_name, synonyms in FIELD_SYNONYMS.items():
        for index, synonym in enumerate(synonyms):
            if len(synonym) >= 4 and folded.startswith(synonym):
                return field_name, index + 50
    return None, 0


def _category_from(
    text: Optional[str], hints: tuple[tuple[str, str], ...]
) -> Optional[str]:
    if not text:
        return None
    folded = sanitize.fold(text)
    for hint, category in hints:
        if hint in folded:
            return category
    return None


@dataclass
class _Columns:
    single: dict[str, int]
    notes: list[int]
    removed: list[tuple[int, str]]
    ignored: list[int]
    headers: dict[int, str]


def _map_columns(header_cells: list[Any]) -> _Columns:
    best: dict[str, tuple[int, int]] = {}
    notes: list[int] = []
    removed: list[tuple[int, str]] = []
    ignored: list[int] = []
    headers: dict[int, str] = {}
    losers: list[tuple[int, str]] = []
    for index, raw in enumerate(header_cells):
        header = _cell_text(raw)
        if not header:
            continue
        headers[index] = header
        reason = sanitize.secret_header_reason(header)
        if reason:
            removed.append((index, reason))
            continue
        field_name, rank = match_header(header)
        if field_name is None:
            ignored.append(index)
            continue
        if field_name == "notes":
            notes.append(index)
            continue
        current = best.get(field_name)
        if current is None or rank < current[1]:
            if current is not None:
                losers.append((current[0], field_name))
            best[field_name] = (index, rank)
        else:
            losers.append((index, field_name))
    for index, field_name in losers:
        # Druga kolumna stanu to dalej opis egzemplarza; drugi „użytkownik” — nie.
        if field_name in ("status", "project"):
            notes.append(index)
        else:
            ignored.append(index)
    return _Columns(
        single={name: index for name, (index, _rank) in best.items()},
        notes=notes,
        removed=removed,
        ignored=ignored,
        headers=headers,
    )


def _header_score(cells: list[Any]) -> int:
    mapped = {match_header(_cell_text(c) or "")[0] for c in cells if _cell_text(c)}
    mapped.discard(None)
    mapped.discard("notes")
    has_key = bool(mapped & {*IDENTIFIER_FIELDS, "model"})
    return len(mapped) if has_key else 0


def _scan_unknown(values: Iterable[Any]) -> Optional[str]:
    for value in values:
        text = _cell_text(value)
        if not text:
            continue
        if sanitize.is_bitlocker_key(text):
            return "klucz BitLocker"
        if sanitize.is_pesel(text):
            return "PESEL"
        if sanitize.is_private_email(text):
            return "prywatny e-mail"
    return None


def _clean_identifier(
    field_name: str, raw: Optional[str], row: ParsedRow
) -> Optional[str]:
    if not raw:
        return None
    cleaned = sanitize.clean_text(raw, identifier=True)
    row.reasons.extend(f"usunięto: {r}" for r in cleaned.reasons)
    if not cleaned.value:
        return None
    value = cleaned.value
    if field_name == "serial_number":
        value = _PREFIX.sub("", value).strip() or None
        if value is None:
            return None
    try:
        return identifiers.normalize(field_name, value)
    except identifiers.IdentifierTooLong as exc:
        row.needs_review.append(
            f"{str(exc).rstrip('.')} — numer pominięty, popraw go w pliku"
        )
        return None
    except identifiers.IdentifierError:
        row.reasons.append(
            f"{identifiers.FIELD_LABELS[field_name]} w złym formacie — pominięty"
        )
        return None


def _clean_free(raw: Optional[str], row: ParsedRow) -> Optional[str]:
    if not raw:
        return None
    cleaned = sanitize.clean_text(raw)
    row.reasons.extend(f"usunięto: {r}" for r in cleaned.reasons)
    return cleaned.value


def _price(raw: Optional[str]) -> Optional[str]:
    if not raw:
        return None
    text = raw.replace("zł", "").replace(" ", "").replace(",", ".")
    try:
        value = Decimal(text)
    except InvalidOperation:
        return None
    return str(value.quantize(Decimal("0.01"))) if value >= 0 else None


def parse_workbook(content: bytes) -> ParseResult:
    """Czysty parser pliku (bez bazy): wiersze z polami i powodami usunięć."""
    if len(content) > MAX_FILE_BYTES:
        raise EquipmentError("Plik jest większy niż 10 MB.")
    if not content.startswith(b"PK"):
        raise EquipmentError("Wgraj plik Excel (.xlsx).")
    try:
        # Strażnik „bomby ZIP” przed openpyxl (sharedStrings ładuje się w całości
        # nawet w read_only) — jak w pozostałych parserach OOXML (R9-N7-13).
        assert_safe_ooxml(content)
        workbook = openpyxl.load_workbook(
            io.BytesIO(content), read_only=True, data_only=True
        )
    except (zipfile.BadZipFile, KeyError, OSError, ValueError) as exc:
        raise EquipmentError("Nie da się odczytać pliku Excel (.xlsx).") from exc
    rows: list[ParsedRow] = []
    sheets: list[SheetReport] = []
    removed_columns: list[dict[str, str]] = []
    ignored_columns: list[dict[str, str]] = []
    removed_values: dict[str, int] = {}
    try:
        for sheet in workbook.worksheets:
            report = SheetReport(name=str(sheet.title).strip()[:120])
            sheets.append(report)
            if sheet.sheet_state != "visible":
                report.skipped = "arkusz ukryty"
                continue
            head = [
                list(r)
                for r in sheet.iter_rows(max_row=HEADER_SCAN_ROWS, values_only=True)
            ]
            scores = [(_header_score(cells), index) for index, cells in enumerate(head)]
            best_score, best_index = max(scores, default=(0, 0))
            if best_score < 2:
                report.skipped = "brak wiersza nagłówków"
                continue
            report.header_row = best_index + 1
            columns = _map_columns(head[best_index])
            unknown_values: dict[int, list[Any]] = {i: [] for i in columns.ignored}
            for index, reason in columns.removed:
                removed_columns.append(
                    {
                        "sheet": report.name,
                        "column": columns.headers[index],
                        "reason": reason,
                    }
                )
            sheet_hint = _category_from(report.name, _SHEET_HINTS)
            for number, values in enumerate(
                sheet.iter_rows(min_row=report.header_row + 1, values_only=True),
                start=report.header_row + 1,
            ):
                if number > report.header_row + MAX_ROWS_PER_SHEET:
                    report.skipped = f"przycięto do {MAX_ROWS_PER_SHEET} wierszy"
                    break
                cells = list(values)

                def get(field_name: str) -> Optional[str]:
                    index = columns.single.get(field_name)
                    if index is None or index >= len(cells):
                        return None
                    return _cell_text(cells[index])

                for index in columns.ignored:
                    if index < len(cells) and len(unknown_values[index]) < 400:
                        unknown_values[index].append(cells[index])
                if not any(_cell_text(c) for c in cells):
                    continue
                parsed = ParsedRow(sheet=report.name, row_number=number, data={})
                data = parsed.data
                serial_raw = get("serial_number")
                imei_raw = get("imei")
                if serial_raw and not imei_raw:
                    found = _IMEI_IN_TEXT.search(serial_raw)
                    if found:
                        imei_raw, serial_raw = found.group(1), None
                for field_name in IDENTIFIER_FIELDS:
                    raw = (
                        imei_raw
                        if field_name == "imei"
                        else (
                            serial_raw
                            if field_name == "serial_number"
                            else get(field_name)
                        )
                    )
                    data[field_name] = _clean_identifier(field_name, raw, parsed)
                category = (
                    _category_from(get("category"), _VALUE_HINTS)
                    or sheet_hint
                    or ("karta_sim" if data["phone_number"] or data["iccid"] else None)
                    or ("karta_dostepu" if data["access_card_number"] else None)
                    or ("telefon" if data["imei"] else None)
                    or "inne"
                )
                data["category"] = category
                model = _clean_free(get("model"), parsed)
                data["model"] = (model or identifiers.CATEGORY_LABELS[category])[:200]
                data["holder_text"] = _clean_free(get("holder"), parsed)
                data["project"] = _clean_free(get("project"), parsed)
                data["status_text"] = _clean_free(get("status"), parsed)
                notes = [
                    _clean_free(
                        _cell_text(cells[i]) if i < len(cells) else None, parsed
                    )
                    for i in columns.notes
                ]
                data["notes"] = " · ".join(n for n in notes if n) or None
                purchase = get("purchase_date")
                data["purchase_date"] = (
                    purchase
                    if purchase and re.fullmatch(r"\d{4}-\d{2}-\d{2}", purchase)
                    else None
                )
                data["price"] = _price(get("price"))
                for reason in parsed.reasons:
                    if reason.startswith("usunięto: "):
                        key = reason.removeprefix("usunięto: ")
                        removed_values[key] = removed_values.get(key, 0) + 1
                report.rows += 1
                rows.append(parsed)
            for index in columns.ignored:
                found = _scan_unknown(unknown_values[index])
                entry = {"sheet": report.name, "column": columns.headers[index]}
                if found:
                    removed_columns.append({**entry, "reason": found})
                else:
                    ignored_columns.append(entry)
    finally:
        workbook.close()
    return ParseResult(rows, sheets, removed_columns, ignored_columns, removed_values)


def holder_state(data: dict[str, Any]) -> tuple[str, Optional[str]]:
    """(stan, tekst posiadacza): stan z kolumny stanu, „wolny” w kolumnie osoby itd."""
    status_text = sanitize.fold(data.get("status_text") or "")
    for word, status in _STATUS_WORDS:
        if word in status_text:
            return status, None
    holder = data.get("holder_text")
    folded = sanitize.fold(holder or "")
    if holder and folded.startswith(_FREE_PREFIXES):
        return "na_stanie", None
    if holder:
        return "wydany", holder
    if status_text.startswith(_FREE_STATUS):
        return "na_stanie", None
    return "nieznany", None


def _name_keys(name: str) -> set[str]:
    folded = sanitize.fold(name)
    parts = folded.split()
    keys = {folded}
    if len(parts) == 2:
        keys.add(f"{parts[1]} {parts[0]}")
    return keys


async def _people_index(
    db: AsyncSession,
) -> tuple[dict[str, set[int]], dict[str, set[int]]]:
    users: dict[str, set[int]] = {}
    for user_id, name in (
        await db.execute(select(User.id, User.name).where(User.is_active.is_(True)))
    ).all():
        if name:
            users.setdefault(sanitize.fold(name), set()).add(user_id)
    contracts: dict[str, set[int]] = {}
    for contract_id, first, last in (
        await db.execute(
            select(Contract.id, Candidate.name, Candidate.lastname)
            .join(Candidate, Candidate.id == Contract.candidate_id)
            .where(Contract.status.in_(ISSUABLE_CONTRACT_STATUSES))
        )
    ).all():
        full = " ".join(p for p in (first, last) if p)
        if full:
            contracts.setdefault(sanitize.fold(full), set()).add(contract_id)
    return users, contracts


def _unique(index: dict[str, set[int]], text: str) -> Optional[int]:
    found: set[int] = set()
    for key in _name_keys(text):
        found |= index.get(key, set())
    return next(iter(found)) if len(found) == 1 else None


async def _existing_numbers(db: AsyncSession) -> dict[tuple[str, str], int]:
    out: dict[tuple[str, str], int] = {}
    columns = [getattr(EquipmentItem, f) for f in IDENTIFIER_FIELDS]
    for row in (await db.execute(select(EquipmentItem.id, *columns))).all():
        for field_name, value in zip(IDENTIFIER_FIELDS, row[1:]):
            if value:
                out[(field_name, value)] = row[0]
    return out


def _is_current_sheet(name: str) -> bool:
    """Arkusz z aktualną listą (np. „Aktualne”) — w plikach recepcji obok stoją jej kopie."""
    return "aktualn" in sanitize.fold(name)


async def create_batch(
    db: AsyncSession, *, filename: str, content: bytes, actor: User
) -> EquipmentImportBatch:
    # Parsowanie poza pętlą zdarzeń: duży arkusz nie zatrzymuje innych żądań.
    parsed = await asyncio.to_thread(parse_workbook, content)
    users, contracts = await _people_index(db)
    existing = await _existing_numbers(db)
    # Numer → (gdzie pierwszy raz, czy z aktualnej listy). Wiersze aktualnej listy
    # idą pierwsze: kopia stojąca w pliku PRZED „Aktualne” nie może jej wygrać
    # (próba na pliku kart 16.10: 66 aktualnych kart dostawało „konflikt”).
    seen: dict[tuple[str, str], tuple[str, bool]] = {}
    counts: dict[str, int] = {d: 0 for d in schema.IMPORT_DECISIONS}
    ordered = sorted(parsed.rows, key=lambda r: not _is_current_sheet(r.sheet))
    batch = EquipmentImportBatch(
        file_name=(filename or "plik.xlsx")[:255],
        file_sha256=hashlib.sha256(content).hexdigest(),
        created_by=actor.id,
    )
    db.add(batch)
    await db.flush()
    for row in ordered:
        data = row.data
        numbers = [(f, data[f]) for f in IDENTIFIER_FIELDS if data.get(f)]
        holder_user = holder_contract = None
        status, holder_text = holder_state(data)
        reasons = list(row.reasons)
        if not numbers:
            # Jedyny numer był za długi: egzemplarz istnieje, ktoś musi go poprawić.
            decision = "do_wyjasnienia" if row.needs_review else "pomin"
            if not row.needs_review:
                reasons.insert(0, "brak numeru egzemplarza")
        elif hit := next((key for key in numbers if key in existing), None):
            decision = "konflikt"
            reasons.insert(
                0,
                f"{identifiers.FIELD_LABELS[hit[0]]} jest już w rejestrze (pozycja #{existing[hit]})",
            )
        elif dup := next((key for key in numbers if key in seen), None):
            where, from_current = seen[dup]
            if from_current and not _is_current_sheet(row.sheet):
                decision = "pomin"
                reasons.insert(0, f"kopia — liczy się aktualna lista ({where})")
            else:
                decision = "konflikt"
                reasons.insert(
                    0, f"ten sam {identifiers.FIELD_LABELS[dup[0]]} co {where}"
                )
        else:
            for key in numbers:
                seen[key] = (
                    f"arkusz „{row.sheet}”, wiersz {row.row_number}",
                    _is_current_sheet(row.sheet),
                )
            if status == "wydany" and holder_text:
                holder_user = _unique(users, holder_text)
                if holder_user is None:
                    holder_contract = _unique(contracts, holder_text)
                if holder_user or holder_contract:
                    decision = "przyjmij"
                else:
                    decision = "do_wyjasnienia"
                    reasons.insert(0, "nie znaleziono tej osoby w NEXUSIE")
            elif status == "nieznany" and data["category"] == "laptop":
                decision = "do_wyjasnienia"
                reasons.insert(0, "laptop bez posiadacza i bez stanu")
            else:
                decision = "przyjmij"
        if row.needs_review:
            if decision == "przyjmij":
                decision = "do_wyjasnienia"
            reasons[0:0] = row.needs_review
        if holder_user or holder_contract:
            data["item_status"] = "wydany"
        elif status in ("serwis", "wycofany", "zgubiony"):
            data["item_status"] = status
        else:
            data["item_status"] = "na_stanie"
        counts[decision] += 1
        db.add(
            EquipmentImportRow(
                batch_id=batch.id,
                sheet=row.sheet,
                row_number=row.row_number,
                data=data,
                decision=decision,
                reason="; ".join(dict.fromkeys(reasons)) or None,
                holder_user_id=holder_user,
                holder_contract_id=holder_contract,
            )
        )
    batch.summary = {
        "sheets": [s.as_dict() for s in parsed.sheets],
        "removed_columns": parsed.removed_columns,
        "ignored_columns": parsed.ignored_columns,
        "removed_values": parsed.removed_values,
        "counts": counts,
    }
    await db.flush()
    audit.record(
        db,
        actor_user_id=actor.id,
        subject_user_id=None,
        entity_type="equipment_import",
        entity_id=batch.id,
        action="import_uploaded",
        details={"counts": counts},
    )
    return batch


def batch_payload(
    batch: EquipmentImportBatch, *, pending: int | None = None
) -> dict[str, Any]:
    return {
        "id": batch.id,
        "file_name": batch.file_name,
        "status": batch.status,
        "summary": batch.summary or {},
        "created_at": batch.created_at.isoformat() if batch.created_at else None,
        "applied_at": batch.applied_at.isoformat() if batch.applied_at else None,
        "pending": pending,
    }


def row_payload(
    row: EquipmentImportRow, holder_name: Optional[str] = None
) -> dict[str, Any]:
    return {
        "id": row.id,
        "sheet": row.sheet,
        "row_number": row.row_number,
        "data": row.data or {},
        "decision": row.decision,
        "reason": row.reason,
        "holder_user_id": row.holder_user_id,
        "holder_contract_id": row.holder_contract_id,
        "holder_name": holder_name,
        "item_id": row.item_id,
        "applied": row.applied_at is not None,
    }


async def _pending(db: AsyncSession, batch_id: int) -> int:
    return int(
        await db.scalar(
            select(func.count()).where(
                EquipmentImportRow.batch_id == batch_id,
                EquipmentImportRow.decision == "przyjmij",
                EquipmentImportRow.applied_at.is_(None),
            )
        )
        or 0
    )


async def list_batches(db: AsyncSession) -> list[dict[str, Any]]:
    batches = (
        await db.scalars(
            select(EquipmentImportBatch)
            .order_by(EquipmentImportBatch.id.desc())
            .limit(50)
        )
    ).all()
    return [batch_payload(b, pending=await _pending(db, b.id)) for b in batches]


async def get_batch(db: AsyncSession, batch_id: int) -> EquipmentImportBatch:
    batch = await db.get(EquipmentImportBatch, batch_id)
    if batch is None:
        raise EquipmentNotFound("Nie ma takiej partii importu.")
    return batch


async def batch_rows(
    db: AsyncSession, batch_id: int, *, decision: Optional[str] = None
) -> list[dict[str, Any]]:
    holder = func.coalesce(
        User.name,
        func.nullif(func.concat_ws(" ", Candidate.name, Candidate.lastname), ""),
    )
    query = (
        select(EquipmentImportRow, holder)
        .outerjoin(User, User.id == EquipmentImportRow.holder_user_id)
        .outerjoin(Contract, Contract.id == EquipmentImportRow.holder_contract_id)
        .outerjoin(Candidate, Candidate.id == Contract.candidate_id)
        .where(EquipmentImportRow.batch_id == batch_id)
        .order_by(EquipmentImportRow.sheet, EquipmentImportRow.row_number)
    )
    if decision:
        query = query.where(EquipmentImportRow.decision == decision)
    return [row_payload(row, name) for row, name in (await db.execute(query)).all()]


async def update_row(
    db: AsyncSession,
    batch_id: int,
    row_id: int,
    *,
    decision: str,
    holder_user_id: Optional[int],
    holder_contract_id: Optional[int],
    actor: User,
) -> EquipmentImportRow:
    row = await db.get(EquipmentImportRow, row_id, with_for_update=True)
    if row is None or row.batch_id != batch_id:
        raise EquipmentNotFound("Nie ma takiego wiersza.")
    if row.applied_at is not None:
        raise EquipmentConflict("Ten wiersz jest już w rejestrze.")
    if decision not in schema.IMPORT_DECISIONS:
        raise EquipmentError("Nieznana decyzja.")
    if holder_user_id and holder_contract_id:
        raise EquipmentError("Wybierz jedną osobę: konto albo kontrakt.")
    if holder_user_id and await db.get(User, holder_user_id) is None:
        raise EquipmentNotFound("Nie ma takiego konta.")
    if holder_contract_id and await db.get(Contract, holder_contract_id) is None:
        raise EquipmentNotFound("Nie ma takiego kontraktu.")
    if decision == "przyjmij" and not any(row.data.get(f) for f in IDENTIFIER_FIELDS):
        raise EquipmentError("Wiersz bez numeru egzemplarza nie trafi do rejestru.")
    row.decision = decision
    row.holder_user_id = holder_user_id
    row.holder_contract_id = holder_contract_id
    data = dict(row.data or {})
    if holder_user_id or holder_contract_id:
        data["item_status"] = "wydany"
    elif data.get("item_status") == "wydany":
        data["item_status"] = "na_stanie"
    row.data = data
    await db.flush()
    audit.record(
        db,
        actor_user_id=actor.id,
        subject_user_id=None,
        entity_type="equipment_import",
        entity_id=batch_id,
        action="import_row_decided",
        details={"row_id": row.id, "decision": decision},
    )
    return row


async def apply_batch(
    db: AsyncSession, batch_id: int, *, actor: User
) -> dict[str, int]:
    """Zapis wierszy „przyjmij”, które nie są jeszcze w rejestrze."""
    batch = await db.get(EquipmentImportBatch, batch_id, with_for_update=True)
    if batch is None:
        raise EquipmentNotFound("Nie ma takiej partii importu.")
    rows = (
        await db.scalars(
            select(EquipmentImportRow)
            .where(
                EquipmentImportRow.batch_id == batch.id,
                EquipmentImportRow.decision == "przyjmij",
                EquipmentImportRow.applied_at.is_(None),
            )
            .order_by(EquipmentImportRow.id)
        )
    ).all()
    created = conflicts = to_review = 0
    now = datetime.now(timezone.utc)
    for row in rows:
        data = row.data or {}
        numbers = {f: data.get(f) for f in IDENTIFIER_FIELDS}
        try:
            # Partia sprzed sprawdzania długości przy wgraniu: za długi numer
            # wywracał całe zatwierdzenie (500) bez wskazania wiersza.
            for name, value in numbers.items():
                identifiers.check_length(name, value)
        except identifiers.IdentifierTooLong as exc:
            row.decision = "do_wyjasnienia"
            row.reason = str(exc)
            to_review += 1
            continue
        problems = await registry.identifier_conflicts(db, numbers)
        if problems:
            row.decision = "konflikt"
            row.reason = problems[0]
            conflicts += 1
            continue
        model = await registry.get_or_create_model(
            db,
            category=data.get("category") or "inne",
            name=data.get("model") or "",
            price=Decimal(data["price"]) if data.get("price") else None,
            actor=actor,
        )
        holder = row.holder_user_id or row.holder_contract_id
        status = "wydany" if holder else data.get("item_status") or "na_stanie"
        if status == "wydany" and not holder:
            status = "na_stanie"
        note_parts = [
            data.get("notes"),
            data.get("project") and f"projekt: {data['project']}",
        ]
        item = EquipmentItem(
            model_id=model.id,
            status=status,
            ownership="nasz",
            purchase_date=date.fromisoformat(data["purchase_date"])
            if data.get("purchase_date")
            else None,
            notes=" · ".join(p for p in note_parts if p) or None,
            created_by=actor.id,
            updated_by=actor.id,
            **numbers,
        )
        db.add(item)
        await db.flush()
        if holder:
            db.add(
                EquipmentAssignment(
                    item_id=item.id,
                    user_id=row.holder_user_id,
                    contract_id=row.holder_contract_id
                    if not row.holder_user_id
                    else None,
                    issued_at=now,
                    issued_by=actor.id,
                )
            )
        row.item_id = item.id
        row.applied_at = now
        created += 1
    await db.flush()
    await registry.link_contract_equipment(db)
    batch.status = "zatwierdzona"
    batch.applied_at = now
    batch.applied_by = actor.id
    await db.flush()
    audit.record(
        db,
        actor_user_id=actor.id,
        subject_user_id=None,
        entity_type="equipment_import",
        entity_id=batch.id,
        action="import_applied",
        details={"created": created, "conflicts": conflicts, "to_review": to_review},
    )
    # Wiersze cofnięte „do wyjaśnienia” widać w partii; kształt odpowiedzi bez zmian.
    return {"created": created, "conflicts": conflicts}


_SHEET_TITLES = (
    ("przyjmij", "Do przyjęcia"),
    ("do_wyjasnienia", "Do wyjaśnienia"),
    ("konflikt", "Konflikty"),
    ("pomin", "Pominięte"),
)
_COLUMNS = (
    ("Arkusz", lambda r: r["sheet"]),
    ("Wiersz", lambda r: r["row_number"]),
    (
        "Kategoria",
        lambda r: identifiers.CATEGORY_LABELS.get(r["data"].get("category"), ""),
    ),
    ("Model", lambda r: r["data"].get("model")),
    ("Nr ewid.", lambda r: r["data"].get("inventory_number")),
    ("Nr seryjny", lambda r: r["data"].get("serial_number")),
    ("IMEI", lambda r: r["data"].get("imei")),
    ("ICCID", lambda r: r["data"].get("iccid")),
    ("Telefon", lambda r: r["data"].get("phone_number")),
    ("Nr karty", lambda r: r["data"].get("access_card_number")),
    ("Osoba w pliku", lambda r: r["data"].get("holder_text")),
    ("Osoba w NEXUSIE", lambda r: r["holder_name"]),
    (
        "Stan po imporcie",
        lambda r: identifiers.STATUS_LABELS.get(r["data"].get("item_status"), ""),
    ),
    ("Uwagi", lambda r: r["data"].get("notes")),
    ("Powód", lambda r: r["reason"]),
)


async def control_sheet(db: AsyncSession, batch_id: int) -> bytes:
    """Arkusz kontrolny XLSX: podsumowanie i wiersze według decyzji."""
    batch = await get_batch(db, batch_id)
    rows = await batch_rows(db, batch_id)
    workbook = openpyxl.Workbook()
    summary = workbook.active
    summary.title = "Podsumowanie"
    summary.append(safe_row(["Plik", batch.file_name]))
    summary.append(safe_row(["Partia", batch.id]))
    summary.append(safe_row([]))
    counts = (batch.summary or {}).get("counts", {})
    for decision, title in _SHEET_TITLES:
        summary.append(safe_row([title, counts.get(decision, 0)]))
    summary.append(safe_row([]))
    summary.append(safe_row(["Kolumny usunięte (dane, których NEXUS nie przechowuje)"]))
    for entry in (batch.summary or {}).get("removed_columns", []):
        summary.append(safe_row([entry["sheet"], entry["column"], entry["reason"]]))
    summary.append(safe_row([]))
    summary.append(safe_row(["Wartości usunięte z wierszy"]))
    for reason, count in (batch.summary or {}).get("removed_values", {}).items():
        summary.append(safe_row([reason, count]))
    summary.append(safe_row([]))
    summary.append(safe_row(["Arkusze"]))
    for sheet in (batch.summary or {}).get("sheets", []):
        summary.append(
            safe_row([sheet["name"], sheet["rows"], sheet.get("skipped") or ""])
        )
    for decision, title in _SHEET_TITLES:
        sheet = workbook.create_sheet(title)
        sheet.append(safe_row([name for name, _ in _COLUMNS]))
        for row in rows:
            if row["decision"] == decision:
                sheet.append(safe_row([getter(row) for _, getter in _COLUMNS]))
        sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
