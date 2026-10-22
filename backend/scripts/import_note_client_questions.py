"""Pytania klienta z notatek po rozmowie u klienta → debrief i bank pytań (20.10.2026).

Do 20.10.2026 rozmowa u klienta bez terminu w kalendarzu NEXUSA nie miała
okna na pytania (88 z 99 par w 30 dniach), więc rekruterzy wpisywali je
w notatki („klient pytał o…”). Ten skrypt przenosi je tam, gdzie powinny być:
odbyta rozmowa pary + debrief z pytaniami (jak ``PUT …/pairs/{c}/{j}/debrief``)
i bank pytań klienta (źródło ``client_debrief``, pin w rekrutacji) — dzięki
temu trafiają do prepów kolejnych kandydatów, a para przechodzi bramkę przed
„Umową”. Notatki zostają bez zmian.

Przebieg (w kontenerze backendu)::

    docker exec <backend> python -m scripts.import_note_client_questions \\
        --plan /tmp/note_q_plan.json --review /tmp/note_q_review.xlsx
    # człowiek przegląda arkusz, dopiero potem (opcjonalnie bez części notatek):
    docker exec <backend> python -m scripts.import_note_client_questions \\
        --apply /tmp/note_q_plan.json [--skip-notes 101,102:7]

* ``--plan`` czyta notatki i woła model (ten sam prompt i ta sama kontrola
  cytatu co archiwum z Excela, ``import_legacy_interview_questions``); niczego
  nie zapisuje.
* ``--apply`` zapisuje dokładnie przejrzany plan, bez modelu; para, która ma
  już debrief z pytaniami albo z „klient nie pytał”, jest pomijana.
  ``--skip-notes`` przyjmuje id notatki (cała) albo ``id_notatki:id_rekrutacji``.
* Notatka bez rekrutacji trafia do pary tylko, gdy kandydat ma w oknie
  dokładnie jedną; inaczej jest w „Odrzuconych” jako „niejednoznaczna
  rekrutacja” (audyt 21.10.2026).
* ``--rollback`` cofa wszystkie przebiegi z paragonu: debriefy i rozmowy
  założone przez import, dopisane pytania i przypięcia — bez pracy ludzi po
  imporcie (pytanie przypięte gdzie indziej, rozmowa z debriefem klienta
  zostają; audyt 22.10.2026).

Paragon w ``app_settings['note_client_questions_import']``: same liczby i ID.
"""

from __future__ import annotations

import argparse
import asyncio
import json
import logging
import re
import sys
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Optional

BACKEND_ROOT = Path(__file__).resolve().parents[1]
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from scripts.import_legacy_interview_questions import (  # noqa: E402
    Entry,
    _model_payload,
    question_hash,
    validate_model_output,
)

logger = logging.getLogger("note_client_questions")

RECEIPT_KEY = "note_client_questions_import"
PLAN_VERSION = 1
DEFAULT_DAYS = 90
# Notatka po rozmowie: od dnia przed przesunięciem na „Rozmowę z klientem”
# (karta bywa przesuwana po fakcie) do trzech tygodni po nim.
NOTE_BEFORE = timedelta(days=1)
NOTE_AFTER = timedelta(days=21)
MAX_NOTE_CHARS = 8_000
PIN_ORDER_BASE = 10_000.0
IMPORT_CONCERNS = "Pytania przeniesione z notatki z {day} (import 20.10.2026)."

# Lustro ``frontend/src/lib/client-question-notes.ts`` (podpowiedź pod notatką).
_ASKED = re.compile(r"pyta\w*", re.IGNORECASE)
_CONTEXT = re.compile(
    r"(?<!\w)(klient\w*|rozmow\w*|interview\w*|spotkani\w*|techniczn\w*|manager\w*|HM)(?!\w)",
    re.IGNORECASE,
)


def looks_like_client_questions(text: str) -> bool:
    value = (text or "").strip()
    if not value:
        return False
    if _ASKED.search(value) and _CONTEXT.search(value):
        return True
    return value.count("?") >= 2 and bool(_CONTEXT.search(value))


@dataclass(frozen=True)
class NoteRow:
    note_id: int
    candidate_id: int
    job_id: int
    job_title: str
    client_id: Optional[int]
    author_id: int
    created_at: datetime
    content: str


async def candidate_notes(db: Any, *, days: int, now: datetime) -> list[NoteRow]:
    """Notatki ludzi przy parach, które weszły w „Rozmowę z klientem”
    w ostatnich ``days`` dniach, napisane w oknie po przesunięciu."""
    from sqlalchemy import and_, func, or_, select

    from app.models.job import Job
    from app.models.note import Note
    from app.models.recruitment_pipeline import CandidateStage, PipelineStage

    entered = (
        select(
            CandidateStage.candidate_id,
            CandidateStage.job_id,
            func.min(CandidateStage.moved_at).label("moved_at"),
        )
        .where(
            CandidateStage.stage == PipelineStage.client_interview,
            CandidateStage.moved_at >= now - timedelta(days=days),
        )
        .group_by(CandidateStage.candidate_id, CandidateStage.job_id)
        .subquery()
    )
    rows = (
        await db.execute(
            select(
                Note.id,
                entered.c.candidate_id,
                entered.c.job_id,
                Job.title,
                Job.client_id,
                Note.author_id,
                Note.created_at,
                Note.content,
            )
            .join(entered, entered.c.candidate_id == Note.candidate_id)
            .join(Job, Job.id == entered.c.job_id)
            .where(
                Note.author_id.isnot(None),
                Note.source_deleted_at.is_(None),
                or_(Note.job_id.is_(None), Note.job_id == entered.c.job_id),
                and_(
                    Note.created_at >= entered.c.moved_at - NOTE_BEFORE,
                    Note.created_at <= entered.c.moved_at + NOTE_AFTER,
                ),
                Note.content.op("~*")("pyta"),
            )
            .order_by(Note.created_at)
        )
    ).all()
    return [
        NoteRow(
            note_id=r[0],
            candidate_id=r[1],
            job_id=r[2],
            job_title=r[3] or "",
            client_id=r[4],
            author_id=r[5],
            created_at=r[6] if r[6].tzinfo else r[6].replace(tzinfo=timezone.utc),
            content=(r[7] or "")[:MAX_NOTE_CHARS],
        )
        for r in rows
        if looks_like_client_questions(r[7] or "")
    ]


AMBIGUOUS_JOB = "niejednoznaczna rekrutacja"


def _pairs_per_note(
    items: list[tuple[int, int, int]],
) -> dict[int, set[tuple[int, int]]]:
    """(notatka, kandydat, rekrutacja) → pary, przy których stoi notatka."""
    out: dict[int, set[tuple[int, int]]] = {}
    for note_id, candidate_id, job_id in items:
        out.setdefault(note_id, set()).add((candidate_id, job_id))
    return out


def split_ambiguous_notes(
    notes: list[NoteRow],
) -> tuple[list[NoteRow], list[dict[str, Any]]]:
    """Notatka bez rekrutacji trafia do pary tylko, gdy pasuje dokładnie jedna.

    Notatka z ``job_id`` łączy się z jedną parą; bez niego — z każdą parą
    kandydata, która weszła w „Rozmowę z klientem” w oknie notatki. Wtedy nie
    wiadomo, o którą rozmowę chodzi, a te same pytania lądowałyby w debriefach
    i bankach kilku klientów (audyt 21.10.2026) — notatka idzie do
    „Odrzuconych” z listą rekrutacji.
    """
    pairs = _pairs_per_note([(n.note_id, n.candidate_id, n.job_id) for n in notes])
    kept = [n for n in notes if len(pairs[n.note_id]) == 1]
    rejected = [
        {
            "note_id": note_id,
            "reason": AMBIGUOUS_JOB,
            "question": "rekrutacje: "
            + ", ".join(str(job) for job in sorted(j for _c, j in found)),
        }
        for note_id, found in sorted(pairs.items())
        if len(found) > 1
    ]
    return kept, rejected


def business_day(value: datetime) -> date:
    from zoneinfo import ZoneInfo

    from app.core.config import settings

    return value.astimezone(ZoneInfo(settings.BUSINESS_TZ)).date()


async def build_plan(notes: list[NoteRow], *, concurrency: int = 4) -> dict[str, Any]:
    """Model + kontrola cytatu dla każdej notatki. Bez zapisu."""
    sem = asyncio.Semaphore(concurrency)
    questions: list[dict[str, Any]] = []
    note_count = len({n.note_id for n in notes})
    notes, rejected = split_ambiguous_notes(notes)
    failed: list[dict[str, Any]] = []

    async def one(note: NoteRow) -> None:
        entry = Entry(
            sheet="notatka",
            row=note.note_id,
            client_ids=(note.client_id,) if note.client_id else (),
            role=note.job_title,
            refs=[],
            text=note.content,
        )
        async with sem:
            try:
                payload = await _model_payload(entry)
            except Exception as exc:  # noqa: BLE001 — jeden wpis nie wywraca planu
                failed.append({"note_id": note.note_id, "error": type(exc).__name__})
                return
        kept, bad = validate_model_output(entry, payload)
        for item in kept:
            questions.append(
                {
                    "note_id": note.note_id,
                    "candidate_id": note.candidate_id,
                    "job_id": note.job_id,
                    "job_title": note.job_title,
                    "client_id": note.client_id,
                    "author_id": note.author_id,
                    "note_day": business_day(note.created_at).isoformat(),
                    "question": item["question"],
                    "quote": item["quote"],
                }
            )
        for r in bad:
            rejected.append(
                {"note_id": note.note_id, "reason": r.reason, "question": r.question}
            )

    await asyncio.gather(*(one(n) for n in notes))
    questions.sort(key=lambda q: (q["candidate_id"], q["job_id"], q["note_id"]))
    return {
        "version": PLAN_VERSION,
        "built_at": datetime.now(timezone.utc).isoformat(),
        "notes": note_count,
        "questions": questions,
        "rejected": rejected,
        "failed": failed,
    }


def write_review(plan: dict[str, Any], path: Path) -> None:
    """Arkusz do przeglądu: pytania (z numerem notatki do ``--skip-notes``),
    odrzucone i błędy."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pytania"
    ws.append(
        [
            "Notatka (id)",
            "Kandydat (id)",
            "Rekrutacja (id)",
            "Rekrutacja",
            "Dzień",
            "Pytanie",
            "Cytat z notatki",
        ]
    )
    for q in plan.get("questions") or []:
        ws.append(
            [
                q["note_id"],
                q["candidate_id"],
                q["job_id"],
                q["job_title"],
                q["note_day"],
                q["question"],
                q["quote"],
            ]
        )
    ws = wb.create_sheet("Odrzucone")
    ws.append(["Notatka (id)", "Powód", "Pytanie"])
    for r in plan.get("rejected") or []:
        ws.append([r["note_id"], r["reason"], r["question"]])
    ws = wb.create_sheet("Błędy")
    ws.append(["Notatka (id)", "Błąd"])
    for f in plan.get("failed") or []:
        ws.append([f["note_id"], f["error"]])
    wb.save(path)


def parse_skip_notes(
    raw: Optional[str],
) -> tuple[frozenset[int], frozenset[tuple[int, int]]]:
    """``--skip-notes 101,102:7`` → notatki pominięte w całości i pary
    (notatka, rekrutacja) pominięte tylko w tej rekrutacji."""
    notes: set[int] = set()
    note_jobs: set[tuple[int, int]] = set()
    for token in (raw or "").split(","):
        token = token.strip()
        if not token:
            continue
        note, sep, job = token.partition(":")
        try:
            if sep:
                note_jobs.add((int(note), int(job)))
            else:
                notes.add(int(note))
        except ValueError as exc:
            raise ValueError(
                f"--skip-notes: „{token}” — podaj id notatki albo id_notatki:id_rekrutacji"
            ) from exc
    return frozenset(notes), frozenset(note_jobs)


def group_by_pair(
    questions: list[dict[str, Any]],
    *,
    skip_notes: frozenset[int] = frozenset(),
    skip_note_jobs: frozenset[tuple[int, int]] = frozenset(),
) -> dict[tuple[int, int], list[dict[str, Any]]]:
    """Pytania planu po parach (kandydat, rekrutacja), bez pominiętych notatek
    i bez powtórzeń w parze (ten sam klucz co debrief).

    Notatka stojąca przy kilku parach (plan sprzed 21.10.2026, notatka bez
    rekrutacji) nie trafia do żadnej — ta sama reguła co ``split_ambiguous_notes``.
    """
    spread = _pairs_per_note(
        [(q["note_id"], q["candidate_id"], q["job_id"]) for q in questions]
    )
    pairs: dict[tuple[int, int], list[dict[str, Any]]] = {}
    seen: set[tuple[int, int, str]] = set()
    for q in questions:
        if q["note_id"] in skip_notes or (q["note_id"], q["job_id"]) in skip_note_jobs:
            continue
        if len(spread[q["note_id"]]) > 1:
            continue
        key = (q["candidate_id"], q["job_id"])
        dedup = (*key, question_hash(q["question"]))
        if dedup in seen:
            continue
        seen.add(dedup)
        pairs.setdefault(key, []).append(q)
    return pairs


async def _bank(
    db: Any, *, texts: list[str], client_id: Optional[int], job_id: int, user_id: int
) -> tuple[list[int], list[int]]:
    """Bank pytań klienta + pin (jak ``_save_client_questions``) — zwraca id
    NOWYCH pytań i NOWYCH przypięć (do cofnięcia)."""
    from sqlalchemy import select

    from app.models.interview_question import (
        InterviewQuestion,
        InterviewQuestionSource,
        JobQuestion,
        JobQuestionAddedBySource,
    )

    inserted: list[int] = []
    pinned: list[int] = []
    if client_id is None:
        return inserted, pinned
    for n, text in enumerate(texts):
        norm = question_hash(text)
        question = await db.scalar(
            select(InterviewQuestion).where(
                InterviewQuestion.client_id == client_id,
                InterviewQuestion.normalized_text_hash == norm,
            )
        )
        if question is None:
            question = InterviewQuestion(
                text=text,
                client_id=client_id,
                source=InterviewQuestionSource.client_debrief,
                normalized_text_hash=norm,
                skill_tags=[],
                created_by=user_id,
            )
            db.add(question)
            await db.flush()
            inserted.append(question.id)
        exists = await db.scalar(
            select(JobQuestion.id).where(
                JobQuestion.job_id == job_id, JobQuestion.question_id == question.id
            )
        )
        if exists is None:
            link = JobQuestion(
                job_id=job_id,
                question_id=question.id,
                is_pinned=True,
                added_by_source=JobQuestionAddedBySource.manual,
                added_by_user_id=user_id,
                order_index=PIN_ORDER_BASE + n,
            )
            db.add(link)
            await db.flush()
            pinned.append(link.id)
    return inserted, pinned


async def apply_plan(
    db: Any,
    plan: dict[str, Any],
    *,
    skip_notes: frozenset[int] = frozenset(),
    skip_note_jobs: frozenset[tuple[int, int]] = frozenset(),
    now: Optional[datetime] = None,
) -> dict[str, Any]:
    """Zapisuje plan. Para z debriefem, który już zamyka rundę, jest pomijana —
    powtórka nic nie dubluje. Commit robi wołający."""
    from sqlalchemy import select

    from app.models.app_setting import AppSetting
    from app.models.calendar_event import CalendarEvent, EventType
    from app.models.interview_feedback import FeedbackSource, InterviewFeedback
    from app.services import client_interview_entry
    from app.services.debrief_gate import missing_debrief

    if plan.get("version") != PLAN_VERSION:
        raise ValueError("nieznana wersja planu")
    now = now or datetime.now(timezone.utc)
    receipt: dict[str, Any] = {
        "applied_at": now.isoformat(),
        "skipped_notes": sorted(skip_notes),
        "skipped_note_jobs": [f"{n}:{j}" for n, j in sorted(skip_note_jobs)],
        "pairs": 0,
        "pairs_already_closed": 0,
        "event_ids": [],
        "feedback_ids": [],
        "updated_feedbacks": [],
        "question_ids": [],
        "pin_ids": [],
    }
    for (candidate_id, job_id), items in group_by_pair(
        plan.get("questions") or [],
        skip_notes=skip_notes,
        skip_note_jobs=skip_note_jobs,
    ).items():
        missing = await missing_debrief(
            db, candidate_id=candidate_id, job_id=job_id, now=now
        )
        events_exist = await db.scalar(
            select(CalendarEvent.id)
            .where(
                CalendarEvent.candidate_id == candidate_id,
                CalendarEvent.job_id == job_id,
                CalendarEvent.event_type == EventType.client_interview,
            )
            .limit(1)
        )
        if missing is None and events_exist is not None:
            receipt["pairs_already_closed"] += 1
            continue
        first = min(items, key=lambda q: q["note_day"])
        author_id = first["author_id"]
        texts = [q["question"] for q in items]
        event = None
        if missing is not None and missing.get("event_id") is not None:
            if missing.get("interview_pending"):
                # Rozmowa jeszcze przed nami — notatka nie jest debriefem po niej.
                continue
            event = await db.get(CalendarEvent, missing["event_id"])
        if event is None:
            day = date.fromisoformat(first["note_day"])
            event = await client_interview_entry.held_interview_on(
                db, candidate_id=candidate_id, job_id=job_id, day=day, now=now
            )
        if event is None:
            start = client_interview_entry.held_interview_start(
                date.fromisoformat(first["note_day"]), None, now=now
            )
            event = await client_interview_entry.create_interview(
                db,
                candidate_id=candidate_id,
                job_id=job_id,
                start=start,
                duration_minutes=client_interview_entry.DEFAULT_DURATION_MINUTES,
                created_by=author_id,
                held=True,
                now=now,
            )
            receipt["event_ids"].append(event.id)
        fb = await db.scalar(
            select(InterviewFeedback).where(
                InterviewFeedback.calendar_event_id == event.id,
                InterviewFeedback.feedback_source == FeedbackSource.candidate_side,
            )
        )
        if fb is None:
            fb = InterviewFeedback(
                calendar_event_id=event.id,
                candidate_id=candidate_id,
                job_id=job_id,
                author_id=author_id,
                feedback_source=FeedbackSource.candidate_side,
                client_questions="\n".join(texts),
                no_client_questions=False,
                concerns=IMPORT_CONCERNS.format(
                    day=date.fromisoformat(first["note_day"]).strftime("%d.%m.%Y")
                ),
            )
            db.add(fb)
            await db.flush()
            receipt["feedback_ids"].append(fb.id)
        else:
            receipt["updated_feedbacks"].append(
                {
                    "id": fb.id,
                    "client_questions": fb.client_questions,
                    "no_client_questions": fb.no_client_questions,
                }
            )
            current = [
                line.strip()
                for line in (fb.client_questions or "").splitlines()
                if line.strip()
            ]
            known = {question_hash(t) for t in current}
            fb.client_questions = "\n".join(
                current + [t for t in texts if question_hash(t) not in known]
            )
            fb.no_client_questions = False
            # Cofnięcie przywraca debrief tylko wtedy, gdy nikt go potem nie zmienił.
            receipt["updated_feedbacks"][-1]["written_client_questions"] = (
                fb.client_questions
            )
            await db.flush()
        inserted, pins = await _bank(
            db,
            texts=texts,
            client_id=first["client_id"] or event.client_id,
            job_id=job_id,
            user_id=author_id,
        )
        receipt["question_ids"].extend(inserted)
        receipt["pin_ids"].extend(pins)
        receipt["pairs"] += 1
    setting = await db.get(AppSetting, RECEIPT_KEY)
    if setting is None:
        db.add(AppSetting(key=RECEIPT_KEY, value={"runs": [receipt]}))
    else:
        value = dict(setting.value or {})
        value["runs"] = [*(value.get("runs") or []), receipt]
        setting.value = value
    return receipt


async def _still_used(db: Any, ids: list[int], columns: tuple[Any, ...]) -> set[int]:
    """Id, do których po cofnięciu części importu wciąż coś się odwołuje —
    praca ludzi po imporcie (przypięcie w innej rekrutacji, ocena, debrief
    klienta, prep). Te wiersze zostają (audyt 22.10.2026)."""
    from sqlalchemy import select

    used: set[int] = set()
    for column in columns:
        used.update(
            (await db.scalars(select(column).where(column.in_(ids)).distinct())).all()
        )
    return used


async def rollback(db: Any) -> dict[str, int]:
    """Cofa wszystkie przebiegi z paragonu (od najnowszego). Commit robi wołający.

    Nie zabiera pracy ludzi po imporcie: pytanie, które ktoś potem przypiął
    albo ocenił, i rozmowa, do której ktoś coś dopisał, zostają
    (``kept_questions`` / ``kept_events``); uzupełniony debrief wraca do stanu
    sprzed importu tylko, gdy nikt go potem nie zmienił (``kept_feedbacks``).
    """
    from sqlalchemy import delete

    from app.models.app_setting import AppSetting
    from app.models.calendar_event import CalendarEvent
    from app.models.candidate_contact import CandidateContactOpportunity
    from app.models.client_interview_slot_request import ClientInterviewSlotRequest
    from app.models.followup_meeting import FollowupMeeting
    from app.models.interview_feedback import InterviewFeedback
    from app.models.interview_question import (
        InterviewQuestion,
        InterviewQuestionRating,
        JobQuestion,
    )
    from app.models.prep_meeting import Prep2Skip, PrepMeeting

    setting = await db.get(AppSetting, RECEIPT_KEY)
    runs = list((setting.value or {}).get("runs") or []) if setting else []
    counts = {
        "feedbacks": 0,
        "events": 0,
        "questions": 0,
        "pins": 0,
        "restored": 0,
        "kept_feedbacks": 0,
        "kept_questions": 0,
        "kept_events": 0,
    }
    for run in reversed(runs):
        for saved in run.get("updated_feedbacks") or []:
            fb = await db.get(InterviewFeedback, saved["id"])
            if fb is None:
                continue
            # Paragon sprzed 22.10.2026 nie zna wpisanej treści — nie zgadujemy.
            if (
                "written_client_questions" not in saved
                or fb.client_questions != saved["written_client_questions"]
                or fb.no_client_questions
            ):
                counts["kept_feedbacks"] += 1
                continue
            fb.client_questions = saved["client_questions"]
            fb.no_client_questions = saved["no_client_questions"]
            counts["restored"] += 1
        if run.get("pin_ids"):
            counts["pins"] += int(
                (
                    await db.execute(
                        delete(JobQuestion).where(JobQuestion.id.in_(run["pin_ids"]))
                    )
                ).rowcount
                or 0
            )
        if run.get("question_ids"):
            used = await _still_used(
                db,
                run["question_ids"],
                (JobQuestion.question_id, InterviewQuestionRating.question_id),
            )
            counts["kept_questions"] += len(used)
            doomed = [qid for qid in run["question_ids"] if qid not in used]
            if doomed:
                counts["questions"] += int(
                    (
                        await db.execute(
                            delete(InterviewQuestion).where(
                                InterviewQuestion.id.in_(doomed)
                            )
                        )
                    ).rowcount
                    or 0
                )
        if run.get("feedback_ids"):
            counts["feedbacks"] += int(
                (
                    await db.execute(
                        delete(InterviewFeedback).where(
                            InterviewFeedback.id.in_(run["feedback_ids"])
                        )
                    )
                ).rowcount
                or 0
            )
        if run.get("event_ids"):
            used = await _still_used(
                db,
                run["event_ids"],
                (
                    InterviewFeedback.calendar_event_id,
                    PrepMeeting.calendar_event_id,
                    Prep2Skip.interview_event_id,
                    FollowupMeeting.calendar_event_id,
                    CandidateContactOpportunity.meeting_event_id,
                    ClientInterviewSlotRequest.event_id,
                ),
            )
            counts["kept_events"] += len(used)
            doomed = [eid for eid in run["event_ids"] if eid not in used]
            if doomed:
                counts["events"] += int(
                    (
                        await db.execute(
                            delete(CalendarEvent).where(CalendarEvent.id.in_(doomed))
                        )
                    ).rowcount
                    or 0
                )
    if setting is not None:
        setting.value = {
            "runs": [],
            "rolled_back": [
                *((setting.value or {}).get("rolled_back") or []),
                {"at": datetime.now(timezone.utc).isoformat(), **counts},
            ],
        }
    return counts


async def _main(args: argparse.Namespace) -> int:
    from app.core.database import AsyncSessionLocal

    if args.rollback:
        async with AsyncSessionLocal() as db:
            counts = await rollback(db)
            await db.commit()
        print(f"KONIEC rollback: {counts}")
        return 0
    if args.apply:
        plan = json.loads(Path(args.apply).read_text(encoding="utf-8"))
        skip, skip_jobs = parse_skip_notes(args.skip_notes)
        async with AsyncSessionLocal() as db:
            receipt = await apply_plan(
                db, plan, skip_notes=skip, skip_note_jobs=skip_jobs
            )
            await db.commit()
        summary = {
            "pary": receipt["pairs"],
            "pary_juz_zamkniete": receipt["pairs_already_closed"],
            "rozmowy": len(receipt["event_ids"]),
            "debriefy_nowe": len(receipt["feedback_ids"]),
            "debriefy_uzupelnione": len(receipt["updated_feedbacks"]),
            "pytania_nowe": len(receipt["question_ids"]),
            "przypiecia": len(receipt["pin_ids"]),
        }
        print(f"KONIEC apply: {summary}")
        return 0
    async with AsyncSessionLocal() as db:
        notes = await candidate_notes(
            db, days=args.days, now=datetime.now(timezone.utc)
        )
    if args.limit:
        notes = notes[: args.limit]
    plan = await build_plan(notes, concurrency=args.concurrency)
    Path(args.plan).write_text(
        json.dumps(plan, ensure_ascii=False, indent=1), encoding="utf-8"
    )
    if args.review:
        write_review(plan, Path(args.review))
    pairs = {(q["candidate_id"], q["job_id"]) for q in plan["questions"]}
    print(
        f"KONIEC plan: notatki={plan['notes']} pary={len(pairs)} "
        f"pytania={len(plan['questions'])} odrzucone={len(plan['rejected'])} "
        f"błędy={len(plan['failed'])}"
    )
    return 0


def main(argv: Optional[list[str]] = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--plan", help="Notatki → plan JSON (bez zapisu)")
    mode.add_argument("--apply", help="Zapisz przejrzany plan JSON")
    mode.add_argument("--rollback", action="store_true", help="Cofnij import")
    parser.add_argument("--review", help="Arkusz XLSX do przeglądu (z --plan)")
    parser.add_argument(
        "--skip-notes",
        help="Id notatek do pominięcia, np. 101,102:7 (102 tylko w rekrutacji 7; z --apply)",
    )
    parser.add_argument("--days", type=int, default=DEFAULT_DAYS)
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--concurrency", type=int, default=4)
    args = parser.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    return asyncio.run(_main(args))


if __name__ == "__main__":
    raise SystemExit(main())
