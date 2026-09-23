"""Buduje formularz Word „Profil Championa” v5.0 z v4.0 (09.2026).

v5 = v4 + dwie sekcje profilu:

* **4. Doświadczenie poza stackiem** — Dziedzina / Certyfikaty / Regulacje
  (po sekcji 3, jak w edytorze),
* **8. Wiedza z rozmów** — „Od klienta” + wiersz „Insight od naszego
  konsultanta” przeniesiony tu z sekcji „O kliencie”.

Kolejne sekcje są przenumerowane (4→5, 5→6, 6→7). Parser formularza
(`champion_document.table_profile`) rozpoznaje pola po ETYKIETACH wierszy
i sekcję projektu po nazwie nagłówka, więc v4 nadal się wczytuje.

Nowe tabele są kopią istniejących (styl, obramowania, szerokości kolumn),
a nie budowane od zera — plik ma wyglądać jak ten sam wzór.

Uruchomienie (z katalogu backend):
    python -m scripts.build_champion_template_v5
"""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

ASSETS = Path(__file__).resolve().parents[1] / "app" / "assets" / "champion"
SOURCE = ASSETS / "Profil_Championa_v4.0.docx"
TARGET = ASSETS / "Profil_Championa_v5.0.docx"

EXPERIENCE_ROWS = (
    "Dziedzina — w czym kandydat pracował (np. płatności kartowe); "
    "dopisz „min. N lat”, a przy mile widzianej „(mile)”:",
    "Certyfikaty — np. ISTQB Foundation; mile widziane oznacz „(mile)”:",
    "Regulacje / standardy — np. PSD2, PCI DSS, RODO:",
)
EXPERIENCE_HINT = (
    "Dziedzina to obszar biznesowy, nie technologia. Wpisy rozdzielaj "
    "przecinkami. Kandydat bez śladu w CV nie znika z wyników — rekruter "
    "dopyta w screeningu."
)
INSIGHTS_HINT = (
    "Smaczki z rozmów: czego klient naprawdę szuka, za co odrzucał, kto "
    "decyduje; od konsultanta — jak wygląda praca na co dzień. Widzi to "
    "wyłącznie zespół rekrutacji."
)
CLIENT_HINT = (
    "Wpisuj ustalenia dotyczące tej roli. Pytania historyczne opieraj na "
    "przebytych rozmowach z klientem. Nieznane informacje pozostaw puste. "
    "Podawaj wyłącznie potwierdzone atuty i warunki oferty."
)
FROM_CLIENT_LABEL = "Od klienta — czego naprawdę szuka, za co odrzucał, kto decyduje:"


def _set_cell_text(cell, text: str) -> None:
    """Tekst komórki z zachowaniem formatu pierwszego runu."""
    first = cell.paragraphs[0]
    runs = first.runs
    if runs:
        runs[0].text = text
        for run in runs[1:]:
            run._r.getparent().remove(run._r)
    else:
        first.add_run(text)
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def _set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    runs = paragraph.runs
    if runs:
        runs[0].text = text
        for run in runs[1:]:
            run._r.getparent().remove(run._r)
    else:
        paragraph.add_run(text)


def _find_paragraph(doc, startswith: str) -> Paragraph:
    for block in doc.iter_inner_content():
        if isinstance(block, Paragraph) and block.text.strip().upper().startswith(
            startswith
        ):
            return block
    raise LookupError(startswith)


def _table_after(paragraph: Paragraph) -> Table:
    node = paragraph._p.getnext()
    while node is not None and not node.tag.endswith("}tbl"):
        node = node.getnext()
    if node is None:
        raise LookupError("brak tabeli po nagłówku")
    return Table(node, paragraph._parent)


def _clone_table(template: Table, labels: tuple[str, ...]) -> Table:
    element = deepcopy(template._tbl)
    table = Table(element, template._parent)
    rows = list(table.rows)
    for extra in rows[len(labels) :]:
        element.remove(extra._tr)
    while len(table.rows) < len(labels):
        element.append(deepcopy(table.rows[-1]._tr))
    for row, label in zip(table.rows, labels):
        _set_cell_text(row.cells[0], label)
        for cell in row.cells[1:]:
            _set_cell_text(cell, "")
    return table


def build() -> Path:
    doc = Document(SOURCE)

    # Przenumerowanie od końca, żeby nie trafić dwa razy w ten sam nagłówek.
    for old, new in (
        ("6. O KLIENCIE", "7. O KLIENCIE"),
        ("5. SCREENING", "6. SCREENING"),
        ("4. O PROJEKCIE", "5. O PROJEKCIE"),
    ):
        _set_paragraph_text(_find_paragraph(doc, old), new)

    # ── 4. Doświadczenie poza stackiem — po tabeli stacku ─────────────────
    stack_heading = _find_paragraph(doc, "3. STACK")
    stack_table = _table_after(stack_heading)
    stack_hint = Paragraph(stack_heading._p.getnext(), stack_heading._parent)

    heading = deepcopy(stack_heading._p)
    hint = deepcopy(stack_hint._p)
    table = _clone_table(stack_table, EXPERIENCE_ROWS)
    anchor = stack_table._tbl.getnext()  # akapit „Gdy wystarczy jedna z alternatyw…”
    anchor.addnext(heading)
    heading.addnext(hint)
    hint.addnext(table._tbl)
    _set_paragraph_text(
        Paragraph(heading, stack_heading._parent), "4. DOŚWIADCZENIE POZA STACKIEM"
    )
    _set_paragraph_text(Paragraph(hint, stack_heading._parent), EXPERIENCE_HINT)

    # ── 8. Wiedza z rozmów — insight konsultanta przeniesiony z „O kliencie”
    client_heading = _find_paragraph(doc, "7. O KLIENCIE")
    client_table = _table_after(client_heading)
    insight_row = next(
        row
        for row in client_table.rows
        if row.cells[0].text.strip().lower().startswith("insight od")
    )
    insight_label = insight_row.cells[0].text.strip().split("\n")[0]
    insights_table = _clone_table(client_table, (FROM_CLIENT_LABEL, insight_label))
    client_table._tbl.remove(insight_row._tr)

    client_hint = Paragraph(client_heading._p.getnext(), client_heading._parent)
    insights_heading = deepcopy(client_heading._p)
    insights_hint = deepcopy(client_hint._p)
    _set_paragraph_text(client_hint, CLIENT_HINT)
    client_table._tbl.addnext(insights_heading)
    insights_heading.addnext(insights_hint)
    insights_hint.addnext(insights_table._tbl)
    _set_paragraph_text(
        Paragraph(insights_heading, client_heading._parent), "8. WIEDZA Z ROZMÓW"
    )
    _set_paragraph_text(Paragraph(insights_hint, client_heading._parent), INSIGHTS_HINT)

    for section in doc.sections:
        for paragraph in section.footer.paragraphs:
            for run in paragraph.runs:
                if "wzór v4.0" in run.text:
                    run.text = run.text.replace("wzór v4.0", "wzór v5.0")

    doc.save(TARGET)
    return TARGET


if __name__ == "__main__":
    print(build())
