"""Builder: szablony trzech aneksów zakładki „Generator aneksów”.

Ticket „Generator aneksów w Generatorze umów B2B” (29.09.2026) podaje PEŁNĄ
treść trzech aneksów (uzupełnienie danych firmy, zmiana daty startu, zmiana
stawki) — wspólny wstęp, komparycję Partnera w trzech wariantach (osoba
fizyczna, JDG, spółka), § 1 każdego typu i wspólny § 2. Dlatego te szablony
NIE powstają ze wzorów działu (``build_b2b_document_templates.py``), tylko
z tego opisu: jedna lista bloków → DOCX (docxtpl) i HTML (podgląd). Zmiana
treści aneksu = zmiana listy bloków tutaj i ponowne uruchomienie.

Oprawa (nagłówek z logo, stopka „B2B.net S.A., Aleje Jerozolimskie 180,
Kopernik Office Building…”, marginesy, style) pochodzi z obecnego
``annex_party_data_pl.docx`` — jedynego wzoru działu ze stopką. Builder
czyści jego treść i wstawia nową, więc ponowne uruchomienie daje ten sam wynik.

    python scripts/build_b2b_annex_templates.py
"""

from __future__ import annotations

import copy
import html as _html
from dataclasses import dataclass
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

OUT_DIR = Path(__file__).resolve().parents[1] / "app" / "templates" / "documents"
SHELL = OUT_DIR / "annex_party_data_pl.docx"

FONT = "Aptos"


# ── Bloki treści ─────────────────────────────────────────────────────────────


@dataclass(frozen=True)
class P:
    """Akapit: lista (tekst, pogrubienie)."""

    runs: tuple[tuple[str, bool], ...]
    align: str = "both"
    size: int = 9
    indent_cm: float = 0.0
    hanging_cm: float = 0.0
    space_before: int = 0


@dataclass(frozen=True)
class Ctl:
    """Sterowanie Jinja: ``if …``, ``else``, ``endif``, ``for …``, ``endfor``."""

    statement: str


@dataclass(frozen=True)
class Signatures:
    pass


def p(*runs: str | tuple[str, bool], **kw) -> P:
    return P(
        tuple((r, False) if isinstance(r, str) else r for r in runs),
        **kw,
    )


def b(text: str) -> tuple[str, bool]:
    return (text, True)


def heading(text: str) -> P:
    return p(b(text), align="center", space_before=6)


def header(action: str) -> list:
    return [
        p("{{ document_place }}, {{ document_date }} r.", align="right"),
        p(
            b("ANEKS DO UMOWY O WSPÓŁPRACĘ B2B"),
            align="center",
            size=10,
            space_before=12,
        ),
        p(
            f"Strony zgodnie postanawiają dokonać {action} w treści umowy nr "
            "{{ base.contract_number }} zawartej dnia {{ base.signing_date }} roku "
            "w Warszawie pomiędzy:",
            space_before=12,
        ),
        p(
            b("{{ company.name }}"),
            " z siedzibą w Warszawie, {{ company.address }}, numer NIP: "
            "{{ company.nip }}, reprezentowaną przez {{ company.representative }},",
        ),
        p("zwaną dalej „Spółka” lub „", b("B2BNET"), "”"),
        p("a"),
    ]


PHYSICAL_PARTY = [
    p(
        b("{{ g.g_pan }} {{ partner.name_instrumental or partner.name or '…' }}"),
        ", {{ g.g_zamieszkaly }} pod adresem: {{ partner.home_address or '…' }}, "
        "{{ g.g_legitymujacy }} dowodem osobistym o numerze: "
        "{{ partner.id_document or '…' }},",
    ),
]

JDG_PARTY = p(
    b("{{ g.g_pan }} {{ partner.name_instrumental or partner.name or '…' }}"),
    " {{ g.g_prowadzacy }} działalność gospodarczą pod firmą: "
    "{{ partner.legal_name or '…' }}, zarejestrowaną w Centralnej Ewidencji "
    "i Informacji o Działalności Gospodarczej pod adresem: "
    "{{ partner.business_address or '…' }}, NIP: {{ partner.nip or '…' }}, "
    "REGON: {{ partner.regon or '…' }},",
)

COMPANY_PARTY = p(
    b("{{ partner.legal_name or '…' }}"),
    " z siedzibą {{ partner.seat_locative or '…' }}, pod adresem "
    "{{ partner.business_address or '…' }}, wpisaną do Rejestru Przedsiębiorców "
    "Krajowego Rejestru Sądowego, prowadzonego przez "
    "{{ partner.registry_court or '…' }}, pod numerem KRS: "
    "{{ partner.krs or '…' }}, NIP: {{ partner.nip or '…' }}, REGON: "
    "{{ partner.regon or '…' }}, o kapitale zakładowym "
    "{{ partner.share_capital or '…' }}, reprezentowaną przez "
    "{{ partner.representation or '…' }},",
)

BUSINESS_PARTY = [
    Ctl("if partner.is_company"),
    COMPANY_PARTY,
    Ctl("else"),
    JDG_PARTY,
    Ctl("endif"),
]

PARTY_CLOSING = [
    p("{{ g.g_zwanym }} dalej „", b("Partnerem"), "”"),
    p("B2BNET i Partner łącznie są zwani „", b("Stronami"), "”"),
]

FINAL = [
    heading("§ 2"),
    p(
        "1.\t{{ doc.effective_clause }}",
        indent_cm=0.63,
        hanging_cm=0.63,
    ),
    p(
        "2.\tStrony ustalają, że niniejszy aneks może być zawarty w formie "
        "pisemnej lub w formie elektronicznej z wykorzystaniem kwalifikowanego "
        "podpisu elektronicznego, zgodnie z art. 78¹ §1 Kodeksu cywilnego.",
        indent_cm=0.63,
        hanging_cm=0.63,
    ),
    p(
        "3.\tPozostałe postanowienia umowy nie ulegają zmianie.",
        indent_cm=0.63,
        hanging_cm=0.63,
    ),
    p(
        "4.\tAneks sporządzono w dwóch jednobrzmiących egzemplarzach po jednym "
        "dla każdej ze stron.",
        indent_cm=0.63,
        hanging_cm=0.63,
    ),
    Signatures(),
]

# Frazę („z dniem 15.10.2026 r.”) składa kontekst — dokument sprzed generatora
# aneksów niesie tryb daty („nie później niż …”) i ma się pobrać bez zmian.
START_QUOTE = (
    "„Partner zobowiązany jest do podjęcia świadczenia usług wynikających "
    "z niniejszej Umowy {phrase}”"
)

TEMPLATES: dict[str, list] = {
    "annex_party_data_pl": [
        *header("uzupełnienia"),
        *PHYSICAL_PARTY,
        *PARTY_CLOSING,
        p(
            "Aneks, zgodnym postanowieniem Stron, wprowadza do umowy nr "
            "{{ base.contract_number }} zawartej dnia {{ base.signing_date }} "
            "następujące treści:",
            space_before=12,
        ),
        heading("§1"),
        # Dokument sprzed generatora aneksów bywał wystawiony dla spółki —
        # ponowne pobranie ma oddać to, co wtedy wydano.
        Ctl("if doc.entity_type == 'company'"),
        p(
            b(
                "„Partnerem” w rozumieniu zapisów umowy o współpracę B2B nr "
                "{{ base.contract_number }} z dnia {{ base.signing_date }} jest "
                "spółka {{ doc.new_legal_name or '…' }} z siedzibą pod adresem: "
                "{{ doc.new_business_address or '…' }}, wpisana do rejestru "
                "przedsiębiorców Krajowego Rejestru Sądowego pod numerem KRS: "
                "{{ doc.company_krs or '…' }}, NIP: {{ doc.new_nip or '…' }}, "
                "REGON: {{ doc.new_regon or '…' }}, reprezentowana przez: "
                "{{ doc.company_representative or '…' }}."
            )
        ),
        Ctl("else"),
        p(
            "„Partnerem” w rozumieniu zapisów umowy o współpracę B2B nr "
            "{{ base.contract_number }} z dnia {{ base.signing_date }} jest "
            "{{ g.g_pan_nom }} {{ partner.name or '…' }} "
            "{{ g.g_prowadzacy_nom }} działalność gospodarczą pod firmą: "
            "{{ doc.new_legal_name or '…' }}, zarejestrowaną w Centralnej "
            "Ewidencji i Informacji o Działalności Gospodarczej pod adresem: "
            "{{ doc.new_business_address or '…' }}, NIP: {{ doc.new_nip or '…' }}, "
            "REGON: {{ doc.new_regon or '…' }}."
        ),
        Ctl("endif"),
        *FINAL,
    ],
    "annex_start_date_pl": [
        *header("zmian"),
        *BUSINESS_PARTY,
        *PARTY_CLOSING,
        heading("§1"),
        p(
            "Aneks zgodnym postanowieniem stron wprowadza do umowy nr "
            "{{ base.contract_number }} następujące zmiany:"
        ),
        p("a) Określoną w {{ doc.paragraph_ref }} treść:"),
        p(START_QUOTE.format(phrase="{{ doc.current_start_phrase }}"), indent_cm=0.63),
        p("zastępuje się treścią:"),
        p(START_QUOTE.format(phrase="{{ doc.new_start_phrase }}"), indent_cm=0.63),
        p(
            "b) W Załączniku nr 3 do Umowy (wzór określający Klienta B2BNET) "
            "pozycję „Data rozpoczęcia świadczenia usług”:"
        ),
        p("{{ doc.current_start_value }}", indent_cm=0.63),
        p("zastępuje się wartością:"),
        p("{{ doc.new_start_value }}", indent_cm=0.63),
        *FINAL,
    ],
    "annex_rate_change_pl": [
        *header("zmian"),
        *BUSINESS_PARTY,
        *PARTY_CLOSING,
        p(
            "Aneks zgodnym postanowieniem stron wprowadza do Umowy nr "
            "{{ base.contract_number }} zawartej dnia {{ base.signing_date }} "
            "następującą zmianę w {{ doc.paragraph_ref }}:",
            space_before=12,
        ),
        heading("§ 1"),
        p("{{ doc.paragraph_ref }} otrzymuje następujące brzmienie:"),
        Ctl("if doc.rate_single"),
        p(
            "Z tytułu realizacji usług wynikających z niniejszej Umowy Partnerowi "
            "przysługiwać będzie {{ doc.rate_sentence }}"
        ),
        Ctl("else"),
        p(
            "Z tytułu realizacji usług wynikających z niniejszej Umowy Partnerowi "
            "przysługiwać będzie:"
        ),
        Ctl("for line in doc.rate_lines"),
        p("•\t{{ line }}", indent_cm=0.63, hanging_cm=0.63),
        Ctl("endfor"),
        Ctl("endif"),
        *FINAL,
    ],
}


# ── DOCX ─────────────────────────────────────────────────────────────────────


def _clear_body(document: docx.document.Document) -> None:
    body = document.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _style_run(run, *, bold: bool, size: int) -> None:
    run.bold = bold
    run.font.name = FONT
    run.font.size = Pt(size)
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is not None:
        fonts.set(qn("w:cs"), "Arial")


def _add_paragraph(document, block: P):
    para = document.add_paragraph(style="Normalny1")
    fmt = para.paragraph_format
    fmt.alignment = {
        "both": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
    }[block.align]
    fmt.line_spacing = 1.5
    fmt.space_after = Pt(0)
    fmt.space_before = Pt(block.space_before)
    if block.indent_cm:
        fmt.left_indent = Cm(block.indent_cm)
    if block.hanging_cm:
        fmt.first_line_indent = Cm(-block.hanging_cm)
        fmt.tab_stops.add_tab_stop(Cm(block.indent_cm))
    for text, bold in block.runs:
        parts = text.split("\t")
        for index, part in enumerate(parts):
            if index:
                run = para.add_run()
                run.add_tab()
                _style_run(run, bold=bold, size=block.size)
            if part:
                run = para.add_run(part)
                _style_run(run, bold=bold, size=block.size)
    return para


def _add_control(document, statement: str) -> None:
    para = document.add_paragraph(style="Normalny1")
    para.add_run("{%p " + statement + " %}")


def _no_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl_pr.append(borders)


def _add_signatures(document) -> None:
    for _ in range(2):
        _add_paragraph(document, p(""))
    table = document.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _no_borders(table)
    for row, text in enumerate(("________________________", "(podpis)")):
        for col in range(2):
            cell = table.cell(row, col)
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run(text)
            _style_run(run, bold=False, size=9)


def build_docx(key: str, blocks: list, shell_bytes: bytes) -> None:
    import io

    document = docx.Document(io.BytesIO(shell_bytes))
    _clear_body(document)
    for block in blocks:
        if isinstance(block, Ctl):
            _add_control(document, block.statement)
        elif isinstance(block, Signatures):
            _add_signatures(document)
        else:
            _add_paragraph(document, block)
    core = document.core_properties
    core.author = "B2B.net S.A."
    core.last_modified_by = "B2B.net S.A."
    core.title = "Aneks do umowy o współpracę B2B"
    core.comments = ""
    document.save(str(OUT_DIR / f"{key}.docx"))


# ── HTML (podgląd) ───────────────────────────────────────────────────────────


def _html_runs(block: P) -> str:
    out = []
    for text, bold in block.runs:
        # Jinja zostaje dosłownie; escapujemy wyłącznie znaki HTML treści.
        chunk = _html.escape(text, quote=False).replace("\t", " ")
        out.append(f"<strong>{chunk}</strong>" if bold else chunk)
    return "".join(out)


def build_html(key: str, blocks: list) -> None:
    lines: list[str] = []
    for block in blocks:
        if isinstance(block, Ctl):
            lines.append("{% " + block.statement + " %}")
        elif isinstance(block, Signatures):
            lines.append(
                '<table class="signatures"><tr><td>________________________</td>'
                "<td>________________________</td></tr>"
                "<tr><td>(podpis)</td><td>(podpis)</td></tr></table>"
            )
        else:
            cls = {"center": "center", "right": "right"}.get(block.align)
            if block.indent_cm:
                cls = f"{cls} indent" if cls else "indent"
            attr = f' class="{cls}"' if cls else ""
            lines.append(f"<p{attr}>{_html_runs(block)}</p>")
    lines.append(
        '<p class="footer">B2B.net S.A., Aleje Jerozolimskie 180, Kopernik Office '
        "Building, 02-486 Warszawa, office@B2Bnetwork.pl, www.B2Bnetwork.pl</p>"
    )
    (OUT_DIR / f"{key}.html").write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    shell_bytes = SHELL.read_bytes()
    for key, blocks in TEMPLATES.items():
        build_docx(key, copy.deepcopy(blocks), shell_bytes)
        build_html(key, blocks)
        print(f"zbudowano {key}")


if __name__ == "__main__":
    main()
