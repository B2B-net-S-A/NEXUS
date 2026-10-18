"""Arkusz do ręcznego uzupełnienia statystyk sprzed 12.10.2026 — lokalnie.

Para do `scripts/import_dr_manual_fill.py` (patrz jego opis): tu powstaje
szablon XLSX z zakładką „Braki” i tu arkusz wypełniony przez zespół zamienia
się na JSON z samymi id kont — nazwiska zostają w arkuszu na komputerze, nie
trafiają na serwer, do repo ani do logów. Wymaga tylko `openpyxl`::

    # luki i listy z produkcji (import_dr_manual_fill.py --gaps / --lists)
    python3 backend/scripts/dr_manual_fill_sheet.py template \\
        --gaps braki.json --lists listy.json --out uzupelnienie.xlsx
    # arkusz wypełniony przez zespół → wejście importu
    python3 backend/scripts/dr_manual_fill_sheet.py to-json uzupelnienie.xlsx > dane.json

Zakładka = tabela `dr_*`, kolumna = pole wejścia importu (`TABLES`). Wiersz bez
liczb (np. przygotowany w „Braki”, ale nieuzupełniony) jest pomijany; wiersz
z liczbą, ale bez osoby albo daty to błąd z numerem wiersza.
"""

from __future__ import annotations

import argparse
import json
import sys
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

PEOPLE_SHEET = "Osoby"
CLIENTS_SHEET = "Klienci"
GAPS_SHEET = "Braki"
GUIDE_SHEET = "Instrukcja"
CUTOVER = date(2026, 10, 12)
# Wiersze przygotowane pod listy wyboru (osoba, klient, liczby).
PREPARED_ROWS = 500


@dataclass(frozen=True)
class Column:
    header: str
    key: str  # pole w JSON importu; "" = kolumna tylko dla zespołu
    kind: str  # person | day | month | count | text | note
    width: int = 16


@dataclass(frozen=True)
class Table:
    section: str
    title: str
    columns: tuple[Column, ...]
    # Wiersz, w którym te pola są puste, nie niesie danych — pomijamy go.
    content: tuple[str, ...]


_PERSON = Column("Osoba", "user_id", "person", 34)
_NOTE = Column("Uwagi (nie są wgrywane)", "", "note", 40)

TABLES = (
    Table(
        "weeks",
        "Tygodnie rekruterów",
        (
            _PERSON,
            Column("Pierwszy dzień wpisu", "report_date", "day", 20),
            Column("Weryfikacje", "verifications", "count"),
            Column("Rekomendacje", "recommendations", "count"),
            Column("Rozmowy u klienta", "interviews", "count", 18),
            Column("Placementy", "placements", "count"),
            Column("Dni przepracowane", "days_worked", "count", 18),
            _NOTE,
        ),
        ("verifications", "recommendations", "interviews", "placements", "days_worked"),
    ),
    Table(
        "placements",
        "Placementy z klientem",
        (
            _PERSON,
            Column("Data placementu", "placement_date", "day", 18),
            Column("Klient", "client", "text", 30),
            Column("Liczba (puste = 1)", "count", "count", 18),
            _NOTE,
        ),
        ("client", "count"),
    ),
    Table(
        "dl_months",
        "Miesiące DL",
        (
            _PERSON,
            Column("Miesiąc (RRRR-MM)", "report_month", "month", 18),
            Column("Zapytania", "requests", "count"),
            Column("Placementy", "placements", "count"),
            Column("Wakaty", "vacancies", "count"),
            Column("Otwarte zapytania", "open_requests", "count", 18),
            Column("Otwarte wakaty", "open_vacancies", "count"),
            _NOTE,
        ),
        ("requests", "placements", "vacancies", "open_requests", "open_vacancies"),
    ),
    Table(
        "board_months",
        "Rada — placementy firmy",
        (
            Column("Miesiąc (RRRR-MM)", "report_month", "month", 18),
            Column("Placementy firmy", "placements", "count", 18),
            _NOTE,
        ),
        ("placements",),
    ),
    Table(
        "board_clients",
        "Rada — klienci",
        (
            Column("Miesiąc (RRRR-MM)", "report_month", "month", 18),
            Column("Klient", "client_name", "text", 30),
            Column("Placementy", "placement_count", "count"),
            _NOTE,
        ),
        ("client_name", "placement_count"),
    ),
)
_BY_SECTION = {t.section: t for t in TABLES}

GUIDE = (
    "Uzupełnienie statystyk sprzed 12.10.2026",
    "",
    "Od 12.10.2026 statystyki liczy NEXUS. Wcześniejsze dni liczy wyłącznie z tabel "
    "DynaReportera — ten arkusz dopisuje do nich to, czego DynaReporter nie ma.",
    "Zakładka „Braki” wylicza luki policzone z produkcji; wiersze do uzupełnienia są "
    "już przygotowane w zakładkach z danymi (osoba i okres, liczby puste).",
    "",
    "Zasady",
    "• Tylko dni do 11.10.2026. Wiersz z 12.10.2026 albo później import odrzuci.",
    "• Osobę wybierz z listy (zakładka „Osoby”: aktywne konta, a osoby, które "
    "odeszły — konto z importu Traffita). Osoby spoza listy nie wgramy.",
    "• Puste pole = bez zmiany (w nowym wierszu 0, dni przepracowane 5). Wpisana "
    "liczba zastępuje zapisaną — także 0.",
    "• Tydzień rekrutera: pierwszy dzień roboczy tygodnia. Tydzień przechodzący przez "
    "koniec miesiąca to dwa wiersze — drugi od 1. dnia miesiąca (jak w "
    "DynaReporterze). Są tylko sumy miesięczne? Jeden wiersz od pierwszego dnia "
    "roboczego miesiąca.",
    "• Październik 2026 w „Miesiące DL” to tylko 1–11.10 — od 12.10 liczy NEXUS.",
    "• Klienci: w 2026 placement z klientem (zakładka „Placementy z klientem”, data "
    "dnia, liczba ustawia ile placementów tej osoby u tego klienta tego dnia; 0 "
    "usuwa). Przed 2026 — „Rada — klienci” (miesiąc, klient, liczba).",
    "• Przed 2026 placementy firmy w miesiącu: „Rada — placementy firmy”. Placementy "
    "z wpisów rekruterów tego miesiąca dostają osobę, reszta zostaje bez osoby.",
    "• Klienta wybierz z listy albo wpisz nazwę klienta z NEXUS-a — nieznaną nazwę "
    "import odrzuci, żeby literówka nie stała się nowym klientem.",
    "• Miesiące 1–9.2026 są rozliczone w konkursach (Liga Mistrzów, wyścigi, Liga "
    "DL Q1–Q3): uzupełnienie zmienia statystyki, nie wypłacone wyniki.",
    "",
    "Wgranie: arkusz → JSON (lokalnie, bez nazwisk) → przebieg na sucho z raportem "
    "„przed → po” → zapis dopiero po zgodzie Artura.",
)


class SheetError(ValueError):
    """Arkusz ma błędy — wypisujemy je wszystkie, nic nie konwertujemy."""


def _people_labels(people: list[dict]) -> list[dict]:
    """Jedna osoba = jedno konto: aktywne, inaczej najstarsze (import Traffita)."""
    chosen: dict[str, dict] = {}
    for person in people:
        key = str(person["name"]).strip().lower()
        best = chosen.get(key)
        if best is None or (person["active"], -person["id"]) > (
            best["active"],
            -best["id"],
        ):
            chosen[key] = person
    labels: dict[str, int] = {}
    out = []
    for person in sorted(
        chosen.values(), key=lambda p: (not p["active"], str(p["name"]).lower())
    ):
        label = str(person["name"]).strip()
        if not person["active"]:
            label = f"{label} (nieaktywne)"
        if label in labels:
            label = f"{label} (#{person['id']})"
        labels[label] = person["id"]
        out.append({**person, "label": label})
    return out


def _style_header(ws, columns: tuple[Column, ...]) -> None:
    fill = PatternFill("solid", fgColor="E7E6E6")
    for i, col in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=i, value=col.header)
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[cell.column_letter].width = col.width
    ws.freeze_panes = "A2"


def _add_validations(ws, table: Table, people_n: int, clients_n: int) -> None:
    last = PREPARED_ROWS + 1
    for i, col in enumerate(table.columns, start=1):
        letter = ws.cell(row=1, column=i).column_letter
        cells = f"{letter}2:{letter}{last}"
        if col.kind == "person":
            dv = DataValidation(
                type="list",
                formula1=f"={PEOPLE_SHEET}!$A$2:$A${people_n + 1}",
                allow_blank=True,
                showErrorMessage=True,
                errorTitle="Osoba spoza listy",
                error="Wybierz osobę z zakładki „Osoby”.",
            )
        elif col.kind == "text" and clients_n:
            # Lista podpowiada, ale wolno wpisać nazwę klienta z NEXUS-a.
            dv = DataValidation(
                type="list",
                formula1=f"={CLIENTS_SHEET}!$A$2:$A${clients_n + 1}",
                allow_blank=True,
                showErrorMessage=False,
            )
        elif col.kind == "count":
            dv = DataValidation(
                type="whole",
                operator="between",
                formula1="0",
                formula2="10000",
                allow_blank=True,
                showErrorMessage=True,
                error="Liczba całkowita od 0.",
            )
        else:
            continue
        ws.add_data_validation(dv)
        dv.add(cells)
    for i, col in enumerate(table.columns, start=1):
        letter = ws.cell(row=1, column=i).column_letter
        for row in range(2, last + 1):
            if col.kind == "day":
                ws[f"{letter}{row}"].number_format = "yyyy-mm-dd"
            elif col.kind == "month":
                ws[f"{letter}{row}"].number_format = "@"


def _gap_rows(gaps: dict, names: dict[int, str]) -> tuple[list[tuple], dict]:
    """Wiersze zakładki „Braki” i wiersze przygotowane w zakładkach z danymi."""
    rows: list[tuple] = []
    prepared: dict[str, list[dict]] = {"weeks": [], "dl_months": []}
    for m in gaps["months"]:
        if m["month"] >= "2026-01":
            continue
        if m["weeks"] == 0:
            board = m["board_placements"]
            rows.append(
                (
                    "Tygodnie rekruterów (KPI)",
                    "wszyscy rekruterzy",
                    m["month"],
                    "Tygodnie rekruterów",
                    "0 wpisów"
                    + (
                        ""
                        if board is None
                        else f"; Rada: {board} placementów bez osoby"
                    ),
                )
            )
    for m in gaps["months"]:
        if m["month"] < "2026-01" and m["dl_entries"] == 0:
            rows.append(
                (
                    "Wpisy Delivery Leadów",
                    "wszyscy DL",
                    m["month"],
                    "Miesiące DL",
                    "0 wpisów",
                )
            )
    october = CUTOVER.isoformat()[:7]
    for gap in gaps["dl_missing"]:
        name = names.get(gap["user_id"], f"konto #{gap['user_id']}")
        is_october = gap["month"] == october
        rows.append(
            (
                "Wpis Delivery Leada" + (" 1–11.10.2026" if is_october else ""),
                name,
                f"{gap['month']} (1–11.10)" if is_october else gap["month"],
                "Miesiące DL (wiersz przygotowany)",
                "brak wpisu",
            )
        )
        prepared["dl_months"].append({"person": name, "report_month": gap["month"]})
    for gap in gaps["recruiter_october_missing"]:
        name = names.get(gap["user_id"], f"konto #{gap['user_id']}")
        start = date.fromisoformat(gap["report_date"])
        span = "1–4.10.2026" if start.day == 1 else "5–11.10.2026"
        rows.append(
            (
                "Tydzień rekrutera",
                name,
                span,
                "Tygodnie rekruterów (wiersz przygotowany)",
                "brak wpisu",
            )
        )
        prepared["weeks"].append({"person": name, "report_date": start})
    rows.append(
        (
            "Bez tabeli w DynaReporterze",
            "—",
            "do 11.10.2026",
            "— (nie do uzupełnienia tym arkuszem)",
            "etapy Tablicy poza weryfikacją, rekomendacją, rozmową i placementem, "
            "źródła kandydatów, czas do zatrudnienia, hit ratio — liczone od 12.10.2026",
        )
    )
    return rows, prepared


def build_template(gaps: dict, lists: dict, out: Path) -> None:
    people = _people_labels(lists["people"])
    names = {p["id"]: p["label"] for p in people}
    clients = sorted({str(c).strip() for c in lists["clients"] if str(c).strip()})
    wb = Workbook()

    guide = wb.active
    guide.title = GUIDE_SHEET
    guide.column_dimensions["A"].width = 120
    for i, line in enumerate(GUIDE, start=1):
        cell = guide.cell(row=i, column=1, value=line)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1 or line in {"Zasady"}:
            cell.font = Font(bold=True, size=13 if i == 1 else 11)

    gap_rows, prepared = _gap_rows(gaps, names)
    ws = wb.create_sheet(GAPS_SHEET)
    _style_header(
        ws,
        (
            Column("Co brakuje", "", "note", 30),
            Column("Osoba", "", "note", 34),
            Column("Okres", "", "note", 18),
            Column("Gdzie wpisać", "", "note", 40),
            Column("Stan w NEXUS-ie", "", "note", 60),
        ),
    )
    for r, values in enumerate(gap_rows, start=2):
        for c, value in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=value)

    for table in TABLES:
        ws = wb.create_sheet(table.title)
        _style_header(ws, table.columns)
        _add_validations(ws, table, len(people), len(clients))
        keys = [c.key if c.kind != "person" else "person" for c in table.columns]
        for r, row in enumerate(prepared.get(table.section, []), start=2):
            for key, value in row.items():
                ws.cell(row=r, column=keys.index(key) + 1, value=value)

    ws = wb.create_sheet(PEOPLE_SHEET)
    _style_header(
        ws,
        (
            Column("Osoba", "", "note", 40),
            Column("Id konta NEXUS", "", "note", 16),
            Column("Rola", "", "note", 22),
            Column("Aktywne", "", "note", 10),
        ),
    )
    for p in people:
        ws.append([p["label"], p["id"], p["role"], "tak" if p["active"] else "nie"])
    ws = wb.create_sheet(CLIENTS_SHEET)
    _style_header(ws, (Column("Klient (DynaReporter)", "", "note", 40),))
    for name in clients:
        ws.append([name])
    wb.save(out)


# ── Arkusz → JSON ───────────────────────────────────────────────────────


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _as_day(value: Any) -> Optional[str]:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    raw = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(raw, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def _as_month(value: Any) -> Optional[str]:
    if isinstance(value, (datetime, date)):
        return f"{value.year:04d}-{value.month:02d}"
    raw = str(value).strip()
    for fmt in ("%Y-%m", "%m.%Y", "%Y-%m-%d", "%m/%Y"):
        try:
            parsed = datetime.strptime(raw, fmt)
        except ValueError:
            continue
        return f"{parsed.year:04d}-{parsed.month:02d}"
    return None


def _as_count(value: Any) -> Optional[int]:
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    raw = str(value).strip()
    return int(raw) if raw.isdigit() else None


def _people_map(wb) -> dict[str, int]:
    if PEOPLE_SHEET not in wb.sheetnames:
        raise SheetError(f"brak zakładki „{PEOPLE_SHEET}”")
    out: dict[str, int] = {}
    for label, account, *_ in wb[PEOPLE_SHEET].iter_rows(min_row=2, values_only=True):
        if _blank(label) or _as_count(account) is None:
            continue
        out[str(label).strip()] = _as_count(account)
    return out


def sheet_to_payload(path: Path) -> tuple[dict[str, list[dict]], dict[str, int]]:
    """Arkusz → wejście `import_dr_manual_fill.py` i liczba pominiętych pustych wierszy."""
    wb = load_workbook(path, data_only=True)
    people = _people_map(wb)
    errors: list[str] = []
    payload: dict[str, list[dict]] = {}
    skipped: dict[str, int] = {}
    for table in TABLES:
        if table.title not in wb.sheetnames:
            raise SheetError(f"brak zakładki „{table.title}”")
        rows: list[dict] = []
        skipped[table.section] = 0
        ws = wb[table.title]
        for r, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            cells = dict(zip(table.columns, values))
            by_key = {c.key: v for c, v in cells.items() if c.key}
            if all(_blank(v) for v in by_key.values()):
                continue
            if all(_blank(by_key.get(k)) for k in table.content):
                skipped[table.section] += 1
                continue
            where = f"{table.title}, wiersz {r}"
            row: dict[str, Any] = {}
            for col, value in cells.items():
                if not col.key:
                    continue
                if col.kind == "person":
                    if isinstance(value, (int, float)) and _as_count(value):
                        row[col.key] = _as_count(value)
                    elif not _blank(value) and str(value).strip() in people:
                        row[col.key] = people[str(value).strip()]
                    else:
                        errors.append(f"{where}: osoby nie ma w zakładce „Osoby”")
                elif _blank(value):
                    row[col.key] = None
                    if col.kind in {"day", "month"}:
                        errors.append(f"{where}: brak — {col.header}")
                elif col.kind == "day":
                    row[col.key] = _as_day(value)
                    if row[col.key] is None:
                        errors.append(f"{where}: {col.header} — data RRRR-MM-DD")
                elif col.kind == "month":
                    row[col.key] = _as_month(value)
                    if row[col.key] is None:
                        errors.append(f"{where}: {col.header} — miesiąc RRRR-MM")
                elif col.kind == "count":
                    row[col.key] = _as_count(value)
                    if row[col.key] is None:
                        errors.append(f"{where}: {col.header} — liczba całkowita od 0")
                else:
                    row[col.key] = str(value).strip()
            rows.append(row)
        payload[table.section] = rows
    if errors:
        raise SheetError("\n".join(errors))
    return payload, skipped


def main(argv: Optional[list[str]] = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    sub = parser.add_subparsers(dest="command", required=True)
    template = sub.add_parser("template", help="szablon XLSX z zakładką „Braki”")
    template.add_argument("--gaps", type=Path, required=True)
    template.add_argument("--lists", type=Path, required=True)
    template.add_argument("--out", type=Path, required=True)
    to_json = sub.add_parser("to-json", help="arkusz → JSON importu (stdout)")
    to_json.add_argument("sheet", type=Path)
    args = parser.parse_args(argv)
    if args.command == "template":
        build_template(
            json.loads(args.gaps.read_text()),
            json.loads(args.lists.read_text()),
            args.out,
        )
        print(f"zapisano {args.out}", file=sys.stderr)
        return 0
    try:
        payload, skipped = sheet_to_payload(args.sheet)
    except SheetError as exc:
        print(str(exc), file=sys.stderr)
        return 2
    print(json.dumps(payload, ensure_ascii=False, indent=1))
    counts = ", ".join(
        f"{_BY_SECTION[s].title}: {len(rows)} (pominięte puste: {skipped[s]})"
        for s, rows in payload.items()
    )
    print(counts, file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())
