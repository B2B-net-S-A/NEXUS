"""Baza startowa „Champion po ludzku”: technologie, role i opisy klientów (29.09.2026).

Research w internecie robimy RAZ, przegląda go człowiek, a wynik trafia do bazy:

* technologie i role → pliki w repo ``app/data/plain_knowledge/terms.json``
  i ``roles.json`` (wiedza ogólna, bez danych klientów; zasiew przy starcie);
* opisy klientów → prosto do kart klientów na produkcji (NIE do repo — repo jest
  publiczne, a lista klientów to informacja handlowa).

Przebieg (w kontenerze backendu produkcji)::

    docker exec <backend> python -m scripts.build_plain_knowledge \\
        --plan /tmp/plain_plan.json --review /tmp/plain_review.xlsx
    # człowiek przegląda arkusz (docker cp na zewnątrz), poprawia plan, potem lokalnie:
    python -m scripts.build_plain_knowledge --write-seed plain_plan.json
    # po wdrożeniu (zasiew technologii i ról już w bazie), na produkcji:
    docker exec <backend> python -m scripts.build_plain_knowledge \\
        --apply-clients /tmp/plain_plan.json
    docker exec <backend> python -m scripts.build_plain_knowledge --assign-roles

* ``--plan`` czyta bazę i woła research (web search + model), NIC nie zapisuje.
  Role powstają z grupowania historii rekrutacji (``job_similarity.similarity_score``,
  grupowanie zachłanne, próg 55, grupy ≥ ``--min-group``).
* ``--write-seed`` zapisuje technologie i role z planu do plików w repo.
* ``--apply-clients`` wpisuje opisy klientów tylko tam, gdzie karta nie ma opisu,
  tylko klientom z rekrutacjami, żywym i bez duplikatu; ``origin=web``, wpis
  w historii karty; paragon w ``app_settings`` (same liczby i ID).
* ``--assign-roles`` dopasowuje rolę każdej rekrutacji (``role_matcher``),
  poza wyborem ręcznym.
"""

from __future__ import annotations

import argparse
import asyncio
import json
import logging
import sys
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Optional

BACKEND_ROOT = Path(__file__).resolve().parents[1]
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

logger = logging.getLogger("build_plain_knowledge")

PLAN_VERSION = 1
RECEIPT_KEY = "plain_knowledge_clients_apply"
CLUSTER_MIN_SCORE = 55
FREQUENT_MUST_MIN = 20
# Jedna reguła czyszczenia nazwy roli z aplikacją (nowe role z researchu).
from app.services.plain_knowledge.role_matcher import clean_role_name  # noqa: E402


def cluster_jobs(jobs: list[dict[str, Any]], min_group: int) -> list[dict[str, Any]]:
    """Grupowanie zachłanne: rekrutacja dołącza do najlepszego lidera ≥ progu."""
    from app.services.job_similarity import similarity_score

    leaders: list[dict[str, Any]] = []
    for job in jobs:
        best: Optional[tuple[int, dict[str, Any]]] = None
        for group in leaders:
            lead = group["leader"]
            score = similarity_score(
                job["skills"],
                job["title"],
                job["cc"],
                lead["skills"],
                lead["title"],
                lead["cc"],
            )
            if score >= CLUSTER_MIN_SCORE and (best is None or score > best[0]):
                best = (score, group)
        if best is None:
            leaders.append({"leader": job, "members": [job]})
        else:
            best[1]["members"].append(job)
    out = []
    for group in leaders:
        members = group["members"]
        if len(members) < min_group:
            continue
        names = Counter(m["role_name"] for m in members if m["role_name"])
        titles = Counter(w for m in members for w in m["title"])
        skills = Counter(s for m in members for s in m["skills"])
        cats = Counter(m["cc_slug"] for m in members if m["cc_slug"])
        out.append(
            {
                "name": names.most_common(1)[0][0] if names else "",
                "jobs": len(members),
                "job_ids": [m["id"] for m in members][:50],
                "title_words": [w for w, _ in titles.most_common(4)],
                "skills": [s for s, _ in skills.most_common(6)],
                "category": cats.most_common(1)[0][0] if cats else None,
                "name_candidates": [n for n, _ in names.most_common(5)],
            }
        )
    return sorted(out, key=lambda g: -g["jobs"])


async def _load_jobs(db) -> list[dict[str, Any]]:
    from sqlalchemy import select

    from app.models.competence_category import CompetenceCategory
    from app.models.job import Job
    from app.services import champion_view
    from app.services.job_similarity import skill_set, title_tokens

    slugs = dict(
        (await db.execute(select(CompetenceCategory.id, CompetenceCategory.slug))).all()
    )
    rows = (
        await db.execute(
            select(
                Job.id,
                Job.title,
                Job.working_title,
                Job.must_skills,
                Job.champion_profile,
                Job.competence_category_id,
            ).order_by(Job.id.desc())
        )
    ).all()
    out = []
    for jid, title, working, must, profile, cc in rows:
        role = champion_view.basics(profile).get("role_name")
        name = clean_role_name(
            role if isinstance(role, str) and role.strip() else (title or "")
        )
        out.append(
            {
                "id": jid,
                "title": title_tokens(name)
                or title_tokens(working)
                or title_tokens(title),
                "skills": skill_set(must, profile),
                "cc": cc,
                "cc_slug": slugs.get(cc),
                "role_name": name,
            }
        )
    return out


async def _frequent_terms(db) -> list[str]:
    from sqlalchemy import select, text

    from app.models.skill import Skill

    names = [n for (n,) in (await db.execute(select(Skill.canonical_name))).all()]
    rows = (
        await db.execute(
            text(
                """
                SELECT min(btrim(COALESCE(e->>'name', e#>>'{}'))) AS name, count(*) AS n
                  FROM jobs j,
                       jsonb_array_elements(CASE WHEN jsonb_typeof(j.must_skills) = 'array'
                                                 THEN j.must_skills ELSE '[]' END) e
                 GROUP BY lower(btrim(COALESCE(e->>'name', e#>>'{}')))
                HAVING count(*) >= :n
                """
            ),
            {"n": FREQUENT_MUST_MIN},
        )
    ).all()
    return names + [name for name, _n in rows if name]


async def _clients_without_about(db) -> list[tuple[int, str]]:
    from sqlalchemy import text

    rows = (
        await db.execute(
            text(
                """
                SELECT c.id, COALESCE(NULLIF(btrim(c.display_name), ''), c.name)
                  FROM clients c
                  LEFT JOIN client_playbooks p ON p.client_id = c.id
                 WHERE c.deleted_at IS NULL AND c.merged_into_client_id IS NULL
                   AND c.hidden = FALSE
                   AND EXISTS (SELECT 1 FROM jobs j WHERE j.client_id = c.id)
                   AND NOT EXISTS (SELECT 1 FROM clients d
                                    WHERE d.merged_into_client_id = c.id AND d.deleted_at IS NULL
                                      AND EXISTS (SELECT 1 FROM client_playbooks dp WHERE dp.client_id = d.id))
                   AND (p.id IS NULL OR COALESCE(btrim(p.about_for_candidate), '') = '')
                 ORDER BY c.id
                """
            )
        )
    ).all()
    return [(cid, name) for cid, name in rows if name]


def load_journal(path: Path) -> dict[tuple[str, str], dict[str, Any]]:
    """Wyniki już zbadane (ostatni zapis per hasło). Brak wyniku = do ponowienia."""
    done: dict[tuple[str, str], dict[str, Any]] = {}
    if not path.exists():
        return done
    for line in path.read_text(encoding="utf-8").splitlines():
        try:
            rec = json.loads(line)
        except ValueError:
            continue
        if isinstance(rec, dict) and rec.get("result"):
            done[(rec["kind"], rec["key"])] = rec
    return done


async def _plan_items(args: argparse.Namespace) -> list[dict[str, Any]]:
    from app.core.database import AsyncSessionLocal
    from app.services.plain_knowledge import knowledge

    async with AsyncSessionLocal() as db:
        terms_src = await _frequent_terms(db)
        jobs = await _load_jobs(db)
        clients = await _clients_without_about(db)
        await db.commit()

    items: list[dict[str, Any]] = []
    seen: set[str] = set()
    for name in terms_src:
        if not knowledge.researchable_term(name):
            continue
        key = knowledge.term_key_for(name)
        if key in seen:
            continue
        seen.add(key)
        items.append({"kind": "term", "key": key, "term_key": key, "name": name})
    terms = [i for i in items if i["kind"] == "term"][: args.limit_terms or None]

    roles: dict[str, dict[str, Any]] = {}
    for group in cluster_jobs(jobs, args.min_group):
        name = group["name"] or " ".join(group["title_words"]).title()
        slug = knowledge.slugify(name)
        if slug in roles:  # dwie grupy o tej samej nazwie = jedna rola
            roles[slug]["jobs"] += group["jobs"]
            continue
        roles[slug] = {**group, "kind": "role", "key": slug, "slug": slug, "name": name}
    role_items = list(roles.values())[: args.limit_roles or None]

    client_items = [
        {"kind": "client", "key": str(cid), "client_id": cid, "name": name}
        for cid, name in clients[: args.limit_clients or None]
    ]
    return terms + role_items + client_items


async def build_plan(args: argparse.Namespace) -> dict[str, Any]:
    """Research wznawialny i równoległy: każdy wynik od razu trafia do dziennika
    ``<plan>.jsonl``, a ponowny bieg pomija hasła już zbadane (deploy ubija
    proces w kontenerze — bieg trwa godziny)."""
    from app.core.config import settings
    from app.core.database import AsyncSessionLocal
    from app.services.plain_knowledge.research import research
    from app.services.skill_taxonomy_loader import refresh_alias_map

    await refresh_alias_map()
    # Jednorazowa budowa bazy startowej: dzienny sufit researchu dotyczy
    # aplikacji, nie tego przebiegu (kilkaset haseł naraz).
    settings.PLAIN_KNOWLEDGE_RESEARCH_DAILY_LIMIT = 100_000
    items = await _plan_items(args)
    journal = Path(f"{args.plan}.jsonl")
    done = load_journal(journal)
    todo = [i for i in items if (i["kind"], i["key"]) not in done]
    logger.info(
        "plan: %s haseł, zrobione %s, do zbadania %s", len(items), len(done), len(todo)
    )
    sem = asyncio.Semaphore(max(1, args.concurrency))
    lock = asyncio.Lock()
    counter = {"n": 0}

    async def run(item: dict[str, Any]) -> None:
        async with sem:
            async with AsyncSessionLocal() as db:
                result = await research(
                    db, item["kind"], item["name"], skills=item.get("skills")
                )
        async with lock:
            with journal.open("a", encoding="utf-8") as fh:
                fh.write(
                    json.dumps({**item, "result": result}, ensure_ascii=False) + "\n"
                )
            counter["n"] += 1
            logger.info(
                "%s/%s %s %s: %s",
                counter["n"],
                len(todo),
                item["kind"],
                item["key"],
                "ok" if result else "brak",
            )

    await asyncio.gather(*(run(i) for i in todo))
    done = load_journal(journal)

    def rows(kind: str) -> list[dict[str, Any]]:
        return [
            {**i, "result": (done.get((kind, i["key"])) or {}).get("result")}
            for i in items
            if i["kind"] == kind
        ]

    return {
        "version": PLAN_VERSION,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "terms": rows("term"),
        "roles": rows("role"),
        "clients": rows("client"),
    }


def write_review(plan: dict[str, Any], path: Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Technologie"
    ws.append(
        [
            "Klucz",
            "Nazwa",
            "Po ludzku",
            "Czym zajmuje się osoba",
            "W CV",
            "Nie myl z",
            "Źródła",
        ]
    )
    for t in plan.get("terms", []):
        r = t.get("result") or {}
        ws.append(
            [
                t["term_key"],
                r.get("display_name") or t["name"],
                r.get("summary"),
                r.get("does"),
                ", ".join(r.get("cv_hints") or []),
                r.get("confused_with"),
                " ".join(s["url"] for s in r.get("sources") or []),
            ]
        )
    ws = wb.create_sheet("Role")
    ws.append(
        [
            "Slug",
            "Nazwa",
            "Rekrutacji",
            "Inne nazwy",
            "Słowa tytułu",
            "Umiejętności",
            "Opis",
            "Przykład z codzienności",
            "Dzień pracy",
            "Pytania kandydatów",
            "Źródła",
        ]
    )
    for g in plan.get("roles", []):
        r = g.get("result") or {}
        ws.append(
            [
                g["slug"],
                g["name"],
                g["jobs"],
                ", ".join(g.get("name_candidates") or []),
                ", ".join(g["title_words"]),
                ", ".join(g["skills"]),
                r.get("summary"),
                r.get("example"),
                " | ".join(r.get("day_to_day") or []),
                " | ".join(r.get("candidate_questions") or []),
                " ".join(s["url"] for s in r.get("sources") or []),
            ]
        )
    ws = wb.create_sheet("Klienci")
    ws.append(["Id", "Klient", "Opis dla kandydata", "Źródła"])
    for c in plan.get("clients", []):
        r = c.get("result") or {}
        ws.append(
            [
                c["client_id"],
                c["name"],
                r.get("about"),
                " ".join(s["url"] for s in r.get("sources") or []),
            ]
        )
    wb.save(path)


def write_seed(plan: dict[str, Any], data_dir: Path) -> tuple[int, int]:
    terms = []
    for t in plan.get("terms", []):
        r = t.get("result") or {}
        if not r.get("summary"):
            continue
        terms.append(
            {
                "term_key": t["term_key"],
                "display_name": r.get("display_name") or t["name"],
                "summary": r.get("summary"),
                "does": r.get("does"),
                "cv_hints": r.get("cv_hints") or [],
                "confused_with": r.get("confused_with"),
                "sources": r.get("sources") or [],
            }
        )
    roles = []
    for g in plan.get("roles", []):
        r = g.get("result") or {}
        if not r.get("summary"):
            continue
        roles.append(
            {
                "slug": g["slug"],
                "name": g["name"],
                "summary": r.get("summary"),
                "example": r.get("example"),
                "day_to_day": r.get("day_to_day") or [],
                "candidate_questions": r.get("candidate_questions") or [],
                "typical_skills": g["skills"],
                "match_rules": {
                    "title_words": g["title_words"],
                    "skills": g["skills"],
                    "category": g.get("category"),
                },
                "sources": r.get("sources") or [],
            }
        )
    for name, key, rows in (
        ("terms.json", "terms", terms),
        ("roles.json", "roles", roles),
    ):
        path = data_dir / name
        current = (
            json.loads(path.read_text(encoding="utf-8"))
            if path.exists()
            else {"version": 1}
        )
        current[key] = sorted(rows, key=lambda x: x.get("term_key") or x.get("slug"))
        path.write_text(
            json.dumps(current, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
    return len(terms), len(roles)


async def apply_clients(plan: dict[str, Any]) -> dict[str, Any]:
    from sqlalchemy import select, text

    from app.core.database import AsyncSessionLocal
    from app.models.client_playbook import ClientPlaybook
    from app.models.client_playbook_event import ClientPlaybookEvent

    wanted = {
        c["client_id"]: c
        for c in plan.get("clients", [])
        if (c.get("result") or {}).get("about")
    }
    applied: list[int] = []
    async with AsyncSessionLocal() as db:
        allowed = {cid for cid, _ in await _clients_without_about(db)}
        for cid, item in wanted.items():
            if cid not in allowed:
                continue
            row = await db.scalar(
                select(ClientPlaybook)
                .where(ClientPlaybook.client_id == cid)
                .with_for_update()
            )
            if row is not None and (row.about_for_candidate or "").strip():
                continue
            if row is None:
                last = await db.scalar(
                    text(
                        "SELECT max(playbook_version) FROM client_playbook_events WHERE client_id = :c"
                    ),
                    {"c": cid},
                )
                row = ClientPlaybook(client_id=cid, version=int(last or 0) + 1)
                db.add(row)
            else:
                row.version = int(row.version or 1) + 1
            about = item["result"]["about"]
            row.about_for_candidate = about
            row.about_for_candidate_origin = "web"
            row.about_for_candidate_sources = item["result"].get("sources") or []
            db.add(
                ClientPlaybookEvent(
                    client_id=cid,
                    playbook_version=row.version,
                    action="web_research",
                    changes={"about_for_candidate": {"from": None, "to": about}},
                    actor_user_id=None,
                    actor_name="Research AI",
                )
            )
            applied.append(cid)
        receipt = {
            "applied_at": datetime.now(timezone.utc).isoformat(),
            "count": len(applied),
            "client_ids": applied,
        }
        await db.execute(
            text(
                "INSERT INTO app_settings (key, value) VALUES (:k, CAST(:v AS jsonb)) "
                "ON CONFLICT (key) DO UPDATE SET value = EXCLUDED.value"
            ),
            {"k": RECEIPT_KEY, "v": json.dumps(receipt)},
        )
        await db.commit()
    return receipt


async def assign_roles() -> dict[str, int]:
    from sqlalchemy import select

    from app.core.database import AsyncSessionLocal
    from app.models.competence_category import CompetenceCategory
    from app.models.job import Job
    from app.models.plain_knowledge import RoleProfile
    from app.services.plain_knowledge import role_matcher
    from app.services.skill_taxonomy_loader import refresh_alias_map

    await refresh_alias_map()
    stats = {"assigned": 0, "unchanged": 0, "no_match": 0}
    async with AsyncSessionLocal() as db:
        roles = [
            role_matcher.rules_of(rid, r)
            for rid, r in (
                await db.execute(
                    select(RoleProfile.id, RoleProfile.match_rules).where(
                        RoleProfile.status == "ready"
                    )
                )
            ).all()
        ]
        slugs = dict(
            (
                await db.execute(select(CompetenceCategory.id, CompetenceCategory.slug))
            ).all()
        )
        jobs = (
            (
                await db.execute(
                    select(Job).where(
                        (Job.role_profile_source.is_(None))
                        | (Job.role_profile_source == "auto")
                    )
                )
            )
            .scalars()
            .all()
        )
        for job in jobs:
            title, skills = role_matcher.job_signals(job)
            best = role_matcher.best_role(
                roles, title, skills, slugs.get(job.competence_category_id)
            )
            if best is None:
                stats["no_match"] += 1
            elif best == job.role_profile_id:
                stats["unchanged"] += 1
            else:
                job.role_profile_id, job.role_profile_source = best, "auto"
                stats["assigned"] += 1
        await db.commit()
    return stats


async def _main(args: argparse.Namespace) -> int:
    if args.plan:
        plan = await build_plan(args)
        Path(args.plan).write_text(
            json.dumps(plan, ensure_ascii=False, indent=1), encoding="utf-8"
        )
        if args.review:
            write_review(plan, Path(args.review))
        print(
            f"plan: terms={len(plan['terms'])} roles={len(plan['roles'])} clients={len(plan['clients'])}"
        )
    if args.write_seed:
        plan = json.loads(Path(args.write_seed).read_text(encoding="utf-8"))
        terms, roles = write_seed(
            plan, BACKEND_ROOT / "app" / "data" / "plain_knowledge"
        )
        print(f"seed: terms={terms} roles={roles}")
    if args.apply_clients:
        plan = json.loads(Path(args.apply_clients).read_text(encoding="utf-8"))
        print(json.dumps(await apply_clients(plan)))
    if args.assign_roles:
        print(json.dumps(await assign_roles()))
    return 0


def main(argv: Optional[list[str]] = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--plan")
    parser.add_argument("--review")
    parser.add_argument("--write-seed")
    parser.add_argument("--apply-clients")
    parser.add_argument("--assign-roles", action="store_true")
    parser.add_argument("--min-group", type=int, default=5)
    parser.add_argument("--limit-terms", type=int, default=0)
    parser.add_argument("--limit-roles", type=int, default=0)
    parser.add_argument("--limit-clients", type=int, default=0)
    parser.add_argument("--concurrency", type=int, default=6)
    args = parser.parse_args(argv)
    logging.basicConfig(level=logging.INFO)
    return asyncio.run(_main(args))


if __name__ == "__main__":
    raise SystemExit(main())
