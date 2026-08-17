"""Pełny parse profili Championa — schemat v3 (komplet 7 sekcji), wznawialny."""
import asyncio, json, sys
sys.path.insert(0, "/app")
from pathlib import Path
from starlette.concurrency import run_in_threadpool

FILES_DIR = Path("/tmp/champions_all")
OUT = Path("/tmp/champ-parsed-v3.jsonl")

PROMPT = """Z dokumentu "Profil Championa" (opis idealnego kandydata uzgodniony z klientem; sekcje: informacje o projekcie, profil kandydata MUST/NICE, kontekst projektu, screening, success strategy, informacja o kliencie, standardy rekrutacji) wyciągnij DOKŁADNIE tę strukturę JSON:
{
 "role_name": str|null,
 "seniority_min_years": int|null,          // np. z "Minimum 10 lat doświadczenia"
 "must_skills": [{"name": str}],           // sekcja MUST-HAVE, punkt=wpis, skracaj do esencji
 "nice_skills": [{"name": str}],           // sekcja NICE-TO-HAVE
 "rate_value": float|null,                 // sama liczba, np. 122.50
 "rate_raw": str|null,                     // stawka dokładnie jak w dokumencie
 "location": str|null,                     // lokalizacja kandydata/biura
 "work_mode": "stacjonarnie"|"hybrydowo"|"zdalnie"|null,
 "basics": {"onsite_days_per_week": int|null, "candidate_location_pref": str|null, "language": str|null},
 "start_date": str|null,
 "contract_length": str|null,              // np. "3-5 miesięcy z możliwością przedłużenia"
 "sectors": [str],                         // branże z wymagań/kontekstu (banking, fintech, płatności...)
 "project_context": {"about": str, "responsibilities": str, "selling_points": str},
   // about: cel+charakter, max 4 zdania; responsibilities: obowiązki (może być lista po przecinkach);
   // selling_points: "Co przekona kandydata do oferty?" + najważniejsze atuty klienta z sekcji 6
 "screening_questions": [{"id": "q1", "question": str, "ideal_answer": str, "deal_breaker": str}],
   // WSZYSTKIE pytania z sekcji SCREENING, z pełnymi idealnymi odpowiedziami i deal-breakerami
 "sourcing": {"keywords": str, "target_companies": str, "notes": str},
   // keywords: frazy z "Kluczowe słowa do wyszukiwania" rozdzielone przecinkami;
   // target_companies: "Firmy docelowe" jako tekst z priorytetami;
   // notes: "Uwagi / plan działania" — ZWŁASZCZA wskazówki typu "nie zawężamy do X"
 "disqualifiers": [str],                   // twarde wykluczenia, jeśli dokument je podaje
 "client_standards": {"priority_rules": str|null, "offlimit": bool|null, "contract_type": str|null, "cv_language": str|null}
   // priority_rules: reguły typu "kandydaci z bankowością w pierwszej kolejności"
}
Zasady: NICZEGO nie wymyślaj — brak sekcji/informacji = null/pusta wartość. Cytuj wiernie, skracaj tylko redakcyjnie. Zwróć SAM JSON.

DOKUMENT:
"""

def extract(path: Path):
    if path.suffix == ".docx":
        try:
            import docx
            d = docx.Document(str(path))
            parts = [p.text for p in d.paragraphs if p.text.strip()]
            for t in d.tables:
                for row in t.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(dict.fromkeys(cells)))
            return "\n".join(parts)
        except Exception:
            return None
    try:
        from app.services.cv_text_extractor import extract_text
        return extract_text(str(path), path.name)
    except Exception:
        return None

async def parse_one(call_claude, p: Path):
    rid, fid = p.stem.split("__")[1:3]
    text = extract(p)
    if not text or len(text) < 200:
        return {"rid": int(rid), "file": int(fid), "error": "no_text"}
    try:
        msg = await run_in_threadpool(
            call_claude,
            model="claude-haiku-4-5-20251001",
            max_tokens=6000,
            temperature=0,
            thinking={"type": "disabled"},
            messages=[{"role": "user", "content": PROMPT + text[:14000]}],
        )
        raw = "".join(b.text for b in msg.content if getattr(b, "type", "") == "text")
        parsed = json.loads(raw[raw.find("{"):raw.rfind("}") + 1])
        return {"rid": int(rid), "file": int(fid), "parsed": parsed,
                "usage": {"in": msg.usage.input_tokens, "out": msg.usage.output_tokens}}
    except Exception as e:
        return {"rid": int(rid), "file": int(fid), "error": repr(e)[:200]}

async def main():
    from app.services.claude_client import call_claude
    done = set()
    if OUT.exists():
        for line in OUT.open(encoding="utf-8"):
            try:
                done.add(json.loads(line)["file"])
            except Exception:
                pass
    files = [p for p in sorted(FILES_DIR.glob("champ__*"))
             if int(p.stem.split("__")[2]) not in done]
    print(f"do przetworzenia: {len(files)} (pominięte, już zrobione: {len(done)})")
    sem = asyncio.Semaphore(3)
    out_lock = asyncio.Lock()
    counter = {"n": 0}

    async def worker(p):
        async with sem:
            row = await parse_one(call_claude, p)
        async with out_lock:
            with OUT.open("a", encoding="utf-8") as f:
                f.write(json.dumps(row, ensure_ascii=False) + "\n")
            counter["n"] += 1
            if counter["n"] % 50 == 0:
                print(f"postęp: {counter['n']}/{len(files)}")

    await asyncio.gather(*(worker(p) for p in files))
    print("KONIEC PARSE V3")

asyncio.run(main())
