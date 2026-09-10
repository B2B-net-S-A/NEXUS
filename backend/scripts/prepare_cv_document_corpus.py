"""Build synthetic full-CV inputs and a review manifest, without provider calls.

python -m scripts.prepare_cv_document_corpus --output /tmp/cv-document-corpus
This prepares benchmark inputs; it does not score or approve generated CVs.
"""

import argparse
import hashlib
import json
from pathlib import Path

from docx import Document


CORPUS_PATH = Path(__file__).parents[1] / "app/data/cv_quality/full_documents_v2.json"


def load_cases():
    raw = CORPUS_PATH.read_bytes()
    corpus = json.loads(raw)
    cases = corpus["cases"]
    if corpus["schema_version"] != 1 or corpus["synthetic"] is not True:
        raise ValueError("Unsupported corpus")
    if len(cases) != 40 or len({case["id"] for case in cases}) != 40:
        raise ValueError("Expected forty unique variants")
    scenarios = {case["scenario"] for case in cases}
    if len(scenarios) != 20:
        raise ValueError("Expected twenty source scenarios")
    for scenario in scenarios:
        pair = [case for case in cases if case["scenario"] == scenario]
        if {case["language"] for case in pair} != {"pl", "en"}:
            raise ValueError("Missing language pair")
        if pair[0]["cv_paragraphs"] != pair[1]["cv_paragraphs"]:
            raise ValueError("Language variants must use identical sources")
    for case in cases:
        if case["source_format"] not in {"docx_table", "docx_paragraphs"}:
            raise ValueError("Unsupported source format")
        if not case["expected"]["review_criterion"] or len(case["cv_paragraphs"]) < 8:
            raise ValueError("Incomplete source or acceptance criteria")
        if case["expected"]["requires_human_review"] is not True:
            raise ValueError("Human quality review must remain explicit")
    return cases, hashlib.sha256(raw).hexdigest()


def write_scanned_pdf(paragraphs, path, font):
    """Rasterize every source paragraph; never clip text at page boundaries."""
    from PIL import Image, ImageDraw

    pages = []
    page = None
    y = 0
    try:
        for paragraph in paragraphs:
            lines, line = [], ""
            for word in paragraph.split():
                if font.getlength(word) > 1080:
                    raise ValueError("Source word exceeds scan page width")
                proposed = f"{line} {word}".strip()
                if font.getlength(proposed) > 1080:
                    lines.append(line)
                    line = word
                else:
                    line = proposed
            if line:
                lines.append(line)
            for line in lines:
                if page is None or y + 36 > 1674:
                    page = Image.new("RGB", (1240, 1754), "white")
                    pages.append(page)
                    y = 80
                ImageDraw.Draw(page).text((80, y), line, font=font, fill="black")
                y += 36
            y += 18
        if not pages:
            raise ValueError("Empty scan source")
        pages[0].save(
            path, "PDF", resolution=150, save_all=True, append_images=pages[1:]
        )
    finally:
        for page in pages:
            page.close()


def prepare(output: Path, *, scan_font: Path | None = None):
    cases, digest = load_cases()
    font = None
    scan_metadata = None
    if scan_font is not None:
        from PIL import ImageFont

        font = ImageFont.truetype(str(scan_font), size=24)
        scan_metadata = {
            "renderer": "synthetic_scan_v1",
            "font_sha256": hashlib.sha256(scan_font.read_bytes()).hexdigest(),
            "dpi": 150,
        }
        digest = hashlib.sha256(
            (digest + json.dumps(scan_metadata, sort_keys=True)).encode()
        ).hexdigest()
    extension = "pdf" if font is not None else "docx"
    output.mkdir(parents=True, exist_ok=True)
    existing = output / "manifest.json"
    if existing.exists():
        manifest = json.loads(existing.read_text())
        if manifest.get("corpus_sha256") != digest or len(
            manifest.get("cases", [])
        ) != len(cases):
            raise ValueError("Existing corpus differs; use a new output directory")
        for original, saved in zip(cases, manifest["cases"], strict=True):
            filename = f"{original['id']}.{extension}"
            if (
                any(saved.get(key) != value for key, value in original.items())
                or saved.get("input_file") != filename
                or not (output / filename).is_file()
                or hashlib.sha256((output / filename).read_bytes()).hexdigest()
                != saved.get("input_sha256")
            ):
                raise ValueError(
                    "Existing corpus input changed; refusing to overwrite evidence"
                )
        return manifest
    if any(output.iterdir()):
        raise ValueError("Incomplete corpus directory; use a new output directory")
    manifest = {"corpus_sha256": digest, "evaluated": False, "cases": []}
    if scan_metadata:
        manifest["scan"] = scan_metadata
    for case in cases:
        path = output / f"{case['id']}.{extension}"
        document = Document()
        paragraphs = case["cv_paragraphs"]
        document.add_heading(paragraphs[0], 0)
        if case["source_format"] == "docx_table":
            table = document.add_table(rows=0, cols=1)
            for paragraph in paragraphs[1:]:
                table.add_row().cells[0].text = paragraph
        else:
            for paragraph in paragraphs[1:]:
                document.add_paragraph(paragraph)
        if font is not None:
            write_scanned_pdf(paragraphs, path, font)
        else:
            document.save(path)
        manifest["cases"].append(
            {
                **case,
                "input_format": "scanned_pdf"
                if font is not None
                else case["source_format"],
                "input_file": path.name,
                "input_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                "outcome": "not_run",
                "human_accepted": None,
            }
        )
    (output / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n"
    )
    return manifest


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument(
        "--scan-font", type=Path, help="Render scan inputs using this font file"
    )
    args = parser.parse_args()
    manifest = prepare(args.output, scan_font=args.scan_font)
    print(
        json.dumps(
            {
                "cases": len(manifest["cases"]),
                "evaluated": False,
                "corpus_sha256": manifest["corpus_sha256"],
            }
        )
    )
