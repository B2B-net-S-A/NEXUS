"""Builder: wzory działu (Word) → szablony dokumentów pochodnych umowy B2B.

Czyta 13 wzorów z katalogu działu i zapisuje do ``app/templates/documents/``
po parze plików na typ × język z rejestru (``services/b2b_documents/registry``):

    <typ>_<język>.docx  — szablon docxtpl (``{{ }}``, ``{%p if %}``)
    <typ>_<język>.html  — lustro Jinja do podglądu w przeglądarce

Zasady (decyzja Artura 23.09.2026 — wzory poprawiamy sami do umowy 2026):

* formatowanie, nagłówek, stopka i logo wzoru zostają; pola do wypełnienia
  (``____``, ``……``, żółte ``w:highlight`` z PRZYKŁADOWYMI danymi, „XX/XX”,
  „[nr umowy]”) zamieniają się w placeholdery, podświetlenie znika;
* komparycja B2B.net → ``{{ company.* }}``, Partnera → ``{{ partner.* }}``
  + formy rodzajowe ``{{ g.* }}``;
* numery paragrafów umowy bazowej → ``{{ refs.* }}`` (wzory cytowały paragrafy
  STAREJ umowy);
* żadne dane ze wzoru (nazwisko, NIP, numer umowy, data, kwota) nie może
  zostać w wyniku — skrypt odrzuca wynik z literalną datą, numerem umowy albo
  niewypełnionym polem, a metadane pakietu (autor, SharePoint, dodatki Worda)
  są czyszczone.

Uruchom po zmianie wzoru; wynik jest commitowany. NIE jest częścią ścieżki
produkcyjnej.

    python scripts/build_b2b_document_templates.py <katalog_wzorów> [katalog_wyjściowy]
"""

from __future__ import annotations

import argparse
import copy
import html as _html
import io
import re
import unicodedata
import zipfile
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING

import docx
from docx.document import Document as _Doc
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run

OUT_DIR = Path(__file__).resolve().parents[1] / "app" / "templates" / "documents"

DOTS = "…"


class BuildError(RuntimeError):
    pass


# ── Segmenty tekstu: (tekst, pogrubienie) — None = formatowanie z wzoru ──────

if TYPE_CHECKING:
    Segment = tuple[str, bool | None]


def _segments(repl: str | list[Segment]) -> list[Segment]:
    return [(repl, None)] if isinstance(repl, str) else list(repl)


# ── Operacje na akapitach ────────────────────────────────────────────────────


def paragraphs(doc: _Doc) -> list[Paragraph]:
    """Wszystkie akapity treści (także w tabelach), w kolejności dokumentu."""
    return [Paragraph(p, doc) for p in doc.element.body.iter(qn("w:p"))]


def _runs_text(p: Paragraph) -> tuple[list[Run], list[str]]:
    runs = list(p.runs)
    return runs, [r.text for r in runs]


_TEXT_ONLY = {qn("w:rPr"), qn("w:t"), qn("w:tab"), qn("w:br"), qn("w:cr")}


def _is_text_run(r: Run) -> bool:
    return all(child.tag in _TEXT_ONLY for child in r._r)


def _new_run_like(src: Run, p: Paragraph, text: str, bold: bool | None) -> Run:
    el = copy.deepcopy(src._r)
    run = Run(el, p)
    run.text = text
    if bold is not None:
        run.bold = bold
    return run


def _replace_span(p: Paragraph, start: int, end: int, repl) -> None:
    """Zastąp znaki [start, end) tekstu akapitu, zachowując formatowanie
    tekstu wokół — wstawka dziedziczy formatowanie runu, w którym zaczyna się
    dopasowanie (chyba że segment mówi inaczej o pogrubieniu)."""
    runs, texts = _runs_text(p)
    full = "".join(texts)
    if full != p.text:
        raise BuildError(f"Akapit ma tekst poza runami (hiperłącze?): {p.text[:60]!r}")
    if not runs:
        p.add_run("")
        runs, texts = _runs_text(p)
    starts, acc = [], 0
    for t in texts:
        starts.append(acc)
        acc += len(t)

    def run_at(pos: int) -> int:
        for k, t in enumerate(texts):
            if starts[k] <= pos < starts[k] + len(t):
                return k
        return len(texts) - 1

    i = run_at(start)
    j = run_at(end - 1) if end > start else i
    ri = runs[i]
    prefix = texts[i][: start - starts[i]]
    suffix = texts[j][end - starts[j] :]
    for k in range(i + 1, j):
        runs[k]._r.getparent().remove(runs[k]._r)
    anchor = ri._r
    new_els = []
    for text, bold in _segments(repl):
        if text:
            new_els.append(_new_run_like(ri, p, text, bold)._r)
    if j > i:
        runs[j].text = suffix
    elif suffix:
        new_els.append(_new_run_like(ri, p, suffix, None)._r)
    for el in new_els:
        anchor.addnext(el)
        anchor = el
    if prefix:
        ri.text = prefix
    elif _is_text_run(ri):
        ri._r.getparent().remove(ri._r)
    else:
        ri.text = ""


def _sub_para(p: Paragraph, rx: re.Pattern, repl) -> int:
    matches = [m for m in rx.finditer(p.text) if m.end() > m.start()]
    for m in reversed(matches):
        value = repl(m) if callable(repl) else repl
        _replace_span(p, m.start(), m.end(), value)
    return len(matches)


def sub(doc: _Doc, pattern: str, repl, *, count: int | None = 1, flags: int = 0) -> int:
    """Podstaw ``repl`` za każde dopasowanie w akapitach; ``count`` = oczekiwana
    liczba trafień (inaczej błąd — wzór się zmienił i reguła nie trafia)."""
    rx = re.compile(pattern, flags | re.S)
    hits = sum(_sub_para(p, rx, repl) for p in paragraphs(doc))
    if count is not None and hits != count:
        raise BuildError(f"Reguła {pattern[:70]!r}: {hits} trafień, oczekiwano {count}")
    return hits


def find(doc: _Doc, pattern: str) -> Paragraph:
    rx = re.compile(pattern, re.S)
    hits = [p for p in paragraphs(doc) if rx.search(p.text)]
    if len(hits) != 1:
        raise BuildError(
            f"Wzorzec akapitu {pattern[:70]!r}: {len(hits)} trafień, oczekiwano 1"
        )
    return hits[0]


def set_text(p: Paragraph, repl) -> Paragraph:
    if not p.runs:
        p.add_run("")
    _replace_span(p, 0, len(p.text), repl)
    return p


def drop_empty_numbered(doc: _Doc) -> None:
    """Pusty akapit listy numerowanej rysuje sam numer — a wzór daty startu
    „wpisywał” tak nową datę (numeracja wielopoziomowa ustawiona na 19.03.2026).
    Taki akapit nie niesie treści, tylko etykietę z przykładowych danych."""
    for p in paragraphs(doc):
        ppr = p._p.pPr
        if ppr is not None and ppr.numPr is not None and not p.text.strip():
            remove(p)


def remove(p: Paragraph) -> None:
    p._p.getparent().remove(p._p)


def clone_after(anchor: Paragraph, repl, *, like: Paragraph | None = None) -> Paragraph:
    """Nowy akapit za ``anchor`` o formatowaniu ``like`` (domyślnie ``anchor``)."""
    el = copy.deepcopy((like or anchor)._p)
    anchor._p.addnext(el)
    new = Paragraph(el, anchor._parent)
    runs = new.runs
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    if runs:
        runs[0].text = "x"
    else:
        new.add_run("x")
    set_text(new, repl)
    return new


def _control(anchor: Paragraph, tag: str, *, before: bool) -> Paragraph:
    el = copy.deepcopy(anchor._p)
    ppr = el.find(qn("w:pPr"))
    if ppr is not None:
        num = ppr.find(qn("w:numPr"))
        if num is not None:
            ppr.remove(num)
    (anchor._p.addprevious if before else anchor._p.addnext)(el)
    ctrl = Paragraph(el, anchor._parent)
    for r in ctrl.runs[1:]:
        r._r.getparent().remove(r._r)
    if ctrl.runs:
        ctrl.runs[0].text = tag
    else:
        ctrl.add_run(tag)
    return ctrl


def wrap_if(first: Paragraph, last: Paragraph, cond: str) -> None:
    """Akapity od ``first`` do ``last`` renderują się tylko przy ``cond``."""
    _control(first, "{%p if " + cond + " %}", before=True)
    _control(last, "{%p endif %}", before=False)


def control_before(anchor: Paragraph, statement: str) -> Paragraph:
    return _control(anchor, "{%p " + statement + " %}", before=True)


def control_after(anchor: Paragraph, statement: str) -> Paragraph:
    return _control(anchor, "{%p " + statement + " %}", before=False)


def paragraph_heading(doc: _Doc, number: int) -> Paragraph:
    return find(doc, r"^\s*§\s*" + str(number) + r"\s*\.?\s*$")


def renumber(doc: _Doc, numbers: dict[int, str]) -> None:
    """Nagłówki „§ N” → wyrażenie Jinja (paragrafy warunkowe przesuwają numerację)."""
    targets = [(n, paragraph_heading(doc, n)) for n in numbers]
    for n, p in targets:
        sub_in(p, r"§\s*" + str(n) + r"\s*\.?", "§ " + numbers[n])


def sub_in(p: Paragraph, pattern: str, repl, count: int = 1) -> None:
    hits = _sub_para(p, re.compile(pattern, re.S), repl)
    if hits != count:
        raise BuildError(
            f"Reguła {pattern[:60]!r} w akapicie {p.text[:50]!r}: {hits} trafień"
        )


# ── Wspólne klocki treści ────────────────────────────────────────────────────

NAME_INSTR = "{{ partner.name_instrumental or partner.name or '" + DOTS + "' }}"
NAME_NOM = "{{ partner.name or '" + DOTS + "' }}"


def dot(expr: str) -> str:
    return "{{ " + expr + " or '" + DOTS + "' }}"


PL_DATE_HEADER = "{{ document_place }}, {{ document_date }} r."
EN_DATE_HEADER = "{{ document_place }}, {{ document_date }}"

COMPANY_PL_SHORT: list[Segment] = [
    ("{{ company.name }}", True),
    (
        " z siedzibą w Warszawie, {{ company.address }}, numer NIP: "
        "{{ company.nip }}, reprezentowaną przez {{ company.representative }},",
        False,
    ),
]
COMPANY_PL_FULL: list[Segment] = [
    ("{{ company.name }}", True),
    (
        " z siedzibą w Warszawie, {{ company.address }}, wpisaną do "
        "Rejestru Przedsiębiorców prowadzonego przez {{ company.court }}, KRS pod "
        "nr {{ company.krs }}, NIP: {{ company.nip }}, o kapitale zakładowym "
        "{{ company.share_capital }}, reprezentowaną przez {{ company.representative }},",
        False,
    ),
]
COMPANY_EN: list[Segment] = [
    ("{{ company.name }}", True),
    (
        " with its registered office in {{ company.place }}, {{ company.address }}, "
        "NIP (Tax Identification Number): {{ company.nip }}, represented by "
        "{{ company.representative }},",
        False,
    ),
]

RX_COMPANY_PL = (
    r"B2B\.net S\.A\.\s*z siedzibą w Warszawie,\s*Aleje Jerozolimskie 180,\s*"
    r"02-486 Warszawa,\s*numer NIP:\s*5711707392,\s*reprezentowaną przez Pana "
    r"Artura Twardowskiego – Prezesa Zarządu,"
)
RX_COMPANY_EN = (
    r"B2B\.net S\.A\.\s*with\s+its\s+registered office in Warsaw,\s*Aleje\s*"
    r"Jerozolimskie\s*180,\s*02-486 Warsaw,\s*NIP\s*\(Tax Identification Number\):"
    r"\s*5711707392,\s*represented by Mr\.\s*Artur Twardowski – President of the "
    r"Management Board,"
)

JDG_PL: list[Segment] = [
    ("{{ g.g_pan }} " + NAME_INSTR, True),
    (
        " {{ g.g_prowadzacy }} działalność gospodarczą pod firmą: "
        + dot("partner.legal_name")
        + ", zarejestrowaną w Centralnej Ewidencji i Informacji o Działalności "
        "Gospodarczej pod adresem: "
        + dot("partner.business_address")
        + ", NIP: "
        + dot("partner.nip")
        + ", REGON: "
        + dot("partner.regon")
        + ",",
        False,
    ),
]
RX_JDG_PL = (
    r"Pan\S*?\s*_+\s*prowadząc\S*\s+działalność gospodarczą pod firmą:.*?"
    r"REGON:\s*_+\s*,?"
)

JDG_EN: list[Segment] = [
    ("{{ g.g_mr }} " + NAME_NOM, True),
    (
        " conducting business activity under the name: "
        + dot("partner.legal_name")
        + ", registered in the Central Register and Information on Economic "
        "Activity at the address: "
        + dot("partner.business_address")
        + ", Tax Identification Number (NIP): "
        + dot("partner.nip")
        + ", National Business Registry Number (REGON): "
        + dot("partner.regon")
        + ",",
        False,
    ),
]
RX_JDG_EN = (
    r"Mr/Ms\s*_+\s*conducting business activity under the name:.*?REGON\)?:\s*_+\s*,?"
)

# Nazwa paragrafu w wersji EN: numery są prowadzone po polsku („§ 6 ust. 1”).
EN_REF = " | replace('ust.', 'section') | replace(' pkt ', ' point ') | replace(' i ', ' and ')"


def en_ref(key: str) -> str:
    return "{{ refs." + key + EN_REF + " }}"


def zwanym(
    doc: _Doc,
    *,
    count: int = 1,
    pattern: str = r"(?m)^\s*[Zz]\s*wan\S*(?=\s+dalej\s+(jako\s+)?„Partnerem”)",
) -> None:
    """„zwany/ą dalej …” → forma narzędnika zgodna z płcią (po „a Panem X,”)."""
    sub(doc, pattern, "{{ g.g_zwanym }}", count=count)


def check_date_header_pl(doc: _Doc) -> None:
    sub(doc, r"^(\s*)Warszawa\s*,.*$", lambda m: m.group(1) + PL_DATE_HEADER)


# ── Szablony ─────────────────────────────────────────────────────────────────


def build_annex_rate_change_pl(doc: _Doc) -> None:
    check_date_header_pl(doc)
    sub(
        doc,
        r"umowy nr\s*\S+\s*zawartej dnia\s*\S+\s*roku",
        "umowy nr {{ base.contract_number }} zawartej dnia {{ base.signing_date }} roku",
    )
    sub(doc, RX_COMPANY_PL, COMPANY_PL_SHORT)
    sub(doc, RX_JDG_PL, JDG_PL)
    zwanym(doc)
    sub(doc, r"łącznie są zwani Stronami", "łącznie są zwani „Stronami”")
    # Wzór podawał numer umowy drugi raz (z pustym polem) — umowa jest już
    # wskazana w preambule.
    sub(
        doc,
        r"^Aneks zgodnym postanowieniem stron wprowadza do Umowy nr.*$",
        "Aneks zgodnym postanowieniem Stron wprowadza do Umowy następującą zmianę "
        "w {{ refs.rate_paragraph }}:",
    )
    sub(doc, r"^§\s*6\s*ust\.\s*1(?= otrzymuje)", "{{ refs.rate_paragraph }}")
    # Brzmienie ust. 1 jak w umowie 2026 (§ 6 ust. 1), z nową stawką.
    sub(
        doc,
        r"^Z tytułu realizacji usług.*$",
        "Z tytułu należytego wykonania Usług Partnerowi przysługuje wynagrodzenie "
        "obliczane jako iloczyn liczby godzin faktycznego świadczenia Usług "
        "w danym miesiącu oraz stawki godzinowej w wysokości {{ doc.new_rate }} "
        "{{ doc.currency }} (słownie: {{ doc.new_rate_words }}) netto + VAT.",
    )
    sub(
        doc,
        r"z dniem\s*_+\s*roku\.",
        "z dniem {{ doc.effective_date }} r.",
    )


def build_annex_rate_change_en(doc: _Doc) -> None:
    sub(doc, r"^Warszawa\s*,.*$", EN_DATE_HEADER)
    sub(
        doc,
        r"^The parties mutually agree to supplement the content of the agreement.*?between:",
        "The Parties mutually agree to amend the agreement No. {{ base.contract_number }} "
        "concluded on {{ base.signing_date }} in Warsaw between:",
    )
    sub(doc, RX_COMPANY_EN, COMPANY_EN)
    sub(doc, RX_JDG_EN, JDG_EN)
    sub(
        doc,
        r"The annex, by mutual decision of the parties, introduces into the contract.*$",
        "The annex, by mutual decision of the Parties, introduces the following "
        "amendment to the Agreement: " + en_ref("rate_paragraph") + " of the "
        "Agreement shall read as follows:",
    )
    # Wzór EN nie wskazywał zmienianego paragrafu; brzmienie jak § 6 ust. 1 umowy 2026.
    sub(
        doc,
        r"^The Partner will be entitled to remuneration.*$",
        "For the proper performance of the Services, the Partner shall be entitled "
        "to remuneration calculated as the product of the number of hours of actual "
        "provision of the Services in a given month and an hourly rate of "
        "{{ doc.new_rate }} {{ doc.currency }} (in words: {{ doc.new_rate_words }}) "
        "net + VAT.",
    )
    sub(
        doc,
        r"shall enter into force on\s*…+\.?",
        "shall enter into force on {{ doc.effective_date }}.",
    )


def build_annex_start_date_pl(doc: _Doc) -> None:
    check_date_header_pl(doc)
    sub(
        doc,
        r"umowy nr\s*\S+\s*zawartej dnia\s*\S+\s*roku",
        "umowy nr {{ base.contract_number }} zawartej dnia {{ base.signing_date }} roku",
    )
    sub(doc, RX_COMPANY_PL, COMPANY_PL_SHORT)
    sub(doc, RX_JDG_PL, JDG_PL)
    # Wzór miał „zwany/ą dalej Partnerem” dwa razy (drugi raz bez formy żeńskiej).
    remove(find(doc, r"^zwanym dalej „Partnerem”\s*$"))
    zwanym(doc)
    sub(
        doc,
        r"^Aneks zgodnym postanowieniem stron wprowadza do umowy nr \S+ następujące zmiany:",
        "Aneks zgodnym postanowieniem Stron wprowadza do Umowy następujące zmiany:",
    )
    # a) — data rozpoczęcia usług: w umowie 2026 to § 13 ust. 2 (wzór: § 12).
    sub(
        doc,
        r"^a\)\s*Określoną w § 12\s*Ust\.\s*2 treść:",
        "a) Określoną w {{ refs.start_paragraph }} treść:",
    )
    old_clause = (
        "„Umowa wchodzi w życie z dniem jej podpisania, a Partner zobowiązany jest "
        "do podjęcia świadczenia Usług {{ base.start_clause }}.”"
    )
    new_clause = (
        "„Umowa wchodzi w życie z dniem jej podpisania, a Partner zobowiązany jest "
        "do podjęcia świadczenia Usług {{ doc.new_start_clause }}.”"
    )
    olds = [
        p
        for p in paragraphs(doc)
        if p.text.startswith("Partner zobowiązany jest do podjęcia")
    ]
    if len(olds) != 2:
        raise BuildError("Aneks daty startu: oczekiwano dwóch brzmień § 13 ust. 2")
    set_text(olds[0], old_clause)
    set_text(olds[1], new_clause)
    # b) — Załącznik nr 3: wzór nie miał nowej wartości po „zastępuje się wartością:”.
    sub(
        doc,
        r"^b\)\s*W Załączniku nr 3 do Umowy.*$",
        "b) W dokumencie „{{ refs.appendix_start }}” do Umowy (wzór określający "
        "Klienta B2BNET) pozycję „Data rozpoczęcia świadczenia usług”:",
    )
    old_value = find(doc, r"^\d{2}\.\d{2}\.\d{4}\s*r\.\s*$")
    set_text(old_value, "{{ base.start_clause }}")
    marker = find(doc, r"^zastępuje się wartością:\s*$")
    clone_after(marker, "{{ doc.new_start_clause }}", like=old_value)
    # Rejestr nie ma pola „wchodzi w życie” dla tego aneksu → dzień podpisania.
    sub(
        doc,
        r"^Ustalone zmiany wchodzą w życie z dniem\s*_+\s*$",
        "Ustalone zmiany wchodzą w życie z dniem podpisania niniejszego aneksu.",
    )


def _party_data_entity_pl() -> str:
    company = (
        "spółka "
        + dot("doc.new_legal_name")
        + " z siedzibą pod adresem: "
        + dot("doc.new_business_address")
        + ", wpisana do rejestru przedsiębiorców Krajowego Rejestru Sądowego pod "
        "numerem KRS: "
        + dot("doc.company_krs")
        + ", NIP: "
        + dot("doc.new_nip")
        + ", REGON: "
        + dot("doc.new_regon")
        + ", reprezentowana przez: "
        + dot("doc.company_representative")
        + "."
    )
    sole = (
        "{{ g.g_pan_nom }} " + NAME_NOM + " {{ g.g_prowadzacy_nom }} działalność "
        "gospodarczą pod firmą: " + dot("doc.new_legal_name") + ", zarejestrowaną "
        "w Centralnej Ewidencji i Informacji o Działalności Gospodarczej pod adresem: "
        + dot("doc.new_business_address")
        + ", NIP: "
        + dot("doc.new_nip")
        + ", REGON: "
        + dot("doc.new_regon")
        + "."
    )
    return (
        "{% if doc.entity_type == 'company' %}"
        + company
        + "{% else %}"
        + sole
        + "{% endif %}"
    )


def _party_data_entity_en() -> str:
    company = (
        "the company "
        + dot("doc.new_legal_name")
        + " with its registered office at "
        + dot("doc.new_business_address")
        + ", entered in the Register of Entrepreneurs of the National Court "
        "Register under KRS No. "
        + dot("doc.company_krs")
        + ", NIP: "
        + dot("doc.new_nip")
        + ", REGON: "
        + dot("doc.new_regon")
        + ", represented by "
        + dot("doc.company_representative")
        + "."
    )
    sole = (
        "{{ g.g_mr }} "
        + NAME_NOM
        + " conducting business activity under the name: "
        + dot("doc.new_legal_name")
        + ", registered in the Central Register and "
        "Information on Economic Activity at the address: "
        + dot("doc.new_business_address")
        + ", Tax Identification Number (NIP): "
        + dot("doc.new_nip")
        + ", National Business Registry Number (REGON): "
        + dot("doc.new_regon")
        + "."
    )
    return (
        "{% if doc.entity_type == 'company' %}"
        + company
        + "{% else %}"
        + sole
        + "{% endif %}"
    )


def build_annex_party_data_pl(doc: _Doc) -> None:
    check_date_header_pl(doc)
    sub(
        doc,
        r"umowy nr\s*…+\s*zawartej dnia\s*…+\s*roku",
        "umowy nr {{ base.contract_number }} zawartej dnia {{ base.signing_date }} roku",
    )
    sub(doc, RX_COMPANY_PL, COMPANY_PL_SHORT)
    sub(doc, r"^dalej jako „Spółka”", "zwaną dalej jako „Spółka”")
    sub(
        doc,
        r"^Panem/Panią.*?dowodem osobistym",
        [
            ("{{ g.g_pan }} " + NAME_INSTR, True),
            (
                ", {{ g.g_zamieszkaly }} w "
                + dot("partner.home_address")
                + ", {{ g.g_legitymujacy }} dowodem osobistym"
                "{% if partner.id_document %} nr {{ partner.id_document }}{% endif %},",
                False,
            ),
        ],
    )
    zwanym(doc)
    sub(
        doc,
        r"^Aneks, zgodnym postanowieniem Stron, wprowadza do umowy nr.*$",
        "Aneks, zgodnym postanowieniem Stron, wprowadza do Umowy następujące postanowienia:",
    )
    sub(
        doc,
        r"^„Partnerem” w rozumieniu zapisów umowy o współpracę B2B.*$",
        "„Partnerem” w rozumieniu zapisów umowy o współpracę B2B nr "
        "{{ base.contract_number }} z dnia {{ base.signing_date }} r. jest "
        + _party_data_entity_pl(),
    )
    sub(doc, r"z dniem\s*…+\.?\s*$", "z dniem {{ doc.effective_date }} r.")


def build_annex_party_data_en(doc: _Doc) -> None:
    sub(doc, r"^Warsaw\s*,.*$", EN_DATE_HEADER)
    sub(
        doc,
        r"^The parties mutually agree to supplement the content of the agreement.*?between:",
        "The Parties mutually agree to supplement the agreement No. "
        "{{ base.contract_number }} concluded on {{ base.signing_date }} in Warsaw between:",
    )
    sub(
        doc,
        r"B2B\.net S\.A\.\s*with\s+its\s+registered office\s*in Warsaw.*?Management Board,",
        COMPANY_EN,
    )
    sub(
        doc,
        r"^Mr/Mrs.*?referred to as the „Partner”",
        [
            ("{{ g.g_mr }} " + NAME_NOM, True),
            (
                ", residing at " + dot("partner.home_address") + ", holding an ID card"
                "{% if partner.id_document %} No. {{ partner.id_document }}{% endif %}, "
                "hereinafter referred to as the “Partner”",
                False,
            ),
        ],
    )
    sub(
        doc,
        r"The annex, by mutual decision of the parties, introduces into the contract.*$",
        "The annex, by mutual decision of the Parties, introduces the following "
        "provisions into the Agreement:",
    )
    sub(
        doc,
        r'^"Partner" within the meaning of the provisions.*$',
        "“Partner” within the meaning of the provisions of the B2B cooperation "
        "agreement No. {{ base.contract_number }} dated {{ base.signing_date }} is "
        + _party_data_entity_en(),
    )
    sub(
        doc,
        r"shall enter into force on\s*…+\.?",
        "shall enter into force on {{ doc.effective_date }}.",
    )


def build_annex_subcontractor_pl(doc: _Doc) -> None:
    check_date_header_pl(doc)
    sub(
        doc,
        r"B2B nr\s*…+\s*zawartej dnia\s*…+\s*roku",
        "B2B nr {{ base.contract_number }} zawartej dnia {{ base.signing_date }} roku",
    )
    sub(doc, RX_COMPANY_PL, COMPANY_PL_SHORT)
    sub(doc, RX_JDG_PL, JDG_PL)
    zwanym(doc)
    # Umowa 2026 nie ma § 2 ust. 5 — zgoda na osoby skierowane żyje w tym aneksie.
    sub(
        doc,
        r"^Na zasadach określonych w § 2 ust\. 5 Umowy Partner może",
        "Partner może",
    )
    sub(
        doc,
        r"wynikającymi z § 6 Umowy oraz Deklaracji Poufności",
        "wynikającymi z {{ refs.confidentiality_paragraph }} Umowy oraz Deklaracji Poufności",
    )
    sub(
        doc,
        r"postanowieniami § 5 Umowy",
        "postanowieniami {{ refs.personal_data_paragraph }} Umowy",
    )
    sub(
        doc,
        r"określonych w § 3 Umowy\.”",
        "określonych w {{ refs.ip_paragraph }} Umowy.”",
    )
    sub(
        doc,
        r"^Pan/i\s*…+,?\s*adres e-mail:[\s….]*$",
        [
            ("{{ dg.g_pan_nom }} " + dot("doc.delegate_name"), True),
            (", adres e-mail: " + dot("doc.delegate_email") + ".", False),
        ],
    )
    sub(doc, r"z dniem\s*_+\s*\.", "z dniem {{ doc.effective_date }} r.")


def build_annex_mandate_pl(doc: _Doc) -> None:
    sub(doc, r"^zawarty dnia\s*_+\s*r\.", "zawarty dnia {{ document_date }} r.")
    sub(doc, RX_COMPANY_PL, COMPANY_PL_SHORT)
    # Spółka jest rodzaju żeńskiego — wzór miał „zwanym dalej Zleceniodawcą”.
    sub(doc, r"^zwanym dalej „Zleceniodawcą”", "zwaną dalej „Zleceniodawcą”")
    sub(
        doc,
        r"^Panem\s*_+.*?Pesel:\s*_+\s*\.",
        [
            ("{{ g.g_pan }} " + NAME_INSTR, True),
            (
                ", {{ g.g_zamieszkaly }} pod adresem: "
                + dot("partner.home_address")
                + ", PESEL: "
                + dot("partner.pesel")
                + ",",
                False,
            ),
        ],
    )
    zwanym(doc, pattern=r"^zwan\S*(?= dalej „Zleceniobiorcą)")
    sub(
        doc,
        r"Umowy zlecenie z dnia \d{2}\.\d{2}\.\d{4} r\.",
        "Umowy zlecenie z dnia {{ doc.base_signing_date }} r.",
    )
    period_f = "(1 if doc.change_period else 0)"
    rate_f = "(1 if doc.change_rate else 0)"
    extra_f = "(1 if extras else 0)"
    # § 1 — okres zlecenia (warunkowo)
    h1 = paragraph_heading(doc, 1)
    control_before(
        h1,
        "set extras = doc.extra_provisions_list if doc.extra_provisions_list is defined else []",
    )
    p1 = find(doc, r"^Zleceniobiorca będzie realizował zlecenie w okresie")
    set_text(
        p1,
        "Zleceniobiorca będzie realizował zlecenie w okresie od {{ doc.period_from }} r. "
        "do {{ doc.period_to }} r.",
    )
    wrap_if(h1, p1, "doc.change_period")
    # § 2 — stawka brutto (warunkowo)
    h2 = paragraph_heading(doc, 2)
    p2 = find(doc, r"^Za wykonanie prac określonych w § 1")
    sub_in(
        p2,
        r"w wysokości .*? brutto",
        "w wysokości {{ doc.gross_hourly_rate }} zł (słownie: "
        "{{ doc.gross_hourly_rate_words }}) brutto",
    )
    sub_in(h2, r"§\s*2", "§ {{ 1 + " + period_f + " }}")
    wrap_if(h2, p2, "doc.change_rate")
    # § 3 — dodatkowe ustępy: wzór miał pusty paragraf bez treści.
    h3 = paragraph_heading(doc, 3)
    intro = find(doc, r"^Dodane zostają następujące ustępy w § 3 Umowy:")
    sub_in(h3, r"§\s*3", "§ {{ 1 + " + period_f + " + " + rate_f + " }}")
    item = clone_after(intro, "{{ item }}", like=p1)
    control_before(item, "for item in extras")
    endfor = control_after(item, "endfor")
    wrap_if(h3, endfor, "extras")
    h4 = paragraph_heading(doc, 4)
    sub_in(
        h4, r"§\s*4", "§ {{ 1 + " + period_f + " + " + rate_f + " + " + extra_f + " }}"
    )
    sub(
        doc,
        r"wchodzą w życie z dniem \d{2}\.\d{2}\.\d{4} r\.",
        "wchodzą w życie z dniem {{ doc.effective_date }} r.",
    )


RELEASE_PL = (
    "B2B.net S.A. zwalnia Partnera z zakazu konkurencji określonego w "
    "{{ refs.non_compete_paragraph }} Umowy i zgadza się na bezpośrednie "
    "zatrudnienie Partnera lub świadczenie przez Partnera usług bezpośrednio na "
    "rzecz: " + dot("doc.non_compete_client_name") + ", spółek macierzystych oraz "
    "innych podmiotów powiązanych."
)
RELEASE_EN = (
    "B2B.net S.A. releases the Partner from the non-competition obligation set out "
    "in " + en_ref("non_compete_paragraph") + " of the Agreement and consents to "
    "the Partner being employed directly by, or providing services directly to: "
    + dot("doc.non_compete_client_name")
    + ", its parent companies and other "
    "affiliated entities."
)


def build_termination_agreement_pl(doc: _Doc) -> None:
    check_date_header_pl(doc)
    sub(
        doc,
        r"^Nr\s*_+\s*zawartej dnia\s*_+\s*r\.",
        "Nr {{ base.contract_number }} zawartej dnia {{ base.signing_date }} r.",
    )
    sub(doc, RX_COMPANY_PL, COMPANY_PL_SHORT)
    sub(doc, RX_JDG_PL, JDG_PL)
    zwanym(doc)
    sub(
        doc,
        r"B2B nr XX/XX zawartą w dniu .*? za porozumieniem Stron\.",
        "B2B nr {{ base.contract_number }} zawartą w dniu {{ base.signing_date }} r. "
        "za porozumieniem Stron.",
    )
    last = find(doc, r"^Strony zgodnie oświadczają, że skutek")
    set_text(
        last,
        "Strony zgodnie oświadczają, że skutek w postaci rozwiązania umowy nastąpi "
        "w dniu {{ doc.termination_date }} r.\nOstatni dzień świadczenia usług dla "
        "B2BNET to {{ doc.last_service_date }} r.",
    )
    r = "doc.release_non_compete"
    renumber(
        doc,
        {
            n: "{{ " + str(n + 1) + " if " + r + " else " + str(n) + " }}"
            for n in (3, 4, 5, 6)
        },
    )
    # Zwolnienie z zakazu konkurencji (dawniej osobny wzór) — § 10 ust. 1 umowy 2026.
    heading = clone_after(last, "§ 3", like=paragraph_heading(doc, 2))
    body = clone_after(
        heading, RELEASE_PL, like=find(doc, r"^Na mocy niniejszego porozumienia")
    )
    wrap_if(heading, body, r)


def build_termination_agreement_en(doc: _Doc) -> None:
    sub(doc, r"^(\s*)Warsaw,\s*\d.*$", lambda m: m.group(1) + EN_DATE_HEADER)
    sub(
        doc,
        r"^No\.\s*_+\s*signed on\s*_+\s*r\.",
        "No. {{ base.contract_number }} signed on {{ base.signing_date }}",
    )
    sub(doc, RX_COMPANY_EN, COMPANY_EN)
    sub(doc, r"^between\n", "and\n")
    sub(doc, RX_JDG_EN, JDG_EN)
    sub(
        doc,
        r"^The Parties unanimously agree to terminate.*$",
        "The Parties unanimously agree to terminate the B2B Cooperation Agreement "
        "No. {{ base.contract_number }} dated {{ base.signing_date }} by mutual "
        "agreement of the Parties as of {{ doc.termination_date }}. The last day of "
        "service provision for B2BNET is {{ doc.last_service_date }}.",
    )
    r = "doc.release_non_compete"
    sub(doc, r"^Podpis(\s+)Podpis", lambda m: "Signature" + m.group(1) + "Signature")
    h2 = paragraph_heading(doc, 2)
    equipment = find(doc, r"^The Partner declares that upon completion")
    # Zwrot sprzętu jak w PL (termin, koszty obciążające Partnera).
    set_text(
        equipment,
        "The Partner declares that upon completion of the provision of services to "
        "B2BNET, it will return all equipment entrusted to it to the B2BNET customer "
        "or to B2BNET in an undamaged condition, taking into account normal wear and "
        "tear, without delay, but no later than on the last day of service provision. "
        "If this deadline is not met or the equipment is returned in a deteriorated "
        "condition, B2BNET shall be entitled to charge the Partner with all costs "
        "charged to B2BNET for this reason by the B2BNET customer.",
    )
    h3 = paragraph_heading(doc, 3)
    sub_in(h2, r"§\s*2", "§ {{ 3 if " + r + " else 2 }}")
    sub_in(h3, r"§\s*3", "§ {{ 5 if " + r + " else 4 }}")
    # Poufność / reputacja — brak we wzorze EN, jest w PL.
    rep_heading = clone_after(equipment, "§ {{ 4 if " + r + " else 3 }}", like=h2)
    clone_after(
        rep_heading,
        "The Partner undertakes to keep confidential any information that could in "
        "any way harm the good name, reputation or interests of B2BNET, its "
        "employees, associates, contractors and affiliated entities, and to refrain "
        "from taking any actions, including public statements or publications, that "
        "could negatively affect the image or reputation of the Company.",
        like=equipment,
    )
    first_body = find(doc, r"^The Parties unanimously agree to terminate")
    rel_heading = clone_after(first_body, "§ 2", like=h2)
    rel_body = clone_after(rel_heading, RELEASE_EN, like=first_body)
    wrap_if(rel_heading, rel_body, r)


def build_termination_agreement_mandate_pl(doc: _Doc) -> None:
    sub(
        doc,
        r"^Zawarte w dniu .*? r\s*\.\s*w Warszawie",
        "Zawarte w dniu {{ document_date }} r. w Warszawie",
    )
    sub(doc, RX_COMPANY_PL, COMPANY_PL_SHORT)
    sub(
        doc,
        r"^Panią/em.*?Pesel:\s*_+",
        [
            ("{{ g.g_pan }} " + NAME_INSTR, True),
            (
                ", {{ g.g_zamieszkaly }} pod adresem: "
                + dot("partner.home_address")
                + ", {{ g.g_legitymujacy }} numerem PESEL: "
                + dot("partner.pesel")
                + ",",
                False,
            ),
        ],
    )
    sub(doc, r"^Zwan\S*(?= dalej Zleceniobiorcą)", "{{ g.g_zwanym }}")
    sub(
        doc,
        r"umowę zlecenie zawartą w dniu\s*_+\s*r\.",
        "umowę zlecenie zawartą w dniu {{ doc.base_signing_date }} r.",
    )
    sub(
        doc,
        r"nastąpi w dniu\s*_+\s*r\.",
        "nastąpi w dniu {{ doc.termination_date }} r.",
    )
    r = "doc.release_non_compete"
    h3 = paragraph_heading(doc, 3)
    body = find(doc, r"^Zleceniodawca zwalnia Zleceniobiorcę")
    # Paragraf zakazu w umowie zlecenie zależy od jej wersji (wzór: § 8 ust. 3) —
    # zwolnienie wskazuje zakaz po treści, nie po numerze; klient z formularza.
    set_text(
        body,
        "Zleceniodawca zwalnia Zleceniobiorcę z zakazu konkurencji określonego "
        "w umowie zlecenie i zgadza się na bezpośrednie zatrudnienie Zleceniobiorcy "
        "lub świadczenie przez Zleceniobiorcę usług bezpośrednio na rzecz: "
        + dot("doc.non_compete_client_name")
        + ", spółek macierzystych oraz innych podmiotów powiązanych.",
    )
    renumber(
        doc,
        {
            n: "{{ " + str(n) + " if " + r + " else " + str(n - 1) + " }}"
            for n in (4, 5, 6)
        },
    )
    wrap_if(h3, body, r)
    sub(doc, r"roszczeń wobec Spółki\s*$", "roszczeń wobec Zleceniodawcy.")
    sub(doc, r"^Umowę sporządzono w 2", "Porozumienie sporządzono w 2")


def build_termination_notice_pl(doc: _Doc) -> None:
    check_date_header_pl(doc)
    sub(
        doc,
        r"o numerze\s*_+\s*zawartą w dniu \d{2}\.\d{2}\.\d{4} r\.\s*",
        "o numerze {{ base.contract_number }} zawartą w dniu {{ base.signing_date }} r. ",
    )
    sub(doc, RX_COMPANY_PL, COMPANY_PL_SHORT)
    sub(doc, RX_JDG_PL, JDG_PL)
    zwanym(doc)
    # Umowa 2026: § 12 ust. 2 pkt 2, miesiąc ze skutkiem na koniec miesiąca
    # (wzór: § 2 pkt 2.2 i 30 dni).
    sub(
        doc,
        r"^B2B\.net S\.A\., działając zgodnie.*$",
        "B2B.net S.A., działając zgodnie z {{ refs.notice_paragraph }} umowy "
        "o współpracy B2B nr {{ base.contract_number }} z dnia {{ base.signing_date }} r., "
        "wypowiada niniejszą umowę z zachowaniem okresu wypowiedzenia "
        "{{ refs.notice_period_pl }}. Wypowiedzenie zostaje doręczone Partnerowi "
        "w dniu {{ doc.delivery_date }} r., a umowa ulega rozwiązaniu z dniem "
        "{{ doc.termination_date }} r.",
    )


def build_notice_withdrawal_pl(doc: _Doc) -> None:
    sub(
        doc,
        r"^Miejscowość, data: Warszawa,\s*_+\s*2026 r\.",
        "Miejscowość, data: " + PL_DATE_HEADER,
    )
    set_text(find(doc, r"^B2B\.net S\.A\.\s*$"), "{{ company.name }}")
    set_text(find(doc, r"^Aleje Jerozolimskie 180\s*$"), "{{ company.address }}")
    remove(find(doc, r"^02-486 Warszawa\s*$"))
    base = "nr {{ base.contract_number }} z dnia {{ base.signing_date }} r."
    sub(
        doc,
        r"nr \d+/\d{4} z dnia \d{2}\.\d{2}\.\d{4} r\.?(?=\s*$|\s*pozostaje)",
        base,
        count=2,
    )
    sub(
        doc,
        r"nr \d+/\d{4} z dnia \d{2}\.\d{2}\.\d{4} r\.?, które przekazałem Państwu "
        r"w dniu \d{2}\.\d{2}\.\d{4} r\.",
        base + ", które {{ g.g_przekazal }} Państwu w dniu "
        "{{ doc.notice_delivery_date }} r.",
    )
    sub(
        doc,
        r"^Imię i nazwisko, nazwa firmy\s*$",
        NAME_NOM + "{% if partner.legal_name %}, {{ partner.legal_name }}{% endif %}",
    )
    sub(doc, r"^NIP:\s*_+\s*$", "NIP: " + dot("partner.nip"))
    sub(
        doc,
        r"^Artur Twardowski – Prezes Zarządu\s*$",
        "{{ company.representative_nom }}",
    )


def build_preliminary_cez_pl(doc: _Doc) -> None:
    sub(doc, r"^Zawarta w dniu\s*\.+\s*r\.", "Zawarta w dniu {{ document_date }} r.")
    sub(
        doc,
        r"B2B\.net S\.A\.\s*z siedzibą w Warszawie,.*?reprezentowaną przez Pana "
        r"Artura Twardowskiego – Prezesa Zarządu,",
        COMPANY_PL_FULL,
    )
    sub(
        doc,
        r"^Panem/Panią.*$",
        [
            ("{{ g.g_pan }} " + NAME_INSTR, True),
            (
                ", {{ g.g_zamieszkaly }} pod adresem: "
                + dot("partner.home_address")
                + ", {{ g.g_legitymujacy }}",
                False,
            ),
        ],
    )
    sub(
        doc,
        r"^dokumentem tożsamości o numerze:.*$",
        "dokumentem tożsamości o numerze: "
        + dot("partner.id_document")
        + ", wydanym przez "
        + dot("partner.id_document_issuer")
        + ",",
    )
    sub(doc, r"^PESEL:\s*_+\s*$", "PESEL: " + dot("partner.pesel") + ",")
    sub(doc, r"^zwan\S*(?= dalej „Partnerem”)", "{{ g.g_zwanym }}")
    sub(doc, r"projektu nr\s*_+\s*,", "projektu nr " + dot("doc.project_number") + ",")
    sub(
        doc,
        r"stawki w wysokości\s*_+\s*zł netto za jedną godzinę świadczenia usług "
        r"\(słownie:\s*_+\s*złotych 00/100\)",
        "stawki w wysokości {{ doc.hourly_rate }} zł netto za jedną godzinę "
        "świadczenia usług (słownie: {{ doc.hourly_rate_words }})",
    )
    sub(doc, r"zasady1 fakturowania", "zasady fakturowania")
    sub(doc, r"do dnia\s*_+\s*\.", "do dnia {{ doc.valid_until }} r.")


@dataclass(frozen=True)
class Spec:
    source: str
    key: str
    build: Callable[[_Doc], None]


SPECS: tuple[Spec, ...] = (
    Spec(
        "aneks zmiana stawki B2B (draft 2026).docx",
        "annex_rate_change_pl",
        build_annex_rate_change_pl,
    ),
    Spec(
        "aneks zmiana stawki B2B_(draft 2026)_EN.docx",
        "annex_rate_change_en",
        build_annex_rate_change_en,
    ),
    Spec(
        "aneks zmiana daty startu B2B (draft 2026).docx",
        "annex_start_date_pl",
        build_annex_start_date_pl,
    ),
    Spec(
        "aneks uzupełnienie danych firmy B2B (draft 2026).docx",
        "annex_party_data_pl",
        build_annex_party_data_pl,
    ),
    Spec(
        "aneks uzupełnienie danych firmy B2B (draft 2026)_EN.docx",
        "annex_party_data_en",
        build_annex_party_data_en,
    ),
    Spec(
        "aneks o oddelegowaniu pracownika.docx",
        "annex_subcontractor_pl",
        build_annex_subcontractor_pl,
    ),
    Spec(
        "Aneks do umowy zlecenie (draft 2026).docx",
        "annex_mandate_pl",
        build_annex_mandate_pl,
    ),
    Spec(
        "Porozumienie o rozwiązaniu umowy (draft 2026) - zachowanie zakazu.docx",
        "termination_agreement_pl",
        build_termination_agreement_pl,
    ),
    Spec(
        "Porozumienie o rozwiązaniu umowy (draft 2026)_ENG.docx",
        "termination_agreement_en",
        build_termination_agreement_en,
    ),
    Spec(
        "rozwiązanie umowy za porozumieniem stron_zlecenie_zdjęcie zakazu.docx",
        "termination_agreement_mandate_pl",
        build_termination_agreement_mandate_pl,
    ),
    Spec(
        "Wypowiedzenie umowy.docx", "termination_notice_pl", build_termination_notice_pl
    ),
    Spec(
        "Oswiadczenie_o_cofnięcie_wypowiedzenia.docx",
        "notice_withdrawal_pl",
        build_notice_withdrawal_pl,
    ),
    Spec(
        "Umowa_przedwstepna_CeZ_B2B.docx",
        "preliminary_cez_pl",
        build_preliminary_cez_pl,
    ),
)


# ── Kontrola wyniku ──────────────────────────────────────────────────────────

_JINJA = re.compile(r"\{\{.*?\}\}|\{%.*?%\}", re.S)
_ONLY_LINES = re.compile(r"^(?:[\s_.…()]|[Pp]odpis|[Ss]ignature)*$")
_LEFTOVERS = [
    (re.compile(r"_{3,}"), "niewypełnione pole ____"),
    (re.compile(r"\.{4,}|…{2,}"), "niewypełnione pole ……"),
    (re.compile(r"XX/XX|\[nr umowy\]|\[data"), "znacznik ze wzoru"),
    (re.compile(r"\brr\.|\.\.\s*\."), "zdublowana końcówka"),
    (re.compile(r"\d{1,2}\.\d{2}\.\d{4}"), "data ze wzoru"),
    # 910/2014 = rozporządzenie eIDAS, nie numer umowy.
    (re.compile(r"\b(?!910/2014)\d{2,4}/20\d\d\b"), "numer umowy ze wzoru"),
    (re.compile(r"\w/(ą|ią|em|cą|a|i|Panią|Ms|Mrs)\b"), "forma rodzajowa ze wzoru"),
]


def verify(doc: _Doc, key: str) -> None:
    problems = []
    for p in paragraphs(doc):
        text = _JINJA.sub("", p.text)
        if _ONLY_LINES.match(text):
            continue
        for rx, label in _LEFTOVERS:
            if rx.search(text):
                problems.append(f"{label}: {p.text[:90]!r}")
    if problems:
        raise BuildError(
            f"[{key}] wynik ma pozostałości wzoru:\n  " + "\n  ".join(problems)
        )


# ── HTML (lustro podglądu) ───────────────────────────────────────────────────


def _esc(s: str) -> str:
    return _html.escape(s, quote=False)


def _run_html(r: Run) -> str:
    text = re.sub(r"[\t\n]+", " ", r.text)
    if not text:
        return ""
    out = _esc(text)
    if r.underline:
        out = f"<u>{out}</u>"
    if r.bold:
        out = f"<strong>{out}</strong>"
    return out


def _is_numbered(p: Paragraph) -> bool:
    ppr = p._p.pPr
    return ppr is not None and ppr.numPr is not None


def _iter_blocks(doc: _Doc):
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def doc_to_html(doc: _Doc) -> str:
    parts: list[str] = []
    in_list = False
    heading_done = False
    for block in _iter_blocks(doc):
        if isinstance(block, Table):
            rows = "".join(
                "<tr>"
                + "".join(f"<td>{_esc(c.text.strip())}</td>" for c in row.cells)
                + "</tr>"
                for row in block.rows
            )
            parts.append(f"<table>{rows}</table>")
            continue
        text = block.text.strip()
        if not text:
            continue
        if text.startswith("{%p"):
            parts.append("{%" + text[3:])
            continue
        numbered = _is_numbered(block)
        if numbered and not in_list:
            parts.append("<ol>")
            in_list = True
        elif not numbered and in_list:
            parts.append("</ol>")
            in_list = False
        inner = "".join(_run_html(r) for r in block.runs).strip()
        if numbered:
            parts.append(f"<li>{inner}</li>")
        elif re.match(r"^§\s*\S+\s*\.?$", _JINJA.sub("1", text)):
            parts.append(f"<h2>{_esc(text)}</h2>")
        elif not heading_done and block.alignment == 1:
            parts.append(f"<h1>{_esc(text)}</h1>")
            heading_done = True
        else:
            cls = ' class="right"' if block.alignment == 2 else ""
            parts.append(f"<p{cls}>{inner}</p>")
    if in_list:
        parts.append("</ol>")
    return "\n".join(parts) + "\n"


# ── Pakiet: podświetlenia i metadane ─────────────────────────────────────────

_CORE_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/'
    'metadata/core-properties" xmlns:dc="http://purl.org/dc/elements/1.1/" '
    'xmlns:dcterms="http://purl.org/dc/terms/" xmlns:dcmitype="http://purl.org/dc/'
    'dcmitype/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
    "<dc:title>{title}</dc:title><dc:creator>B2B.net S.A.</dc:creator>"
    "<cp:lastModifiedBy>NEXUS</cp:lastModifiedBy><cp:revision>1</cp:revision>"
    '<dcterms:created xsi:type="dcterms:W3CDTF">2026-09-23T00:00:00Z</dcterms:created>'
    '<dcterms:modified xsi:type="dcterms:W3CDTF">2026-09-23T00:00:00Z</dcterms:modified>'
    "</cp:coreProperties>"
)


def _dropped(name: str) -> bool:
    return (
        name.startswith("customXml/")
        or name.startswith("word/webextensions/")
        or name == "docProps/custom.xml"
    )


_REL = re.compile(r"<Relationship\b[^>]*/>")
_OVERRIDE = re.compile(r"<Override\b[^>]*/>")


def scrub_package(data: bytes, title: str) -> bytes:
    """Bez autorów wzoru, metadanych SharePointu, dodatków Worda i podświetleń."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            name = info.filename
            if _dropped(name):
                continue
            body = src.read(name)
            if name == "docProps/core.xml":
                body = _CORE_XML.format(title=_esc(title)).encode()
            elif name.endswith(".rels"):
                text = body.decode("utf-8")
                text = _REL.sub(
                    lambda m: (
                        ""
                        if re.search(
                            r'Target="[^"]*(customXml/|webextensions/|custom\.xml)',
                            m.group(0),
                        )
                        else m.group(0)
                    ),
                    text,
                )
                body = text.encode()
            elif name == "[Content_Types].xml":
                text = body.decode("utf-8")
                text = _OVERRIDE.sub(
                    lambda m: (
                        ""
                        if re.search(
                            r'PartName="/(customXml/|word/webextensions/|docProps/custom\.xml)',
                            m.group(0),
                        )
                        else m.group(0)
                    ),
                    text,
                )
                body = text.encode()
            elif name.startswith("word/") and name.endswith(".xml"):
                body = re.sub(rb"<w:highlight\b[^>]*/>", b"", body)
            zi = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(zi, body)
    return out.getvalue()


# ── Główna pętla ─────────────────────────────────────────────────────────────


def _nfc(s: str) -> str:
    return unicodedata.normalize("NFC", s)


def _source(src_dir: Path, name: str) -> Path:
    wanted = _nfc(name)
    for path in src_dir.iterdir():
        if _nfc(path.name) == wanted:
            return path
    raise BuildError(f"Brak wzoru w {src_dir}: {name}")


def build_one(spec: Spec, src_dir: Path, out_dir: Path) -> None:
    doc = docx.Document(str(_source(src_dir, spec.source)))
    spec.build(doc)
    drop_empty_numbered(doc)
    verify(doc, spec.key)
    buf = io.BytesIO()
    doc.save(buf)
    title = next(
        (
            p.text.strip()
            for p in paragraphs(doc)
            if p.alignment == 1 and p.text.strip()
        ),
        spec.key,
    )
    (out_dir / f"{spec.key}.docx").write_bytes(scrub_package(buf.getvalue(), title))
    (out_dir / f"{spec.key}.html").write_text(doc_to_html(doc), encoding="utf-8")
    print(f"[{spec.key}] ✓")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument(
        "source_dir", type=Path, help="katalog ze wzorami działu (.docx)"
    )
    parser.add_argument("out_dir", type=Path, nargs="?", default=OUT_DIR)
    args = parser.parse_args()
    args.out_dir.mkdir(parents=True, exist_ok=True)
    for spec in SPECS:
        build_one(spec, args.source_dir, args.out_dir)
    print(f"Gotowe — {len(SPECS) * 2} plików w {args.out_dir}")


if __name__ == "__main__":
    main()
