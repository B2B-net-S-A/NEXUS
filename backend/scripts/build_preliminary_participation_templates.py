"""Builder: szablony umów przedwstępnych (decyzje Artura, 21.10.2026).

Zapisuje do ``app/templates/documents/``:

* ``preliminary_participation_jdg_pl.{docx,html}`` i
  ``preliminary_participation_person_pl.{docx,html}`` — „Umowa o udziale
  w procesie rekrutacyjnym” (każdy klient poza Centrum e-Zdrowia) w dwóch
  wariantach: kandydat z jednoosobową działalnością i osoba fizyczna. Wzoru
  Word nie ma — treść jest tu, w ``_participation_blocks``, a styl (nagłówek
  z logo, stopka, marginesy, czcionka) bierze się z ``preliminary_cez_pl.docx``;
* ``preliminary_cez_v2_pl.{docx,html}`` — wzór CeZ z nową klauzulą formy
  zawarcia (każda forma podpisu: pisemna, kwalifikowany podpis, podpis
  zaufany, podpis osobisty). Pliki bez rewizji zostają dla wydanych umów.

Tekst prawny jest przepisany ze wzoru działu słowo w słowo; zmienione są
wyłącznie pola do wypełnienia (→ ``{{ }}``) i dwie klauzule formy podpisu
(§ 7 ust. 1 i 4) z decyzji z 21.10.2026.

Każdy akapit to jeden run na segment formatowania, więc znaczniki Jinja nie
są cięte między runy (docxtpl ich wtedy nie widzi). Wynik jest commitowany;
skrypt NIE jest częścią ścieżki produkcyjnej.

    python3 scripts/build_preliminary_participation_templates.py [katalog_wyjściowy]
"""

from __future__ import annotations

import argparse
import copy
import html as _html
import re
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

TEMPLATE_DIR = Path(__file__).resolve().parents[1] / "app" / "templates" / "documents"
CEZ_BASE = "preliminary_cez_pl"

FONT = "Aptos"
SIZE = Pt(10)

#: Klauzula formy zawarcia we wzorze CeZ (v1) i jej brzmienie od 21.10.2026.
CEZ_FORM_OLD = (
    "Niniejsza Umowa może zostać zawarta w formie pisemnej albo w formie "
    "elektronicznej poprzez opatrzenie oświadczeń Stron kwalifikowanymi "
    "podpisami elektronicznymi. Oświadczenie złożone w formie elektronicznej "
    "jest równoważne oświadczeniu złożonemu w formie pisemnej zgodnie z art. "
    "78¹ § 2 Kodeksu cywilnego."
)
CEZ_FORM_NEW = (
    "Niniejsza Umowa może zostać zawarta w formie pisemnej, w formie "
    "elektronicznej poprzez opatrzenie oświadczeń Stron kwalifikowanymi "
    "podpisami elektronicznymi albo w formie dokumentowej, w szczególności "
    "poprzez opatrzenie oświadczeń Stron podpisem zaufanym lub podpisem "
    "osobistym. Oświadczenie złożone w formie elektronicznej z kwalifikowanym "
    "podpisem elektronicznym jest równoważne oświadczeniu złożonemu w formie "
    "pisemnej zgodnie z art. 78¹ § 2 Kodeksu cywilnego."
)


class BuildError(RuntimeError):
    pass


# ── Treść: bloki ─────────────────────────────────────────────────────────────
#
# ("title", tekst) · ("center", tekst) · ("p", [(tekst, pogrubienie), …]) ·
# ("heading", tekst) · ("item", nr, tekst, [podpunkty], dalszy_ciąg | None) ·
# ("signatures", lewa, prawa)

PARTNER_NAME = "{{ g.g_pan }} {{ partner.name_instrumental or partner.name or '…' }}"


def _company(variant: str) -> list[tuple[str, bool]]:
    successor = (
        " oraz każde inne przedsiębiorstwo, które będzie następcą prawnym "
        "{{ company.name }} w tej części działalności, której dotyczy niniejsza "
        "umowa"
        if variant == "person"
        else ""
    )
    return [
        ("{{ company.name }}", True),
        (
            ", {{ company.address }}, wpisaną do Rejestru Przedsiębiorców "
            "prowadzonego przez {{ company.court }}, pod nr {{ company.krs }}, "
            f"NIP: {{{{ company.nip }}}}{successor}, reprezentowaną przez "
            "{{ company.representative }}, zwaną dalej „Zlecającym”",
            False,
        ),
    ]


def _candidate(variant: str) -> list[tuple[str, bool]]:
    if variant == "jdg":
        rest = (
            ", {{ g.g_prowadzacy }} jednoosobową działalność gospodarczą pod "
            "firmą {{ partner.legal_name or '…' }}, adres: "
            "{{ partner.business_address or '…' }}, NIP: {{ partner.nip or '…' }}, "
            "REGON: {{ partner.regon or '…' }}, {{ g.g_wpisany }} do Centralnej "
            "Ewidencji i Informacji o Działalności Gospodarczej (CEIDG), "
            "{{ g.g_zwanym }} dalej „Kandydatem”"
        )
    else:
        rest = (
            ", {{ g.g_zamieszkaly }} pod adresem: {{ partner.home_address or '…' }}, "
            "{{ g.g_legitymujacy }} dowodem osobistym o numerze: "
            "{{ partner.id_document or '…' }}, PESEL: {{ partner.pesel or '…' }}, "
            "{{ g.g_zwanym }} dalej „Kandydatem”"
        )
    return [(PARTNER_NAME, True), (rest, False)]


def _participation_blocks(variant: str) -> list[tuple]:
    person = variant == "person"
    p1_4 = (
        "Strony zgodnie oświadczają, że niniejsza Umowa nie stanowi umowy o pracę "
        "ani zobowiązania do jej zawarcia, a udział Kandydata w procesie "
        "rekrutacyjnym ma charakter dobrowolny."
        if person
        else "Strony zgodnie oświadczają, że niniejsza Umowa nie stanowi umowy "
        "o pracę ani zobowiązania do jej zawarcia, a udział w procesie "
        "rekrutacyjnym oraz ewentualna przyszła współpraca będą realizowane na "
        "zasadach właściwych dla relacji pomiędzy niezależnymi przedsiębiorcami, "
        "bez podporządkowania charakterystycznego dla stosunku pracy, o którym "
        "mowa w art. 22 §1 Kodeksu pracy."
    )
    rate_tail = (
        "powiększone o należny podatek VAT — pod warunkiem rozpoczęcia przez "
        "Kandydata działalności gospodarczej."
        if person
        else "powiększone o należny podatek VAT."
    )
    penalty_cases = (
        [
            "nie przystąpi w wyznaczonym terminie do testu online udostępnionego "
            "przez Zlecającego lub Klienta, lub",
            "nie przystąpi w umówionym terminie do rozmowy online przeprowadzanej "
            "przez Zlecającego lub Klienta, lub",
            "w inny sposób zaprzestanie udziału w procesie rekrutacyjnym "
            "prowadzonym przez Zlecającego lub Klienta przed wykonaniem "
            "zobowiązań, o których mowa w §2 ust. 1,",
        ]
        if person
        else [
            "nie przystąpi do testu online w wyznaczonym terminie, lub",
            "nie przystąpi do rozmowy online w umówionym terminie, lub",
            "w inny sposób zaprzestanie udziału w procesie rekrutacyjnym przed "
            "wykonaniem zobowiązań, o których mowa w §2 ust. 1,",
        ]
    )
    jurisdiction = (
        "W sprawach nieuregulowanych Umową zastosowanie mają przepisy prawa "
        "polskiego, w szczególności Kodeksu cywilnego. Właściwość sądu określają "
        "przepisy ogólne."
        if person
        else "W sprawach nieuregulowanych Umową zastosowanie mają przepisy prawa "
        "polskiego, w szczególności Kodeksu cywilnego. Ewentualne spory Strony "
        "poddają pod rozstrzygnięcie sądu właściwego miejscowo dla siedziby "
        "Zlecającego."
    )
    blocks: list[tuple] = [
        ("title", "Umowa o udziale w procesie rekrutacyjnym"),
        ("center", "(poprzedzająca ewentualną współpracę w modelu B2B)"),
        ("p", [("Zawarta w dniu {{ document_date }} r. w Warszawie pomiędzy:", False)]),
        ("p", _company(variant)),
        ("p", [("a", False)]),
        ("p", _candidate(variant)),
        (
            "p",
            [
                (
                    "Zlecający i Kandydat zwani są w dalszej części Umowy łącznie "
                    "„Stronami”, a każde z osobna „Stroną”.",
                    False,
                )
            ],
        ),
        ("heading", "§1. Przedmiot Umowy"),
        (
            "item",
            "1.",
            "Zlecający prowadzi proces rekrutacyjny mający na celu wyłonienie "
            "specjalisty do świadczenia usług informatycznych na rzecz klienta "
            "Zlecającego {{ doc.project_label or '…' }} (dalej „Klient”), "
            "z docelowym modelem współpracy w formule business-to-business (B2B).",
            [],
            None,
        ),
        (
            "item",
            "2.",
            "Kandydat jest zainteresowany udziałem w tym procesie rekrutacyjnym "
            "i wyraża wolę wzięcia w nim udziału na zasadach określonych "
            "w niniejszej Umowie.",
            [],
            None,
        ),
        (
            "item",
            "3.",
            "Celem Umowy jest zapewnienie, że Kandydat — po przystąpieniu do "
            "procesu rekrutacyjnego — weźmie w nim rzetelny udział aż do "
            "zakończenia etapów wskazanych w §2 ust. 1.",
            [],
            None,
        ),
        ("item", "4.", p1_4, [], None),
        ("heading", "§2. Zobowiązania Kandydata"),
        (
            "item",
            "1.",
            "Kandydat zobowiązuje się wziąć udział w procesie rekrutacyjnym "
            "prowadzonym przez Zlecającego na potrzeby projektu Klienta, "
            "obejmującym w szczególności następujące etapy:",
            [
                "rozwiązanie testu online (testu kompetencyjnego / technicznego) "
                "udostępnionego przez Zlecającego lub Klienta — w wyznaczonym "
                "terminie;",
                "przystąpienie do rozmowy rekrutacyjnej online przeprowadzanej "
                "przez Zlecającego lub Klienta — w umówionym terminie.",
            ],
            None,
        ),
        (
            "item",
            "2.",
            "Kandydat zobowiązuje się uczestniczyć w procesie rekrutacyjnym "
            "w sposób lojalny i zgodny z zasadami dobrej wiary, w tym potwierdzać "
            "wyznaczone terminy, stawiać się w umówionych terminach oraz "
            "informować Zlecającego — z odpowiednim wyprzedzeniem — "
            "o okolicznościach uniemożliwiających udział.",
            [],
            None,
        ),
        (
            "item",
            "3.",
            "Zobowiązania, o których mowa w ust. 1, uważa się za wykonane z chwilą "
            "rozwiązania testu online oraz odbycia rozmowy online.",
            [],
            None,
        ),
        (
            "item",
            "4.",
            "Obowiązek Kandydata dotyczy wyłącznie przystąpienia do etapów "
            "wskazanych w ust. 1 i nie obejmuje osiągnięcia określonego wyniku ani "
            "zawarcia jakiejkolwiek dalszej umowy. Umowa nie zobowiązuje Kandydata "
            "do przyjęcia oferty współpracy ani do nawiązania współpracy "
            "z Zlecającym lub Klientem.",
            [],
            None,
        ),
        ("heading", "§3. Zobowiązania Zlecającego"),
        (
            "item",
            "1.",
            "Zlecający zobowiązuje się umożliwić Kandydatowi udział w procesie "
            "rekrutacyjnym, w szczególności udostępnić Kandydatowi dostęp do testu "
            "online oraz link lub dostęp do rozmowy online, a także wskazać "
            "terminy poszczególnych etapów z wyprzedzeniem co najmniej 2 (dwóch) "
            "dni roboczych.",
            [],
            None,
        ),
        (
            "item",
            "2.",
            "Na żądanie Kandydata Zlecający poinformuje go o wyniku danego etapu "
            "procesu rekrutacyjnego.",
            [],
            None,
        ),
        (
            "item",
            "3.",
            "W przypadku pozytywnego zakończenia procesu rekrutacyjnego i wyboru "
            "Kandydata przez Klienta, Zlecający zobowiązuje się zaproponować "
            "Kandydatowi nawiązanie współpracy w modelu B2B na warunkach "
            "obejmujących wynagrodzenie w wysokości {{ doc.hourly_rate }} zł netto "
            f"za każdą godzinę świadczenia usług, {rate_tail}",
            [],
            None,
        ),
        (
            "item",
            "4.",
            "Zaproponowanie współpracy zgodnie z ust. 3 nie zobowiązuje żadnej ze "
            "Stron do zawarcia umowy o współpracy; jej ostateczne warunki zostaną "
            "określone w odrębnej umowie. Postanowienia niniejszego paragrafu nie "
            "tworzą stosunku pracy w rozumieniu art. 22 §1 Kodeksu pracy.",
            [],
            None,
        ),
        ("heading", "§4. Kara umowna"),
        (
            "item",
            "1.",
            "W przypadku gdy Kandydat, bez uzasadnionej przyczyny:",
            penalty_cases,
            "Zlecający będzie uprawniony do żądania od Kandydata zapłaty kary "
            "umownej w wysokości {{ doc.penalty_amount }} zł (słownie: "
            "{{ doc.penalty_amount_words }}).",
        ),
        (
            "item",
            "2.",
            "Kara umowna, o której mowa w ust. 1, przysługuje jednorazowo, "
            "niezależnie od liczby niewykonanych etapów procesu rekrutacyjnego.",
            [],
            None,
        ),
        (
            "item",
            "3.",
            "Kara umowna będzie płatna w terminie 14 dni od dnia doręczenia "
            "Kandydatowi wezwania do zapłaty.",
            [],
            None,
        ),
        (
            "item",
            "4.",
            "Kara umowna nie należy się w szczególności wówczas, gdy niewykonanie "
            "zobowiązań wynika z:",
            [
                "choroby, zdarzenia losowego lub siły wyższej,",
                "nieudostępnienia przez Zlecającego lub Klienta narzędzi, dostępu "
                "lub terminu niezbędnych do wykonania danego etapu,",
                "zmiany terminu dokonanej przez Zlecającego lub Klienta,",
                "innych uzasadnionych okoliczności niezależnych od Kandydata — pod "
                "warunkiem poinformowania Zlecającego przed wyznaczonym terminem, "
                "o ile było to możliwe.",
            ],
            None,
        ),
        (
            "item",
            "5.",
            "Zapłata kary umownej wyczerpuje roszczenia Zlecającego z tytułu "
            "nieuczestniczenia Kandydata w procesie rekrutacyjnym.",
            [],
            None,
        ),
    ]
    if person:
        blocks.append(
            (
                "item",
                "6.",
                "Postanowienia dotyczące kary umownej zostały indywidualnie "
                "uzgodnione przez Strony i odpowiadają wzajemnym zobowiązaniom "
                "Stron określonym w niniejszej Umowie.",
                [],
                None,
            )
        )
    blocks += [
        ("heading", "§5. Czas trwania Umowy"),
        (
            "item",
            "1.",
            "Umowa wchodzi w życie z dniem jej podpisania i obowiązuje do dnia "
            "wykonania zobowiązań, o których mowa w §2 ust. 1, nie dłużej jednak "
            "niż do dnia {{ doc.valid_until }} r.",
            [],
            None,
        ),
        (
            "item",
            "2.",
            "Po odbyciu testu online oraz rozmowy online zobowiązania Kandydata "
            "określone w §2 wygasają. Brak wyłonienia Kandydata w procesie "
            "rekrutacyjnym lub niepodjęcie dalszej współpracy nie rodzi po żadnej "
            "ze Stron jakichkolwiek roszczeń.",
            [],
            None,
        ),
        ("heading", "§6. Dane osobowe"),
        (
            "item",
            "1.",
            "Administratorem danych osobowych Kandydata jest B2B.net S.A. "
            "z siedzibą w Warszawie przy Al. Jerozolimskich 180, 02-486 Warszawa.",
            [],
            None,
        ),
        (
            "item",
            "2.",
            "Dane osobowe Kandydata będą przetwarzane w celu przeprowadzenia "
            "procesu rekrutacyjnego, realizacji niniejszej Umowy oraz podjęcia "
            "działań zmierzających do ewentualnego nawiązania przyszłej "
            "współpracy.",
            [],
            None,
        ),
        (
            "item",
            "3.",
            "Podstawą przetwarzania danych osobowych jest art. 6 ust. 1 lit. b "
            "RODO (czynności zmierzające do zawarcia i wykonanie niniejszej Umowy) "
            "oraz art. 6 ust. 1 lit. f RODO (prawnie uzasadniony interes "
            "administratora polegający na prowadzeniu rekrutacji na rzecz Klienta "
            "oraz ustaleniu, dochodzeniu lub obronie roszczeń).",
            [],
            None,
        ),
        (
            "item",
            "4.",
            "Dane kontaktowe Inspektora Ochrony Danych: iod@b2bnetwork.pl.",
            [],
            None,
        ),
        (
            "item",
            "5.",
            "Dane osobowe mogą być udostępniane: (a) podmiotom świadczącym na rzecz "
            "administratora usługi księgowe, prawne, IT oraz rekrutacyjne, a także "
            "(b) Klientowi Zlecającego — wyłącznie w zakresie niezbędnym do "
            "przeprowadzenia procesu rekrutacyjnego.",
            [],
            None,
        ),
        (
            "item",
            "6.",
            "Kandydatowi przysługują prawa przewidziane w RODO, w szczególności "
            "prawo dostępu do danych, ich sprostowania, usunięcia, ograniczenia "
            "przetwarzania, a także prawo wniesienia skargi do Prezesa Urzędu "
            "Ochrony Danych Osobowych.",
            [],
            None,
        ),
        (
            "item",
            "7.",
            "Dane osobowe będą przechowywane przez okres niezbędny do realizacji "
            "celów wskazanych w ust. 2 oraz przez okres przedawnienia ewentualnych "
            "roszczeń.",
            [],
            None,
        ),
        ("heading", "§7. Postanowienia końcowe"),
        (
            "item",
            "1.",
            # Decyzja Artura 21.10.2026: każda forma podpisu (wzór działu
            # dopuszczał tylko pisemną albo kwalifikowany podpis).
            "Umowa może zostać zawarta w formie pisemnej, w formie elektronicznej "
            "z wykorzystaniem kwalifikowanego podpisu elektronicznego, zgodnie "
            "z art. 78¹ §1 Kodeksu cywilnego, albo w formie dokumentowej, "
            "w szczególności z wykorzystaniem podpisu zaufanego lub podpisu "
            "osobistego. Strony uznają, że kwalifikowany podpis elektroniczny "
            "użyty przez Stronę jest równoważny z podpisem własnoręcznym dla celów "
            "prawnych i dowodowych.",
            [],
            None,
        ),
        (
            "item",
            "2.",
            "Wszelkie zmiany Umowy wymagają formy dokumentowej pod rygorem "
            "nieważności.",
            [],
            None,
        ),
        ("item", "3.", jurisdiction, [], None),
        (
            "item",
            "4.",
            "Umowa została sporządzona w dwóch jednobrzmiących egzemplarzach, po "
            "jednym dla każdej ze Stron. W przypadku zawarcia Umowy w formie "
            "elektronicznej albo dokumentowej Umowa zostaje sporządzona w jednym "
            "egzemplarzu w postaci elektronicznej, dostępnym dla obu Stron.",
            [],
            None,
        ),
        ("signatures", "Zlecający", "Kandydat"),
    ]
    return blocks


# ── DOCX ─────────────────────────────────────────────────────────────────────


def _run(paragraph, text: str, bold: bool = False):
    run = paragraph.add_run(text)
    run.bold = bold or None
    run.font.size = SIZE
    run.font.name = FONT
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:cs"), FONT)
    return run


def _paragraph(doc, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after: float = 4):
    paragraph = doc.add_paragraph()
    fmt = paragraph.paragraph_format
    fmt.alignment = align
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    fmt.line_spacing = Pt(15)
    return paragraph


def _hanging(paragraph, left: float, hang: float = 0.75) -> None:
    fmt = paragraph.paragraph_format
    fmt.left_indent = Cm(left)
    fmt.first_line_indent = Cm(-hang)
    fmt.tab_stops.add_tab_stop(Cm(left))


def _write_docx(blocks: list[tuple], out: Path) -> None:
    doc = docx.Document(str(TEMPLATE_DIR / f"{CEZ_BASE}.docx"))
    body = doc.element.body
    # Treść wzoru CeZ znika; zostają właściwości sekcji (nagłówek z logo,
    # stopka, marginesy) i style.
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    for block in blocks:
        kind = block[0]
        if kind == "title":
            p = _paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
            run = _run(p, block[1], bold=True)
            run.font.size = Pt(13)
        elif kind == "center":
            _run(_paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=10), block[1])
        elif kind == "p":
            p = _paragraph(doc, after=8)
            for text, bold in block[1]:
                _run(p, text, bold)
        elif kind == "heading":
            p = _paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.keep_with_next = True
            _run(p, block[1], bold=True)
        elif kind == "item":
            _, number, text, subs, tail = block
            p = _paragraph(doc)
            _hanging(p, 0.75)
            _run(p, f"{number}\t{text}")
            for index, sub in enumerate(subs):
                sp = _paragraph(doc)
                _hanging(sp, 1.5)
                _run(sp, f"{chr(ord('a') + index)})\t{sub}")
            if tail:
                tp = _paragraph(doc)
                tp.paragraph_format.left_indent = Cm(0.75)
                _run(tp, tail)
        elif kind == "signatures":
            _, left, right = block
            for _ in range(3):
                _paragraph(doc, after=0)
            for text in (
                "\t________________________________\t________________________________",
                f"\t{left}\t{right}",
            ):
                p = _paragraph(doc, align=WD_ALIGN_PARAGRAPH.LEFT, after=0)
                stops = p.paragraph_format.tab_stops
                stops.add_tab_stop(Cm(4.5), WD_TAB_ALIGNMENT.CENTER)
                stops.add_tab_stop(Cm(13.5), WD_TAB_ALIGNMENT.CENTER)
                _run(p, text)
        else:  # pragma: no cover - błąd w specyfikacji
            raise BuildError(f"nieznany blok {kind!r}")
    _check_tags(doc)
    doc.save(str(out))


def _check_tags(doc) -> None:
    """Każdy znacznik Jinja w całości w jednym runie — docxtpl nie skleja runów."""
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if run.text.count("{{") != run.text.count("}}"):
                raise BuildError(f"znacznik Jinja pocięty: {paragraph.text[:60]!r}")


# ── HTML ─────────────────────────────────────────────────────────────────────

_TAG = re.compile(r"({{.*?}})")


def _esc(text: str) -> str:
    """Escapuje tekst, nie znaczniki Jinja (te escapuje render)."""
    return "".join(
        part if _TAG.fullmatch(part) else _html.escape(part, quote=False)
        for part in _TAG.split(text)
    )


def _segments_html(segments: list[tuple[str, bool]]) -> str:
    return "".join(
        f"<strong>{_esc(text)}</strong>" if bold else _esc(text)
        for text, bold in segments
    )


def _write_html(blocks: list[tuple], out: Path) -> None:
    lines: list[str] = []
    open_list = False
    for block in blocks:
        kind = block[0]
        if kind != "item" and open_list:
            lines.append("</ol>")
            open_list = False
        if kind == "title":
            lines.append(f"<p><strong>{_esc(block[1])}</strong></p>")
        elif kind == "center":
            lines.append(f"<p>{_esc(block[1])}</p>")
        elif kind == "p":
            lines.append(f"<p>{_segments_html(block[1])}</p>")
        elif kind == "heading":
            lines.append(f"<h2>{_esc(block[1])}</h2>")
        elif kind == "item":
            _, _number, text, subs, tail = block
            if not open_list:
                lines.append("<ol>")
                open_list = True
            inner = _esc(text)
            if subs:
                items = "".join(f"<li>{_esc(s)}</li>" for s in subs)
                inner += f'<ol type="a">{items}</ol>'
            if tail:
                inner += _esc(tail)
            lines.append(f"<li>{inner}</li>")
        elif kind == "signatures":
            _, left, right = block
            lines.append(
                "<p>________________________________     "
                "________________________________</p>"
            )
            lines.append(
                f"<p>{_esc(left)}                                    {_esc(right)}</p>"
            )
    if open_list:
        lines.append("</ol>")
    out.write_text("\n".join(lines) + "\n", encoding="utf-8")


# ── CeZ v2 ───────────────────────────────────────────────────────────────────


def _build_cez_v2(out_dir: Path) -> None:
    doc = docx.Document(str(TEMPLATE_DIR / f"{CEZ_BASE}.docx"))
    hits = [p for p in doc.paragraphs if p.text.strip() == CEZ_FORM_OLD]
    if len(hits) != 1:
        raise BuildError(f"klauzula formy we wzorze CeZ: {len(hits)} trafień, nie 1")
    paragraph = hits[0]
    first, *rest = paragraph.runs
    # Pierwszy run niesie formatowanie akapitu (Aptos 10) — reszta znika.
    first.text = CEZ_FORM_NEW
    for run in rest:
        run._element.getparent().remove(run._element)
    doc.save(str(out_dir / "preliminary_cez_v2_pl.docx"))

    source = (TEMPLATE_DIR / f"{CEZ_BASE}.html").read_text(encoding="utf-8")
    if source.count(CEZ_FORM_OLD) != 1:
        raise BuildError("klauzula formy w HTML CeZ: nie jedno trafienie")
    (out_dir / "preliminary_cez_v2_pl.html").write_text(
        source.replace(CEZ_FORM_OLD, CEZ_FORM_NEW), encoding="utf-8"
    )


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("out_dir", nargs="?", default=str(TEMPLATE_DIR))
    out_dir = Path(parser.parse_args().out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    for variant in ("jdg", "person"):
        blocks = _participation_blocks(variant)
        key = f"preliminary_participation_{variant}_pl"
        _write_docx(copy.deepcopy(blocks), out_dir / f"{key}.docx")
        _write_html(blocks, out_dir / f"{key}.html")
        print(f"zapisano {key}.docx i .html")
    _build_cez_v2(out_dir)
    print("zapisano preliminary_cez_v2_pl.docx i .html")


if __name__ == "__main__":
    main()
